"""The registry's records and the hospitals' 2026 operating-room record, written from the clean truth in world.py:
the registry's prospective export and legacy enrolment workbook, its revisions, the state death-index linkage, the
follow-up status, the component crosswalk, the registry clinic visits, the 2026 case list and implant log, and the
HIM merge list."""
import csv
import datetime as dt

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import common as K

R_SEED = 26517


# ---------------------------------------------------------------------------------------------- catalogue numbers
STEM_SIZES = {"EDV": range(1, 13), "ASK": range(1, 11), "HOL": range(0, 12), "KLV": range(1, 12)}
CUP_SIZES = {"KEL": range(44, 66, 2), "QRV": range(46, 66, 2), "PLV": range(44, 68, 2), "MVN": range(46, 64, 2),
             "ORS": range(44, 60, 2), "FNM": range(46, 62, 2)}


def cat_no(fam, size, old=False):
    if fam == "EDV":
        return "4410-01-%03d" % size
    if fam == "KEL":
        return "5120-02-%03d" % size
    if fam == "FNM":
        return "5130-02-%03d" % size
    if fam == "ASK":
        return ("71-302-%02d" % size) if old else ("TK-A302-%02d" % size)
    if fam == "QRV":
        return ("71-118-%02d" % size) if old else ("TK-Q118-%02d" % size)
    if fam == "ORS":
        return "71-140-%02d" % size
    if fam == "HOL":
        return "HB-1100-%02d" % size
    if fam == "PLV":
        return "PV-2200-%02d" % size
    if fam == "KLV":
        return "KV-410-%02d" % size
    if fam == "MVN":
        return "MV-520-%02d" % size
    raise KeyError(fam)


def families(c):
    for code, s, p in K.CONSTRUCTS + K.RETIRED:
        if code == c:
            return s, p
    raise KeyError(c)


def kessick_old(d):
    return d < K.THORVIK_RECODE


def mrn_system(rng):
    return "E%07d" % rng.integers(1000000, 9999999)


def mrn_kittelstad(rng):
    return "KO%06d" % rng.integers(100000, 999999)


KITTELSTAD_EPIC = dt.date(2024, 3, 4)


# ---------------------------------------------------------------------------------------------- the registry's view
def registry(T):
    """Attach the registry's identifiers, components and MRNs to the hips and patients; build its tables."""
    rng = np.random.default_rng(R_SEED)
    pats, hips = T["patients"], T["hips"]
    for p in pats:
        p["mrn"] = mrn_system(rng)
        p["mrn_ko"] = mrn_kittelstad(rng)
    for h in hips:
        s, c = families(h["construct"])
        h["stem_size"] = int(rng.choice(list(STEM_SIZES[s])))
        h["cup_size"] = int(rng.choice(list(CUP_SIZES[c])))
        h["head"] = int(rng.choice([28, 32, 36], p=[.14, .52, .34]) if h["cup_size"] >= 48 else 28)
        h["bearing"] = str(rng.choice(["CoP", "MoP", "CoC"], p=[.62, .31, .07]))
        h["approach"] = str(rng.choice(["Posterior", "Anterior", "Lateral"], p=[.58, .33, .09]))
        h["surgeon"] = "S%s-%02d" % (h["hospital"][0], int(rng.integers(1, 9)))
        h["asa"] = int(rng.choice([1, 2, 3, 4], p=[.07, .55, .35, .03]))
        h["diag"] = str(rng.choice(["Osteoarthritis", "Avascular necrosis", "Developmental dysplasia",
                                    "Inflammatory arthritis", "Femoral neck fracture"], p=[.86, .05, .03, .02, .04]))
    on = [h for h in hips if h["entry"] is not None]
    # registry identifiers in the order the registry met the hips
    order = sorted(on, key=lambda h: (h["entry"], h["hospital"], h["hip"]))
    pid = {}
    for i, h in enumerate(order):
        h["reg_hip"] = "H%06d" % (100231 + i * 3 + int(rng.integers(0, 3)))
        if h["patient"] not in pid:
            pid[h["patient"]] = "P%06d" % (500117 + len(pid) * 2 + int(rng.integers(0, 2)))
        h["reg_pat"] = pid[h["patient"]]
    for p in pats:
        p["reg_pat"] = pid.get(p["patient"])
    T["on"] = on
    T["reg_pid"] = pid
    return T


def mrn_at(T, h, when):
    p = T["patients"][h["patient"]]
    if h["hospital"] == "KOH" and when < KITTELSTAD_EPIC:
        return p["mrn_ko"]
    return p.get("mrn_at_op", {}).get(h["hip"], p["mrn"])


def _head(ws, header, widths, fill="1F4E79"):
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i) if i < 26 else "A" + chr(65 + i - 26)].width = w
    ws.freeze_panes = "A2"


def write_prospective(path, T):
    rows = [h for h in T["on"] if h["op"] >= K.REG_START]
    rows.sort(key=lambda h: h["reg_hip"])
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["registry_hip_id", "registry_patient_id", "hospital_code", "hospital_mrn", "operation_date", "side",
                    "stem_catalogue_no", "cup_catalogue_no", "head_diameter_mm", "bearing", "approach", "surgeon_id",
                    "age_at_operation", "sex", "asa_grade", "primary_diagnosis"])
        for h in rows:
            s, c = families(h["construct"])
            old = kessick_old(h["op"])
            p = T["patients"][h["patient"]]
            w.writerow([h["reg_hip"], h["reg_pat"], h["hospital"], mrn_at(T, h, h["op"]), h["op"].isoformat(), h["side"],
                        cat_no(s, h["stem_size"], old), cat_no(c, h["cup_size"], old), h["head"], h["bearing"],
                        h["approach"], h["surgeon"], int(h["age"]), p["sex"], h["asa"], h["diag"]])
    return len(rows)


def write_legacy(path, T):
    rows = [h for h in T["on"] if h["op"] < K.REG_START]
    rows.sort(key=lambda h: (h["entry"], h["reg_hip"]))
    rng = np.random.default_rng(R_SEED + 1)
    initials = ["MB", "KL", "TD", "RS", "JW", "AH", "CN", "PV"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Enrolled hips"
    _head(ws, ["Registry hip ID", "Registry patient ID", "Hospital", "MRN", "Side", "Date of operation",
               "Stem catalogue no.", "Cup catalogue no.", "Head (mm)", "Bearing", "Approach", "Surgeon",
               "Age at operation", "Sex", "Enrolled on", "Clinic", "Abstracted by", "Notes"],
          [13, 13, 28, 11, 7, 13, 16, 16, 8, 8, 10, 9, 9, 6, 12, 26, 9, 30])
    for h in rows:
        s, c = families(h["construct"])
        p = T["patients"][h["patient"]]
        note = None
        if rng.random() < 0.032:
            k = int(rng.integers(0, 10))
            note = ["Op note scanned from paper chart", "Components from implant sticker sheet",
                    "Implant stickers missing; catalogue numbers from the op note",
                    "Side confirmed on radiograph at visit", "Patient brought implant card",
                    "Surgeon from theatre register (initials only on op note)",
                    "Head size from radiograph template, not on op note", "Op date from discharge summary",
                    "Records retrieved from offsite storage",
                    "Bearing from implant card"][k]
        ws.append([h["reg_hip"], h["reg_pat"], K.HOSP_NAME[h["hospital"]], mrn_at(T, h, h["entry"]),
                   "Right" if h["side"] == "R" else "Left", h["op"].strftime("%m/%d/%Y"),
                   cat_no(s, h["stem_size"], True), cat_no(c, h["cup_size"], True), h["head"], h["bearing"],
                   h["approach"], h["surgeon"], int(h["age"]), p["sex"], h["entry"],
                   K.HOSP_NAME[h["hospital"]].replace("Hospital", "").replace("Medical Center", "").strip()
                   + " registry clinic", str(rng.choice(initials)), note])
        ws.cell(ws.max_row, 15).number_format = "yyyy-mm-dd"
    wb.save(path)
    return len(rows)


REASONS_EARLY = (["Dislocation", "Infection", "Periprosthetic fracture", "Aseptic loosening", "Other"],
                 [.34, .29, .19, .10, .08])
REASONS_LATE = (["Aseptic loosening", "Osteolysis or wear", "Infection", "Periprosthetic fracture", "Dislocation",
                 "Other"], [.47, .14, .14, .11, .09, .05])
COMPONENTS = {"Dislocation": (["Head and liner", "Cup", "Head and liner"], [.6, .3, .1]),
              "Infection": (["All components", "Head and liner"], [.7, .3]),
              "Periprosthetic fracture": (["Stem", "All components"], [.8, .2]),
              "Aseptic loosening": (["Cup", "Stem", "All components"], [.45, .35, .2]),
              "Osteolysis or wear": (["Cup", "Head and liner"], [.55, .45]),
              "Other": (["Head and liner", "Cup", "Stem"], [.5, .3, .2])}


def revisions(T):
    """The registry's revision records: first and later revisions of registry hips while the patient lived in
    Wisconsin, at a system hospital (registry form) or at another Wisconsin hospital (state discharge data)."""
    rng = np.random.default_rng(R_SEED + 2)
    out = []
    for h in sorted(T["on"], key=lambda x: x["reg_hip"]):
        p = T["patients"][h["patient"]]
        end = min(p["death"], p["moved"], K.EXTRACT)
        for n, when in ((1, h["rev"]), (2, h["rerev"])):
            if when is None or when > end or when < h["entry"]:
                continue
            t = (when - h["op"]).days / K.YEAR
            names, probs = REASONS_EARLY if t < 1 else REASONS_LATE
            reason = str(rng.choice(names, p=probs))
            comp = str(rng.choice(COMPONENTS[reason][0], p=COMPONENTS[reason][1]))
            where = h["rev_where"] if n == 1 else ("system" if rng.random() < .8 else "other")
            out.append(dict(hip=h, n=n, date=when, reason=reason, comp=comp,
                            hosp=(h["hospital"] if rng.random() < .82 else K.HOSP[int(rng.integers(0, 3))])
                            if where == "system" else "OTH",
                            source="Registry form" if where == "system" else "State discharge data"))
    out.sort(key=lambda r: (r["date"], r["hip"]["reg_hip"], r["n"]))
    for i, r in enumerate(out):
        r["id"] = "R%05d" % (20113 + i)
    T["revisions"] = out
    return out


def write_revisions(path, T):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["revision_id", "registry_hip_id", "revision_date", "revision_number", "hospital_code",
                    "components_changed", "reason", "source"])
        for r in T["revisions"]:
            w.writerow([r["id"], r["hip"]["reg_hip"], r["date"].isoformat(), r["n"], r["hosp"], r["comp"], r["reason"],
                        r["source"]])
    return len(T["revisions"])


def deaths(T):
    """The state death-index linkage: deaths of registry patients while resident in Wisconsin, to the extract
    date; some certificates in the linkage carry the month and year only."""
    rng = np.random.default_rng(R_SEED + 3)
    out = []
    for p in T["patients"]:
        if p["reg_pat"] is None:
            continue
        if p["death"] <= K.EXTRACT and p["death"] < p["moved"]:
            month_only = rng.random() < 0.075
            out.append(dict(p=p, date=p["death"], month_only=month_only))
    out.sort(key=lambda r: r["p"]["reg_pat"])
    T["deaths"] = out
    return out


def write_deaths(path, T):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["registry_patient_id", "date_of_death", "linkage_match", "linkage_run"])
        rng = np.random.default_rng(R_SEED + 4)
        for r in T["deaths"]:
            d = r["date"]
            run = dt.date(d.year + (1 if d.month > 9 else 0), (d.month + 2) % 12 + 1, 1)
            run = min(max(run, dt.date(2018, 4, 1)), dt.date(2027, 1, 22))
            w.writerow([r["p"]["reg_pat"], d.strftime("%Y-%m") if r["month_only"] else d.isoformat(),
                        "Deterministic" if rng.random() < .93 else "Probabilistic, reviewed", run.isoformat()])
    return len(T["deaths"])


def write_status(path, T):
    rng = np.random.default_rng(R_SEED + 5)
    rows = []
    for p in T["patients"]:
        if p["reg_pat"] is None:
            continue
        moved = p["moved"] if (p["moved"] <= K.EXTRACT and p["moved"] < p["death"]) else None
        if moved is not None and moved < min(T["hips"][i]["entry"] or K.EXTRACT for i in p["hips"]):
            moved = None
        rows.append((p["reg_pat"], moved, p))
    rows.sort()
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["registry_patient_id", "residence_state", "left_wisconsin_on", "address_source", "last_verified"])
        for pid, moved, p in rows:
            if moved is not None:
                w.writerow([pid, str(rng.choice(["MN", "IL", "MI", "IA", "AZ", "FL", "TX", "CO"],
                                               p=[.24, .22, .14, .10, .12, .10, .04, .04])),
                            moved.isoformat(), "EHR address history",
                            min(moved + dt.timedelta(days=int(rng.integers(5, 200))), dt.date(2027, 1, 5)).isoformat()])
            else:
                dead = p["death"] <= K.EXTRACT
                last = min(p["death"], K.EXTRACT) if dead else K.EXTRACT - dt.timedelta(days=int(rng.integers(0, 700)))
                w.writerow([pid, "WI", "", "EHR address history", last.isoformat()])
    return len(rows)


def write_crosswalk(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Components"
    _head(ws, ["Maker", "Family", "Component", "Fixation", "Size", "Catalogue no.", "Catalogue no. before 1 July 2020"],
          [22, 10, 18, 12, 7, 15, 20])
    for fam in ("EDV", "KEL", "FNM", "ASK", "QRV", "ORS", "HOL", "PLV", "KLV", "MVN"):
        f = K.FAM[fam]
        sizes = STEM_SIZES.get(fam) or CUP_SIZES.get(fam)
        for s in sizes:
            if fam in ("ASK", "QRV"):
                ws.append([K.MAKERS[f[2]], f[1], "Femoral stem" if f[3] == "stem" else "Acetabular shell", f[4].title(),
                           s, cat_no(fam, s), cat_no(fam, s, True)])
            else:
                ws.append([K.MAKERS[f[2]], f[1], "Femoral stem" if f[3] == "stem" else "Acetabular shell", f[4].title(),
                           s, cat_no(fam, s), None])
    ws2 = wb.create_sheet("Constructs")
    _head(ws2, ["Construct", "Femoral stem", "Acetabular shell", "Hip agreement 2022 to 2027", "Used at the hospitals since"],
          [20, 14, 16, 26, 20])
    import world as W
    first = lambda code: min(y for y in range(2005, 2027) if W.volume(code, y) > 0)
    for code, s, c in K.CONSTRUCTS:
        ws2.append([K.CNAME[code], K.FAM[s][1], K.FAM[c][1], "On the agreement", first(code)])
    ws2.append(["Askvane/Orsinne", "Askvane", "Orsinne", "Not on the agreement (Orsinne discontinued 2012)", first("MC")])
    ws2.append(["Edvalt/Fennimark", "Edvalt", "Fennimark", "Not on the agreement (Fennimark discontinued 2010)", first("TI")])
    wb.save(path)


def write_visits(path, T):
    """Registry clinic visits: the legacy enrolment visit, then the scheduled reviews the patient attended."""
    rng = np.random.default_rng(R_SEED + 6)
    rows = []
    for h in T["on"]:
        p = T["patients"][h["patient"]]
        end = min(p["death"], p["moved"], K.EXTRACT, h["rev"])
        plan = []
        if h["op"] < K.REG_START:
            plan.append((h["entry"], "Legacy enrolment"))
            for yrs in (10, 15, 20):
                d = h["op"] + dt.timedelta(days=int(yrs * K.YEAR) + int(rng.integers(-40, 60)))
                if d > h["entry"] + dt.timedelta(days=200) and rng.random() < .62:
                    plan.append((d, "%d years" % yrs))
        else:
            for yrs, pr in ((1, .74), (2, .46), (5, .41)):
                d = h["op"] + dt.timedelta(days=int(yrs * K.YEAR) + int(rng.integers(-30, 75)))
                if rng.random() < pr:
                    plan.append((d, "%d year%s" % (yrs, "" if yrs == 1 else "s")))
        if rng.random() < .09:
            d = h["entry"] + dt.timedelta(days=int(rng.integers(30, 3000)))
            plan.append((d, "Unscheduled"))
        for d, kind in plan:
            while d.weekday() >= 5:
                d += dt.timedelta(days=1)
            if d > end or (K.PAUSE[0] <= d <= K.PAUSE[1]):
                continue
            hhs = int(np.clip(rng.normal(91 if kind != "Unscheduled" else 78, 7), 31, 100))
            rows.append((d, h["reg_hip"], h["hospital"], kind, "Y" if rng.random() < .71 else "N", hhs))
    rows.sort()
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["visit_id", "registry_hip_id", "clinic", "visit_date", "visit_type", "radiograph", "harris_hip_score"])
        n = 300417
        for d, hip, clin, kind, xr, hhs in rows:
            n += int(rng.integers(1, 3))
            w.writerow(["V%07d" % n, hip, clin, d.isoformat(), kind, xr, hhs])
    return len(rows)
