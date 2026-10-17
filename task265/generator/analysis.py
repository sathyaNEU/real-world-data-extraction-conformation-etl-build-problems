"""Every figure the build owes, computed forward from the shipped pack (never from the generator's state): the
ten-year cumulative percent revision of each construct on each rung of the ladder and under each convention, and
the asks' figures on the golden path and on the cracker's and the mirror's paths."""
import csv
import datetime as dt
import os
from collections import Counter, defaultdict

import numpy as np
from openpyxl import load_workbook

import common as K


def _csv(path):
    with open(path, encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def _xlsx(path, sheet, header_row=1):
    ws = load_workbook(path, read_only=True, data_only=True)[sheet]
    it = ws.iter_rows(values_only=True)
    for _ in range(header_row - 1):
        next(it)
    head = next(it)
    return [dict(zip(head, r)) for r in it if any(v is not None for v in r)]


def crosswalk(P):
    fam, pair = {}, {}
    for r in _xlsx(P(K.FILES["crosswalk"]), "Components"):
        fam[r["Catalogue no."]] = r["Family"]
        if r["Catalogue no. before 1 July 2020"]:
            fam[r["Catalogue no. before 1 July 2020"]] = r["Family"]
    for r in _xlsx(P(K.FILES["crosswalk"]), "Constructs"):
        pair[(r["Femoral stem"], r["Acetabular shell"])] = r["Construct"]
    return fam, pair


def registry_hips(target, death_day=15):
    """One record per registry hip: construct, operation, registry entry, first revision, death, move."""
    P = lambda f: os.path.join(target, f)
    fam, pair = crosswalk(P)
    hips = {}
    for r in _csv(P(K.FILES["prosp"])):
        op = dt.date.fromisoformat(r["operation_date"])
        hips[r["registry_hip_id"]] = dict(pat=r["registry_patient_id"], op=op, entry=op, legacy=False,
                                          construct=pair.get((fam[r["stem_catalogue_no"]], fam[r["cup_catalogue_no"]])),
                                          hosp=r["hospital_code"])
    names = {v: k for k, v in K.HOSP_NAME.items()}
    for r in _xlsx(P(K.FILES["legacy"]), "Enrolled hips"):
        op = dt.datetime.strptime(r["Date of operation"], "%m/%d/%Y").date()
        en = r["Enrolled on"]
        en = en.date() if isinstance(en, dt.datetime) else en
        assert r["Registry hip ID"] not in hips
        hips[r["Registry hip ID"]] = dict(pat=r["Registry patient ID"], op=op, entry=en, legacy=True,
                                          construct=pair.get((fam[r["Stem catalogue no."]], fam[r["Cup catalogue no."]])),
                                          hosp=names[r["Hospital"]])
    rev = {}
    for r in _csv(P(K.FILES["revisions"])):
        d = dt.date.fromisoformat(r["revision_date"])
        if r["revision_number"] == "1":
            assert r["registry_hip_id"] not in rev
            rev[r["registry_hip_id"]] = d
    death = {}
    for r in _csv(P(K.FILES["deaths"])):
        v = r["date_of_death"]
        if len(v) == 7:
            y, m = map(int, v.split("-"))
            last = (dt.date(y + (m == 12), m % 12 + 1, 1) - dt.timedelta(days=1)).day
            death[r["registry_patient_id"]] = dt.date(y, m, {"first": 1, "last": last}.get(death_day, death_day)
                                                      if isinstance(death_day, str) else min(death_day, last))
        else:
            death[r["registry_patient_id"]] = dt.date.fromisoformat(v)
    moved = {}
    for r in _csv(P(K.FILES["status"])):
        if r["left_wisconsin_on"]:
            moved[r["registry_patient_id"]] = dt.date.fromisoformat(r["left_wisconsin_on"])
    for k, h in hips.items():
        end = min(death.get(h["pat"], K.EXTRACT), moved.get(h["pat"], K.EXTRACT), K.EXTRACT)
        r = rev.get(k)
        h["event"] = r is not None and r <= end
        h["stop"] = r if h["event"] else end
        h["rev"] = r
    return hips


def km(rows, horizon=K.TEN, inclusive=False):
    a = np.array(rows, float)
    e, s, d = a[:, 0], a[:, 1], a[:, 2].astype(bool)
    S = 1.0
    for t in np.unique(s[d & (s <= horizon)]):
        at = ((e <= t) if inclusive else (e < t)) & (s >= t)
        n = at.sum()
        if n:
            S *= 1 - (d & (s == t) & at).sum() / n
    return 100 * (1 - S)


def rows_for(hips, name, rung):
    out = []
    for h in hips.values():
        if h["construct"] != name:
            continue
        if rung == "prosp" and h["legacy"]:
            continue
        s = (h["stop"] - h["op"]).days
        if rung == "golden" or rung == "prosp":
            e = (h["entry"] - h["op"]).days
        elif rung == "op":
            e = 0
        elif rung == "reg":
            e = max(0, (K.REG_START - h["op"]).days)
        out.append((e, s, h["event"]))
    return out


def call(target, death_day=15, horizon=K.TEN, inclusive=False):
    hips = registry_hips(target, death_day)
    out = {}
    for code in K.CODES:
        nm = K.CNAME[code]
        f = {r: km(rows_for(hips, nm, r), horizon, inclusive) for r in ("golden", "op", "reg")}
        rows = rows_for(hips, nm, "golden")
        f["crude"] = 100 * sum(1 for x in rows if x[2]) / len(rows)
        pr = rows_for(hips, nm, "prosp")
        f["prosp_max_years"] = max(x[1] for x in pr) / K.YEAR
        f["n"] = len(rows)
        f["legacy"] = sum(1 for h in hips.values() if h["construct"] == nm and h["legacy"])
        f["at_risk_10"] = sum(1 for x in rows if x[0] < horizon <= x[1])
        out[code] = f
    return out, hips


def keeps(fig, rung):
    return [c for c in K.CODES if fig[c][rung] <= K.BAR]


# ---------------------------------------------------------------------------------------------- the asks
def or_hips(target, honour_side=True):
    """2026 primary total hips from the operating-room record: (case, side) with an implanted stem and an
    implanted shell that were not explanted, mapped to a construct; with honour_side the sides the case list
    says were operated on only."""
    P = lambda f: os.path.join(target, f)
    fam, pair = crosswalk(P)
    cases = {r["Case number"]: r for r in _xlsx(P(K.FILES["cases"]), "Cases 2026")}
    names = {v: k for k, v in K.HOSP_NAME.items()}
    parts = defaultdict(Counter)
    for r in _csv(P(K.FILES["implants"])):
        if r["procedure"] != "ARTHROPLASTY, HIP, TOTAL":
            continue
        f = fam.get(r["catalogue_no"])
        if f is None:
            continue
        q = int(r["quantity"])
        if r["action"] == "Implanted":
            parts[(r["case_number"], r["implant_side"])][f] += q
        elif r["action"] == "Explanted":
            parts[(r["case_number"], r["implant_side"])][f] -= q
    out = []
    for (case, side), fams in parts.items():
        c = cases[case]
        if honour_side and c["Side"] != "Bilateral" and c["Side"] != side:
            continue
        stems = [f for f, n in fams.items() if n > 0 and f in ("Edvalt", "Askvane", "Holbrin", "Kalvoran")]
        cups = [f for f, n in fams.items() if n > 0 and f not in ("Edvalt", "Askvane", "Holbrin", "Kalvoran")]
        if len(stems) != 1 or len(cups) != 1:
            continue
        sd = c["Surgery date"]
        sd = sd.date() if isinstance(sd, dt.datetime) else sd
        out.append(dict(case=case, side=side[0], hosp=names[c["Hospital"]], date=sd, mrn=c["MRN"],
                        construct=pair.get((stems[0], cups[0]))))
    return out


def asks(target, honour_side=True, merges=True):
    P = lambda f: os.path.join(target, f)
    hs = or_hips(target, honour_side)
    A = {code: {h: 0 for h in K.HOSP} for code in K.CODES}
    by_name = {K.CNAME[c]: c for c in K.CODES}
    for h in hs:
        A[by_name[h["construct"]]][h["hosp"]] += 1
    reg = set()
    for r in _csv(P(K.FILES["prosp"])):
        if r["operation_date"].startswith("2026"):
            reg.add((r["hospital_mrn"], r["side"], r["operation_date"]))
    alias = {}
    if merges:
        for r in _csv(P(K.FILES["merges"])):
            alias.setdefault(r["surviving_mrn"], []).append(r["retired_mrn"])
    Ch = {h: 0 for h in K.HOSP}
    Cc = {c: 0 for c in K.CODES}
    for h in hs:
        keys = [h["mrn"]] + alias.get(h["mrn"], [])
        if not any((m, h["side"], h["date"].isoformat()) in reg for m in keys):
            Ch[h["hosp"]] += 1
            Cc[by_name[h["construct"]]] += 1
    return dict(A=A, C_hosp=Ch, C_construct=Cc)
