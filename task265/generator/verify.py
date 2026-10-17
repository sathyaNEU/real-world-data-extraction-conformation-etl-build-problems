"""The independent verifier: reads only the shipped bundle (python3 verify.py <target>), recomputes the call, every
rung and every graded ask figure on its own code path (the csv module and openpyxl, a product-limit estimator
written from counts of entries and exits, no generator state, no pandas, no analysis.py), and compares everything
with generator/figures.json. Exit 0 when every figure agrees."""
import bisect
import csv
import datetime as dt
import json
import os
import re
import sys
from collections import defaultdict

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
FAIL = []
TEN = 3652.5
REG = dt.date(2018, 1, 1)
END = dt.date(2026, 12, 31)


def same(label, a, b):
    if a != b:
        FAIL.append("%s: verifier %r, figures.json %r" % (label, a, b))


def rows_csv(path):
    with open(path, encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def rows_xlsx(path, sheet):
    ws = load_workbook(path, read_only=True, data_only=True)[sheet]
    it = ws.iter_rows(values_only=True)
    head = next(it)
    return [dict(zip(head, r)) for r in it if any(v is not None for v in r)]


def day(v):
    return v.date() if isinstance(v, dt.datetime) else v


def round1(x):
    return int(x * 10 + 0.5) / 10


def product_limit(spells):
    """spells: (entry, exit, revised) in days since the operation. At each revision time t the hips at risk are
    those entered before t (entry < t) that have not left before t (exit >= t)."""
    entries = sorted(e for e, x, r in spells)
    exits = sorted(x for e, x, r in spells)
    events = defaultdict(int)
    for e, x, r in spells:
        if r and x <= TEN:
            events[x] += 1
    s = 1.0
    for t in sorted(events):
        at_risk = bisect.bisect_left(entries, t) - bisect.bisect_left(exits, t)
        s *= 1 - events[t] / at_risk
    return 100 * (1 - s)


def main(target):
    P = lambda name: os.path.join(target, name)
    files = os.listdir(target)
    find = lambda pat: [f for f in files if re.search(pat, f)][0]
    F = json.load(open(os.path.join(HERE, "figures.json")))

    # ------------------------------------------------------------------ catalogue numbers to constructs
    fam, cons = {}, {}
    for r in rows_xlsx(P(find(r"crosswalk")), "Components"):
        fam[r["Catalogue no."]] = r["Family"]
        if r["Catalogue no. before 1 July 2020"]:
            fam[r["Catalogue no. before 1 July 2020"]] = r["Family"]
    for r in rows_xlsx(P(find(r"crosswalk")), "Constructs"):
        if r["Hip agreement 2022 to 2027"] == "On the agreement":
            cons[(r["Femoral stem"], r["Acetabular shell"])] = r["Construct"]

    # ------------------------------------------------------------------ the registry's hips
    hip = {}
    for r in rows_csv(P(find(r"prospective"))):
        op = dt.date.fromisoformat(r["operation_date"])
        hip[r["registry_hip_id"]] = (r["registry_patient_id"], op, op,
                                     cons.get((fam[r["stem_catalogue_no"]], fam[r["cup_catalogue_no"]])))
    for r in rows_xlsx(P(find(r"legacy")), "Enrolled hips"):
        m, d, y = map(int, r["Date of operation"].split("/"))
        hip[r["Registry hip ID"]] = (r["Registry patient ID"], dt.date(y, m, d), day(r["Enrolled on"]),
                                     cons.get((fam[r["Stem catalogue no."]], fam[r["Cup catalogue no."]])))
    first_rev = {}
    for r in rows_csv(P(find(r"revisions"))):
        d = dt.date.fromisoformat(r["revision_date"])
        k = r["registry_hip_id"]
        first_rev[k] = min(first_rev.get(k, d), d)
    gone = {}
    for r in rows_csv(P(find(r"death_index"))):
        v = r["date_of_death"]
        gone[r["registry_patient_id"]] = dt.date(int(v[:4]), int(v[5:7]), 15) if len(v) == 7 else dt.date.fromisoformat(v)
    for r in rows_csv(P(find(r"followup_status"))):
        if r["left_wisconsin_on"]:
            d = dt.date.fromisoformat(r["left_wisconsin_on"])
            gone[r["registry_patient_id"]] = min(gone.get(r["registry_patient_id"], d), d)
    spells = defaultdict(lambda: defaultdict(list))
    for k, (pat, op, entry, c) in hip.items():
        if c is None:
            continue
        stop = min(gone.get(pat, END), END)
        rv = first_rev.get(k)
        revised = rv is not None and rv <= stop
        x = ((rv if revised else stop) - op).days
        legacy = op < REG
        spells[c]["golden"].append(((entry - op).days, x, revised))
        spells[c]["op"].append((0, x, revised))
        spells[c]["reg"].append(((REG - op).days if legacy else 0, x, revised))
        spells[c]["crude"].append(revised)
    names = sorted(spells)
    rung = {r: {c: round1(product_limit(spells[c][r])) for c in names} for r in ("golden", "op", "reg")}
    rung["crude"] = {c: round1(100 * sum(spells[c]["crude"]) / len(spells[c]["crude"])) for c in names}
    for r in ("golden", "op", "reg", "crude"):
        same("rung %s" % r, rung[r], F["rungs"][r])
        same("rung %s set" % r, sorted(c for c in names if rung[r][c] <= 5.0), sorted(F["rung_sets"][r]))
    same("ten-year figures", rung["golden"], F["call"]["cpr10"])
    same("constructs that stay", sorted(c for c in names if rung["golden"][c] <= 5.0), sorted(F["call"]["stays"]))

    # ------------------------------------------------------------------ the 2026 operating-room record
    side_of = {}
    hosp_of, date_of, mrn_of = {}, {}, {}
    for r in rows_xlsx(P(find(r"case_list")), "Cases 2026"):
        side_of[r["Case number"]] = r["Side"]
        hosp_of[r["Case number"]] = r["Hospital"]
        date_of[r["Case number"]] = day(r["Surgery date"])
        mrn_of[r["Case number"]] = r["MRN"]
    net = defaultdict(lambda: defaultdict(int))
    for r in rows_csv(P(find(r"implant_log"))):
        if r["procedure"] != "ARTHROPLASTY, HIP, TOTAL" or r["catalogue_no"] not in fam:
            continue
        sign = {"Implanted": 1, "Explanted": -1}.get(r["action"], 0)
        net[(r["case_number"], r["implant_side"])][fam[r["catalogue_no"]]] += sign * int(r["quantity"])
    stems = {"Edvalt", "Askvane", "Holbrin", "Kalvoran"}
    hips26 = []
    for (case, side), parts in net.items():
        if side_of[case] not in ("Bilateral", side):
            continue
        s = [f for f, n in parts.items() if n > 0 and f in stems]
        c = [f for f, n in parts.items() if n > 0 and f not in stems]
        if len(s) == 1 and len(c) == 1:
            hips26.append((cons[(s[0], c[0])], hosp_of[case], side[0], date_of[case], mrn_of[case]))
    ask_a = {c: {h: 0 for h in F["ask_a_hospital_totals"]} for c in names}
    for c, h, *_ in hips26:
        ask_a[c][h] += 1
    same("Ask A", ask_a, F["ask_a"])
    on_reg = {(r["hospital_mrn"], r["side"], r["operation_date"]) for r in rows_csv(P(find(r"prospective")))}
    earlier = defaultdict(set)
    for r in rows_csv(P(find(r"mrn_merges"))):
        earlier[r["surviving_mrn"]].add(r["retired_mrn"])
    miss_h = {h: 0 for h in F["ask_c_hospital"]}
    miss_c = {c: 0 for c in names}
    for c, h, s, d, m in hips26:
        if not any((x, s, d.isoformat()) in on_reg for x in {m} | earlier[m]):
            miss_h[h] += 1
            miss_c[c] += 1
    same("Ask C by hospital", miss_h, F["ask_c_hospital"])
    same("Ask C by construct", miss_c, F["ask_c_construct"])
    if FAIL:
        print("verify: %d mismatches" % len(FAIL))
        for x in FAIL:
            print("  " + x)
        sys.exit(1)
    print("verify: call, rungs and every ask figure agree with figures.json (%d registry hips, %d hips in 2026)"
          % (len(hip), len(hips26)))


if __name__ == "__main__":
    main(sys.argv[1])
