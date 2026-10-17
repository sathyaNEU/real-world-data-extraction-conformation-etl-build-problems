"""The hospitals' 2026 operating-room record (the case list and the perioperative implant log), the HIM merge list
and the registry clinic visits, written from the clean truth. Every primary total hip done in 2026 is a case here
whether or not the registry captured it."""
import csv
import datetime as dt

import numpy as np
from openpyxl import Workbook

import common as K
import records as R

O_SEED = 26529
ROOMS = {"BVMC": ["MOR %02d" % i for i in range(1, 15)], "TLH": ["OR %d" % i for i in range(1, 7)],
         "KOH": ["OR %d" % i for i in range(1, 11)]}
PROC = {"THA": "ARTHROPLASTY, HIP, TOTAL", "RTHA": "ARTHROPLASTY, HIP, REVISION", "HEMI": "HEMIARTHROPLASTY, HIP",
        "TKA": "ARTHROPLASTY, KNEE, TOTAL", "RTKA": "ARTHROPLASTY, KNEE, REVISION",
        "ORIF": "OPEN REDUCTION INTERNAL FIXATION", "IMN": "INTRAMEDULLARY NAILING, FEMUR OR TIBIA",
        "FUS": "FUSION, SPINE, LUMBAR, POSTERIOR"}
MAKER_OF = {"EDV": "VNH", "KEL": "VNH", "FNM": "VNH", "ASK": "THV", "QRV": "THV", "ORS": "THV", "HOL": "JUV",
            "PLV": "JUV", "KLV": "BRY", "MVN": "BRY"}
HEAD_CAT = {"VNH": "4600-%02d-%s", "THV": "TK-H55-%02d%s", "JUV": "HD-3300-%02d%s", "BRY": "BYH-610-%02d%s"}
LINER_CAT = {"VNH": "5180-%02d-%03d", "THV": "TK-L18-%02d%02d", "JUV": "LN-2290-%02d%02d", "BRY": "BYL-525-%02d%02d"}


def _lot(rng, maker, year):
    if maker == "VNH":
        return "%d%s%05d" % (year % 100, "ABCDEFGHJK"[int(rng.integers(0, 10))], rng.integers(0, 99999))
    if maker == "THV":
        return "K%07d" % rng.integers(1000000, 9999999)
    if maker == "JUV":
        return "%dT%06d" % (year % 10, rng.integers(0, 999999))
    return "BY-%04d-%03d" % (rng.integers(1000, 9999), rng.integers(0, 999))


def hip_lines(rng, fam_s, fam_c, ss, cs, head, bearing, year, old=False):
    """The four components of a total hip, as (maker, catalogue no., description, qty)."""
    ms, mc = MAKER_OF[fam_s], MAKER_OF[fam_c]
    out = [(ms, R.cat_no(fam_s, ss, old), "%s FEMORAL STEM SZ %d %s" % (K.FAM[fam_s][1].upper(), ss,
                                                                      K.FAM[fam_s][4].upper()), 1),
           (mc, R.cat_no(fam_c, cs, old), "%s ACETABULAR SHELL %dMM" % (K.FAM[fam_c][1].upper(), cs), 1)]
    mat = "CERAMIC" if bearing in ("CoP", "CoC") else "COCR"
    out.append((mc, LINER_CAT[mc] % (head, cs) if mc != "VNH" else LINER_CAT[mc] % (head, cs),
                ("CERAMIC LINER %d/%d" if bearing == "CoC" else "HXL POLY LINER %d/%d") % (head, cs), 1))
    out.append((ms, HEAD_CAT[ms] % (head, "C" if mat == "CERAMIC" else "M"),
                "%s FEMORAL HEAD %dMM +%d" % (mat, head, int(rng.choice([0, 0, 3, 4, 5]))), 1))
    return out


def or_record(T):
    rng = np.random.default_rng(O_SEED)
    pats = T["patients"]
    h26 = [h for h in T["hips"] if h["op"].year == 2026]
    assert all(h["construct"] in K.CODES for h in h26)
    # one case per patient and day: bilateral hips in one operation share it
    by = {}
    for h in sorted(h26, key=lambda x: (x["op"], x["hip"])):
        by.setdefault((h["patient"], h["op"]), []).append(h)
    cases = []
    for (pid, day), hs in by.items():
        side = "Bilateral" if len(hs) == 2 else ("Right" if hs[0]["side"] == "R" else "Left")
        cases.append(dict(kind="THA", day=day, hosp=hs[0]["hospital"], hips=hs, side=side, patient=pid,
                          surgeon=hs[0]["surgeon"]))
    # revisions of hips the system replaced, done at a system hospital in 2026
    revs = []
    reg_rev = {(r["hip"]["hip"], r["n"]): r for r in T["revisions"]}
    for h in T["hips"]:
        p = pats[h["patient"]]
        end = min(p["death"], p["moved"], K.EXTRACT)
        for n, when in ((1, h["rev"]), (2, h["rerev"])):
            if when is None or when.year != 2026 or when > end:
                continue
            r = reg_rev.get((h["hip"], n))
            if r is not None:
                if r["hosp"] == "OTH":
                    continue
                revs.append(dict(kind="RTHA", day=when, hosp=r["hosp"], hip=h, comp=r["comp"], patient=h["patient"]))
            elif (h["rev_where"] == "system" if n == 1 else rng.random() < .8):
                revs.append(dict(kind="RTHA", day=when, hosp=h["hospital"], hip=h,
                                 comp=str(rng.choice(["Head and liner", "Cup", "Stem", "All components"],
                                                     p=[.3, .3, .15, .25])), patient=h["patient"]))
    for r in revs:
        r["side"] = "Right" if r["hip"]["side"] == "R" else "Left"
        r["surgeon"] = "S%s-%02d" % (r["hosp"][0], int(rng.integers(1, 9)))
    cases += revs
    T["or_cases"] = cases
    return cases


def choose_devices(T):
    """Existing hips recorded in a 2026 case (Ask A's primary device) and MRN merges after the operation (Ask C's)."""
    rng = np.random.default_rng(O_SEED + 1)
    pats = T["patients"]
    tha = [c for c in T["or_cases"] if c["kind"] == "THA"]
    single = [c for c in tha if c["side"] != "Bilateral" and len(pats[c["patient"]]["hips"]) == 1]
    used = set()
    phantoms = []
    for code in K.CODES:
        for hosp in K.HOSP:
            pool = [c for c in single if c["hosp"] == hosp and id(c) not in used]
            k = max(1, int(rng.poisson(2.3)))
            k = min(k, 4)
            for i in rng.choice(len(pool), size=k, replace=False):
                c = pool[int(i)]
                used.add(id(c))
                s, cp = R.families(code)
                yr = int(rng.integers(2006, 2018)) if code in ("TA", "MF", "VS") else int(rng.integers(2012, 2018))
                if code == "TS":
                    yr = int(rng.integers(2009, 2018))
                old = yr < 2020
                ph = dict(case=c, construct=code, year=yr, stem=s, cup=cp,
                          ss=int(rng.choice(list(R.STEM_SIZES[s]))), cs=int(rng.choice(list(R.CUP_SIZES[cp]))),
                          head=int(rng.choice([28, 32, 36], p=[.3, .5, .2])),
                          bearing=str(rng.choice(["CoP", "MoP"], p=[.5, .5])), old=old,
                          side="Left" if c["side"] == "Right" else "Right", lot=rng.random() < 0.3)
                c["phantom"] = ph
                phantoms.append(ph)
    T["phantoms"] = phantoms
    # MRN merges: duplicate registrations merged after the operation; the registry kept the MRN on the form
    reg26 = [c for c in tha if all(h["entry"] is not None for h in c["hips"]) and c["side"] != "Bilateral"
             and not c.get("phantom")]
    merged = []
    need = [(code, hosp) for code in K.CODES for hosp in K.HOSP]
    order = list(rng.permutation(len(need)))
    covered_c, covered_h = set(), set()
    for i in order:
        code, hosp = need[int(i)]
        if code in covered_c and hosp in covered_h and len(merged) >= 12:
            continue
        pool = [c for c in reg26 if c["hosp"] == hosp and c["hips"][0]["construct"] == code and id(c) not in used]
        if not pool:
            continue
        c = pool[int(rng.integers(0, len(pool)))]
        used.add(id(c))
        merged.append(c)
        covered_c.add(code)
        covered_h.add(hosp)
    assert covered_c == set(K.CODES) and covered_h == set(K.HOSP)
    merges = []
    for c in merged:
        p = pats[c["patient"]]
        retired = R.mrn_system(rng)
        when = c["day"] + dt.timedelta(days=int(rng.integers(3, 150)))
        when = min(when, dt.date(2026, 12, 30))
        assert when > c["day"]
        p.setdefault("mrn_at_op", {})[c["hips"][0]["hip"]] = retired
        merges.append(dict(on=when, retired=retired, surviving=p["mrn"], site=c["hosp"], device=True))
    # duplicates merged before the operation (the form already carried the surviving number)
    early = [c for c in reg26 if id(c) not in used]
    for i in rng.choice(len(early), size=7, replace=False):
        c = early[int(i)]
        p = pats[c["patient"]]
        when = dt.date(2026, 1, 2) + dt.timedelta(days=int(rng.integers(0, max(1, (c["day"] - dt.date(2026, 1, 2)).days))))
        if when >= c["day"]:
            continue
        merges.append(dict(on=when, retired=R.mrn_system(rng), surviving=p["mrn"], site=c["hosp"], device=False))
    # everyone else HIM merged in 2026
    for _ in range(362):
        when = dt.date(2026, 1, 2) + dt.timedelta(days=int(rng.integers(0, 362)))
        merges.append(dict(on=when, retired=R.mrn_system(rng), surviving=R.mrn_system(rng),
                           site=K.HOSP[int(rng.choice(3, p=[.45, .22, .33]))], device=False))
    merges.sort(key=lambda m: (m["on"], m["retired"]))
    T["merges"] = merges
    T["merged_cases"] = merged
    return T


def write_merges(path, T):
    rng = np.random.default_rng(O_SEED + 2)
    staff = ["HIM%04d" % x for x in (2207, 2231, 2318, 2402, 2466)]
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["merge_id", "merged_on", "retired_mrn", "surviving_mrn", "registration_site", "reason",
                    "merged_by"])
        for i, m in enumerate(T["merges"]):
            w.writerow(["MRG26-%04d" % (i + 1), m["on"].isoformat(), m["retired"], m["surviving"], m["site"],
                        str(rng.choice(["Duplicate registration", "Duplicate registration", "Duplicate registration",
                                        "Newborn or unidentified patient record"], p=[.4, .3, .25, .05])),
                        str(rng.choice(staff))])
    return len(T["merges"])


# ---------------------------------------------------------------------------------------------- the other 2026 cases
def _weekday(rng):
    while True:
        d = dt.date(2026, 1, 1) + dt.timedelta(days=int(rng.integers(0, 365)))
        if d.weekday() < 5 and d not in (dt.date(2026, 1, 1), dt.date(2026, 5, 25), dt.date(2026, 7, 3),
                                         dt.date(2026, 9, 7), dt.date(2026, 11, 26), dt.date(2026, 12, 25)):
            return d


def other_cases(T):
    rng = np.random.default_rng(O_SEED + 3)
    out = []
    plan = [("HEMI", 172), ("TKA", 1648), ("RTKA", 109), ("ORIF", 1243), ("IMN", 297), ("FUS", 604), ("RTHA_OUT", 36)]
    mix = {"HEMI": (.52, .33, .15), "TKA": (.36, .21, .43), "RTKA": (.45, .15, .40), "ORIF": (.55, .30, .15),
           "IMN": (.60, .30, .10), "FUS": (.62, .00, .38), "RTHA_OUT": (.40, .15, .45)}
    for kind, n in plan:
        for _ in range(n):
            hosp = K.HOSP[int(rng.choice(3, p=mix[kind]))]
            side = "N/A" if kind == "FUS" else ("Right" if rng.random() < .52 else "Left")
            out.append(dict(kind="RTHA" if kind == "RTHA_OUT" else kind, outside=kind == "RTHA_OUT", day=_weekday(rng),
                            hosp=hosp, side=side, patient=None, mrn=R.mrn_system(rng),
                            surgeon="S%s-%02d" % (hosp[0], int(rng.integers(1, 13)))))
    T["or_cases"] += out
    return out


def _maker(rng, p=(.3, .3, .25, .15)):
    return ["VNH", "THV", "JUV", "BRY"][int(rng.choice(4, p=p))]


def case_lines(rng, c, pats):
    """Implant documentation for one case: (side, action, maker, catalogue no., description, qty, lot, minutes)."""
    L = []
    t = 25

    def add(side, action, maker, cat, desc, qty=1, lot=None, minute=None):
        nonlocal t
        if minute is None:
            t += int(rng.integers(2, 14))
        L.append((side, action, maker, cat, desc, qty, _lot(rng, maker, 2026) if lot is None else lot,
                  t if minute is None else minute))

    if c["kind"] == "THA":
        for h in c["hips"]:
            side = "Right" if h["side"] == "R" else "Left"
            s, cp = R.families(h["construct"])
            parts = hip_lines(rng, s, cp, h["stem_size"], h["cup_size"], h["head"], h["bearing"], 2026)
            mc = MAKER_OF[cp]
            if rng.random() < .03:
                add(side, "Implanted", parts[0][0], R.cat_no(s, max(min(R.STEM_SIZES[s]), h["stem_size"] - 1)),
                    parts[0][2].replace("SZ %d" % h["stem_size"], "SZ %d" % max(min(R.STEM_SIZES[s]), h["stem_size"] - 1)))
                add(side, "Explanted", L[-1][2], L[-1][3], L[-1][4], lot=L[-1][6])
            if rng.random() < .05:
                ws = h["stem_size"] + (1 if h["stem_size"] < max(R.STEM_SIZES[s]) else -1)
                add(side, "Wasted", parts[0][0], R.cat_no(s, ws), parts[0][2].replace("SZ %d" % h["stem_size"], "SZ %d" % ws))
            add(side, "Implanted", *parts[1])
            for _ in range(int(rng.choice([0, 0, 1, 2, 3], p=[.45, .15, .15, .15, .10])) if cp != "ORS" else 0):
                add(side, "Implanted", mc, {"VNH": "5190-65-025", "THV": "TK-SC65-25", "JUV": "SC-6525",
                                            "BRY": "BYS-65-25"}[mc], "CANCELLOUS BONE SCREW 6.5X25MM")
            add(side, "Implanted", *parts[2])
            if s == "ASK":
                add(side, "Implanted", "THV", "TK-BC40", "BONE CEMENT 40G", qty=2)
            add(side, "Implanted", *parts[0])
            if rng.random() < .03:
                add(side, "Wasted", *parts[3])
            add(side, "Implanted", *parts[3])
        ph = c.get("phantom")
        if ph:
            m0 = int(rng.integers(-12, 4))
            for k, part in enumerate(hip_lines(rng, ph["stem"], ph["cup"], ph["ss"], ph["cs"], ph["head"], ph["bearing"],
                                               ph["year"], old=ph["old"])):
                lot = _lot(rng, part[0], ph["year"]) if ph["lot"] else ""
                L.append((ph["side"], "Implanted", part[0], part[1], part[2], 1, lot, m0 + k))
    elif c["kind"] == "RTHA":
        side = c["side"]
        if c.get("outside"):
            for desc in ("FEMORAL STEM, MANUFACTURER UNKNOWN", "ACETABULAR SHELL, MANUFACTURER UNKNOWN",
                         "POLY LINER, MANUFACTURER UNKNOWN", "FEMORAL HEAD, MANUFACTURER UNKNOWN"):
                if rng.random() < .8:
                    add(side, "Explanted", "", "", desc, lot="")
            comp = str(rng.choice(["Cup", "Stem", "All components"], p=[.4, .25, .35]))
            old = None
        else:
            h = c["hip"]
            comp = c["comp"]
            s, cp = R.families(h["construct"])
            old = hip_lines(rng, s, cp, h["stem_size"], h["cup_size"], h["head"], h["bearing"], h["op"].year,
                            old=h["op"] < K.THORVIK_RECODE)
            take = {"Head and liner": (2, 3), "Cup": (1, 2, 3), "Stem": (0, 3), "All components": (0, 1, 2, 3)}[comp]
            for k in take:
                add(side, "Explanted", old[k][0], old[k][1], old[k][2], lot="")
        head = int(rng.choice([32, 36, 40], p=[.3, .5, .2]))
        cs = int(rng.choice(list(range(50, 68, 2))))
        if comp in ("Cup", "All components"):
            if rng.random() < .45:
                add(side, "Implanted", "JUV", R.cat_no("PLV", min(cs, 66)), "PELLVORO ACETABULAR SHELL %dMM" % min(cs, 66))
            else:
                add(side, "Implanted", "JUV", "PV-2280-%02d" % cs, "PELLVORO MULTIHOLE REVISION SHELL %dMM" % cs)
            for _ in range(int(rng.integers(2, 5))):
                add(side, "Implanted", "JUV", "SC-6530", "CANCELLOUS BONE SCREW 6.5X30MM")
        if comp in ("Stem", "All components"):
            if comp == "All components" and rng.random() < .3:
                ss = int(rng.integers(6, 12))
                add(side, "Implanted", "JUV", R.cat_no("HOL", ss), "HOLBRIN FEMORAL STEM SZ %d UNCEMENTED" % ss)
            else:
                add(side, "Implanted", "JUV", "HB-1180-%02d" % int(rng.integers(13, 19)), "HOLBRIN REVISION STEM 200MM")
        add(side, "Implanted", "JUV", LINER_CAT["JUV"] % (head, min(cs, 66)), "HXL POLY LINER %d/%d" % (head, min(cs, 66)))
        add(side, "Implanted", "JUV", HEAD_CAT["JUV"] % (head, "C"), "CERAMIC FEMORAL HEAD %dMM +%d" % (head, int(rng.choice([0, 4]))))
    elif c["kind"] == "HEMI":
        side = c["side"]
        fam = "HOL" if rng.random() < .6 else "KLV"
        ss = int(rng.choice(list(R.STEM_SIZES[fam])))
        if rng.random() < .5:
            add(side, "Implanted", MAKER_OF[fam], {"HOL": "JV-BC40", "KLV": "BY-BC40"}[fam],
                "BONE CEMENT 40G", qty=2)
        add(side, "Implanted", MAKER_OF[fam], R.cat_no(fam, ss), "%s FEMORAL STEM SZ %d %s" % (
            K.FAM[fam][1].upper(), ss, K.FAM[fam][4].upper()))
        bh = int(rng.integers(42, 56))
        add(side, "Implanted", MAKER_OF[fam], {"HOL": "BH-3500-%02d", "KLV": "BYB-700-%02d"}[fam] % bh,
            "BIPOLAR HEAD %dMM" % bh)
    elif c["kind"] in ("TKA", "RTKA"):
        side = c["side"]
        m = _maker(rng)
        pre = {"VNH": "6210", "THV": "TK-K", "JUV": "KN-40", "BRY": "BYK"}[m]
        if c["kind"] == "RTKA":
            for d in ("KNEE FEMORAL COMPONENT", "KNEE TIBIAL BASEPLATE", "KNEE TIBIAL INSERT"):
                add(side, "Explanted", "", "", d + ", PRIOR IMPLANT", lot="")
        fz, tz = int(rng.integers(1, 10)), int(rng.integers(1, 10))
        add(side, "Implanted", m, "%s-01-%02d" % (pre, fz), "KNEE FEMORAL COMPONENT %s SZ %d" % (
            "REVISION" if c["kind"] == "RTKA" else str(rng.choice(["CR", "PS"])), fz))
        add(side, "Implanted", m, "%s-02-%02d" % (pre, tz), "KNEE TIBIAL BASEPLATE SZ %d" % tz)
        add(side, "Implanted", m, "%s-03-%02d%02d" % (pre, tz, int(rng.choice([9, 10, 11, 12, 14]))), "KNEE TIBIAL INSERT")
        if rng.random() < .82:
            add(side, "Implanted", m, "%s-04-%02d" % (pre, int(rng.integers(26, 41))), "ALL-POLY PATELLA")
        add(side, "Implanted", m, {"VNH": "VH-BC40", "THV": "TK-BC40", "JUV": "JV-BC40", "BRY": "BY-BC40"}[m],
            "BONE CEMENT 40G", qty=2)
        if rng.random() < .04:
            add(side, "Wasted", m, "%s-03-%02d10" % (pre, tz), "KNEE TIBIAL INSERT")
    elif c["kind"] == "ORIF":
        side = c["side"]
        m = _maker(rng, (.3, .4, .2, .1))
        plate = str(rng.choice(["DISTAL RADIUS VOLAR PLATE", "LATERAL ANKLE PLATE", "PROXIMAL HUMERUS PLATE",
                                "CLAVICLE PLATE", "TIBIAL PLATEAU PLATE"]))
        add(side, "Implanted", m, "%s-PL-%03d" % (m, int(rng.integers(100, 400))), "LOCKING " + plate)
        for _ in range(int(rng.integers(4, 11))):
            ln = int(rng.choice([12, 14, 16, 18, 20, 22, 24, 26, 30, 36, 40]))
            add(side, "Implanted", m, "%s-SC-%s%02d" % (m, str(rng.choice(["L", "C"])), ln),
                "%s SCREW 3.5X%dMM" % (str(rng.choice(["LOCKING", "CORTICAL"])), ln), qty=int(rng.choice([1, 1, 1, 2])))
        if rng.random() < .06:
            add(side, "Wasted", m, "%s-SC-L%02d" % (m, 20), "LOCKING SCREW 3.5X20MM")
    elif c["kind"] == "IMN":
        side = c["side"]
        m = _maker(rng, (.25, .45, .2, .1))
        add(side, "Implanted", m, "%s-NL-%03d" % (m, int(rng.integers(300, 460))), "INTRAMEDULLARY NAIL %dX%dMM" % (
            int(rng.choice([9, 10, 11, 12])), int(rng.choice([300, 340, 360, 380, 400]))))
        for _ in range(int(rng.integers(2, 5))):
            add(side, "Implanted", m, "%s-LS-%02d" % (m, int(rng.integers(30, 80))), "LOCKING SCREW 5.0MM")
        add(side, "Implanted", m, "%s-EC-00" % m, "NAIL END CAP")
    elif c["kind"] == "FUS":
        m = "JUV" if rng.random() < .5 else "VNH"
        for _ in range(int(rng.choice([4, 4, 6, 6, 8]))):
            add("", "Implanted", m, "%s-PS-%02d%02d" % (m, int(rng.choice([55, 65, 75])), int(rng.choice([40, 45, 50]))),
                "PEDICLE SCREW %s" % str(rng.choice(["5.5X40MM", "6.5X45MM", "6.5X50MM", "7.5X45MM"])))
        add("", "Implanted", m, "%s-RD-%03d" % (m, int(rng.integers(40, 120))), "ROD 5.5MM PRE-BENT", qty=2)
        add("", "Implanted", m, "%s-SS-00" % m, "SET SCREW", qty=int(rng.choice([4, 6, 8])))
        for _ in range(int(rng.choice([1, 1, 2]))):
            add("", "Implanted", m, "%s-CG-%02d" % (m, int(rng.integers(8, 15))), "INTERBODY CAGE PEEK")
    return L


def build_log(T):
    rng = np.random.default_rng(O_SEED + 4)
    pats = T["patients"]
    cases = sorted(T["or_cases"], key=lambda c: (c["day"], c["hosp"], c["kind"], c.get("surgeon", ""),
                                                 c.get("mrn") or pats[c["patient"]]["mrn"],
                                                 c.get("side", "")))
    seq = {h: int(rng.integers(1100, 1900)) for h in K.HOSP}
    rooms_used = {}
    for c in cases:
        seq[c["hosp"]] += int(rng.integers(1, 4))
        c["number"] = "26-%s-%05d" % (c["hosp"], seq[c["hosp"]])
        c["room"] = str(rng.choice(ROOMS[c["hosp"]]))
        k = (c["hosp"], c["room"], c["day"])
        rooms_used[k] = rooms_used.get(k, 0) + 1
        start = 7 * 60 + 15 + (rooms_used[k] - 1) * int(rng.integers(95, 150))
        c["start"] = min(start, 18 * 60)
        if c["patient"] is not None:
            c["mrn"] = pats[c["patient"]]["mrn"]
        c["klass"] = ("Outpatient" if rng.random() < .34 else "Inpatient") if c["kind"] in ("THA", "TKA") else \
            ("Outpatient" if c["kind"] == "ORIF" and rng.random() < .55 else "Inpatient")
    lines = []
    nurses = {h: ["RN%05d" % x for x in rng.integers(10000, 99999, size=14)] for h in K.HOSP}
    for c in cases:
        nurse = str(rng.choice(nurses[c["hosp"]]))
        for side, action, maker, cat, desc, qty, lot, minute in case_lines(rng, c, pats):
            if action == "Implanted" and lot and rng.random() < .03:
                lot = ""
            at = dt.datetime.combine(c["day"], dt.time(0, 0)) + dt.timedelta(minutes=c["start"] + minute)
            lines.append(dict(case=c, side=side, action=action, maker=maker, cat=cat, desc=desc, qty=qty, lot=lot,
                              at=at, by=nurse))
    lines.sort(key=lambda x: (x["at"], x["case"]["number"]))
    n = int(rng.integers(4100000, 4200000))
    for x in lines:
        n += int(rng.integers(1, 3))
        x["id"] = "IMP%07d" % n
    T["or_cases"] = cases
    T["log"] = lines
    return lines


def write_log(path, T):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["log_line_id", "case_number", "hospital_code", "surgery_date", "procedure", "implant_side", "action",
                    "manufacturer", "catalogue_no", "description", "lot_no", "quantity", "documented_at",
                    "documented_by"])
        for x in T["log"]:
            c = x["case"]
            w.writerow([x["id"], c["number"], c["hosp"], c["day"].isoformat(), PROC[c["kind"]], x["side"], x["action"],
                        K.MAKERS.get(x["maker"], "Unknown"), x["cat"], x["desc"], x["lot"], x["qty"],
                        x["at"].strftime("%Y-%m-%d %H:%M"), x["by"]])
    return len(T["log"])


def write_cases(path, T):
    wb = Workbook()
    ws = wb.active
    ws.title = "Cases 2026"
    R._head(ws, ["Case number", "Hospital", "Room", "Surgery date", "Procedure", "Side", "Primary surgeon", "MRN",
                 "Patient class", "Case status"], [16, 30, 9, 12, 40, 10, 14, 11, 13, 12])
    for c in T["or_cases"]:
        ws.append([c["number"], K.HOSP_NAME[c["hosp"]], c["room"], c["day"], PROC[c["kind"]], c["side"],
                   c["surgeon"], c["mrn"], c["klass"], "Completed"])
        ws.cell(ws.max_row, 4).number_format = "mm/dd/yyyy"
    wb.save(path)
    return len(T["or_cases"])
