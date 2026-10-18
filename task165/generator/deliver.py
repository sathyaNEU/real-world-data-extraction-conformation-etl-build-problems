"""task165 deliverables: the golden workbook and memo, generator/figures.json and submission.md, every figure computed
from the shipped pack by calc.py."""
from __future__ import annotations

import datetime as dt
import io
import json
import os
import sys

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt                                      # noqa: E402
import numpy as np                                                   # noqa: E402
from docx import Document                                            # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH                         # noqa: E402
from docx.shared import Mm, Pt, RGBColor                              # noqa: E402
from openpyxl import Workbook                                        # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill              # noqa: E402
from openpyxl.utils import get_column_letter                         # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import calc as C            # noqa: E402
import common as K          # noqa: E402
import postprocess as PP    # noqa: E402

WINTERS3 = ["2023-24", "2024-25", "2025-26"]
YNAME = {"KLV": "Kelvin Road", "MPF": "Mapleford", "RDG": "Ridge", "LKS": "Lakeshore"}
SERIES = ["#2a78d6", "#eb6834", "#1baf7a"]      # reference categorical slots 1 to 3
F_WB = "salt_commitment_2026-27.xlsx"
F_MEMO = "seasonal_salt_commitment_memo.docx"


def r1(v):
    return float("%.1f" % (v + 0.0))


def compute(target):
    P = C.Pack(target)
    raw = C.golden(P)
    nn = P.new_rates()
    L = float(nn["lane_km"].sum())
    x0 = float((nn["p2"] * nn["lane_km"]).sum() / L)
    b = C.gain_coefficients(P)
    d = P.prepost()
    stock, room = C.ask_a(target)
    art = C.arterial_salt(target)
    per_seg = []
    reg = P.reg.set_index("segment_id")
    for _, r in nn.iterrows():
        g = (b[0] + b[1] * r["p2"]) * r["lane_km"] * C.PLAN_SEV
        per_seg.append((r["segment_id"], reg.loc[r["segment_id"], "street"], reg.loc[r["segment_id"], "from_street"],
                        reg.loc[r["segment_id"], "to_street"], r["yard"], float(r["lane_km"]),
                        float(r["p2"] * C.PLAN_SEV), float(g)))
    return {"raw": raw, "commitment": int(round(raw / 100.0) * 100), "lane_km": L, "own_p2": x0 * C.PLAN_SEV,
            "gain_per_lkm": raw / L, "b0": float(b[0]), "b1": float(b[1]), "n_moves": len(d),
            "moves_lkm": float(d["lane_km"].sum()), "stock": stock, "room": room, "art": art, "per_seg": per_seg,
            "wsi": P.wsi}


# ---------------------------------------------------------------------------------------------- workbook
BOLD = Font(bold=True)
HEAD = PatternFill("solid", fgColor="DCE6F1")


def _head(ws, row, vals, widths=None):
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=row, column=j, value=v)
        c.font = BOLD
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="top")
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w


def workbook(g, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Commitment"
    ws["A1"] = "Commit {:,} t of rock salt above the rolled-over 39,100 t for November 2026 to March 2027".format(
        g["commitment"])
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = "Roads Division, winter operations. Prepared 28 September 2026 for Council."
    rows = [
        ("Contract quantity rolled over on 1 October 2026 (t)", 39100),
        ("Service change since the 2025-26 winter closed", "Collectors moved to Priority 1 on 30 August 2026 (QRT frequent network)"),
        ("Lane-km moved", round(g["lane_km"], 2)),
        ("Planning severity (WSI points)", 61.7),
        ("Their Priority 2 rock salt in a planning-severity winter (t per lane-km)", r1(g["own_p2"])),
        ("Gain on moving to Priority 1 in a planning-severity winter (t per lane-km)", r1(g["gain_per_lkm"])),
        ("Seasonal tonnage to commit (t, nearest 100)", g["commitment"]),
    ]
    fmts = ["#,##0", "@", "0.0", "0.0", "0.0", "0.0", "#,##0"]
    for i, (k, v) in enumerate(rows, 4):
        ws.cell(row=i, column=1, value=k)
        c = ws.cell(row=i, column=2, value=v)
        c.number_format = fmts[i - 4]
        if isinstance(v, (int, float)):
            c.alignment = Alignment(horizontal="right")
        if i == 4 + len(rows) - 1:
            ws.cell(row=i, column=1).font = BOLD
            c.font = BOLD
    ws.column_dimensions["A"].width = 72
    ws.column_dimensions["B"].width = 62
    r0 = 4 + len(rows) + 1
    ws.cell(row=r0, column=1, value="Gain line from the %d collectors moved to Priority 1 from 1 November 2022 to "
                                    "1 November 2025 (%.1f lane-km), each street's salt per lane-km per WSI point "
                                    "after its move less before, against before, weighted by lane-km"
            % (g["n_moves"], g["moves_lkm"]))
    ws.cell(row=r0 + 1, column=1, value="Gain at no Priority 2 salt, planning-severity winter (t per lane-km)")
    ws.cell(row=r0 + 1, column=2, value=round(g["b0"] * C.PLAN_SEV, 2)).number_format = "0.00"
    ws.cell(row=r0 + 2, column=1, value="Change in gain per tonne a lane-km of Priority 2 salt")
    ws.cell(row=r0 + 2, column=2, value=round(g["b1"], 3)).number_format = "0.000"
    ws2 = wb.create_sheet("By segment")
    _head(ws2, 1, ["Segment", "Street", "From", "To", "Yard", "Lane-km", "Priority 2 salt, planning winter (t per lane-km)",
                   "Gain, planning winter (t)"], [11, 24, 24, 24, 7, 9, 22, 14])
    for i, row in enumerate(sorted(g["per_seg"], key=lambda r: (r[4], r[1], r[0])), 2):
        sid, st, fr, to, y, lk, p2, gain = row
        for j, v in enumerate([sid, st, fr, to, y, lk, round(p2, 2), round(gain, 1)], 1):
            c = ws2.cell(row=i, column=j, value=v)
            if j >= 6:
                c.number_format = ["0.00", "0.00", "0.0"][j - 6]
    n = len(g["per_seg"]) + 2
    ws2.cell(row=n, column=1, value="Total").font = BOLD
    c = ws2.cell(row=n, column=6, value=round(g["lane_km"], 2))
    c.font, c.number_format = BOLD, "0.00"
    c = ws2.cell(row=n, column=8, value=round(sum(r[7] for r in g["per_seg"]), 1))
    c.font, c.number_format = BOLD, "#,##0.0"
    ws2.freeze_panes = "B2"
    ws2.print_title_rows = "1:1"
    ws2.page_setup.orientation = "landscape"
    ws2.page_setup.fitToWidth = 1
    ws2.auto_filter.ref = "A1:H%d" % (n - 1)
    ws3 = wb.create_sheet("Domes 1 Nov 2026")
    _head(ws3, 1, ["Yard", "Rock salt in the dome on 1 November 2026 (t)", "Tonnes more the dome can take (t)"],
          [16, 28, 28])
    for i, y in enumerate(C.YARDS, 2):
        ws3.cell(row=i, column=1, value=YNAME[y])
        ws3.cell(row=i, column=2, value=r1(g["stock"][y])).number_format = "#,##0.0"
        ws3.cell(row=i, column=3, value=r1(g["room"][y])).number_format = "#,##0.0"
    ws3.cell(row=7, column=1, value="Survey stock brought forward by every weighing in to and out of each dome after "
                                    "its capture, and the October bookings in the 21 September dispatch issue.")
    ws4 = wb.create_sheet("Arterial salt")
    _head(ws4, 1, ["Yard"] + ["%s (t)" % w for w in WINTERS3], [16, 14, 14, 14])
    for i, y in enumerate(C.YARDS, 2):
        ws4.cell(row=i, column=1, value=YNAME[y])
        for j, w in enumerate(WINTERS3, 2):
            ws4.cell(row=i, column=j, value=r1(g["art"][(y, w)])).number_format = "#,##0.0"
    ws4.cell(row=7, column=1, value="Rock salt spread on each yard's Priority 1 arterial circuits, 1 November to 31 "
                                    "March, contract circuits included.")
    for wsx in (ws3, ws4):
        wsx.freeze_panes = "A2"
    ws5 = wb.create_sheet("Notes")
    notes = [
        "Data pulled 25 September 2026: legacy AVL segment-by-shift summaries (November 2021 to March 2023, kilograms, "
        "local time) and telematics segment passes (from 1 October 2023, tonnes, UTC, taken to Eastern time).",
        "Legacy segment ids taken to the current segments through the 2023 GIS crosswalk by length share.",
        "Winters run 1 November to 31 March. Each winter's salt is divided by its Winter Severity Index from the storm "
        "log (2021-22 63.4, 2022-23 50.1, 2023-24 38.9, 2024-25 67.2, 2025-26 52.6 points).",
        "Planning severity 61.7 WSI points (winter maintenance plan 2026-27, section 4). Seasonal tonnage per RO-PR-07 "
        "4.2, rounded to the nearest 100 t.",
        "Collector figures are rock salt (RS) only; the collector and local circuits spread dry salt.",
        "Dome stock: survey report TG-2026-0917 brought forward by the scale house weighings after each capture, loads "
        "weighed at another yard's scale counted at the dome named in the remarks, and the 21 September dispatch issue.",
        "Arterial salt: older-controller runs counted once with the spinner brine taken off at 1.18 kg a litre; Kettlewood "
        "Road Services circuits from its monthly returns; each circuit under the yard that ran it that winter.",
    ]
    ws5.column_dimensions["A"].width = 120
    for i, t in enumerate(notes, 1):
        ws5.cell(row=i, column=1, value=t).alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)
    PP.normalize_ooxml(path, K.PEOPLE["winter_manager"], K.PEOPLE["winter_manager"], "2026-09-25T18:20:00Z",
                       "2026-09-28T19:05:00Z", excel=True, company=K.ORG)


# ---------------------------------------------------------------------------------------------- chart and memo
def chart_png(g):
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 8.5, "svg.hashsalt": "task165"})
    fig, ax = plt.subplots(figsize=(6.3, 2.9), dpi=150)
    x = np.arange(len(C.YARDS))
    wd = 0.26
    for k, w in enumerate(WINTERS3):
        vals = [g["art"][(y, w)] for y in C.YARDS]
        bars = ax.bar(x + (k - 1) * (wd + 0.02), vals, wd, color=SERIES[k], label=w, zorder=2)
        for b_, v in zip(bars, vals):
            ax.text(b_.get_x() + b_.get_width() / 2, v + 120, "{:,.1f}".format(v), ha="center", va="bottom",
                    fontsize=6.2, color="#33332f", rotation=90)
    ax.set_xticks(x)
    ax.set_xticklabels([YNAME[y] for y in C.YARDS])
    ax.set_ylabel("Rock salt (t)", color="#52514e")
    ax.set_ylim(0, max(g["art"].values()) * 1.22)
    ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, p: "{:,.0f}".format(v)))
    ax.grid(axis="y", color="#e2e1dc", linewidth=0.6, zorder=0)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.spines["left"].set_color("#b9b8b2")
    ax.spines["bottom"].set_color("#b9b8b2")
    ax.tick_params(colors="#52514e")
    ax.set_title("Priority 1 arterial rock salt by yard: 2024-25 the heaviest winter at all four yards", fontsize=9,
                 color="#0b0b0b", loc="left")
    ax.legend(frameon=False, ncol=3, loc="upper left", fontsize=7.5)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", metadata={"Software": None})
    plt.close(fig)
    return buf.getvalue()


def memo(g, path):
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Mm(215.9), Mm(279.4)
    sec.left_margin = sec.right_margin = Mm(22)
    sec.top_margin, sec.bottom_margin = Mm(18), Mm(16)
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)

    def para(text, bold=False, size=None, after=6, color=None):
        p = d.add_paragraph()
        r = p.add_run(text)
        r.bold = bold
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        p.paragraph_format.space_after = Pt(after)
        return p

    para("City of Quellbrook, Public Works, Roads Division", bold=True, size=9, after=0, color="1F4E79")
    para("Report to Council", bold=True, size=15, after=4)
    for k, v in (("To", "Mayor and Members of Council"),
                 ("From", "%s, manager, winter operations" % K.PEOPLE["winter_manager"]),
                 ("Date", "28 September 2026"),
                 ("Subject", "{:,} t of seasonal rock salt for 2026-27, for approval before 1 October".format(
                     g["commitment"]))):
        p = d.add_paragraph()
        r = p.add_run("%s: " % k)
        r.bold = True
        p.add_run(v)
        p.paragraph_format.space_after = Pt(0)
    para("", after=2)
    para("Recommendation", bold=True, size=10.5, after=1, color="1F4E79")
    para("That Council approve a seasonal commitment of {:,} t of rock salt with Northgrit above the rolled-over "
         "39,100 t contract, for November 2026 to March 2027.".format(g["commitment"]), bold=True, after=6)
    para("Analysis", bold=True, size=10.5, after=1, color="1F4E79")
    para("The commitment covers the collectors the route audit moved to Priority 1 on 30 August for Quellbrook "
         "Regional Transit's frequent network, {:.1f} lane-km, sized as RO-PR-07 4.2 sets it out: the extra salt "
         "those streets will need over what they took at Priority 2, in a winter of 61.7 WSI points.".format(
             g["lane_km"]), after=6)
    para("At Priority 2 our trucks salted these streets at their intersections and little else, {:.1f} t a "
         "lane-km in a winter of that severity. The collectors we moved to Priority 1 between 2022 and 2025 were "
         "mostly hill, school and hospital streets that Priority 2 already salted along much of their length, and "
         "their spreader records show the gain from the move falling the more Priority 2 salt a street already took. "
         "Read at these streets' own Priority 2 salt, the move adds {:.1f} t a lane-km.".format(
             r1(g["own_p2"]), r1(g["gain_per_lkm"])), after=6)
    para("The working, each yard's dome stock and room for 1 November, and the arterial figures charted below are in "
         "salt_commitment_2026-27.xlsx.", after=6)
    pic = io.BytesIO(chart_png(g))
    d.add_picture(pic, width=Mm(168))
    d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para("Contract circuits are included; each yard's figure is for the arterial circuits it ran that winter. WSI is "
         "the Roads Operations Centre's Winter Severity Index (plan section 4).", size=8.5, after=0, color="52514E")
    d.save(path)
    res = PP.normalize_docx(path, K.PEOPLE["winter_manager"], K.PEOPLE["winter_manager"], "2026-09-28T13:40:00Z",
                            "2026-09-28T19:10:00Z", revision=5, total_time=48, seed="memo-salt-2026",
                            title="Seasonal rock salt 2026-27", company=K.ORG)
    assert res["pages"] == 1, res["pages"]


# ---------------------------------------------------------------------------------------------- figures and write-up
def figures(g):
    f = {"commitment_t": g["commitment"], "rollover_t": 39100, "lane_km": round(g["lane_km"], 1),
         "segments": len(g["per_seg"]), "planning_wsi": 61.7, "own_p2_t_per_lkm": r1(g["own_p2"]),
         "gain_t_per_lkm": r1(g["gain_per_lkm"]), "earlier_moves": g["n_moves"], "procedure_clause": "4.2",
         "rounding_unit_t": 100}
    for y in C.YARDS:
        f["stock_%s" % y] = r1(g["stock"][y])
        f["room_%s" % y] = r1(g["room"][y])
        for w in WINTERS3:
            f["art_%s_%s" % (y, w)] = r1(g["art"][(y, w)])
    return f


def write_submission(path, g):
    f = figures(g)
    L = []
    L.append("## Tags\n")
    L.append("**Domain:** Urban & Regional Science (municipal service demand and resource allocation: a city's winter "
             "road salt for a change of service level, forecast from the spreader records of the streets that went "
             "through the same change before).")
    L.append("**Analytical objective:** Forecasting & Predictive Modeling.\n")
    L.append("## 1. Final Recommendation\n")
    L.append("**Commit {:,} t of rock salt above the rolled-over 39,100 t contract for November 2026 to March 2027.** "
             "The redesign's collectors took {:.1f} t a lane-km of Priority 2 salt in a 61.7-point winter, far less "
             "than the collectors moved to Priority 1 before them, and read off those earlier moves at each "
             "street's own Priority 2 salt the move adds {:.1f} t a lane-km over {:.1f} lane-km. Not the Priority 1 "
             "collectors' rate less these streets' own salt, which carries the higher Priority 1 level of streets with "
             "more intersections, hills and curves onto flat ones; not the earlier moves' average gain, which belongs "
             "to streets Priority 2 had already salted along much of their length.\n".format(
                 f["commitment_t"], f["own_p2_t_per_lkm"], f["gain_t_per_lkm"], g["lane_km"]))
    L.append("## 2. Critical Components\n")
    L.append("1. The {:.1f} lane-km's own Priority 2 rock salt in a winter of 61.7 WSI points: **{:.1f} t per "
             "lane-km**.".format(g["lane_km"], f["own_p2_t_per_lkm"]))
    L.append("2. Their gain on moving to Priority 1 in that winter: **{:.1f} t per lane-km**.\n".format(
        f["gain_t_per_lkm"]))
    L.append("## 3. Step-by-Step Solution\n")
    L.append("1. Took the {} segments route_audit_revision_log.csv moves from priority 2 to 1 on 2026-08-30, {:.1f} "
             "lane-km in winter_route_register_2026-09-14.csv.".format(f["segments"], g["lane_km"]))
    L.append("2. Summed each collector segment's rock salt from 1 November to 31 March in each winter, from "
             "avl_legacy_segment_shift_2021-11-01_2023-03-31.csv (kilograms, legacy ids through "
             "gis_segment_crosswalk_2023.csv) and telematics_segment_passes_2023-10-01_2026-09-23.parquet (UTC to local), "
             "per lane-km and per point of the winter's WSI in storm_log_2021-22_to_2025-26.xlsx.")
    L.append("3. For the {} collectors the revision log moves from priority 2 to 1 on 1 November 2022, 2023, 2024 and "
             "2025, took each street's salt per lane-km per point in its winters before and after the move.".format(
                 f["earlier_moves"]))
    L.append("4. Fitted each moved street's gain against its own Priority 2 salt, weighted by lane-km: the gain falls "
             "as a street's Priority 2 salt rises.")
    L.append("5. Read the line at the new segments' own Priority 2 salt and scaled to 61.7 WSI points "
             "(winter_maintenance_plan_2026-27.pdf, section 4) over their {:.1f} lane-km, the additional salt "
             "road_salt_procurement_procedure_RO-PR-07.pdf 4.2 commits: {:,} t to the nearest 100 t.\n".format(
                 g["lane_km"], f["commitment_t"]))
    L.append("## 4. Deliverable Answers\n")
    L.append("### %s\n" % F_WB)
    L.append("1. Seasonal tonnage above the rolled-over contract: {:,} t.".format(f["commitment_t"]))
    L.append("2. Rock salt in each yard's dome on 1 November 2026 once the October deliveries booked with the supplier "
             "have landed, and the tonnes more the dome can then take:")
    for y in C.YARDS:
        L.append("   - {}: {:,.1f} t in the dome, {:,.1f} t more".format(YNAME[y], f["stock_%s" % y], f["room_%s" % y]))
    L.append("3. Rock salt spread on each yard's Priority 1 arterial circuits:")
    for y in C.YARDS:
        L.append("   - {}: 2023-24 {:,.1f} t, 2024-25 {:,.1f} t, 2025-26 {:,.1f} t".format(
            YNAME[y], f["art_%s_2023-24" % y], f["art_%s_2024-25" % y], f["art_%s_2025-26" % y]))
    L.append("\n### %s\n" % F_MEMO)
    L.append("1. One-page report to council stating the seasonal commitment of {:,} t, with a grouped bar "
             "chart of the arterial figures, each yard's three winters as labelled bars.".format(f["commitment_t"]))
    open(path, "w", encoding="utf-8").write("\n".join(L) + "\n")
    return f


def main(target, golden_dir, figures_path, submission_path):
    os.makedirs(golden_dir, exist_ok=True)
    g = compute(target)
    workbook(g, os.path.join(golden_dir, F_WB))
    memo(g, os.path.join(golden_dir, F_MEMO))
    f = write_submission(submission_path, g)
    json.dump(f, open(figures_path, "w"), indent=1, sort_keys=True)
    return g, f


if __name__ == "__main__":
    task = os.path.dirname(HERE)
    main(os.path.join(task, "target"), os.path.join(task, "golden"), os.path.join(HERE, "figures.json"),
         os.path.join(task, "submission.md"))
