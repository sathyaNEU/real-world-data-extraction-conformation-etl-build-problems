"""task165 independent verifier. Reads only the shipped bundle (target/), the goldens, figures.json and submission.md,
on its own code path (csv, pyarrow, openpyxl, pypdf and python-docx readers, its own season, clock and regression
code; nothing from the generator), and recomputes the call, both components, the two rivals the write-up names,
every ask figure and each figure's distance to its rounding edge. Every format the pack ships is read back, with a
positive control on the same extractor for each document that carries a filed sentence.

    python3 verify.py [target]"""
from __future__ import annotations

import csv
import datetime as dt
import json
import os
import re
import sys
from zoneinfo import ZoneInfo

import numpy as np

HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)
TARGET = sys.argv[1] if len(sys.argv) > 1 else os.path.join(TASK, "target")
TZ = ZoneInfo("America/Toronto")
WINTERS = ["2021-22", "2022-23", "2023-24", "2024-25", "2025-26"]
OUT = []


def ok(label, cond, detail=""):
    OUT.append(bool(cond))
    print("%s %-72s %s" % ("PASS" if cond else "FAIL", label[:72], detail))


def season(d: dt.date):
    if d.month >= 11:
        return "%d-%02d" % (d.year, (d.year + 1) % 100)
    if d.month <= 3:
        return "%d-%02d" % (d.year - 1, d.year % 100)
    return None


def edge(v, u):
    return u / 2 - abs(v - round(v / u) * u)


def rows_csv(name):
    with open(os.path.join(TARGET, name), encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def pdf(name):
    import pypdf
    r = pypdf.PdfReader(os.path.join(TARGET, name))
    return re.sub(r"\s+", " ", " ".join(p.extract_text() for p in r.pages))


def wfit(x, y, w):
    """Weighted least squares line y = a + b x."""
    sw = w.sum()
    mx, my = (w * x).sum() / sw, (w * y).sum() / sw
    b = (w * (x - mx) * (y - my)).sum() / (w * (x - mx) ** 2).sum()
    return my - b * mx, b


def main():
    import openpyxl
    import pyarrow.parquet as pq

    # ---------------------------------------------------------------------------------------- documents
    plan = pdf("winter_maintenance_plan_2026-27.pdf")
    proc = pdf("road_salt_procurement_procedure_RO-PR-07.pdf")
    surv = pdf("dome_stock_survey_2026-09.pdf")
    ok("PDF plan read back (positive control: its title)", "Winter maintenance plan 2026-27" in plan)
    m = re.search(r"planning severity of ([\d.]+) WSI points", plan)
    sev = float(m.group(1)) if m else None
    ok("PDF plan files the planning severity", sev == 61.7, str(sev))
    ok("PDF plan files the dome stock rule and the per lane-km rates",
       "brought forward by the loads weighed in to it and out of it at the scale houses" in plan and
       "set in kilograms per lane-km" in plan)
    m = re.search(r"plant brine at (\d+) litres to the tonne of salt \(23 per cent salt by weight, ([\d.]+) kg a litre\)",
                  plan)
    prewet, brine = (float(m.group(1)), float(m.group(2))) if m else (None, None)
    ok("PDF plan files the pre-wet rate and the brine's weight", prewet == 40.0 and brine == 1.18,
       "%s, %s" % (prewet, brine))
    ok("PDF procedure read back (positive control: its title)", "Road salt procurement" in proc)
    ok("PDF procedure 4.2 files the estimand and the rounding",
       "the additional rock salt the roads the change moves will need from 1 November to 31 March, over what they "
       "needed before it, in a winter at the planning severity" in proc and "rounded to the nearest 100 t" in proc)
    ok("PDF survey read back (positive control: its title)", "Salt dome stock survey" in surv)

    # ---------------------------------------------------------------------------------------- register, log, WSI
    reg = {r["segment_id"]: r for r in rows_csv("winter_route_register_2026-09-14.csv")}
    rev = rows_csv("route_audit_revision_log.csv")
    moves = [(r["segment_id"], dt.date.fromisoformat(r["effective_date"])) for r in rev
             if r["field"] == "PRIORITY" and r["old_value"] == "2" and r["new_value"] == "1"
             and reg[r["segment_id"]]["func_class"] == "COLLECTOR"]
    new = [s for s, e in moves if e == dt.date(2026, 8, 30)]
    old = [(s, e) for s, e in moves if e < dt.date(2026, 8, 30)]
    lk = {s: float(r["lane_km"]) for s, r in reg.items()}
    figs0 = json.load(open(os.path.join(HERE, "figures.json")))
    ok("CSV register and revision log: the new segments, 424.0 lane-km, and the earlier moves",
       len(new) == figs0["segments"] and abs(sum(lk[s] for s in new) - 424.0) < 1e-6 and
       len(old) == figs0["earlier_moves"], "%d new, %d earlier" % (len(new), len(old)))
    wb = openpyxl.load_workbook(os.path.join(TARGET, "storm_log_2021-22_to_2025-26.xlsx"), read_only=True)
    wsi, season_rows = {}, {}
    for w in WINTERS:
        rows = list(wb[w].iter_rows(values_only=True))
        hi = next(i for i, r in enumerate(rows) if r and any(str(c).strip().lower() in ("points", "wsi_points")
                                                             for c in r if c is not None))
        col = [str(c).strip().lower() if c is not None else "" for c in rows[hi]].index(
            "points" if "points" in [str(c).strip().lower() for c in rows[hi] if c is not None] else "wsi_points")
        tot = 0.0
        for r in rows[hi + 1:]:
            if r[0] is None or str(r[0]).lower().startswith("season"):
                if r[0] is not None:
                    season_rows[w] = float(r[col])
                continue
            tot += float(r[col])
        wsi[w] = round(tot, 6)
    wb.close()
    ok("XLSX storm log: five winters' WSI, Season WSI rows tie",
       all(abs(season_rows[w] - wsi[w]) < 1e-6 for w in season_rows) and len(season_rows) == 3,
       str({w: round(v, 1) for w, v in wsi.items()}))

    # ---------------------------------------------------------------------------------------- salt per segment-winter
    xw = {}
    for r in rows_csv("gis_segment_crosswalk_2023.csv"):
        xw.setdefault(r["legacy_segment_id"], []).append((r["segment_id"], float(r["length_share"])))
    salt = {}
    for r in rows_csv("avl_legacy_segment_shift_2021-11-01_2023-03-31.csv"):
        if r["MATL"] != "SALT":
            continue
        w = season(dt.date.fromisoformat(r["SHIFT_DATE"]))
        if w is None:
            continue
        for sid, sh in xw[r["SEG_ID"]]:
            salt[(sid, w)] = salt.get((sid, w), 0.0) + float(r["KG"]) / 1000.0 * sh
    t = pq.read_table(os.path.join(TARGET, "telematics_segment_passes_2023-10-01_2026-09-23.parquet"))
    cols = {c: t.column(c).to_pylist() for c in t.column_names}
    n_pass = len(cols["segment_id"])
    loc = [e.astimezone(TZ) for e in cols["entry_utc"]]
    for i in range(n_pass):
        if cols["material_code"][i] != "RS" or "-A" in cols["route_code"][i]:
            continue
        w = season(loc[i].date())
        if w is None:
            continue
        k = (cols["segment_id"][i], w)
        salt[k] = salt.get(k, 0.0) + cols["applied_t"][i]
    ok("PARQUET pass export read back (UTC stamps)", n_pass > 25000 and
       cols["entry_utc"][0].tzinfo is not None, "%d rows" % n_pass)

    def rate(sid, ws):
        return sum(salt.get((sid, w), 0.0) for w in ws) / (lk[sid] * sum(wsi[w] for w in ws))

    # ---------------------------------------------------------------------------------------- the call
    X, G, Wt = [], [], []
    for sid, e in old:
        first = "%d-%02d" % (e.year, (e.year + 1) % 100)
        pre = WINTERS[:WINTERS.index(first)]
        post = WINTERS[WINTERS.index(first):]
        X.append(rate(sid, pre))
        G.append(rate(sid, post) - rate(sid, pre))
        Wt.append(lk[sid])
    X, G, Wt = map(np.array, (X, G, Wt))
    a, b = wfit(X, G, Wt)
    Ln = sum(lk[s] for s in new)
    x0 = sum(rate(s, WINTERS) * lk[s] for s in new) / Ln
    gold = (a + b * x0) * Ln * sev
    call = int(round(gold / 100.0) * 100)
    figs = json.load(open(os.path.join(HERE, "figures.json")))
    ok("CALL the seasonal tonnage recomputes", call == figs["commitment_t"], "%.1f t -> %d" % (gold, call))
    ok("CALL a quarter-bin or more from its rounding edge", edge(gold, 100) >= 25, "%.1f t" % edge(gold, 100))
    own, gain = x0 * sev, gold / Ln
    ok("CALL components recompute (own Priority 2 salt in all and a lane-km, gain per lane-km)",
       int(round(own * Ln / 100.0) * 100) == figs["own_p2_t"] and own < figs["own_p2_t_per_lkm_under"] and
       int(round(gain)) == figs["gain_t_per_lkm"] and edge(own * Ln, 100) >= 15 and edge(gain, 1.0) >= 0.2,
       "%.1f t, %.3f, %.3f" % (own * Ln, own, gain))
    per_w = [sum(salt.get((s, w), 0.0) for s in new) / (Ln * wsi[w]) * sev for w in WINTERS]
    bases = {"pooled": own * Ln, "mean of the winters": float(np.mean(per_w)) * Ln,
             "mean of the streets": float(np.mean([rate(s, WINTERS) for s in new])) * sev * Ln}
    ok("CALL own Priority 2 salt over the five winters files the same figure on every basis",
       all(int(round(v / 100.0) * 100) == figs["own_p2_t"] and v / Ln < figs["own_p2_t_per_lkm_under"]
           for v in bases.values()), str({k: round(v, 1) for k, v in bases.items()}))
    # the rivals block 1 names
    tot = den = 0.0
    oldd = dict(old)
    for sid, r in reg.items():
        if r["func_class"] != "COLLECTOR" or r["priority"] != "1" or sid in new:
            continue
        ws = WINTERS
        if sid in oldd:
            e = oldd[sid]
            ws = WINTERS[WINTERS.index("%d-%02d" % (e.year, (e.year + 1) % 100)):]
        for w in ws:
            tot += salt.get((sid, w), 0.0)
            den += lk[sid] * wsi[w]
    r1 = (tot / den - x0) * Ln * sev
    r2 = (G * Wt).sum() / Wt.sum() * Ln * sev
    ok("RIVAL the Priority 1 collectors' rate less own salt files another bin", round(r1 / 100) * 100 == 2600,
       "%.1f t" % r1)
    ok("RIVAL the earlier moves' average gain files another bin (1,500 to 1,600)",
       round(r2 / 100) * 100 in (1500, 1600), "%.1f t" % r2)
    ok("RIVAL both rivals at least 250 t from the golden", min(abs(r1 - gold), abs(gold - r2)) >= 250, "")

    # ---------------------------------------------------------------------------------------- Ask A
    yards = ["KLV", "MPF", "RDG", "LKS"]
    names = {"Kelvin Road Yard": "KLV", "Mapleford Yard": "MPF", "Ridge Yard": "RDG", "Lakeshore Yard": "LKS"}
    stock0, sday, cap_end = {}, {}, {}
    for nm, y in names.items():
        m = re.search(re.escape(nm) + r" (\d{1,2} September 2026) [\d,]+ [\d.]+ ([\d,]+\.\d)", surv)
        sday[y] = dt.datetime.strptime(m.group(1), "%d %B %Y").date()
        stock0[y] = float(m.group(2).replace(",", ""))
        m2 = re.search(r"TG-26-\d{3} " + re.escape(nm) + r" (2026-09-\d{2}) (\d{2}:\d{2}) (\d{2}:\d{2})", surv)
        end_utc = dt.datetime.strptime(m2.group(1) + " " + m2.group(3), "%Y-%m-%d %H:%M").replace(tzinfo=dt.timezone.utc)
        cap_end[y] = end_utc.astimezone(TZ).replace(tzinfo=None)
    ok("PDF survey: four domes' stock, survey day and capture end", len(stock0) == 4 and
       all(cap_end[y].date() == sday[y] for y in yards), str({y: str(cap_end[y])[5:16] for y in yards}))
    wb = openpyxl.load_workbook(os.path.join(TARGET, "scale_house_log_2026-08-24_to_2026-09-24.xlsx"), read_only=True)
    recs = []
    for y in yards:
        rows = list(wb[y].iter_rows(values_only=True))
        h = list(rows[0])
        for r in rows[1:]:
            rec = dict(zip(h, r))
            rec["sheet"] = y
            rec["when"] = dt.datetime.combine(rec["Date"].date(), dt.datetime.strptime(rec["Time in"], "%H:%M").time())
            recs.append(rec)
    wb.close()
    # on its own path: a B-train carries one docket per trailer, so the same truck weighed in at two scale houses on
    # one day under the two dockets of one load is one load, which tipped at the later scale house
    sup_in = sorted((x for x in recs if x["Direction"] == "IN" and str(x["Carrier"]).startswith("Northgrit")),
                    key=lambda x: x["when"])
    early = set()
    for i, a_ in enumerate(sup_in):
        for b_ in sup_in[i + 1:]:
            if (b_["Vehicle"], b_["Date"].date()) == (a_["Vehicle"], a_["Date"].date()) and b_["sheet"] != a_["sheet"] \
                    and abs(int(b_["Supplier ticket"][2:]) - int(a_["Supplier ticket"][2:])) == 1:
                early.add(id(a_))
    st = dict(stock0)
    for rec in recs:
        if id(rec) in early:
            continue
        rem = (rec["Remarks"] or "").upper()
        dome = "LKS" if ("LKS" in rem or "LAKESHORE" in rem) else ("MPF" if ("MPF" in rem or "MAPLEFORD" in rem)
                                                                   else rec["sheet"])
        if rec["when"] <= cap_end[dome]:
            continue
        st[dome] += (1 if rec["Direction"] == "IN" else -1) * rec["Net kg"] / 1000.0
    ok("ASK A the scale houses' earlier weighings of B-trains that tipped at another yard found by their dockets",
       len(early) == 5, "%d weighings" % len(early))
    # and on the consistency check a grader runs instead: a supplier truck is never in at a City scale again sooner
    # than an hour and a half, a drive back to the terminal and out again, except where one load was weighed twice
    by = {}
    for x in sup_in:
        by.setdefault((x["Vehicle"], x["Date"].date()), []).append(x)
    gaps = [((b_["when"] - a_["when"]).total_seconds() / 60.0, a_, b_) for v in by.values() for a_, b_ in zip(v, v[1:])]
    quick = {id(a_) for m_, a_, b_ in gaps if m_ < 90 and a_["sheet"] != b_["sheet"] and a_["Gross kg"] == b_["Gross kg"]}
    ok("ASK A the same five found by the trucks' times between scale houses (one gross, under 90 minutes apart; every "
       "other same-truck day 2 hours or more)", quick == early and
       min(m_ for m_, a_, b_ in gaps if id(a_) not in early) >= 120,
       "shortest other %.0f min" % min(m_ for m_, a_, b_ in gaps if id(a_) not in early))
    for r in rows_csv("northgrit_dispatch_2026-10_issued_2026-09-21.csv"):
        if r["status"] == "BOOKED" and r["delivery_date"][:7] == "2026-10":
            st[r["ship_to"][3:]] += float(r["planned_net_t"])
    capd = {r["yard"]: float(r["rated_capacity_t"]) for r in rows_csv("salt_dome_register.csv")}
    good = all(round(st[y], 1) == figs["stock_%s" % y] and round(capd[y] - st[y], 1) == figs["room_%s" % y]
               for y in yards)
    ok("ASK A every dome's stock and room recompute", good, str({y: round(st[y], 1) for y in yards}))
    ok("ASK A every figure clear of its rounding edge (from the recomputed values)",
       min(min(edge(st[y], 0.1), edge(capd[y] - st[y], 0.1)) for y in yards) >= 0.025,
       "%.4f" % min(min(edge(st[y], 0.1), edge(capd[y] - st[y], 0.1)) for y in yards))

    # ---------------------------------------------------------------------------------------- Ask B
    # the yard each circuit belonged to that winter: the platform carries a renamed circuit's current code, so the
    # revision log's ROUTE entries take a code back to the one it replaced before the change's effective date
    back = {}
    for r in rev:
        if r["field"] == "ROUTE" and "-A" in r["old_value"] and "-A" in r["new_value"]:
            back[r["new_value"]] = (r["old_value"], dt.date.fromisoformat(r["effective_date"]))
    w3 = ("2023-24", "2024-25", "2025-26")
    arts = []
    for i in range(n_pass):
        rc = cols["route_code"][i]
        if "-A" not in rc or cols["material_code"][i] != "RS":
            continue
        d = loc[i].date()
        w = season(d)
        if w not in w3:
            continue
        code = back[rc][0] if rc in back and d < back[rc][1] else rc
        arts.append((cols["unit_id"][i], w, code[:3], cols["run_ref"][i], cols["applied_t"][i], cols["prewet_l"][i]))
    # the older controllers, on their own path: a unit-winter whose rows carry fewer litres a tonne than the plan's
    # pre-wet rate (the salt and brine reported as one mass)
    lit = {}
    for u, w, y, ref, t_, l_ in arts:
        if t_ >= 0.1:
            lit.setdefault((u, w), []).append(l_ / t_)
    older = {k for k, v in lit.items() if float(np.median(v)) < prewet - 0.9}
    art, seen, art_l = {}, set(), {}
    for u, w, y, ref, t_, l_ in arts:
        k = (y, w)
        if (u, w) in older:
            if ref in seen:
                continue
            seen.add(ref)
            v, vl = t_ - l_ * brine / 1000.0, l_ / prewet
        else:
            v, vl = t_, t_
        art[k] = art.get(k, 0.0) + v
        art_l[k] = art_l.get(k, 0.0) + vl
    ok("ASK B older controllers found unit-winter by unit-winter on their litres a tonne (the plan's rate over salt "
       "and brine)", len(older) > 0 and all(abs(float(np.median(lit[k])) - prewet / (1 + prewet * brine / 1000)) < 0.1
                                            for k in older), "%d unit-winters" % len(older))
    # the maintenance record joined by unit and date: each controller the work orders replaced was an older one in
    # every winter before the replacement and in none after
    wo = [r for r in rows_csv("fleet_work_orders_2023-07_to_2026-09.csv") if "M410 to M720" in r["description"]]
    agree = all(((r["unit_id"], w) in older) == (dt.date(int(w[:4]), 11, 1) < dt.date.fromisoformat(r["closed"]))
                for r in wo for w in w3 if (r["unit_id"], w) in lit)
    ok("ASK B the work orders' controller replacements agree with the litres a tonne, unit by unit and winter by winter",
       len(wo) == 4 and agree, "%d replacements" % len(wo))
    wb = openpyxl.load_workbook(os.path.join(TARGET, "kettlewood_material_returns_2023-24_to_2025-26.xlsx"),
                                read_only=True)
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        h = list(rows[2])
        for r in rows[3:]:
            rec = dict(zip(h, r))
            if rec["Month"] is None:
                continue
            w = season(dt.date.fromisoformat(rec["Month"] + "-01"))
            if w == ws.title:
                for dd in (art, art_l):
                    dd[(rec["City circuit"][:3], w)] = dd.get((rec["City circuit"][:3], w), 0.0) + \
                        float(rec["Rock salt (t)"])
    wb.close()
    good = all(int(round(art[(y, w)] / 50.0) * 50) == figs["art_%s_%s" % (y, w)] for y in yards for w in w3)
    ok("ASK B every yard-winter arterial figure recomputes to the nearest 50 t", good, "")
    ok("ASK B every figure a quarter-bin or more from its 50 t rounding edge (from the recomputed values)",
       min(edge(v, 50.0) for v in art.values()) >= 12.5, "%.1f t" % min(edge(v, 50.0) for v in art.values()))
    ok("ASK B the older controllers' salt read back from their litres at the plan's rate files the same 50 t figures",
       all(int(round(art_l[k] / 50.0) * 50) == int(round(art[k] / 50.0) * 50) for k in art),
       "largest gap %.2f t" % max(abs(art_l[k] - art[k]) for k in art))

    # ---------------------------------------------------------------------------------------- the other formats
    n_csv = sum(1 for f in os.listdir(TARGET) if f.endswith(".csv"))
    for f in sorted(os.listdir(TARGET)):
        if f.endswith(".csv"):
            rows_csv(f)
    ok("CSV every shipped CSV parses with a header", n_csv >= 10, "%d files" % n_csv)
    qb = openpyxl.load_workbook(os.path.join(TARGET, "qrt_frequent_network_2026_streets.xlsx"), read_only=True)
    streets = {r[1] for r in list(qb.worksheets[0].iter_rows(values_only=True))[4:] if r[1]}
    qb.close()
    ok("XLSX transit street list covers the 216 segments' streets", streets == {reg[s]["street"] for s in new},
       "%d streets" % len(streets))

    # ---------------------------------------------------------------------------------------- goldens and write-up
    gwb = openpyxl.load_workbook(os.path.join(TASK, "golden", "salt_commitment_2026-27.xlsx"), read_only=True)
    vals = [c for r in gwb["Commitment"].iter_rows(values_only=True) for c in r if c is not None]
    domes = [r for r in gwb["Domes 1 Nov 2026"].iter_rows(values_only=True)][1:5]
    arts = [r for r in gwb["Arterial salt"].iter_rows(values_only=True)][1:5]
    gwb.close()
    ok("GOLDEN workbook carries the call, both components and every ask figure as figures.json",
       call in vals and figs["own_p2_t"] in vals and figs["gain_t_per_lkm"] in vals and
       all(r[1] == figs["stock_%s" % y] and r[2] == figs["room_%s" % y] for r, y in zip(domes, yards)) and
       all(list(r[1:4]) == [figs["art_%s_%s" % (y, w)] for w in ("2023-24", "2024-25", "2025-26")]
           for r, y in zip(arts, yards)), "")
    from docx import Document
    memo = " ".join(p.text for p in Document(os.path.join(TASK, "golden",
                                                          "seasonal_salt_commitment_council_report.docx")).paragraphs)
    ok("GOLDEN council report read back: the commitment, its components, one chart",
       "{:,} t of rock salt".format(call) in memo and "about %d t a lane-km" % figs["gain_t_per_lkm"] in memo and
       "about {:,} t".format(figs["own_p2_t"]) in memo and
       len(Document(os.path.join(TASK, "golden", "seasonal_salt_commitment_council_report.docx")).inline_shapes) == 1,
       "")
    sub = open(os.path.join(TASK, "submission.md"), encoding="utf-8").read()
    want = ["{:,}".format(call), "about {:,} t, under {:d} t per lane-km".format(figs["own_p2_t"],
                                                                                figs["own_p2_t_per_lkm_under"]),
            "about %d t per lane-km" % figs["gain_t_per_lkm"]] + \
        ["{:,.1f}".format(figs[k]) for k in figs if k.startswith(("stock_", "room_"))] + \
        ["{:,} t".format(figs[k]) for k in figs if k.startswith("art_")]
    ok("WRITEUP every graded figure appears in submission.md at its stated rounding", all(x in sub for x in want),
       str([x for x in want if x not in sub][:4]))

    n_fail = OUT.count(False)
    print("\nverify: %d checks, %d failed" % (len(OUT), n_fail))
    sys.exit(1 if n_fail else 0)


if __name__ == "__main__":
    main()
