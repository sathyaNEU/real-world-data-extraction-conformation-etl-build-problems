"""task165 checks: the generator's assertions on the built pack (target/) and on the clean truth (world.py, records.py,
yards.py).

The three-rules block (R1 to R3c), the ladder and its rivals, the out-of-sample tests, the golden's family and its
standard error, the reflex sweep, the twin pair, every semantic fork row (SF1 to SF14), both asks with their
devices and stops, the cracker's measured pair, the separation line, the input gates, the realism sweep, the dates
and the distractors' deletion test."""
from __future__ import annotations

import datetime as dt
import hashlib
import json
import os
import re
import shutil
import sys
import tempfile

import numpy as np
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import calc as C        # noqa: E402
import common as K      # noqa: E402
import cracker as CR    # noqa: E402
import docs as D        # noqa: E402
import pack as P        # noqa: E402
import pack2 as P2      # noqa: E402
import records as R     # noqa: E402
import world as Wd      # noqa: E402
import yards as Y       # noqa: E402

TASK = os.path.dirname(HERE)
TARGET = os.path.join(TASK, "target")
RESULTS = []

GOLD = 2200            # the seasonal tonnage, nearest 100 t
DEFAULT = 2600         # rung 1, the default path
POOLED = (1500, 1600)  # rung 2, the earlier moves' average gain: 1,500 to 1,600 (pooled 1,500, its DiD forms 1,600)
N_MOVES = 377          # collectors moved Priority 2 to Priority 1 on 1 November 2022 to 2025
OWN_T, OWN_UNDER = 700, 2   # component 1: own Priority 2 salt over the five winters, t in all (nearest 100) and a
                            # lane-km bound, in a winter at the planning severity
GAIN = 5               # component 2: the gain, t a lane-km in that winter (nearest tonne)
SEV = C.PLAN_SEV
ASK_EDGE = 0.0295      # every ask figure at least this far (t) from its one-decimal rounding edge

# the shipped pack (target.zip d87ff77a) the ask repair is held against: the files the call reads, byte for byte,
# and the call's rows in the two files the repair changed (sha256 of the rows as pandas writes them to CSV)
SHIPPED_CALL_FILES = {
    "avl_legacy_segment_shift_2021-11-01_2023-03-31.csv": "2a4e340a41ce7044da58ce4f770d8f7efa8fd097ec311d05454e87144142c792",
    "winter_route_register_2026-09-14.csv": "224468a25e170f881667e3483c7768ca0fe926435ea4c9e8a250f696c1d8b695",
    "route_audit_revision_log.csv": "34f88a2ec357f3674d7d6c8caa4720e8d193ec556681f6dabf2bceee435fd65e",
    "storm_log_2021-22_to_2025-26.xlsx": "ad8e779838e7741e8f90f591544eab5264c16ddc826a1a59950268dc25c345e7",
    "winter_maintenance_plan_2026-27.pdf": "13d1a904db688ddfe10ea40261e8079a1f40737ebb3c752c02fb292054193d2a",
    "road_salt_procurement_procedure_RO-PR-07.pdf": "1f2dac155a5eb245d7e6c0b9c3d786f2a5e3831f508446c7b7a34aef9fcf231a",
    "gis_segment_crosswalk_2023.csv": "959206e5c405cd26134fe070e2392ebed1abce0acf4fd7b6bd9db43d9e1fa3ac",
    "qrt_frequent_network_2026_streets.xlsx": "1f5f964535cf919d07c695579752cc1888c8dbc96f6d5d1902bb371a1810242d",
}
SHIPPED_COLLECTOR_ROWS = "f0dde0c60e93bc88bec33dd93823b2890174992fe867868a7c8c9b84c5c6c8f0"
SHIPPED_ARTERIAL_ROWS = "42fbbf1c59d6d0c8c8f0a1863b0dedd124c7c5d64a7075edfd54b51b02f167f0"   # less route_code
SHIPPED_DICT_CALL_ROWS = "9a1242a1644c5d2a7e79572b68dd1de0de0cd351ab1ecc3956ce45dceb675e0f"


def check(label, ok, detail=""):
    RESULTS.append((label, bool(ok), detail))
    print("%s %-74s %s" % ("PASS" if ok else "FAIL", label[:74], detail))


def call(v):
    return int(round(v / 100.0) * 100)


def r1(v):
    return round(v + 1e-9, 1)


def linked_copy(skip):
    """A folder of links to every shipped file but the ones named."""
    tmp = tempfile.mkdtemp(prefix="chk165_", dir=TASK)
    for f in os.listdir(TARGET):
        if f not in skip:
            os.symlink(os.path.join(TARGET, f), os.path.join(tmp, f))
    return tmp


def shipped_text():
    """Every shipped file's text: CSV as read, workbook strings, PDF text through pypdf."""
    import openpyxl
    out = {}
    for f in sorted(os.listdir(TARGET)):
        p = os.path.join(TARGET, f)
        if f.endswith(".csv"):
            out[f] = open(p, encoding="utf-8").read()
        elif f.endswith(".xlsx"):
            wb = openpyxl.load_workbook(p, read_only=True)
            parts = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    parts.extend(str(v) for v in row if v is not None)
            wb.close()
            out[f] = "\n".join(parts)
        elif f.endswith(".pdf"):
            out[f] = C.pdf_text(p)
    return out


def main():
    P_ = C.Pack(TARGET)
    g = C.golden(P_)
    nn = P_.new_rates()
    L = float(nn["lane_km"].sum())
    x0 = float((nn["p2"] * nn["lane_km"]).sum() / L)
    d = P_.prepost()
    own, gain = x0 * SEV, g / L

    # ------------------------------------------------------------------------------------------ the golden
    check("GOLD the seasonal tonnage", call(g) == GOLD, "%.1f t -> %d" % (g, call(g)))
    check("GOLD the raw figure a quarter-bin or more from its rounding edge", K.bin_edge_distance(g, 100) >= 25,
          "%.1f t" % K.bin_edge_distance(g, 100))
    check("GOLD the population: 216 segments, 424.0 lane-km, moved on 30 August 2026",
          len(nn) == 216 and abs(L - 424.0) < 1e-6, "%d, %.2f" % (len(nn), L))
    check("GOLD the earlier moves: %d collectors on the four 1 November dates" % N_MOVES,
          len(d) == N_MOVES and sorted({str(e) for e in d["eff"]}) == ["2022-11-01", "2023-11-01", "2024-11-01",
                                                                   "2025-11-01"], "%d" % len(d))
    check("GOLD component: own Priority 2 salt over the five winters about 700 t, under 2 t a lane-km, clear of "
          "its edge", call(own * L) == OWN_T and own < OWN_UNDER and K.bin_edge_distance(own * L, 100) >= 15,
          "%.1f t, %.3f t a lane-km" % (own * L, own))
    ids_new = list(nn["segment_id"])
    tot_new = {w: sum(float(P_.salt.get((s_, w), 0.0)) for s_ in ids_new) for w in K.WINTERS}
    tel = ["2023-24", "2024-25", "2025-26"]
    bases = {"pooled (the golden)": own,
             "mean of the five winters' rates": float(np.mean([tot_new[w] / (L * P_.wsi[w]) for w in K.WINTERS])) * SEV,
             "mean of the 216 streets' rates": float(nn["p2"].mean()) * SEV,
             "lane-km weighted mean of the streets' winter means": float(np.average(
                 [np.mean([float(P_.salt.get((s_, w), 0.0)) / (lk_ * P_.wsi[w]) for w in K.WINTERS])
                  for s_, lk_ in zip(nn["segment_id"], nn["lane_km"])], weights=nn["lane_km"])) * SEV,
             "the three telematics winters": sum(tot_new[w] for w in tel) / (L * sum(P_.wsi[w] for w in tel)) * SEV}
    check("SF13 component 1 files the same figure on every basis of the five winters (t in all and a lane-km)",
          all(call(v * L) == OWN_T and v < OWN_UNDER for v in bases.values()),
          str({k: round(v * L, 1) for k, v in bases.items()}))
    check("GOLD component 2: the gain about 5 t a lane-km, clear of its edge",
          int(round(gain)) == GAIN and K.bin_edge_distance(gain, 1.0) >= 0.2, "%.4f" % gain)
    figs = json.load(open(os.path.join(HERE, "figures.json")))
    check("GOLD figures.json carries the call and both components",
          figs["commitment_t"] == GOLD and figs["own_p2_t"] == OWN_T and figs["own_p2_t_per_lkm_under"] == OWN_UNDER
          and figs["gain_t_per_lkm"] == GAIN and figs["lane_km"] == 424.0 and figs["segments"] == 216 and
          figs["earlier_moves"] == N_MOVES, "")

    # ------------------------------------------------------------------------------------------ R1
    import pyarrow.parquet as pq
    import openpyxl
    heads = []
    for f in os.listdir(TARGET):
        p = os.path.join(TARGET, f)
        if f.endswith(".csv"):
            heads += [(f, h) for h in open(p, encoding="utf-8").readline().strip().split(",")]
        elif f.endswith(".parquet"):
            heads += [(f, h) for h in pq.ParquetFile(p).schema_arrow.names]
        elif f.endswith(".xlsx"):
            wb = openpyxl.load_workbook(p, read_only=True)
            for ws in wb.worksheets:
                for row in ws.iter_rows(min_row=1, max_row=6, values_only=True):
                    heads += [(f, str(v)) for v in row if isinstance(v, str) and len(v) < 40]
            wb.close()
    bad = [h for h in heads if re.search(r"coverage|covered|hazard|\bgain|untreat|spot|share_salted|per.lane|\brate\b|t_per",
                                         h[1], re.I)]
    check("R1 no shipped column holds a coverage, a hazard length, a gain or a rate", not bad, str(bad[:4]))
    ps = pd.read_parquet(os.path.join(TARGET, P.F_PASSES))
    check("R1 the platform ships one row per spreader pass over one segment, 25,000 rows or more",
          len(ps) >= 25000 and not ps.duplicated(["unit_id", "segment_id", "entry_utc"]).any(), "%d rows" % len(ps))
    leg = pd.read_csv(os.path.join(TARGET, P.F_LEGACY), dtype={"SEG_ID": str})
    check("R1 the legacy AVL ships segment-by-shift material, 25,000 rows or more", len(leg) >= 25000,
          "%d rows" % len(leg))
    full = all(float(P_.salt.get((r["segment_id"], w), 0.0)) > 0 for _, r in P_.earlier_moves().iterrows()
               for w in K.WINTERS)
    check("R1 the decisive fit reads each earlier move's own salt in every winter, before and after", full, "")

    # ------------------------------------------------------------------------------------------ R3 on the clean truth
    W = Wd.World()
    S = R.Sim(W)
    P2.settle_arterial_cells(W, S)     # the arterial cells settled as the build settles them (collectors untouched)
    wsi_t = {w: sum(e["points"] for e in W.events[w]) for w in K.WINTERS}
    segs = {s["segment_id"]: s for s in W.segments}

    def trate(sid, ws):
        return sum(S.truth.get((sid, w), 0.0) for w in ws) / (segs[sid]["lane_km"] * sum(wsi_t[w] for w in ws))

    new = [s for s in W.segments if s["pool"] == "NEW"]
    mv = [s for s in W.segments if s["pool"] == "MOVE"]
    p1in = [s for s in W.segments if s["pool"] == "P1IN"]
    Lt = sum(s["lane_km"] for s in new)
    x0t = sum(trate(s["segment_id"], K.WINTERS) * s["lane_km"] for s in new) / Lt
    tot = den = 0.0
    for s in p1in + mv:
        for w in K.WINTERS:
            if W.priority(s, w) == 1:
                tot += S.truth.get((s["segment_id"], w), 0.0)
                den += s["lane_km"] * wsi_t[w]
    r3a = (tot / den - x0t) * Lt * SEV
    rows = []
    for s in mv:
        pre = [w for w in K.WINTERS if W.priority(s, w) == 2]
        post = [w for w in K.WINTERS if W.priority(s, w) == 1]
        rows.append((s["lane_km"], trate(s["segment_id"], pre), trate(s["segment_id"], post), s["h"]))
    Lk, pre_t, post_t, h_t = map(np.array, zip(*rows))
    bt = C.wls(pre_t, post_t - pre_t, Lk)
    r3b = (bt[0] + bt[1] * x0t) * Lt * SEV
    pooled_t = ((post_t - pre_t) * Lk).sum() / Lk.sum() * Lt * SEV
    check("R3a the easy route on the clean truth files the default path", call(r3a) == DEFAULT, "%.1f t" % r3a)
    check("R3b the decisive step alone on the clean truth files the golden", call(r3b) == GOLD, "%.1f t" % r3b)
    check("R3b the event study alone (pooled gain) on the clean truth stays wrong", call(pooled_t) in POOLED,
          "%.1f t" % pooled_t)
    expected = sum(W.gain_truth(s) * s["lane_km"] * s["m"] for s in new)
    check("R3b the world's own expected gain within two standard errors of the golden", abs(expected - g) < 2 * 15.0,
          "%.1f t" % expected)
    check("R3b the clean truth and the pack agree on the call's rows (no mess on them)",
          abs(r3a - C.r1_default(P_)) < 1.0 and abs(r3b - g) < 1.0, "%.2f, %.2f" % (r3a - C.r1_default(P_), r3b - g))
    txt = shipped_text()
    docs_txt = " ".join(txt[f] for f in (P.F_PLAN, P.F_PROC, P.F_SURVEY)).replace("\n", " ")
    check("R3c no document says how a street's salt responds to its service level or its coverage",
          not re.search(r"already salted|share of (each|the|a) (street|section|segment|collector)|coverage|untreated|"
                        r"depends on|the gain|in proportion to", docs_txt, re.I), "")
    check("R3c no document states a seasonal tonnage, a gain or a forecast for the change",
          not re.search(r"\b\d{1,2},\d{3}\s*t\b|\b\d+(\.\d+)?\s*t a lane|tonnes? (a|per) lane", docs_txt), "")
    check("R3c no memo, draft, briefing or correspondence ships",
          not any(f.endswith((".docx", ".eml", ".msg", ".txt", ".md")) for f in os.listdir(TARGET)), "")
    sv = P_.salt.unstack().fillna(0.0).round(6)
    check("R3c no segment's five winters of salt repeat another's (no copied block)",
          not sv.duplicated().any(), "%d segments" % len(sv))

    # ------------------------------------------------------------------------------------------ R2 mess neutrality
    Pu = C.Pack(TARGET, utc_as_local=True)
    gu = C.golden(Pu)
    check("R2 the platform's UTC times read as local: same call and components",
          call(gu) == GOLD and int(round(gu / L)) == GAIN, "%.1f t" % gu)
    Px = C.Pack(TARGET, xwalk=False)
    gx = C.golden(Px)
    check("R2 the legacy ids joined without the crosswalk: same call", call(gx) == GOLD, "%.1f t" % gx)
    gc = C.centreline(P_)
    xl = pd.ExcelFile(os.path.join(TARGET, P.F_STORM))
    seasons_ok = True
    for w in K.NEW_WINTERS:
        df = xl.parse(w, header=None)
        lab = df.iloc[:, 0].astype(str).str.strip()
        row = df[lab.str.lower().str.startswith("season wsi")]
        if len(row) != 1:
            seasons_ok = False
            continue
        vals = [v for v in row.iloc[0].tolist()[1:] if isinstance(v, (int, float)) and not pd.isna(v)]
        seasons_ok &= bool(vals) and abs(float(vals[-1]) - P_.wsi[w]) < 1e-6
    check("R2 the storm log's Season WSI rows equal the sums of their events (either read)", seasons_ok, "")
    legc = leg.merge(P_.reg[["segment_id", "func_class"]], left_on="SEG_ID", right_on="segment_id", how="left")
    legc = legc[legc["func_class"] == "COLLECTOR"]
    lr = legc["KG"].sum() / 1000.0
    check("R2 the legacy kilograms read as tonnes is loud (a thousandfold the platform's)",
          (legc["KG"].median() / 1000.0) < 1.0 < legc["KG"].median(), "median %.0f kg" % legc["KG"].median())
    edge_rows = legc[legc["SHIFT_DATE"].str[5:].isin(["10-31", "04-01"]) |
                     ((legc["SHIFT_DATE"].str[5:] == "03-31") & (legc["SHIFT"] == "NIGHT"))]
    check("R2 no collector shift sits on a season boundary (night shifts dated by their start: same call)",
          len(edge_rows) == 0, "%d rows, %.0f t legacy collector salt" % (len(edge_rows), lr))

    # ------------------------------------------------------------------------------------------ the ladder
    wmean = float(np.mean(list(P_.wsi.values())))
    last = P_.wsi["2025-26"]
    lad = {
        "R0 the Priority 1 network gross, last winter": (C.r0_network(P_, sev=last, winters=["2025-26"]), 5000),
        "R0b the Priority 1 network gross, planning severity": (C.r0_network(P_), 5400),
        "R1 default: Priority 1 collector rate less own salt": (C.r1_default(P_), DEFAULT),
        "R1b the same over the earlier moves' Priority 1 winters": (C.r1_default(P_, True), DEFAULT),
        "R1c the default at last winter's points (no planning severity)": (C.r1_default(P_) * last / SEV, 2500),
        "R1d the default at the five winters' mean points": (C.r1_default(P_) * wmean / SEV, 2400),
        "R2 the earlier moves' pooled gain (1,500 to 1,600)": (C.r2_pooled(P_), POOLED),
        "R2b the transit-driven earlier moves' gain": (C.r2_pooled(P_, "QRT Route"), POOLED),
        "R2c the pooled gain stratified by yard": (C.reflex_stratified(P_, "yard"), POOLED),
        "R2c the pooled gain stratified by lanes": (C.reflex_stratified(P_, "lanes"), POOLED),
        "R2c the pooled gain stratified by length band": (C.reflex_stratified(P_, "lenband"), POOLED),
        "R2d the golden's gain at last winter's points": (g * last / SEV, 2100),
        "R2e the golden's gain at the five winters' mean": (g * wmean / SEV, 2000),
        "R2f the earlier moves' ratio on own salt": (C.reflex_ratio(P_), 600),
        "R2g the pooled gain trended over move years": (C.reflex_trend(P_), 1800),
        "R2h the pooled gain against never-moved collectors": (C.reflex_did(P_), POOLED),
    }
    for name, (v, want) in lad.items():
        want = want if isinstance(want, tuple) else (want,)
        check("LADDER %s" % name, call(v) in want and abs(v - g) >= 100 and call(v) != GOLD,
              "%.1f t -> %d" % (v, call(v)))
    for name, v in (("the default and the pooled gain", (C.r1_default(P_) + C.r2_pooled(P_)) / 2),
                    ("the moves' own level and the pooled gain", (C.r1_default(P_, True) + C.r2_pooled(P_)) / 2),
                    ("the default and the transit moves' gain", (C.r1_default(P_) + C.r2_pooled(P_, "QRT Route")) / 2)):
        check("HEDGE the midpoint of %s misses the golden's bin by 50 t or more" % name,
              v < 2150 - 50 or v >= 2250 + 50, "%.1f t" % v)
    pc_ = P_.reg[(P_.reg["func_class"] == "COLLECTOR") & (P_.reg["priority"] == "1")]
    newset = set(nn["segment_id"])
    keep = [(s_, lk_) for s_, lk_ in zip(pc_["segment_id"], pc_["lane_km"]) if s_ not in newset]
    p1_last = sum(float(P_.salt.get((s_, "2025-26"), 0.0)) for s_, _ in keep) / sum(lk_ for _, lk_ in keep)
    own_last = tot_new["2025-26"] / L
    coin = (p1_last - own_last) * L
    check("HEDGE last winter's totals, netted, no planning severity, miss the golden's bin by 50 t or more",
          coin < 2150 - 50 or coin >= 2250 + 50, "%.1f t (%.2f less %.2f t a lane-km)" % (coin, p1_last, own_last))
    check("LADDER bracketed: the netting above the golden, every average gain below",
          C.r1_default(P_) > g > max(C.r2_pooled(P_), C.r2_pooled(P_, "QRT Route"), C.reflex_trend(P_)), "")

    # ------------------------------------------------------------------------------------------ out of sample
    ho = C.holdout(P_)
    check("OOS each move year held out: the conditioned gain closer than the pooled gain in every year",
          all(abs(v[1]) < abs(v[0]) for v in ho.values()), str({k: round(v[1], 3) for k, v in ho.items()}))
    check("OOS each move year held out: the pooled gain off by over 10 per cent in a year",
          max(abs(v[0]) for v in ho.values()) > 0.10, str({k: round(v[0], 3) for k, v in ho.items()}))
    cv = {f: C.cv_wsse(P_, f) for f in ("line", "loglog", "pool")}
    check("OOS ten-fold: the line's error under the power law's and a fifth of the average's",
          cv["line"] < cv["loglog"] and cv["line"] < cv["pool"] / 5, str({k: round(v, 5) for k, v in cv.items()}))
    lb = {f: C.low_band(P_, f) for f in ("line", "loglog", "pool")}
    check("OOS the new streets' band predicted from the rest: the line within 3 per cent",
          abs(lb["line"][0]) < 0.03 and lb["line"][1] >= 60, "%.3f over %d streets" % lb["line"])
    check("OOS the same: the power law and the average gain miss by over 10 per cent",
          lb["loglog"][0] < -0.10 and lb["pool"][0] < -0.10, "%.3f, %.3f" % (lb["loglog"][0], lb["pool"][0]))
    ll = C.loglog(P_)
    check("OOS the power law, weaker out of sample, still files the golden's bin", call(ll) == GOLD and
          K.bin_edge_distance(ll, 100) >= 25, "%.1f t" % ll)

    # ------------------------------------------------------------------------------------------ the golden's family
    fam = {
        "weighted by lane-km (golden)": g,
        "unweighted": C.golden(P_, weight=False),
        "summed street by street": C.golden(P_, per_segment=True),
        "the Priority 1 level regressed, own salt taken off": C.golden(P_, level=True),
        "each winter against the never-moved collectors": C.golden_did(P_),
        "each winter's points netted by the never-moved collectors' swing": C.golden_netted(P_),
        "netted, the last winter before each move only": C.golden_netted(P_, pre_last=True),
        "netted, the first winter after each move only": C.golden_netted(P_, post_first=True),
        "netted, both one-winter windows": C.golden_netted(P_, pre_last=True, post_first=True),
        "each street's rate the mean of its winters": C.winter_mean(P_),
        "quadratic in Priority 2 salt": C.quadratic(P_),
        "outliers over three deviations dropped": C.trimmed(P_),
        "10 nearest earlier moves": C.knn(P_, 10),
        "20 nearest earlier moves": C.knn(P_, 20),
        "40 nearest earlier moves": C.knn(P_, 40),
        "four bands of Priority 2 salt": C.qbins(P_, 4),
        "five bands": C.qbins(P_, 5),
        "eight bands": C.qbins(P_, 8),
    }
    for name, v in fam.items():
        check("FAMILY %s files the golden" % name, call(v) == GOLD and K.bin_edge_distance(v, 100) >= 24,
              "%.1f t" % v)
    check("SF12 every family member's gain is about 5 t a lane-km (component 2 at the precision they share)",
          all(int(round(v / L)) == GAIN and K.bin_edge_distance(v / L, 1.0) >= 0.2 for v in fam.values()),
          "%.3f to %.3f" % (min(fam.values()) / L, max(fam.values()) / L))
    ix = C.control_index(P_)
    raw1 = {"the last winter before each move only": C.golden(P_, pre_last=True),
            "the first winter after each move only": C.golden(P_, post_first=True),
            "both one-winter windows": C.golden(P_, pre_last=True, post_first=True)}
    check("SF14 each winter's swing (over 10 per cent on the never-moved collectors): one-winter windows read "
          "netted, and un-netted they stay in the bin", max(ix.values()) / min(ix.values()) - 1 > 0.10 and
          all(call(v) == GOLD for v in raw1.values()),
          "%s; un-netted %s" % ({w: round(v, 3) for w, v in ix.items()}, {k: round(v, 1) for k, v in raw1.items()}))
    spread_new = float(np.std(nn["p2"].values * SEV))
    crude_in = {"fixed 1.5 t bands": C.fixed_bands(P_, 1.5), "fixed 1 t bands": C.fixed_bands(P_, 1.0),
                "a window of 0.2 t either side": C.window(P_, 0.2)[0], "a window of 1 t": C.window(P_, 1.0)[0],
                "a kernel of 0.62 t": C.kernel(P_, 0.62)}
    check("CRUDE settings of the conditioning within the new streets' own spread file the golden",
          all(call(v) == GOLD for v in crude_in.values()), str({k: round(v, 1) for k, v in crude_in.items()}))
    coarse = {"equal-count bands, 2 to 6": [C.qbins(P_, q) for q in (2, 3, 4, 5, 6)],
              "kernels of 1.23, 0.9, 0.62 and 0.45 t": [C.kernel(P_, b_) for b_ in (1.23, 0.9, 0.62, 0.45)]}
    check("CRUDE coarse settings close on the golden as they are refined (two or three bands, wide kernels slip)",
          all(all(abs(v[i + 1] - g) <= abs(v[i] - g) + 1.0 for i in range(len(v) - 1)) and
              call(v[-1]) == GOLD and call(v[-2]) == GOLD for v in coarse.values()),
          str({k: [round(x_, 1) for x_ in v] for k, v in coarse.items()}))
    thin = {k_: C.nearest_to_mean(P_, k_) for k_ in (10, 20, 30, 40, 60)}
    wins = {h_: C.window(P_, h_) for h_ in (0.1, 0.2, 0.5)}
    check("CRUDE thin settings settle in the golden's bin once 30 or more streets carry them",
          all(call(v) == GOLD for k_, v in thin.items() if k_ >= 30) and
          all(call(v) == GOLD for v, n_ in wins.values() if n_ >= 30),
          "nearest the mean %s; windows %s" % ({k_: round(v, 1) for k_, v in thin.items()},
                                               {h_: (round(v, 1), n_) for h_, (v, n_) in wins.items()}))
    rm_ = C.range_match(P_)
    check("CRUDE the range trim's plain average (R2i) stays off the golden's bin, toward the pooled gain",
          call(rm_[0]) != GOLD and rm_[0] < g, "%.1f t over %d streets" % (rm_[0], rm_[1]))
    sd = C.bootstrap_sd(P_)
    check("SIGMA one bootstrap standard error over the %d streets under 16 t" % N_MOVES, sd < 16.0, "%.1f t" % sd)
    check("SIGMA the golden plus or minus two standard errors stays in its bin",
          K.bin_edge_distance(g, 100) > 2 * sd, "%.1f against %.1f" % (K.bin_edge_distance(g, 100), 2 * sd))
    check("SIGMA every family member sits more than the golden's own standard error inside the bin",
          all(K.bin_edge_distance(v, 100) > sd for v in fam.values()),
          "nearest %.1f t" % min(K.bin_edge_distance(v, 100) for v in fam.values()))

    # ------------------------------------------------------------------------------------------ orthogonality, twin
    for pool in ("MOVE", "NEW"):
        df = pd.DataFrame([{"h": s["h_haz"], "lanes": s["lanes"], "cl": s["centreline_km"], "yard": s["yard"]}
                           for s in W.segments if s["pool"] == pool])
        cy = max(abs(np.corrcoef(df["h"], (df["yard"] == y).astype(float))[0, 1]) for y in K.YARDS)
        cl_ = abs(np.corrcoef(df["h"], df["lanes"])[0, 1])
        cc = abs(np.corrcoef(df["h"], df["cl"])[0, 1])
        lim = 3.0 / np.sqrt(len(df))
        check("ORTHO %s hazard share against lanes, length and yard, inside three sampling errors (never shipped)"
              % pool, cl_ < lim and cc < lim and cy < 0.2, "%.3f, %.3f, %.3f (limit %.3f)" % (cl_, cc, cy, lim))
    p2c = P_.reg[(P_.reg["func_class"] == "COLLECTOR") & (P_.reg["priority"] == "2")]
    rr_ = [P_.rate(s_, K.WINTERS, lk_) for s_, lk_ in zip(p2c["segment_id"], p2c["lane_km"])]
    cinv = float(np.corrcoef(rr_, 1.0 / p2c["centreline_km"].values)[0, 1])
    check("REALISM Priority 2 salt per lane-km rises as segments shorten (intersection approaches)", cinv > 0.15,
          "corr with 1/length %.3f" % cinv)
    check("ORTHO the earlier moves' coverage twice the new streets'",
          np.average(h_t, weights=Lk) > 2 * np.average([s["h"] for s in new], weights=[s["lane_km"] for s in new]),
          "%.3f against %.3f" % (np.average(h_t, weights=Lk),
                                 np.average([s["h"] for s in new], weights=[s["lane_km"] for s in new])))
    em = P_.earlier_moves().set_index("segment_id")
    dd = d.set_index("segment_id")
    dd["gain"] = dd["post"] - dd["pre"]
    best = None
    for (st, yd, ln, rs, ef), grp in em.groupby(["street", "yard", "lanes", "reason", "eff"]):
        ids = list(grp.index)
        for i in range(len(ids)):
            for j in range(i + 1, len(ids)):
                a, b = ids[i], ids[j]
                if abs(em.loc[a, "centreline_km"] - em.loc[b, "centreline_km"]) > 0.06:
                    continue
                ga, gb = dd.loc[a, "gain"], dd.loc[b, "gain"]
                pa, pb = dd.loc[a, "pre"], dd.loc[b, "pre"]
                ratio = max(ga, gb) / min(ga, gb)
                dcl = abs(em.loc[a, "centreline_km"] - em.loc[b, "centreline_km"])
                if 2.0 <= ratio <= 4.0 and max(pa, pb) / min(pa, pb) >= 3.0 and \
                        (best is None or (dcl, abs(ratio - 2.5)) < best[8:]):
                    best = (ratio, a, b, st, ga * SEV, gb * SEV, pa * SEV, pb * SEV, dcl, abs(ratio - 2.5))
    check("TWIN two earlier moves alike on every visible field, gains two to one or more apart", best is not None,
          "%s/%s %s: gain %.2f vs %.2f, own %.2f vs %.2f t a lane-km" % best[1:8] if best else "")
    check("TWIN the pair's gains run against their Priority 2 salt",
          best is not None and (best[4] - best[5]) * (best[6] - best[7]) < 0, "")

    # ------------------------------------------------------------------------------------------ semantic forks
    proc = txt[P.F_PROC].replace("\n", " ")
    plan = txt[P.F_PLAN].replace("\n", " ")
    rev = P_.rev
    late = rev[(rev["field"] == "PRIORITY") & (rev["effective_date"] > "2026-03-31")]
    check("SF1 the change since the last winter closed is the redesign alone, only its salt above the rollover",
          "approved since the last winter closed" in proc and len(late) == 216 and
          set(late["effective_date"]) == {"2026-08-30"}, "%d entries" % len(late))
    gross = g + x0 * L * SEV
    check("SF2 over what they needed before: their whole Priority 1 salt files another figure",
          "over what they needed before it" in proc and call(gross) != GOLD and call(gross + 0) >= 2800,
          "%.1f t" % gross)
    check("SF3 the 424 lane-km are the register's lane_km; centreline rates are off the plan's per lane-km basis",
          abs(L - 424.0) < 1e-6 and call(gc) != GOLD and
          "application rates, which are set in kilograms per lane-km" in plan, "centreline throughout %.1f t" % gc)
    col_ids = set(P_.reg.loc[P_.reg["func_class"] == "COLLECTOR", "segment_id"])
    pc = ps[ps["segment_id"].isin(col_ids)].copy()
    loc = pc["entry_utc"].dt.tz_convert("America/Toronto").dt.tz_localize(None)
    near = 0
    for yr in range(2023, 2027):
        for b in (dt.datetime(yr, 11, 1), dt.datetime(yr, 4, 1), dt.datetime(yr, 10, 1)):
            near += int(((loc - b).abs() < pd.Timedelta(hours=6)).sum())
            near += int(((pc["entry_utc"].dt.tz_convert(None) - b).abs() < pd.Timedelta(hours=6)).sum())
    octo = int(((loc.dt.month == 10) | (loc.dt.month == 4)).sum())
    check("SF4 the winter by local date, by UTC or from the supply season's 1 October: same call",
          near == 0 and octo == 0 and call(gu) == GOLD, "%d near, %d in October or April" % (near, octo))
    check("SF5 the planning severity is filed: 61.7 WSI points",
          "Material planning uses the planning severity of 61.7 WSI points" in plan and
          call(g * last / SEV) != GOLD and call(g * wmean / SEV) != GOLD, "")
    ctrl = P_.reg[(P_.reg["func_class"] == "COLLECTOR") & (P_.reg["priority"] == "2")]
    ws_ = np.array([P_.wsi[w] for w in K.WINTERS])
    st_ = np.array([sum(float(P_.salt.get((s, w), 0.0)) for s in ctrl["segment_id"]) for w in K.WINTERS])
    bi = np.polyfit(ws_, st_, 1)
    check("SF6 salt against severity: per WSI point, as the plan plans it (five winters cannot pin an intercept)",
          "plans each road's salt per WSI point" in plan,
          "a line through the never-moved collectors' five winters: intercept %.1f t of %.1f" % (bi[1], st_.mean()))
    check("SF7 the collectors carry rock salt only (RS), no sand and salt mix",
          set(pc["material_code"]) == {"RS"} and set(legc["MATL"]) == {"SALT"}, "")
    qrt = pd.read_excel(os.path.join(TARGET, P.F_QRT), header=3)
    reg_new = P_.reg[P_.reg["segment_id"].isin(set(nn["segment_id"]))]
    by_street = set(qrt["Street"])
    check("SF8 the transit agency's street list covers exactly the revision log's 30 August segments",
          by_street == set(reg_new["street"]) and
          set(P_.reg[P_.reg["street"].isin(by_street) & (P_.reg["func_class"] == "COLLECTOR") &
                     (P_.reg["frequent_transit"].fillna("").str.match(r"^(41|43|46|48|52|55|57|61|64)"))]["segment_id"])
          == set(nn["segment_id"]), "%d streets" % len(by_street))
    dic = pd.read_excel(os.path.join(TARGET, P.F_DICT), sheet_name="Fields")
    kg = dic[(dic["File"] == P.F_LEGACY) & (dic["Field"] == "KG")]["Description"].iloc[0]
    check("SF9 the legacy export's KG field is filed in kilograms", "kilograms" in kg, kg)
    check("SF10 the total rounded once, or the gain per lane-km rounded first: same call",
          call(r1(gain) * L) == GOLD and "rounded to the nearest 100 t" in proc, "%.1f t" % (r1(gain) * L))
    moves_all = rev[(rev["field"] == "PRIORITY") & (rev["old_value"] == "2") & (rev["new_value"] == "1")]
    check("SF11 the earlier moves are every collector moved 2 to 1 before the redesign (transit only is a rung)",
          len(moves_all) == N_MOVES + 216 and call(C.r2_pooled(P_, "QRT Route")) != GOLD, "%d entries" % len(moves_all))

    # ------------------------------------------------------------------------------------------ Ask A
    Yd = Y.Yards(W, S)
    pr = P.pass_rows(W, S)
    P2.settle_dome_edges(Yd)
    gs, gr = C.ask_a(TARGET)
    check("ASK A the golden equals the clean truth (stock on 1 November, every dome)",
          all(abs(gs[y] - Yd.stock[y]) < 1e-6 for y in K.YARDS), str({y: r1(gs[y]) for y in K.YARDS}))
    cap = pd.read_csv(os.path.join(TARGET, P.F_DOMES)).set_index("yard")["rated_capacity_t"]
    check("ASK A room is capacity less stock, every dome has room", all(abs(gr[y] - (cap[y] - gs[y])) < 1e-9 and
                                                                     gr[y] > 0 for y in K.YARDS), "")
    ea = min(min(K.bin_edge_distance(gs[y], 0.1), K.bin_edge_distance(gr[y], 0.1)) for y in K.YARDS)
    check("ASK A every figure clear of its rounding edge", ea >= ASK_EDGE, "%.4f t" % ea)
    # the changeover weighings (A3 and A4): a B-train weighed on its way at the yard that had weighed for its own
    sl = C.scale_log(TARGET)
    pw = C.preweighings(sl)
    pw_same = C.preweighings(sl, ("Vehicle", "date", "Gross kg", "Tare kg", "Net kg"))
    pairs = []
    for i in sorted(pw):
        a_ = sl.loc[i]
        b_ = sl[(sl["Vehicle"] == a_["Vehicle"]) & (sl["date"] == a_["date"]) & (sl["Gross kg"] == a_["Gross kg"]) &
                (sl.index != i)].iloc[0]
        pairs.append((a_, b_))
    dwell = lambda r: (dt.datetime.strptime(r["Time out"], "%H:%M") - dt.datetime.strptime(r["Time in"], "%H:%M")).seconds / 60
    check("ASK A five B-trains weighed on their way, each in again the same day at its own yard on the same gross "
          "under the other trailer's docket, a few minutes without tipping first and the stay at its own yard after",
          len(pw) == 5 and all(abs(int(a_["Supplier ticket"][2:]) - int(b_["Supplier ticket"][2:])) == 1 and
                                dwell(a_) <= 9 and dwell(b_) >= 14 and b_["t_in"] > a_["t_in"] and
                                a_["sheet"] in ("RDG", "KLV") and b_["sheet"] in ("LKS", "MPF") for a_, b_ in pairs),
          "%d weighings" % len(pw))
    check("ASK A A3: three of them on the receiving yard's stored tare with the outage remark kept (identical weights)",
          len(pw_same) == 3 and pw_same <= pw and all(str(sl.loc[i]["Remarks"]) != "" for i in pw_same),
          "%d" % len(pw_same))
    check("ASK A A4: two on the weighing scale house's own account and its own stored tare for the truck, no remark, "
          "the net 240 and 260 kg above the receiving yard's",
          sorted(int(a_["Net kg"] - b_["Net kg"]) for a_, b_ in pairs if a_.name not in pw_same) == [240, 260] and
          all(str(a_["Remarks"]) == "" for a_, b_ in pairs if a_.name not in pw_same), "")
    sup_ = sl[(sl["Direction"] == "IN") & sl["Carrier"].astype(str).str.startswith("Northgrit")].drop(index=sorted(pw))
    check("ASK A no two loads that tipped share a truck, a day and a gross across the scale houses (the rule merges "
          "no genuine pair)", int((sup_.groupby(["Vehicle", "date", "Gross kg"])["sheet"].nunique() > 1).sum()) == 0, "")
    inb = sl[sl["Supplier ticket"].notna() & (sl["Direction"] == "IN")]
    check("ASK A every supplier docket is keyed once on a weighing in across the four scale houses (the wet load's "
          "weighing out carries its delivery's)", inb["Supplier ticket"].is_unique and
          set(sl.loc[sl["Supplier ticket"].notna() & (sl["Direction"] == "OUT"), "Supplier ticket"]) <=
          set(inb["Supplier ticket"]), "")
    # the consistency check that exposes A3 and A4: a supplier truck cannot be back at a City scale with another load
    # sooner than its round trip from the terminal
    ng = sl[(sl["Direction"] == "IN") & sl["Carrier"].astype(str).str.startswith("Northgrit")]
    gaps = []
    for _, g_ in ng.groupby(["Vehicle", "date"]):
        g_ = g_.sort_values("t_in")
        for i1, i2 in zip(g_.index[:-1], g_.index[1:]):
            gaps.append(((g_.loc[i2, "t_in"] - g_.loc[i1, "t_in"]).total_seconds() / 60.0, i1, i2))
    inside = {i1 for m_, i1, i2 in gaps if m_ < C.ROUND_TRIP_MIN}
    other = min(m_ for m_, i1, i2 in gaps if i1 not in pw)
    check("ASK A the five are the only supplier trucks in at two scale houses inside a round trip from the terminal "
          "(the consistency check finds them and nothing else); every other same-truck day is two hours or more apart",
          inside == pw and other >= C.ROUND_TRIP_MIN and max(m_ for m_, i1, i2 in gaps if i1 in pw) < 60,
          "pairs %s min; shortest other %.0f min" % (sorted(int(m_) for m_, i1, i2 in gaps if i1 in pw), other))
    acct = ng.assign(acct=[C.dome_of(r_) for _, r_ in ng.iterrows()]).groupby(["acct", "Vehicle"])["Tare kg"].nunique()
    check("ASK A each scale house weighs a truck on one stored tare for each account (a weighing for another yard's "
          "dome on that yard's), so only where it tipped does a ticket carry the receiving yard's tare unremarked",
          (acct == 1).all(), "%d account-truck tares" % len(acct))
    lead = {}
    for _, r_ in ng.iterrows():
        n_ = int(r_["Supplier ticket"][2:]) - (1 if r_.name in pw else 0)
        lo_ = r_["t_in"] - dt.timedelta(minutes=Y.TERMINAL_MIN[r_["sheet"]])
        lead[n_] = min(lead.get(n_, lo_), lo_)
    seq = [lead[k_] for k_ in sorted(lead)]
    bt = ng[ng["Gross kg"] > 50000]
    check("ASK A the terminal's dockets run in loadout order, a B-train taking two (the lead trailer's keyed at the "
          "yard, the second only on the five)", all(a_ <= b_ for a_, b_ in zip(seq, seq[1:])) and
          all(int(t_[2:]) + 1 not in lead for t_ in bt["Supplier ticket"] if int(t_[2:]) in lead), "%d loads" % len(lead))
    stopsA = {"day_after": K.YARDS, "day_after_dome": K.YARDS, "day_whole": K.YARDS, "utc_local": K.YARDS,
              "sheets": K.YARDS, "supplier_only": ["MPF", "RDG", "LKS"], "first_issue": ["KLV", "MPF", "RDG"],
              "both_weighings": K.YARDS, "same_weights": ["KLV", "RDG"], "reader": K.YARDS}
    for m, want in stopsA.items():
        s_, r_ = C.ask_a(TARGET, m)
        moved = [y for y in K.YARDS if r1(s_[y]) != r1(gs[y])]
        gap = min(abs(s_[y] - gs[y]) for y in want)
        check("ASK A the %s route moves %s" % (m, ",".join(want)), moved == want and gap >= 0.15,
              "smallest %.1f t" % gap)
    hit = 0
    for f in np.arange(0.50, 0.95, 0.001):
        hit = max(hit, sum(1 for y in K.YARDS if r1(cap[y] * f) == r1(gs[y])))
    check("ASK A no single fill share of capacity recovers more than one dome", hit <= 1, "%d" % hit)

    # ------------------------------------------------------------------------------------------ Ask B
    ga = C.arterial_salt(TARGET)
    cells = P2.arterial_cells(W, S, pr)
    check("ASK B the golden equals the records' own arterial salt (runs once, brine off, contractor in)",
          all(abs(ga[k] - cells[k]) < 1e-6 for k in ga), "")
    truthB = {}
    for p in list(S.passes) + list(S.patrols):
        if p["kind"] != "ART":
            continue
        for rr in p["rows"]:
            w = K.winter_of(rr["entry"].date())
            if w in K.NEW_WINTERS:
                key = (p["route_code"][:3], w)
                truthB[key] = truthB.get(key, 0.0) + rr["t"]
    dB = max(abs(ga[k] - truthB.get(k, 0.0)) for k in ga)
    check("ASK B the golden within a tonne of the simulated truth in every cell", dB < 1.0, "%.3f t" % dB)
    eb = min(K.bin_edge_distance(ga[k], C.ART_UNIT) for k in ga)
    check("ASK B every cell a quarter-bin or more from its 50 t edge", eb >= 12.5, "nearest edge %.1f t" % eb)
    # B1 and B1b re-sited: nothing names the older controllers; their own rows do
    artp = ps[ps["route_code"].str.contains("-A") & (ps["material_code"] == "RS")].copy()
    artp["local"] = artp["entry_utc"].dt.tz_convert("America/Toronto").dt.tz_localize(None)
    artp["winter"] = [C.season_of_local(t_) for t_ in artp["local"]]
    artp = artp[artp["winter"].isin(K.NEW_WINTERS)]
    found = C.older_unit_winters(artp)
    truth_uw = {(u, w) for u, d_, w in zip(artp["unit_id"], artp["local"].dt.date, artp["winter"])
                if W.controller_on(u, d_) == "M410"}
    check("ASK B the older controllers found from their own rows (one total on every row of a run) are exactly the "
          "unit-winters that ran an M410", found == truth_uw, "%d unit-winters" % len(found))
    fcols = list(pd.read_csv(os.path.join(TARGET, P.F_FLEET), nrows=0).columns)
    fl0 = pd.read_csv(os.path.join(TARGET, P.F_FLEET), dtype=str)
    wo = pd.read_csv(os.path.join(TARGET, P.F_WORKORDERS), dtype=str)
    named = wo[wo["description"].str.contains("M410|M720")]
    check("ASK B no column isolates the older controllers: the fleet list carries none, no loader log ships, and the "
          "work orders name a controller model only on the four replacements",
          not any("controller" in c_ for c_ in fcols) and not os.path.exists(os.path.join(TARGET, P.F_LOADING)) and
          len(named) == 4, str(fcols))
    # the maintenance record joined by unit and date, and the plan's pre-wet rate, expose B1 and B1b
    artl = artp[artp["prewet_l"] > 0].copy()
    artl["day"] = artl["local"].dt.date
    artl["lpt"] = artl["prewet_l"] / artl["applied_t"]
    flips = []
    for u, closed in zip(named["unit_id"], named["closed"]):
        c_ = dt.date.fromisoformat(closed)
        a_ = artl[artl["unit_id"] == u]
        pre, post = a_[a_["day"] <= c_], a_[a_["day"] > c_]
        rpt = lambda x: float((x.groupby("run_ref")["applied_t"].nunique() == 1)[
            x.groupby("run_ref").size() > 1].mean())
        flips.append((u, round(float(pre["lpt"].median()), 2), round(float(post["lpt"].median()), 2),
                      round(rpt(pre), 2), round(rpt(post), 3),
                      W.unit_by_id[u]["retrofit"] is not None and c_ < W.unit_by_id[u]["retrofit"]))
    check("ASK B the four controller replacements (M410 to M720), signed off before each unit's first run on its newer "
          "controller, take its rows from 38.2 to 40.0 litres a tonne and end its repeated run totals",
          len(flips) == 4 and all(a_ == 38.2 and b_ == 40.0 and c_ == 1.0 and d_ < 0.01 and e_
                                  for u, a_, b_, c_, d_, e_ in flips), str(flips))
    rest = {u for u, w in found} - set(named["unit_id"])
    lr = artl[artl["unit_id"].isin(rest)].groupby(["unit_id", "winter"])["lpt"].median()
    check("ASK B the other eight units whose runs repeat one total carry 38.2 litres a tonne in every winter",
          len(rest) == 8 and ((lr - 40.0 / (1 + 40.0 * 1.18 / 1000)).abs() < 0.02).all() and len(lr) == 24,
          "%d units, %.2f to %.2f" % (len(rest), lr.min(), lr.max()))
    tand = set(fl0.loc[fl0["class"] == "TANDEM", "unit_id"])
    calw = wo[wo["description"].str.startswith("Pre-season spreader calibration") &
              wo["description"].str.contains("pre-wet checked at 40 L/t")]
    yrs = calw.assign(y=calw["opened"].str[:4]).groupby("unit_id")["y"].apply(set)
    check("ASK B every City tandem's spreader is calibrated each autumn with its pre-wet checked at 40 litres a tonne, "
          "the older controllers' units included (the 38.2 is no lower setting)",
          all({"2023", "2024", "2025"} <= yrs.get(u, set()) for u in tand), "%d tandems" % len(tand))
    lpt = CR.litres_per_tonne(TARGET)
    check("ASK B the litres a tonne admit one handling of the older rows: as written 38.2, the brine's salt off 38.6, "
          "the brine's weight off 40.0 (the plan's rate); the brine's salt off the newer rows puts them at 40.4",
          abs(lpt["older as written"] - 38.20) < 0.01 and abs(lpt["older, brine's salt off"] - 38.60) < 0.01 and
          abs(lpt["older, brine's weight off"] - 40.0) < 0.005 and abs(lpt["newer as written"] - 40.0) < 0.005 and
          abs(lpt["newer, brine's salt off"] - 40.44) < 0.01, "; ".join("%s %.2f" % kv for kv in lpt.items()))
    busy = set(zip(ps["unit_id"], ps["entry_utc"].dt.tz_convert("America/Toronto").dt.date)) | \
        set(zip(ps["unit_id"], ps["exit_utc"].dt.tz_convert("America/Toronto").dt.date))
    clash = 0
    for u, o_, c_ in zip(wo["unit_id"], wo["opened"], wo["closed"]):
        d0, d1_ = dt.date.fromisoformat(o_), dt.date.fromisoformat(c_)
        while d0 <= d1_:
            clash += (u, d0) in busy
            d0 += dt.timedelta(days=1)
    check("ASK B no work order books a unit in on a day it spread or patrolled; every unit is in the fleet list; the "
          "garage works weekdays", clash == 0 and set(wo["unit_id"]) <= set(fl0["unit_id"]) and
          (pd.to_datetime(wo["opened"]).dt.weekday < 5).all() and (pd.to_datetime(wo["closed"]).dt.weekday < 5).all(),
          "%d work orders" % len(wo))
    # HB2 re-sited: the platform carries a renamed circuit's current code on every pass it made
    hist = C.circuit_history(TARGET)
    codes = set(artp["route_code"])
    first = artp.groupby("route_code")["local"].min()
    check("ASK B the platform shows the realigned circuits under their current codes back to October 2023, and the "
          "revision log takes them back", not ({"MPF-A01", "MPF-A02"} & codes) and
          {hist.get("KLV-A15", ("",))[0], hist.get("KLV-A16", ("",))[0]} == {"MPF-A01", "MPF-A02"} and
          all(first[c_] < pd.Timestamp("2023-12-01") for c_ in ("KLV-A15", "KLV-A16")), str(sorted(hist)))
    gp = C.arterial_salt(TARGET, "prewet_rate")
    check("ASK B the salt read back from the brine at 40 litres a tonne files the same 50 t figures (C1)",
          all(C.r50(gp[k]) == C.r50(ga[k]) for k in ga), "largest gap %.3f t" % max(abs(gp[k] - ga[k]) for k in ga))
    art = ps[ps["route_code"].str.contains("-A") & (ps["material_code"] == "RS")].copy()
    art["day"] = art["entry_utc"].dt.tz_convert("America/Toronto").dt.date
    m410 = [W.controller_on(u, d) == "M410" for u, d in zip(art["unit_id"], art["day"])]
    art["m410"] = m410
    r720 = art[~art["m410"]]
    r410 = art[art["m410"]]
    check("ASK B newer controllers' rows carry 40 litres a tonne exactly, older ones about 38.2 (salt and brine)",
          (abs(r720["prewet_l"] - 40.0 * r720["applied_t"]) <= 0.051).all() and
          abs(float((r410["prewet_l"] / r410["applied_t"]).median()) - 40.0 / (1 + 40.0 * 1.18 / 1000)) < 0.05,
          "%.2f" % float((r410["prewet_l"] / r410["applied_t"]).median()))
    runs = r410.groupby("run_ref")["applied_t"].nunique()
    check("ASK B every older-controller run repeats one total on each of its rows", (runs == 1).all(),
          "%d runs" % len(runs))
    two = [(y, w) for y in ("KLV", "MPF") for w in K.NEW_WINTERS]
    stopsB = {"rows": list(ga), "runs_once": list(ga), "net_brine": list(ga), "no_contract": two,
              "register_yard": two, "code_yard": two}
    for m, want in stopsB.items():
        a_ = C.arterial_salt(TARGET, m)
        moved = sorted(k for k in ga if C.r50(a_[k]) != C.r50(ga[k]))
        gap = min(abs(a_[k] - ga[k]) for k in want)
        check("ASK B the %s route moves %d figures at 50 t" % (m, len(want)), moved == sorted(want),
              "smallest %.1f t" % gap)
    ro = C.arterial_salt(TARGET, "runs_once")
    hv = C.arterial_salt(TARGET, "halve")
    us = C.arterial_salt(TARGET, "utc_season")
    check("ASK B the brine left in (B1b) lands a full 50 t cell or more above in every cell",
          all(C.r50(ro[k]) - C.r50(ga[k]) >= 50 for k in ga),
          "smallest excess %.1f t; halving moves %d cells (%.0f to %+.0f t); the UTC season moves %d" % (
              min(ro[k] - ga[k] for k in ga), sum(C.r50(hv[k]) != C.r50(ga[k]) for k in ga),
              min(hv[k] - ga[k] for k in ga), max(hv[k] - ga[k] for k in ga),
              sum(C.r50(us[k]) != C.r50(ga[k]) for k in ga)))
    lk = P_.reg[P_.reg["func_class"] == "ARTERIAL"].copy()
    lk["y"] = lk["route_code"].str[:3]
    tot_l = lk["lane_km"].sum()
    blind = 0
    for w in K.NEW_WINTERS:
        rate_ = sum(ga[(y, w)] for y in K.YARDS) / tot_l
        for y in K.YARDS:
            blind += int(C.r50(rate_ * lk.loc[lk["y"] == y, "lane_km"].sum()) == C.r50(ga[(y, w)]))
    check("ASK B one network rate times each yard's arterial lane-km recovers no cell", blind == 0, "%d" % blind)

    # ------------------------------------------------------------------------------------------ the cracker
    cr = CR.main(TARGET, quiet=True)
    check("CRACKER both blind readers' own paths, re-run on this pack, keep no ask figure (ceiling a fifth, 4 of 20)",
          all(v == 0 for v in cr["reader_hits"].values()), str(cr["reader_hits"]))
    rd_ = cr["readings"]
    check("CRACKER the judge's literal readings that keep the double weighings, the brine, or take the brine's salt "
          "off the older runs keep no figure", rd_["judge, double weighings kept"] == 0 and
          rd_["judge, brine kept"] == 0 and rd_["judge, brine's salt off the older runs"] == 0, str(rd_))
    check("CRACKER pair ceiling 40 or under on the readers' paths", cr["pair"] <= 40.0, "%.1f" % cr["pair"])
    check("CRACKER mirror leakage under 35 per cent of ask weight", cr["leak"] < 35.0, "%.1f" % cr["leak"])
    for name, (k_, tot_, pr_) in cr["single"].items():
        print("INFO %-74s %d of 20, pair %.1f" % (name[:74], k_, pr_))
    print("INFO %-74s %d of 12 (the newer rows then carry 40.4 litres a tonne)" % (
        "the judge's brine's salt off every row", rd_["judge, brine's salt off every row"]))

    # ------------------------------------------------------------------------------------------ separation
    fl = pd.read_csv(os.path.join(TARGET, P.F_FLEET), dtype=str)
    old_units = {u for u, w in found}
    check("SEPARATION every older-controller unit is an arterial tandem whose every pass is on an arterial circuit",
          set(fl.loc[fl["unit_id"].isin(old_units), "class"]) == {"TANDEM"} and
          ps[ps["unit_id"].isin(old_units)]["route_code"].str.contains("-A").all(), "%d units" % len(old_units))
    # the repair's scope against the shipped commit (d87ff77a): the files the call reads, the call's rows in the
    # platform export and the dictionary, and the two network rungs that read the arterial rows as written
    sha = lambda f: hashlib.sha256(open(os.path.join(TARGET, f), "rb").read()).hexdigest()
    rows_sha = lambda df_: hashlib.sha256(df_.to_csv(index=False).encode()).hexdigest()
    check("SEPARATION the files the call reads are byte-identical to the shipped pack",
          all(sha(f) == h for f, h in SHIPPED_CALL_FILES.items()), "%d files" % len(SHIPPED_CALL_FILES))
    check("SEPARATION the platform export's collector rows are the shipped rows, and its arterial rows the shipped rows "
          "in every column but the realigned circuits' codes",
          rows_sha(ps[~ps["route_code"].str.contains("-A")]) == SHIPPED_COLLECTOR_ROWS and
          rows_sha(ps[ps["route_code"].str.contains("-A")].drop(columns=["route_code"])) == SHIPPED_ARTERIAL_ROWS, "")
    dic_f = pd.read_excel(os.path.join(TARGET, P.F_DICT), sheet_name="Fields")
    check("SEPARATION the dictionary's rows for the call's files are the shipped rows",
          rows_sha(dic_f[dic_f["File"].isin(list(SHIPPED_CALL_FILES) + [P.F_PASSES]) & (dic_f["Field"] != "route_code")])
          == SHIPPED_DICT_CALL_ROWS, "")
    r0, r0b = C.r0_network(P_, sev=P_.wsi["2025-26"], winters=["2025-26"]), C.r0_network(P_)
    check("SEPARATION the network rungs that read the arterial rows recompute to the shipped figures (R0 5,045.3, "
          "R0b 5,392.8)", round(r0, 1) == 5045.3 and round(r0b, 1) == 5392.8, "%.1f, %.1f" % (r0, r0b))
    check("SEPARATION the collectors spread dry salt (no brine on any collector row)",
          int(pc["prewet_l"].sum()) == 0, "")
    xw = pd.read_csv(os.path.join(TARGET, P.F_XWALK), dtype=str)
    split = set(xw.loc[xw["change"] == "SPLIT", "segment_id"])
    check("SEPARATION no collector segment the call reads is split by the crosswalk",
          not (split & (set(d["segment_id"]) | set(nn["segment_id"]))), "%d split" % len(split))
    con = P_.reg[P_.reg["operator"] == "CONTRACT"]
    yardch = rev[rev["field"].isin(["YARD", "OPERATOR"])].merge(P_.reg[["segment_id", "func_class"]], on="segment_id")
    check("SEPARATION the contractor and the realignment touch arterials only",
          set(con["func_class"]) == {"ARTERIAL"} and set(yardch["func_class"]) == {"ARTERIAL"}, "")
    asks_files = [P.F_SCALE, P.F_DISPATCH1, P.F_DISPATCH2, P.F_SURVEY, P.F_DOMES, P.F_PELLARD, P.F_FLEET,
                  P.F_WORKORDERS]
    tmp = linked_copy(asks_files)
    try:
        g_del = C.golden(C.Pack(tmp))
    finally:
        shutil.rmtree(tmp)
    check("SEPARATION the golden reads none of the asks' files (deleted: unchanged)", abs(g_del - g) < 1e-9, "")
    check("SEPARATION Ask B reads arterial route codes only, Ask A no AVL row", True,
          "by construction in calc.arterial_salt and calc.ask_a")

    # ------------------------------------------------------------------------------------------ input gates
    meta = json.load(open(os.path.join(TASK, "metadata.json")))
    files = sorted(os.listdir(TARGET))
    fmts = {f.rsplit(".", 1)[1] for f in files}
    check("GATES ten or more files", len(files) >= 10, "%d" % len(files))
    check("GATES three or more formats", len(fmts) >= 3, str(sorted(fmts)))
    check("GATES a file of 25,000 rows or more", meta["input_gates"]["largest_file_rows"] >= 25000,
          "%s %d" % (meta["input_gates"]["largest_file"], meta["input_gates"]["largest_file_rows"]))
    dis = meta["distractor_files"]
    check("GATES two distractors, shipped, none named as such", len(dis) >= 2 and all(x in files for x in dis) and
          not any("distract" in x.lower() for x in files), str(dis))
    tmp = linked_copy(dis)
    try:
        Pd = C.Pack(tmp)
        same = abs(C.golden(Pd) - g) < 1e-9 and C.ask_a(tmp) == (gs, gr) and C.arterial_salt(tmp) == ga
    finally:
        shutil.rmtree(tmp)
    check("DISTRACTORS deleted: the golden and both asks unchanged", same, "")
    man = pd.read_csv(os.path.join(TARGET, P.F_MANIFEST), dtype=str)
    check("REGISTERS the manifest lists every shipped file but itself (set equality)",
          set(man["file"]) == set(files) - {P.F_MANIFEST}, str(sorted(set(files) - set(man["file"]))))
    dic_files = set(dic["File"])
    data_files = {f for f in files if f.endswith((".csv", ".parquet", ".xlsx"))} - {P.F_DICT}
    check("REGISTERS the dictionary covers every shipped data file (set equality)", dic_files == data_files,
          str(sorted(data_files ^ dic_files)))
    check("GATES every shipped file has a source, a date and a licence in metadata.json",
          sorted(x["path"] for x in meta["files"]) == files and all(x["source"] and x["date"] and x["license"]
                                                                   for x in meta["files"]), "")

    # ------------------------------------------------------------------------------------------ realism
    dash = [f for f, t in txt.items() if "\u2014" in t or "\u2013" in t]
    check("REALISM no shipped file carries an em or en dash", not dash, str(dash))
    sub = open(os.path.join(TASK, "submission.md"), encoding="utf-8").read()
    pt = os.path.join(TASK, "prompt.txt")
    prompt = open(pt, encoding="utf-8").read() if os.path.exists(pt) else ""
    check("REALISM the submission and prompt carry no em or en dash",
          not any(c in sub + prompt for c in ("\u2014", "\u2013")), "")
    alltxt = " ".join(txt.values())
    check("REALISM every persona from guard.py names appears, and no placeholder identity",
          all(n in alltxt for n in K.PEOPLE.values()) and
          not re.search(r"John Doe|Jane Doe|Acme|Company [AB]\b|Lorem|Example Corp", alltxt), "")
    tot_rs = float(ps.loc[ps["material_code"] == "RS", "applied_t"].sum())
    check("REALISM no round headline totals (platform rock salt, WSI, ask figures)",
          abs(tot_rs - round(tot_rs, -1)) > 0.5 and all(abs(v - round(v)) > 0.01 for v in list(ga.values()) +
                                                         list(gs.values())), "%.1f t" % tot_rs)
    tk = []
    for y in K.YARDS:
        df = pd.read_excel(os.path.join(TARGET, P.F_SCALE), sheet_name=y, dtype=str)
        tk += list(df["Ticket"])
    check("REALISM scale tickets unique across the four scale houses", len(tk) == len(set(tk)), "%d" % len(tk))
    units = set(fl["unit_id"])
    check("REALISM every unit in the passes is in the fleet list", set(ps["unit_id"]) <= units, "")
    tr = pd.read_excel(os.path.join(TARGET, P.F_SCALE), sheet_name="RDG", dtype=str)
    unit_tr = set(tr.loc[tr["Carrier"] == "CITY OF QUELLBROOK", "Vehicle"].str.replace("Unit ", ""))
    check("REALISM the transfer truck is a Ridge unit in the fleet list",
          unit_tr <= set(fl.loc[fl["yard"] == "RDG", "unit_id"]), str(unit_tr))
    mvr = rev[(rev["field"] == "PRIORITY") & (rev["old_value"] == "2")]
    check("REALISM each move reason is one programme on one effective date",
          all(mvr[mvr["reason"] == r_]["effective_date"].nunique() == 1 for r_ in mvr["reason"].unique()),
          "%d reasons" % mvr["reason"].nunique())
    lagd = (pd.to_datetime(rev["effective_date"]) - pd.to_datetime(rev["entered_on"])).dt.days
    check("REALISM route audit entries lag their effective dates unevenly, some entered after",
          lagd.min() < 0 and lagd.max() - lagd.min() >= 40, "%d to %d days" % (lagd.min(), lagd.max()))
    batch = rev[rev["reason"] != "Assumed from developer, subdivision collector"]
    sit = batch.groupby(["effective_date", "reason"])["entered_on"].nunique()
    pair = rev[rev["field"].isin(["PRIORITY", "ROUTE"]) & rev["reason"].isin(mvr["reason"].unique())]
    same = pair.groupby(["segment_id", "effective_date"])["entered_on"].nunique()
    red = rev[rev["effective_date"] == "2026-08-30"]
    wk = pd.to_datetime(rev["entered_on"]).dt.weekday
    check("REALISM the audit keys each change in a few sittings, a segment's priority and route rows together, on "
          "working days", sit.max() <= 6 and (same == 1).all() and red["entered_on"].nunique() <= 6 and (wk < 5).all(),
          "redesign on %d days; at most %d sittings a change" % (red["entered_on"].nunique(), sit.max()))
    cnt = d.groupby([str(e) for e in d["eff"]]).size()
    check("REALISM the earlier moves' yearly counts follow their programmes (uneven)",
          cnt.max() >= 2 * cnt.min(), str(cnt.to_dict()))
    p1w = {}
    for w in K.WINTERS:
        tt = dd_ = 0.0
        for s_ in W.segments:
            if s_["func_class"] == "COLLECTOR" and s_["pool"] != "NEW" and W.priority(s_, w) == 1:
                tt += float(P_.salt.get((s_["segment_id"], w), 0.0))
                dd_ += s_["lane_km"] * P_.wsi[w]
        p1w[w] = tt / dd_ * SEV
    check("REALISM the Priority 1 collector rate per point swings between winters (10 to 25 per cent)",
          0.10 <= max(p1w.values()) / min(p1w.values()) - 1 <= 0.25, str({w: round(v, 2) for w, v in p1w.items()}))
    lenc = {}
    for pool in ("NEW", "MOVE"):
        ss = [s_ for s_ in W.segments if s_["pool"] == pool]
        if pool == "NEW":
            rr = [P_.rate(s_["segment_id"], K.WINTERS, s_["lane_km"]) for s_ in ss]
        else:
            pre_ = d.set_index("segment_id")["pre"]
            rr = [float(pre_[s_["segment_id"]]) for s_ in ss]
        lenc[pool] = float(np.corrcoef(rr, [1.0 / s_["centreline_km"] for s_ in ss])[0, 1])
    check("REALISM short segments carry more Priority 2 salt a lane-km: plain on the redesign's flat streets, "
          "swamped on the earlier moves by their hills, frontages and approaches", lenc["NEW"] > 0.5 and
          abs(lenc["MOVE"]) < 0.3, "corr with 1/length: new %.2f, earlier moves %.2f" % (lenc["NEW"], lenc["MOVE"]))
    upc = ps.groupby(["route_code"])["unit_id"].nunique()
    check("REALISM trucks rotate: most circuits run by more than one unit over the three winters",
          (upc > 1).mean() > 0.5, "%.0f per cent of circuits" % (100 * (upc > 1).mean()))

    # ------------------------------------------------------------------------------------------ dates
    end_local = ps["exit_utc"].dt.tz_convert("America/Toronto").max()
    check("DATES the platform export ends before the 24 September pull window",
          end_local.tz_localize(None) < pd.Timestamp("2026-09-24"), str(end_local))
    sc_max = max(pd.read_excel(os.path.join(TARGET, P.F_SCALE), sheet_name=y)["Date"].max() for y in K.YARDS)
    check("DATES the scale houses' last weighing on 24 September", str(sc_max)[:10] == "2026-09-24", str(sc_max))
    check("DATES the revision log entered no later than the register date",
          rev["entered_on"].max() <= "2026-09-14" and rev["effective_date"].max() <= "2026-09-14", "")
    d1 = pd.read_csv(os.path.join(TARGET, P.F_DISPATCH2))
    check("DATES only the supplier's dispatch schedules carry dates after the pull (forward bookings)",
          d1["delivery_date"].min() >= "2026-10-01", "")
    check("DATES every document issued before the pull",
          max(D.PLAN_DATE, D.PROC_DATE, D.SURVEY_DATE) <= K.AS_OF and "Issued 18 September 2026" in plan and
          "issued 16 September 2026" in txt[P.F_SURVEY].replace("\n", " ") and "Approved 12 May 2025" in proc, "")
    sv_ = C.survey(TARGET)
    ok_cap = all(7 <= (sv_[y]["capture_end_utc"] - dt.timedelta(hours=4)).hour <= 9 and
                 sv_[y]["capture_end_utc"].date() == Y.SURVEY_DAY[y] for y in K.YARDS)
    check("DATES each survey capture ends in the morning of its survey day, local", ok_cap, "")

    # ------------------------------------------------------------------------------------------ documents read back
    check("DOCS the procedure files the estimand (4.2), positive control on the same extractor",
          "the additional rock salt the roads the change moves will need from 1 November to 31 March" in proc and
          "Road salt procurement" in proc, "")
    check("DOCS the plan files the Priority 2 service level and the dome stock rule",
          "Salted at intersections, hills and curves during and after each event" in plan and
          "brought forward by the loads weighed in to it and out of it at the scale houses" in plan, "")
    check("DOCS the plan files the pre-wet brine and its weight", "1.18 kg a litre" in plan, "")
    check("DOCS the survey report's stock table reads back", all(abs(sv_[y]["stock"] - Y.SURVEY_STOCK[y]) < 1e-9
                                                                 for y in K.YARDS), "")

    n_fail = sum(1 for _l, ok, _d in RESULTS if not ok)
    print("\n%d checks, %d failed" % (len(RESULTS), n_fail))
    sys.exit(1 if n_fail else 0)


if __name__ == "__main__":
    main()
