"""task229 generator: the filed documents. The hydrometric data standard, the 2027 data plan, the rating book,
Corriston Hydro's 2027 operating plan and the 2025 flood frequency update (a distractor)."""
import os
from datetime import date

from docx import Document
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import params as P
import postprocess
import world as W

ORG = "Corriston Basin Hydrometric Cooperative"
PDF_ORG = "Corriston Basin Cooperative"
HYDRO = "Corriston Hydro Ltd."


def _styles():
    ss = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=ss["BodyText"], fontName="Helvetica", fontSize=9.5, leading=13, spaceAfter=4)
    h = ParagraphStyle("h", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=10.5, spaceBefore=8, spaceAfter=3)
    t = ParagraphStyle("t", parent=ss["Title"], fontName="Helvetica-Bold", fontSize=14, alignment=0, spaceAfter=4)
    small = ParagraphStyle("s", parent=body, fontSize=8, leading=10.5, textColor=colors.HexColor("#444444"))
    return body, h, t, small


def _footer(text):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(20 * mm, 10 * mm, text)
        canvas.drawRightString(190 * mm, 10 * mm, "Page %d" % doc.page)
        canvas.restoreState()
    return draw


def _pdf(path, title, author, footer, blocks, producer, stamp):
    body, h, t, small = _styles()
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=18 * mm,
                            bottomMargin=18 * mm, title=title, author=author, creator=author, invariant=1)
    flow = []
    for kind, val in blocks:
        if kind == "title":
            flow.append(Paragraph(val, t))
        elif kind == "h":
            flow.append(Paragraph(val, h))
        elif kind == "p":
            flow.append(Paragraph(val, body))
        elif kind == "small":
            flow.append(Paragraph(val, small))
        elif kind == "sp":
            flow.append(Spacer(1, val))
        elif kind == "table":
            rows, widths = val
            tb = Table(rows, colWidths=widths, repeatRows=1)
            tb.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", 7.6),
                                    ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 7.6),
                                    ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#9a9a9a")),
                                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e6eaef")),
                                    ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            flow.append(tb)
    doc.build(flow, onFirstPage=_footer(footer), onLaterPages=_footer(footer))
    postprocess.scrub(path, producer, stamp)


# ------------------------------------------------------------------------------------------------ the standard
def standard(path):
    _pdf(path, "Hydrometric data standard, edition 2", "Hydrometric section", "Corriston Hydrometric Network, hydrometric data standard, edition 2", [
        ("title", "Corriston Hydrometric Network: hydrometric data standard, edition 2"),
        ("small", "Corriston Basin Hydrometric Cooperative, hydrometric section. Adopted by the members on 14 March 2025; "
                  "replaces edition 1 (2016). Owner: the data quality manager."),
        ("h", "1. Scope"),
        ("p", "This standard governs the water level and discharge records of the Corriston Hydrometric Network: how a "
              "day and a water year are reckoned, how a month of a station's published discharge is graded, and how "
              "gaugings and ratings are kept. It applies to every station in the station register."),
        ("h", "2. The day and the water year"),
        ("p", "A day runs from midnight to midnight local standard time (UTC-8) throughout the year; clock changes do "
              "not move it. A water year runs from 1 October to 30 September and is named for the year in which it "
              "ends."),
        ("h", "3. Grade of a station-month"),
        ("p", "3.1 A station-month is the record of one station for one calendar month."),
        ("p", "3.2 A station-month meets grade B when the mean of the station's published daily discharges for the "
              "month is within 5 per cent of the month's true mean discharge, the mean over the month's days of the "
              "discharge that passed the station."),
        ("p", "3.3 A station-month that does not meet grade B is graded C."),
        ("p", "3.4 A station-month graded C is published with the qualifier <i>below grade B</i>."),
        ("h", "4. Extrapolation"),
        ("p", "A station-month whose mean discharge exceeds one and a half times the largest discharge gauged at the "
              "station is graded C whatever its error, because the rating above that point rests on no measurement. "
              "The largest discharge gauged at a station is the top of the gauged range printed under its rating in "
              "the rating book."),
        ("h", "5. Gaugings"),
        ("p", "5.1 A gauging measures the discharge passing the station's section and records the gauge height at the "
              "measurement."),
        ("p", "5.2 Gaugings are made at steady flow: the gauge height may not move by more than 1 cm over the "
              "measurement. A section that splits into channels is measured channel by channel and the gauging is "
              "their sum."),
        ("p", "5.3 A gauging's difference from a rating is the measured discharge less the rating's discharge at the "
              "gauging's gauge height, stated as a percentage of the rating's discharge."),
        ("h", "6. Ratings"),
        ("p", "6.1 A station's rating relates its water level to its discharge. Since 2014 every network rating has "
              "been drawn as the curve Q = C (h - e)<super>b</super>, where h is the water level in metres above the "
              "station's gauge datum, e the water level at which flow ceases, and C and b are fitted to the station's "
              "gaugings."),
        ("p", "6.2 A rating takes effect on the date the rating book gives and stays in force until a later rating "
              "takes effect. Every gauging is compared with the rating in force on its day."),
        ("h", "7. Records kept"),
        ("p", "The network keeps each station's daily record, its gaugings, its ratings and its field visits for the "
              "life of the station, and a station moved to a new section keeps its record under its new number."),
    ], PDF_ORG, "2025-03-14")


# ------------------------------------------------------------------------------------------------ the data plan
def _docx(path, build, who, created, modified):
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    build(d)
    d.core_properties.author = who
    d.core_properties.last_modified_by = who
    d.core_properties.title = ""
    d.save(path)
    postprocess.scrub(path, who, modified[:10])
    postprocess.finish_office(path, created, modified, ORG)


def data_plan(path):
    def build(d):
        d.add_heading("Corriston Hydrometric Network: 2027 data plan", level=1)
        d.add_paragraph("Corriston Basin Hydrometric Cooperative. For the provincial water licence "
                        "inspector's annual review, due 30 November 2026. Prepared 5 November 2026 by Phyllis "
                        "Thompson, data quality manager.")
        d.add_heading("1. The network in 2027", level=2)
        d.add_paragraph("The network runs the 23 stations of the station register through 2027. The 2026 telemetry "
                        "migration finished on 28 September 2026: every station now reports on a cellular logger or, "
                        "in the Talbot Canyon and Merrin Creek gorges where there is no cellular coverage, on an "
                        "Iridium platform. No station changes platform or rating in 2027. Daily values are transmitted "
                        "and processed in 2027 as the data dictionary describes, and the ratings in force are those of "
                        "the rating book, issue 2.")
        d.add_heading("2. The 2027 grade review", level=2)
        d.add_paragraph("The review grades each station-month of 2027 under sections 3 and 4 of the hydrometric data "
                        "standard, on the flows that the 2027 water supply outlook (issued 2 November 2026) and "
                        "Corriston Hydro's 2027 operating plan put past each station.")
        d.add_paragraph("From the first month of 2027 in which the review expects a station-month to be graded C, the "
                        "network publishes that station's daily discharge with the qualifier below grade B until the "
                        "station is reviewed again.")
        d.add_heading("3. The 2026 water year", level=2)
        d.add_paragraph("The review covers the 2026 water year's record as well, its completeness and its gaugings. "
                        "The approved record runs to 30 June 2026; July to September 2026 stay provisional until the "
                        "annual review closes in January.")
        d.add_heading("4. Changes in 2027", level=2)
        d.add_paragraph("Merrin Falls Generating Station enters service in 2027 (Corriston Hydro's plan gives the date "
                        "and its operation). No new station opens and none closes.")
    _docx(path, build, "Phyllis Thompson", "2026-10-21T16:40:00Z", "2026-11-05T22:15:00Z")


# ------------------------------------------------------------------------------------------------ the rating book
def rating_book(path, GR):
    hdr = ["Station", "Name", "Rating", "In force from", "C", "e (m)", "b", "Gauged range (m3/s)", "Top gauged on"]
    rows = [hdr]
    for no in W.CURRENT:
        C, e, b = W.RATING[no]
        g = GR[no]
        frm = "4 Jun 2026" if no == "C122" else "1 Apr 2026"
        rows.append([no, W.ST[no]["name"], W.RATING_NO[no], frm, "%.4g" % C, "%.3f" % e, "%.2f" % b,
                     "%s to %s" % (_q(g["low"]), _q(g["high"])), g["high_date"].strftime("%d %b %Y").lstrip("0")])
    closed = [hdr]
    C, e, b = W.RATING["C110"]
    g = GR["C110"]
    closed.append(["C110", W.ST["C110"]["name"], W.RATING_NO["C110"], "1 Apr 2026", "%.4g" % C, "%.3f" % e, "%.2f" % b,
                   "%s to %s" % (_q(g["low"]), _q(g["high"])), g["high_date"].strftime("%d %b %Y").lstrip("0")])
    hs = [0.4, 0.6, 0.8, 1.0, 1.4, 1.8, 2.4]
    trows = [["Station"] + ["h = %.1f m" % h for h in hs]]
    for no in W.CURRENT + ["C110"]:
        par = W.RATING[no]
        trows.append([no] + [("%.4g" % float(W.q_of_h(par, h))) if h > par[1] + 0.02 else "" for h in hs])
    w9 = [13 * mm, 50 * mm, 15 * mm, 18 * mm, 11 * mm, 11 * mm, 9 * mm, 25 * mm, 18 * mm]
    _pdf(path, "Rating book, issue 2", "Hydrometric section", "Corriston Hydrometric Network, rating book, issue 2 (30 June 2026)", [
        ("title", "Corriston Hydrometric Network: rating book, issue 2"),
        ("small", "Ratings in force from 1 April 2026, refitted to each station's 2016 to 2026 gaugings at the move to "
                  "the new hydrometric database; Mattin Creek at Quarry Road (C122) from 4 June 2026, fitted to the "
                  "section survey and its first gaugings. Issued 30 June 2026. Replaces the ratings of the 2019 rating "
                  "review."),
        ("p", "Each rating is the curve Q = C (h - e)<super>b</super>, Q in cubic metres a second and h the water "
              "level in metres above the station's gauge datum. The gauged range is the lowest and the highest "
              "discharge among the gaugings the rating was fitted to."),
        ("table", (rows, w9)),
        ("h", "Closed sites"),
        ("p", "Mattin Creek at the gravel pit (C110) closed on 8 June 2026 when the station moved 300 m upstream to "
              "Quarry Road; its record continues under C122."),
        ("table", (closed, w9)),
        ("h", "Rating tables (extract)"),
        ("small", "Discharge in cubic metres a second at the water levels shown; blank below the level at which flow "
                  "ceases."),
        ("table", (trows, [16 * mm] + [21 * mm] * len(hs))),
    ], PDF_ORG, "2026-06-30")


def _q(x):
    return ("%#.3g" % x) if x < 10 else ("%.1f" % x)


# ------------------------------------------------------------------------------------------------ the operating plan
def operating_plan(path):
    T_, M_ = W.TALBOT, W.MERRIN
    plants = [["Plant", "River", "Tailrace (river km)", "Capacity (m3/s)", "Minimum release (m3/s)", "In service"],
              [T_["name"], T_["river"], "%.1f" % T_["tailrace_km"], "%.0f" % T_["cap"], "%.1f" % T_["fmin"], "2009"],
              [M_["name"], M_["river"], "%.1f" % M_["tailrace_km"], "%.0f" % M_["cap"], "%.1f" % M_["fmin"], "1 June 2027"],
              [W.FENN["name"], W.FENN["river"], "%.1f" % W.FENN["tailrace_km"], "7", "none", "2012"]]
    mon = ["Plant"] + W.M
    tal = ["Talbot Canyon"] + ["continuous" if m in W.CONTINUOUS_MONTHS else "peaking" for m in range(1, 13)]
    mer = ["Merrin Falls"] + [("not in service" if m < 6 else ("continuous" if m in W.CONTINUOUS_MONTHS else "peaking")) for m in range(1, 13)]
    fen = ["Fenn Creek"] + ["run-of-river"] * 12
    _pdf(path, "2027 operating plan", "Brandon Holden", "Corriston Hydro Ltd., 2027 operating plan, Corriston and Merrin systems", [
        ("title", "Corriston Hydro Ltd.: 2027 operating plan, Corriston and Merrin systems"),
        ("small", "Issued 19 October 2026 to the Comptroller of Water Rights and the Corriston Basin Hydrometric Cooperative. "
                  "Brandon Holden, operations planning."),
        ("h", "1. Plants"),
        ("table", (plants, [52 * mm, 27 * mm, 22 * mm, 20 * mm, 26 * mm, 20 * mm])),
        ("p", "Talbot Canyon has operated under the peaking terms of its 2017 water licence amendment since that year. "
              "Merrin Falls enters service on 1 June 2027 under a licence on the same terms."),
        ("h", "2. Operating modes in 2027"),
        ("table", ([mon, tal, mer, fen], [24 * mm] + [12.2 * mm] * 12)),
        ("h", "3. Peaking"),
        ("p", "On a peaking day a plant generates in two windows, 06:00 to 10:00 and 16:00 to 21:00 local time, every "
              "day of the week, holidays included. Between the windows it holds its minimum release. Each day's inflow "
              "is released within the day: in the windows the plant generates at the rate that passes the day's "
              "inflow less the minimum release held outside them. Under the licence's ramping condition, output moves "
              "between the minimum release and the generating rate over 30 minutes at the start and at the end of each "
              "window, inside the window. When a "
              "day's inflow is more than the windows can pass at capacity, the plant generates continuously that day."),
        ("h", "4. Continuous generation"),
        ("p", "From 1 May to 30 June the forebays cannot hold the freshet back for the windows, and each plant passes "
              "its inflow at an even rate through the day."),
        ("h", "5. Fenn Creek"),
        ("p", "Fenn Creek Generating Station is a run-of-river plant with no pondage: it passes the creek's flow as it "
              "arrives."),
        ("h", "6. Works"),
        ("p", "No works in the river channel below any tailrace are planned for 2027. Forebay sediment flushing at "
              "Talbot Canyon is scheduled for the continuous-generation period."),
    ], HYDRO, "2026-10-19")


# ------------------------------------------------------------------------------------------------ distractor
def flood_frequency(path, peaks):
    rows = [["Station", "Name", "Years", "2-year", "10-year", "50-year", "200-year"]]
    for no, name, n, q2, q10, q50, q200 in peaks:
        rows.append([no, name, str(n), _q(q2), _q(q10), _q(q50), _q(q200)])
    _pdf(path, "Flood frequency update 2025", "Hydrometric section", "Corriston Hydrometric Network, flood frequency update 2025", [
        ("title", "Corriston Hydrometric Network: flood frequency update, 2025"),
        ("small", "Hydrometric section, 9 December 2025. For the flood management programme and the culvert and bridge "
                  "design guide."),
        ("p", "Annual maximum instantaneous discharge at the network's unregulated stations, from each station's "
              "logger record through water year 2025, fitted with a generalised extreme value distribution by L-moments. "
              "Stations below generating stations and the spring-fed creeks are not analysed. Discharges in cubic "
              "metres a second."),
        ("table", (rows, [15 * mm, 55 * mm, 12 * mm, 18 * mm, 18 * mm, 18 * mm, 18 * mm])),
        ("p", "The update supersedes the 2019 estimates. Estimates beyond the 50-year event rest on few peaks and "
              "carry wide intervals; the design guide applies its own allowance for climate change."),
    ], PDF_ORG, "2025-12-09")
