"""Write every file under target/ from the simulated world."""
import csv
import datetime as dt
import hashlib
import io
import json
import os

import numpy as np
import openpyxl
from openpyxl.styles import Font, Alignment

import world as W
import sim as S
import records as R

FILES = {
    "mdm": "mdm_interval_receipts_20261005_20261011.csv",
    "hes": "hes_meter_firmware_history.json",
    "register": "mp_installation_register_20261012.xlsx",
    "catalogue": "fw_release_catalogue.csv",
    "wave8": "wave8_schedule.xlsx",
    "standard": "MDS-STD-114_firmware_read_health_r3.pdf",
    "fieldref": "mds_extract_field_reference.md",
    "log": "extract_log_20261012.csv",
    "census": "hes_firmware_census_20261012.csv",
    "telco": "carrier_maintenance_notices_oct2026.xlsx",
}
DISTRACTORS = ["census", "telco"]


def _hex(*parts, n=10):
    return hashlib.sha1("|".join(str(p) for p in parts).encode()).hexdigest()[:n]


# ------------------------------------------------------------------ MDM receipts

BASE_KWH = {"E1": (4.0, 3.1), "E2": (3.0, 2.2), "B1": (3.0, 4.4), "Q1": (4.0, 6.0)}


def write_mdm(path, insts, rows):
    r = S.rng("kwh")
    prof = {i.nmi: i.profile for i in insts}
    level = {}
    for i in sorted(insts, key=lambda x: x.nmi):
        cls = W.PROFILES[i.profile][2]
        for sx in W.PROFILES[i.profile][0]:
            k, th = BASE_KWH[sx]
            lv = r.gamma(k, th) * (9.0 if cls == "SME" and sx in ("E1", "Q1") else 1.0)
            level[(i.nmi, sx)] = lv
    day_fx = [1.04, 0.97, 1.01, 0.99, 1.06, 0.92, 0.95]
    out = []
    for (nmi, sx, rd, L, t, a, s_, f, e, vi) in rows:
        dd = (rd - W.READ_DATES[0]).days
        base = level[(nmi, sx)] * day_fx[dd]
        wob = 1.0 + ((int(_hex(nmi, sx, rd, n=4), 16) % 200) - 100) / 4000.0
        tot = base * wob * (1.0 + 0.004 * vi)
        uom = "KVARH" if sx == "Q1" else "KWH"
        ta = t.astimezone(W.AEST)
        did = "D%s.%s.%02d" % (ta.strftime("%y%m%d"), ta.strftime("%H%M"), (int(_hex(nmi, ta, n=4), 16) % 23) + 1)
        out.append((R.utc(t), nmi, sx, rd.strftime("%Y%m%d"), L, uom, did, a, s_, f, e, "%.3f" % tot))
    out.sort(key=lambda x: (x[0], x[1], x[2], x[3]))
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["NMI", "NMI_SUFFIX", "READ_DATE", "INTERVAL_LENGTH", "UOM", "DELIVERY_ID", "RECEIVED_UTC",
                    "QTY_A", "QTY_S", "QTY_F", "QTY_E", "DAY_TOTAL"])
        for (tu, nmi, sx, rd, L, uom, did, a, s_, f, e, tot) in out:
            w.writerow([nmi, sx, rd, L, uom, did, tu, a, s_, f, e, tot])
    return len(out)


# ------------------------------------------------------------------ head-end history

REPLAY_FROM = dt.datetime(2026, 10, 7, 1, 40, tzinfo=W.AEDT)
REPLAY_TO = dt.datetime(2026, 10, 7, 2, 55, tzinfo=W.AEDT)
FAILOVER_INGEST = dt.datetime(2026, 10, 7, 3, 12, tzinfo=W.AEDT)


def write_hes(path, insts, ev, bench):
    r = S.rng("hes")
    meta = {}
    for i in insts:
        for m in i.meters:
            meta[m[0]] = (m[2], m[1], "REMOVED" if m[4] is not None else "IN_SERVICE")
    for serial, (fam, _) in bench.items():
        meta[serial] = (fam, "M1PH-P5" if fam == "PX5" else "M1PH-E3B", "BENCH")
        ev[serial] = bench[serial][1]
    allev = []
    for serial in sorted(ev):
        for e in ev[serial]:
            allev.append((e["at"], serial, e))
    allev.sort(key=lambda x: (x[0], x[1]))
    seq = 1000000
    per_meter = {s: [] for s in ev}
    replays = []
    for at, serial, e in allev:
        seq += int(r.integers(1, 4))
        eid = "evt-" + _hex(serial, e["type"], at.isoformat(), n=12)
        ingest = at + dt.timedelta(seconds=int(r.integers(2, 90)))
        rec = {"seq": seq, "event_id": eid, "type": e["type"], "image": e["image"], "job": e["job"],
               "ts_utc": R.utc(at), "ingested_utc": R.utc(ingest)}
        per_meter[serial].append(rec)
        if REPLAY_FROM <= at <= REPLAY_TO:
            replays.append((serial, rec))
    for serial, rec in replays:
        seq += int(r.integers(1, 3))
        dup = dict(rec)
        dup["seq"] = seq
        dup["ingested_utc"] = R.utc(FAILOVER_INGEST + dt.timedelta(seconds=int(r.integers(0, 600))))
        per_meter[serial].append(dup)
    meters = []
    for serial in sorted(per_meter):
        fam, model, status = meta[serial]
        meters.append({"meter_id": "MTR-" + serial, "module": fam, "model": model, "status": status,
                       "events": sorted(per_meter[serial], key=lambda x: x["seq"])})
    doc = {"export": {"system": "HVM head-end (HES) firmware service", "exported_utc": R.utc(W.EXTRACT_AT - dt.timedelta(minutes=22)),
                      "scope": "every meter registered to the head-end, firmware events since commissioning",
                      "meters": len(meters)},
           "meters": meters}
    with open(path, "w") as fh:
        json.dump(doc, fh, separators=(",", ":"))
    return len(meters), sum(len(m["events"]) for m in meters), len(replays)


# ------------------------------------------------------------------ register workbook

def _local_str(t_naive):
    return t_naive.strftime("%d/%m/%Y %H:%M")


def write_register(path, insts):
    wb = openpyxl.Workbook()
    hdr = "Halvane Metering - MP installation register extract - as at 12/10/2026 06:00 (local time)"
    ws = wb.active
    ws.title = "Installations"
    ws.append([hdr])
    ws.append([])
    ws.append(["NMI", "Suburb", "Postcode", "Customer class", "Phases", "Solar", "Controlled load"])
    for i in sorted(insts, key=lambda x: x.nmi):
        suff, L, cls = W.PROFILES[i.profile]
        ws.append([i.nmi + S.nmi_checksum(i.nmi), i.suburb, int(i.postcode), cls,
                   3 if cls == "SME" else 1, "Y" if "B1" in suff else "N", "Y" if "E2" in suff else "N"])
    ws2 = wb.create_sheet("Meters")
    ws2.append([hdr])
    ws2.append([])
    ws2.append(["NMI", "Meter serial", "Model", "Comms module", "Installed", "Removed"])
    for i in sorted(insts, key=lambda x: x.nmi):
        for (serial, model, module, inst_t, rem_t) in i.meters:
            # serials migrated from the old CIS before 2024 came across as numbers
            cell = int(serial) if inst_t.year < 2024 else serial
            ws2.append([i.nmi + S.nmi_checksum(i.nmi), cell, model, module, _local_str(inst_t),
                        _local_str(rem_t) if rem_t else None])
    ws3 = wb.create_sheet("Data streams")
    ws3.append([hdr])
    ws3.append([])
    ws3.append(["NMI", "Suffix", "Interval (min)", "UOM", "Active from"])
    for i in sorted(insts, key=lambda x: x.nmi):
        suff, L, cls = W.PROFILES[i.profile]
        for sx in suff:
            ws3.append([i.nmi + S.nmi_checksum(i.nmi), sx, L, "kVArh" if sx == "Q1" else "kWh",
                        i.meters[0][3].strftime("%d/%m/%Y")])
    for w in (ws, ws2, ws3):
        w["A1"].font = Font(bold=True)
        for c in w[3]:
            c.font = Font(bold=True)
        w.freeze_panes = "A4"
    wb.properties.creator = "MP Registry"
    wb.properties.created = dt.datetime(2026, 10, 12, 6, 5)
    wb.properties.modified = dt.datetime(2026, 10, 12, 6, 5)
    wb.save(path)


# ------------------------------------------------------------------ catalogue

CAT_NOTES = {
    "EM3.9.6": "superseded", "EM3.10.2": "superseded",
    "EM3.10.9": "retiring; no new moves onto this image",
    "EM3.11.4": "", "EM3.12.0": "",
    "EM3.12.1": "hotfix: mesh rejoin timer after a collector restart; deployed to the Feb-26 hotfix list only",
    "EM3.12.2": "rollout target for EM3 fleet (waves 6 to 9)",
    "PX5.1.9": "superseded", "PX5.2.3": "superseded",
    "PX5.2.7": "", "PX5.3.0": "",
    "PX5.3.1": "pilot; wave 8 extends the pilot to the PX5.3.0 schedule list",
}


def write_catalogue(path):
    rel = [("EM3.9.6", "em3-fw-030906-b1874", "EM3"), ("EM3.10.2", "em3-fw-031002-b2011", "EM3")]
    rel += [(b[0], b[1], b[2]) for b in W.BUILDS if b[2] == "EM3"]
    rel += [("PX5.1.9", "px5-fw-050109-r0455", "PX5"), ("PX5.2.3", "px5-fw-050203-r0610", "PX5")]
    rel += [(b[0], b[1], b[2]) for b in W.BUILDS if b[2] == "PX5"]
    chan = {b[0]: b[3] for b in W.BUILDS}
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["release", "module", "hes_image", "vendor_build", "released", "channel", "notes"])
        for name, img, mod in rel:
            d = dt.date.fromisoformat(R.RELEASED[name])
            w.writerow([name, mod, img, img.split("-")[-1], d.strftime("%d/%m/%Y"), chan.get(name, "superseded"),
                        CAT_NOTES[name]])


# ------------------------------------------------------------------ Wave 8 schedule

DEPOTS = ["Seven Hills", "Wyong", "Unanderra", "Bathurst"]


def write_wave8(path, rows):
    r = S.rng("wave8_sheet")
    wb = openpyxl.Workbook()
    sheets = {"A": wb.active, "B": wb.create_sheet("Batch B 14-10")}
    sheets["A"].title = "Batch A 13-10"
    for b, ws in sheets.items():
        ws.append(["WAVE 8 - comms firmware - batch %s - night of %s October (AEDT)" % (b, "13/14" if b == "A" else "14/15")])
        ws.append(["Seq", "Meter No", "NMI", "Depot", "Current FW", "Target FW", "Window"])
    seq = {"A": 0, "B": 0}
    for row in rows:
        b = row["batch"]
        seq[b] += 1
        i = row["inst"]
        cur = row["frm"].replace(".", " ", 1)
        meter = int(row["serial"]) if b == "A" else row["serial"]
        nmi = "" if (b == "A" and r.random() < 0.31) else i.nmi
        sheets[b].append([seq[b], meter, nmi, DEPOTS[int(_hex(i.nmi, n=2), 16) % 4], cur, row["tgt"], "23:30-03:30"])
    for ws in sheets.values():
        ws["A1"].font = Font(bold=True)
        for c in ws[2]:
            c.font = Font(bold=True)
    n = wb.create_sheet("Notes")
    for line in ["Cut 01/10/2026 from the HES inventory (meters then on the source images).",
                 "Batch B rebuilt 09/10/2026 after the Wyong depot swap.",
                 "Windows are local time. Field support on call 23:00-04:00.",
                 "The HES job skips any meter already on the target image and any meter not in service at its installation on the night.",
                 "Go/no-go per target image comes from the read-health gate (MDS-STD-114) on Monday."]:
        n.append([line])
    wb.properties.creator = "Rollout Desk"
    wb.properties.created = dt.datetime(2026, 10, 1, 14, 20)
    wb.properties.modified = dt.datetime(2026, 10, 9, 16, 48)
    wb.save(path)


# ------------------------------------------------------------------ distractors

def write_census(path, insts, ev, bench):
    """HES inventory census at 00:00 AEDT 12 Oct: meters by reported image and HES status."""
    at = dt.datetime(2026, 10, 12, 0, 0, tzinfo=W.AEDT)
    counts = {}
    status = {}
    for i in insts:
        for m in i.meters:
            status[m[0]] = "REMOVED" if m[4] is not None else "IN_SERVICE"
    for s in bench:
        status[s] = "BENCH"
    allev = dict(ev)
    for s, (fam, e) in bench.items():
        allev[s] = e
    for serial, es in allev.items():
        img = None
        for e in sorted(es, key=lambda x: x["at"]):
            if e["at"] > at:
                break
            if e["type"] in ("COMMISSIONED", "FW_ACTIVATED"):
                img = e["image"]
        st = status[serial]
        counts[(img, st)] = counts.get((img, st), 0) + 1
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["census_at_local", "hes_image", "hes_status", "meters"])
        for (img, st) in sorted(counts):
            w.writerow(["12/10/2026 00:00", img, st, counts[(img, st)]])


def write_telco(path):
    r = S.rng("telco")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Notices"
    ws.append(["Carrier planned maintenance - IoT (LTE-M / NB-IoT) - notices received October 2026"])
    ws.append(["Notice", "Start (local)", "End (local)", "Area", "Service", "Expected impact", "Status"])
    areas = ["Penrith / St Marys", "Blacktown / Seven Hills", "Parramatta CBD", "Liverpool / Casula", "Gosford / Erina",
             "Newcastle West", "Wollongong / Unanderra", "Bathurst", "Orange", "Camden / Oran Park", "Hornsby",
             "Castle Hill / Rouse Hill", "Maitland", "Dubbo"]
    for k in range(23):
        day = dt.date(2026, 10, 1) + dt.timedelta(days=int(r.integers(0, 24)))
        st = dt.datetime(day.year, day.month, day.day, int(r.choice([0, 1, 2, 23])), int(r.choice([0, 30])))
        en = st + dt.timedelta(hours=int(r.integers(2, 5)))
        ws.append(["CMN-%06d" % (481200 + k * 7 + int(r.integers(0, 6))), st.strftime("%d/%m/%Y %H:%M"),
                   en.strftime("%d/%m/%Y %H:%M"), areas[int(r.integers(0, len(areas)))],
                   str(r.choice(["LTE-M", "LTE-M, NB-IoT", "NB-IoT"])),
                   str(r.choice(["Intermittent data sessions", "Data sessions may drop up to 20 min",
                                 "No impact expected", "Brief loss of service during cutover"])),
                   "Completed" if st < dt.datetime(2026, 10, 12, 7, 55) else "Scheduled"])
    ws["A1"].font = Font(bold=True)
    wb.properties.creator = "Network Ops"
    wb.properties.created = dt.datetime(2026, 10, 1, 8, 0)
    wb.properties.modified = dt.datetime(2026, 10, 12, 7, 55)
    wb.save(path)


# ------------------------------------------------------------------ provenance

def write_log(path, counts):
    rows = [
        (FILES["mdm"], "MDM (meter data management)", "12/10/2026 15:40", "every data stream HVM is MDP for; read dates 05/10/2026 to 11/10/2026; every delivery received up to the extract", counts["mdm"]),
        (FILES["hes"], "HES firmware service", "12/10/2026 15:18", "every meter registered to the head-end, all firmware events since commissioning", counts["hes_meters"]),
        (FILES["register"], "MP registry", "12/10/2026 06:05", "installations, meters (with removal times) and data streams", counts["register"]),
        (FILES["catalogue"], "Firmware engineering", "09/10/2026 11:30", "releases and the HES image each one ships as", counts["catalogue"]),
        (FILES["wave8"], "Rollout desk", "09/10/2026 16:48", "Wave 8 schedule, both batches", counts["wave8"]),
        (FILES["census"], "HES reporting", "12/10/2026 00:05", "nightly inventory census by image and status", counts["census"]),
        (FILES["telco"], "Network operations", "12/10/2026 07:55", "carrier maintenance notices received in October", 23),
        (FILES["standard"], "Metering Data Services", "", "controlled document, revision 3", ""),
    ]
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["file", "source_system", "extracted_local", "coverage", "records"])
        for row in rows:
            w.writerow(row)
