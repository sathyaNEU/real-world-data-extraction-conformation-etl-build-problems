"""task229 generator: the golden deliverables and figures.json, written from the checks' record.

discharge_grades_2027.png: the 22 stations against the months of 2027, every month short of grade B shaded.
record_review_wy2026.xlsx: the 2026 water year's measured station-days by month, and each station's mean gauging
difference from the rating in force on the gauging day."""
import json
import math
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Patch, Rectangle
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import params as P
import postprocess
import world as W
import golden as GD

ORG = "Corriston Basin Hydrometric Cooperative"
ORDER = ["C104", "C105", "C107", "C101", "C102", "C103", "C106", "C108", "C109", "C111", "C112", "C113", "C114",
         "C115", "C116", "C117", "C118", "C119", "C120", "C121", "C122", "C123"]
MON = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]


def pct1(x):
    s = "%.1f" % x
    return s.replace("-", "−") if False else s


def chart(rec, path):
    short = rec["call"]["short"]
    order = list(W.CURRENT)
    fig, ax = plt.subplots(figsize=(10.5, 8.4), dpi=150)
    reg_c, ext_c, ok_c = "#b03a2e", "#e59866", "#f2f4f5"
    for i, no in enumerate(order):
        y = len(order) - 1 - i
        for m in range(1, 13):
            if m in short[no]:
                col = ext_c if no not in W.PLANT_OF else reg_c
            else:
                col = ok_c
            ax.add_patch(Rectangle((m - 0.5, y - 0.45), 0.96, 0.9, facecolor=col, edgecolor="white", linewidth=1.0))
        first = rec["call"]["first"][no]
        if first:
            ax.text(12.75, y, "flag from %s" % MON[first - 1], va="center", ha="left", fontsize=8.5, color="#222222")
    ax.set_xlim(0.4, 14.9)
    ax.set_ylim(-0.7, len(order) - 0.3)
    ax.set_yticks(range(len(order)))
    ax.set_yticklabels(["%s  %s" % (no, W.ST[no]["name"]) for no in reversed(order)], fontsize=8)
    ax.set_xticks(range(1, 13))
    ax.set_xticklabels([m[:3] for m in MON], fontsize=8.5)
    ax.xaxis.tick_top()
    ax.tick_params(length=0)
    for s in ("top", "right", "left", "bottom"):
        ax.spines[s].set_visible(False)
    fig.text(0.012, 0.965, "Five stations fall short of grade B in 2027", fontsize=13, fontweight="bold", color="#1b2330")
    fig.text(0.012, 0.938, "The two gauges below Talbot Canyon from January, Sable Creek in the freshet, and the two gauges below "
             "Merrin Falls from July.", fontsize=9, color="#333333")
    fig.text(0.012, 0.917, "Published daily discharge by station-month of 2027, graded under the hydrometric data standard, "
             "sections 3 and 4. Shaded: short of grade B.", fontsize=8.5, color="#555555")
    ax.legend(handles=[Patch(facecolor=reg_c, label="published more than 5 per cent below the month's mean flow"),
                       Patch(facecolor=ext_c, label="month's flow above one and a half times the largest gauging"),
                       Patch(facecolor=ok_c, edgecolor="#cccccc", label="meets grade B")],
              loc="upper left", bbox_to_anchor=(-0.415, -0.012), ncol=3, fontsize=7.8, frameon=False, handlelength=1.4)
    fig.text(0.012, 0.012, "Sources: 2027 water supply outlook (2 November 2026); Corriston Hydro Ltd. 2027 operating plan; "
             "rating book, issue 2. Corriston Basin Hydrometric Cooperative.", fontsize=7.5, color="#555555")
    fig.subplots_adjust(left=0.30, right=0.985, top=0.86, bottom=0.08)
    fig.savefig(path, dpi=150, metadata={"Software": None})
    plt.close(fig)


def workbook(rec, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Measured station-days"
    ws["A1"] = "2026 water year: station-days with a water level measured on the day"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = "22 stations; Mattin Creek counted once across its June 2026 move (C110 to C122)."
    for j, h in enumerate(["Month", "Station-days"], start=1):
        c = ws.cell(row=4, column=j, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDE4EE")
    r = 5
    for lab in GD.MLAB:
        ws.cell(row=r, column=1, value=lab)
        c = ws.cell(row=r, column=2, value=int(rec["A"][lab]))
        c.number_format = "#,##0"
        r += 1
    ws.cell(row=r, column=1, value="Water year").font = Font(bold=True)
    c = ws.cell(row=r, column=2, value=int(sum(rec["A"].values())))
    c.number_format = "#,##0"
    c.font = Font(bold=True)
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 14
    ws.freeze_panes = "A5"
    w2 = wb.create_sheet("Gauging agreement")
    w2["A1"] = "2026 water year: mean of each station's gauging percentage differences from the rating in force on the gauging day"
    w2["A1"].font = Font(bold=True, size=12)
    w2["A2"] = "A gauging's difference is the measured discharge less the rating's, as a percentage of the rating's (standard, 5.3); a section gauged in two channels counts as one gauging."
    hdr = ["Station", "Name", "Gaugings", "Mean difference (per cent)"]
    for j, h in enumerate(hdr, start=1):
        c = w2.cell(row=4, column=j, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDE4EE")
    r = 5
    for no in W.CURRENT:
        w2.cell(row=r, column=1, value=no)
        w2.cell(row=r, column=2, value=W.ST[no]["name"])
        w2.cell(row=r, column=3, value=int(rec["B_n"][no]))
        c = w2.cell(row=r, column=4, value=round(rec["B"][no], 1))
        c.number_format = "0.0"
        r += 1
    for col, wd in zip("ABCD", (9, 36, 10, 20)):
        w2.column_dimensions[col].width = wd
    w2.freeze_panes = "A5"
    w3 = wb.create_sheet("Notes")
    notes = [
        "Notes",
        "Daily record: daily_records_wy2023_wy2026.csv as exported 6 November 2026; October 2025 to June 2026 approved, July to September 2026 provisional.",
        "A day the GOES ingest filled with the platform's previous message (received before the day it covers had ended) holds no water level measured on that day.",
        "Gaugings before 1 April 2026 are compared with the 2019 review's ratings, which the record's published discharges before that date reproduce; from 1 April with the rating book, issue 2.",
        "Pellam and Garrow creeks are gauged in litres a second and converted. Days and gauging dates in local standard time.",
    ]
    for i, t in enumerate(notes, start=1):
        c = w3.cell(row=i, column=1, value=t)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 1:
            c.font = Font(bold=True)
    w3.column_dimensions["A"].width = 120
    wb.save(path)
    postprocess.scrub(path, "Phyllis Thompson", "2026-11-07")
    postprocess.finish_office(path, "2026-11-06T18:05:00Z", "2026-11-07T19:40:00Z", ORG)


def figures(rec):
    first = rec["call"]["first"]
    errs = rec["call"]["errs"]
    fail = [(no, m, errs[no][str(m)] if isinstance(next(iter(errs[no])), str) else errs[no][m]) for no in W.PLANT_OF for m in rec["call"]["short"][no]]
    least = max(fail, key=lambda x: x[2])
    most = min(fail, key=lambda x: x[2])
    rng_lo = round(-least[2] * 100)
    rng_hi = round(-most[2] * 100)
    tops = rec["call"]["tops"]
    fig = {
        "call": {no: (MON[first[no] - 1] if first[no] else None) for no in W.CURRENT},
        "call_first": {"C104": "January", "C105": "January", "C107": "May", "C112": "July", "C113": "July"},
        "holding_stations": 18,
        "stations": 23,
        "c124_max_pct": math.ceil(max(abs(rec["call"]["errs"]["C124"][m]) for m in range(1, 13)) * 1000) / 10,
        "bias_range_pct": [rng_lo, rng_hi],
        "bias_least": {"station": least[0], "month": MON[least[1] - 1], "pct": round(-least[2] * 100, 1)},
        "sable": {"may": W.outlook_2027("C107", 5), "june": W.outlook_2027("C107", 6), "top": tops["C107"]},
        "grade_b_pct": 5,
        "short_months": {no: [MON[m - 1] for m in rec["call"]["short"][no]] for no in W.CURRENT if rec["call"]["short"][no]},
        "ask_a": {k: int(v) for k, v in rec["A"].items()},
        "ask_a_total": int(sum(rec["A"].values())),
        "ask_b": {no: round(rec["B"][no], 1) for no in W.CURRENT},
        "ask_b_n": rec["B_n"],
        "windows": "06:00 to 10:00 and 16:00 to 21:00",
        "min_release": {"talbot": W.TALBOT["fmin"], "merrin": W.MERRIN["fmin"]},
        "ramp_min": 30,
        "corridor_worst_pct": round(rec["corridor_worst"] * 100, 1),
        "flip_damp_pct": round(rec["flip_damp_C105_jan"] * 100),
    }
    return fig


def write(rec, gdir, figpath):
    os.makedirs(gdir, exist_ok=True)
    chart(rec, os.path.join(gdir, "discharge_grades_2027.png"))
    workbook(rec, os.path.join(gdir, "record_review_wy2026.xlsx"))
    fig = figures(rec)
    with open(figpath, "w") as fh:
        json.dump(fig, fh, indent=1, ensure_ascii=False, sort_keys=True)
        fh.write("\n")
    return rec, fig
