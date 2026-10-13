"""task229 generator: deterministic writers for the data files under target/."""
from datetime import date, datetime, timedelta

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import params as P
import postprocess
import world as W

ORG = "Corriston Basin Hydrometric Cooperative"


def csv(df, path):
    df.to_csv(path, index=False, lineterminator="\n")


def _xlsx_save(wb, path, who, created, modified):
    wb.save(path)
    postprocess.scrub(path, who, modified[:10])
    postprocess.finish_office(path, created, modified, ORG)


def _hdr(ws, cols, row):
    for j, c in enumerate(cols, start=1):
        cell = ws.cell(row=row, column=j, value=c)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDE4EE")
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def _widths(ws, widths):
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w


# ------------------------------------------------------------------------------------------------ daily record
def daily_record(df, path):
    out = df.copy()
    out["date"] = [d.isoformat() for d in out.date]
    out["water_level_m"] = [("%.2f" % v) if W.ST[s]["res"] == 0.01 else ("%.3f" % v) for s, v in zip(out.station, out.water_level_m)]
    out["discharge_m3s"] = [_g4(v) for v in out.discharge_m3s]
    csv(out, path)


def _g4(v):
    s = "%.4g" % v
    if "e" in s:
        s = "%.6f" % v
    return s


# ------------------------------------------------------------------------------------------------ gaugings
def gaugings(G, path):
    out = G[["gauging_no", "station", "measured_utc", "gh_m", "discharge", "unit", "method", "rating", "crew", "notes"]].copy()
    out["gh_m"] = ["%.3f" % v for v in out.gh_m]
    out["discharge"] = ["%.3g" % v for v in out.discharge]
    csv(out, path)


# ------------------------------------------------------------------------------------------------ the register
SITE_NOTES = {
    "C101": "gravel riffle control; wire-weight gauge on the road bridge",
    "C102": "lake level at the outlet weir; stop logs set for the summer level; float and shaft encoder",
    "C103": "boulder control; staff gauge on the left bank",
    "C104": "bedrock canyon reach; staff gauge and pressure transducer on the left wall",
    "C105": "bedrock control under the bridge; staff gauge on the downstream pier",
    "C106": "culvert outlet pool; staff gauge on the wingwall",
    "C107": "no bridge or cableway; gauged by wading",
    "C108": "wide gravel-bed section; staff gauge on the right bank",
    "C109": "concrete sill below the powerhouse tailrace",
    "C110": "gravel pit reach; site closed when the station moved to Quarry Road",
    "C111": "spring pool outlet; float and shaft encoder",
    "C112": "bedrock reach below the falls; staff gauge on the right bank",
    "C113": "riffle control upstream of the road crossing",
    "C114": "glacial silt; boulder control",
    "C115": "spring pool outlet; float and shaft encoder",
    "C116": "boulder cascade above the municipal intake",
    "C117": "clay bank, gravel bar control",
    "C118": "log jam upstream; staff gauge on the left bank",
    "C119": "wide gravel-bed section at the old bridge piers",
    "C120": "rock ledge control",
    "C121": "fish fence apron control",
    "C122": "bedrock ledge control, 300 m upstream of the old gravel pit site",
    "C123": "canyon mouth; boulder control",
    "C124": "gravel bar control below the Vardy River confluence; cableway",
}


def register(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Stations"
    cols = ["station", "name", "river", "river_km", "drainage_area_km2", "established", "status", "closed", "replaced_by",
            "telemetry", "on_platform_since", "goes_slot_utc", "wl_resolution_m", "field_unit", "site_notes"]
    _hdr(ws, cols, 1)
    r = 2
    for no in sorted(W.ST):
        st = W.ST[no]
        closed = no == "C110"
        vals = [no, st["name"], st["river"], st["km"], st["area"], W.ESTABLISHED[no], "closed" if closed else "active",
                P.RELOC_END if closed else None, "C122" if closed else None,
                None if closed else st["new_platform"], None if closed else st["cutover"],
                st["slot"] or None, st["res"], st["unit"], SITE_NOTES[no]]
        if closed:
            vals[9], vals[10] = st["new_platform"], st["cutover"]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v)
            if isinstance(v, date):
                c.number_format = "yyyy-mm-dd"
        r += 1
    ws.freeze_panes = "B2"
    _widths(ws, [9, 34, 16, 9, 11, 10, 8, 11, 11, 10, 13, 11, 11, 9, 60])
    _xlsx_save(wb, path, "Phyllis Thompson", "2019-03-04T17:20:00Z", "2026-10-02T16:05:00Z")


# ------------------------------------------------------------------------------------------------ the outlook
def outlook(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Outlook 2027"
    ws["A1"] = "2027 water supply outlook: monthly mean flow, most probable"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = "Corriston Basin Hydrometric Cooperative, hydrometric section. Issued 2 November 2026. Teresa Chen, hydrologist."
    ws["A3"] = "Cubic metres a second. Rows TAL-FB and MER-HP are the inflows to the Talbot Canyon forebay and the Merrin Falls headpond."
    cols = ["station", "name"] + W.M
    _hdr(ws, cols, 5)
    r = 6
    for no in W.CURRENT:
        ws.cell(row=r, column=1, value=no)
        ws.cell(row=r, column=2, value=W.ST[no]["name"])
        for m in range(1, 13):
            c = ws.cell(row=r, column=2 + m, value=W.outlook_2027(no, m))
            c.number_format = "0.000" if W.outlook_2027(no, m) < 0.1 else ("0.00" if W.outlook_2027(no, m) < 10 else "0.0")
        r += 1
    for code, name, ser in (("TAL-FB", "Talbot Canyon forebay inflow", W.TALBOT_2027), ("MER-HP", "Merrin Falls headpond inflow", W.MERRIN_2027)):
        ws.cell(row=r, column=1, value=code)
        ws.cell(row=r, column=2, value=name)
        for m in range(1, 13):
            c = ws.cell(row=r, column=2 + m, value=float(ser[m - 1]))
            c.number_format = "0.00" if ser[m - 1] < 10 else "0.0"
        r += 1
    ws.freeze_panes = "C6"
    _widths(ws, [9, 34] + [8] * 12)
    m2 = wb.create_sheet("Method")
    lines = [
        "Method",
        "Monthly volumes from the 1 November snow courses, the seasonal climate outlook and each basin's regression on its index stations, spread over the months by the station's recession and melt curves.",
        "Stations below Talbot Canyon and Merrin Falls carry the forebay or headpond inflow plus the inflow from the drainage between the plant and the station, the Vardy River included at Merrin Creek near the mouth.",
        "Most probable: the median of the outlook's ensemble. The 10 and 90 per cent exceedance flows are issued to the water licence holders separately.",
        "Next issue: 1 February 2027, after the 1 February snow survey.",
    ]
    for i, t in enumerate(lines, start=1):
        c = m2.cell(row=i, column=1, value=t)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 1:
            c.font = Font(bold=True)
    m2.column_dimensions["A"].width = 120
    _xlsx_save(wb, path, "Teresa Chen", "2026-10-26T18:00:00Z", "2026-11-02T17:35:00Z")


# ------------------------------------------------------------------------------------------------ migration tracker
def migration(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Cut-overs"
    cols = ["station", "name", "from", "to", "installed", "first_day_new", "goes_removed", "notes"]
    _hdr(ws, cols, 1)
    r = 2
    for no in sorted(W.ST):
        if no == "C122":
            continue
        st = W.ST[no]
        cut = st["cutover"]
        note = {"C104": "no cellular coverage in the canyon", "C105": "no cellular coverage at the bridge",
                "C112": "no cellular coverage in the gorge", "C113": "no cellular coverage in the gorge",
                "C110": "site closed 8 June 2026; Quarry Road (C122) opened on cellular 4 June"}.get(no, "")
        vals = [no, st["name"], "GOES", st["new_platform"], cut - timedelta(days=W.install_lead(no, cut)), cut, cut + timedelta(days=1), note]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v)
            if isinstance(v, date):
                c.number_format = "yyyy-mm-dd"
        r += 1
    ws.freeze_panes = "A2"
    _widths(ws, [9, 34, 7, 9, 11, 13, 12, 46])
    _xlsx_save(wb, path, "Phyllis Thompson", "2026-01-12T19:30:00Z", "2026-09-29T21:10:00Z")


# ------------------------------------------------------------------------------------------------ visit log
def visit_log(rows, path):
    df = pd.DataFrame(rows, columns=["date", "station", "crew", "work"])
    df = df.sort_values(["date", "station", "work"], kind="mergesort")
    df["date"] = [d.isoformat() for d in df.date]
    csv(df, path)


# ------------------------------------------------------------------------------------------------ dictionary
DICT = [
    ("daily_records_wy2023_wy2026.csv", "station", "station number"),
    ("daily_records_wy2023_wy2026.csv", "date", "the day the row's water level covers: the UTC day for GOES platforms, the local standard-time day for cellular and Iridium platforms"),
    ("daily_records_wy2023_wy2026.csv", "water_level_m", "the day's mean water level, the mean of the platform's readings over the day, in metres above the station's gauge datum, to the station's resolution"),
    ("daily_records_wy2023_wy2026.csv", "discharge_m3s", "the day's discharge, from water_level_m through the station's rating in force on the date, in cubic metres a second"),
    ("daily_records_wy2023_wy2026.csv", "status", "provisional as received; approved once the annual review has passed the day"),
    ("daily_records_wy2023_wy2026.csv", "platform", "the telemetry platform that sent the row: GOES, cellular or Iridium"),
    ("daily_records_wy2023_wy2026.csv", "received_utc", "when the network's data system received the row"),
    ("gaugings_2016_2026.csv", "gauging_no", "measurement number"),
    ("gaugings_2016_2026.csv", "station", "station number"),
    ("gaugings_2016_2026.csv", "measured_utc", "mid-time of the measurement, UTC"),
    ("gaugings_2016_2026.csv", "gh_m", "gauge height at the measurement, metres above the station's gauge datum"),
    ("gaugings_2016_2026.csv", "discharge", "measured discharge, in the unit column's unit"),
    ("gaugings_2016_2026.csv", "unit", "unit of the discharge column"),
    ("gaugings_2016_2026.csv", "method", "ADCP, current meter or volumetric"),
    ("gaugings_2016_2026.csv", "rating", "rating number"),
    ("gaugings_2016_2026.csv", "crew", "field crew"),
    ("gaugings_2016_2026.csv", "notes", "the crew's field notes"),
    ("station_register.xlsx", "river_km", "distance from the river's mouth along the channel"),
    ("station_register.xlsx", "replaced_by", "for a closed station, the number its record continues under: the closed number's daily rows and gaugings are that station's"),
    ("station_register.xlsx", "telemetry", "the station's platform since its 2026 cut-over"),
    ("station_register.xlsx", "on_platform_since", "first day on the current platform"),
    ("station_register.xlsx", "goes_slot_utc", "the GOES transmission slot the station held before its cut-over"),
    ("station_register.xlsx", "wl_resolution_m", "resolution of the station's water level"),
    ("station_register.xlsx", "field_unit", "the unit the field software records the station's gaugings in"),
    ("telemetry_migration_2026.xlsx", "first_day_new", "the first day the new platform reports; the GOES platform reports that day too and is removed the next day"),
    ("water_supply_outlook_2027.xlsx", "Jan to Dec", "monthly mean flow in 2027, cubic metres a second, most probable"),
    ("field_visits_wy2026.csv", "work", "what the crew did on the visit"),
    ("water_quality_grab_samples_2026.csv", "water_temperature_degc", "water temperature at the sample, degrees C"),
    ("water_quality_grab_samples_2026.csv", "turbidity_ntu", "nephelometric turbidity"),
    ("water_quality_grab_samples_2026.csv", "conductivity_us_cm", "specific conductance at 25 degrees C"),
]


def dictionary(path):
    csv(pd.DataFrame(DICT, columns=["file", "column", "description"]), path)


# ------------------------------------------------------------------------------------------------ distractor
def water_quality(rng, TR, path):
    import truth as T
    rows = []
    sites = ["C101", "C103", "C106", "C107", "C109", "C114", "C116", "C117", "C118", "C121"]
    for no in sites:
        for k in range(12):
            y, m = (2025, 10 + k) if k < 3 else (2026, k - 2)
            d = date(y, m, int(rng.integers(5, 26)))
            hh = int(rng.integers(9, 15))
            i = (d - T.T0).days
            q = float(TR["S"][no][i * T.SPD + hh * 12])
            temp = 3.0 + 9.0 * np.sin(np.pi * ((m - 3) % 12) / 12.0) ** 2 + rng.normal(0, 0.6)
            cond = 180.0 / (1 + 0.6 * np.log1p(q)) * (1 + rng.normal(0, 0.05))
            turb = max(0.3, 0.8 * q ** 0.6 * np.exp(rng.normal(0, 0.4)))
            rows.append([d.isoformat(), no, "%02d:%02d" % (hh, int(rng.integers(0, 60))), round(temp, 1),
                         int(round(cond)), round(turb, 1), round(7.2 + rng.normal(0, 0.25), 2)])
    df = pd.DataFrame(rows, columns=["date", "station", "time_local", "water_temperature_degc", "conductivity_us_cm", "turbidity_ntu", "ph"])
    csv(df.sort_values(["date", "station"], kind="mergesort"), path)


def manifest(path, names, desc, dated):
    rows = [(n, desc[n][0], desc[n][1], dated[n]) for n in names]
    csv(pd.DataFrame(rows, columns=["file", "contents", "made_by", "dated"]), path)
