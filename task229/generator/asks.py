"""The ask paths: Thornacre's line catalogue and its assignment to faces, the stock snapshot and batch
register, the line master, the base cost list and the period-10 change file (ask A); the store list, the
servicing map, the trading calendar, the category table and the daily forecast (ask B); the two
distractors. Every device and hazard lives here, never on the main path."""
import csv
import datetime as dt
import json
import os

import numpy as np

import world as W

D = dt.date
REGIME_CODE = {"Ambient": "AMB", "Chilled": "CHL", "Frozen": "FRZ"}

CATEGORIES = [
    # code, name, regime, forecast unit, units per case (forecast conversion), share of the regime's lines
    ("A01", "Grocery tins, jars and packets", "Ambient", "CS", None, 0.38),
    ("A02", "Soft drinks and mixers", "Ambient", "CS", None, 0.14),
    ("A03", "Household and paper", "Ambient", "CS", None, 0.18),
    ("A04", "Health and beauty singles", "Ambient", "EA", 12, 0.18),
    ("A05", "Beers, wines and spirits singles", "Ambient", "EA", 6, 0.12),
    ("C01", "Dairy, eggs and spreads", "Chilled", "CS", None, 0.30),
    ("C02", "Meat, fish and poultry", "Chilled", "CS", None, 0.28),
    ("C03", "Ready meals, deli and pizza", "Chilled", "CS", None, 0.27),
    ("C04", "Cut flowers and plants", "Chilled", "EA", 8, 0.15),
    ("F01", "Frozen vegetables, chips and potato", "Frozen", "CS", None, 0.40),
    ("F02", "Ice cream and desserts", "Frozen", "CS", None, 0.38),
    ("F03", "Frozen party food singles", "Frozen", "EA", 10, 0.22),
]
CAT = {c[0]: c for c in CATEGORIES}
N_LINES = {"Ambient": 4700, "Chilled": 4450, "Frozen": 3900}
SECOND_FACE = {"Thurrock": 0.11}   # share of A faces that are a fast line's second face
DATE_CODED = {"Ambient": 0.14, "Chilled": 1.0, "Frozen": 0.45}

WORDS = {
    "A01": ["Chopped Tomatoes", "Baked Beans", "Basmati Rice", "Penne Pasta", "Plain Flour", "Tuna Chunks", "Peanut Butter",
            "Strawberry Jam", "Porridge Oats", "Cornflakes", "Tomato Ketchup", "Mayonnaise", "Chicken Soup", "Coconut Milk",
            "Garden Peas", "Sweetcorn", "Digestive Biscuits", "Ginger Nuts", "Teabags", "Instant Coffee", "Gravy Granules"],
    "A02": ["Cola", "Diet Cola", "Lemonade", "Orange Squash", "Tonic Water", "Sparkling Water", "Still Water", "Ginger Beer",
            "Apple Juice", "Energy Drink"],
    "A03": ["Toilet Tissue", "Kitchen Roll", "Washing Up Liquid", "Bio Liquid", "Fabric Conditioner", "Bin Bags",
            "Dishwasher Tabs", "Bleach", "Foil", "Cling Film", "Surface Spray"],
    "A04": ["Shampoo", "Conditioner", "Shower Gel", "Toothpaste", "Deodorant", "Hand Wash", "Razors", "Paracetamol",
            "Ibuprofen", "Lip Balm", "Sun Cream", "Moisturiser"],
    "A05": ["Prosecco", "Sauvignon Blanc", "Merlot", "Rioja", "London Dry Gin", "Blended Scotch", "Spiced Rum", "Vodka",
            "Port", "Sherry", "Cream Liqueur", "Champagne"],
    "C01": ["Whole Milk", "Semi Skimmed Milk", "Mature Cheddar", "Greek Yoghurt", "Free Range Eggs", "Salted Butter",
            "Single Cream", "Double Cream", "Spreadable", "Halloumi", "Mozzarella"],
    "C02": ["Chicken Breast Fillets", "Beef Mince", "Pork Sausages", "Smoked Bacon", "Salmon Fillets", "Turkey Crown",
            "Lamb Leg", "Cooked Ham", "Cod Fillets", "Pigs In Blankets", "Gammon Joint"],
    "C03": ["Lasagne", "Chicken Tikka Masala", "Cottage Pie", "Margherita Pizza", "Houmous", "Coleslaw", "Sushi Box",
            "Fresh Soup", "Quiche", "Sandwich Selection", "Scotch Eggs"],
    "C04": ["Roses", "Tulips", "Mixed Bouquet", "Poinsettia", "Lilies", "Chrysanthemums", "Christmas Wreath", "Orchid"],
    "F01": ["Garden Peas", "Oven Chips", "Roast Potatoes", "Mixed Vegetables", "Sweetcorn", "Hash Browns", "Yorkshire Puddings",
            "Spinach", "Broccoli Florets"],
    "F02": ["Vanilla Ice Cream", "Choc Ices", "Cheesecake", "Arctic Roll", "Sorbet", "Lollies", "Profiteroles", "Trifle"],
    "F03": ["Prawn Rings", "Party Platter", "Mini Quiches", "Spring Rolls", "Vol Au Vents", "Chicken Bites", "Mince Pies"],
}
SIZES = {"A01": ["400g", "500g", "1kg", "250g", "340g"], "A02": ["2L", "6x330ml", "1L", "12x330ml"],
         "A03": ["9 Roll", "4 Roll", "500ml", "1.3L", "20 Pack"], "A04": ["250ml", "400ml", "100ml", "16 Pack", "75ml"],
         "A05": ["75cl", "70cl", "1L", "35cl"], "C01": ["2 Pints", "4 Pints", "400g", "500g", "6 Pack", "300ml"],
         "C02": ["500g", "1kg", "400g", "300g", "2.4kg"], "C03": ["400g", "450g", "200g", "300g", "6 Pack"],
         "C04": ["Bunch", "Single", "Large", "12 Stem"], "F01": ["1kg", "1.5kg", "907g", "750g"],
         "F02": ["1L", "4 Pack", "500ml", "6 Pack"], "F03": ["Box 20", "Box 12", "Tray 30", "Box 24"]}


def catalogue():
    """Thornacre lines by regime: category, stock unit, case-size history, date coding, cost per case."""
    rng = np.random.default_rng(W.ASK_SEED)
    pool = rng.permutation(np.arange(203000, 899999, 7))
    lines, k = [], 0
    for regime, n in N_LINES.items():
        cats = [c for c in CATEGORIES if c[2] == regime]
        p = np.array([c[5] for c in cats]); p = p / p.sum()
        pick = rng.choice(len(cats), size=n, p=p)
        for i in range(n):
            c = cats[pick[i]]
            code = c[0]
            ea = c[3] == "EA"
            size = int(rng.choice([4, 6, 6, 8, 10, 12, 12, 16, 20, 24])) if not ea else int(rng.choice([6, 6, 8, 10, 12, 12, 18, 24]))
            hist = []
            if rng.random() < (0.30 if ea else 0.06):
                ch = D(2026, 6, 20) + dt.timedelta(days=int(rng.integers(0, 150)))
                old = size + int(rng.choice([2, 2, 4, 6]))
                old = max(old, 3)
                if old == size:
                    old = size + 2
                start = D(2023, 1, 9) + dt.timedelta(days=int(rng.integers(0, 900)))
                hist = [(old, start, ch - dt.timedelta(days=1)), (size, ch, None)]
            else:
                start = D(2022, 3, 7) + dt.timedelta(days=int(rng.integers(0, 1500)))
                hist = [(size, start, None)]
            if ea:
                cost = round(float(rng.lognormal(np.log(3.2), 0.5)) * size, 2)
            else:
                cost = round(float(rng.lognormal(np.log(14.0), 0.55)), 2)
            words = WORDS[code]
            desc = "TH %s %s" % (words[int(rng.integers(len(words)))], SIZES[code][int(rng.integers(len(SIZES[code])))])
            lines.append(dict(line_no=str(int(pool[k])), category=code, regime=regime, unit="EA" if ea else "CS",
                              hist=hist, case_size=size, date_coded=bool(rng.random() < DATE_CODED[regime]),
                              cost=cost, desc=desc))
            k += 1
    # period-10 cost changes, effective 6 December (mostly increases)
    for ln in lines:
        if rng.random() < 0.19:
            up = rng.random() < 0.86
            f = 1 + (rng.uniform(0.025, 0.09) if up else -rng.uniform(0.02, 0.06))
            ln["p10"] = round(ln["cost"] * f, 2)
            if ln["p10"] == ln["cost"]:
                ln["p10"] = round(ln["cost"] + 0.05, 2)
        else:
            ln["p10"] = None
    return lines


def assign(sim, lines):
    """Each simulated face gets a location code and a Thornacre line; some fast lines take a second face."""
    rng = np.random.default_rng(W.ASK_SEED + 1)
    by_regime = {}
    for ln in lines:
        by_regime.setdefault(ln["regime"], []).append(ln)
    out = {}
    for c in W.CENTRES:
        faces, _ = sim[c["name"]]
        n = len(faces["cls"])
        pool = by_regime[c["regime"]]
        order = rng.permutation(len(pool))
        line_idx = np.full(n, -1)
        a_faces = np.where(faces["cls"] == "A")[0]
        second = set(rng.choice(a_faces, size=int(round(SECOND_FACE.get(c["name"], 0.075) * len(a_faces))), replace=False).tolist())
        firsts = [i for i in range(n) if i not in second]
        for j, i in enumerate(firsts):
            line_idx[i] = order[j]
        a_first = [i for i in a_faces if i not in second]
        for i in sorted(second):
            line_idx[i] = line_idx[a_first[int(rng.integers(len(a_first)))]]
        # location codes: pick levels 1 and 2 of the pick aisles, shuffled
        aisles = int(np.ceil(n / 110.0)) + 1
        codes = ["%02d-%03d-%d" % (a, b, lv) for a in range(11, 11 + aisles) for b in range(1, 56) for lv in (1, 2)]
        pick = rng.choice(len(codes), size=n, replace=False)
        out[c["name"]] = dict(line=line_idx, loc=[codes[i] for i in pick], aisles=aisles)
    return out


FILES = {
    "snapshot": "WMS_stock_by_location_THN_20261207_0600.csv",
    "batches": "WMS_batch_register_THN_20261207_0600.csv",
    "linemaster": "Thornacre_line_master_20261204.xlsx",
    "costbase": "THN_cost_price_list_20261128.csv",
    "costp10": "THN_cost_changes_P10_2026.csv",
    "forecast": "THN_store_replen_forecast_run_20261123.csv",
    "servicing": "THN_store_servicing_map.csv",
    "calendar": "THN_store_trading_calendar_2026_27.csv",
    "stores": "THN_store_list.json",
    "categories": "THN_category_table.csv",
    "labour": "Corderoy_peak_labour_plan_2026.xlsx",
    "goodsin": "goods_in_bookings_summary_7Dec_3Jan.csv",
}
DISTRACTORS = ["labour", "goodsin"]
SNAP_TS = "2026-12-07T06:00:00Z"
P10_FROM = D(2026, 12, 6)
VAL_DATE = D(2026, 12, 7)


def write_csv(path, cols, rows, eol="\n"):
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh, lineterminator=eol)
        w.writerow(cols)
        for r in rows:
            w.writerow([r[k] for k in cols] if isinstance(r, dict) else r)


def batch_code(rng):
    return "L%05d%s" % (int(rng.integers(10000, 99999)), "ABCDEFGHJK"[int(rng.integers(10))])


def stock(sim, assign, lines, loc_rows, salt=0):
    """Stock on every Thornacre location at 06:00 on 7 December, batch by batch (private structure)."""
    rng = np.random.default_rng(W.ASK_SEED + 11 + salt)
    by_no = {ln["line_no"]: ln for ln in lines}
    cur = {}
    for r in loc_rows:
        if r["client"] != "THN" or not r["line_no"]:
            continue
        if r["eff_from"] <= "2026-12-07" and (r["eff_to"] == "" or r["eff_to"] >= "2026-12-07"):
            cur[(r["site"], r["location"])] = r
    recs = []
    for (site, loc), r in sorted(cur.items()):
        ln = by_no[r["line_no"]]
        pick = r["location_type"] == "PICK"
        if pick and rng.random() < 0.06:
            continue                              # an empty face has no stock record
        if ln["unit"] == "EA":
            qty = int(rng.integers(6, 150 if pick else 400))
        else:
            qty = int(rng.integers(2, 46 if pick else 120))
        nb = 1
        if ln["date_coded"]:
            nb = int(rng.choice([1, 2, 3], p=[0.52, 0.36, 0.12]))
            nb = min(nb, qty)
        cuts = np.sort(rng.choice(np.arange(1, qty), size=nb - 1, replace=False)) if nb > 1 else np.array([], int)
        parts = np.diff(np.r_[0, cuts, qty]).astype(int)
        horizon = {"Chilled": (3, 14), "Frozen": (60, 400), "Ambient": (90, 500)}[ln["regime"]]
        bb = sorted(VAL_DATE + dt.timedelta(days=int(rng.integers(*horizon))) for _ in range(nb))
        bat = [(batch_code(rng) if ln["date_coded"] else "", bb[i].isoformat() if ln["date_coded"] else "", int(parts[i]))
               for i in range(nb)]
        recs.append(dict(site=site, location=loc, location_type=r["location_type"], line_no=r["line_no"], qty=qty,
                         unit=ln["unit"], status={"ACTIVE": "AVAILABLE", "HOLD": "QA_HOLD", "RECALL": "RECALL"}[r["status"]],
                         batches=bat))
    return recs


def write_stock(out, recs):
    rng = np.random.default_rng(W.ASK_SEED + 12)
    snap, breg = [], []
    for r in recs:
        for b, bb, q in r["batches"]:
            snap.append(dict(extract_ts_utc=SNAP_TS, site=r["site"], location=r["location"], location_type=r["location_type"],
                             line_no=r["line_no"], batch_no=b, best_before=bb, qty=r["qty"], stock_unit=r["unit"],
                             stock_status=r["status"]))
            if b:
                rec = VAL_DATE - dt.timedelta(days=int(rng.integers(1, 40)), minutes=int(rng.integers(0, 1440)))
                breg.append(dict(site=r["site"], location=r["location"], line_no=r["line_no"], batch_no=b, best_before=bb,
                                 batch_qty=q, received_utc=dt.datetime(rec.year, rec.month, rec.day, int(rng.integers(5, 23)),
                                                                       int(rng.integers(0, 60))).strftime("%Y-%m-%dT%H:%M:%SZ")))
    write_csv(os.path.join(out, FILES["snapshot"]), ["extract_ts_utc", "site", "location", "location_type", "line_no",
              "batch_no", "best_before", "qty", "stock_unit", "stock_status"], snap)
    breg.sort(key=lambda x: (x["site"], x["location"], x["best_before"], x["batch_no"]))
    write_csv(os.path.join(out, FILES["batches"]), ["site", "location", "line_no", "batch_no", "best_before", "batch_qty",
              "received_utc"], breg)
    return len(snap), len(breg)


def write_line_master(out, lines):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Line master"
    ws.append(["Line No", "Description", "Category", "Supply Chain", "Stock Unit", "Case Size", "Valid From", "Valid To",
               "Date Coded"])
    for ln in sorted(lines, key=lambda x: x["line_no"]):
        for size, f, t in ln["hist"]:
            ws.append([int(ln["line_no"]), ln["desc"], ln["category"], REGIME_CODE[ln["regime"]], ln["unit"], size,
                       f, t, "Y" if ln["date_coded"] else "N"])
    for row in ws.iter_rows(min_row=2, min_col=7, max_col=8):
        for cell in row:
            if cell.value is not None:
                cell.number_format = "dd/mm/yyyy"
    ws.freeze_panes = "A2"
    for col, wdt in zip("ABCDEFGHI", (10, 42, 9, 11, 10, 9, 11, 11, 10)):
        ws.column_dimensions[col].width = wdt
    wb.properties.creator = "Thornacre Merchandising"
    wb.properties.created = dt.datetime(2021, 3, 8, 9, 12, 0)
    wb.properties.modified = dt.datetime(2026, 12, 4, 8, 31, 12)
    wb.save(os.path.join(out, FILES["linemaster"]))


def write_costs(out, lines):
    rng = np.random.default_rng(W.ASK_SEED + 13)
    base = []
    for ln in sorted(lines, key=lambda x: x["line_no"]):
        eff = D(2026, 11, 28) - dt.timedelta(days=int(rng.integers(5, 400)))
        base.append(dict(line_no=ln["line_no"], cost_per_case="%.2f" % ln["cost"], currency="GBP", effective_from=eff.isoformat(),
                         extracted="2026-11-28"))
    write_csv(os.path.join(out, FILES["costbase"]), ["line_no", "cost_per_case", "currency", "effective_from", "extracted"], base)
    ch = []
    k = 0
    for ln in sorted(lines, key=lambda x: x["line_no"]):
        if ln["p10"] is None:
            continue
        k += 1
        reason = "SUPPLIER COST INCREASE" if ln["p10"] > ln["cost"] else str(rng.choice(["COST REVIEW", "SUPPLIER FUNDING"]))
        ch.append(dict(change_ref="P10-%05d" % (k * 3 + int(rng.integers(0, 3))), line_no=ln["line_no"],
                       old_cost_per_case="%.2f" % ln["cost"], new_cost_per_case="%.2f" % ln["p10"],
                       effective_from=P10_FROM.isoformat(), period="2026-P10", reason=reason))
    write_csv(os.path.join(out, FILES["costp10"]), ["change_ref", "line_no", "old_cost_per_case", "new_cost_per_case",
              "effective_from", "period", "reason"], ch)


# ---- ask B: stores, servicing, calendar, categories and the forecast run --------------------------------

REGIONS = [  # region, country, stores, ambient, chilled, frozen depot (centre names)
    ("London", "England", 58, "Thurrock", "Daventry", "Doncaster"),
    ("South East", "England", 40, "Thurrock", "Daventry", "Magor"),
    ("East of England", "England", 27, "Thurrock", "Daventry", "Doncaster"),
    ("East Midlands", "England", 24, "Hinckley", "Daventry", "Doncaster"),
    ("West Midlands", "England", 28, "Hinckley", "Daventry", "Magor"),
    ("North West", "England", 32, "Hinckley", "Warrington", "Doncaster"),
    ("Yorkshire and the Humber", "England", 24, "Hinckley", "Warrington", "Doncaster"),
    ("North East", "England", 14, "Bellshill", "Warrington", "Doncaster"),
    ("Scotland", "Scotland", 33, "Bellshill", "Warrington", "Doncaster"),
    ("Wales", "Wales", 18, "Avonmouth", "Warrington", "Magor"),
    ("South West", "England", 27, "Avonmouth", "Daventry", "Magor"),
]
TOWNS = {
    "London": ["Wandsworth", "Lewisham", "Ealing", "Croydon", "Enfield", "Bromley", "Hackney", "Harrow", "Stratford", "Kingston",
               "Islington", "Wembley", "Romford", "Sutton", "Hounslow", "Barnet", "Walthamstow", "Greenwich", "Camden", "Tooting"],
    "South East": ["Reading", "Guildford", "Crawley", "Maidstone", "Brighton", "Basingstoke", "Slough", "Ashford", "Woking",
                   "Chichester", "Tunbridge Wells", "Milton Keynes", "High Wycombe", "Oxford", "Eastbourne"],
    "East of England": ["Chelmsford", "Colchester", "Norwich", "Ipswich", "Cambridge", "Peterborough", "Luton", "Stevenage",
                        "Basildon", "Bedford", "Watford"],
    "East Midlands": ["Leicester", "Nottingham", "Derby", "Lincoln", "Northampton", "Loughborough", "Mansfield", "Kettering",
                      "Chesterfield"],
    "West Midlands": ["Birmingham", "Coventry", "Wolverhampton", "Solihull", "Walsall", "Stoke", "Worcester", "Telford",
                      "Shrewsbury", "Redditch", "Nuneaton"],
    "North West": ["Manchester", "Liverpool", "Preston", "Bolton", "Stockport", "Wigan", "Blackburn", "Chester", "Lancaster",
                   "Southport", "Oldham", "Altrincham"],
    "Yorkshire and the Humber": ["Leeds", "Sheffield", "Bradford", "Hull", "York", "Wakefield", "Huddersfield", "Harrogate",
                                 "Rotherham", "Halifax"],
    "North East": ["Newcastle", "Sunderland", "Durham", "Middlesbrough", "Gateshead", "Darlington", "Hexham"],
    "Scotland": ["Glasgow", "Edinburgh", "Aberdeen", "Dundee", "Paisley", "Livingston", "Stirling", "Ayr", "Inverness",
                 "Perth", "Falkirk", "Kilmarnock"],
    "Wales": ["Cardiff", "Swansea", "Newport", "Wrexham", "Bridgend", "Llanelli", "Aberystwyth", "Bangor"],
    "South West": ["Bristol", "Exeter", "Plymouth", "Bath", "Gloucester", "Taunton", "Swindon", "Cheltenham", "Truro",
                   "Bournemouth", "Yeovil"],
}
SUFFIX = ["", " Retail Park", " High Street", " Central", " North", " South", " East", " West", " Road", " Extra"]
# re-zoning PZ26-04 (agreed 1 December): (regime, from, to, stores, of which return on 27 December, region of the stores)
REZONE = [
    ("Ambient", "Thurrock", "Hinckley", 18, 6, ("East of England", "London")),
    ("Ambient", "Hinckley", "Bellshill", 5, 3, ("Yorkshire and the Humber",)),
    ("Ambient", "Avonmouth", "Hinckley", 4, 3, ("South West",)),
    ("Ambient", "Thurrock", "Avonmouth", 3, 1, ("South East",)),
    ("Ambient", "Hinckley", "Thurrock", 3, 1, ("East Midlands",)),
    ("Chilled", "Daventry", "Warrington", 6, 5, ("West Midlands", "East Midlands")),
    ("Chilled", "Warrington", "Daventry", 2, 1, ("Wales",)),
    ("Frozen", "Doncaster", "Magor", 6, 3, ("London", "East Midlands")),
    ("Frozen", "Magor", "Doncaster", 2, 1, ("West Midlands",)),
]
REZONE_FROM, REZONE_BACK = D(2026, 12, 7), D(2026, 12, 27)
FCST_DAYS = [D(2026, 12, 1) + dt.timedelta(days=i) for i in range(39)]   # 1 Dec .. 8 Jan
BOXING = D(2026, 12, 26)
N_BOXING_CLOSED = 41
RUN_ID, RUN_DATE = "RF-2026-W48-01", "2026-11-23"
BASE_DEMAND = {"A01": 118, "A02": 64, "A03": 58, "A04": 260, "A05": 150, "C01": 40, "C02": 30, "C03": 32, "C04": 60,
               "F01": 27, "F02": 22, "F03": 70}
FORMAT = [("Superstore", 1.9, 0.30), ("Supermarket", 1.0, 0.50), ("Local", 0.32, 0.20)]


def stores():
    rng = np.random.default_rng(W.ASK_SEED + 21)
    out = []
    nos = rng.choice(np.arange(1012, 2990), size=sum(r[2] for r in REGIONS), replace=False)
    k = 0
    for reg, country, n, amb, chl, frz in REGIONS:
        towns = TOWNS[reg]
        for i in range(n):
            f = FORMAT[int(rng.choice(3, p=[x[2] for x in FORMAT]))]
            name = towns[i % len(towns)] + (SUFFIX[(i // len(towns)) % len(SUFFIX)] if i >= len(towns) else "")
            out.append(dict(store_no=int(nos[k]), name="Thornacre " + name, format=f[0], size=f[1] * float(rng.uniform(0.8, 1.25)),
                            region=reg, country=country, depots={"Ambient": amb, "Chilled": chl, "Frozen": frz},
                            opened=(D(1998, 1, 1) + dt.timedelta(days=int(rng.integers(0, 9500)))).isoformat()))
            k += 1
    return out


def rezoning(st):
    """Store-regime moves: list of (store_no, regime, from, to, returns_27_dec)."""
    rng = np.random.default_rng(W.ASK_SEED + 22)
    moves, used = [], set()
    for regime, a, b, n, back, regions in REZONE:
        cand = [s for s in st if s["region"] in regions and s["depots"][regime] == a and (s["store_no"], regime) not in used]
        pick = rng.choice(len(cand), size=n, replace=False)
        for j, i in enumerate(sorted(pick)):
            s = cand[i]
            used.add((s["store_no"], regime))
            moves.append((s["store_no"], regime, a, b, j < back))
    return moves


def depot_on(store, regime, day, moves_ix):
    m = moves_ix.get((store["store_no"], regime))
    base = store["depots"][regime]
    if m is None or day < REZONE_FROM:
        return base
    if m[4] and day >= REZONE_BACK:
        return base
    return m[3]


def boxing_closed(st):
    rng = np.random.default_rng(W.ASK_SEED + 23)
    pool = [s for s in st if s["format"] != "Local"]
    return set(int(x) for x in rng.choice([s["store_no"] for s in pool], size=N_BOXING_CLOSED, replace=False))


def new_year_closed(st):
    return set(s["store_no"] for s in st if s["format"] == "Local" and s["store_no"] % 3 == 0)


def season(day):
    wd = day.weekday()
    f = [0.92, 0.88, 0.95, 1.02, 1.12, 1.18, 0.82][wd]
    if D(2026, 12, 18) <= day <= D(2026, 12, 24):
        f *= 1.0 + 0.07 * (day - D(2026, 12, 17)).days
    if day in (D(2026, 12, 26), D(2027, 1, 1)):
        f *= 0.55
    if D(2026, 12, 27) <= day <= D(2026, 12, 31):
        f *= 1.08
    if day >= D(2027, 1, 2):
        f *= 0.86
    return f


def forecast(st, salt=0):
    rng = np.random.default_rng(W.ASK_SEED + 24 + salt)
    nyc = new_year_closed(st)
    rows = []
    for s in st:
        for code, name, regime, unit, upc, _ in CATEGORIES:
            lvl = BASE_DEMAND[code] * s["size"] * float(rng.uniform(0.85, 1.15))
            for day in FCST_DAYS:
                if day == D(2026, 12, 25) or (day == D(2027, 1, 1) and s["store_no"] in nyc):
                    continue
                q = lvl * season(day) * float(rng.lognormal(0, 0.08))
                rows.append([s["store_no"], s["depots"][regime], code, day, round(q, 2) if unit == "CS" else round(q, 0), unit])
    return rows


def write_stores(out, st):
    data = {"source": "Thornacre estate service", "extractedAt": "2026-12-03T07:12:44Z",
            "stores": [{"storeNumber": s["store_no"], "storeName": s["name"], "format": s["format"], "region": s["region"],
                        "country": s["country"], "openedDate": s["opened"], "status": "TRADING"}
                       for s in sorted(st, key=lambda x: x["store_no"])]}
    with open(os.path.join(out, FILES["stores"]), "w") as fh:
        json.dump(data, fh, indent=1)
        fh.write("\n")


def write_servicing(out, st, moves):
    rng = np.random.default_rng(W.ASK_SEED + 25)
    dep = {c["name"]: c["depot"] for c in W.CENTRES}
    mix = {(m[0], m[1]): m for m in moves}
    rows = []
    for s in sorted(st, key=lambda x: x["store_no"]):
        for regime in ("Ambient", "Chilled", "Frozen"):
            rc = REGIME_CODE[regime]
            since = (D(2019, 3, 3) + dt.timedelta(days=int(rng.integers(0, 2000)))).isoformat()
            base = s["depots"][regime]
            m = mix.get((s["store_no"], regime))
            if m is None:
                rows.append([s["store_no"], rc, dep[base], since, "", "BASE"])
                continue
            rows.append([s["store_no"], rc, dep[base], since, "2026-12-06", "BASE"])
            if m[4]:
                rows.append([s["store_no"], rc, dep[m[3]], "2026-12-07", "2026-12-26", "PZ26-04"])
                rows.append([s["store_no"], rc, dep[base], "2026-12-27", "", "PZ26-04"])
            else:
                rows.append([s["store_no"], rc, dep[m[3]], "2026-12-07", "", "PZ26-04"])
    write_csv(os.path.join(out, FILES["servicing"]), ["store_no", "supply_chain", "depot", "valid_from", "valid_to", "change_ref"], rows)


def write_calendar(out, st, closed26):
    nyc = new_year_closed(st)
    rows = []
    for s in sorted(st, key=lambda x: x["store_no"]):
        n = s["store_no"]
        rows.append([n, "2026-12-24", "EARLY_CLOSE", "07:00", "16:00", "2026-09-14T10:22:05Z"])
        rows.append([n, "2026-12-25", "CLOSED", "", "", "2026-09-14T10:22:05Z"])
        if n in closed26:
            rows.append([n, "2026-12-26", "CLOSED", "", "", "2026-12-01T16:48:31Z"])
        else:
            rows.append([n, "2026-12-26", "LATE_OPEN", "10:00", "18:00", "2026-09-14T10:22:05Z"])
        rows.append([n, "2026-12-31", "EARLY_CLOSE", "07:00", "18:00", "2026-09-14T10:22:05Z"])
        if n in nyc:
            rows.append([n, "2027-01-01", "CLOSED", "", "", "2026-10-02T09:15:40Z"])
        else:
            rows.append([n, "2027-01-01", "LATE_OPEN", "10:00", "17:00", "2026-09-14T10:22:05Z"])
    write_csv(os.path.join(out, FILES["calendar"]), ["store_no", "trading_date", "status", "open_time", "close_time", "updated_utc"], rows)


def write_categories(out):
    rows = [[c[0], c[1], REGIME_CODE[c[2]], c[3], c[4] if c[4] else ""] for c in CATEGORIES]
    write_csv(os.path.join(out, FILES["categories"]), ["category", "category_name", "supply_chain", "forecast_unit",
              "units_per_case"], rows)


def write_forecast(out, frows):
    dep = {c["name"]: c["depot"] for c in W.CENTRES}
    rows = [[RUN_ID, RUN_DATE, r[0], dep[r[1]], r[2], r[3].isoformat(), ("%.2f" % r[4]) if r[5] == "CS" else "%d" % r[4], r[5]]
            for r in frows]
    write_csv(os.path.join(out, FILES["forecast"]), ["run_id", "run_date", "store_no", "depot", "category", "trading_date",
              "fcst_qty", "uom"], rows)
    return len(rows)


# ---- distractors ------------------------------------------------------------------------------------------

def write_labour(out):
    from openpyxl import Workbook
    rng = np.random.default_rng(W.ASK_SEED + 31)
    wb = Workbook()
    ws = wb.active
    ws.title = "Agency hours"
    weeks = [D(2026, 11, 30) + dt.timedelta(days=7 * i) for i in range(6)]
    ws.append(["Site", "Centre"] + ["w/c %s" % w.strftime("%d/%m") for w in weeks] + ["Total"])
    for c in W.CENTRES:
        base = c["n"] * float(rng.uniform(0.9, 1.3))
        vals = [int(round(base * f * float(rng.uniform(0.94, 1.06)), -1)) for f in (0.85, 1.15, 1.35, 1.2, 0.7, 0.55)]
        ws.append([c["code"], c["name"]] + vals + [sum(vals)])
    ws2 = wb.create_sheet("Headcount")
    ws2.append(["Site", "Permanent FTE", "Agency FTE peak", "Agency supplier", "Induction complete by"])
    for c in W.CENTRES:
        ws2.append([c["code"], int(c["n"] * float(rng.uniform(0.11, 0.15))), int(c["n"] * float(rng.uniform(0.04, 0.07))),
                    str(rng.choice(["Staffline", "Pertemps", "Blue Arrow", "Adecco"])), "27/11/2026"])
    wb.properties.creator = "Corderoy Peak Planning"
    wb.properties.created = dt.datetime(2026, 10, 19, 14, 2, 0)
    wb.properties.modified = dt.datetime(2026, 11, 27, 17, 5, 48)
    wb.save(os.path.join(out, FILES["labour"]))


def write_goodsin(out):
    rng = np.random.default_rng(W.ASK_SEED + 32)
    rows = []
    for day in W.FREEZE_DAYS:
        if day == D(2026, 12, 25):
            continue
        for c in W.CENTRES:
            for typ in ("PRIMARY", "IMPORT", "RETURNS"):
                if typ == "IMPORT" and c["regime"] != "Ambient":
                    continue
                b = int(c["n"] / {"PRIMARY": 55, "IMPORT": 400, "RETURNS": 300}[typ] * float(rng.uniform(0.7, 1.3)) * season(day))
                pal = int(b * float(rng.uniform(14, 24)))
                rows.append([day.isoformat(), c["code"], typ, b, pal, int(b * float(rng.uniform(0.02, 0.09)))])
    write_csv(os.path.join(out, FILES["goodsin"]), ["booking_date", "site", "booking_type", "bookings", "pallets_booked",
              "late_bookings"], rows)


# ---- the ask computations, read from the shipped files ----------------------------------------------------

def _read(target, key, **kw):
    import pandas as pd
    return pd.read_csv(os.path.join(target, FILES[key]), dtype=str, keep_default_na=False, **kw)


def line_master(target):
    import pandas as pd
    lm = pd.read_excel(os.path.join(target, FILES["linemaster"]), dtype={"Line No": str})
    lm["Valid From"] = pd.to_datetime(lm["Valid From"]).dt.date
    lm["Valid To"] = pd.to_datetime(lm["Valid To"]).dt.date
    return lm


def ask_a(target, fanout=True, units=True, case_version="current", p10=True, overclean=False):
    """Thornacre stock value on each centre's in-scope faces at 06:00 on 7 December, GBP (unrounded)."""
    import pandas as pd
    import estimate as E
    loc = pd.read_csv(os.path.join(target, "wms_location_master_THN_eff.csv"), dtype=str, keep_default_na=False)
    faces = E.in_scope(loc)[["site", "location", "line_no"]]
    snap = _read(target, "snapshot")
    snap["qty"] = snap.qty.astype(float)
    s = snap.merge(faces, on=["site", "location", "line_no"], how="inner")
    if fanout:
        s = s.drop_duplicates(["site", "location", "line_no"])
    if overclean:
        s = s.drop_duplicates(["site", "line_no"])
    lm = line_master(target)
    if case_version == "current":
        cur = lm[(lm["Valid From"] <= VAL_DATE) & (lm["Valid To"].isna() | (lm["Valid To"] >= VAL_DATE))]
    else:
        cur = lm.sort_values("Valid From").drop_duplicates("Line No", keep="first")
    size = dict(zip(cur["Line No"], cur["Case Size"]))
    base = _read(target, "costbase")
    cost = dict(zip(base.line_no, base.cost_per_case.astype(float)))
    if p10:
        ch = _read(target, "costp10")
        ch = ch[ch.effective_from <= VAL_DATE.isoformat()]
        cost.update(dict(zip(ch.line_no, ch.new_cost_per_case.astype(float))))
    cases = s.qty
    if units:
        cases = s.qty.where(s.stock_unit == "CS", s.qty / s.line_no.map(size).astype(float))
    s = s.assign(value=cases * s.line_no.map(cost))
    names = {c["code"]: c["name"] for c in W.CENTRES}
    v = s.groupby("site").value.sum()
    return {names[k]: float(v.get(k, 0.0)) for k in names}


def ask_b(target, basis="daily", units=True, closures=True, overclean=False):
    """Cases Thornacre's stores are forecast to draw from each centre over the freeze (unrounded)."""
    import pandas as pd
    f = _read(target, "forecast")
    f = f[(f.trading_date >= "2026-12-07") & (f.trading_date <= "2027-01-03")].copy()
    f["q"] = f.fcst_qty.astype(float)
    cat = _read(target, "categories")
    f = f.merge(cat[["category", "supply_chain", "units_per_case"]], on="category", how="left")
    if units:
        ea = f.uom == "EA"
        f.loc[ea, "q"] = f.loc[ea, "q"] / f.loc[ea, "units_per_case"].astype(float)
    if closures:
        cal = _read(target, "calendar")
        shut = set(zip(cal[cal.status == "CLOSED"].store_no, cal[cal.status == "CLOSED"].trading_date))
        f = f[[(a, b) not in shut for a, b in zip(f.store_no, f.trading_date)]]
    sm = _read(target, "servicing")
    if overclean:
        moved = set(zip(sm[sm.change_ref != "BASE"].store_no, sm[sm.change_ref != "BASE"].supply_chain))
        f = f[[(a, b) not in moved for a, b in zip(f.store_no, f.supply_chain)]]
    if basis == "run":
        f["dep"] = f.depot
    else:
        sm = sm.copy()
        sm["valid_to"] = sm.valid_to.replace("", "9999-12-31")
        if basis == "start":
            on = sm[(sm.valid_from <= "2026-12-07") & (sm.valid_to >= "2026-12-07")]
            f = f.merge(on[["store_no", "supply_chain", "depot"]].rename(columns={"depot": "dep"}), on=["store_no", "supply_chain"])
        else:
            m = f.merge(sm[["store_no", "supply_chain", "depot", "valid_from", "valid_to"]].rename(columns={"depot": "dep"}),
                        on=["store_no", "supply_chain"])
            f = m[(m.valid_from <= m.trading_date) & (m.valid_to >= m.trading_date)]
    names = {c["depot"]: c["name"] for c in W.CENTRES}
    v = f.groupby("dep").q.sum()
    return {names[k]: float(v.get(k, 0.0)) for k in names}
