"""Cut the pack for task251 architecture 5 from the world."""
import datetime as dt
import os
from xml.sax.saxutils import escape

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

import docs
import ops
import params as P
import postprocess
import records as R
import sim

FILES = dict(
    standard="TDP-03_permanent_count_stations_rev4.pdf",
    hourly="trafficvault_hourly_counts_2027.csv",
    validation="Northline_validation_counts_2027.xlsx",
    register="count_station_register.xlsx",
    network="regional_road_network_2027.xlsx",
    mail="Carrick_spring_recalibration_round_2028.eml",
    events="trafficvault_station_events_2027.csv",
    orders="Carrick_work_orders_2027.xlsx",
    weather="weather_daily_2027.csv",
    collisions="collisions_regional_roads_2027.xlsx",
    pavement="regional_roads_pavement_condition_2026.xlsx",
)
PACK = list(FILES)
DISTRACTORS = ("collisions", "pavement")
CONTEXT = ("mail",)


def _sheet(wb, title, df, widths=None, first=False, fmts=None):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(list(df.columns))
    col_int = {i: c in ("lanes_each_way", "built", "in_service_since") for i, c in enumerate(df.columns)}
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and np.isnan(v)) else
                   (int(v) if isinstance(v, float) and float(v).is_integer() and col_int.get(i) else
                    (v.item() if hasattr(v, "item") else v)) for i, v in enumerate(row)])
    fmts = dict(fmts or {})
    for i, col in enumerate(df.columns):
        if len(df) and isinstance(df[col].iloc[0], dt.date):
            fmts.setdefault(col, "yyyy-mm-dd")
    for col, f in fmts.items():
        j = list(df.columns).index(col) + 1
        for r in range(2, len(df) + 2):
            ws.cell(row=r, column=j).number_format = f
    ws.freeze_panes = "A2"
    for i, col in enumerate(df.columns, 1):
        w = (widths or {}).get(col) or max(10, min(40, int(df[col].astype(str).str.len().max() if len(df) else 10) + 2))
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws


def weather(W, rng, O):
    """the airport's daily data as the weather service's bulk export gives it: the day's temperatures with their mean,
    degree days, rain, snow and precipitation, snow on the ground and the strongest gust, each value with its flag; the
    storm days behind the outages; one day the station's temperature record is missing"""
    storm = {o["start"].date() for o in O["outs"]}
    days = W["C"]["days"]
    lost = days[int(rng.integers(20, len(days) - 20))]
    snow_ground = 0.0
    rows = []
    for day in days:
        tmax, tmin = W["temp"][day]
        mean = round((tmax + tmin) / 2 + 1e-9, 1)
        wet = rng.random() < 0.34 or day in storm
        precip = float(rng.gamma(1.4, 5.0)) + (rng.uniform(8, 30) if day in storm else 0.0)
        rain = round(precip, 1) if wet and tmax > 1 else 0.0
        snow = round(float(rng.gamma(1.3, 4.0)) + (rng.uniform(5, 15) if day in storm else 0.0), 1) \
            if wet and tmax <= 1 else 0.0
        trace_r = (not wet) and tmax > 1 and rng.random() < 0.06
        trace_s = (not wet) and tmax <= 1 and rng.random() < 0.10
        snow_ground = max(0.0, snow_ground + snow * 0.8 - max(0.0, tmax) * 1.6 - 0.3)
        gust = rng.normal(38, 9) + (rng.uniform(25, 45) if day in storm else 0.0)
        r = {"Longitude (x)": -80.71, "Latitude (y)": 43.27, "Station Name": "HALBROOK A", "Climate ID": "6152988",
             "Date/Time": day.isoformat(), "Year": day.year, "Month": day.month, "Day": day.day, "Data Quality": "",
             "Max Temp (\u00b0C)": tmax, "Max Temp Flag": "", "Min Temp (\u00b0C)": tmin, "Min Temp Flag": "",
             "Mean Temp (\u00b0C)": mean, "Mean Temp Flag": "",
             "Heat Deg Days (\u00b0C)": round(max(0.0, 18 - mean), 1), "Heat Deg Days Flag": "",
             "Cool Deg Days (\u00b0C)": round(max(0.0, mean - 18), 1), "Cool Deg Days Flag": "",
             "Total Rain (mm)": rain, "Total Rain Flag": "T" if trace_r else "",
             "Total Snow (cm)": snow, "Total Snow Flag": "T" if trace_s else "",
             "Total Precip (mm)": round(rain + snow * 0.9, 1), "Total Precip Flag": "T" if (trace_r or trace_s) else "",
             "Snow on Grnd (cm)": int(round(snow_ground)) if snow_ground >= 0.5 else (0 if day.month in (11, 12, 1, 2, 3) else ""),
             "Snow on Grnd Flag": "",
             "Dir of Max Gust (10s deg)": int(rng.integers(1, 37)) if gust >= 31 else "", "Dir of Max Gust Flag": "",
             "Spd of Max Gust (km/h)": int(gust) if gust >= 31 else "<31", "Spd of Max Gust Flag": ""}
        if day == lost:
            for k in ("Max Temp", "Min Temp", "Mean Temp", "Heat Deg Days", "Cool Deg Days"):
                col = [c for c in r if c.startswith(k + " (")][0]
                r[col] = ""
                r[k + " Flag"] = "M"
        rows.append(r)
    return pd.DataFrame(rows)


def pavement(W, rng):
    """the Region's pavement management system: each section of the Regional Roads with its surface, last
    rehabilitation, and the condition measured in the autumn 2026 survey"""
    rows = []
    roads = sorted({x["road"] for x in W["st"]})
    for road in roads:
        km, n = 0.0, 1
        end = float(rng.uniform(14.5, 19.5))
        while km < end - 0.2:
            length = min(end - km, float(rng.uniform(0.4, 2.6)))
            built = int(rng.integers(1994, 2025))
            age = 2026 - built
            pci = float(np.clip(rng.normal(96 - 2.1 * age, 6.5), 18, 100))
            iri = round(float(np.clip(0.9 + (100 - pci) * 0.045 + rng.normal(0, 0.18), 0.7, 6.5)), 2)
            rut = round(float(np.clip((100 - pci) * 0.12 + rng.normal(0, 1.1), 0.5, 18.0)), 1)
            cond = "Excellent" if pci >= 85 else "Good" if pci >= 70 else "Fair" if pci >= 55 else \
                "Poor" if pci >= 40 else "Very poor"
            surf = str(rng.choice(["HMA", "HMA", "HMA", "HMA over concrete", "surface treatment"],
                                  p=[0.55, 0.15, 0.1, 0.12, 0.08]))
            survey = dt.date(2026, 9, 14) + dt.timedelta(days=int(rng.integers(0, 45)))
            rows.append(dict(section="RR%02d-%02d" % (road, n), road="RR %d" % road, from_km=round(km, 2),
                             to_km=round(km + length, 2), length_km=round(length, 2), surface=surf,
                             last_rehabilitated=built, pci=int(round(pci)), iri_m_per_km=iri, rut_mm=rut,
                             condition=cond, surveyed=survey))
            km += length
            n += 1
    return pd.DataFrame(rows)


def collisions(W, rng):
    rows = []
    roads = sorted({x["road"] for x in W["st"]})
    for k in range(int(rng.integers(260, 320))):
        day = W["C"]["days"][int(rng.integers(len(W["C"]["days"])))]
        road = int(rng.choice(roads))
        rows.append(dict(date=day.isoformat(), time="%02d:%02d" % (int(rng.integers(0, 24)), int(rng.integers(0, 60))),
                         road="RR %d" % road, km=round(float(rng.uniform(0.3, 18.0)), 1),
                         severity=str(rng.choice(["property damage", "property damage", "property damage", "injury",
                                                  "injury", "fatal"], p=[.3, .25, .2, .13, .11, .01])),
                         vehicles=int(rng.choice([1, 2, 2, 2, 3, 4]))))
    return pd.DataFrame(rows).sort_values(["date", "time"]).reset_index(drop=True)


def run(target, seed=P.SEED, W=None):
    os.makedirs(target, exist_ok=True)
    W = W or sim.world(seed)
    F = {k: os.path.join(target, v) for k, v in FILES.items()}
    rng = np.random.default_rng(seed + 900)
    O = ops.ops(W)
    jn, sn = R.network(W)
    st, hist = R.register(W, jn)
    docs.standard(F["standard"], FILES)
    docs.mail(F["mail"], FILES)
    R.hourly(W, O).to_csv(F["hourly"], index=False)
    R.events(W, O).to_csv(F["events"], index=False)
    orders, labour = R.work_orders(W, O)
    wb = Workbook()
    _sheet(wb, "Orders", orders, first=True, widths=dict(description=44, resolution=44, stations=34))
    _sheet(wb, "Labour", labour, fmts=dict(hours="0.00"))
    wb.save(F["orders"])
    wb = Workbook()
    _sheet(wb, "Counts", R.validation(W), first=True)
    wb.save(F["validation"])
    wb = Workbook()
    _sheet(wb, "Stations", st, first=True, fmts=dict(km="0.00"))
    _sheet(wb, "Code history", hist)
    wb.save(F["register"])
    wb = Workbook()
    _sheet(wb, "Junctions", jn, first=True, fmts=dict(km="0.00"))
    _sheet(wb, "Structures", sn, fmts=dict(km_from="0.00", km_to="0.00"))
    wb.save(F["network"])
    weather(W, rng, O).to_csv(F["weather"], index=False)
    wb = Workbook()
    _sheet(wb, "Collisions", collisions(W, rng), first=True)
    wb.save(F["collisions"])
    wb = Workbook()
    _sheet(wb, "Sections", pavement(W, rng), first=True, fmts=dict(from_km="0.00", to_km="0.00", length_km="0.00"))
    wb.save(F["pavement"])
    postprocess.scrub(F["standard"], "Microsoft Word 2019", "2026-01-15T10:22:00")
    for k, who, made, saved in (("orders", P.CONTRACTOR_MAINT, "2028-01-06T14:02:00Z", "2028-01-07T09:31:00Z"),
                                ("validation", P.CONTRACTOR_VAL, "2027-09-13T15:40:00Z", "2027-09-17T11:08:00Z"),
                                ("register", P.ORG, "2021-03-02T13:15:00Z", "2028-01-05T10:44:00Z"),
                                ("network", P.ORG, "2024-05-21T09:12:00Z", "2028-01-05T10:52:00Z"),
                                ("collisions", P.ORG, "2028-01-04T08:20:00Z", "2028-01-04T08:26:00Z"),
                                ("pavement", P.ORG, "2026-11-30T15:05:00Z", "2026-12-02T10:18:00Z")):
        who = escape(who)
        postprocess.scrub(F[k], who, saved[:10])
        postprocess.finish_office(F[k], made, saved, who)
    return W, O


if __name__ == "__main__":
    import sys
    run(sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                                                           "target"), seed=int(sys.argv[2]) if len(sys.argv) > 2 else P.SEED)
