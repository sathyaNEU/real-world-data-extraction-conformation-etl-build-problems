"""Independent verifier for task238: recomputes the band, the response rates and the valid answers from the shipped
files alone (no generator module), and checks figures.json, the golden workbook and submission.md against them."""
import datetime as dt
import json
import math
import os
import re
import sys
from zoneinfo import ZoneInfo

import numpy as np
import pandas as pd

MT = ZoneInfo("Pacific/Auckland")      # Halvreth time: UTC+12, UTC+13 in daylight time (CST-27 plan, section 4)
CLOSE = {"2027Q2": dt.date(2027, 6, 30), "2027Q3": dt.date(2027, 9, 30)}
OUTAGE = (dt.date(2027, 8, 9), dt.date(2027, 8, 12))


def _pdf(path):
    from pypdf import PdfReader
    return " ".join(" ".join(p.extract_text().split()) for p in PdfReader(path).pages)


def _range(text):
    a, b = [int(x) for x in re.findall(r"\d+", text)[:2]]
    return list(range(a, b + 1))


def main(target, golden, figures_path, submission_path):
    f = {n: os.path.join(target, n) for n in os.listdir(target)}
    std = _pdf(f["mso_ms14_tracker_continuity_standard.pdf"])
    tol = {m.group(1): float(m.group(2)) for m in re.finditer(r"\b([A-Z]{2}) (\d\.\d\d) (?:percentage )?points", std)}
    assert len(tol) == 12, tol
    ix = pd.read_excel(f["mwt_questionnaire_codebook_v9.xlsx"], sheet_name="Indexes", dtype=str).dropna(subset=["code"])
    iss = pd.read_csv(f["mwt_cst27_issued_addresses.csv"], dtype=str)
    web = pd.read_csv(f["mwt_cst27_web_returns.csv"], dtype={"login_code": str})
    cati = pd.read_csv(f["mwt_cst27_cati_interviews.csv"], dtype={"ADDRESS_ID": str})
    addr = iss.set_index("login_code")["address_id"]
    q_of = iss.set_index("address_id")["quarter"]

    w = web[web["last_section"] == 6].copy()
    w["address_id"] = addr.loc[w["login_code"]].values
    w["t"] = pd.to_datetime(w["submitted_utc"], utc=True)
    close_utc = w["address_id"].map(q_of).map(
        lambda q: pd.Timestamp(dt.datetime.combine(CLOSE[q], dt.time(23, 59), tzinfo=MT)).tz_convert("UTC"))
    w = w[(w["t"] <= close_utc) & ~(w["age"] < 18)]
    loc = w["t"].dt.tz_convert(MT).dt.date
    w["mode"] = np.where((w["opened_by"] == "staff") & (loc >= OUTAGE[0]) & (loc <= OUTAGE[1]), "phone", "online")
    c = cati[cati["RESULT"] == "C"].copy()
    c["address_id"] = c["ADDRESS_ID"]
    c["t"] = [pd.Timestamp(dt.datetime.fromisoformat(x).replace(tzinfo=MT)).tz_convert("UTC") for x in c["END_TIME"]]
    c["mode"] = "phone"
    items = [x for x in w.columns if re.fullmatch(r"Q\d\d", x)]
    s = pd.concat([w[["address_id", "t", "mode"] + items], c[["address_id", "t", "mode"] + items]])
    s = s.sort_values("t").drop_duplicates("address_id").merge(iss[["address_id", "letter", "quarter"]], on="address_id")
    online = s["mode"] == "online"
    A = s["letter"] == "A"

    widest, valid = {}, {}
    for r in ix.itertuples():
        base = _range(r._6)
        v = s[r.question]
        ok = v.isin(base)
        if str(r.measure).startswith("mean"):
            y = v[ok].astype(float)
        else:
            counted = _range(r._5) if "to" in str(r._5) else [int(x) for x in re.findall(r"\d+", str(r._5).split("(")[0])]
            y = v[ok].isin(counted).astype(float) * 100
        a, o_ = A[ok], online[ok]
        eff = (y[a].mean() - y[~a].mean()) / (o_[a].mean() - o_[~a].mean())
        widest[r.code] = 100 * tol[r.code] / abs(eff)
        valid[r.code] = (int((ok & online).sum()), int((ok & ~online).sum()))
    bind = min(widest, key=widest.get)
    band = math.floor(widest[bind])

    rates = {}
    done = set(s["address_id"])
    for q in ("2027Q2", "2027Q3"):
        for L in ("A", "B"):
            g = iss[(iss["quarter"] == q) & (iss["letter"] == L)]
            has = g["address_id"].isin(done)
            elig = ~g["final_outcome"].isin(["4.1", "4.2", "4.3"]) | has
            rates[(q, L)] = round(100 * has.sum() / elig.sum(), 1)

    fig = json.load(open(figures_path))
    assert fig["call"]["band"] == band == 4 and bind == "BD", (band, bind)
    for (q, L), v in rates.items():
        key = "%s quarter 2027 %s" % ({"2027Q2": "second", "2027Q3": "third"}[q], L)
        assert fig["response"][key] == "%.1f" % v, (key, v)
    for k, (on, ph) in valid.items():
        assert fig["valid"][k] == [format(on, ","), format(ph, ",")], (k, on, ph)
    from openpyxl import load_workbook
    wb = load_workbook(os.path.join(golden, "mwt_mode_mix_band_2027q4.xlsx"), data_only=True)
    assert str(wb["Band"]["B3"].value).startswith("%d percentage points" % band)
    tr = [(r[2].value, r[3].value, r[4].value) for r in wb["Test response"].iter_rows(min_row=2, max_row=5)]
    assert [round(x[2], 1) for x in tr] == [rates[k] for k in sorted(rates)], tr
    va = {r[1].value: (r[3].value, r[4].value) for r in wb["Valid answers"].iter_rows(min_row=2, max_row=13)}
    assert va == valid, (va, valid)
    sub = open(submission_path).read()
    assert "band of %d percentage points" % band in sub
    for (q, L), v in rates.items():
        assert "%s quarter 2027, letter %s: %.1f per cent" % ({"2027Q2": "second", "2027Q3": "third"}[q], L, v) in sub
    for k, (on, ph) in valid.items():
        assert "(%s): %s online, %s by telephone" % (k, format(on, ","), format(ph, ",")) in sub
    print("verify: band %d (%s %.2f), %d completes, response rates %s, valid answers for %d indexes: all agree"
          % (band, bind, widest[bind], len(s), [rates[k] for k in sorted(rates)], len(valid)))
    return dict(band=band, bind=bind, rates=rates, valid=valid)


if __name__ == "__main__":
    t = sys.argv[1]
    main(os.path.join(t, "target"), os.path.join(t, "golden"), os.path.join(t, "generator", "figures.json"),
         os.path.join(t, "submission.md"))
