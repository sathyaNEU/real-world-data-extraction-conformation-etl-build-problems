"""task238 documents: the office's continuity standard MS-14 and the CST-27 test plan (pdf), the close-out (docx), the
questionnaire codebook (xlsx), the data dictionary and the readme (md)."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import params as P
import postprocess

OFFICE = "Halvreth Statistics Office"
TRACKER = "Halvreth Wellbeing Tracker"
PEOPLE = {"method": "Tania Nicholas", "field": "Bernard Johnstone", "proc": "Leonard McMurray"}
QUESTIONS = [
    ("Q01", 1, "Taking everything into account, how satisfied are you with your life at the moment?",
     "0 (not at all satisfied) to 10 (completely satisfied)"),
    ("Q02", 1, "How much do you feel the things you do in your life are worthwhile?",
     "0 (not at all worthwhile) to 10 (completely worthwhile)"),
    ("Q03", 1, "How happy did you feel yesterday, overall?", "0 (not at all happy) to 10 (completely happy)"),
    ("Q04", 1, "How anxious did you feel yesterday, overall?", "0 (not at all anxious) to 10 (completely anxious)"),
    ("Q10", 2, "How often do you feel lonely?",
     "1 hardly ever or never; 2 occasionally; 3 some of the time; 4 often; 5 always"),
    ("Q11", 2, "Would you say that most people can be trusted, or that you can't be too careful in dealing with people?",
     "1 most people can be trusted; 2 you can't be too careful; 3 it depends"),
    ("Q12", 2, "How safe do you feel walking alone in your neighbourhood after dark?",
     "1 very safe; 2 fairly safe; 3 a bit unsafe; 4 very unsafe; 5 I never walk alone after dark"),
    ("Q13", 3, "In general, how is your health?", "1 very good; 2 good; 3 fair; 4 bad; 5 very bad"),
    ("Q14", 3, "Do you smoke cigarettes or other tobacco nowadays?",
     "1 yes, every day; 2 yes, but not every day; 3 no"),
    ("Q15", 3, "In the last seven days, on how many days did you have six or more standard drinks?",
     "0 to 7 days (0 also for people who did not drink)"),
    ("Q16", 4, "In the last 12 months, have you done any unpaid work for a group or organisation, such as a club, "
               "school, church or charity?", "1 yes; 2 no"),
    ("Q17", 4, "How well would you say you are managing financially these days?",
     "1 living comfortably; 2 doing alright; 3 just about getting by; 4 finding it quite difficult; 5 finding it very "
     "difficult"),
]
INDEX_ROWS = [
    ("LS", "Life satisfaction", "Q01", "mean score, 0 to 10", "", "0 to 10"),
    ("WW", "Worthwhile", "Q02", "mean score, 0 to 10", "", "0 to 10"),
    ("HY", "Happiness yesterday", "Q03", "mean score, 0 to 10", "", "0 to 10"),
    ("AX", "Anxiety yesterday", "Q04", "mean score, 0 to 10", "", "0 to 10"),
    ("LN", "Loneliness", "Q10", "per cent", "4 or 5 (often, always)", "1 to 5"),
    ("TR", "Trust in others", "Q11", "per cent", "1 (most people can be trusted)", "1 to 3"),
    ("SF", "Safe after dark", "Q12", "per cent", "1 or 2 (very or fairly safe)", "1 to 5"),
    ("GH", "Good health", "Q13", "per cent", "1 or 2 (very good or good)", "1 to 5"),
    ("SM", "Smoking", "Q14", "per cent", "1 or 2 (every day, not every day)", "1 to 3"),
    ("BD", "Heavy drinking", "Q15", "per cent", "1 to 7 days", "0 to 7"),
    ("VL", "Volunteering", "Q16", "per cent", "1 (yes)", "1 or 2"),
    ("FS", "Financial strain", "Q17", "per cent", "4 or 5 (quite or very difficult)", "1 to 5"),
]


def _styles():
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontName="Helvetica", fontSize=9.4, leading=13, spaceAfter=5)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName="Helvetica-Bold", fontSize=14, leading=18, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=10.5, leading=14, spaceBefore=8,
                        spaceAfter=4, keepWithNext=1)
    small = ParagraphStyle("small", parent=body, fontSize=8, leading=10.5, textColor=colors.HexColor("#444444"))
    return dict(body=body, h1=h1, h2=h2, small=small)


def _footer(text):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(18 * mm, 10 * mm, text)
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, "Page %d" % doc.page)
        canvas.restoreState()
    return draw


def _table(rows, widths, size=8.2):
    t = Table(rows, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", size), ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", size),
                           ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9AA3A8")),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E6EEF1")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    return t


def _doc(path, title, author, foot, story, stamp):
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm,
                            bottomMargin=18 * mm, title=title, author=author, creator=author, invariant=1)
    doc.build(story, onFirstPage=_footer(foot), onLaterPages=_footer(foot))
    postprocess.scrub(path, author, stamp)


def _p(text, style):
    return Paragraph(text, style)


# ------------------------------------------------------------------ MS-14, the continuity standard
def standard(path):
    S = _styles(); b, h1, h2, sm = S["body"], S["h1"], S["h2"], S["small"]
    tol = {x[0]: x[4] for x in P.INDEXES}
    st = [_p("Continuity of the %s through the move to web-first collection" % TRACKER, h1),
          _p("%s. Methods Standard MS-14, issued 8 October 2027; applies from the fourth quarter of 2027. Owner: %s, "
             "Head of Survey Methodology." % (OFFICE, PEOPLE["method"]), sm),
          Spacer(1, 4),
          _p("1. Purpose", h2),
          _p("The %s has been collected telephone-first since its first quarter in 2019. From the fourth quarter of 2027 "
             "every sampled address receives a web-first letter, and each quarter's estimates are made from all of the "
             "quarter's completes, online and by telephone together. People can answer the same question differently "
             "online and to an interviewer, so a quarter in which more or fewer households answer online can move an "
             "index although nothing about the households has changed. This standard sets how that movement is watched "
             "and when it is corrected, so that the published series stay comparable from one quarter to the next."
             % TRACKER, b),
          _p("2. Scope", h2),
          _p("The twelve headline indexes in Table 1, as the questionnaire codebook scores them. Other tracker estimates "
             "follow the same rules where a release says so.", b)]
    rows = [["Index", "Code", "Question", "Measure"]]
    for code, name, item, measure, counted, base in INDEX_ROWS:
        rows.append([name, code, item, measure if counted == "" else "per cent answering " + counted])
    st += [_table(rows, [42 * mm, 14 * mm, 18 * mm, 98 * mm]), _p("Table 1. Headline indexes.", sm),
           _p("3. Terms", h2),
           _p("<b>Web share.</b> The percentage of a quarter's completes made online.", b),
           _p("<b>Reference share.</b> The web share of the first web-first quarter, the fourth quarter of 2027.", b),
           _p("<b>Mode effect.</b> For a headline index, the difference between the index as it would stand if every "
              "household in the population it is measured on answered online and the index as it would stand if every "
              "one of those households answered by telephone. It is given in the index's own units: points on the 0 to "
              "10 scores, percentage points on the percentage indexes.", b),
           _p("<b>Mode-mix artefact.</b> The movement in an index that comes from the web share alone. A quarter whose "
              "web share departs from the reference share by <i>d</i> percentage points carries an artefact of "
              "<i>d</i>/100 times the index's mode effect.", b),
           _p("4. The mode-mix monitor", h2),
           _p("Before each quarter's estimates are released, the processing team compares the quarter's web share with "
              "the reference share. Where the departure is no wider than the monitor's band, the indexes are released as "
              "estimated. Where it is wider, the quarter's completes are calibrated to the reference share by mode "
              "before the indexes are estimated, and the release notes the adjustment.", b),
           _p("5. Setting the band", h2),
           _p("The band is a whole number of percentage points of web share: the widest whole-point band at which the "
              "mode-mix artefact on every headline index, taken at the central estimate of its mode effect, stays within "
              "that index's continuity tolerance (Table 2). The Head of Survey Methodology sets it before the "
              "fourth-quarter estimates are processed and reviews it each year.", b),
           _p("The mode effects are measured on the 2027 contact strategy test (test plan CST-27): on its completes from "
              "both quarters and under both letters.", b)]
    rows = [["Index", "Code", "Tolerance"]]
    for code, name, item, measure, counted, base in INDEX_ROWS:
        unit = "points" if measure.startswith("mean") else "percentage points"
        rows.append([name, code, "%.2f %s" % (tol[code], unit)])
    st += [KeepTogether([_table(rows, [50 * mm, 16 * mm, 50 * mm]),
                         _p("Table 2. Continuity tolerances, agreed with the tracker's main users in 2026.", sm)]),
           _p("6. Records", h2),
           _p("The band, the mode effects behind it and the test's response are filed in the tracker's methods log with "
              "the date the band was set.", b)]
    _doc(path, "MS-14 Continuity of the Halvreth Wellbeing Tracker", PEOPLE["method"],
         "%s | MS-14 | issued 8 October 2027" % OFFICE, st, "2027-10-08")


# ------------------------------------------------------------------ CST-27, the test plan
def plan(path):
    S = _styles(); b, h1, h2, sm = S["body"], S["h1"], S["h2"], S["small"]
    st = [_p("Contact strategy test 2027 (CST-27): test plan", h1),
          _p("%s, %s. Prepared by %s, Fieldwork Manager; approved by %s, Head of Survey Methodology, 8 March 2027."
             % (OFFICE, TRACKER, PEOPLE["field"], PEOPLE["method"]), sm),
          Spacer(1, 4),
          _p("1. Why the test", h2),
          _p("Interviewer hours per complete on the tracker have risen every year since 2021 as fewer households answer "
             "the phone. The office will choose the tracker's contact strategy for the fourth quarter of 2027 onward. "
             "This test runs the two candidate strategies side by side for two quarters so that the choice rests on the "
             "response and the cost of a complete under each.", b),
          _p("2. Sample and allocation", h2),
          _p("Each quarter's sample is the tracker's usual 38,000 addresses, drawn from the Halvreth Address Register's "
             "residential addresses with a contact number and released in six fortnightly batches. Before the first "
             "batch is mailed, the addresses in each region are allocated at random, half to letter A and half to letter "
             "B. The allocation is held on the sample management system and is not changed after it is made.", b),
          _p("3. The two strategies", h2),
          _p("<b>Letter A, web-first.</b> A letter inviting the person in the household aged 18 or over with the next "
             "birthday to complete the tracker online with the household's login code; a reminder with the same code "
             "eight days later; from the fifteenth day, calls to households that have not completed.", b),
          _p("<b>Letter B, telephone-first</b> (the tracker's current design). An advance letter saying an interviewer "
             "will call; calls from the third day; after six calls without an interview, a reminder letter with the "
             "household's login code for the online questionnaire.", b),
          _p("Both routes ask the same questions in the same order, and interviewers ask for the same person the letter "
             "asks for. The letters are in Appendix A.", b),
          _p("4. Fieldwork calendar", h2)]
    rows = [["Quarter", "Batches mailed", "Fieldwork closes"],
            ["Second", "Thu 1 Apr, 15 Apr, 29 Apr, 13 May, 27 May, 10 Jun", "23:59, Wednesday 30 June 2027"],
            ["Third", "Thu 1 Jul, 15 Jul, 29 Jul, 12 Aug, 26 Aug, 9 Sep", "23:59, Thursday 30 September 2027"]]
    st += [_table(rows, [22 * mm, 92 * mm, 58 * mm]),
           _p("Times are Halvreth time. Daylight time ends at 03:00 on Sunday 4 April and begins at 02:00 on Sunday 26 "
              "September; interviewer shifts follow the clock.", sm),
           _p("5. What counts", h2),
           _p("A web return counts when the respondent reaches the end of the questionnaire before the quarter's close; "
              "a telephone interview counts when the interviewer codes it complete. A return or interview counts only when "
              "it comes from the person the letters ask for. Returns and interviews that stop before the end are partials "
              "and do not count. An address counts once: where it has a complete both online and by telephone, the one "
              "submitted first counts. The web platform keeps accepting returns after a close; they do not count for the "
              "quarter.", b),
           _p("An address's outcome is its highest result: a complete before a partial, and a partial before any "
              "non-response or eligibility code. The response rate is completes as a percentage of eligible addresses. "
              "Eligible addresses are those issued, less addresses found to be vacant, not a private dwelling or not "
              "found (outcome codes 4.1, 4.2 and 4.3); addresses whose eligibility was never settled (3.1) stay in.", b),
           _p("6. Systems", h2),
           _p("Telephone interviews are held in the CATI system, stamped in Halvreth time. Web returns are held on the "
              "web platform, stamped in UTC, and carry the household's login code. The sample management system holds "
              "each address's allocation, mailing date and final outcome. Each night at 02:00 the web platform's new "
              "completes are passed to the CATI system and taken out of the next day's call lists. When a household's "
              "login code will not work, the help line opens its online return from the office and the household "
              "completes it online.", b),
           _p("7. Releases during the test", h2),
           _p("The second- and third-quarter releases are estimated from the letter B half of each quarter, the "
              "tracker's current design. The completes from both halves are kept for the office's methods work.", b),
           _p("Appendix A. The letters", h2),
           _p("<b>Letter A.</b> <i>Your household has been chosen to take part in the %s, the %s's survey of how "
              "people in Halvreth are getting on. Please ask the person in your household aged 18 or over who has the "
              "next birthday to complete the survey online, using the login code at the top of this letter. It takes "
              "about 15 minutes. If we have not heard from you in two weeks, one of our interviewers may phone to help.</i>"
              % (TRACKER, OFFICE), b),
           _p("<b>Letter B.</b> <i>Your household has been chosen to take part in the %s, the %s's survey of how "
              "people in Halvreth are getting on. One of our interviewers will phone you in the next few days and ask "
              "to speak to the person in your household aged 18 or over who has the next birthday. The interview takes "
              "about 20 minutes.</i>" % (TRACKER, OFFICE), b)]
    _doc(path, "CST-27 test plan", PEOPLE["field"], "%s | CST-27 test plan | 8 March 2027" % OFFICE, st, "2027-03-08")


# ------------------------------------------------------------------ the close-out (docx)
def close_out(path, figures):
    import docx
    from docx.shared import Pt
    doc = docx.Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    head = doc.add_paragraph()
    r = head.add_run("%s | %s" % (OFFICE, TRACKER))
    r.bold = True
    r.font.size = Pt(12)
    for k, v in (("To", "%s, Head of Survey Methodology" % PEOPLE["method"]), ("Copy", "Executive team"),
                 ("From", "%s, Fieldwork Manager" % PEOPLE["field"]), ("Date", "15 October 2027"),
                 ("Subject", "CST-27 close-out")):
        p = doc.add_paragraph()
        a = p.add_run(k + ": ")
        a.bold = True
        p.add_run(v)
        p.paragraph_format.space_after = Pt(0)
    doc.add_paragraph()
    paras = [
        "The contact strategy test ran through the second and third quarters as planned. All 76,000 addresses were "
        "allocated and mailed on their dates, and none was moved from one letter to the other.",
        "Response. The two letters drew much the same response in both quarters.",
        "Uptake and cost. About %s in every seven letter A completes came in online, against about one in five under "
        "letter B. Interviewer hours per complete under letter A were about a third of letter B's, and the extra "
        "reminder's printing and postage came to well under the hours it saved." % figures["sevenths"],
        "Fieldwork. The one interruption came in August, when the CATI system was down from Monday 9 to Thursday 12 "
        "August after the upgrade. The interviewers kept working through it, entering their interviews on the web "
        "platform with each household's login code, and the help line team joined them on the phones; the calls "
        "themselves went into a paper log.",
        "Decision. On 6 October the executive team agreed to move the tracker to letter A from the fourth quarter. The "
        "fourth quarter's first batch went out on 1 October.",
    ]
    for t in paras:
        p = doc.add_paragraph(t)
        p.paragraph_format.space_after = Pt(6)
    p = doc.add_paragraph("Thanks to the interviewers and the help line for getting through August.")
    p.paragraph_format.space_before = Pt(6)
    doc.add_paragraph(PEOPLE["field"])
    doc.core_properties.author = PEOPLE["field"]
    doc.core_properties.title = "CST-27 close-out"
    doc.save(path)
    postprocess.scrub(path, PEOPLE["field"], "2027-10-15")
    postprocess.finish_office(path, "2027-10-13T21:05:00Z", "2027-10-14T23:40:00Z", OFFICE)


# ------------------------------------------------------------------ the questionnaire codebook (xlsx)
def codebook(path):
    wb = Workbook(); ws = wb.active; ws.title = "Questions"
    ws.append(["question", "section", "wording", "answers", "web platform export", "CATI export"])
    for q, sec, text, ans in QUESTIONS:
        web = "as answered; blank if left unanswered" + ("; -7 prefer not to say" if q in ("Q14", "Q15", "Q17") else "")
        ws.append([q, sec, text, ans, web, "as answered; -8 don't know; -9 refused"])
    ws.append(["age", 5, "How old are you?", "age in years", "as typed", "AGE, as given"])
    ws.append(["sex", 5, "What is your sex?", "male; female; another", "1 male; 2 female; 3 another", "SEX: M; F; X"])
    ws.append(["qual", 5, "What is your highest qualification?",
               "no qualification; school qualification; post-school certificate or diploma; bachelor's degree or higher",
               "1 to 4, in that order", "QUAL: 1 to 4, in that order"])
    ws.append(["emp", 5, "Which of these best describes what you are doing at present?",
               "employed or self-employed; unemployed; retired; student; home or caring duties; other",
               "1 to 6, in that order", "EMP: 1 to 6, in that order"])
    ws.append(["tenure", 5, "Is your home owned or rented?", "owned outright; owned with a mortgage; rented; other",
               "1 to 4, in that order", "TENURE: 1 to 4, in that order"])
    ws.append(["hhsize", 5, "How many people usually live in your household, including you?", "number of people",
               "as typed", "HHSIZE, as given"])
    ws.append(["(end)", 6, "Thank you. Your answers have been recorded.", "", "reaching this page completes the return",
               "the interviewer codes the result"])
    w2 = wb.create_sheet("Indexes")
    w2.append(["index", "code", "question", "measure", "counted answers", "valid answers (the index's base)"])
    for code, name, item, measure, counted, base in INDEX_ROWS:
        w2.append([name, code, item, measure, counted, base])
    w2.append([])
    w2.append(["A respondent without a valid answer to the question (don't know, refused, prefer not to say, left "
               "unanswered, or not reached) is left out of that index."])
    w3 = wb.create_sheet("Sections")
    for row in (["section", "content", "questions"], [1, "Your wellbeing", "Q01 to Q04"],
                [2, "People and place", "Q10 to Q12"], [3, "Health", "Q13 to Q15"],
                [4, "Community and money", "Q16, Q17"], [5, "About you", "age, sex, qual, emp, tenure, hhsize"],
                [6, "End", ""]):
        w3.append(row)
    for sh, widths in ((ws, [10, 9, 62, 56, 36, 32]), (w2, [20, 8, 10, 22, 32, 30]), (w3, [9, 24, 40])):
        for c in sh[1]:
            c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="E6EEF1")
        for i, w in enumerate(widths, 1):
            sh.column_dimensions[get_column_letter(i)].width = w
        for row in sh.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        sh.freeze_panes = "A2"
    wb.save(path)
    postprocess.scrub(path, PEOPLE["proc"], "2027-03-22")
    postprocess.finish_office(path, "2027-03-15T02:10:00Z", "2027-03-22T04:35:00Z", OFFICE)


# ------------------------------------------------------------------ the dictionary and the readme
def dictionary(files):
    return """# Data dictionary

Dates are Halvreth dates. Halvreth time is UTC+12, and UTC+13 under daylight time.

## {issued}

The sample management system's extract for CST-27, taken 11 October 2027: one row per issued address.

| field | meaning |
|---|---|
| address_id | the address, as on the Halvreth Address Register |
| login_code | the household's login code for the online questionnaire, printed on its letters |
| quarter | the tracker quarter the address was issued in |
| release | the fortnightly batch (1 to 6) |
| region_code | region, as in {regions} |
| dwelling_type | house, flat or rural, from the register |
| letter | the letter the address was allocated: A (web-first) or B (telephone-first) |
| letter_mailed | the date the first letter was mailed |
| final_outcome | the system's final outcome code for the address (below) |
| outcome_date | the date the final outcome was recorded |

Outcome codes: 1.1 complete; 1.2 partial; 2.1 refusal; 2.2 non-contact; 2.3 other non-response; 3.1 eligibility not
settled; 4.1 vacant; 4.2 not a private dwelling; 4.3 address not found.

## {web}

The web platform's export for CST-27 logins, taken 11 October 2027: one row per session that saved an answer.

| field | meaning |
|---|---|
| response_id | the platform's id for the session |
| login_code | the login code the session was opened with |
| opened_by | respondent (opened with the login code from the letter) or staff (opened by office staff) |
| device_type | the device type the platform recorded |
| started_utc | when the session opened (UTC) |
| submitted_utc | when the return was submitted at the end of the questionnaire (UTC); blank if it was not |
| last_saved_utc | when the session last saved an answer (UTC) |
| last_section | the last questionnaire section reached (6 is the end) |
| Q01 to Q17 | answers, on the codes in {codebook} |
| age, sex, qual, emp, tenure, hhsize | section 5 answers, on the codes in {codebook} |

## {cati}

The CATI system's interview export for CST-27, taken 11 October 2027: one row per interview started.

| field | meaning |
|---|---|
| CASE_ID | the CATI case |
| ADDRESS_ID | the address |
| INTERVIEWER | the interviewer |
| START_TIME, END_TIME | the interview's start, and the time the interviewer coded its result, the CATI system's clock (Halvreth time) |
| RESULT | C complete; P partial (broken off) |
| Q01 to Q17 | answers, on the codes in {codebook} |
| AGE, SEX, QUAL, EMP, TENURE, HHSIZE | section 5 answers, on the codes in {codebook} |

## {calls}

The CATI system's call records for CST-27: one row per call attempt (CASE_ID, ATTEMPT, CALL_TIME in Halvreth time,
RESULT: NA no answer, BZ busy, AM message left, AP appointment, CB call back, RF refusal, WN wrong number, IC interview
complete, IP interview partial).

## {published}

The tracker's published headline estimates from the first quarter of 2019 to the second quarter of 2027, telephone-first
throughout: quarter, index, estimate, its 95 per cent margin of error, completes, the percentage of completes made
online, and the release date.

## {regions}

Regions, with the register's 2027 counts of households and adults.

## {codebook}

Sheet Questions: each question's wording, answers and export codes in each system. Sheet Indexes: how each headline
index is scored. Sheet Sections: the questionnaire's sections.
""".format(**files)


def readme(files):
    return """# CST-27 files

The methods team's copy of the 2027 contact strategy test's files, with the standard and plan they were run under,
compiled by {proc} (Data Processing Lead) on Monday 18 October 2027.

- `{standard}`: Methods Standard MS-14, continuity of the {tracker} through the move to web-first collection.
- `{plan}`: the CST-27 test plan.
- `{close}`: the fieldwork manager's close-out.
- `{issued}`: the sample management system's issued addresses and outcomes.
- `{web}`: the web platform's export.
- `{cati}`: the CATI system's interview export.
- `{calls}`: the CATI system's call records.
- `{codebook}`: the questionnaire codebook, version 9.
- `{published}`: the tracker's published headline estimates, 2019 to the second quarter of 2027.
- `{regions}`: regions.
- `DATA_DICTIONARY.md`: the files and their fields.
""".format(proc=PEOPLE["proc"], tracker=TRACKER, **files)
