"""The golden deliverables: mwt_mode_mix_band_2027q4.xlsx and mwt_mode_mix_band_2027q4.png."""
import os

import matplotlib
import matplotlib.ticker

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import params as P
import postprocess

XLSX = "mwt_mode_mix_band_2027q4.xlsx"
PNG = "mwt_mode_mix_band_2027q4.png"
INK = "#24323C"
BIND = "#B5462F"
BAR = "#7C97A8"
HEAD = PatternFill("solid", fgColor="DCE6EB")
THIN = Border(bottom=Side(style="thin", color="9AA8B0"))


def _head(ws, row, values, widths=None):
    for j, v in enumerate(values, 1):
        c = ws.cell(row=row, column=j, value=v)
        c.font = Font(bold=True)
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="bottom")
        c.border = THIN
    if widths:
        for j, w in enumerate(widths):
            ws.column_dimensions[chr(65 + j)].width = w


def workbook(path, A):
    wb = Workbook()
    ws = wb.active
    ws.title = "Band"
    ws["A1"] = "Halvreth Wellbeing Tracker: mode-mix monitor band from the fourth quarter of 2027"
    ws["A1"].font = Font(bold=True, size=13)
    rows = [("Band", "%d percentage points of web share" % A["band"]),
            ("Binding index", "%s (%s): its widest band within tolerance is the narrowest of the twelve, %s"
             % (A["bind_name"], A["bind"], A["bind_range"])),
            ("Next", "%s (%s): %s" % (A["runner_name"], A["runner"], A["runner_range"])),
            ("Mode effects measured on", "CST-27 completes, second and third quarters, letters A and B: %s"
             % format(A["completes"], ",")),
            ("Read through", "the random letter allocation: each index's difference between the letter A and letter B "
                             "completes, divided by the letters' difference in the share of completes made online"),
            ("Standard", "MS-14 section 5: the widest whole-point band keeping every headline index's mode-mix "
                         "artefact, at the central estimate of its mode effect, within its continuity tolerance"),
            ("Prepared", "20 October 2027, before the fourth-quarter estimates are processed")]
    for k, (a, b) in enumerate(rows, 3):
        ws.cell(row=k, column=1, value=a).font = Font(bold=True)
        c = ws.cell(row=k, column=2, value=b)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 104

    w2 = wb.create_sheet("Index evidence")
    _head(w2, 1, ["Index", "Code", "Units", "Letter A less letter B", "Online share, A less B (points)", "Mode effect",
                  "Tolerance", "Widest band (points of web share)"], [22, 7, 18, 14, 15, 12, 11, 16])
    for k, r in enumerate(A["indexes"], 2):
        vals = [r["name"], r["code"], r["units"], r["itt"], r["dw"], r["effect"], r["tol"], r["widest"]]
        for j, v in enumerate(vals, 1):
            c = w2.cell(row=k, column=j, value=v)
            if j in (4, 6):
                c.number_format = "0.00" if r["kind"] == "score" else "0.0"
            elif j == 5:
                c.number_format = "0.0"
            elif j == 7:
                c.number_format = "0.00"
            elif j == 8:
                c.number_format = "0.0"
            if r["code"] == A["bind"]:
                c.font = Font(bold=True, color="9C3A26")
    w2.freeze_panes = "C2"
    n = len(A["indexes"]) + 3
    w2.cell(row=n, column=1, value="Mode effect: web less telephone, points on the 0 to 10 scores and percentage points "
                                   "on the percentage indexes. Widest band: 100 times the tolerance over the mode "
                                   "effect's size. Two-stage least squares with the letter as the instrument and the "
                                   "recorded characteristics as controls sets the same band.")

    w3 = wb.create_sheet("Test response")
    _head(w3, 1, ["Quarter", "Letter", "Eligible addresses", "Completes", "Response rate (%)"], [11, 22, 16, 12, 16])
    for k, r in enumerate(A["response"], 2):
        for j, v in enumerate([r["quarter"], r["letter_name"], r["eligible"], r["completes"], r["rate"]], 1):
            c = w3.cell(row=k, column=j, value=v)
            if j in (3, 4):
                c.number_format = "#,##0"
            if j == 5:
                c.number_format = "0.0"
    w3.cell(row=7, column=1, value="Completes under CST-27 section 5; eligible addresses are those issued less outcome "
                                   "codes 4.1 to 4.3, an address with a complete counted eligible whatever its code.")

    w4 = wb.create_sheet("Valid answers")
    _head(w4, 1, ["Index", "Code", "Question", "Online", "By telephone"], [22, 7, 10, 12, 14])
    for k, r in enumerate(A["valid"], 2):
        for j, v in enumerate([r["name"], r["code"], r["item"], r["web"], r["cati"]], 1):
            c = w4.cell(row=k, column=j, value=v)
            if j in (4, 5):
                c.number_format = "#,##0"
    w4.cell(row=15, column=1, value="Test completes giving the index's question an answer on its scale (codebook v9, "
                                    "sheet Indexes); the 9 to 12 August staff-entry interviews counted by telephone.")

    w5 = wb.create_sheet("Notes")
    notes = ["How the test's completes were built (CST-27 section 5):",
             "Web returns that reached section 6 and were submitted by 23:59 Halvreth time on the quarter's last day "
             "(UTC+12 on 30 June; UTC+13 on 30 September, daylight time having begun on 26 September).",
             "Telephone interviews coded C. An address counts once, at the complete submitted first, the CATI clock "
             "read as Halvreth time.",
             "Returns giving an age under 18 are not the selected adult's and do not count.",
             "Staff-entry returns of 9 to 12 August are the telephone interviews keyed through the web platform during "
             "the CATI outage (close-out, 15 October 2027) and count as telephone.",
             "Sources: the CST-27 exports and issued addresses (extracts of 11 October 2027), MS-14, the CST-27 plan, "
             "codebook v9."]
    for k, t in enumerate(notes, 1):
        c = w5.cell(row=k, column=1, value=t)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if k == 1:
            c.font = Font(bold=True)
    w5.column_dimensions["A"].width = 120
    for sh in wb.worksheets:
        sh.page_setup.orientation = "landscape"
        sh.page_setup.paperSize = sh.PAPERSIZE_A4
        sh.page_setup.fitToWidth = 1
        sh.page_setup.fitToHeight = 0
        sh.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(path)
    postprocess.scrub(path, "Survey Methodology, Halvreth Statistics Office", "2027-10-20")
    postprocess.finish_office(path, "2027-10-19T21:40:00Z", "2027-10-20T03:15:00Z", "Halvreth Statistics Office")


def chart(path, A):
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 9})
    rows = sorted(A["indexes"], key=lambda r: r["widest"], reverse=True)
    fig, ax = plt.subplots(figsize=(10.2, 6.7), dpi=150)
    y = range(len(rows))
    cols = [BIND if r["code"] == A["bind"] else BAR for r in rows]
    ax.barh(list(y), [r["widest"] for r in rows], color=cols, height=0.62, zorder=2)
    ax.scatter([r["adj_widest"] for r in rows], list(y), s=40, facecolors="white", edgecolors=INK, linewidths=1.2,
               zorder=3)
    ax.axvline(A["band"], color=INK, linestyle="--", linewidth=1.1, zorder=4)
    ax.text(A["band"] * 1.04, len(rows) - 0.55, "band set: %d points" % A["band"], fontsize=8.6, color=INK,
            va="center")
    ax.set_xscale("log")
    ax.set_xlim(1.5, 80)
    ax.set_xticks([2, 3, 4, 5, 6, 8, 10, 15, 20, 30, 50])
    ax.set_xticklabels(["2", "3", "4", "5", "6", "8", "10", "15", "20", "30", "50"])
    ax.set_yticks(list(y))
    ax.set_yticklabels(["%s (%s)" % (r["name"], r["code"]) for r in rows], fontsize=8.8)
    ax.set_ylim(-0.6, len(rows) - 0.2)
    ax.set_xlabel("widest departure of the quarter's web share that keeps the index within its tolerance "
                  "(percentage points, log scale)", fontsize=8.6)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.grid(axis="x", color="#E3E8EB", zorder=0)
    handles = [Line2D([], [], color=BAR, linewidth=7, label="mode effect read through the random letter allocation"),
               Line2D([], [], color=BIND, linewidth=7, label="binding index"),
               Line2D([], [], marker="o", linestyle="", markerfacecolor="white", markeredgecolor=INK, markersize=7,
                      label="web against telephone completes, recorded characteristics held fixed (households chose "
                            "their mode)")]
    ax.xaxis.set_minor_formatter(matplotlib.ticker.NullFormatter())
    ax.legend(handles=handles, loc="upper center", bbox_to_anchor=(0.42, -0.1), fontsize=7.8, frameon=False, ncol=1)
    fig.subplots_adjust(left=0.2, right=0.97, top=0.87, bottom=0.22)
    fig.text(0.012, 0.975, "Run the mode-mix monitor at a band of %d points: %s binds" % (A["band"],
             A["bind_name"].lower()), fontsize=12.5, fontweight="bold", va="top", color=INK)
    fig.text(0.012, 0.925, "Halvreth Wellbeing Tracker, headline indexes. Mode effects from the CST-27 completes "
                           "(second and third quarters of 2027, both letters); tolerances from MS-14 Table 2.",
             fontsize=8.3, va="top", color="#333333")
    fig.savefig(path, metadata={"Software": None})
    plt.close(fig)


def write(golden_dir, A):
    os.makedirs(golden_dir, exist_ok=True)
    workbook(os.path.join(golden_dir, XLSX), A)
    chart(os.path.join(golden_dir, PNG), A)
