"""task308 documents: the brokers' quotations, the board's valuation basis, the catastrophe event definitions, the 2025
claims operations report and the group programme summary (pdf), the business plan, the rebuild cost index and the
Major Loss Unit's monthly reports (xlsx)."""
import datetime as dt

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import params as P
import world as W

F = P.FILES
NAVY = colors.HexColor("#1F3A5F")


def _styles():
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontName="Helvetica", fontSize=9.5, leading=13.2, spaceAfter=5)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName="Helvetica-Bold", fontSize=13.2, leading=16.5,
                        spaceAfter=4, textColor=NAVY)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=10.0, leading=13.5,
                        spaceBefore=7, spaceAfter=3, keepWithNext=1)
    small = ParagraphStyle("small", parent=body, fontSize=8.0, leading=10.6, textColor=colors.HexColor("#4A4A4A"))
    return body, h1, h2, small


def _footer(left, right=None):
    def foot(canvas, d):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.2)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(18 * mm, 10 * mm, left)
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, (right + "   " if right else "") + "Page %d" % d.page)
        canvas.restoreState()
    return foot


def _table(rows, widths, size=8.2, align_right=(), pad=6):
    cell = ParagraphStyle("cell", fontName="Helvetica", fontSize=size, leading=size + 2.2)
    cellr = ParagraphStyle("cellr", parent=cell, alignment=2)
    head = ParagraphStyle("head", parent=cell, fontName="Helvetica-Bold")
    out = []
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            font = "Helvetica-Bold" if r == 0 else "Helvetica"
            for word in str(v).replace("&amp;", "&").split():
                # no word may be wider than its column, or the renderer splits it mid-word
                assert stringWidth(word, font, size) <= widths[c] * mm - 2 * pad, (word, widths[c])
        out.append([Paragraph(str(v), head if r == 0 else (cellr if c in align_right else cell))
                    for c, v in enumerate(row)])
    t = Table(out, colWidths=[w * mm for w in widths], repeatRows=1, hAlign="LEFT")
    t.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), pad), ("RIGHTPADDING", (0, 0), (-1, -1), pad),
                           ("LINEBELOW", (0, 0), (-1, 0), 0.6, NAVY),
                           ("LINEBELOW", (0, 1), (-1, -1), 0.25, colors.HexColor("#B8B8B8")),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EDF3")),
                           ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    return t


def _pdf(path, title, subject, story, foot, author, right=None):
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm,
                            bottomMargin=18 * mm, title=title, author=author, creator=author, subject=subject,
                            invariant=1)
    doc.build(story, onFirstPage=_footer(foot, right), onLaterPages=_footer(foot, right))


def quotes(path):
    body, h1, h2, small = _styles()
    rows = [["Region", "Lead reinsurer", "Rate on line", "Annual premium (GBP)"]]
    for r in P.REGIONS:
        rows.append([r, P.LEAD[r], "%.2f%%" % (100 * P.PREMIUM[r] / P.LIMIT), "{:,}".format(P.PREMIUM[r])])
    sec = [["Reinsurer", "Rating (S&amp;P)", "Lines offered"],
           ["Halvorsen Re", "AA-", "Severn Vale, Coastal North (lead); Pennine, Fenland (following)"],
           ["Castellan Re", "A+", "Pennine (lead); Severn Vale (following)"],
           ["Merrow Re (Lloyd's syndicate 4417)", "A+ (Lloyd's)", "Fenland (lead); Coastal North (following)"],
           ["Ostrander Ruck SE", "A", "all four regions (following)"]]
    s = [Paragraph("%s | Treaty broking" % P.BROKER, small),
         Paragraph("%s: 2027 regional catastrophe excess of loss, quotations" % P.ORG, h1),
         Paragraph("Prepared for %s by %s, %s. Our reference WCR/ARD/27/014." %
                   (P.ORG, P.BROKER, P.QUOTE_DATE.strftime("%-d %B %Y")), small), Spacer(1, 4),
         Paragraph("The layer", h2),
         Paragraph("GBP 15,000,000 each and every loss occurrence excess of GBP 10,000,000 each and every loss "
                   "occurrence, on the home buildings and contents book of one region. One full reinstatement at 100 "
                   "per cent additional premium, pro rata as to the amount of cover reinstated only, so the most the "
                   "layer pays in the year is GBP 30,000,000 and the reinstatement premium in a year is at most the "
                   "annual premium. Losses occurring during 1 January 2027 to 31 December 2027. Perils: windstorm, "
                   "flood, and freeze (escape of water caused by freezing). Loss occurrence as defined in the "
                   "company's catastrophe event definitions, which follow the London market hours clause.", body),
         Paragraph("Quotations", h2),
         Paragraph("Each quotation is for the named region's book alone. Markets have quoted all four regions so that "
                   "the company can place the one layer its 2027 plan provides for.", body),
         _table(rows, [34, 62, 26, 40], align_right=(2, 3)), Spacer(1, 6),
         Paragraph("Premium is payable in four equal instalments, the first on 1 January 2027. Reinstatement premium "
                   "falls due when a loss is paid. Quotations are firm until 15 December 2026 and subject to no "
                   "deterioration in the claims record before inception.", body),
         Paragraph("Security", h2),
         _table(sec, [62, 28, 76]), Spacer(1, 6),
         Paragraph("Ratings as at 13 November 2026. Following lines are subject to signing down if the order is "
                   "oversubscribed.", small)]
    _pdf(path, "2027 regional catastrophe excess of loss: quotations", "Reinsurance quotations", s,
         "%s | WCR/ARD/27/014 | Confidential" % P.BROKER, P.BROKER)


def basis(path):
    body, h1, h2, small = _styles()
    s = [Paragraph("%s | Board policy RC-07" % P.ORG, small),
         Paragraph("How a catastrophe layer is valued", h1),
         Paragraph("Adopted by the Board on the recommendation of the Risk Committee, March 2024. Reviewed without "
                   "change, September 2026. Owner: Head of Reinsurance.", small), Spacer(1, 4),
         Paragraph("1. Purpose", h2),
         Paragraph("When the company buys catastrophe excess of loss protection for one part of its home book, it "
                   "places the layer where it is worth most to the company. This policy sets out how that worth is "
                   "measured, so that quotations for different books are compared on one footing.", body),
         Paragraph("2. The measure", h2),
         Paragraph("A layer's value for the cover year is its expected recoveries in the cover year less the premium "
                   "quoted and the expected reinstatement premium.", body),
         Paragraph("Expected recoveries for the cover year are the average, over the sixteen years 2010 to 2025, of "
                   "what the layer would have recovered in each year had that year's windstorm, flood and freeze "
                   "catastrophes struck the region's home book as it will stand in the cover year: each catastrophe's "
                   "claims at the cover year's policies in force in each postcode district, at what repairing them "
                   "would cost in the cover year, with the layer's retention, limit, reinstatement and loss occurrence "
                   "definitions applied to each year's catastrophes.", body),
         Paragraph("The expected reinstatement premium is the average, over the same sixteen years, of the "
                   "reinstatement premium each year's recoveries would have triggered.", body),
         Paragraph("3. The records", h2),
         Paragraph("Claims are taken at their incurred value, paid plus outstanding, on the claims platform at the "
                   "extract date (31 October 2026 for the 2027 placement). The cover year's book is the business "
                   "plan's policy count for the region, spread over its postcode districts in their mid-year "
                   "proportions for the latest year of the exposure extract. The business plan states the rebuild "
                   "cost index it expects for the cover year. The index follows the price of rebuilding a standard "
                   "home in ordinary times; it does not carry what else made a catastrophe's repairs cost what they "
                   "did.", body),
         Paragraph("4. Reporting", h2),
         Paragraph("Values are reported to the Board in GBP million to one decimal place, for every book quoted.", body)]
    _pdf(path, "How a catastrophe layer is valued", "Board policy RC-07", s, "%s | Board policy RC-07 | Internal" % P.ORG,
         P.ORG)


def events(path):
    body, h1, h2, small = _styles()
    s = [Paragraph("%s | Reinsurance" % P.ORG, small),
         Paragraph("Catastrophe event definitions (loss occurrence)", h1),
         Paragraph("Version 3, effective 1 January 2024. These definitions apply to every catastrophe excess of loss "
                   "contract the company buys, and to the Board's valuation of them under policy RC-07.", small),
         Spacer(1, 4),
         _table([["Peril", "One loss occurrence is all the peril's individual losses dated within"],
                 ["Windstorm", "72 consecutive hours"],
                 ["Flood (river, surface water and storm surge)", "168 consecutive hours"],
                 ["Freeze (escape of water caused by freezing)", "168 consecutive hours"]], [80, 86]), Spacer(1, 6),
         Paragraph("The company may choose the date and time at which any period of consecutive hours begins, but no "
                   "period may begin earlier than the first individual loss arising out of the catastrophe, no two "
                   "periods may overlap, and no loss may be counted in more than one occurrence. An occurrence is "
                   "measured region by region for a regional contract.", body),
         Paragraph("Each peril's losses are measured under its own hours: a storm that brings both wind and flood "
                   "damage gives rise to a windstorm occurrence and a flood occurrence. Losses arising out of two "
                   "different storms are not added together.", body),
         Paragraph("The claims team gives a claim an event code when it declares an event, usually within a day of a "
                   "named storm or a severe flood or freeze warning. The code does not decide the occurrence: claims "
                   "with no code belong to an occurrence by their date of loss, and claims notified late are often "
                   "entered without one.", body)]
    _pdf(path, "Catastrophe event definitions", "Loss occurrence", s, "%s | Reinsurance | Version 3" % P.ORG, P.ORG)


def ops(path, claims_at_report):
    """the 2025 claims operations report, written in February 2026 from the platform as it stood on 31 January"""
    body, h1, h2, small = _styles()
    c25 = claims_at_report
    vol = c25.groupby(["region", "peril"]).size().unstack(fill_value=0)
    order = ["escape of water", "accidental damage", "theft", "fire", "subsidence", "windstorm", "flood", "freeze"]
    rows = [["Region"] + [p.capitalize() for p in order] + ["All"]]
    for r in P.REGIONS:
        v = [int(vol.loc[r, p]) if p in vol.columns else 0 for p in order]
        rows.append([r] + ["{:,}".format(x) for x in v] + ["{:,}".format(sum(v))])
    tot = [int(vol[p].sum()) if p in vol.columns else 0 for p in order]
    rows.append(["All regions"] + ["{:,}".format(x) for x in tot] + ["{:,}".format(sum(tot))])
    cash = c25.groupby(c25.peril.isin(P.CAT_PERILS))["cash"].mean()
    s = [Paragraph("%s | Home claims" % P.ORG, small),
         Paragraph("Home claims operations: 2025 annual report", h1),
         Paragraph("Prepared by Nicholas Hardy, Head of Home Claims, for the Executive Committee. February 2026.", small),
         Spacer(1, 4),
         Paragraph("1. Volumes", h2),
         Paragraph("Home claims with a date of loss in 2025, as recorded on the claims platform on 31 January 2026. "
                   "Storm Eowyn on 24 January accounted for most of the windstorm claims in Coastal North and "
                   "Pennine, and the January freeze for most of the freeze claims.", body),
         _table(rows, [22] + [17] * 8 + [14], size=7.2, align_right=tuple(range(1, 10)), pad=3), Spacer(1, 6),
         Paragraph("2. Settlement", h2),
         Paragraph("About %d per cent of windstorm, flood and freeze claims, and %d per cent of other claims, were "
                   "settled in cash at the customer's request; the rest were repaired. Repairs go through our approved "
                   "repair panel in each region, and a claim's record carries the panel firm that did the work and the "
                   "date the repair was signed off." % (round(100 * cash.get(True, 0.12)), round(100 * cash.get(False, 0.3))), body),
         Paragraph("3. The Calder and Colne book", h2),
         Paragraph("On 1 January 2021 the home policies we took on from %s moved onto our policy and claims "
                   "platforms. Their claims have been handled by our Pennine team since then." % P.TRANSFEROR, body),
         Paragraph("4. The Lowther and Pike scheme", h2),
         Paragraph("Homes insured through the %s scheme are bound for us by %s under a delegated authority and "
                   "reported to us in a monthly risk bordereau: each month's bordereau lists every term that %s's "
                   "records at the month end show in force at some time in the month, and every cancellation "
                   "processed in the month with the day it took effect. Their claims come to our claims teams like "
                   "any other, and they are counted in the exposure extract with the rest of the book." %
                   (P.SOCIETY, P.COVERHOLDER, P.COVERHOLDER.split()[0]), body),
         Paragraph("5. Service", h2),
         Paragraph("Customer satisfaction on settled claims was 84 per cent, against 86 per cent in 2024. Complaints "
                   "upheld by the Financial Ombudsman Service fell from 61 to 48. Average time to first contact after "
                   "notification was 1.6 working days.", body)]
    _pdf(path, "Home claims operations: 2025 annual report", "Claims operations", s,
         "%s | Home Claims | Annual report 2025" % P.ORG, P.ORG)


def group(path):
    body, h1, h2, small = _styles()
    s = [Paragraph("%s | Reinsurance" % P.ORG, small),
         Paragraph("Group catastrophe programme 2026: summary of cover", h1),
         Paragraph("The group's whole-account catastrophe excess of loss programme for 2026 protects the home, small "
                   "commercial and farm books together.", body),
         _table([["Layer", "Cover (GBP m)", "Reinstatements", "Rate on line"],
                 ["1", "60 xs 50", "1 at 100%", "11.8%"], ["2", "90 xs 110", "1 at 100%", "6.1%"],
                 ["3", "150 xs 200", "1 at 100%", "3.4%"]], [20, 40, 40, 30], align_right=(3,)), Spacer(1, 6),
         Paragraph("The programme renews on 1 January 2027 on the same structure. A regional layer bought for 2027 sits "
                   "inside the group retention of GBP 50 million and inures to the benefit of the group programme.",
                   body)]
    _pdf(path, "Group catastrophe programme 2026", "Summary of cover", s, "%s | Reinsurance" % P.ORG, P.ORG)


def _xl_head(ws, row, ncol):
    thin = Side(style="thin", color="1F3A5F")
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3A5F")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)


def plan(path, pol25):
    wb = Workbook()
    ws = wb.active
    ws.title = "Policies"
    ws.append(["Region", "Policies in force, mid-2025", "Policies in force, 2027 (plan)"])
    _xl_head(ws, 1, 3)
    for r in P.REGIONS:
        ws.append([r, int(pol25[r]), int(W.PLAN_2027[r] * 1000)])
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=3):
        for c in row:
            c.number_format = "#,##0"
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 28
    ws.freeze_panes = "A2"
    wa = wb.create_sheet("Assumptions")
    wa.append(["Item", "Value", "Note"])
    _xl_head(wa, 1, 3)
    i27 = round(sum(W.cost_index(2027, m) for m in range(1, 13)) / 12, 1)
    for it in [("Rebuild cost index, 2027 average", i27, "%s rebuild cost index, 2015 = 100" % P.INDEX_HOUSE),
               ("Contents claims inflation, 2027", 0.031, "annual"),
               ("Gross written premium, home, 2027 (GBP m)", 412.6, "all regions"),
               ("Expense ratio, home, 2027", 0.284, ""),
               ("Regional catastrophe layer, 2027", "one region", "GBP 15m xs 10m, one reinstatement")]:
        wa.append(list(it))
    wa["B3"].number_format = "0.0%"
    wa["B5"].number_format = "0.0%"
    wa.column_dimensions["A"].width = 42
    wa.column_dimensions["B"].width = 12
    wa.column_dimensions["C"].width = 44
    wv = wb.create_sheet("About")
    for line in ["%s: business plan 2027, home account" % P.ORG,
                 "Version 4, approved by the Board on 29 October 2026.",
                 "Prepared by the actuarial function (Anne Shah, Chief Actuary) with Finance.",
                 "Plan targets and cost assumptions for the home account in 2027, by region."]:
        wv.append([line])
    wv["A1"].font = Font(bold=True, size=11, color="1F3A5F")
    wv.column_dimensions["A"].width = 92
    wb.save(path)


def index(path, CI):
    wb = Workbook()
    ws = wb.active
    ws.title = "Index"
    ws.append(["Year", "Month", "Rebuild cost index (2015 = 100)"])
    _xl_head(ws, 1, 3)
    for r in CI.itertuples(index=False):
        ws.append([int(r.year), int(r.month), float(r.index)])
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        for c in row:
            c.number_format = "0.0"
    ws.column_dimensions["C"].width = 30
    ws.freeze_panes = "A2"
    wn = wb.create_sheet("Notes")
    for line in ["%s: house rebuilding cost index, United Kingdom, monthly." % P.INDEX_HOUSE,
                 "Base: 2015 average = 100.",
                 "Measures the cost of rebuilding a typical house after a total loss: labour, materials and plant at "
                 "market rates, excluding demolition and professional fees.",
                 "Published to September 2026. The latest three months are provisional and may be revised.",
                 "Licensed to %s for internal use." % P.ORG]:
        wn.append([line])
    wn.column_dimensions["A"].width = 110
    wb.save(path)


def mlu(path, sheets):
    wb = Workbook()
    wb.remove(wb.active)
    cols = ["MLU reference", "Policy number", "Postcode district", "Peril", "Date of loss", "Transferred to MLU",
            "Loss adjuster", "Paid to date", "Outstanding", "Incurred", "Status"]
    widths = [15, 14, 10, 18, 13, 13, 22, 14, 14, 14, 9]
    for ms, me, df in sheets:
        ws = wb.create_sheet(ms.strftime("%b %Y"))
        ws.append(["Major Loss Unit: open and closed claims, month ending %s" % me.strftime("%-d %B %Y")])
        ws["A1"].font = Font(bold=True, size=12, color="1F3A5F")
        ws.append([])
        ws.append(cols)
        _xl_head(ws, 3, len(cols))
        for r in df.itertuples(index=False):
            ws.append([r[0], int(r[1]), r[2], r[3], r[4], r[5], r[6], float(r[7]), float(r[8]), float(r[9]), r[10]])
        n = len(df) + 3
        for c in ("E", "F"):
            for row in ws["%s4:%s%d" % (c, c, n)]:
                for cell in row:
                    cell.number_format = "dd/mm/yyyy"
        for c in ("H", "I", "J"):
            for row in ws["%s4:%s%d" % (c, c, n)]:
                for cell in row:
                    cell.number_format = "#,##0.00"
        for k, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(k)].width = w
        ws.freeze_panes = "A4"
        ws.append([])
        ws.append(["%d claims listed. A claim closed in the month is listed once with its settled figures. Incurred "
                   "is paid to date plus outstanding, in GBP, at the month end." % len(df)])
    wb.save(path)
