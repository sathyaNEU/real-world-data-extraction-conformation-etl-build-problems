"""The two goldens, from figures.json: capacity_award_2027.xlsx (the evaluation and the member figures) and
capacity_award_2027.docx (the committee paper, with its chart drawn by matplotlib)."""
import datetime
import io
import json
import os
import sys

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Inches, Pt, RGBColor  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import params as P  # noqa: E402

COUNTERPARTY = {"Ketchum Ridge": "Halvard Energy Partners LLC", "Red Willow import": "Red Willow Hydro Authority",
                "Sandhill Draw": "Sandhill Draw Wind LLC", "Coulee Flats": "Coulee Flats Solar LLC"}
RESOURCE = {"Ketchum Ridge": "Gas peaker, Northgate IT", "Red Willow import": "Firm import, Red Willow tie",
            "Sandhill Draw": "Wind", "Coulee Flats": "Solar PV"}


def _thin():
    s = Side(style="thin", color="A6A6A6")
    return Border(left=s, right=s, top=s, bottom=s)


def workbook(F, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Evaluation"
    ws["A1"] = "2027 capacity RFO: evaluation under RA-3 (Rev. 4)"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = "Planning year 2027, weather years 2016 to 2025, value of lost load USD 20,000 per MWh"
    ws["A2"].font = Font(italic=True, size=9, color="555555")
    ws["A4"] = "2027 fleet expected unserved energy (MWh a year)"
    ws["E4"] = F["base_eue_mwh"]
    ws["E4"].number_format = "#,##0"
    hdr = ["Offer", "Counterparty", "Resource", "Capacity (MW)", "2027 capacity payment (USD M)",
           "EUE removed (MWh a year)", "Net value (USD M)", "Rank"]
    ws.append([])
    ws.append(hdr)
    for c in ws[6]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DCE6DE")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.border = _thin()
    order = sorted(P.OFFERS, key=lambda o: -F["net_exact"][o])
    for i, o in enumerate(P.OFFERS):
        ws.append([o, COUNTERPARTY[o], RESOURCE[o], F["offer_mw"][o], F["payment_musd"][o], F["removed_mwh"][o],
                   F["net_musd"][o], order.index(o) + 1])
        for c in ws[ws.max_row]:
            c.border = _thin()
        ws.cell(ws.max_row, 5).number_format = "0.0"
        ws.cell(ws.max_row, 6).number_format = "#,##0"
        ws.cell(ws.max_row, 7).number_format = "#,##0;-#,##0"
    ws.append([])
    ws.append(["Award", F["award"]])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.cell(ws.max_row, 2).font = Font(bold=True)
    for col, w in zip("ABCDEFGH", [20, 28, 28, 13, 18, 18, 15, 7]):
        ws.column_dimensions[col].width = w
    ws.row_dimensions[6].height = 45
    m = wb.create_sheet("Member allocation")
    m["A1"] = "2025 member figures for the cost allocation (MS-12)"
    m["A1"].font = Font(bold=True, size=12)
    m.append([])
    m.append(["Member", "2025 net energy delivered (MWh)", "2025 12-CP average (MW)"])
    for c in m[3]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DCE6DE")
        c.border = _thin()
        c.alignment = Alignment(wrap_text=True)
    for k in P.MEMBERS:
        m.append([k, F["ask_a_mwh"][k], F["ask_b_mw"][k]])
        m.cell(m.max_row, 2).number_format = "#,##0"
        m.cell(m.max_row, 3).number_format = "0.0"
        for c in m[m.max_row]:
            c.border = _thin()
    for col, w in zip("ABC", [32, 24, 20]):
        m.column_dimensions[col].width = w
    m.row_dimensions[3].height = 30
    n = wb.create_sheet("Notes")
    notes = [
        "Sources: weather_years_2027_hourly.csv (2019 delivered in kW, converted), outage_event_log_2016_2025.csv "
        "(latest revision of each event, forced events only), fleet_list_2027.xlsx (Sandoval 1 out), "
        "red_willow_tie_outages_2016_2025.csv, offer_profiles_2016_2025.parquet, the RFO offers sheet.",
        "Hours aligned on UTC: the planning table by hour ending in local prevailing time (the autumn 2* hour is the "
        "standard-time 01:00), the event log in UTC, the tie log in CST/CDT, the profiles by hour beginning in UTC.",
        "Ketchum Ridge counted in each hour at 300 MW times the share of the ten Northgate IT peakers' capability "
        "available in that hour. Every hour with a shortfall in 2016 to 2025 fell inside a Northgate curtailment, "
        "when all ten were out on lack of fuel (cause 9130), so the offer removes no unserved energy.",
        "Red Willow import counted at 300 MW in every hour the tie is in service; wind and solar at their modelled "
        "output for the hour.",
        "Member figures: highest version of each interval; delivered less received; multipliers corrected to the "
        "field's current transformer ratios from the date of each replacement; the meters found on standard time "
        "moved forward an hour from the spring change to their reset for the coincident peaks; the delivery point "
        "that moved from Lone Elm to Marais Electric is Marais's from 1 July 2025.",
    ]
    for t in notes:
        n.append([t])
        n.cell(n.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    n.column_dimensions["A"].width = 120
    wb.properties.creator = "Joseph Graham"
    wb.properties.lastModifiedBy = "Joseph Graham"
    wb.properties.created = datetime.datetime(2026, 11, 3, 10, 12)
    wb.properties.modified = datetime.datetime(2026, 11, 6, 16, 48)
    wb.save(path)


def chart(F):
    fig, ax = plt.subplots(figsize=(6.4, 3.0), dpi=150)
    names = P.OFFERS
    vals = [F["net_musd"][o] for o in names]
    cols = ["#2E6B4F" if o == F["award"] else "#9AA5A0" for o in names]
    ax.bar(range(len(names)), vals, color=cols, width=0.55)
    ax.axhline(0, color="#444444", lw=0.8)
    for i, v in enumerate(vals):
        ax.text(i, v + (3 if v >= 0 else -9), "%+d" % v, ha="center", fontsize=8.5, color="#222222")
    ax.set_xticks(range(len(names)))
    ax.set_xticklabels(["Ketchum Ridge\npeaker", "Red Willow\nfirm import", "Sandhill Draw\nwind",
                        "Coulee Flats\nsolar"], fontsize=8.5)
    ax.set_ylabel("Net value, 2027 (USD M)", fontsize=8.5)
    ax.tick_params(axis="y", labelsize=8)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.set_ylim(min(vals) - 15, max(vals) + 15)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", metadata={"Software": None})
    plt.close(fig)
    buf.seek(0)
    return buf


def paper(F, path, dp_transfer):
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    t = d.add_paragraph()
    r = t.add_run("Resource Planning Committee, 19 November 2026")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    h = d.add_paragraph()
    hr = h.add_run("Item 4: award of the 2027 capacity contract")
    hr.bold = True
    hr.font.size = Pt(14)
    p = d.add_paragraph()
    p.add_run("Recommendation. ").bold = True
    p.add_run("Award the 2027 contract to Red Willow Hydro Authority for the 300 MW firm import over the Red Willow "
              "tie. Under RA-3 it removes %s MWh of the 2027 fleet's %s MWh of expected unserved energy, a net value "
              "of USD %d million after its USD %.1f million capacity payment."
              % (format(F["removed_mwh"]["Red Willow import"], ","), format(F["base_eue_mwh"], ","),
                 F["net_musd"]["Red Willow import"], F["payment_musd"]["Red Willow import"]))
    d.add_paragraph("The 2027 fleet's shortfalls all fall in Northgate curtailments. In 2016 to 2025 every hour with "
                    "unserved energy came while Northgate had curtailed interruptible transportation and the ten "
                    "peakers at Elk Prairie, Torrance and Wexford were out together on lack of fuel. Ketchum Ridge "
                    "would take gas on the same interruptible service, so in those hours it would have been out with "
                    "them: it removes no unserved energy, and its USD %.1f million payment buys nothing under the "
                    "standard. Rated at its 4.5 per cent EFORd in every hour it would look like the best offer, "
                    "which is the reading the committee should not rely on." % F["payment_musd"]["Ketchum Ridge"])
    tb = d.add_table(rows=1, cols=5)
    tb.style = "Light Grid Accent 1"
    for i, txt in enumerate(["Offer", "MW", "Payment (USD M)", "EUE removed (MWh)", "Net value (USD M)"]):
        tb.rows[0].cells[i].text = txt
    for o in P.OFFERS:
        row = tb.add_row().cells
        row[0].text = o
        row[1].text = str(F["offer_mw"][o])
        row[2].text = "%.1f" % F["payment_musd"][o]
        row[3].text = format(F["removed_mwh"][o], ",")
        row[4].text = "%d" % F["net_musd"][o]
    d.add_paragraph()
    d.add_picture(chart(F), width=Inches(5.8))
    cap = d.paragraphs[-1]
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.add_paragraph("Figure 1. Net value of each offer in 2027: the EUE it removes at USD 20,000 per MWh less its "
                    "capacity payment.").runs[0].font.size = Pt(9)
    d.add_paragraph("Sandhill Draw is the only other offer with a positive value, USD %d million: its output in the "
                    "curtailment hours is low, as the contracted farms' is. Coulee Flats produces nothing in the "
                    "winter evening hours that carry the shortfall and comes out at USD %d million."
                    % (F["net_musd"]["Sandhill Draw"], F["net_musd"]["Coulee Flats"]))
    h2 = d.add_paragraph()
    h2.add_run("Member allocation basis (MS-12)").bold = True
    d.add_paragraph("The contract's cost is allocated on the members' 2025 12-CP averages. The 2025 figures below "
                    "correct the meter data export for the current transformer replacements the MDM picked up late "
                    "and for the meters that stayed on standard time after the spring change.")
    mt = d.add_table(rows=1, cols=3)
    mt.style = "Light Grid Accent 1"
    for i, txt in enumerate(["Member", "2025 net energy (MWh)", "2025 12-CP average (MW)"]):
        mt.rows[0].cells[i].text = txt
    for k in P.MEMBERS:
        row = mt.add_row().cells
        row[0].text = k
        row[1].text = format(F["ask_a_mwh"][k], ",")
        row[2].text = "%.1f" % F["ask_b_mw"][k]
    d.add_paragraph()
    d.add_paragraph("Prepared by Joseph Graham, Resource Planning. Workbook: capacity_award_2027.xlsx.").runs[0].font.size = Pt(9)
    cp = d.core_properties
    cp.author = "Joseph Graham"
    cp.last_modified_by = "Joseph Graham"
    cp.created = datetime.datetime(2026, 11, 3, 9, 30)
    cp.modified = datetime.datetime(2026, 11, 6, 17, 5)
    cp.revision = 6
    cp.title = "Item 4: award of the 2027 capacity contract"
    d.save(path)


def main(out_dir):
    F = json.load(open(os.path.join(HERE, "figures.json")))
    os.makedirs(out_dir, exist_ok=True)
    import build
    workbook(F, os.path.join(out_dir, "capacity_award_2027.xlsx"))
    paper(F, os.path.join(out_dir, "capacity_award_2027.docx"), None)
    build.office(os.path.join(out_dir, "capacity_award_2027.xlsx"), "2026-11-03T10:12:00Z", "2026-11-06T16:48:00Z",
                 "Joseph Graham")
    build.office(os.path.join(out_dir, "capacity_award_2027.docx"), "2026-11-03T09:30:00Z", "2026-11-06T17:05:00Z",
                 "Joseph Graham")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(HERE), "golden"))
