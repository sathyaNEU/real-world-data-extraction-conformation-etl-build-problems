"""Independent verifier for task151: reads only the shipped bundle (target/), the goldens and figures.json, on its own
code path (pandas frames, pdftotext, no generator module), and recomputes the accrual and each player's figure, the
default path's figure, the write-up's step figures, the look-back for four bases, and the ask (the playing-time bonuses
other than plate appearances paid for 2023 to 2025)."""
from __future__ import annotations

import datetime as dt
import json
import math
import os
import re
import subprocess
import sys
import unicodedata

import numpy as np
import openpyxl
import pandas as pd

MIN = {"7-day": 7, "10-day": 10, "15-day": 15, "60-day": 60}


def txt(path):
    return re.sub(r"\s+", " ", subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True,
                                              check=True).stdout.decode("utf-8"))


def ascii_upper(s):
    return "".join(c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c)).upper()


def batting(target):
    leg = pd.read_csv(os.path.join(target, "lsf_batting_2021_2024.txt"), sep="|", dtype=str, keep_default_na=False)
    xw = pd.read_csv(os.path.join(target, "lsf_player_id_crosswalk.csv"), dtype=str)
    leg = leg.merge(xw[["legacy_id", "player_id"]], left_on="PLAYER_CD", right_on="legacy_id", how="left")
    assert leg.player_id.notna().all()
    for c in ("AB", "BB", "HBP", "SH", "SF", "CI", "LINEUP", "SEQ"):
        leg[c] = leg[c].astype(int)
    leg["pa"] = leg.AB + leg.BB + leg.HBP + leg.SH + leg.SF + leg.CI
    leg["season"] = pd.to_datetime(leg.GAME_DT, format="%m/%d/%Y").dt.year
    leg["slot"] = leg.LINEUP
    leg["starter"] = leg.SEQ == 1
    leg["pid"] = leg.player_id.astype(int)
    leg["game"] = leg.GAME_ID
    new = pd.read_csv(os.path.join(target, "lsf_batting_2025_2026.csv"))
    new["season"] = pd.to_datetime(new.official_date).dt.year
    new["slot"] = new.batting_order // 100
    new["starter"] = new.batting_order % 100 == 0
    new["pid"] = new.player_id
    new["game"] = new.game_pk.astype(str)
    cols = ["season", "game", "pid", "slot", "starter", "pa"]
    return pd.concat([leg[cols], new[cols]], ignore_index=True)


def plan(target):
    wb = openpyxl.load_workbook(os.path.join(target, "hvbd_winter_depth_charts_2021_2026.xlsx"), data_only=True)
    out = []
    for ws in wb.worksheets:
        season = int(re.search(r"(\d{4})", ws.title).group(1))
        rows = list(ws.iter_rows(values_only=True))
        h = next(i for i, r in enumerate(rows) if r and r[0] == "Tm")
        hdr = list(rows[h])
        for r in rows[h + 1:]:
            if r and r[0]:
                d = dict(zip(hdr, r))
                out.append(dict(season=season, role=d["Role"], pos=d["Pos"], name=d["Player"], pid=int(d["LG_ID"]),
                                ord=d["Ord"]))
    return pd.DataFrame(out)


def five(target):
    t = txt(os.path.join(target, "term_sheets_2027_free_agent_hitters.pdf"))
    wb = openpyxl.load_workbook(os.path.join(target, "ral_depth_chart_2027.xlsx"), data_only=True)
    slots = {}
    on = False
    for r in wb.active.iter_rows(values_only=True):
        if r and r[0] == "Lineup":
            on = True
            continue
        if r and r[0] in ("Bench", "Rotation", "Bullpen"):
            on = False
        if on and r and isinstance(r[0], int):
            slots[r[2]] = r[0]
    out = []
    for name, slot in slots.items():
        i = t.find(name + " (")
        if i < 0:
            continue
        seg = t[i:i + 2500]
        m = re.search(r"2027 season: \$([\d,]+) each for (\d+), (\d+) and (\d+) plate appearances", seg)
        sg = re.search(r"Signed\s+([A-Z][a-z]+ \d{1,2}, \d{4})", seg)
        out.append(dict(name=name, slot=slot, amount=int(m.group(1).replace(",", "")),
                        T=[int(m.group(2)), int(m.group(3)), int(m.group(4))],
                        signed=dt.datetime.strptime(sg.group(1), "%B %d, %Y").date()))
    return sorted(out, key=lambda x: x["signed"])


def accrual(target):
    bl = batting(target)
    st = bl[bl.starter]
    rate = (st.groupby("slot").pa.sum() / st.groupby("slot").size()).to_dict()
    starts = st.groupby(["season", "pid"]).size()
    pl = plan(target)
    reg = pl[(pl.role == "Lineup") & (pl.pos != "C")]
    s = np.array([starts.get((a, b), 0) for a, b in zip(reg.season, reg.pid)], float)
    F = five(target)
    per = {}
    for f in F:
        x = s * rate[f["slot"]]
        per[f["name"]] = sum(f["amount"] * float((x >= T).mean()) for T in f["T"])
    return per, sum(per.values()), dict(lines=len(bl), planned=len(reg), rate=rate, s=s, F=F, starts=starts, pl=pl)


def default(target, D):
    pr = pd.read_csv(os.path.join(target, "hvbd_hitter_projections_2022_2027.csv"))
    bl = batting(target)
    pa = bl.groupby(["season", "pid"]).pa.sum()
    reg = D["pl"][(D["pl"].role == "Lineup") & (D["pl"].pos != "C")]
    past = pr[pr.season <= 2026].merge(reg[["season", "pid"]], left_on=["season", "lg_id"], right_on=["season", "pid"])
    err = np.array([pa.get((a, b), 0) - p for a, b, p in zip(past.season, past.lg_id, past.pa)], float)
    sd = err.std()
    last = D["pl"][D["pl"].season == 2026]
    tot = 0.0
    for f in D["F"]:
        pid = int(last[last.name == f["name"]].pid.iloc[0])
        p = int(pr[(pr.season == 2027) & (pr.lg_id == pid)].pa.iloc[0])
        tot += sum(f["amount"] * (1 - 0.5 * (1 + math.erf((T - p) / (sd * math.sqrt(2))))) for T in f["T"])
    return tot


def paid(target):
    """The playing-time bonuses other than plate appearances the club paid for 2023 to 2025, by player and season: for
    every player whose register row carries other playing-time terms, the season-end (October) payments and the
    January payments after them, any bonus code, all statuses summed, less an award term the October side exceeds the
    January side by, and less, in the October after a contract's last season, what its salary lines paid short of the
    register's salary over its seasons; a legacy line is the player whose name it carries, an HRIS line the register's
    ID."""
    pay = pd.read_csv(os.path.join(target, "ral_payroll_player_earnings_2023_2026.csv"), dtype={"employee_id": str})
    wb = openpyxl.load_workbook(os.path.join(target, "ral_player_contract_register_2023_2026.xlsx"), data_only=True)
    rows = list(wb["Register"].iter_rows(values_only=True))
    h = next(i for i, r in enumerate(rows) if r and r[0] == "Player")
    reg = pd.DataFrame([dict(zip(rows[h], r)) for r in rows[h + 1:] if r and r[0]])
    oth = reg[reg["Other playing-time bonuses (each season)"].notna()].copy()
    def span(s):
        a = str(s).split("-")
        return set(range(int(a[0]), int(a[-1]) + 1))
    seasons = {}
    for _, r in oth.iterrows():
        seasons.setdefault(r["Player"], set()).update(span(r["Seasons"]))
    award = {}
    for _, r in reg.iterrows():
        if isinstance(r["Award bonuses"], str):
            award.setdefault(r["Player"], set()).update(
                float(x.replace(",", "")) for x in re.findall(r"\$([\d,]+) for", r["Award bonuses"]))
    by_legacy = {ascii_upper(n.split(", ")[0] + " " + n.split(", ")[1]): n for n in reg["Player"]}
    by_id = dict(zip(reg["Employee ID"].astype(str), reg["Player"]))
    legacy_line = pay.employee_id.str.fullmatch(r"\d+")
    pay["player"] = np.where(legacy_line, pay.employee_name.map(by_legacy), pay.employee_id.map(by_id))
    pay["date"] = pd.to_datetime(pay.pay_date)
    sal = pay[pay.earning_code.isin(["E01", "SAL"])].groupby([pay.player, pay.date.dt.year]).amount.sum()
    held = {}
    for _, r in oth.iterrows():
        yrs = span(r["Seasons"])
        gap = sum(float(r["Salary %d" % y]) - float(sal.get((r["Player"], y), 0.0)) for y in yrs
                  if 2023 <= y <= 2026 and pd.notna(r.get("Salary %d" % y)))
        if gap > 0.5:
            held[(max(yrs), r["Player"])] = gap
    pt = pay[pay.player.isin(seasons) & pay.earning_code.isin(["E40", "E45", "BPA"]) &
             pay.date.dt.month.isin([10, 1])].copy()
    pt["season"] = np.where(pt.date.dt.month == 10, pt.date.dt.year, pt.date.dt.year - 1)
    pt["side"] = np.where(pt.date.dt.month == 10, "oct", "jan")
    g = pt.groupby(["season", "player", "side"]).amount.sum().unstack(fill_value=0.0)
    # a contract season with an October side and no January side, on a contract running past the next season: the
    # next October side exceeds the January side after it by this season's October half (net of a listed award), and
    # that excess was this season's January half, paid late
    g = g.copy()
    for player, yrs in seasons.items():
        for y in sorted(yrs):
            if y == max(yrs) or y + 1 == max(yrs) or y + 1 not in yrs:
                continue
            if (y, player) not in g.index or (y + 1, player) not in g.index:
                continue
            a, b = g.loc[(y, player)], g.loc[(y + 1, player)]
            if a.get("jan", 0.0) > 0.5 or a.get("oct", 0.0) < 0.5:
                continue
            excess = round(b.get("oct", 0.0) - held.get((y + 1, player), 0.0) - b.get("jan", 0.0), 2)
            halves = {round(a.get("oct", 0.0) - w, 2) for w in award.get(player, set()) | {0.0}}
            if excess > 0.5 and excess in halves:
                g.loc[(y + 1, player), "oct"] = b.get("oct", 0.0) - excess
                g.loc[(y, player), "jan"] = excess
    out = {}
    for (season, player), r in g.iterrows():
        if season not in (2023, 2024, 2025) or season not in seasons[player]:
            continue
        o = r.get("oct", 0.0) - held.get((season, player), 0.0)
        j = r.get("jan", 0.0)
        diff = round(o - j, 2)
        v = o + j - (diff if j > 0.5 and diff > 0.5 and diff in award.get(player, set()) else 0.0)
        if abs(v) > 0.5:
            out[(int(season), player)] = v
    return out


def lookback(target, D):
    """FIN-14 4.4 for four bases, on this path: the regulars' seasons of starts at the slot, the vendor's projection
    with a normal band from its misses, the player's own last three seasons, every lineup regular's plate appearances;
    and the near rival that reads the slot's rate with its substitutes' plate appearances, per start."""
    bl = batting(target)
    st = bl[bl.starter]
    starts = st.groupby(["season", "pid"]).size()
    pa = bl.groupby(["season", "pid"]).pa.sum()
    sp = st.groupby(["season", "slot"]).pa.sum()
    sn = st.groupby(["season", "slot"]).size()
    sp_all = bl.groupby(["season", "slot"]).pa.sum()
    pr = pd.read_csv(os.path.join(target, "hvbd_hitter_projections_2022_2027.csv"))
    proj = {(a, b): c for a, b, c in zip(pr.season, pr.lg_id, pr.pa)}
    lu = D["pl"][D["pl"].role == "Lineup"].copy()
    lu["c"] = lu.pos == "C"
    lu["st"] = [starts.get((a, b), 0) for a, b in zip(lu.season, lu.pid)]
    lu["pa"] = [pa.get((a, b), 0) for a, b in zip(lu.season, lu.pid)]
    thr = np.array([450, 500, 550, 600, 650], float)
    sq = {k: [] for k in ("golden", "vendor", "own", "nine", "thin")}
    for Y in (2023, 2024, 2025, 2026):
        past = lu[lu.season < Y]
        rate = {s: sp[[(y, s) for y in range(2021, Y)]].sum() / sn[[(y, s) for y in range(2021, Y)]].sum()
                for s in range(1, 10)}
        rate_sub = {s: sp_all[[(y, s) for y in range(2021, Y)]].sum() / sn[[(y, s) for y in range(2021, Y)]].sum()
                    for s in range(1, 10)}
        est = np.array([pa.get((Y - 1, p), 0) >= 100 and pa.get((Y - 2, p), 0) >= 100 for p in lu.pid])
        tests = lu[(lu.season == Y).values & est]
        for _, t in tests.iterrows():
            out = (t.pa >= thr).astype(float)
            pool = past[past.c == t.c]
            g = np.array([(pool.st.values * rate[int(t.ord)] >= x).mean() for x in thr])
            e = np.array([q - proj[(a, b)] for a, b, q in zip(pool.season, pool.pid, pool.pa) if (a, b) in proj], float)
            mu = proj[(Y, t.pid)]
            v = np.array([1 - 0.5 * (1 + math.erf((x - mu) / (e.std() * math.sqrt(2)))) for x in thr])
            own = np.array([pa.get((Y - j, t.pid), 0) for j in (1, 2, 3)], float)
            own = own[own > 0]
            o = np.array([(own >= x).mean() for x in thr])
            n9 = np.array([(past.pa.values >= x).mean() for x in thr])
            th = np.array([(pool.st.values * rate_sub[int(t.ord)] >= x).mean() for x in thr])
            for k, p in (("golden", g), ("vendor", v), ("own", o), ("nine", n9), ("thin", th)):
                sq[k].append((p - out) ** 2)
    return {k: float(np.mean(v)) for k, v in sq.items()}, len(sq["golden"])


def r(x, unit):
    return int(math.floor(x / unit + 0.5)) * unit


def main(target, golden, fig_path):
    per, tot, D = accrual(target)
    fig = json.load(open(fig_path, encoding="utf-8"))
    bad = []
    if "$%s" % format(r(tot, 250000), ",") != fig["call"]["total"]:
        bad.append(("total", tot))
    for n, v in per.items():
        if "$%s" % format(r(v, 100000), ",") != fig["critical_components"]["per_player"][n]:
            bad.append((n, v))
    if "{:,}".format(D["lines"]) != fig["steps"]["batting_lines"] or "{:,}".format(D["planned"]) != fig["steps"]["planned_seasons"]:
        bad.append(("steps", D["lines"], D["planned"]))
    if "%.2f" % D["rate"][1] != fig["steps"]["rate_first"] or "%.2f" % D["rate"][9] != fig["steps"]["rate_ninth"]:
        bad.append(("rates", D["rate"][1], D["rate"][9]))
    if str(int((D["s"] < 100).sum())) != fig["steps"]["under_100"].replace(",", "") or \
            str(int((D["s"] >= 155).sum())) != fig["steps"]["full_155"].replace(",", ""):
        bad.append(("starts counts",))
    dft = default(target, D)
    if "$%s" % format(r(dft, 250000), ",") != fig["rivals"]["default"]:
        bad.append(("default", dft))
    B = paid(target)
    want_b = {}
    for blk in fig["deliverables"]["paid"]:
        for n, val in blk["rows"]:
            want_b[(int(blk["season"]), n)] = val
    got_b = {}
    for (y, n), val in B.items():
        last, first = n.split(", ")
        got_b[(y, "%s %s" % (first, last))] = "$%s" % format(int(round(val)), ",")
    if got_b != want_b:
        bad.append(("paid", got_b, want_b))
    for blk in fig["deliverables"]["paid"]:
        t = sum(val for (y, n), val in B.items() if y == int(blk["season"]))
        if "$%s" % format(int(round(t)), ",") != blk["total"]:
            bad.append(("paid total", blk["season"], t))
    bl = batting(target)
    pa_all = bl.groupby(["season", "pid"]).pa.sum()
    lu = D["pl"][D["pl"].role == "Lineup"]
    nine = np.array([pa_all.get((a, b), 0) for a, b in zip(lu.season, lu.pid)], float)
    at_hand = sum(f["amount"] * float((nine >= T).mean()) for f in D["F"] for T in f["T"])
    if "$%s" % format(r(at_hand, 250000), ",") != fig["rivals"]["default_at_hand"]:
        bad.append(("default at hand", at_hand))
    LB, n = lookback(target, D)
    lab = {"golden": "Non-catcher regulars' seasons of starts at the player's 2027 slot (used)",
           "vendor": "Vendor's projection with a normal band from its misses",
           "own": "The player's own last three seasons",
           "nine": "Every lineup regular's plate appearances, catchers in, any slot"}
    for k, v in LB.items():
        if k in lab and "%.4f" % v != fig["lookback"]["scores"][lab[k]]:
            bad.append(("look-back", k, v))
    if "{:,}".format(n) != fig["lookback"]["hitter_seasons"] or min(LB, key=LB.get) != "golden":
        bad.append(("look-back hitters or ranking", n, LB))
    wb = openpyxl.load_workbook(os.path.join(golden, "pa_bonus_accrual_2027.xlsx"), data_only=True)
    cells = [c for row in wb["Accrual"].iter_rows(values_only=True) for c in row if isinstance(c, (int, float))]
    if r(tot, 250000) not in cells:
        bad.append(("workbook total",))
    if bad:
        print("verify.py: FAIL", bad)
        raise SystemExit(1)
    print("verify.py: the accrual %s (unrounded %.0f), five player figures, the vendor's band %s, the look-back's pick "
          "among the bases at hand %s, the step figures, the FIN-14 4.4 look-back for four bases (%s hitter-seasons; the "
          "substitutes' slot rate loses, %.5f against %.5f) and the other playing-time bonuses paid for 2023 to 2025 "
          "reproduce from the bundle on an independent path" % (fig["call"]["total"], tot, fig["rivals"]["default"], fig["rivals"]["default_at_hand"],
                                      fig["lookback"]["hitter_seasons"], LB["thin"], LB["golden"]))


if __name__ == "__main__":
    t = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "target")
    g = sys.argv[2] if len(sys.argv) > 2 else os.path.join(os.path.dirname(t), "golden")
    f = sys.argv[3] if len(sys.argv) > 3 else os.path.join(os.path.dirname(os.path.abspath(__file__)), "figures.json")
    main(t, g, f)
