"""The golden deliverables, computed from the shipped pack only (calc): pa_bonus_accrual_2027.xlsx (the accrual by
bonus line, player and in total; the playing-time bonuses other than plate-appearance bonuses paid for the 2023 to 2025
seasons) and pa_bonus_accrual_memo_2027.docx (the one-page note for the audit file), with generator/figures.json and
submission.md written from the same objects."""
from __future__ import annotations

import datetime as dt
import io
import json
import os
import re
import zipfile

import docx
import openpyxl
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import calc
import common as K

F_BOOK = "pa_bonus_accrual_2027.xlsx"
F_MEMO = "pa_bonus_accrual_memo_2027.docx"
STAMP = dt.datetime(2026, 12, 11, 17, 35, 0)
AUTHOR = K.PEOPLE["controller"]
ORD = {1: "first", 2: "second", 3: "third", 4: "fourth", 5: "fifth", 6: "sixth", 7: "seventh", 8: "eighth", 9: "ninth"}


def compute(target):
    P = calc.load(target)
    T = calc.seasons(P)
    R, G, NEAR, ACC = calc.rivals(P, T)
    LB, n_lb, _ = calc.lookback(P, T)
    names = {}
    for r in calc.read_csv(os.path.join(target, calc.F_PAY)):
        if r["employee_id"].isdigit():
            names.setdefault(r["employee_id"], set()).add(r["employee_name"])
    reissued = sorted(e for e, n in names.items() if len(n) > 1)
    return dict(P=P, T=T, G=G, R=R, NEAR=NEAR, ACC=ACC, LB=LB, n_lb=n_lb, B=calc.ask(P),
                BD=calc.ask(P, detail=True), reissued=reissued)


def usd(x):
    return "${:,.0f}".format(x)


def surname(name):
    return name.split(" ", 1)[1]


def club_name(P, code):
    rows = calc.read_csv(os.path.join(P.target, "league_club_directory.csv"))
    return [r["club_name"] for r in rows if r["abbrev"] == code][0]


def paid_rows(R):
    """Each season's rows in the register's order (surname), and the season's total."""
    B = R["B"]
    out = []
    for y in calc.ASK_SEASONS:
        rows = sorted(((n, v) for (yy, n), v in B.items() if yy == y and v), key=lambda kv: kv[0])
        out.append((y, rows, sum(v for _, v in rows)))
    return out


def terms_for(R, name, y):
    for r in calc.register(R["P"].target):
        if r["Player"] == name and r[calc.OTHER_COL] and y in calc.seasons_of(r["Seasons"]):
            return r[calc.OTHER_COL]
    return ""


def plain(name_hris):
    last, first = name_hris.split(", ")
    return "%s %s" % (first, last)


# -------------------------------------------------------------------------------------------------------- figures
def figures(R):
    G, P, T = R["G"], R["P"], R["T"]
    F = G["five"]
    rt = G["rate"]
    st = G["starts"]
    rv = R["R"]
    f = dict(
        call=dict(total=usd(G["total_r"]), total_unrounded=usd(G["total"]),
                  sum_of_player_figures=usd(sum(G["per_r"].values()))),
        critical_components=dict(
            per_player={f_["name"]: usd(G["per_r"][f_["name"]]) for f_ in F},
            slots_2027={f_["name"]: ORD[f_["slot"]] for f_ in F},
            slots_old={f_["name"]: ORD[f_["old"]] for f_ in F}),
        steps=dict(
            batting_lines="{:,}".format(len(P.lines)),
            planned_seasons="{:,}".format(G["n"]),
            first_season=str(min(T.years)), last_season=str(max(T.years)),
            rate_first="%.2f" % rt[1], rate_ninth="%.2f" % rt[9],
            rates={ORD[f_["slot"]]: "%.2f" % rt[f_["slot"]] for f_ in F},
            under_100="{:,}".format(int((st < 100).sum())),
            full_155="{:,}".format(int((st >= 155).sum())),
            thresholds=["100", "155"],
            rate_first_1dp="%.1f" % rt[1], rate_ninth_1dp="%.1f" % rt[9],
            line_probabilities={f_["name"]: ["%.1f%%" % (100 * p) for a, t, p in G["lines"][f_["name"]]] for f_ in F},
            per_player_unrounded={f_["name"]: usd(G["per"][f_["name"]]) for f_ in F}),
        references=dict(fin14_sections=["4.2", "4.3", "4.4", "6"], working_paper="CF-14.4", rounding_total="$250,000",
                        rounding_player="$100,000"),
        lookback=dict(hitter_seasons="{:,}".format(R["n_lb"]), seasons="2023 to 2026", min_pa=str(calc.LB_MIN_PA),
                      thresholds=[str(int(t)) for t in calc.LB_THRESHOLDS],
                      scores={LB_LABEL[n]: "%.4f" % R["LB"][n] for n in LB_ROWS},
                      above_chosen={LB_LABEL[n]: "%.1f%%" % (100 * (R["LB"][n] / R["LB"]["golden"] - 1))
                                    for n in LB_ROWS if n != "golden"},
                      accrual_each={LB_LABEL[n]: usd(calc.r250(sum(R["ACC"][n].values()))) for n in LB_ROWS}),
        rivals=dict(
            default=usd(calc.r250(rv["default: vendor projection with a normal band"])),
            default_at_hand=usd(calc.r250(rv["the full nine's own plate appearances"])),
            opening_day=usd(calc.r250(rv["opening-day lineups"])),
            old_slot=usd(calc.r250(rv["regulars at the old slot's rate"]))),
        deliverables=dict(
            paid=[dict(season=str(y), rows=[(plain(n), usd(v)) for n, v in rows], total=usd(tot))
                  for y, rows, tot in paid_rows(R)],
            reissued_legacy_numbers=R["reissued"]),
    )
    return f


# the bases the look-back sheet and the memo report, in the controller's words
LB_ROWS = ["golden", "the last three seasons", "plate appearances rescaled to the slot",
           "default: vendor projection with a normal band", "vendor projection with its past misses",
           "own three seasons as the draws", "the full nine's own plate appearances",
           "planned regulars' own plate appearances", "catchers kept in", "opening-day lineups",
           "planned regulars who opened the season", "hitters with 400 or more plate appearances",
           "own record given a weight of 10 per cent", "regression on the projection, the slot and catching"]
LB_LABEL = {
    "golden": "Non-catcher regulars' seasons of starts at the player's 2027 slot (used)",
    "the last three seasons": "The same, the last three seasons only",
    "plate appearances rescaled to the slot": "The same, each season's plate appearances rescaled to the slot",
    "default: vendor projection with a normal band": "Vendor's projection with a normal band from its misses",
    "vendor projection with its past misses": "Vendor's projection plus each of its past misses",
    "own three seasons as the draws": "The player's own last three seasons",
    "the full nine's own plate appearances": "Every lineup regular's plate appearances, catchers in, any slot",
    "planned regulars' own plate appearances": "Non-catcher regulars' plate appearances, any slot",
    "catchers kept in": "Every lineup regular's starts at the player's slot, catchers in",
    "opening-day lineups": "Opening-day starters' seasons at the player's slot",
    "planned regulars who opened the season": "Regulars who started their club's first game, at the slot",
    "hitters with 400 or more plate appearances": "Hitters with 400 or more plate appearances, at the slot",
    "own record given a weight of 10 per cent": "The regulars' seasons moved 10 per cent toward the player's record",
    "regression on the projection, the slot and catching": "Regression on the vendor's projection, slot and catching",
}

# ------------------------------------------------------------------------------------------------------- workbook
THIN = Side(style="thin", color="999999")
HEAD_FILL = PatternFill("solid", fgColor="DCE6F0")


def _hdr(ws, row, values, widths=None):
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=i + 1, value=v)
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.border = Border(bottom=THIN)


def write_book(R, path):
    G, P = R["G"], R["P"]
    F = G["five"]
    rt = G["rate"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Accrual"
    ws["A1"] = "Plate-appearance bonuses, 2027 free-agent hitters: accrual at signing %s" % usd(G["total_r"])
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = ("Expected amount of each bonus (FIN-14 4.2): the bonus times the probability that the player reaches "
                "its threshold in the 2027 championship season, assessed at signing with the player in his planned "
                "2027 lineup slot (depth chart of December 10, 2026).")
    ws["A3"] = ("Probability: share of the %s seasons of the non-catcher regulars in the January 31 depth charts, %d "
                "to %d, whose starts times the plate appearances per start of the player's 2027 slot reach the "
                "threshold." % ("{:,}".format(G["n"]), min(R["T"].years), max(R["T"].years)))
    row = 5
    _hdr(ws, row, ["Player", "Signed", "2027 slot", "Threshold (PA)", "Bonus", "Probability", "Expected amount"])
    row += 1
    for f_ in F:
        first = True
        for a, t, p in G["lines"][f_["name"]]:
            ws.cell(row=row, column=1, value=f_["name"] if first else None)
            ws.cell(row=row, column=2, value=f_["signed"] if first else None)
            if first:
                ws.cell(row=row, column=2).number_format = "mm/dd/yyyy"
            ws.cell(row=row, column=3, value=f_["slot"] if first else None)
            ws.cell(row=row, column=4, value=t)
            ws.cell(row=row, column=5, value=a).number_format = "$#,##0"
            ws.cell(row=row, column=6, value=round(p, 4)).number_format = "0.0%"
            ws.cell(row=row, column=7, value=round(a * p, 2)).number_format = "$#,##0"
            first = False
            row += 1
        ws.cell(row=row, column=1, value="%s, accrual" % f_["name"]).font = Font(italic=True)
        ws.cell(row=row, column=7, value=round(G["per"][f_["name"]], 2)).number_format = "$#,##0"
        ws.cell(row=row, column=8, value=G["per_r"][f_["name"]]).number_format = "$#,##0"
        row += 1
    ws.cell(row=row, column=1, value="Total").font = Font(bold=True)
    ws.cell(row=row, column=7, value=round(G["total"], 2)).number_format = "$#,##0"
    c = ws.cell(row=row, column=8, value=G["total_r"])
    c.number_format = "$#,##0"
    c.font = Font(bold=True)
    ws.cell(row=5, column=8, value="Accrual (player to $100,000; total to $250,000)").font = Font(bold=True)
    ws.cell(row=5, column=8).fill = HEAD_FILL
    ws.cell(row=5, column=8).alignment = Alignment(wrap_text=True, vertical="top")
    row += 2
    ws.cell(row=row, column=1, value="Plate appearances per start by lineup slot (starters' lines, %d to %d)" % (
        min(R["T"].years), max(R["T"].years))).font = Font(bold=True)
    row += 1
    _hdr(ws, row, ["Slot"] + [str(s) for s in calc.SLOTS])
    row += 1
    ws.cell(row=row, column=1, value="PA per start")
    for i, s in enumerate(calc.SLOTS):
        ws.cell(row=row, column=i + 2, value=round(rt[s], 3)).number_format = "0.00"
    for col, w in zip("ABCDEFGHIJ", (30, 12, 10, 14, 12, 12, 16, 22, 8, 8)):
        ws.column_dimensions[col].width = w
    ws.row_dimensions[5].height = 30
    ws.freeze_panes = "B6"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "5:5"
    # the look-back
    wl = wb.create_sheet("Look-back (FIN-14 4.4)")
    wl["A1"] = ("FIN-14 4.4 look-back (working paper CF-14.4): %s hitter-seasons, 2023 to 2026; the basis with the "
                "lowest score is used" %
                "{:,}".format(R["n_lb"]))
    wl["A1"].font = Font(bold=True)
    wl["A2"] = ("Each basis applied to each season with what was published before it began, to every hitter in a "
                "starting lineup of that season's January 31 depth charts with 100 plate appearances in each of the two "
                "seasons before; score: the average squared difference between the probability of 450, 500, 550, 600 "
                "and 650 plate appearances and the outcome (1 reached, 0 not).")
    _hdr(wl, 4, ["Basis", "Score", "Above the basis used", "2027 accrual on this basis"])
    gb = R["LB"]["golden"]
    for i, n in enumerate(LB_ROWS):
        r = 5 + i
        wl.cell(row=r, column=1, value=LB_LABEL[n]).font = Font(bold=(n == "golden"))
        wl.cell(row=r, column=2, value=round(R["LB"][n], 4)).number_format = "0.0000"
        c = wl.cell(row=r, column=3, value=None if n == "golden" else round(R["LB"][n] / gb - 1, 4))
        c.number_format = "0.0%"
        wl.cell(row=r, column=4, value=calc.r250(sum(R["ACC"][n].values()))).number_format = "$#,##0"
    r = 5 + len(LB_ROWS) + 1
    wl.cell(row=r, column=1, value="The basis used prices each tested hitter from the seasons of his own group: a "
                                   "catcher from the catcher regulars' seasons, every other hitter from the non-catcher "
                                   "regulars'. The bases that tie it (the last three seasons; plate appearances "
                                   "rescaled) give the same accrual. Accruals to the nearest $250,000.")
    wl.cell(row=r + 1, column=1, value="Sources: hvbd_winter_depth_charts_2021_2026.xlsx; "
                                       "hvbd_hitter_projections_2022_2027.csv; the batting lines 2021 to 2026.")
    for col, w in zip("ABCD", (66, 10, 14, 18)):
        wl.column_dimensions[col].width = w
    wl.freeze_panes = "A5"
    # the other playing-time bonuses paid
    ws3 = wb.create_sheet("Other PT bonuses 2023-2025")
    ws3["A1"] = ("Other playing-time bonuses paid: %s for 2023, %s for 2024, %s for 2025 (payroll register through "
                 "December 11, 2026)" % tuple(usd(tot) for _, _, tot in paid_rows(R)))
    ws3["A1"].font = Font(bold=True)
    _hdr(ws3, 3, ["Season", "Player", "Contract terms (each season)", "Season-end payment", "January payment",
                  "Award in season-end line", "Deferred salary in season-end line", "Paid for the season"])
    r = 4
    for y, rows, tot in paid_rows(R):
        for n, v in rows:
            dd = R["BD"][(y, n)]
            ws3.cell(row=r, column=1, value=y)
            ws3.cell(row=r, column=2, value=plain(n))
            ws3.cell(row=r, column=3, value=terms_for(R, n, y))
            ws3.cell(row=r, column=4, value=round(dd["oct"])).number_format = "$#,##0"
            ws3.cell(row=r, column=5, value=round(dd["jan"]) or None).number_format = "$#,##0"
            c = ws3.cell(row=r, column=6, value=-round(dd["award"]) if dd["award"] else None)
            c.number_format = "$#,##0;($#,##0)"
            c = ws3.cell(row=r, column=7, value=-round(dd["deferral"]) if dd["deferral"] else None)
            c.number_format = "$#,##0;($#,##0)"
            ws3.cell(row=r, column=8, value=round(v)).number_format = "$#,##0"
            r += 1
        ws3.cell(row=r, column=2, value="Total %d" % y).font = Font(bold=True)
        parts = [R["BD"][(y, n)] for n, _ in rows]
        for col, v in ((4, sum(p["oct"] for p in parts)), (5, sum(p["jan"] for p in parts)),
                       (6, -sum(p["award"] for p in parts) or None), (7, -sum(p["deferral"] for p in parts) or None),
                       (8, tot)):
            c = ws3.cell(row=r, column=col, value=round(v) if v is not None else None)
            c.number_format = "$#,##0;($#,##0)"
            c.font = Font(bold=True)
            c.border = Border(top=THIN)
        r += 2
    ws3.cell(row=r, column=1, value="The pitchers' bonuses are paid half in the season-end run (October 15) and half "
                                    "on the following January 15; after the last season of a contract the season-end "
                                    "run pays the whole season. Before the January 2026 HRIS cutover the legacy payroll "
                                    "put the season-end payments on E45 OTHER BONUS and the January halves on E40 PERF "
                                    "BONUS; from January 2026 both are BPA.")
    ws3.cell(row=r + 1, column=1, value="Taken out of the season-end line: an All-Star bonus paid in it (Ruiz 2023, "
                                        "Livingston 2024, Pietersz 2025), and after a contract's last season the salary "
                                        "the contract deferred to its end (the register's salary less the salary paid, "
                                        "each season of the contract). White's Cy Young bonus (November 15, 2024) is an "
                                        "award, not playing time.")
    ws3.cell(row=r + 2, column=1, value="Legacy number %s was reissued: Yordy Benítez's lines in 2023, Gilmar "
                                        "Pietersz's from 2025; legacy lines are taken by the name on the line. The "
                                        "October 2025 void and reissue to Reed counts once." % R["reissued"][0])
    ws3.cell(row=r + 3, column=1, value="Sources: ral_payroll_player_earnings_2023_2026.csv; "
                                        "ral_hris_legacy_employee_map.csv; ral_player_contract_register_2023_2026.xlsx.")
    for col, w in zip("ABCDEFGH", (8, 20, 62, 14, 14, 14, 14, 14)):
        ws3.column_dimensions[col].width = w
    ws3.row_dimensions[3].height = 30
    ws3.freeze_panes = "C4"
    ws3.page_setup.orientation = "landscape"
    ws3.page_setup.fitToWidth = 1
    ws3.sheet_properties.pageSetUpPr.fitToPage = True
    ws3.print_title_rows = "3:3"
    # notes
    wn = wb.create_sheet("Notes")
    for line in [
        "As of December 11, 2026 (files pulled that afternoon). Prepared by %s, Controller." % AUTHOR,
        "Accrual under FIN-14 4.2 and 4.3: each plate-appearance bonus at its amount times the probability that the "
        "player reaches its threshold in the 2027 championship season, assessed at signing, the contract taken as "
        "standing all season with the player in his planned role (lineup slot in ral_depth_chart_2027.xlsx) and his "
        "health and playing time left in the probability. Basis chosen by the FIN-14 4.4 look-back (sheet Look-back).",
        "Plate appearances as the term sheets define them: championship season, for the Club; postseason and "
        "rehabilitation-assignment plate appearances excluded.",
        "Comparison seasons: the non-catchers in the lineups of hvbd_winter_depth_charts_2021_2026.xlsx, %s seasons; "
        "starts counted for every club the player started for that season, from lsf_batting_2021_2024.txt (through "
        "lsf_player_id_crosswalk.csv) and lsf_batting_2025_2026.csv." % "{:,}".format(G["n"]),
        "Rounding (FIN-14 6): each player's accrual to the nearest $100,000 and the total to the nearest $250,000, "
        "both from the unrounded line amounts; the rounded player figures add to %s." % usd(sum(G["per_r"].values())),
    ]:
        wn.append([line])
    wn.column_dimensions["A"].width = 130
    for row in wn.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    K.save_workbook(wb, path, STAMP, AUTHOR, "Plate-appearance bonus accrual, 2027")


# ----------------------------------------------------------------------------------------------------------- memo
def _docx_fix(path, stamp, author):
    data = open(path, "rb").read()
    K.repack_ooxml(data, path, stamp, app="Microsoft Office Word", creator=author)
    import zipfile as Z
    z = Z.ZipFile(path)
    bits = {n: z.read(n) for n in z.namelist()}
    z.close()
    core = bits["docProps/core.xml"].decode("utf-8")
    core = re.sub(r"<dc:description>[^<]*</dc:description>", "<dc:description></dc:description>", core)
    core = re.sub(r"<cp:keywords>[^<]*</cp:keywords>", "<cp:keywords></cp:keywords>", core)
    bits["docProps/core.xml"] = core.encode("utf-8")
    out = io.BytesIO()
    zo = Z.ZipFile(out, "w", Z.ZIP_DEFLATED)
    for n, b in bits.items():
        zi = Z.ZipInfo(n, date_time=(stamp.year, stamp.month, stamp.day, stamp.hour, stamp.minute, 0))
        zi.compress_type = Z.ZIP_DEFLATED
        zi.external_attr = 0o600 << 16
        zo.writestr(zi, b)
    zo.close()
    open(path, "wb").write(out.getvalue())


def write_memo(R, path, fig):
    G = R["G"]
    F = G["five"]
    d = docx.Document()
    cp = d.core_properties
    cp.author = AUTHOR
    cp.last_modified_by = AUTHOR
    cp.title = "Plate-appearance bonus accrual at signing"
    cp.comments = ""
    cp.keywords = ""
    cp.subject = ""
    cp.category = ""
    cp.revision = 3
    cp.created = STAMP
    cp.modified = STAMP
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    for s in d.sections:
        s.left_margin = s.right_margin = Inches(0.9)
        s.top_margin = s.bottom_margin = Inches(0.8)
    p = d.add_paragraph()
    r = p.add_run("Raleigh Oaks Finance: memo to the December 2026 close file")
    r.bold = True
    r.font.size = Pt(12)
    for line in ["To: Audit file, December 2026 close", "From: %s, Controller" % AUTHOR, "Date: December 11, 2026",
                 "Re: Plate-appearance bonus accrual at signing, five free-agent hitters (FIN-14 section 4)"]:
        d.add_paragraph(line).paragraph_format.space_after = Pt(0)
    d.add_paragraph()
    d.add_paragraph(
        "We accrue %s at signing for the plate-appearance bonuses in the 2027 contracts of the five free-agent hitters "
        "signed between November 24 and December 10, 2026. Each bonus is accrued at its expected amount under FIN-14 4.2 "
        "and 4.3: the bonus times the probability, at the signing date, that the player reaches its threshold in the 2027 "
        "championship season, with the player in the lineup slot Baseball Operations plans for him." % fig["call"]["total"])
    t = d.add_table(rows=1, cols=4)
    t.style = "Light Grid Accent 1"
    hdr = t.rows[0].cells
    for c, v in zip(hdr, ["Player", "Signed", "2027 slot", "Accrual"]):
        c.text = v
    for f_ in F:
        cells = t.add_row().cells
        cells[0].text = f_["name"]
        cells[1].text = f_["signed"].strftime("%B %-d, %Y")
        cells[2].text = "%s (%s for his old club in 2026)" % (ORD[f_["slot"]].capitalize(), ORD[f_["old"]])
        cells[3].text = fig["critical_components"]["per_player"][f_["name"]]
    cells = t.add_row().cells
    cells[0].text = "Total"
    cells[3].text = fig["call"]["total"]
    for row in t.rows:
        row.cells[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    src = d.add_paragraph("Source: term sheets of the five contracts; Baseball Operations depth chart of December 10, "
                          "2026; pa_bonus_accrual_2027.xlsx.")
    src.runs[0].font.size = Pt(8.5)
    src.runs[0].italic = True
    lb = fig["lookback"]
    d.add_paragraph(
        "Basis for the probabilities (FIN-14 4.4 and working paper CF-14.4). We scored each basis on the %s "
        "hitter-seasons of 2023 to 2026 the working paper names, each basis given only what was published before the "
        "season. The lowest score went to the share of the %s seasons of the non-catcher regulars in the vendor's "
        "January 31 depth charts, 2021 to 2026, whose starts, times the plate appearances per start of the player's "
        "2027 slot, reach the threshold; plate "
        "appearances per start fall from %s batting first to %s batting ninth. The vendor's projection with a band "
        "from its past misses and the players' own last three seasons both scored worse." % (
            lb["hitter_seasons"], fig["steps"]["planned_seasons"], fig["steps"]["rate_first"], fig["steps"]["rate_ninth"]))
    d.add_paragraph(
        "The five come to us off full seasons batting higher for their old clubs; each will bat lower here, and in the "
        "look-back the regulars' seasons at the new slot forecast better than each hitter's own recent seasons. Player "
        "figures are rounded to the nearest $100,000 and the "
        "total to the nearest $250,000 from the unrounded amounts (FIN-14 6), so the player figures add to %s. The "
        "accrual is remeasured at each 2027 quarter end (FIN-14 4.5); pa_bonus_accrual_2027.xlsx carries each bonus "
        "line and the look-back." % usd(sum(G["per_r"].values())))
    nt = d.add_paragraph("Plate appearances are counted as the contracts define them: championship season, for the "
                         "Club, excluding postseason and rehabilitation-assignment plate appearances. Catchers are left "
                         "out of the comparison seasons; none of the five catches.")
    nt.runs[0].font.size = Pt(8.5)
    d.save(path)
    _docx_fix(path, STAMP, AUTHOR)


# ----------------------------------------------------------------------------------------------------- submission
def submission(R, fig, path):
    G = R["G"]
    F = G["five"]
    S = fig["steps"]
    per = fig["critical_components"]["per_player"]
    lines = []
    lines.append("## Tags\n")
    lines.append("**Domain:** General Statistics & Data Science (resampling and simulation methods: the expected value "
                 "of stepped plate-appearance bonuses priced from the distribution of comparable whole seasons, each "
                 "season's starts read at the player's new lineup slot's plate appearances per start).")
    lines.append("**Analytical objective:** Forecasting & Predictive Modeling.\n")
    lines.append("## 1. Final Recommendation\n")
    lines.append("**Accrue %s at signing for the plate-appearance bonuses in the five free-agent hitters' 2027 "
                 "contracts.**" % fig["call"]["total"])
    LBk = fig["lookback"]
    lines.append("Each bonus is priced at the share of the %s seasons of the non-catcher regulars in the vendor's "
                 "January 31 depth charts, %s to %s, whose starts, read at the plate appearances per start of the "
                 "player's 2027 lineup slot, reach its threshold: the basis with the lowest score in FIN-14 4.4's "
                 "look-back over the %s hitter-seasons of 2023 to 2026. Not the vendor's 2027 projections with a band "
                 "from their past misses, which carry each player's full recent seasons at a higher slot for his old "
                 "club; not the players' own seasons; and not the starting nines' seasons taken whole, which keep the "
                 "catchers in and ignore where each hitter bats; each scores higher in the look-back." % (
                     S["planned_seasons"], S["first_season"], S["last_season"], LBk["hitter_seasons"]))
    lines.append("")
    lines.append("## 2. Critical Components\n")
    lines.append("1. Accrual by player: %s." % "; ".join("%s **%s**" % (f_["name"], per[f_["name"]]) for f_ in F))
    lines.append("2. Plate appearances per start fall from **%s** batting first to **%s** batting ninth, and the "
                 "depth chart bats the five %s." % (S["rate_first_1dp"], S["rate_ninth_1dp"], ", ".join(
                     "%s %s" % (surname(f_["name"]), fig["critical_components"]["slots_2027"][f_["name"]]) for f_ in F)))
    lines.append("3. FIN-14 4.4's look-back over the **%s** hitter-seasons of 2023 to 2026 ranks the regulars' seasons of "
                 "starts at the slot ahead of the vendor's projection with a band and of the players' own seasons." %
                 LBk["hitter_seasons"])
    lines.append("")
    lines.append("## 3. Step-by-Step Solution\n")
    lines.append("1. Joined lsf_batting_2021_2024.txt (plate appearances AB + BB + HBP + SH + SF + CI, per lsf_feed_notes.txt) "
                 "through lsf_player_id_crosswalk.csv to lsf_batting_2025_2026.csv: %s batting lines, %s to %s." % (
                     S["batting_lines"], S["first_season"], S["last_season"]))
    lines.append("2. Took the lineup rows of hvbd_winter_depth_charts_2021_2026.xlsx other than catchers, %s seasons, and "
                 "counted each player's starts that season for every club (a start is a starter's line): %s of the "
                 "seasons had 155 starts or more and %s had fewer than 100." % (S["planned_seasons"], S["full_155"],
                                                                            S["under_100"]))
    lines.append("3. Plate appearances per start by lineup slot over all starters' lines, %s to %s, fall from %s batting "
                 "first to %s batting ninth; at the five's 2027 slots they are %s." % (
                     S["first_season"], S["last_season"], S["rate_first"], S["rate_ninth"],
                     ", ".join("%s batting %s" % (v, k) for k, v in S["rates"].items())))
    sc = LBk["scores"]
    lines.append("4. Ran the look-back of FIN-14 4.4 (working paper CF-14.4) on the %s hitter-seasons of 2023 to 2026 (a "
                 "starting lineup of the January 31 depth charts and 100 plate appearances in each of the two seasons "
                 "before), each basis built from what was published before the season: the regulars' seasons of starts "
                 "at the slot, a tested catcher priced from the catcher regulars' seasons and every other hitter from "
                 "the non-catcher regulars', scored %s, the lowest; the vendor's projection with a band, the players' "
                 "own seasons and the starting nines' seasons taken whole each scored higher." % (
                     LBk["hitter_seasons"], sc[LB_LABEL["golden"]]))
    lines.append("5. For each bonus line in term_sheets_2027_free_agent_hitters.pdf, the probability under FIN-14 4.2 and "
                 "4.3 is the share of those seasons whose starts times the player's 2027 slot rate reach the threshold, "
                 "and the expected amount is the bonus times that share.")
    lines.append("6. Summed each player's expected amounts and rounded to the nearest $100,000: %s." % (
        "; ".join("%s %s" % (f_["name"], per[f_["name"]]) for f_ in F)))
    lines.append("7. Summed the five players' unrounded amounts and rounded to the nearest $250,000 (FIN-14 6): %s; the "
                 "rounded player figures add to %s." % (fig["call"]["total"], usd(sum(G["per_r"].values()))))
    pd_ = fig["deliverables"]["paid"]
    lines.append("8. For each player whose contract in ral_player_contract_register_2023_2026.xlsx carries other "
                 "playing-time terms, added the season-end payment (October 15) and the January 15 payment after it in "
                 "ral_payroll_player_earnings_2023_2026.csv, whatever the code (the legacy payroll paid the pitchers' "
                 "season-end payments on E45 OTHER BONUS), net of the October 2025 void and reissue, less the All-Star "
                 "bonus a legacy season-end line carries and, in the line after a contract's last season, the salary "
                 "the contract deferred to its end (its register salary less the E01 salary lines paid, each season), "
                 "each legacy line read by its name (number %s is Yordy Benítez in 2023 and Gilmar Pietersz in 2025), "
                 "and where a January run paid no half on a contract running on (Max White's for 2023, Joe Livingston's "
                 "for 2024) the next season-end line carries it, so its excess over the January half after it goes back "
                 "to that season: %s." % (fig["deliverables"]["reissued_legacy_numbers"][0],
                          "; ".join("%s %s" % (b["season"], b["total"]) for b in pd_)))
    lines.append("")
    lines.append("## 4. Deliverable Answers\n")
    lines.append("### %s\n" % F_BOOK)
    lines.append("1. Accrual for each player and in total:")
    for f_ in F:
        lines.append("   - %s: %s" % (f_["name"], per[f_["name"]]))
    lines.append("   - Total: %s" % fig["call"]["total"])
    lines.append("2. Playing-time bonuses other than plate-appearance bonuses paid for each season, player by player "
                 "and in total:")
    for blk in fig["deliverables"]["paid"]:
        for n, v in blk["rows"]:
            lines.append("   - %s season, %s: %s" % (blk["season"], n, v))
        lines.append("   - %s season, total: %s" % (blk["season"], blk["total"]))
    lines.append("")
    lines.append("### %s\n" % F_MEMO)
    lines.append("1. One-page note for the audit file stating the accrual of %s and each player's amount: %s." % (
        fig["call"]["total"], "; ".join("%s %s" % (f_["name"], per[f_["name"]]) for f_ in F)))
    txt = "\n".join(lines) + "\n"
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write(txt)
    return txt


def main(target, golden, fig_path, sub_path):
    R = compute(target)
    os.makedirs(golden, exist_ok=True)
    fig = figures(R)
    write_book(R, os.path.join(golden, F_BOOK))
    write_memo(R, os.path.join(golden, F_MEMO), fig)
    with open(fig_path, "w", encoding="utf-8", newline="\n") as f:
        json.dump(fig, f, indent=1, ensure_ascii=False)
        f.write("\n")
    submission(R, fig, sub_path)
    return R, fig
