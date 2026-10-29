"""Every figure from the shipped files alone: the accrual (the golden), the default path and every rival, the
constructions that must round alike, the bootstrap, and the two asks. Reads only <target>; never the world."""
from __future__ import annotations

import csv
import datetime as dt
import io
import json
import os
import re
from collections import defaultdict
from statistics import NormalDist

import numpy as np
import openpyxl

NORM = NormalDist()
SLOTS = range(1, 10)
MIN_DAYS = {"7-day": 7, "10-day": 10, "15-day": 15, "60-day": 60}


# ------------------------------------------------------------------------------------------------------------ loading
def pdf_text(path):
    from pdfminer.high_level import extract_text
    return extract_text(path)


def read_csv(path, delimiter=","):
    with open(path, encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f, delimiter=delimiter))


def mdy(s):
    return dt.datetime.strptime(s, "%m/%d/%Y").date() if s else None


def ymd(s):
    return dt.date.fromisoformat(s) if s else None


def fold(s):
    import unicodedata
    return "".join(c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c))


class Pack:
    pass


def load(target, mishandle=()):
    """The shipped files read into one shape. `mishandle` names Layer 1 items handled the wrong way (for the R2
    neutrality checks): 'no_ci' plate appearances without catcher's interference in the legacy layout; 'subs_as_starts'
    every platform line a start; 'dh_once' one start a date (doubleheaders merged)."""
    P = Pack()
    P.target = target
    f = lambda n: os.path.join(target, n)
    xw = read_csv(f("lsf_player_id_crosswalk.csv"))
    P.legacy_to_id = {r["legacy_id"]: int(r["player_id"]) for r in xw}
    P.xwalk_name = {int(r["player_id"]): "%s %s" % (r["first_name"], r["last_name"]) for r in xw}
    clubs = read_csv(f("league_club_directory.csv"))
    P.legacy_club = {r["legacy_code"]: r["abbrev"] for r in clubs}
    P.team_club = {int(r["team_id"]): r["abbrev"] for r in clubs}
    P.season = {}
    if os.path.exists(f("league_calendar_2021_2027.csv")):
        cal = read_csv(f("league_calendar_2021_2027.csv"))
        P.season = {int(r["season"]): (ymd(r["opening_day"]), ymd(r["last_scheduled_day"])) for r in cal}
    # batting lines in one shape: season, game key, date, club, player, slot, starter, plate appearances, position
    lines = []
    comps = ("AB", "BB", "HBP", "SH", "SF") if "no_ci" in mishandle else ("AB", "BB", "HBP", "SH", "SF", "CI")
    for r in read_csv(f("lsf_batting_2021_2024.txt"), delimiter="|"):
        d = mdy(r["GAME_DT"])
        pa = sum(int(r[c]) for c in comps)
        lines.append((d.year, r["GAME_ID"], d, P.legacy_club[r["TEAM_CD"]], P.legacy_to_id[r["PLAYER_CD"]],
                      int(r["LINEUP"]), r["SEQ"] == "1", pa, int(r["POS"])))
    posmap = {"C": 2, "1B": 3, "2B": 4, "3B": 5, "SS": 6, "LF": 7, "CF": 8, "RF": 9, "DH": 10, "PH": 11, "P": 1}
    for r in read_csv(f("lsf_batting_2025_2026.csv")):
        d = ymd(r["official_date"])
        bo = int(r["batting_order"])
        start = True if "subs_as_starts" in mishandle else bo % 100 == 0
        lines.append((d.year, r["game_pk"], d, P.team_club[int(r["team_id"])], int(r["player_id"]), bo // 100,
                      start, int(r["pa"]), posmap[r["position"]]))
    if "dh_once" in mishandle:
        seen, kept = set(), []
        for ln in lines:
            if ln[6]:
                key = (ln[2], ln[4])
                if key in seen:
                    continue
                seen.add(key)
            kept.append(ln)
        lines = kept
    P.lines = lines
    # the vendor's winter depth charts
    wb = openpyxl.load_workbook(f("hvbd_winter_depth_charts_2021_2026.xlsx"), read_only=True)
    P.plan = []
    for ws in wb.worksheets:
        y = int(ws.title.split()[-1])
        hdr = None
        for row in ws.iter_rows(values_only=True):
            if row and row[0] == "Tm":
                hdr = list(row)
                continue
            if hdr and row and row[0]:
                d = dict(zip(hdr, row))
                P.plan.append(dict(season=y, team=d["Tm"], role=d["Role"], ord=d["Ord"], pos=d["Pos"], name=d["Player"],
                                   lg=int(d["LG_ID"]), age=int(d["Age"])))
    # the club's 2027 depth chart
    wb = openpyxl.load_workbook(f("ral_depth_chart_2027.xlsx"), read_only=True)
    ws = wb.worksheets[0]
    P.lineup27 = {}
    mode = None
    for row in ws.iter_rows(values_only=True):
        if row and row[0] == "Lineup":
            mode = "lineup"
            continue
        if row and row[0] in ("Bench", "Rotation", "Bullpen"):
            mode = None
        if mode == "lineup" and row and isinstance(row[0], int):
            P.lineup27[row[2]] = (row[0], row[1])
    # the term sheets
    P.terms = term_sheets(f("term_sheets_2027_free_agent_hitters.pdf"))
    P.proj = read_csv(f("hvbd_hitter_projections_2022_2027.csv")) if os.path.exists(
        f("hvbd_hitter_projections_2022_2027.csv")) else []
    return P


def term_sheets(path):
    txt = " ".join(pdf_text(path).split())
    heads = list(re.finditer(r"([A-Z][^()]{2,40}?) \((\w+), bats (\w), born ([A-Za-z]+ \d+, \d{4})\)", txt))
    out = {}
    for i, m in enumerate(heads):
        seg = txt[m.end():heads[i + 1].start() if i + 1 < len(heads) else len(txt)]
        name = m.group(1).split("Salary ")[-1].strip()
        signed = re.search(r"([A-Z][a-z]+ \d+, \d{4})", seg).group(1)
        b = re.search(r"2027 season: \$([\d,]+) each for ([\d, and]+?) plate appearances", seg)
        amt = int(b.group(1).replace(",", ""))
        ts = [int(x) for x in re.findall(r"\d+", b.group(2))]
        out[name] = dict(pos=m.group(2), signed=dt.datetime.strptime(signed, "%B %d, %Y").date(),
                         born=dt.datetime.strptime(m.group(4), "%B %d, %Y").date(), lines=[(amt, t) for t in ts])
    return out


# --------------------------------------------------------------------------------------------------- season tables
def seasons(P):
    """Starts and plate appearances by player and season (every club), slot rates, and the planned regulars."""
    T = Pack()
    T.starts = defaultdict(int)
    T.pa = defaultdict(int)
    T.start_pa = defaultdict(int)
    T.slot_pa = defaultdict(int)
    T.slot_n = defaultdict(int)
    T.first_game = {}
    T.pos_starts = defaultdict(lambda: defaultdict(int))
    T.club_starts = defaultdict(lambda: defaultdict(int))
    for (y, g, d, club, pid, slot, st, pa, pos) in P.lines:
        T.pa[(y, pid)] += pa
        if st:
            T.starts[(y, pid)] += 1
            T.start_pa[(y, pid)] += pa
            T.slot_pa[(y, slot)] += pa
            T.slot_n[(y, slot)] += 1
            key = (y, club)
            if key not in T.first_game or (d, g) < T.first_game[key][0]:
                T.first_game[key] = ((d, g), [])
            T.pos_starts[(y, club, pos)][pid] += 1
            T.club_starts[(y, club)][pid] += 1
    T.opening, T.opening_c = set(), set()
    for (y, g, d, club, pid, slot, st, pa, pos) in P.lines:
        if st and T.first_game[(y, club)][0] == (d, g):
            (T.opening if pos != 2 else T.opening_c).add((y, pid))
    T.years = sorted({k[0] for k in T.starts})
    T.planned = [r for r in P.plan if r["role"] == "Lineup"]
    return T


def rate(T, years=None):
    years = years or T.years
    return {s: sum(T.slot_pa[(y, s)] for y in years) / sum(T.slot_n[(y, s)] for y in years) for s in SLOTS}


def planned_starts(T, catchers=False, years=None, ages=None):
    rows = [r for r in T.planned if (catchers or r["pos"] != "C") and (years is None or r["season"] in years)
            and (ages is None or ages[0] <= r["age"] <= ages[1])]
    return np.array([T.starts[(r["season"], r["lg"])] for r in rows], float), rows


def five(P):
    """The five signings: name, terms, new slot (2027 depth chart), league ID and old slot (the vendor's last winter
    chart, 2026)."""
    out = []
    last = max(r["season"] for r in P.plan)
    for name, t in P.terms.items():
        slot, pos = P.lineup27[name]
        prior = [r for r in P.plan if r["name"] == name and r["season"] == last and r["role"] == "Lineup"]
        lg = prior[0]["lg"] if prior else [r for r in P.plan if r["name"] == name][-1]["lg"]
        old = prior[0]["ord"] if prior else None
        age27 = 2027 - t["born"].year - (1 if (t["born"].month, t["born"].day) > (7, 1) else 0)
        out.append(dict(name=name, lines=t["lines"], slot=slot, pos=pos, lg=lg, old=old, age=age27, signed=t["signed"]))
    out.sort(key=lambda x: x["signed"])
    return out


def ev(lines, sample):
    """Expected amount of a set of bonus lines over a sample of season plate-appearance totals."""
    sample = np.asarray(sample, float)
    return sum(a * float(np.mean(sample >= T)) for a, T in lines)


def ev_normal(lines, mu, sd):
    return sum(a * (1 - NORM.cdf((T - mu) / sd)) for a, T in lines)


def ev_point(lines, x):
    return sum(a * float(x >= T) for a, T in lines)


def r250(x):
    return int(round(x / 250000.0)) * 250000


def r100(x):
    return int(round(x / 100000.0)) * 100000


# -------------------------------------------------------------------------------------------------------- the call
def golden(P, T=None):
    T = T or seasons(P)
    F = five(P)
    rt = rate(T)
    st, _ = planned_starts(T)
    per = {f["name"]: ev(f["lines"], st * rt[f["slot"]]) for f in F}
    lines = {f["name"]: [(a, Tt, float(np.mean(st * rt[f["slot"]] >= Tt))) for a, Tt in f["lines"]] for f in F}
    tot = sum(per.values())
    return dict(per=per, total=tot, total_r=r250(tot), per_r={k: r100(v) for k, v in per.items()}, lines=lines,
                rate=rt, n=len(st), five=F, starts=st)


def vendor_errors(P, T):
    """The vendor's past misses: actual season plate appearances less the preseason projection, for the regulars in
    its winter depth charts (2022 to 2026)."""
    planned = {(r["season"], r["lg"]) for r in T.planned if r["pos"] != "C"}
    errs = []
    for r in P.proj:
        y = int(r["season"])
        if y > 2026:
            continue
        key = (y, int(r["lg_id"]))
        if key in planned:
            errs.append(T.pa[key] - int(r["pa"]))
    return np.array(errs, float)


# ------------------------------------------------------------------------------------- the bases and the look-back
LB_SEASONS = (2023, 2024, 2025, 2026)          # FIN-14 4.4: the four championship seasons before the signings
LB_THRESHOLDS = np.array([450.0, 500.0, 550.0, 600.0, 650.0])
LB_MIN_PA = 100
FAMILY = ("golden", "the last three seasons", "plate appearances rescaled to the slot")
LANDS = ("regulars planned at the same slot", "ages within three years")   # lose the look-back, file the golden


def emp(vals, thr):
    """The share of a set of season plate-appearance totals at or above each threshold."""
    v = np.sort(np.asarray(vals, float))
    return 1 - np.searchsorted(v, np.asarray(thr, float), side="left") / len(v)


def normal(mu, sd, thr):
    return np.array([1 - NORM.cdf((t - mu) / sd) for t in thr])


def projections(P):
    return {(int(r["season"]), int(r["lg_id"])): int(r["pa"]) for r in P.proj}


def hitter_seasons(P, T):
    """Every hitter-season in the box scores, flagged a catcher's when catcher is his most common starting position."""
    if getattr(T, "_hs", None) is None:
        posc = defaultdict(lambda: defaultdict(int))
        for (y, g, d, club, pid, slot, st, pa, pos) in P.lines:
            if st:
                posc[(y, pid)][pos] += 1
        T._hs = {k: max(d, key=d.get) == 2 for k, d in posc.items()}
    return T._hs


def appearances(P, T):
    if getattr(T, "_app", None) is None:
        app = defaultdict(int)
        ap, an = defaultdict(int), defaultdict(int)
        for (y, g, d, club, pid, slot, st, pa, pos) in P.lines:
            app[(y, pid)] += 1
            ap[(y, slot)] += pa
            an[(y, slot)] += 1
        T._app = (app, ap, an)
    return T._app


class Ctx:
    pass


def context(P, T, Y):
    """What a basis may use for season Y: the box scores of the seasons before it, the depth charts and the vendor's
    projections published before it began."""
    C = Ctx()
    C.Y = Y
    C.years = [y for y in T.years if y < Y]
    C.rt = rate(T, C.years)
    rows = [r for r in T.planned if r["season"] < Y]
    C.st = np.array([T.starts[(r["season"], r["lg"])] for r in rows], float)
    C.pa = np.array([T.pa[(r["season"], r["lg"])] for r in rows], float)
    C.yr = np.array([r["season"] for r in rows])
    C.ord = np.array([r["ord"] for r in rows])
    C.isc = np.array([r["pos"] == "C" for r in rows])
    C.age = np.array([r["age"] for r in rows])
    C.team = np.array([r["team"] for r in rows])
    C.est = np.array([all(T.pa[(r["season"] - j, r["lg"])] >= LB_MIN_PA for j in (1, 2)) for r in rows])
    opened = T.opening | T.opening_c
    C.opened = np.array([(r["season"], r["lg"]) in opened for r in rows])
    C.od = {False: np.array([T.starts[k] for k in sorted(T.opening) if k[0] < Y], float),
            True: np.array([T.starts[k] for k in sorted(T.opening_c) if k[0] < Y], float)}
    hs = hitter_seasons(P, T)
    C.hs = {c: np.array([(T.starts[k], T.pa[k]) for k, ic in sorted(hs.items()) if k[0] < Y and ic == c], float)
            for c in (False, True)}
    proj = projections(P)
    iv = [i for i, r in enumerate(rows) if (r["season"], r["lg"]) in proj]
    C.vp = np.array([proj[(rows[i]["season"], rows[i]["lg"])] for i in iv], float)
    C.vpa, C.visc, C.vord = C.pa[iv], C.isc[iv], C.ord[iv]
    C.err = {c: C.vpa[C.visc == c] - C.vp[C.visc == c] for c in (False, True)}
    C.fit = {}
    for c in (False, True):
        m = C.visc == c
        if m.sum() > 10:
            b, a = np.polyfit(C.vp[m], C.vpa[m], 1)
            C.fit[c] = (a, b, C.vpa[m] - (a + b * C.vp[m]))
    sr = np.array([C.rt[o] for o in C.vord])
    X = np.column_stack([np.ones(len(iv)), C.vp, sr, C.visc.astype(float)])
    C.beta, *_ = np.linalg.lstsq(X, C.vpa, rcond=None)
    C.res = C.vpa - X @ C.beta
    C.beta0, *_ = np.linalg.lstsq(X[:, [0, 2, 3]], C.vpa, rcond=None)
    C.res0 = C.vpa - X[:, [0, 2, 3]] @ C.beta0
    app, ap, an = appearances(P, T)
    C.rta = {s: sum(ap[(y, s)] for y in C.years) / sum(an[(y, s)] for y in C.years) for s in SLOTS}
    # the slot's plate appearances with its substitutes' lines, per start (the round 3 judge's near rival)
    C.rts = {s: sum(ap[(y, s)] for y in C.years) / sum(T.slot_n[(y, s)] for y in C.years) for s in SLOTS}
    C.app = np.array([app[(r["season"], r["lg"])] for r in rows], float)
    return C


def _g(C, t):
    return C.isc == t["is_c"]


def _own(t, n=3):
    return np.array([x for x in t["own"][:n] if x is not None], float)


BASES = {
    # the golden basis and the constructions of it that FIN-14 4.4 scores alike
    "golden": lambda C, t, x: emp(C.st[_g(C, t)] * C.rt[t["slot"]], x),
    "the last three seasons": lambda C, t, x: emp(C.st[_g(C, t) & (C.yr >= C.Y - 3)] * C.rt[t["slot"]], x),
    "plate appearances rescaled to the slot": lambda C, t, x: emp(
        C.pa[_g(C, t)] * C.rt[t["slot"]] / np.array([C.rt[o] for o in C.ord[_g(C, t)]]), x),
    "regulars planned at the same slot": lambda C, t, x: _same_slot(C, t, x),
    "ages within three years": lambda C, t, x: emp(C.st[_g(C, t) & (np.abs(C.age - t["age"]) <= 3)] * C.rt[t["slot"]], x),
    # the vendor's projection
    "default: vendor projection with a normal band": lambda C, t, x: normal(t["proj"], C.err[t["is_c"]].std(), x),
    "vendor projection with its past misses": lambda C, t, x: emp(t["proj"] + C.err[t["is_c"]], x),
    "vendor point": lambda C, t, x: (t["proj"] >= np.asarray(x)).astype(float),
    "vendor projection plus its mean miss, normal band": lambda C, t, x: normal(
        t["proj"] + C.err[t["is_c"]].mean(), C.err[t["is_c"]].std(), x),
    "vendor projection rescaled to the new slot, normal band": lambda C, t, x: normal(
        t["proj"] * C.rt[t["slot"]] / C.rt[t["prev"] or t["slot"]], C.err[t["is_c"]].std(), x),
    "vendor projection rescaled to the new slot, its past misses": lambda C, t, x: emp(
        t["proj"] * C.rt[t["slot"]] / C.rt[t["prev"] or t["slot"]] + C.err[t["is_c"]], x),
    "vendor projection calibrated by regression, normal band": lambda C, t, x: normal(
        C.fit[t["is_c"]][0] + C.fit[t["is_c"]][1] * t["proj"], C.fit[t["is_c"]][2].std(), x),
    "vendor projection calibrated by regression, its residuals": lambda C, t, x: emp(
        C.fit[t["is_c"]][0] + C.fit[t["is_c"]][1] * t["proj"] + C.fit[t["is_c"]][2], x),
    # the player's own record
    "own three seasons as the draws": lambda C, t, x: emp(_own(t), x),
    "own three-season mean, normal band": lambda C, t, x: normal(_own(t).mean(), C.err[t["is_c"]].std(), x),
    "own seasons' starts at the new slot": lambda C, t, x: emp(
        np.array([s for s in t["ownst"][:3] if s is not None], float) * C.rt[t["slot"]], x),
    "own seasons since 2021": lambda C, t, x: emp(_own(t, 9), x),
    # the regulars' seasons read another way
    "regulars at the old slot's rate": lambda C, t, x: emp(C.st[_g(C, t)] * C.rt[t["prev"] or t["slot"]], x),
    "old slot, normal band": lambda C, t, x: normal(C.rt[t["prev"] or t["slot"]] * C.st[_g(C, t)].mean(),
                                                    C.rt[t["prev"] or t["slot"]] * C.st[_g(C, t)].std(), x),
    "new slot, normal band": lambda C, t, x: normal(C.rt[t["slot"]] * C.st[_g(C, t)].mean(),
                                                    C.rt[t["slot"]] * C.st[_g(C, t)].std(), x),
    "new slot at the mean starts": lambda C, t, x: (C.rt[t["slot"]] * C.st[_g(C, t)].mean() >= np.asarray(x)).astype(float),
    "new slot at the median starts": lambda C, t, x: (C.rt[t["slot"]] * np.median(C.st[_g(C, t)]) >= np.asarray(x)).astype(float),
    "new slot, starts game by game": lambda C, t, x: _binomial(C.st[_g(C, t)].mean() / 162.0, C.rt[t["slot"]], x),
    "catchers kept in": lambda C, t, x: emp(C.st * C.rt[t["slot"]], x),
    "opening-day lineups": lambda C, t, x: emp(C.od[t["is_c"]] * C.rt[t["slot"]], x),
    "planned regulars who opened the season": lambda C, t, x: emp(C.st[_g(C, t) & C.opened] * C.rt[t["slot"]], x),
    "hitters with 100 or more starts": lambda C, t, x: emp(_hs(C, t, 0, 100) * C.rt[t["slot"]], x),
    "hitters with 400 or more plate appearances": lambda C, t, x: emp(_hs(C, t, 1, 400) * C.rt[t["slot"]], x),
    "hitters with 502 or more plate appearances": lambda C, t, x: emp(_hs(C, t, 1, 502) * C.rt[t["slot"]], x),
    "planned regulars' own plate appearances": lambda C, t, x: emp(C.pa[_g(C, t)], x),
    "plate appearances per appearance, appearances for starts": lambda C, t, x: emp(C.app[_g(C, t)] * C.rta[t["slot"]], x),
    "the slot's plate appearances with its substitutes', per start": lambda C, t, x: emp(
        C.st[_g(C, t)] * C.rts[t["slot"]], x),
    "the full nine's own plate appearances": lambda C, t, x: emp(C.pa, x),
    "the full nine at the slot, own plate appearances": lambda C, t, x: emp(C.pa[C.ord == t["slot"]], x),
    "the look-back's own hitters' seasons": lambda C, t, x: emp(C.pa[C.est], x),
    "the club's own regulars": lambda C, t, x: emp(C.st[_g(C, t) & (C.team == t["team"])] * C.rt[t["slot"]], x),
    "every game started": lambda C, t, x: (162 * C.rt[t["slot"]] >= np.asarray(x)).astype(float),
    "own record given a weight of 2 per cent": lambda C, t, x: _shift(C, t, x, 0.02, "own"),
    "own record given a weight of 5 per cent": lambda C, t, x: _shift(C, t, x, 0.05, "own"),
    "own record given a weight of 10 per cent": lambda C, t, x: _shift(C, t, x, 0.10, "own"),
    "vendor projection given a weight of 2 per cent": lambda C, t, x: _shift(C, t, x, 0.02, "vendor"),
    "vendor projection given a weight of 5 per cent": lambda C, t, x: _shift(C, t, x, 0.05, "vendor"),
    "own seasons blended in at 5 per cent": lambda C, t, x: 0.95 * emp(C.st[_g(C, t)] * C.rt[t["slot"]], x) +
    0.05 * emp(_own(t), x),
    "own seasons blended in at 10 per cent": lambda C, t, x: 0.90 * emp(C.st[_g(C, t)] * C.rt[t["slot"]], x) +
    0.10 * emp(_own(t), x),
    "vendor projection given a weight of 10 per cent": lambda C, t, x: _shift(C, t, x, 0.10, "vendor"),
    "regression on the projection, the slot and catching": lambda C, t, x: emp(
        C.beta @ np.array([1.0, t["proj"], C.rt[t["slot"]], float(t["is_c"])]) + C.res, x),
    "regression on the slot and catching": lambda C, t, x: emp(
        C.beta0 @ np.array([1.0, C.rt[t["slot"]], float(t["is_c"])]) + C.res0, x),
}
# within a bin of the golden: each loses the look-back (FIN-14 4.4)
NEAR = ("regression on the projection, the slot and catching", "regression on the slot and catching",
        "own record given a weight of 2 per cent", "vendor projection given a weight of 2 per cent",
        "own seasons blended in at 5 per cent")
# read at the signings only (FIN-14 4.4's own hitters need two earlier seasons, which 2023 has not got)
AT_SIGNING_ONLY = ("the look-back's own hitters' seasons",)


def _same_slot(C, t, x):
    m = _g(C, t) & (C.ord == t["slot"])
    return emp(C.pa[m], x) if m.any() else emp(C.st[_g(C, t)] * C.rt[t["slot"]], x)


def _binomial(p, r, x):
    return np.array([1 - NORM.cdf((t / r - 0.5 - 162 * p) / np.sqrt(162 * p * (1 - p))) for t in x])


def _hs(C, t, col, th):
    h = C.hs[t["is_c"]]
    return h[h[:, col] >= th, 0]


def _shift(C, t, x, w, src):
    base = C.st[_g(C, t)] * C.rt[t["slot"]]
    dev = (_own(t).mean() - base.mean()) if src == "own" else (t["proj"] - C.vp.mean())
    return emp(base + w * dev, x)


def test_row(T, proj, r, Y):
    """A depth-chart lineup row as FIN-14 4.4 tests it for season Y."""
    prev = [q["ord"] for q in T.planned if q["season"] == Y - 1 and q["lg"] == r["lg"]]
    return dict(lg=r["lg"], is_c=r["pos"] == "C", slot=r["ord"], age=r["age"], team=r["team"],
                prev=prev[0] if prev else None, proj=proj.get((Y, r["lg"])),
                own=[T.pa[(Y - j, r["lg"])] or None if Y - j >= min(T.years) else None for j in range(1, 7)],
                ownst=[T.starts[(Y - j, r["lg"])] if T.pa[(Y - j, r["lg"])] and Y - j >= min(T.years) else None
                       for j in range(1, 4)])


def five_rows(P, T):
    proj = projections(P)
    out = []
    for f in five(P):
        out.append(dict(name=f["name"], lines=f["lines"], lg=f["lg"], is_c=False, slot=f["slot"], age=f["age"],
                        team="RAL", prev=f["old"], proj=proj.get((2027, f["lg"])),
                        own=[T.pa[(2027 - j, f["lg"])] for j in range(1, 7)],
                        ownst=[T.starts[(2027 - j, f["lg"])] for j in range(1, 4)]))
    return out


def lookback(P, T=None, names=None, seasons_=None, tests_nc=False, est_club=False, strict=False):
    """FIN-14 4.4 executed: each basis applied to each of the 2023 to 2026 seasons with what was published before it,
    to every hitter in a starting lineup of that season's January 31 depth charts who made at least 100 plate
    appearances in each of the two seasons before; the average over every hitter, season and threshold (450, 500, 550,
    600 and 650) of the squared difference between the probability and the outcome."""
    T = T or seasons(P)
    proj = projections(P)
    names = names or [n for n in BASES if n not in AT_SIGNING_ONLY]
    sq = defaultdict(list)
    n_tests = 0
    if est_club:
        club_pa = defaultdict(int)
        for (y, g, d, club, pid, slot, st, pa, pos) in P.lines:
            club_pa[(y, pid, club)] += pa
        inv = {v: k for k, v in {"LVS": "LV", "SAT": "SA", "NOL": "NO", "SLC": "SL", "HRD": "HR", "JAX": "JAC",
                                 "SAC": "SCR"}.items()}
    for Y in seasons_ or LB_SEASONS:
        if est_club:
            ok = lambda r: all(club_pa[(Y - j, r["lg"], inv.get(r["team"], r["team"]))] >= LB_MIN_PA for j in (1, 2))
        else:
            ok = lambda r: all(T.pa[(Y - j, r["lg"])] >= LB_MIN_PA for j in (1, 2))
        tests = [test_row(T, proj, r, Y) for r in T.planned if r["season"] == Y and ok(r) and
                 not (tests_nc and r["pos"] == "C")]
        if not tests:           # a season whose hitters have no two earlier seasons in the box scores (2022)
            continue
        C = context(P, T, Y)
        n_tests += len(tests)
        for t in tests:
            out = ((T.pa[(Y, t["lg"])] > LB_THRESHOLDS) if strict else (T.pa[(Y, t["lg"])] >= LB_THRESHOLDS)).astype(float)
            for nm in names:
                sq[nm].append((BASES[nm](C, t, LB_THRESHOLDS) - out) ** 2)
    score = {nm: float(np.mean(v)) for nm, v in sq.items()}
    per_row = {nm: np.array(v).mean(axis=1) for nm, v in sq.items()}
    return score, n_tests, per_row


def accruals(P, T=None, names=None):
    """Every basis applied at the signings (the 2021 to 2026 seasons): each player's expected amount and the total."""
    T = T or seasons(P)
    C = context(P, T, 2027)
    out = {}
    for nm in names or BASES:
        per = {}
        for t in five_rows(P, T):
            thr = [x for _, x in t["lines"]]
            per[t["name"]] = float(sum(a * p for (a, _), p in zip(t["lines"], BASES[nm](C, t, thr))))
        out[nm] = per
    return out


def rivals(P, T=None):
    """Every rival route's total at the signings (each at least a bin from the golden), the near routes (within a bin,
    losing the look-back) and the golden."""
    T = T or seasons(P)
    G = golden(P, T)
    A = accruals(P, T)
    out = {nm: sum(v.values()) for nm, v in A.items() if nm not in FAMILY + NEAR + LANDS}
    # windows of the full nine's own plate appearances, and the 2024 to 2026 seasons free of an injured-list stint
    for y0, lab in ((2022, "2022-2026"), (2024, "2024-2026")):
        v = np.array([T.pa[(r["season"], r["lg"])] for r in T.planned if r["season"] >= y0], float)
        out["the full nine's own plate appearances, %s" % lab] = sum(ev(f["lines"], v) for f in G["five"])
    hurt = injured_seasons(P)
    v = np.array([T.starts[(r["season"], r["lg"])] for r in T.planned if r["pos"] != "C" and r["season"] >= 2024 and
                  (r["season"], r["lg"]) not in hurt], float)
    out["seasons with no injured-list stint, 2024-2026"] = sum(ev(f["lines"], v * G["rate"][f["slot"]]) for f in G["five"])
    near = {nm: sum(A[nm].values()) for nm in NEAR}
    return out, G, near, A


def injured_seasons(P):
    tx = json.load(open(os.path.join(P.target, "lsf_transactions_2024_2026.json"), encoding="utf-8"))["transactions"]
    return {(int(t["date"][:4]), int(t["player_id"])) for t in tx if t["type"] == "IL_PLACEMENT"}


def hitters_nonc(P, T):
    """Every non-catcher hitter-season in the box scores (a player's most common position not catcher)."""
    posc = defaultdict(lambda: defaultdict(int))
    for (y, g, d, club, pid, slot, st, pa, pos) in P.lines:
        if st:
            posc[(y, pid)][pos] += 1
    return sorted(k for k, d in posc.items() if max(d, key=d.get) != 2)


def constructions(P, T=None):
    """Constructions of the golden method that must round alike (C1): windows, slot-rate windows, ages, plate
    appearances rescaled or kept at the slot, and the loose regular definitions that select almost the same seasons."""
    T = T or seasons(P)
    G = golden(P, T)
    F = G["five"]
    rt = G["rate"]
    out = {}
    for y0 in (2022, 2023, 2024, 2025):
        st, _ = planned_starts(T, years=set(range(y0, 2027)))
        out["seasons from %d" % y0] = {f["name"]: ev(f["lines"], st * rt[f["slot"]]) for f in F}
    st = G["starts"]
    for ys in ((2024, 2025, 2026), (2025, 2026)):
        r2 = rate(T, list(ys))
        out["slot rates %d-%d" % (ys[0], ys[-1])] = {f["name"]: ev(f["lines"], st * r2[f["slot"]]) for f in F}
    _, rows = planned_starts(T)
    ages = np.array([r["age"] for r in rows])
    for w in (2, 3):
        out["ages within %d" % w] = {f["name"]: ev(f["lines"], st[np.abs(ages - f["age"]) <= w] * rt[f["slot"]]) for f in F}
    out["ages 30 to 34"] = {f["name"]: ev(f["lines"], st[(ages >= 30) & (ages <= 34)] * rt[f["slot"]]) for f in F}
    slot_of = np.array([r["ord"] for r in rows])
    pas = np.array([T.pa[(r["season"], r["lg"])] for r in rows], float)
    spas = np.array([T.start_pa[(r["season"], r["lg"])] for r in rows], float)
    sr = np.array([rt[int(q)] for q in slot_of])
    out["own plate appearances rescaled to the new slot"] = {f["name"]: ev(f["lines"], pas * rt[f["slot"]] / sr) for f in F}
    out["starting plate appearances rescaled"] = {f["name"]: ev(f["lines"], spas * rt[f["slot"]] / sr) for f in F}
    out["regulars planned at the new slot"] = {f["name"]: ev(f["lines"], pas[slot_of == f["slot"]]) for f in F}
    own = rate_of_club(P, "RAL")
    out["the club's own plate appearances per start"] = {f["name"]: ev(f["lines"], st * own[f["slot"]]) for f in F}
    top = []
    for (y, club, pos), d in T.pos_starts.items():
        if pos in (2, 11):
            continue
        top.append(T.starts[(y, max(sorted(d), key=lambda q: d[q]))])
    out["most starts at each position"] = {f["name"]: ev(f["lines"], np.array(top) * rt[f["slot"]]) for f in F}
    top8 = []
    for (y, club), d in T.club_starts.items():
        ks = sorted(d, key=lambda q: (-d[q], q))
        ks = [q for q in ks if q not in catchers_of(P, T, y, club)][:8]
        top8 += [T.starts[(y, q)] for q in ks]
    out["eight non-catchers with most starts per club"] = {f["name"]: ev(f["lines"], np.array(top8) * rt[f["slot"]]) for f in F}
    poss = np.array([r["pos"] for r in rows])
    out["the same position"] = {f["name"]: ev(f["lines"], st[poss == f["pos"]] * rt[f["slot"]]) for f in F}
    est = np.array([all(T.pa[(r["season"] - j, r["lg"])] >= LB_MIN_PA for j in (1, 2)) for r in rows])
    out["regulars with two earlier seasons of 100 plate appearances"] = {f["name"]: ev(f["lines"], st[est] * rt[f["slot"]])
                                                                       for f in F}
    nonc = hitters_nonc(P, T)
    d81 = np.array([T.starts[k] for k in nonc if T.starts[k] >= 81], float)
    out["hitters with 81 or more starts"] = {f["name"]: ev(f["lines"], d81 * rt[f["slot"]]) for f in F}
    d300 = np.array([T.starts[k] for k in nonc if T.pa[k] >= 300], float)
    out["hitters with 300 or more plate appearances"] = {f["name"]: ev(f["lines"], d300 * rt[f["slot"]]) for f in F}
    return out, G


def rate_of_club(P, club):
    sp, sn = defaultdict(int), defaultdict(int)
    for (y, g, d, c, pid, slot, st, pa, pos) in P.lines:
        if st and c == club:
            sp[slot] += pa
            sn[slot] += 1
    return {s: sp[s] / sn[s] for s in SLOTS}


_CATCHERS = {}


def catchers_of(P, T, y, club):
    if not _CATCHERS:
        posc = defaultdict(lambda: defaultdict(int))
        for (yy, g, d, c, pid, slot, st, pa, pos) in P.lines:
            if st:
                posc[(yy, pid)][pos] += 1
        for k, d in posc.items():
            if max(d, key=d.get) == 2:
                _CATCHERS.setdefault(k[0], set()).add(k[1])
    return _CATCHERS.get(y, set())


def bootstrap(P, T=None, n=300, seed=151):
    T = T or seasons(P)
    G = golden(P, T)
    st = G["starts"]
    F = G["five"]
    rng = np.random.default_rng(seed)
    tot, per = [], defaultdict(list)
    for _ in range(n):
        i = rng.integers(0, len(st), len(st))
        vals = {f["name"]: ev(f["lines"], st[i] * G["rate"][f["slot"]]) for f in F}
        tot.append(sum(vals.values()))
        for k, v in vals.items():
            per[k].append(v)
    return float(np.std(tot)), {k: float(np.std(v)) for k, v in per.items()}


# ---------------------------------------------------------------------------------------------------------- the ask
F_PAY = "ral_payroll_player_earnings_2023_2026.csv"
F_HRIS = "ral_hris_legacy_employee_map.csv"
F_REG = "ral_player_contract_register_2023_2026.xlsx"
OTHER_COL = "Other playing-time bonuses (each season)"
AWARD_COL = "Award bonuses"
ASK_SEASONS = (2023, 2024, 2025)
ASK_MODES = ("no_catchup", "cracker", "payyear", "e45_whole", "e45_all", "map", "register_id", "gross", "settle_whole", "double",
             "cap", "roundest")
ROUND_UNITS = (1000000, 500000, 250000, 100000, 50000, 25000, 10000, 5000, 1000)
SALARY_CODES = ("E01", "SAL")


def report_name(name):
    nm = fold(name)
    first, last = nm.split(" ", 1)
    return "%s, %s" % (last.upper(), first.upper())


def register(target):
    """The contract register's rows as dicts (the header row is the one that starts with Player)."""
    wb = openpyxl.load_workbook(os.path.join(target, F_REG), read_only=True)
    hdr, out = None, []
    for row in wb["Register"].iter_rows(values_only=True):
        if row and row[0] == "Player":
            hdr = list(row)
            continue
        if hdr and row and row[0]:
            out.append(dict(zip(hdr, row)))
    return out


def parse_terms(text):
    """'Innings pitched: $100,000 each for 150, 165 and 180; games started: $45,000 for 30' as
    [(measure, amount, (lines))]."""
    out = []
    for part in (text or "").split(";"):
        m = re.match(r"\s*([A-Za-z ]+): \$([\d,]+) (?:each )?for ([\d, and]+)\s*$", part)
        if m:
            out.append((m.group(1).strip().lower(), int(m.group(2).replace(",", "")),
                        tuple(int(x) for x in re.findall(r"\d+", m.group(3)))))
    return out


def parse_awards(text):
    return [int(a.replace(",", "")) for a in re.findall(r"\$([\d,]+) for", text or "")]


def seasons_of(text):
    a, b = (str(text).split("-") + [str(text)])[:2]
    return set(range(int(a), int(b) + 1))


def legacy_key(name_hris):
    """The register's 'Last, First' as the legacy payroll writes it: 'LAST FIRST', upper case, no accents."""
    last, first = name_hris.split(", ")
    return "%s %s" % (fold(last).upper(), fold(first).upper())


def combos(terms):
    """Every amount a season's terms allow: [(amount per line, lines)] -> the set of sums."""
    sums = {0}
    for a, n in terms:
        sums = {x + a * j for x in sums for j in range(n + 1)}
    return sums


def ask(P, mode="golden", detail=False):
    """The playing-time bonuses other than plate-appearance bonuses the club paid for each of the 2023, 2024 and 2025
    seasons, by player (the register's name) and season. The golden: every playing-time payment to a player whose
    contract carries other playing-time terms, the season-end (October) payment to its season and the January payment
    to the season before, whatever the payroll code (the legacy payroll paid the season-end halves of the pitchers'
    bonuses on its catch-all code), net of voids, less an award bonus the legacy season-end line carries (the October
    payment exceeding the January one by an award term of his), and less the salary a contract deferred to its end,
    which the season-end line after its last season also carries (the register's salary for each of its seasons less
    the salary lines the payroll paid in it); each legacy line is the player its legacy name gives.
    Modes give the stops: 'cracker' (the playing-time codes the earning-code table gives, E40 and BPA, each legacy
    line the player its legacy name gives, as a key-uniqueness sweep finds), 'payyear' (the same by the year paid),
    'e45_whole' (the catch-all lines taken whole), 'e45_all' (every bonus line to those players, any month, taken
    whole), 'map' (every legacy number through the HRIS map), 'register_id' (every legacy number to the register's own
    Employee ID), 'gross' (the void's reversal left out), 'settle_whole' (the full handling of the catch-all with every
    hazard, the season-end line after a contract's last season taken whole), 'double' (twice the January half), 'cap'
    ('settle_whole' with each row cut to the largest amount his terms allow at or below it), 'roundest' ('settle_whole'
    with a line his terms cannot hold split into a bonus they allow and the roundest remainder), 'golden'."""
    pay = read_csv(os.path.join(P.target, F_PAY))
    hmap = {r["legacy_employee_no"]: r["employee_id"] for r in read_csv(os.path.join(P.target, F_HRIS))}
    reg = register(P.target)
    other, awards, by_id, by_legacy_name = defaultdict(set), defaultdict(set), {}, {}
    contracts, terms = [], {}
    for r in reg:
        name = r["Player"]
        by_id.setdefault(str(r["Employee ID"]), name)
        by_legacy_name[legacy_key(name)] = name
        tt = parse_terms(r[OTHER_COL])
        if tt:
            ss = seasons_of(r["Seasons"])
            other[name] |= ss
            contracts.append((name, ss, {y: r.get("Salary %d" % y) for y in range(2023, 2027)}))
            for y in ss:
                terms[(name, y)] = [(a, len(lines)) for _, a, lines in tt]
        for a in parse_awards(r[AWARD_COL]):
            awards[name].add(a)

    def who(r):
        e = r["employee_id"]
        if not e.isdigit():                       # an HRIS line
            return by_id.get(e)
        if mode == "register_id":
            return by_id.get(e) or by_id.get(hmap.get(e, ""))
        if mode in ("e45_whole", "e45_all", "map"):
            return by_id.get(hmap.get(e, "")) or by_id.get(e)
        return by_legacy_name.get(r["employee_name"])

    def season_of(d):
        return d.year if d.month >= 7 else d.year - 1

    # the salary each contract deferred to its end: its register salary less the salary lines paid, season by season
    paid_salary = defaultdict(float)
    for r in pay:
        if r["earning_code"] in SALARY_CODES:
            paid_salary[(who(r), ymd(r["pay_date"]).year)] += float(r["amount"])
    deferral = {}
    if mode not in ("settle_whole", "cap", "roundest", "e45_whole", "e45_all", "double", "cracker", "payyear"):
        for name, ss, sal in contracts:
            short = sum(sal[y] - paid_salary[(name, y)] for y in ss if sal.get(y))
            if short > 0.5:
                deferral[(max(ss), name)] = round(short, 2)

    parts = defaultdict(lambda: defaultdict(float))
    for r in pay:
        name = who(r)
        if name not in other:
            continue
        d = ymd(r["pay_date"])
        amt = float(r["amount"])
        code = r["earning_code"]
        if mode == "gross" and amt < 0:
            continue
        if mode in ("cracker", "payyear", "double"):
            if code not in ("E40", "BPA"):
                continue
            if mode == "double":
                if d.month == 1:
                    parts[(d.year - 1, name)]["all"] += 2 * amt
                continue
            y = d.year if mode == "payyear" else season_of(d)
            parts[(y, name)]["all"] += amt
            continue
        if code not in ("E40", "E45", "BPA"):
            continue
        if mode == "e45_all":
            parts[(season_of(d), name)]["all"] += amt
            continue
        if d.month == 10:
            parts[(d.year, name)]["oct"] += amt
        elif d.month == 1:
            parts[(d.year - 1, name)]["jan"] += amt
    # a season whose January half never came, on a contract running on past it: the next season-end line exceeds the
    # January half after it by exactly that half, so the excess goes back to the season it was earned in
    if mode not in ("no_catchup", "settle_whole", "cap", "roundest", "e45_whole", "e45_all", "double", "cracker",
                    "payyear", "gross", "map", "register_id"):
        for name, ss, sal in contracts:
            for y in sorted(ss):
                if y == max(ss) or y + 1 not in ss:
                    continue
                a, b = parts.get((y, name)), parts.get((y + 1, name))
                if not a or not b or a["jan"] > 0.5 or a["oct"] < 0.5:
                    continue
                last_next = (y + 1) == max(ss)
                if last_next:
                    continue
                excess = round(b["oct"] - deferral.get((y + 1, name), 0.0) - b["jan"], 2)
                half = next((a["oct"] - w for w in sorted(awards[name] | {0.0}) if abs(a["oct"] - w - excess) < 0.5), None)
                if half is not None and excess > 0.5:
                    b["oct"] -= excess
                    a["jan"] += excess
    out = {}
    for (y, name), p in parts.items():
        if y not in ASK_SEASONS or y not in other[name]:
            continue
        award = 0.0
        dfr = deferral.get((y, name), 0.0)
        if "all" in p:
            v = p["all"]
        else:
            v = p["oct"] - dfr + p["jan"]
            diff = round(p["oct"] - dfr - p["jan"], 2)
            if mode != "e45_whole" and p["jan"] > 0.5 and diff > 0.5 and diff in awards[name]:
                award = diff
                v -= diff
        if mode == "cap":
            v = max(x for x in combos(terms[(name, y)]) if x <= v + 0.5)
        if mode == "roundest" and v > max(combos(terms[(name, y)])) + 0.5:
            v = max((next((u for u in ROUND_UNITS if round(v - b) % u == 0), 1), b)
                    for b in combos(terms[(name, y)]) if b <= v + 0.5)[1]
        if abs(v) > 0.5:
            out[(y, name)] = dict(oct=p["oct"], jan=p["jan"], award=award, deferral=dfr, paid=v) if detail else v
    return out


def ask_totals(d):
    tot = defaultdict(float)
    for (y, n), v in d.items():
        tot[y] += v
    return dict(tot)
