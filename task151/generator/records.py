"""The shipped records. The League statistics feed: box-score batting lines in the legacy layout (2021-2024) and the
platform layout (2025-2026), the player ID crosswalk built at the platform change, the club directory, the transactions
feed and the calendar. The League office's weekly injured-list report and its 2027 schedule. The projection vendor's
winter depth charts and hitter projections. The club's own files: the 2027 depth chart, the player contract register,
the payroll export, the HRIS employee map and the earning codes. Each writer formats the clean world the way the system
that owns the file writes it; nothing here changes a fact."""
from __future__ import annotations

import datetime as dt
import json
import os
from collections import defaultdict

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

import common as K
import names as NM
import world as WD

D = dt.timedelta

F_LEGACY = "lsf_batting_2021_2024.txt"
F_BAT = "lsf_batting_2025_2026.csv"
F_XWALK = "lsf_player_id_crosswalk.csv"
F_CLUBS = "league_club_directory.csv"
F_TRANS = "lsf_transactions_2024_2026.json"
F_ILREP = "league_office_il_report_2024_2026.csv"
F_CAL = "league_calendar_2021_2027.csv"
F_SCHED27 = "league_schedule_2027.csv"
F_DC = "hvbd_winter_depth_charts_2021_2026.xlsx"
F_PROJ = "hvbd_hitter_projections_2022_2027.csv"
F_DC27 = "ral_depth_chart_2027.xlsx"
F_REG = "ral_player_contract_register_2023_2026.xlsx"
F_PAY = "ral_payroll_player_earnings_2023_2026.csv"
F_HRIS = "ral_hris_legacy_employee_map.csv"
F_CODES = "ral_payroll_earning_codes.csv"

LEGACY_POS = {"P": 1, "C": 2, "1B": 3, "2B": 4, "3B": 5, "SS": 6, "LF": 7, "CF": 8, "RF": 9, "DH": 10, "PH": 11}
# clubs whose legacy statistics code differed from the platform abbreviation
LEGACY_CLUB = {"LVS": "LAS", "SAT": "SAN", "HRD": "NFK", "CGY": "CAL", "SLC": "SLK", "OKC": "OKL", "NOL": "NOR"}
# the vendor's own club abbreviations where they differ from the League's
VENDOR_CLUB = {"LVS": "LV", "SAT": "SA", "NOL": "NO", "SLC": "SL", "HRD": "HR", "JAX": "JAC", "SAC": "SCR"}


# ----------------------------------------------------------------------------------------------------- identifiers
def club_ids():
    r = K.rng("team_ids")
    ids = list(range(101, 131))
    r.shuffle(ids)
    return dict(zip(WD.CODES, ids))


def vendor_ids(L):
    r = K.rng("hv_ids")
    keys = sorted(k for k in L.players if not k.startswith("q"))
    nums = r.choice(np.arange(10000, 99999), len(keys), replace=False)
    return {k: "hv%05d" % int(n) for k, n in zip(keys, nums)}


def legacy_codes(L):
    """The legacy statistics system's player codes, for every player who appeared from 2021 to 2024."""
    seen = []
    for y in K.LEGACY_SEASONS:
        for ln in L.lines[y]:
            seen.append(ln["key"])
    out, taken = {}, set()
    for k in sorted(set(seen), key=lambda q: (L.players[q].pid, q)):
        out[k] = NM.legacy_code(L.players[k].name, taken)
    return out


# ------------------------------------------------------------------------------------------------------- box scores
def write_legacy(L, out):
    lc = legacy_codes(L)
    rows = []
    for y in K.LEGACY_SEASONS:
        for ln in sorted(L.lines[y], key=lambda d: (d["date"], d["gnum"], d["home"], d["ha"] == "H", d["slot"], d["seq"])):
            hc = LEGACY_CLUB.get(ln["home"], ln["home"])
            gid = "%s%s%d" % (hc, ln["date"].strftime("%Y%m%d"), ln["gnum"])
            rows.append([gid, K.us(ln["date"]), ln["gnum"], hc, LEGACY_CLUB.get(ln["club"], ln["club"]),
                         LEGACY_CLUB.get(ln["opp"], ln["opp"]), lc[ln["key"]], ln["slot"], ln["seq"] + 1,
                         LEGACY_POS[ln["pos"]], ln["ab"], ln["r"], ln["h"], ln["d2"], ln["t3"], ln["hr"], ln["rbi"],
                         ln["bb"], ln["so"], ln["hbp"], ln["sh"], ln["sf"], ln["ci"], K.us(ln["completed"])])
    hdr = ["GAME_ID", "GAME_DT", "GAME_CT", "HOME_CD", "TEAM_CD", "OPP_CD", "PLAYER_CD", "LINEUP", "SEQ", "POS", "AB", "R",
           "H", "D2", "T3", "HR", "RBI", "BB", "SO", "HBP", "SH", "SF", "CI", "COMPL_DT"]
    K.write_csv(os.path.join(out, F_LEGACY), hdr, rows, delimiter="|", crlf=True)
    return len(rows), lc


def write_platform(L, out, ids):
    rows = []
    for y in K.NEW_SEASONS:
        for ln in sorted(L.lines[y], key=lambda d: (d["date"], d["gnum"], d["pk"], d["ha"] == "H", d["slot"], d["seq"])):
            rows.append([ln["pk"], K.iso(ln["date"]), max(ln["gnum"], 1), ids[ln["home"]], ids[ln["club"]],
                         ids[ln["opp"]], L.players[ln["key"]].pid, ln["slot"] * 100 + ln["seq"],
                         "true" if ln["seq"] else "false", ln["pos"], ln["pa"],
                         ln["ab"], ln["r"], ln["h"], ln["d2"], ln["t3"], ln["hr"], ln["rbi"], ln["bb"], ln["so"],
                         ln["hbp"], ln["sh"], ln["sf"], ln["ci"], K.iso(ln["completed"])])
    hdr = ["game_pk", "official_date", "game_number", "venue_team_id", "team_id", "opponent_team_id", "player_id",
           "batting_order", "is_substitute", "position", "pa", "ab", "r", "h", "doubles", "triples", "hr", "rbi", "bb",
           "so", "hbp",
           "sac_bunts", "sac_flies", "catchers_interference", "resumed_date"]
    K.write_csv(os.path.join(out, F_BAT), hdr, rows)
    return len(rows)


def write_crosswalk(L, out, lc):
    rows = []
    for k, code in sorted(lc.items(), key=lambda kv: kv[1]):
        p = L.players[k]
        nm = NM.ascii_fold(p.name)
        suffix = ""
        if nm.endswith(" Jr."):
            nm, suffix = nm[:-4], "Jr."
        first, last = nm.split(" ", 1)
        rows.append([code, p.pid, last, first, suffix, K.iso(p.born)])
    K.write_csv(os.path.join(out, F_XWALK), ["legacy_id", "player_id", "last_name", "first_name", "suffix", "birth_date"],
                rows)
    return len(rows)


def write_clubs(out, ids):
    rows = []
    for code, city, nick, tz in sorted(WD.CLUBS, key=lambda c: ids[c[0]]):
        rows.append([ids[code], code, LEGACY_CLUB.get(code, code), "%s %s" % (city, nick), city, tz,
                     WD.AFFILIATE[code]])
    K.write_csv(os.path.join(out, F_CLUBS), ["team_id", "abbrev", "legacy_code", "club_name", "city", "time_zone",
                                             "triple_a_affiliate"], rows)
    return len(rows)


def write_calendar(out):
    rows = []
    for y in K.HIST + [2027]:
        od, last, asb = K.SEASON[y]
        rows.append([y, K.iso(od), K.iso(asb), K.iso(asb + D(days=3)), K.iso(last), K.iso(last + D(days=3))])
    K.write_csv(os.path.join(out, F_CAL), ["season", "opening_day", "all_star_break_start", "all_star_break_end",
                                           "last_scheduled_day", "postseason_start"], rows)
    return len(rows)


def schedule_2027():
    """The League's 2027 schedule as published in August 2026 (162 games a club, the same pattern as every season)."""
    r = K.rng("schedule", 2027)
    od, last, asb = K.SEASON[2027]
    asb_days = {asb + D(days=k) for k in range(4)}
    days = [d for d in K.daterange(od, last) if d not in asb_days]
    gaps = set(int(x) for x in r.choice(np.arange(1, 54), 20, replace=False))
    slots, i = [], 0
    for s in range(54):
        slots.append(days[i:i + 3])
        i += 3
        if s + 1 in gaps:
            i += 1
    home_ct = defaultdict(int)
    games = []
    for s, sl in enumerate(slots):
        order = list(WD.CODES)
        r.shuffle(order)
        for j in range(0, 30, 2):
            a, b = order[j], order[j + 1]
            h, v = (a, b) if (home_ct[a] < home_ct[b] or (home_ct[a] == home_ct[b] and r.random() < 0.5)) else (b, a)
            home_ct[h] += 1
            for d in sl:
                t = ["13:05", "18:40", "19:05", "19:10", "19:40", "20:10", "16:10", "13:35"][int(r.integers(8))]
                if d.weekday() == 6:
                    t = ["13:05", "13:35", "13:10", "16:10"][int(r.integers(4))]
                games.append((d, h, v, t))
    games.sort()
    return games


def write_schedule_2027(out, ids):
    rows = []
    for n, (d, h, v, t) in enumerate(schedule_2027()):
        rows.append([K.iso(d), d.strftime("%a"), h, v, ids[h], ids[v], t, dict((c[0], c[3]) for c in WD.CLUBS)[h]])
    K.write_csv(os.path.join(out, F_SCHED27), ["game_date", "day", "home", "away", "home_team_id", "away_team_id",
                                               "first_pitch_local", "venue_time_zone"], rows)
    return len(rows)


# ------------------------------------------------------------------------------------------------ transactions feed
def team_name(code):
    c = [x for x in WD.CLUBS if x[0] == code][0]
    return "%s %s" % (c[1], c[2])


def pos_of(L, k):
    p = L.players[k].pos
    return "RHP" if p == "P" and int(L.players[k].pid) % 3 else ("LHP" if p == "P" else p)


def transactions(L, ids):
    """Every move the League's transaction wire carries for 2024 to 2026, dated the day the club made it."""
    ev = []
    for s in L.stints:
        if s["season"] not in K.NEW_SEASONS + [2024]:
            continue
        k = s["key"]
        p = L.players[k]
        for (d, club, kind, lst, txt) in s["moves"]:
            pos = pos_of(L, k)
            if kind == "PLACED":
                desc = "%s placed %s %s on the %s injured list. %s." % (team_name(club), pos, p.name, lst.replace(" IL", ""),
                                                                       txt[0].upper() + txt[1:])
                typ = "IL_PLACEMENT"
            elif kind == "REHAB ASSIGNMENT":
                desc = "%s sent %s %s on a rehab assignment to %s." % (team_name(club), pos, p.name,
                                                                       WD.AFFILIATE[club])
                typ = "REHAB_ASSIGNMENT"
            elif kind == "ACTIVATED":
                desc = "%s activated %s %s from the %s injured list." % (team_name(club), pos, p.name,
                                                                          lst.replace(" IL", ""))
                typ = "IL_ACTIVATION"
            else:
                desc = "%s reinstated %s %s from the %s injured list." % (team_name(club), pos, p.name,
                                                                           lst.replace(" IL", ""))
                typ = "IL_REINSTATEMENT"
            ev.append((d, club, k, typ, desc))
    for ls in L.lists:
        if ls["season"] < 2024:
            continue
        k = ls["key"]
        p = L.players[k]
        nm = ls["list"]
        ev.append((ls["start"], ls["club"], k, "LIST_PLACEMENT",
                   "%s placed %s %s on the %s." % (team_name(ls["club"]), pos_of(L, k), p.name, nm)))
        ev.append((ls["back"], ls["club"], k, "LIST_ACTIVATION",
                   "%s activated %s %s from the %s." % (team_name(ls["club"]), pos_of(L, k), p.name, nm)))
    for t in L.trades:
        if t["season"] < 2024:
            continue
        p = L.players[t["key"]]
        ev.append((t["date"], t["frm"], t["key"], "TRADE", "%s traded %s %s to %s." % (
            team_name(t["frm"]), pos_of(L, t["key"]), p.name, team_name(t["to"]))))
    for (y, c, k, d0, d1) in L.callups:
        if y < 2024:
            continue
        p = L.players[k]
        ev.append((d0, c, k, "RECALL", "%s recalled %s %s from %s." % (team_name(c), pos_of(L, k), p.name,
                                                                               WD.AFFILIATE[c])))
        if d1 < K.SEASON[y][1] - D(days=3):
            ev.append((d1 + D(days=1), c, k, "OPTION", "%s optioned %s %s to %s." % (
                team_name(c), pos_of(L, k), p.name, WD.AFFILIATE[c])))
    ev.sort(key=lambda e: (e[0], ids[e[1]], L.players[e[2]].pid, e[3]))
    out = []
    r = K.rng("trans_ids")
    tid = 1880000
    for d, club, k, typ, desc in ev:
        tid += int(r.integers(1, 40))
        out.append(dict(transaction_id=tid, date=K.iso(d), team_id=ids[club], team=team_name(club),
                        player_id=L.players[k].pid, player=L.players[k].name, type=typ, description=desc))
    return out


def write_transactions(L, out, ids):
    ev = transactions(L, ids)
    doc = dict(feed="League Statistics Feed: transactions", extracted_at=K.PULL_STAMP.strftime("%Y-%m-%dT%H:%M:%S-05:00"),
               start_date="2024-01-01", end_date=K.iso(K.ASOF), count=len(ev), transactions=ev)
    with open(os.path.join(out, F_TRANS), "w", encoding="utf-8", newline="\n") as f:
        json.dump(doc, f, ensure_ascii=False, indent=1)
        f.write("\n")
    return len(ev)


# ------------------------------------------------------------------------------------------- weekly injured-list report
def il_report(L):
    """The League office's Monday report: every player on an injured list that day, the list, the date of the
    transaction that put him on it and the first date he may be reinstated."""
    rows = []
    for y in (2024, 2025, 2026):
        od, last, _ = K.SEASON[y]
        mondays = [d for d in K.daterange(od, last) if d.weekday() == 0]
        for s in L.stints:
            if s["season"] != y:
                continue
            k = s["key"]
            mv = s["moves"]
            end = s["reinst"] or [m[0] for m in mv if m[2] == "REINSTATED"][0]
            placed = [m for m in mv if m[2] == "PLACED"]
            for d in mondays:
                if not (s["eff"] <= d < end) or placed[0][0] > d:
                    continue
                m = [m for m in placed if m[0] <= d][-1]
                lst = m[3].replace(" IL", "")
                elig = s["eff"] + D(days=WD.MIN_DAYS[lst])
                rows.append((d, WD.club_on(L, y, k, d), k, lst, m[0], elig, s["injury"]))
    rows.sort(key=lambda r: (r[0], r[1], NM.ascii_fold(L.players[r[2]].name.split(" ", 1)[1]), L.players[r[2]].pid))
    return rows


def report_name(name):
    nm = NM.ascii_fold(name)
    sfx = ""
    if nm.endswith(" Jr."):
        nm, sfx = nm[:-4], " Jr."
    first, last = nm.split(" ", 1)
    return "%s%s, %s" % (last.upper(), sfx.upper(), first.upper())


def write_il_report(L, out):
    rows = []
    for d, club, k, lst, placed, elig, inj in il_report(L):
        rows.append([K.iso(d), club, report_name(L.players[k].name), pos_of(L, k), lst, K.us(placed), K.us(elig), inj])
    K.write_csv(os.path.join(out, F_ILREP), ["report_date", "club", "player", "pos", "injured_list", "placed",
                                             "eligible", "injury"], rows)
    return len(rows)


# ---------------------------------------------------------------------------------------------- vendor files
def write_depth_charts(L, out):
    hv = vendor_ids(L)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    n = 0
    for y in K.HIST:
        ws = wb.create_sheet("DC %d" % y)
        ws.append(["%s: projected depth charts, snapshot %s" % (K.VENDOR, K.WINTER_SNAPSHOT[y].strftime("%B %-d, %Y"))])
        ws.append(["Hitters. Lineup = projected starting nine and batting order; Bench = reserves on the 26-man "
                   "projection. Age as of July 1."])
        ws.append([])
        ws.append(["Tm", "Role", "Ord", "Pos", "Player", "HV_ID", "LG_ID", "Bats", "Age"])
        for c in sorted(WD.CODES, key=lambda q: VENDOR_CLUB.get(q, q)):
            R = L.rosters[(y, c)]
            for pos, k in sorted(R["starters"].items(), key=lambda kv: R["slots"][kv[0]]):
                p = L.players[k]
                ws.append([VENDOR_CLUB.get(c, c), "Lineup", R["slots"][pos], pos, p.name, hv[k], p.pid, p.bats, p.age(y)])
                n += 1
            for k in R["bench"]:
                p = L.players[k]
                ws.append([VENDOR_CLUB.get(c, c), "Bench", None, p.pos if p.pos != "DH" else "1B", p.name, hv[k], p.pid,
                           p.bats, p.age(y)])
                n += 1
        ws.column_dimensions["E"].width = 24
        for col in "ABCDFGHI":
            ws.column_dimensions[col].width = 9
        ws["A1"].font = Font(bold=True)
    K.save_workbook(wb, os.path.join(out, F_DC), dt.datetime(2026, 2, 2, 9, 30), K.VENDOR, "Projected depth charts")
    return n, hv


def season_pa(L):
    pa = defaultdict(int)
    for y in K.HIST:
        for ln in L.lines[y]:
            pa[(y, ln["key"])] += ln["pa"]
    return pa


def projections(L, hv):
    """The vendor's preseason hitter projections (2022 to 2027): playing time from the player's last three seasons
    (weights 5, 4, 3) taken 80 per cent of the way from a baseline for his projected role, an age curve, and rate stats
    from the same seasons."""
    pa = season_pa(L)
    out = []
    for y in range(2022, 2028):
        cohort = set()
        for (yy, k), v in pa.items():
            if yy in (y - 1, y - 2, y - 3) and v >= 60:
                cohort.add(k)
        role = {}
        if y <= 2026:
            for c in WD.CODES:
                for k in L.rosters[(y, c)]["starters"].values():
                    role[k] = "reg"
        else:
            for c in WD.CODES:
                for k in L.rosters_end[(2026, c)]["starters"].values():
                    role[k] = "reg"
            for k in K.SIGN_ORDER:
                role[k] = "reg"
        for k in sorted(cohort, key=lambda q: L.players[q].pid):
            p = L.players[k]
            if p.pos == "P":
                continue
            h = [pa.get((y - j, k), 0) for j in (1, 2, 3)]
            w = [5, 4, 3]
            w = [w[j] if y - 1 - j >= 2021 else 0 for j in range(3)]
            base = sum(a * b for a, b in zip(w, h)) / sum(w)
            target = 600 if role.get(k) == "reg" else 260
            age = p.age(y)
            af = 1.0 - 0.012 * max(0, age - 31)
            r = K.rng("proj", y, k)
            proj = (0.8 * base + 0.2 * target) * af + r.normal(0, 6)
            proj = int(round(max(proj, 40)))
            q = p.rates
            bb = int(round(proj * q["bb"] * 0.97))
            hbp = int(round(proj * q["hbp"]))
            ab = proj - bb - hbp - int(round(proj * 0.009))
            hits = int(round(ab * (q["avg"] * 0.97 + 0.255 * 0.03)))
            hr = int(round(hits * q["hr"]))
            so = int(round((ab - hits) * q["k"]))
            out.append(dict(season=y, k=k, pa=proj, ab=ab, h=hits, hr=hr, bb=bb, so=so, hbp=hbp, age=age))
    return out


def write_projections(L, out, hv):
    rows = []
    P = projections(L, hv)
    for d in P:
        p = L.players[d["k"]]
        team = ""
        y = d["season"]
        if y <= 2026:
            team = L.club_of.get((y, d["k"]), "")
        else:
            for c in WD.CODES:
                R = L.rosters_end[(2026, c)]
                if d["k"] in R["starters"].values() or d["k"] in R["bench"]:
                    team = c
            if d["k"] in K.SIGN_ORDER:
                team = K.CLUB if K.SIGNINGS[d["k"]]["signed"] <= dt.date(2026, 12, 1) else "FA"
        team = team or "FA"
        avg = d["h"] / d["ab"] if d["ab"] else 0
        obp = (d["h"] + d["bb"] + d["hbp"]) / d["pa"] if d["pa"] else 0
        slg = (d["h"] + d["hr"] * 2.1 + d["h"] * 0.22) / d["ab"] if d["ab"] else 0
        pub = dt.date(y, 2, 10) if y <= 2026 else dt.date(2026, 12, 1)
        rows.append([y, K.iso(pub), hv[d["k"]], p.pid, p.name, VENDOR_CLUB.get(team, team), p.pos, d["age"], d["pa"],
                     d["ab"], d["h"], d["hr"], d["bb"], d["so"], "%.3f" % avg, "%.3f" % obp, "%.3f" % slg])
    K.write_csv(os.path.join(out, F_PROJ), ["season", "published", "hv_id", "lg_id", "player", "team", "pos", "age",
                                            "pa", "ab", "h", "hr", "bb", "so", "avg", "obp", "slg"], rows)
    return len(rows)


# ------------------------------------------------------------------------------------------------ the club's files
def pitchers_of_club(L, y):
    """The club's pitching staff for a season (names for the depth chart, the register and the payroll)."""
    key = ("ral_pitchers", y)
    if key in L.design:
        return L.design[key]
    prev = L.design.get(("ral_pitchers", y - 1), [])
    r = K.rng("ral_pitchers", y)
    staff = [k for k in prev if r.random() < 0.62]
    n = 0
    while len(staff) < 13:
        age = int(np.clip(round(r.normal(28.5, 3.3)), 22, 38))
        n += 1
        p = WD.new_player(L, r, "P", y - age, key="qRAL%d%02d" % (y, n))
        staff.append(p.key)
    L.design[key] = staff
    return staff


def ral_roster(L, y):
    R = L.rosters[(y, K.CLUB)]
    hitters = list(R["starters"].values()) + list(R["bench"])
    return hitters, pitchers_of_club(L, y)


def write_depth_chart_2027(L, out):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2027"
    ws.append(["Raleigh Oaks Baseball Operations: 2027 depth chart"])
    ws.append(["Updated December 10, 2026 by %s. Lineup, batting order and roles as planned for the 2027 "
               "season." % K.PEOPLE["agm"]])
    ws.append([])
    ws.append(["Lineup"])
    ws.append(["Ord", "Pos", "Player", "Bats", "Acquired"])
    for slot, pos, k in L.plan["lineup"]:
        p = L.players[k]
        acq = "Free agent, signed %s" % K.us(K.SIGNINGS[k]["signed"]) if k in K.SIGNINGS else ""
        ws.append([slot, pos, p.name, p.bats, acq])
    ws.append([])
    ws.append(["Bench"])
    ws.append(["Pos", "Player", "Bats", "Note"])
    for k in L.plan["bench"]:
        p = L.players[k]
        note = "1B/DH; under contract 2027" if k == "INC" else ""
        ws.append([p.pos if p.pos != "DH" else "1B", p.name, p.bats, note])
    ws.append([])
    staff = pitchers_of_club(L, 2027)
    ws.append(["Rotation"])
    for i, k in enumerate(staff[:5]):
        ws.append([i + 1, "SP", L.players[k].name])
    ws.append([])
    ws.append(["Bullpen"])
    roles = ["CL", "SU", "SU", "MR", "MR", "MR", "LR", "MR"]
    for i, k in enumerate(staff[5:13]):
        ws.append([roles[i], "RP", L.players[k].name])
    ws["A1"].font = Font(bold=True)
    for col, w in (("A", 8), ("B", 8), ("C", 24), ("D", 6), ("E", 30)):
        ws.column_dimensions[col].width = w
    K.save_workbook(wb, os.path.join(out, F_DC27), dt.datetime(2026, 12, 10, 17, 12), K.PEOPLE["agm"],
                    "2027 depth chart")
    return len(L.plan["lineup"]) + len(L.plan["bench"])


# ---------------------------------------------------------------------------------------- contracts, payroll, HRIS
# The club's veterans' contracts (2023 to 2027) with plate-appearance bonuses and award terms, and the club's pitchers'
# contracts with their other playing-time bonuses (FIN-14 4.1's innings pitched, games started, relief appearances and
# days on the Active List) and award terms. Keys are the generator's; every amount is the contract's, and the lines a
# pitcher reached each season are the season as it happened (nothing keyed ships; the payroll and the register do).
# 'deferred' is the salary a contract defers to its end (FIN-14 3.2), by the season it is earned: the payroll pays that
# season's salary less it, and the season-end run after the contract's last season pays it, one line a code a run.
VET_CONTRACTS = {
    "V1": dict(signed=dt.date(2022, 12, 14), seasons=(2023, 2024), salary=9_500_000, signing_bonus=1_000_000,
               pa=(140_000, (500, 550, 600, 650)), awards=[("Silver Slugger Award", 75_000)],
               deferred={2023: 350_000, 2024: 350_000}),
    "V3": dict(signed=dt.date(2022, 11, 30), seasons=(2023, 2024), salary=6_000_000, signing_bonus=0,
               pa=(100_000, (450, 500, 550)), awards=[("All-Star selection", 50_000)],
               deferred={2023: 200_000, 2024: 200_000}),
    "V2": dict(signed=dt.date(2023, 12, 6), seasons=(2024, 2025, 2026), salary=8_000_000, signing_bonus=2_000_000,
               pa=(125_000, (500, 550, 600)), awards=[("All-Star selection", 50_000)]),
    "V4": dict(signed=dt.date(2024, 12, 12), seasons=(2025,), salary=7_250_000, signing_bonus=500_000,
               pa=(225_000, (375, 475, 575)), awards=[]),
    "INC": dict(signed=dt.date(2024, 12, 19), seasons=(2025, 2026, 2027), salary=5_500_000, signing_bonus=1_500_000,
                pa=(100_000, (450, 525, 600)), awards=[("All-Star selection", 50_000)]),
    "V5": dict(signed=dt.date(2026, 1, 13), seasons=(2026,), salary=6_500_000, signing_bonus=750_000,
               pa=(175_000, (400, 475, 550)), awards=[]),
}
IP, GS, RA, DAL = "innings pitched", "games started", "relief appearances", "days on the Active List"
CY5 = "finishing in the top five of the Cy Young Award voting"
# per pitcher (staff key), his contracts: seasons, salary, terms (measure, amount per line, lines), award terms, and the
# number of lines of each measure he reached each season (lowest lines first)
PITCHER_CONTRACTS = {
    "qRAL202301": [dict(signed=dt.date(2021, 12, 2), seasons=(2022, 2023), salary={2022: 11_500_000, 2023: 12_000_000},
                        terms=[(IP, 70_000, (150, 170, 190))], awards=[], deferred={2023: 430_000},
                        reached={2022: {IP: 1}, 2023: {IP: 2}}),
                   dict(signed=dt.date(2023, 12, 12), seasons=(2024, 2025, 2026), salary=14_500_000,
                        terms=[(IP, 100_000, (150, 165, 180)), (GS, 50_000, (25, 30))], awards=[],
                        reached={2024: {IP: 2, GS: 2}, 2025: {IP: 2, GS: 0}, 2026: {IP: 1, GS: 2}})],
    "qRAL202304": [dict(signed=dt.date(2022, 11, 21), seasons=(2023, 2024, 2025), salary=8_250_000,
                        terms=[(IP, 60_000, (160, 180)), (GS, 60_000, (28, 32))], awards=[(CY5, 75_000)],
                        deferred={2023: 230_000, 2024: 230_000, 2025: 230_000},
                        reached={2023: {IP: 2, GS: 0}, 2024: {IP: 2, GS: 0}, 2025: {IP: 2, GS: 1}})],
    "qRAL202303": [dict(signed=dt.date(2023, 11, 26), seasons=(2024, 2025), salary=6_750_000,
                        terms=[(IP, 85_000, (150, 175)), (GS, 85_000, (28, 32))], awards=[],
                        deferred={2024: 310_000, 2025: 310_000},
                        reached={2024: {IP: 2, GS: 0}, 2025: {IP: 2, GS: 1}})],
    "qRAL202312": [dict(signed=dt.date(2023, 11, 3), seasons=(2024, 2025, 2026), salary=1_600_000,
                        terms=[(RA, 50_000, (50, 60, 70))], awards=[("All-Star selection", 40_000)],
                        reached={2024: {RA: 2}, 2025: {RA: 2}, 2026: {RA: 1}})],
    "qRAL202309": [dict(signed=dt.date(2023, 11, 19), seasons=(2024, 2025), salary=3_100_000,
                        terms=[(RA, 45_000, (50, 60))], awards=[],
                        reached={2024: {RA: 2}, 2025: {RA: 2}})],
    "qRAL202310": [dict(signed=dt.date(2022, 12, 16), seasons=(2023, 2024), salary=4_200_000,
                        terms=[(DAL, 75_000, (120, 150)), (RA, 75_000, (45, 55))], awards=[("All-Star selection", 50_000)],
                        deferred={2023: 175_000, 2024: 175_000},
                        reached={2023: {DAL: 2, RA: 0}, 2024: {DAL: 2, RA: 1}})],
    "qRAL202313": [dict(signed=dt.date(2022, 12, 5), seasons=(2023,), salary=1_150_000,
                        terms=[(RA, 30_000, (40, 50, 60))], awards=[],
                        reached={2023: {RA: 2}})],
    "qRAL202305": [dict(signed=dt.date(2022, 12, 20), seasons=(2023,), salary=5_500_000,
                        terms=[(IP, 55_000, (140, 160))], awards=[], deferred={2023: 255_000},
                        reached={2023: {IP: 1}})],
    "qRAL202404": [dict(signed=dt.date(2024, 11, 15), seasons=(2025, 2026, 2027), salary=9_000_000,
                        terms=[(GS, 70_000, (28, 32))], awards=[],
                        reached={2025: {GS: 2}, 2026: {GS: 1}})],
    "qRAL202504": [dict(signed=dt.date(2024, 12, 3), seasons=(2025, 2026), salary=2_400_000,
                        terms=[(RA, 60_000, (55, 65))], awards=[("All-Star selection", 50_000)],
                        reached={2025: {RA: 2}, 2026: {RA: 1}})],
}
AWARDS_WON = {("V1", 2024): "Silver Slugger Award", ("V3", 2024): "All-Star selection",
              ("INC", 2026): "All-Star selection", ("qRAL202310", 2023): "All-Star selection",
              ("qRAL202312", 2024): "All-Star selection", ("qRAL202304", 2024): CY5,
              ("qRAL202504", 2025): "All-Star selection"}
# the legacy payroll reissued a departed player's employee number to a later signing (the departed pitcher, the later)
REISSUED = ("qRAL202305", "qRAL202504")
VOID_REISSUE = ("qRAL202303", dt.date(2025, 10, 15))    # a season-end payment voided (bank details) and reissued


def award_terms(L, k):
    if k in VET_CONTRACTS:
        return dict(VET_CONTRACTS[k]["awards"])
    out = {}
    for c in PITCHER_CONTRACTS.get(k, []):
        out.update(dict(c["awards"]))
    return out


def pa_by_club(L):
    out = defaultdict(int)
    for y in K.HIST:
        for ln in L.lines[y]:
            out[(y, ln["key"], ln["club"])] += ln["pa"]
    return out


def pa_reached_with(L, k, y, T):
    """The club the player was with on the game in which his season's plate appearances reached T (None if never)."""
    cum = 0
    for ln in sorted((l for l in L.lines[y] if l["key"] == k), key=lambda l: (l["date"], l["gnum"], l["pk"])):
        cum += ln["pa"]
        if cum >= T:
            return ln["club"], ln["date"]
    return None, None


def vet_bonuses(L):
    """Every plate-appearance bonus line each veteran reached, the club holding the contract when he reached it, and the
    line's amount; the club pays only the lines reached while he was its player."""
    out = []
    for k, c in VET_CONTRACTS.items():
        amt, lines = c["pa"]
        for y in c["seasons"]:
            if y > 2026:
                continue
            for T in lines:
                club, d = pa_reached_with(L, k, y, T)
                if club is not None:
                    out.append(dict(k=k, season=y, T=T, amount=amt, club=club, date=d, kind="PA"))
    return out


def pitcher_bonuses():
    """(pitcher, season) -> the other playing-time bonuses he earned for the season (the truth the payroll pays)."""
    out = {}
    for k, cs in PITCHER_CONTRACTS.items():
        for c in cs:
            for y, rch in c["reached"].items():
                out[(k, y)] = sum(a * rch.get(m, 0) for m, a, lines in c["terms"])
    return out


def contract_salary(k, y):
    for c in PITCHER_CONTRACTS.get(k, []):
        if y in c["seasons"]:
            s = c["salary"]
            return s[y] if isinstance(s, dict) else s
    return None


def salary_of(L, k, y, r):
    if k in VET_CONTRACTS and y in VET_CONTRACTS[k]["seasons"]:
        return VET_CONTRACTS[k]["salary"]
    explicit = contract_salary(k, y)
    if explicit is not None:
        return explicit
    core = {"RC": (2_150_000, 3_400_000, 4_900_000, 6_100_000), "RSS": (1_100_000, 3_250_000, 5_600_000, 7_400_000),
            "RCF": (745_000, 2_600_000, 4_350_000, 5_800_000), "RRF": (4_800_000, 6_900_000, 8_200_000, 9_100_000)}
    if k in core:
        return core[k][y - 2023]
    minimum = {2023: 720_000, 2024: 740_000, 2025: 760_000, 2026: 780_000}[y]
    u = r.random()
    if u < 0.45:
        return minimum + int(r.integers(0, 6)) * 5_000
    if u < 0.8:
        return int(round(r.uniform(1.0, 4.5), 2) * 1_000_000 / 25_000) * 25_000
    return int(round(r.uniform(4.5, 14.0), 2) * 1_000_000 / 50_000) * 50_000


def deferred_of(k, y):
    """The salary the contract in force in season y defers to its end (FIN-14 3.2), out of that season's salary."""
    if k in VET_CONTRACTS and y in VET_CONTRACTS[k]["seasons"]:
        return VET_CONTRACTS[k].get("deferred", {}).get(y, 0)
    for c in PITCHER_CONTRACTS.get(k, []):
        if y in c["seasons"]:
            return c.get("deferred", {}).get(y, 0)
    return 0


def deferral_due(k, y):
    """The deferred salary the season-end run after season y pays the player: the whole deferral of every contract of
    his whose last season is y."""
    cs = ([VET_CONTRACTS[k]] if k in VET_CONTRACTS else []) + PITCHER_CONTRACTS.get(k, [])
    return sum(sum(c.get("deferred", {}).values()) for c in cs if c["seasons"][-1] == y)


def last_season(k, y):
    """True when y is the last season of the pitcher's contract in force in y."""
    for c in PITCHER_CONTRACTS.get(k, []):
        if y in c["seasons"]:
            return c["seasons"][-1] == y
    return False


def paydays(y):
    """Semi-monthly championship-season paydays, April 15 to September 30."""
    out = []
    for m in range(4, 10):
        out.append(dt.date(y, m, 15))
        nxt = dt.date(y, m + 1, 1) - D(days=1)
        out.append(nxt)
    return out


def club_people(L):
    """Everyone the club paid from January 2023: the season-start rosters (hitters and pitchers) and the veterans."""
    for y in (2023, 2024, 2025, 2026, 2027):
        pitchers_of_club(L, y)
    seen = {}
    for y in (2023, 2024, 2025, 2026):
        hitters, staff = ral_roster(L, y)
        for k in hitters + staff:
            seen.setdefault(k, y)
    for k in VET_CONTRACTS:
        seen.setdefault(k, VET_CONTRACTS[k]["seasons"][0])
    return seen


def employee_numbers(L, people):
    """Legacy payroll numbers and HRIS IDs. The legacy system reissued a departed pitcher's number to a later signing,
    so one legacy number carries two players' lines."""
    r = K.rng("hris")
    legacy, new = {}, {}
    nums = r.choice(np.arange(1200, 4800), len(people), replace=False)
    ids = r.choice(np.arange(100, 999), len(people), replace=False)
    for (k, first), a, b in zip(sorted(people.items(), key=lambda kv: (kv[1], L.players[kv[0]].pid)), nums, ids):
        legacy[k] = "%05d" % int(a)
        new[k] = "E10%03d%02d" % (int(b), int(r.integers(10, 99)))
    legacy[REISSUED[1]] = legacy[REISSUED[0]]
    return legacy, new


def payroll(L):
    """The payroll rows the HRIS export carries for players, January 2023 to December 11, 2026. Rows paid before the
    January 1, 2026 migration keep the legacy employee number, name format and earning code. Internal codes: SAL salary
    (the season's salary less what the contract defers to its end), SGN signing bonus, PA a plate-appearance bonus, PT a
    pitcher's playing-time bonus paid in the season-end run (after a contract's last season, the whole season's bonus
    with the salary the contract deferred), PT2 a second half paid in the January run, AW an award bonus, AS an All-Star
    bonus, DEF deferred salary paid alone."""
    people = club_people(L)
    legacy, new = employee_numbers(L, people)
    rows = []
    r = K.rng("payroll")

    def add(d, k, code, amount, status="PAID"):
        rows.append(dict(date=d, k=k, code=code, amount=amount, status=status))

    for y in (2023, 2024, 2025, 2026):
        hitters, staff = ral_roster(L, y)
        for k in hitters + staff:
            sal = salary_of(L, k, y, K.rng("salary", y, k)) - deferred_of(k, y)
            per = round(sal / 12, 2)
            for i, d in enumerate(paydays(y)):
                if y == 2025 and k == "V4" and d > dt.date(2025, 7, 28):
                    continue
                amt = per if i < 11 else round(sal - per * 11, 2)
                add(d, k, "SAL", amt)
    for k, c in VET_CONTRACTS.items():
        if c["signing_bonus"]:
            d = dt.date(c["signed"].year + 1, 1, 15) if c["signed"].month >= 11 else dt.date(c["signed"].year, 2, 13)
            add(d, k, "SGN", c["signing_bonus"])
    # playing-time bonuses through the 2025 season: the hitters' plate-appearance bonuses in full in the season-end run
    # (October 15), the pitchers' half in it and half in the January 15 run after it, except that the season-end run
    # after a contract's last season settles what the contract still owes (the season's bonuses in full and the salary
    # it deferred to its end); the 2026 season's are payable in full on January 15, 2027
    earned = defaultdict(int)
    for b in vet_bonuses(L):
        if b["club"] == K.CLUB:
            earned[(b["k"], b["season"])] += b["amount"]
    for (k, y), amt in sorted(earned.items()):
        if y <= 2025:
            add(dt.date(y, 10, 15), k, "PA", amt)
    star = {}
    for (k, y), award in sorted(AWARDS_WON.items()):
        amt = award_terms(L, k)[award]
        if y >= 2026:
            add(dt.date(y, 7, 31), k, "AS" if "All-Star" in award else "AW", amt)
        elif "All-Star" in award and k in PITCHER_CONTRACTS:
            star[(k, y)] = amt                 # the legacy season-end run, one line a code, with his playing-time half
        elif "All-Star" in award:
            add(dt.date(y, 10, 15), k, "AS", amt)
        else:
            add(dt.date(y, 11, 15), k, "AW", amt)
    # a veteran's deferred salary in the season-end run after his contract's last season, on the legacy catch-all, one
    # line a code a run (with an All-Star bonus paid in the same run)
    for k, c in sorted(VET_CONTRACTS.items()):
        due = deferral_due(k, c["seasons"][-1])
        if not due:
            continue
        d = dt.date(c["seasons"][-1], 10, 15)
        assert d < K.HRIS_LIVE
        same = [x for x in rows if x["k"] == k and x["date"] == d and LEGACY_CODE[x["code"]][0] == "E45"]
        if same:
            same[0]["amount"] += due
        else:
            add(d, k, "DEF", due)
    for (k, y), amt in sorted(pitcher_bonuses().items()):
        if y >= 2026:
            continue
        if last_season(k, y):
            owed = amt + star.pop((k, y), 0) + deferral_due(k, y)
            if owed and y >= 2023:
                add(dt.date(y, 10, 15), k, "PT", owed)
            continue
        if not amt:
            continue
        if y >= 2023:
            add(dt.date(y, 10, 15), k, "PT", amt / 2 + star.pop((k, y), 0))
        add(dt.date(y + 1, 1, 15), k, "PT2", amt / 2)
    assert not star
    # a January run that missed one pitcher's half: the legacy payroll settled it in his next season-end run, on the
    # catch-all line with that season's half (one line a code a run)
    for k, y in CATCH_UP:
        jan = [x for x in rows if x["k"] == k and x["code"] == "PT2" and x["date"] == dt.date(y + 1, 1, 15)]
        nxt = [x for x in rows if x["k"] == k and x["code"] == "PT" and x["date"] == dt.date(y + 1, 10, 15)]
        assert len(jan) == 1 and len(nxt) == 1, (k, y)
        nxt[0]["amount"] += jan[0]["amount"]
        rows.remove(jan[0])
    for x in list(rows):
        if x["k"] == VOID_REISSUE[0] and x["code"] == "PT" and x["date"] == VOID_REISSUE[1]:
            x["status"] = "VOIDED"
            add(VOID_REISSUE[1] + D(days=5), x["k"], "PT", -x["amount"], "VOID REVERSAL")
            add(VOID_REISSUE[1] + D(days=7), x["k"], "PT", x["amount"])
    rows.sort(key=lambda x: (x["date"], x["code"], L.players[x["k"]].pid, x["amount"]))
    pid = 400000
    for x in rows:
        pid += int(r.integers(1, 9))
        x["pay_id"] = pid
    migrated = {x["k"] for x in rows if x["date"] >= dt.date(2025, 1, 1)}
    return rows, legacy, new, people, migrated


# (pitcher, season) whose January half the legacy run missed and settled in his next season-end run
CATCH_UP = (("qRAL202304", 2023), ("qRAL202312", 2024))

# the contract terms the register states for the pitchers' other playing-time bonuses, the schedule payroll() pays to
# (the two January runs in CATCH_UP excepted, each half paid late)
PAY_NOTE = ("Other playing-time bonuses for the 2025 season and earlier were payable one half in the October 15 "
            "payroll run after the season and one half in the January 15 run after that, or in full in the October 15 "
            "run after a contract's last season.")

LEGACY_CODE = {"SAL": ("E01", "SALARY-PLAYER"), "SGN": ("E10", "SIGNING BONUS"), "PA": ("E40", "PERF BONUS"),
               "PT": ("E45", "OTHER BONUS"), "PT2": ("E40", "PERF BONUS"), "AW": ("E45", "OTHER BONUS"),
               "AS": ("E45", "OTHER BONUS"), "DEF": ("E45", "OTHER BONUS")}
NEW_CODE = {"SAL": ("SAL", "Player salary"), "SGN": ("SGN", "Signing bonus"), "PA": ("BPA", "Playing-time bonus"),
            "PT": ("BPA", "Playing-time bonus"), "PT2": ("BPA", "Playing-time bonus"), "AW": ("BAW", "Award bonus"),
            "AS": ("BAS", "All-Star bonus"), "DEF": ("BOT", "Other bonus")}


def legacy_name(name):
    nm = NM.ascii_fold(name).upper()
    first, last = nm.split(" ", 1)
    return "%s %s" % (last, first)


def hris_name(name):
    first, last = name.split(" ", 1)
    return "%s, %s" % (last, first)


def write_payroll(L, out):
    rows, legacy, new, people, migrated = payroll(L)
    w = []
    for x in rows:
        p = L.players[x["k"]]
        pre = x["date"] < K.HRIS_LIVE
        code, desc = (LEGACY_CODE if pre else NEW_CODE)[x["code"]]
        w.append([K.iso(x["date"]), legacy[x["k"]] if pre else new[x["k"]], legacy_name(p.name) if pre else hris_name(p.name),
                  code, desc, "%.2f" % x["amount"], x["pay_id"], x["status"]])
    K.write_csv(os.path.join(out, F_PAY), ["pay_date", "employee_id", "employee_name", "earning_code",
                                           "earning_description", "amount", "payment_id", "status"], w)
    return len(w), legacy, new, people, migrated


def write_hris_map(L, out, legacy, new, migrated):
    rows = []
    for k in sorted(migrated, key=lambda q: legacy[q]):
        p = L.players[k]
        rows.append([legacy[k], new[k], hris_name(p.name), K.iso(K.HRIS_LIVE)])
    K.write_csv(os.path.join(out, F_HRIS), ["legacy_employee_no", "employee_id", "employee_name", "migrated"], rows)
    return len(rows)


def write_earning_codes(out):
    rows = [["E01", "SALARY-PLAYER", "legacy payroll", "2009-01-01", "2025-12-31", "SAL"],
            ["E10", "SIGNING BONUS", "legacy payroll", "2009-01-01", "2025-12-31", "SGN"],
            ["E40", "PERF BONUS", "legacy payroll", "2009-01-01", "2025-12-31", "BPA"],
            ["E45", "OTHER BONUS", "legacy payroll", "2009-01-01", "2025-12-31", "BOT"],
            ["E55", "MOVING ALLOW", "legacy payroll", "2011-03-01", "2025-12-31", "MOV"],
            ["E70", "POSTSEASON SHARE", "legacy payroll", "2009-01-01", "2025-12-31", "PSS"],
            ["SAL", "Player salary", "HRIS", "2026-01-01", "", ""],
            ["SGN", "Signing bonus", "HRIS", "2026-01-01", "", ""],
            ["BPA", "Playing-time bonus", "HRIS", "2026-01-01", "", ""],
            ["BAW", "Award bonus", "HRIS", "2026-01-01", "", ""],
            ["BAS", "All-Star bonus", "HRIS", "2026-01-01", "", ""],
            ["BOT", "Other bonus", "HRIS", "2026-01-01", "", ""],
            ["MOV", "Relocation allowance", "HRIS", "2026-01-01", "", ""],
            ["PSS", "Postseason share", "HRIS", "2026-01-01", "", ""]]
    K.write_csv(os.path.join(out, F_CODES), ["earning_code", "description", "system", "valid_from", "valid_to",
                                             "migrated_as"], rows)
    return len(rows)


def money_terms(amt, lines):
    if len(lines) == 1:
        return "$%s for %d" % ("{:,}".format(amt), lines[0])
    return "$%s each for %s and %d" % ("{:,}".format(amt), ", ".join(str(t) for t in lines[:-1]), lines[-1])


def other_terms_text(terms):
    bits = []
    for i, (m, a, lines) in enumerate(terms):
        lab = m[0].upper() + m[1:] if i == 0 else m
        bits.append("%s: %s" % (lab, money_terms(a, lines)))
    return "; ".join(bits)


def award_text(awards):
    return "; ".join("$%s for %s" % ("{:,}".format(a), nm) for nm, a in awards)


def register_rows(L, legacy, new, people, migrated):
    """The club's player contract register, one row per contract in force in any of the 2023 to 2026 seasons. A
    player's Employee ID is his HRIS ID, or his legacy payroll number if the club has not paid him since 2024."""
    rows = []
    staff27 = set(pitchers_of_club(L, 2027))
    plan27 = {k for _, _, k in L.plan["lineup"]} | set(L.plan["bench"])
    for k in sorted(people, key=lambda q: hris_name(L.players[q].name)):
        p = L.players[k]
        on = [y for y in (2023, 2024, 2025, 2026) if k in ral_roster(L, y)[0] + ral_roster(L, y)[1]]
        contracts = []
        if k in VET_CONTRACTS:
            c = VET_CONTRACTS[k]
            seasons = list(c["seasons"])
            amt, lines = c["pa"]
            status = "Active" if seasons[-1] >= 2027 else "Expired"
            if k == "V4":
                status = "Assigned to New Orleans Levees 07/28/2025"
            contracts.append(dict(signed=c["signed"], seasons=seasons, sal={y: c["salary"] for y in seasons},
                                  sb=c["signing_bonus"], pa_terms=money_terms(amt, lines) + " plate appearances",
                                  other="", awards=award_text(c["awards"]), status=status))
            on = [y for y in on if y not in seasons]
        for c in PITCHER_CONTRACTS.get(k, []):
            seasons = list(c["seasons"])
            contracts.append(dict(signed=c["signed"], seasons=seasons,
                                  sal={y: contract_salary(k, y) for y in seasons}, sb=0, pa_terms="",
                                  other=other_terms_text(c["terms"]), awards=award_text(c["awards"]),
                                  status="Active" if seasons[-1] >= 2027 else "Expired"))
            on = [y for y in on if y not in seasons]
        if on:
            rr = K.rng("signed", k, on[0])
            signed = dt.date(on[0] - 1, int(rr.integers(10, 13)), int(rr.integers(1, 29)))
            seasons = list(on)
            if 2026 in on and (k in staff27 or k in plan27):
                seasons.append(2027)
            contracts.append(dict(signed=signed, seasons=seasons, sb=0, pa_terms="", other="", awards="",
                                  sal={y: salary_of(L, k, y, K.rng("salary", y, k)) for y in on},
                                  status="Active" if seasons[-1] >= 2027 else "Expired"))
        for c in contracts:
            s = c["seasons"]
            rows.append(dict(k=k, name=hris_name(p.name), emp=new[k] if k in migrated else legacy[k],
                             pos=("P" if p.pos == "P" else p.pos), signed=c["signed"],
                             seasons="%d-%d" % (s[0], s[-1]) if len(s) > 1 else str(s[0]), sal=c["sal"], sb=c["sb"],
                             pa_terms=c["pa_terms"], other=c["other"], awards=c["awards"], status=c["status"]))
    return rows


def write_register(L, out, legacy, new, people, migrated):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Register"
    ws.append(["Raleigh Oaks: player contract register, 2023 to 2026 seasons (extract)"])
    ws.append(["Maintained by Baseball Operations and Finance; payroll contact %s. Amounts in US dollars. Employee ID: "
               "HRIS ID, or the legacy payroll number for a player not paid since 2024. %s" % (K.PEOPLE["payroll"],
                                                                                              PAY_NOTE)])
    ws.append([])
    hdr = ["Player", "Employee ID", "Pos", "Signed", "Seasons", "Salary 2023", "Salary 2024", "Salary 2025",
           "Salary 2026", "Signing bonus", "Plate-appearance bonuses (each season)",
           "Other playing-time bonuses (each season)", "Award bonuses", "Status"]
    ws.append(hdr)
    rr = register_rows(L, legacy, new, people, migrated)
    for x in rr:
        ws.append([x["name"], x["emp"], x["pos"], K.us(x["signed"]) if x["signed"] else "", x["seasons"],
                   x["sal"].get(2023), x["sal"].get(2024), x["sal"].get(2025), x["sal"].get(2026), x["sb"] or None,
                   x["pa_terms"] or None, x["other"] or None, x["awards"] or None, x["status"]])
    for i, w in enumerate([24, 12, 5, 11, 10, 12, 12, 12, 12, 12, 40, 60, 44, 34]):
        ws.column_dimensions[get_column_letter(i + 1)].width = w
    for row in ws.iter_rows(min_row=5, min_col=6, max_col=10):
        for c in row:
            c.number_format = "#,##0"
    ws["A1"].font = Font(bold=True)
    for c in ws[4]:
        c.font = Font(bold=True)
    K.save_workbook(wb, os.path.join(out, F_REG), dt.datetime(2026, 12, 8, 11, 25), K.PEOPLE["payroll"],
                    "Player contract register")
    return len(rr)
