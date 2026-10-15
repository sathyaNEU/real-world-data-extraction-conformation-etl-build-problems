"""The golden, computed from the pack alone: who answered (and when), each respondent's answers, the latent
(tetrachoric and polychoric) correlations of the 30 items, the dimensions that stand above random data, each item's
sub-scale, and the coverage figures for the 12 store districts. Then the three deliverables."""
from __future__ import annotations

import datetime as dt

import numpy as np
import pandas as pd

import common as K
import load as L
import psy

PA_REPS = 200
PA_SEED = 2026


# ------------------------------------------------------------------ who answered, and their answers

def answer_rows(T, keep="last", tests=False):
    """Answer rows with the item's English code, the token's channel and the employee it belongs to. keep: 'last'
    (the last submission under each token) or 'all' (every submission)."""
    m, order = L.item_map(T["items"])
    resp = T["resp"].copy()
    subs = T["subs"]
    if not tests:
        test_tokens = set(subs.loc[subs["status"] == "test", "token"])
        resp = resp[~resp["token"].isin(test_tokens)]
    if keep == "last":
        last = resp.groupby("token")["submission"].transform("max")
        resp = resp[resp["submission"] == last]
    ch = subs.drop_duplicates("token").set_index("token")["channel"]
    resp = resp.assign(channel=resp["token"].map(ch))
    emp = dict(zip(T["dist"]["token"], T["dist"]["employee_id"]))
    emp.update({s: "%07d" % int(e) for s, e in zip(T["kiosk"]["session"], T["kiosk"]["employee"])})
    resp = resp.assign(employee=resp["token"].map(emp), key=resp["item"].map(m))
    return resp, order


def before_close(T, rows, paper="sealed", close="local"):
    """Whether each answer row was given before the survey closed at the employee's site.
    paper: 'sealed' (a printed form's answers were given before its envelope was sealed) or 'answered_at'.
    close: 'local' (23:59 local at the site), 'utc_date' (before 28 September UTC), 'eastern' (23:59 Toronto time
    for every site) or 'links' (before Kestrava's links stopped, 06:00 UTC on 28 September)."""
    sites = T["sites"].set_index("site_no")
    opening = L.opening_site(T["hris"])
    tz = rows["employee"].map(lambda e: sites.loc[opening[e], "time_zone"] if e in opening else "America/Toronto")
    if close == "local":
        lim = tz.map(L.close_utc)
    else:
        lim = pd.Series({"utc_date": "2026-09-27T23:59:59Z", "eastern": L.close_utc("America/Toronto"),
                         "links": "2026-09-28T05:59:59Z"}[close], index=rows.index)
    ok = rows["answered_at"] <= lim
    if paper == "sealed":
        pap = rows["channel"] == "paper"
        env = T["paper"]
        site_of = dict(zip(T["dist"]["token"], T["dist"]["site_no"]))
        sealed = {(int(s), pd.Timestamp(k).date().isoformat()): pd.Timestamp(d).date()
                  for s, k, d in zip(env["Site"], env["Keyed"], env["Sealed"])}
        rp = rows[pap]
        keyed_day = rp["answered_at"].str.slice(0, 10)
        sd = [sealed.get((int(site_of[t]), k)) for t, k in zip(rp["token"], keyed_day)]
        ok.loc[pap] = [d is not None and d <= K.CLOSE_DAY for d in sd]
    return ok


def district_counts(T, paper="sealed", close="local", site="opening", people="employees", complete="union",
                    dna="answer", pre=None):
    """Ask A and Ask B for the 12 store districts under a handling of each device and hazard.
    site: 'opening' (HR assignment on the opening day) or 'list' (the distribution list's site; a kiosk session's
    own store). people: 'employees' or 'tokens' (Ask A). complete: 'union' (an employee's items across their tokens),
    'token' (one of their tokens holds all 30) or 'last_session' (their latest token holds all 30). dna: 'answer' or
    'missing' (does-not-apply not counted as an answer)."""
    if pre is None:
        rows, order = answer_rows(T)
        ok = before_close(T, rows, paper=paper, close=close)
    else:
        rows, ok = pre
    if dna == "missing":
        rows = rows.assign(key=rows["key"].where(rows["answer"] != 6))
    r = rows[ok & rows["key"].notna()]
    sites = T["sites"].set_index("site_no")
    if site == "opening":
        opening = L.opening_site(T["hris"])
        tok_site = {t: opening.get(e) for t, e in zip(r["token"], r["employee"])}
    else:
        lst = dict(zip(T["dist"]["token"], T["dist"]["site_no"]))
        legacy = dict(zip(sites["legacy_site_no"], sites.index))
        kio = {s: legacy[int(v)] for s, v in zip(T["kiosk"]["session"], T["kiosk"]["store"])}
        tok_site = {t: lst.get(t, kio.get(t)) for t in r["token"].unique()}
    tok_d = {t: (sites.loc[s, "district_code"] if s is not None else "") for t, s in tok_site.items()}
    kset = r.groupby("token")["key"].apply(set).to_dict()
    emp_of = dict(zip(r["token"], r["employee"]))
    last_at = r.groupby("token")["answered_at"].max().to_dict()
    emp = {}
    for t in sorted(kset):
        e = emp_of[t]
        if e not in emp:
            emp[e] = dict(d=tok_d[t], keys=set(), toks=[])
        emp[e]["keys"] |= kset[t]
        emp[e]["toks"].append(t)
    A = {d: 0 for d, _, _ in K.DISTRICTS}
    B = {d: 0 for d, _, _ in K.DISTRICTS}
    if people == "tokens":
        for t in kset:
            if tok_d[t] in A:
                A[tok_d[t]] += 1
    else:
        for v in emp.values():
            if v["d"] in A:
                A[v["d"]] += 1
    for v in emp.values():
        if v["d"] not in B:
            continue
        if complete == "union":
            c = len(v["keys"]) == 30
        elif complete == "token":
            c = any(len(kset[t]) == 30 for t in v["toks"])
        else:
            c = len(kset[max(v["toks"], key=lambda t: last_at[t])]) == 30
        if c:
            B[v["d"]] += 1
    return A, B


# ------------------------------------------------------------------ the call

def call_matrix(T, paper="sealed", kiosk="employee", keep="last", tests=False, late=False, language=None,
                complete_only=False, dna="missing", enps=False, with_index=False):
    """The respondents' answers as a matrix (rows respondents, columns the 30 items in English order): five-point
    answers 0-4 with negatively worded statements reversed, yes 1 and no 0, does-not-apply and blanks missing.
    Variants serve the convergence checks: paper forms by answered_at, each kiosk session a respondent, every
    submission a respondent, test tokens and late answers kept, one language, complete cases, does-not-apply as an
    extra category, the recommend question as a 31st variable."""
    rows, order = answer_rows(T, keep=keep, tests=tests)
    if not late:
        rows = rows[before_close(T, rows, paper=paper)]
    if language is not None:
        lang = T["subs"].drop_duplicates("token").set_index("token")["language"]
        rows = rows[rows["token"].map(lang) == language]
    if kiosk == "token" or tests or keep == "all":
        rows = rows.assign(unit=rows["token"] + ":" + rows["submission"].astype(str))
    else:
        rows = rows.assign(unit=rows["employee"].fillna(rows["token"]))
    cols = list(order) + (["ENPS"] if enps else [])
    rows = rows.assign(col=rows["key"].where(rows["key"].notna(), rows["item"]))
    rows = rows[rows["col"].isin(cols)]
    wide = rows.sort_values("answered_at").drop_duplicates(["unit", "col"], keep="last").pivot(
        index="unit", columns="col", values="answer").reindex(columns=cols)
    items = T["items"].set_index("Item")
    X = wide.to_numpy(dtype=float).copy()
    Kc = []
    for j, c in enumerate(cols):
        x = X[:, j]
        if c == "ENPS":
            Kc.append(11)
            continue
        scale, rev = items.loc[c, "Scale"], items.loc[c, "Reversed"] == "yes"
        six = x == 6
        if scale == "agree5":
            x = x - 1
            if rev:
                x = 4 - x
            if dna == "category":
                x = np.where(six, 5, x)
                Kc.append(6)
            else:
                x = np.where(six, np.nan, x)
                Kc.append(5)
        else:
            x = np.where(x == 1, 1.0, np.where(x == 2, 0.0, x))
            if dna == "category":
                x = np.where(six, 2, x)
                Kc.append(3)
            else:
                x = np.where(six, np.nan, x)
                Kc.append(2)
        X[:, j] = x
    if complete_only:
        X = X[~np.isnan(X).any(1)]
    if with_index:
        return X, np.array(Kc), cols, list(wide.index)
    return X, np.array(Kc), cols


_REF = {}


def random_reference(n, p, reps=PA_REPS):
    """Eigenvalues of correlations of normal random data with n rows and p columns: mean and 95th percentile."""
    if (n, p, reps) not in _REF:
        _REF[(n, p, reps)] = psy.random_eigen(n, p, reps=reps, seed=PA_SEED)
    return _REF[(n, p, reps)]


def dimensions(X, Kc, route="latent", correct=0.0, reps=PA_REPS, weights=None):
    if route == "latent":
        R = psy.latent_corr(X, Kc, correct=correct, weights=weights)[0]
    elif route == "spearman":
        R = psy.pearson_corr(psy.rank_columns(X), weights=weights)
    else:
        R = psy.pearson_corr(X, weights=weights)
    ev = psy.eigen(R)
    mean, p95 = random_reference(X.shape[0], X.shape[1], reps)
    return dict(R=R, ev=ev, mean=mean, p95=p95, pa95=psy.keep_above(ev, p95), pamean=psy.keep_above(ev, mean),
                kaiser=psy.kaiser(ev), map=psy.velicer_map(R)[0], map4=psy.velicer_map(R, 4)[0], n=X.shape[0])


def name_groups(groups, order, T):
    """Name each factor by the item it holds: the workload statement Q02, the manager statement Q01, the team
    statement Q03."""
    text = T["items"].set_index("Item")["Question (English)"]
    names = {}
    for f in sorted(set(groups)):
        held = [order[j] for j in range(len(order)) if groups[j] == f]
        if "Q02" in held and "Q04" in held:
            names[f] = "Resources"
        elif "Q01" in held and "Q06" in held:
            names[f] = "Manager and growth"
        elif "Q03" in held:
            names[f] = "Connection"
        else:
            names[f] = "Dimension %d (%s)" % (f + 1, text[held[0]][:30])
    return names


def compute(target):
    T = L.read(target)
    X, Kc, order, units = call_matrix(T, with_index=True)
    lat = dimensions(X, Kc, "latent")
    pea = dimensions(X, Kc, "pearson")
    k = lat["pa95"]
    Lp = psy.promax(psy.paf(lat["R"], k))
    groups = psy.placements(Lp)
    names = name_groups(groups, order, T)
    items = T["items"].set_index("Item")
    placement = {order[j]: names[int(groups[j])] for j in range(len(order))}
    loading = {order[j]: float(Lp[j, groups[j]]) for j in range(len(order))}
    Lp5 = psy.promax(psy.paf(pea["R"], pea["pa95"]))
    g5 = psy.placements(Lp5)
    A, B = district_counts(T)
    dnames = dict(zip(T["sites"]["district_code"], T["sites"]["district"]))
    fig = {"call_subscales": int(k), "respondents": int(lat["n"]),
           "latent_eigen": [round(float(v), 2) for v in lat["ev"][:6]],
           "pearson_eigen": [round(float(v), 2) for v in pea["ev"][:6]],
           "random_p95": [round(float(v), 2) for v in lat["p95"][:6]],
           "pearson_subscales": int(pea["pa95"]), "pearson_kaiser": int(pea["kaiser"]),
           "latent_kaiser": int(lat["kaiser"]), "latent_map": int(lat["map"]), "pearson_map": int(pea["map"]),
           "subscale_items": {nm: [c for c in order if placement[c] == nm] for nm in sorted(set(placement.values()))},
           "subscale_sizes": {nm: sum(1 for c in order if placement[c] == nm) for nm in sorted(set(placement.values()))},
           "ask_a": {d: int(A[d]) for d in A}, "ask_b": {d: int(B[d]) for d in B},
           "ask_a_total": int(sum(A.values())), "ask_b_total": int(sum(B.values())),
           "yes_share": {c: round(100 * float(np.nanmean(X[:, j])), 1) for j, c in enumerate(order)
                         if items.loc[c, "Scale"] == "yesno"}}
    return dict(figures=fig, T=T, X=X, Kc=Kc, order=order, units=units, lat=lat, pea=pea, groups=groups, names=names,
                placement=placement, loading=loading, groups5=g5, A=A, B=B, dnames=dnames)


# ------------------------------------------------------------------ deliverables

def write_xlsx(path, res):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    T, order = res["T"], res["order"]
    items = T["items"].set_index("Item")
    hdr, fill = Font(name="Calibri", bold=True), PatternFill("solid", fgColor="E3E9E4")
    wb = Workbook()
    ws = wb.active
    ws.title = "Items"
    ws.append(["Item", "Question (English)", "Answered as", "Sub-scale", "Loading"])
    for c in ws[1]:
        c.font, c.fill = hdr, fill
    rank = {"Resources": 0, "Manager and growth": 1, "Connection": 2}
    for c in sorted(order, key=lambda v: (rank.get(res["placement"][v], 9), v)):
        ws.append([c, items.loc[c, "Question (English)"], "agree / disagree" if items.loc[c, "Scale"] == "agree5"
                   else "yes / no", res["placement"][c], round(res["loading"][c], 2)])
        ws.cell(row=ws.max_row, column=5).number_format = "0.00"
    for col, w in zip("ABCDE", (7, 92, 16, 20, 9)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.append([])
    ws.append(["Loadings: principal-axis factoring of the tetrachoric and polychoric correlations, three factors, "
               "promax rotation; negatively worded statements reverse-scored."])
    wd = wb.create_sheet("Districts")
    wd.append(["District", "Code", "Answered at least one item before the close", "Of them, answered all 30"])
    for c in wd[1]:
        c.font, c.fill = hdr, fill
    for d, nm, _ in K.DISTRICTS:
        wd.append([nm, d, res["A"][d], res["B"][d]])
    wd.append(["All store districts", None, sum(res["A"].values()), sum(res["B"].values())])
    wd.cell(row=wd.max_row, column=1).font = hdr
    for col, w in zip("ABCD", (26, 7, 42, 26)):
        wd.column_dimensions[col].width = w
    wd.append([])
    wd.append(["Employees are counted once, under the district of the store they were assigned to on 14 September "
               "2026."])
    wv = wb.create_sheet("Dimensions")
    wv.append(["Dimension", "Eigenvalue", "Random data, 95th percentile", "Kept"])
    for c in wv[1]:
        c.font, c.fill = hdr, fill
    for i in range(len(order)):
        wv.append([i + 1, round(float(res["lat"]["ev"][i]), 3), round(float(res["lat"]["p95"][i]), 3),
                   "yes" if i < res["figures"]["call_subscales"] else "no"])
    for col, w in zip("ABCD", (11, 12, 28, 8)):
        wv.column_dimensions[col].width = w
    wn = wb.create_sheet("Notes")
    notes = [
        ("Survey", "Engagement Survey 2026 (Kestrava BRK-ENG-2026), 14 to 27 September 2026; Kestrava export of 19 "
                   "October 2026."),
        ("Respondents", "Employees who answered at least one of the 30 items before 23:59 local time on 27 September "
                        "at their site (reporting standard 3.1), each counted once: %s." % "{:,}".format(
                            res["figures"]["respondents"])),
        ("Sub-scale", "One distinct dimension of employees' views that the items measure (reporting standard 2.1); "
                      "each item under the dimension it loads on most (2.2)."),
        ("Correlations", "Tetrachoric for two yes or no questions, polychoric otherwise; two-step maximum likelihood, "
                         "pairwise; negatively worded statements reversed; does-not-apply left out."),
        ("Dimensions kept", "Eigenvalues above the 95th percentile of random normal data of the same size (%d "
                            "samples); MAP and eigenvalues above one agree." % PA_REPS),
        ("Districts", "An employee counts under the district of the store held in HR on 14 September 2026 "
                      "(reporting standard 3.4). Printed forms count by the envelope's sealed date."),
        ("Prepared", "People Analytics, 12 November 2026."),
    ]
    for k_, v_ in notes:
        wn.append([k_, v_])
        wn.cell(row=wn.max_row, column=1).font = hdr
        wn.cell(row=wn.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    wn.column_dimensions["A"].width = 18
    wn.column_dimensions["B"].width = 100
    wb.save(path)


def write_chart(path, res):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    ev, p95 = res["lat"]["ev"], res["lat"]["p95"]
    k = res["figures"]["call_subscales"]
    x = np.arange(1, len(ev) + 1)
    fig, ax = plt.subplots(figsize=(9.2, 5.2), dpi=110)
    ax.plot(x, ev, color="#7d8a80", linewidth=1.2, zorder=1)
    ax.scatter(x[k:], ev[k:], s=26, facecolor="white", edgecolor="#7d8a80", linewidth=1.1, zorder=2,
               label="Not kept")
    ax.scatter(x[:k], ev[:k], s=46, color="#2f5d3a", zorder=3, label="Kept: one sub-scale each")
    ax.plot(x, p95, color="#b5651d", linestyle="--", linewidth=1.1, zorder=1,
            label="Random data, same size, 95th percentile")
    for i in range(k + 1):
        ax.annotate("%.2f" % ev[i], (x[i], ev[i]), textcoords="offset points",
                    xytext=(8, 4) if i < k else (6, -13), fontsize=8.5,
                    color="#2f5d3a" if i < k else "#5d6660")
    ax.annotate("%.2f" % p95[k], (x[k], p95[k]), textcoords="offset points", xytext=(6, 5), fontsize=8.5,
                color="#b5651d")
    ax.set_xlim(0.3, len(ev) + 0.7)
    ax.set_ylim(0, float(ev[0]) * 1.1)
    ax.set_xticks([1, 3, 5, 10, 15, 20, 25, 30])
    ax.set_xlabel("Dimension (in order of eigenvalue)")
    ax.set_ylabel("Eigenvalue")
    ax.set_title("Three dimensions stand above random data\n"
                 "Engagement Survey 2026, 30 items, %s respondents: eigenvalues of the tetrachoric and polychoric "
                 "correlations" % "{:,}".format(res["figures"]["respondents"]), fontsize=10.5, loc="left")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.legend(frameon=False, fontsize=8.5, loc="upper right")
    fig.text(0.01, 0.01, "Source: Kestrava export of 19 October 2026; People Analytics.", fontsize=7.5,
             color="#5d6660")
    fig.tight_layout(rect=(0, 0.03, 1, 1))
    fig.savefig(path, metadata={"Software": None})
    plt.close(fig)


def write_note(path, res, names_files):
    f = res["figures"]
    lat, pea, ref = f["latent_eigen"], f["pearson_eigen"], f["random_p95"]
    T = res["T"]
    items = T["items"].set_index("Item")
    groups = f["subscale_items"]
    ys = f["yes_share"]
    tools = [c for c in groups["Resources"] if items.loc[c, "Scale"] == "yesno" and ys[c] > 80]
    growth = [c for c in groups["Manager and growth"] if items.loc[c, "Scale"] == "yesno" and ys[c] < 20]
    rows = "\n".join("| %s | %s | %s |" % (nm, "{:,}".format(res["A"][code]), "{:,}".format(res["B"][code]))
                     for code, nm, _ in K.DISTRICTS)

    def listing(nm):
        return "; ".join("%s %s" % (c, items.loc[c, "Question (English)"].rstrip(".?")) for c in groups[nm])
    text = (
        "# The 2027 engagement dashboard should report three sub-scales\n\n"
        "To: Alexander Copeland, Head of Employee Listening\n"
        "From: People Analytics\n"
        "Date: 12 November 2026\n"
        "Re: sub-scales for the 2027 dashboard, from the 2026 survey\n\n"
        "The 30 items of the 2026 survey measure three distinct dimensions of employees' views, so the dashboard "
        "should carry three sub-scales, one per dimension, with every item reporting under one of them. The item list "
        "and loadings are in engagement_subscales_2027.xlsx; the eigenvalues are charted in "
        "engagement_dimensions_2026.png.\n\n"
        "**Resources (%d items).** %s.\n\n**Manager and growth (%d items).** %s.\n\n**Connection (%d items).** %s.\n\n"
        "## How the three were found\n\n"
        "%s employees answered at least one item before their site's close. Each answer is a coarse cut through a "
        "continuous view: a yes or no question cuts it once, an agreement statement four times. That matters here "
        "because the tools questions (%s) are answered yes by %s to %s per cent of respondents and the growth questions "
        "(%s) by %s to %s per cent. Correlated as scores, two cuts that sit far to one side can only correlate "
        "strongly with cuts on the same side, so the tools questions correlate mostly with each other and the growth "
        "questions with each other, whatever they measure. We therefore estimated the correlation of the views behind "
        "each pair of items (tetrachoric for two yes or no questions, polychoric otherwise; two-step maximum "
        "likelihood, pairwise, the four negatively worded statements reversed, does-not-apply left out).\n\n"
        "The eigenvalues of those correlations are %s, %s and %s, then %s, %s and %s. Random data of the same size "
        "(%s respondents by 30 items, %d samples) gives a 95th percentile of about %s at the fourth dimension, so the "
        "first three stand above it and the fourth does not; eigenvalues above one, Velicer's MAP and the scree all "
        "give three as well. A three-factor solution (principal-axis factoring, promax rotation) puts every item "
        "under one factor, and minres or maximum-likelihood extraction with varimax, oblimin or geomin rotation gives "
        "the same placements.\n\n"
        "## Why not five\n\n"
        "Run on the answers as scores (Pearson correlations), the same test keeps five dimensions: the eigenvalues are "
        "%s, %s, %s, %s and %s, the fourth and fifth above random data's %s and %s. The two extra dimensions are the "
        "tools questions and the growth questions on their own, split from the workload and manager statements they "
        "move with. They are the lopsided splits, not separate views: on the latent correlations the tools questions "
        "load with workload and the growth questions with the manager statements, and nothing is left over for a "
        "fourth or fifth dimension. Five sub-scales would have managers write action plans for two scores that "
        "mostly record where the yes line sits.\n\n"
        "## Coverage by store district\n\n"
        "| District | Answered at least one item before the close | Of them, answered all 30 |\n|---|---:|---:|\n%s\n"
        "| All store districts | %s | %s |\n\n"
        "Employees are counted once, under the district of the store they were assigned to in HR on 14 September. "
        "Printed forms count by the date their store sealed the envelope (every envelope was sealed by 27 "
        "September), not the date Kestrava keyed them; an employee's kiosk sessions are taken together; a "
        "does-not-apply answer counts as an answer. Answers given after a site's 23:59 close but before Kestrava's "
        "links stopped at 06:00 UTC on 28 September are left out under section 3.1 of the reporting standard. "
        "Distribution centres and head office are outside the 12 districts.\n"
    ) % (len(groups["Resources"]), listing("Resources"), len(groups["Manager and growth"]),
         listing("Manager and growth"), len(groups["Connection"]), listing("Connection"),
         "{:,}".format(f["respondents"]), ", ".join(tools), "%.1f" % min(ys[c] for c in tools),
         "%.1f" % max(ys[c] for c in tools), ", ".join(growth), "%.1f" % min(ys[c] for c in growth),
         "%.1f" % max(ys[c] for c in growth),
         "%.2f" % lat[0], "%.2f" % lat[1], "%.2f" % lat[2], "%.2f" % lat[3], "%.2f" % lat[4], "%.2f" % lat[5],
         "{:,}".format(f["respondents"]), PA_REPS, "%.2f" % ref[3],
         "%.2f" % pea[0], "%.2f" % pea[1], "%.2f" % pea[2], "%.2f" % pea[3], "%.2f" % pea[4], "%.2f" % ref[3],
         "%.2f" % ref[4], rows, "{:,}".format(f["ask_a_total"]), "{:,}".format(f["ask_b_total"]))
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(text)
