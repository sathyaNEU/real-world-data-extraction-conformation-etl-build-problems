"""The pack's record files: Kestrava's long export, submissions and distribution list, the codebook, the kiosk app's
session log, Kestrava's paper returns log, the HR assignment history, the site directory, the pulse survey's site
results and Kestrava's retail benchmark."""
from __future__ import annotations

import csv
import datetime as dt
import json

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import common as K

HDR_FONT = Font(name="Calibri", bold=True)
HDR_FILL = PatternFill("solid", fgColor="DDE4EA")


def _rng(k):
    return np.random.default_rng(K.SEED + 7000 + k)


def _emp(n):
    return "%07d" % n


def _code(it, lang):
    return it["fr_code"] if lang == "FR" else it["en_code"]


def write_responses(path, S):
    items = S["items"]
    rows = []
    for z in S["sub"]:
        for (j, c, t) in z["answers"]:
            rows.append((z["token"], z["n"], _code(items[j], z["lang"]), c, t))
        if z["enps"] is not None:
            rows.append((z["token"], z["n"], "ENPS", z["enps"][0], z["enps"][1]))
    rows.sort(key=lambda r: (r[0], r[1], r[4], r[2]))
    with open(path, "w", newline="", encoding="utf-8") as fh:
        fh.write("# Kestrava Listening | data export | answers, one row per answer\n")
        fh.write("# Client: Brackenmere Markets | Survey: BRK-ENG-2026, Engagement Survey 2026 (EN, FR)\n")
        fh.write("# Run 2026-10-19 13:42 UTC by mvang | %d answer rows | times in UTC\n" % len(rows))
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["token", "submission", "item", "answer", "answered_at"])
        for r in rows:
            w.writerow([r[0], r[1], r[2], r[3], K.iso_z(r[4])])
    return len(rows)


def write_submissions(path, S):
    rows = []
    for z in S["sub"]:
        times = [t for (_, _, t) in z["answers"]] + ([z["enps"][1]] if z["enps"] is not None else [])
        rows.append((z["token"], z["n"], z["channel"], z["lang"], z["status"], min(times), max(times),
                     len(z["answers"])))
    rows.sort(key=lambda r: (r[0], r[1]))
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["token", "submission", "channel", "language", "status", "first_answered_at", "last_answered_at",
                    "items_answered"])
        for r in rows:
            w.writerow([r[0], r[1], r[2], r[3], r[4], K.iso_z(r[5]), K.iso_z(r[6]), r[7]])
    return len(rows)


def write_distribution_list(path, S):
    rows = sorted(S["invite"], key=lambda v: v["token"])
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["token", "employee_id", "site_no", "language", "invitation", "list_loaded"])
        for v in rows:
            w.writerow([v["token"], _emp(v["emp"]), v["site"], v["lang"], v["channel"], K.INVITE_LOADED.isoformat()])
    return len(rows)


def write_kiosk(path, S):
    rows = sorted(S["kiosk"], key=lambda v: (v["signed_in"], v["session"]))
    with open(path, "w", encoding="utf-8") as fh:
        for v in rows:
            off = dt.timedelta(hours=K.TZ_OFFSET[v["tz"]])
            rec = {"session": v["session"], "store": v["store"], "kiosk": v["kiosk"], "employee": v["emp_typed"],
                   "signed_in": (v["signed_in"] + off).strftime("%Y-%m-%d %H:%M:%S"),
                   "closed": (v["closed"] + off).strftime("%Y-%m-%d %H:%M:%S"), "close_reason": v["reason"],
                   "app": "listening-kiosk 3.4.2"}
            fh.write(json.dumps(rec, ensure_ascii=False) + "\n")
    return len(rows)


def write_paper(path, S):
    by_no = {s["site_no"]: s for s in S["sites"]}
    wb = Workbook()
    ws = wb.active
    ws.title = "Returns"
    ws["A1"] = "Kestrava Listening, data entry"
    ws["A1"].font = Font(name="Calibri", bold=True, size=12)
    ws["A2"] = "Brackenmere Markets, Engagement Survey 2026 (BRK-ENG-2026): printed forms returned by stores"
    ws["A3"] = "Prepared 9 October 2026. Forms are keyed under the token printed on each form."
    head = ["Envelope", "Site", "Store", "Sealed", "Received", "Keyed", "Forms received", "Forms keyed",
            "Blank or spoiled"]
    ws.append([])
    ws.append(head)
    for c in ws[5]:
        c.font = HDR_FONT
        c.fill = HDR_FILL
    n = 0
    for key in sorted(S["envelopes"]):
        e = S["envelopes"][key]
        s = by_no[e["site"]]
        keyed = len(e["forms"])
        ws.append(["BKM-%d-%d" % (e["site"], e["n"]), e["site"], s["name"], e["sealed"], e["received"], e["keyed"],
                   keyed + e["blank"], keyed, e["blank"]])
        for col in (4, 5, 6):
            ws.cell(row=ws.max_row, column=col).number_format = "yyyy-mm-dd"
        n += 1
    for col, w in zip("ABCDEFGHI", (15, 7, 34, 12, 12, 12, 15, 12, 16)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A6"
    wb.save(path)
    return n


def write_hris(path, S):
    rows = []
    for p in S["people"]:
        for h in p["history"]:
            rows.append((_emp(p["emp"]), h["site"], h["dept"], h["dept_name"], h["start"], h["end"], h["action"]))
    rows.sort(key=lambda r: (r[0], r[4], r[6] == "termination"))
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["employee_id", "site_no", "department_code", "department", "effective_from", "effective_to",
                    "action"])
        for r in rows:
            w.writerow([r[0], r[1], r[2], r[3], r[4].isoformat(), r[5].isoformat() if r[5] else "", r[6]])
    return len(rows)


def write_sites(path, S):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["site_no", "legacy_site_no", "site_name", "site_type", "format", "city", "province", "time_zone",
                    "district_code", "district", "opened"])
        for s in sorted(S["sites"], key=lambda v: v["site_no"]):
            w.writerow([s["site_no"], s["legacy_no"], s["name"], s["type"], s["format"], s["city"], s["province"],
                        s["tz"], s["district"], s["district_name"], s["opened"].isoformat()])
    return len(S["sites"])


def write_codebook(path, S):
    items = S["items"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Survey"
    meta = [("Survey", "BRK-ENG-2026, Engagement Survey 2026"), ("Client", K.CHAIN),
            ("Languages", "English (EN), French (FR)"), ("Opens", "14 September 2026"),
            ("Closes", "27 September 2026, 23:59 local time"), ("Channels", "personal link, store kiosk, printed form"),
            ("Account manager", K.PEOPLE["vendor"]), ("Codebook version", "2026.3, 2 September 2026")]
    for k, v in meta:
        ws.append([k, v])
        ws.cell(row=ws.max_row, column=1).font = HDR_FONT
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 62
    wi = wb.create_sheet("Items")
    wi.append(["Item", "French code", "Kestrava library item", "Question (English)", "Question (French)", "Scale",
               "Reversed"])
    for c in wi[1]:
        c.font = HDR_FONT
        c.fill = HDR_FILL
    for key in K.EN_ORDER:
        it = next(x for x in items if x["key"] == key)
        wi.append([it["en_code"], it["fr_code"], it["library"], it["en"], it["fr"],
                   "agree5" if it["kind"] == "a5" else "yesno", "yes" if it["rev"] else "no"])
    wi.append(["ENPS", "ENPS", "KL-0001", K.ENPS_EN, K.ENPS_FR, "nps11", "no"])
    for col, w in zip("ABCDEFG", (7, 12, 20, 70, 74, 9, 9)):
        wi.column_dimensions[col].width = w
    for row in wi.iter_rows(min_row=2):
        for c in row[3:5]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wc = wb.create_sheet("Response codes")
    wc.append(["Scale", "Code", "Label (English)", "Label (French)"])
    for c in wc[1]:
        c.font = HDR_FONT
        c.fill = HDR_FILL
    for code, en, fr in [(1, "Strongly disagree", "Fortement en désaccord"), (2, "Disagree", "En désaccord"),
                         (3, "Neither agree nor disagree", "Ni en accord ni en désaccord"), (4, "Agree", "En accord"),
                         (5, "Strongly agree", "Fortement en accord"), (6, "Does not apply to me",
                                                                        "Ne s'applique pas à moi")]:
        wc.append(["agree5", code, en, fr])
    for code, en, fr in [(1, "Yes", "Oui"), (2, "No", "Non"), (6, "Does not apply to me", "Ne s'applique pas à moi")]:
        wc.append(["yesno", code, en, fr])
    wc.append(["nps11", "0-10", "0 not at all likely to 10 extremely likely", "0 pas du tout probable à 10 "
                                                                               "extrêmement probable"])
    wc.append([])
    wc.append(["An item left blank has no row in the answers export."])
    for col, w in zip("ABCD", (9, 7, 40, 40)):
        wc.column_dimensions[col].width = w
    wb.save(path)


def write_pulse(path, S):
    """The quarterly pulse survey's site results for Q2 2026 (five pulse items, favourable share)."""
    r = _rng(1)
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["site_no", "quarter", "fielded", "invited", "responses", "p1_fav_pct", "p2_fav_pct", "p3_fav_pct",
                    "p4_fav_pct", "p5_fav_pct"])
        n = 0
        for s in sorted(S["sites"], key=lambda v: v["site_no"]):
            inv = int(s["size"] * r.uniform(0.86, 0.97))
            resp = int(inv * r.uniform(0.22, 0.47))
            fav = [r.uniform(48, 86), r.uniform(55, 90), r.uniform(38, 79), r.uniform(44, 83), r.uniform(51, 88)]
            w.writerow([s["site_no"], "2026-Q2", "2026-05-11 to 2026-05-22", inv, resp] +
                       ["%.1f" % v for v in fav])
            n += 1
    return n


# Kestrava's library wording where Brackenmere's questionnaire adapted an item to its own stores and programmes
LIBRARY_WORDING = {
    "T1": "In the past four weeks, was the equipment you use most working whenever you needed it?",
    "T5": "Can you get into the scheduling system when you need to?",
    "G5": "In the past twelve months, has someone outside your own team talked with you about your career?",
    "G6": "Have you been invited to apply to a leadership development programme?",
}


def write_benchmark(path, S):
    """Kestrava's retail benchmark: favourable share of each library item in its retail clients' 2025 surveys."""
    r = _rng(2)
    items = S["items"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Retail 2025"
    ws["A1"] = "Kestrava Listening, retail benchmark 2026 edition (surveys fielded January to December 2025)"
    ws["A1"].font = Font(name="Calibri", bold=True, size=12)
    ws["A2"] = "Favourable: agree or strongly agree on agreement items; yes on yes or no items, favourable sense."
    ws.append([])
    ws.append(["Library item", "Item (Kestrava library wording)", "Clients", "Respondents", "Favourable %",
               "25th percentile client", "75th percentile client"])
    for c in ws[4]:
        c.font = HDR_FONT
        c.fill = HDR_FILL
    order = sorted(items, key=lambda it: it["library"])
    for it in order:
        if it["kind"] == "yn":
            fav = float(np.clip(100 * it["p"] + r.normal(0, 2.5), 1.5, 98.5))
        else:
            fav = float(np.clip(r.normal(48, 7) if it["rev"] else r.normal(63, 8), 20, 90))
        clients = int(r.integers(9, 31))
        lg = np.log(fav / (100 - fav))
        q25 = 100 / (1 + np.exp(-(lg - r.uniform(0.18, 0.42))))
        q75 = 100 / (1 + np.exp(-(lg + r.uniform(0.18, 0.42))))
        ws.append([it["library"], LIBRARY_WORDING.get(it["key"], it["en"]), clients, int(clients * r.uniform(1800, 6400)),
                   round(fav, 1), round(float(q25), 1), round(float(q75), 1)])
    for col, w in zip("ABCDEFG", (14, 84, 9, 12, 13, 22, 22)):
        ws.column_dimensions[col].width = w
    wb.save(path)
    return len(order)
