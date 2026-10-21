"""The pack's workbooks: the inspection program's station register, the Department's calendar year 2027 NOx rate
table and the pilot's voucher register, each laid out the way the office that keeps it keeps it."""
import datetime as dt

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import pilot as P
import stations as S
import world as W

HEAD = Font(bold=True)
FILL = PatternFill("solid", fgColor="DCE6F1")


def _sheet(ws, header, rows, widths, numfmt=None):
    ws.append(header)
    for c in ws[1]:
        c.font = HEAD
        c.fill = FILL
        c.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append([None if v == "" else v for v in r])
    for k, w in enumerate(widths):
        ws.column_dimensions[chr(65 + k)].width = w
    if numfmt:
        for col, fmt in numfmt.items():
            for cell in ws[col][1:]:
                cell.number_format = fmt
    ws.freeze_panes = "A2"


STATION_COLS = ["station_no", "station_name", "street", "town", "county", "county_no", "licensed_since", "lanes",
                "licence_status"]


def write_stations(path):
    reg = S.register()
    rows = []
    for x in sorted((x for x in reg.s if x["end"] is None), key=lambda x: x["no"]):
        rows.append([x["no"], x["name"], x["street"], x["town"], x["county"], "%02d" % W.CNO[x["county"]], x["start"],
                     x["lanes"], "ACTIVE"])
    wb = Workbook()
    ws = wb.active
    ws.title = "Licensed stations"
    _sheet(ws, STATION_COLS, rows, [11, 30, 22, 14, 13, 10, 14, 7, 13], {"G": "yyyy-mm-dd"})
    wb.save(path)
    return len(rows)


def write_rates(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Table 2"
    rows = [[my, pc, lt] for my, (pc, lt) in sorted(W.RATES.items())]
    _sheet(ws, ["model_year", "passenger_car_g_per_mi", "light_truck_g_per_mi"], rows, [12, 24, 22],
           {"B": "0.000", "C": "0.000"})
    ns = wb.create_sheet("Notes")
    notes = [["Table 2. Running exhaust NOx, grams per mile, calendar year 2027"],
             ["Gasoline cars (program class PC) and gasoline light trucks of 8,500 pounds gross vehicle weight or less "
              "(program class LT), by model year."],
             ["Annual average rates from the Section's MOVES run for the Halvering Basin: calendar year 2027, the area's "
              "summer and winter gasoline, its temperature and humidity, the Enhanced Inspection Program's OBD credits, "
              "all road types weighted by the area's 2025 travel."],
             ["Each rate is the model year's fleet average at its calendar year 2027 age and includes deterioration; "
              "start and evaporative emissions are not included."],
             ["Flex-fuel vehicles take the gasoline rate. Rates are not provided for diesel, hybrid or electric vehicles."],
             ["Issued 15 April 2026, Mobile Sources Section."]]
    for r in notes:
        ns.append(r)
    ns.column_dimensions["A"].width = 120
    for row in ns.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ns["A1"].font = HEAD
    wb.save(path)


def write_vouchers(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Vouchers"
    rows = P.voucher_rows()
    _sheet(ws, P.VOUCHER_COLS, rows, [13, 10, 20, 10, 11, 8, 16, 10, 12, 14, 12, 14, 10, 12, 8],
           {"I": "yyyy-mm-dd", "J": "yyyy-mm-dd", "N": "yyyy-mm-dd", "M": "#,##0"})
    ws.auto_filter.ref = "A1:O%d" % (len(rows) + 1)
    wb.save(path)
    return len(rows)
