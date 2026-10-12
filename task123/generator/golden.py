"""The goldens, computed from target/ only: the January 2026 county series (CSV), the 2026-27 grant split
(XLSX) and its chart (PNG), and generator/figures.json with every figure the write-up states.

The split calendarizes every residential bill that touches January 2026: a bill's therms are spread over the gas
days of its service period in proportion to the mean daily use of Pine Harbor's AMI meters on each day, and the
January days' part is summed by the premise's county.
"""
import datetime as dt
import json
import os
import sys
from collections import defaultdict

import numpy as np
import pandas as pd

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt   # noqa: E402
import matplotlib.ticker           # noqa: E402,F401

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import postprocess as PP   # noqa: E402

GRANT = 18_640_000
COUNTIES = ["Ashwell", "Brennock", "Farrish", "Halvern", "Keswin", "Merrion", "Ottercreek", "Pellow", "Rennick",
            "Sauvain", "Tillery", "Wynnell"]
CODE = {c: "N%02d" % (i + 1) for i, c in enumerate(COUNTIES)}
JAN0, JAN1 = dt.date(2026, 1, 1), dt.date(2026, 1, 31)
CSV_OUT = "jan2026_county_gas_series.csv"
XLSX_OUT = "heating_grant_split_2026_27.xlsx"
PNG_OUT = "heating_grant_split_2026_27.png"


def find(target, prefix):
    hits = [f for f in os.listdir(target) if f.startswith(prefix)]
    assert len(hits) == 1, (prefix, hits)
    return os.path.join(target, hits[0])


def r5k(x):
    return int(np.floor(x / 5000.0 + 0.5)) * 5000


def r10k(x):
    return int(np.floor(x / 10000.0 + 0.5)) * 10000


# ------------------------------------------------------------------ loading
def load_counties(target):
    sp = pd.read_csv(find(target, "service_points_"), dtype=str)
    return dict(zip(sp.PREMISE_ID, sp.COUNTY.str.title()))


def load_bills(target):
    """Every residential billing line with its first and last gas day and its therms."""
    k = pd.read_csv(find(target, "KVG_RES_BILLING"), dtype={"ACCOUNT_NO": str, "BILL_NO": str})
    kvg = pd.DataFrame({"util": "KVG", "premise": k.PREMISE_NO, "account": k.ACCOUNT_NO,
                        "first": pd.to_datetime(k.SVC_FROM, format="%m/%d/%Y").dt.date,
                        "last": pd.to_datetime(k.SVC_TO, format="%m/%d/%Y").dt.date, "therms": k.THERMS})
    rows = []
    with open(find(target, "NNG_R1_READS"), encoding="ascii") as fh:
        for ln in fh:
            if ln[0] != "D":
                continue
            prev = dt.datetime.strptime(ln[22:30], "%Y%m%d").date()
            cur = dt.datetime.strptime(ln[30:38], "%Y%m%d").date()
            rows.append(("NNG", ln[1:9].strip(), ln[9:17], prev + dt.timedelta(1), cur, int(ln[60:67])))
    nng = pd.DataFrame(rows, columns=["util", "premise", "account", "first", "last", "therms"])
    p = pd.read_csv(find(target, "PHE_RESIDENTIAL_BILLS"), dtype={"ACCOUNT": str})
    phe = pd.DataFrame({"util": "PHE", "premise": p.PREMISE, "account": p.ACCOUNT,
                        "first": pd.to_datetime(p.SERVICE_START).dt.date, "last": pd.to_datetime(p.SERVICE_END).dt.date,
                        "therms": p.THERMS})
    return pd.concat([kvg, nng, phe], ignore_index=True)


def ami_shape(target):
    m = pd.read_parquet(find(target, "PHE_MDMS_DAILY"))
    s = m.groupby("read_date")["usage_ccf"].mean()
    return {d: float(v) for d, v in s.items()}


# ------------------------------------------------------------------ the split
def january(bills, weights, county_of, method="shape"):
    out = defaultdict(float)
    touch = bills[(bills["last"] >= JAN0) & (bills["first"] <= JAN1)]
    for b in touch.itertuples(index=False):
        days = [b.first + dt.timedelta(i) for i in range((b.last - b.first).days + 1)]
        if method == "days":
            w = [1.0] * len(days)
        else:
            w = [weights[d] for d in days]
        jan = sum(x for d, x in zip(days, w) if JAN0 <= d <= JAN1)
        out[county_of[b.premise]] += b.therms * jan / sum(w)
    return dict(out)


def split(totals):
    t = sum(totals.values())
    return {c: GRANT * totals[c] / t for c in COUNTIES}


# ------------------------------------------------------------------ the assistance indicators
def disconnections(target, county_of):
    out = defaultdict(int)
    acct_prem = {}
    k = pd.read_csv(find(target, "KVG_RES_BILLING"), dtype={"ACCOUNT_NO": str}, usecols=["ACCOUNT_NO", "PREMISE_NO"])
    for a, p in zip(k.ACCOUNT_NO, k.PREMISE_NO):
        acct_prem[int(a)] = p
    o = pd.read_csv(find(target, "KVG_SERVICE_ORDERS"), dtype={"ACCOUNT_NO": str})
    o = o[(o.REASON_CD == "NP") & (o.ORDER_TYPE == "OFF") & (o.STATUS == "COMP")]
    done = set()
    for a, ts in zip(o.ACCOUNT_NO, o.COMPLETED_TS):
        d = dt.datetime.strptime(ts, "%m/%d/%Y %H:%M").date()
        done.add((int(a), d))
    nl = pd.read_excel(find(target, "Northline_"), header=4)
    nl = nl[nl["Job #"].astype(str).str.startswith("NL-")]
    for a, d in zip(nl["Acct #"], nl["Date"]):
        done.add((int(a), pd.Timestamp(d).date()))
    for a, d in done:
        if JAN0 <= d <= JAN1:
            out[county_of[acct_prem[a]]] += 1
    act = pd.read_csv(find(target, "NNG_FIELD_ACTIVITY"), dtype=str)
    crew = pd.read_csv(find(target, "NNG_CREW_COMPLETIONS"), dtype=str)
    j = act[(act.ACT_CD == "DNP") & (act.RESULT_CD == "C")].merge(crew, on="ACTIVITY_ID")
    for p, ts in zip(j.PREMISE_ID, j.COMPLETED_AT):
        if JAN0 <= dt.datetime.strptime(ts, "%Y-%m-%d %H:%M").date() <= JAN1:
            out[county_of[p]] += 1
    with open(find(target, "PHE_WORK_ORDERS")) as fh:
        wo = json.load(fh)["workOrders"]
    for w in wo:
        if w["woType"] == "Collections Shutoff" and w["status"] == "Completed":
            if JAN0 <= dt.date.fromisoformat(w["completedAt"][:10]) <= JAN1:
                out[county_of[w["premiseId"]]] += 1
    mh = pd.read_csv(find(target, "PHE_METER_ACCOUNT"), dtype=str)
    meter_prem = dict(zip(mh.METER_NO, mh.PREMISE_ID))
    ev = pd.read_csv(find(target, "PHE_HES_METER_EVENTS"), dtype=str)
    ev = ev[(ev.EVENT_CODE == "VALVE_CLOSE") & (ev.REASON_CD == "NP")]
    for m, ts in zip(ev.METER_ID, ev.EVENT_TS_UTC):
        local = (dt.datetime.strptime(ts, "%Y-%m-%dT%H:%M:%SZ") - dt.timedelta(hours=5)).date()
        if JAN0 <= local <= JAN1:
            out[county_of[meter_prem[m]]] += 1
    return {c: out.get(c, 0) for c in COUNTIES}


def on_plan(target, county_of):
    out = defaultdict(int)
    k = pd.read_csv(find(target, "KVG_RES_BILLING"), dtype={"ACCOUNT_NO": str}, usecols=["ACCOUNT_NO", "PREMISE_NO"])
    acct_prem = dict(zip(k.ACCOUNT_NO, k.PREMISE_NO))
    e = pd.read_csv(find(target, "KVG_WCP_ENROLLMENTS"), dtype=str, keep_default_na=False)
    for a, eff, end in zip(e.ACCOUNT_NO, e.EFFECTIVE_DT, e.END_DT):
        eff = dt.datetime.strptime(eff, "%m/%d/%Y").date()
        end = dt.datetime.strptime(end, "%m/%d/%Y").date() if end else None
        if eff <= JAN1 and (end is None or end >= JAN1):
            out[county_of[acct_prem[a]]] += 1
    hist = pd.read_csv(find(target, "NNG_ACCOUNT_PREMISE"), dtype=str, keep_default_na=False)
    acct_end = {a: (dt.date.fromisoformat(en) if en else None) for a, en in zip(hist.ACCOUNT_NO, hist.END_DT)}
    fl = pd.read_csv(find(target, "NNG_HHAP_PREMISE"), dtype=str)
    for p, a in zip(fl.PREMISE_ID, fl.ACCOUNT_NO):
        if acct_end[a] is None or acct_end[a] >= JAN1:
            out[county_of[p]] += 1
    mh = pd.read_csv(find(target, "PHE_METER_ACCOUNT"), dtype=str, keep_default_na=False)
    live = {}
    for a, p, s, en in zip(mh.ACCOUNT_NO, mh.PREMISE_ID, mh.LINKED_FROM, mh.LINKED_TO):
        if dt.date.fromisoformat(s) <= JAN1 and (not en or dt.date.fromisoformat(en) >= JAN1):
            live[a] = p
    pl = pd.read_csv(find(target, "PHE_PAYMENT_ARR"), dtype=str, keep_default_na=False)
    pl = pl[pl.PLAN_TYPE == "IBPP"]
    for a, s, en in zip(pl.ACCOUNT_NO, pl.START_DATE, pl.END_DATE):
        if dt.date.fromisoformat(s) <= JAN1 and (not en or dt.date.fromisoformat(en) >= JAN1) and a in live:
            out[county_of[live[a]]] += 1
    return {c: out.get(c, 0) for c in COUNTIES}


# ------------------------------------------------------------------ compute
def compute(target):
    county_of = load_counties(target)
    bills = load_bills(target)
    w = ami_shape(target)
    jan = january(bills, w, county_of)
    aw = split(jan)
    k1 = disconnections(target, county_of)
    k2 = on_plan(target, county_of)
    snap = [w[dt.date(2026, 1, d)] for d in range(1, 32)]
    return dict(jan=jan, awards=aw, k1=k1, k2=k2, weights=w, bills=bills, county_of=county_of,
                snap_ratio=float(np.mean(snap[18:]) / np.mean(snap[:18])))


def figures(c):
    aw, jan = c["awards"], c["jan"]
    order = sorted(COUNTIES, key=lambda x: -aw[x])
    return {
        "awards_usd": {k: "{:,}".format(r5k(aw[k])) for k in COUNTIES},
        "awards_order": order,
        "therms_jan2026": {k: "{:,}".format(r10k(jan[k])) for k in COUNTIES},
        "region_therms_jan2026": "{:,}".format(r10k(sum(jan.values()))),
        "awards_total_usd": "{:,}".format(sum(r5k(aw[k]) for k in COUNTIES)),
        "np_disconnections_jan2026": {k: str(c["k1"][k]) for k in COUNTIES},
        "lipp_accounts_2026_01_31": {k: str(c["k2"][k]) for k in COUNTIES},
        "cold_snap_ratio": "%.1f" % c["snap_ratio"],
        "largest": order[0], "smallest": order[-1],
        "rounding": {"award_usd": "5,000", "therms": "10,000"},
        "standard_clauses": ["2.4", "2.5", "2.6", "3.3"],
    }


# ------------------------------------------------------------------ deliverables
def write_csv(c, path):
    rows = ["county_code,county,residential_therms_jan2026,np_disconnections_completed_jan2026,"
            "accounts_on_low_income_plan_2026-01-31"]
    for k in COUNTIES:
        rows.append("%s,%s,%d,%d,%d" % (CODE[k], k, r10k(c["jan"][k]), c["k1"][k], c["k2"][k]))
    with open(path, "w", encoding="utf-8", newline="") as fh:
        fh.write("\r\n".join(rows) + "\r\n")


def write_xlsx(c, path):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    aw, jan = c["awards"], c["jan"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Split"
    ws["A1"] = "Northern Heat Assistance Grant 2026-27: county split"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("$18,640,000 split in proportion to each county's residential natural gas use in January 2026 "
                "(County Energy Data Standard, Ed. 3)")
    ws["A2"].font = Font(italic=True, size=9, color="555555")
    hdr = ["Code", "County", "Residential use, January 2026 (therms)", "Award 2026-27 ($)"]
    for j, h in enumerate(hdr):
        cell = ws.cell(row=4, column=j + 1, value=h)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="left" if j < 2 else "right", wrap_text=True, vertical="bottom")
        cell.fill = PatternFill("solid", fgColor="E8ECEF")
        cell.border = Border(bottom=Side(style="thin", color="333333"))
    for i, k in enumerate(COUNTIES):
        r = 5 + i
        ws.cell(row=r, column=1, value=CODE[k])
        ws.cell(row=r, column=2, value=k)
        a = ws.cell(row=r, column=3, value=r10k(jan[k]))
        a.number_format = "#,##0"
        b = ws.cell(row=r, column=4, value=r5k(aw[k]))
        b.number_format = "#,##0"
    tr = 5 + len(COUNTIES)
    ws.cell(row=tr, column=2, value="Region").font = Font(bold=True)
    t1 = ws.cell(row=tr, column=3, value=r10k(sum(jan.values())))
    t1.number_format = "#,##0"
    t1.font = Font(bold=True)
    t2 = ws.cell(row=tr, column=4, value=sum(r5k(aw[k]) for k in COUNTIES))
    t2.number_format = "#,##0"
    t2.font = Font(bold=True)
    for j in range(1, 5):
        ws.cell(row=tr, column=j).border = Border(top=Side(style="thin", color="333333"))
    ws.column_dimensions["A"].width = 7
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 17
    ws.row_dimensions[4].height = 31
    ws.freeze_panes = "A5"
    ws.cell(row=tr + 2, column=1, value=("Use is rounded to the nearest 10,000 therms and awards to the nearest $5,000; "
                                         "awards are computed from unrounded county use (standard 3.3) and the region "
                                         "row sums the rounded awards.")).font = Font(size=9, color="555555")
    ns = wb.create_sheet("Notes")
    notes = [
        ("Basis", "Residential (RS-1, R1, RES) use on the gas days dated 1 to 31 January 2026, at the county of each "
                  "premise in the utilities' service-point records."),
        ("Calendarization", "Bills read on cycles or every second month are spread over the gas days of their "
                            "service periods in proportion to the mean daily use of Pine Harbor's AMI meters on each "
                            "day; Pine Harbor's calendar-month AMI bills are January as billed."),
        ("Sources", "Kestrel Valley Gas, Northbank Natural Gas and Pine Harbor Energy billing extracts of 15 April 2026; "
                    "Pine Harbor MDMS daily reads to 14 April 2026."),
        ("Prepared", "Research team, Ridgeline Home Energy Fund, May 2026."),
    ]
    ns["A1"] = "Notes"
    ns["A1"].font = Font(bold=True, size=12)
    for i, (a, b) in enumerate(notes):
        ns.cell(row=3 + i, column=1, value=a).font = Font(bold=True)
        cell = ns.cell(row=3 + i, column=2, value=b)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ns.row_dimensions[3 + i].height = 46
    ns.column_dimensions["A"].width = 16
    ns.column_dimensions["B"].width = 92
    wb.save(path)
    PP.normalize_ooxml(path, "Douglas Tyler", "Regina Fowler", "2026-05-04T14:21:00Z", "2026-05-06T19:02:00Z",
                       excel=True, company="Ridgeline Home Energy Fund")


def write_png(c, path):
    aw = c["awards"]
    order = sorted(COUNTIES, key=lambda x: -aw[x])
    vals = [r5k(aw[k]) for k in order]
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 9, "svg.hashsalt": "t123", "text.parse_math": False})
    fig, ax = plt.subplots(figsize=(8.6, 4.9), dpi=150)
    ax.bar(range(len(order)), [v / 1e6 for v in vals], color="#2b5c8a", width=0.66, zorder=3)
    for i, v in enumerate(vals):
        ax.text(i, v / 1e6 + 0.03, "$" + format(v, ","), ha="center", va="bottom", fontsize=7.2, color="#1a1a1a")
    ax.set_xticks(range(len(order)))
    ax.set_xticklabels(order, rotation=35, ha="right")
    ax.set_ylabel("Award ($ million)")
    ax.set_ylim(0, max(vals) / 1e6 * 1.14)
    ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda y, _: "%.1f" % y))
    ax.yaxis.grid(True, color="#dddddd", linewidth=0.7, zorder=0)
    for s_ in ("top", "right"):
        ax.spines[s_].set_visible(False)
    ax.set_title("%s's $%s is the largest 2026-27 award and %s's $%s the smallest" % (
        order[0], format(vals[0], ","), order[-1], format(vals[-1], ",")), loc="left", fontsize=10, color="#1a1a1a")
    fig.text(0.012, 0.014, "Northern Heat Assistance Grant 2026-27: $18,640,000 split in proportion to each county's "
             "residential gas use on the gas days of January 2026,\nawards rounded to the nearest $5,000. Source: utility "
             "billing extracts of 15 April 2026 and Pine Harbor MDMS daily reads.", fontsize=6.8, color="#555555",
             linespacing=1.4)
    fig.tight_layout(rect=(0, 0.06, 1, 1))
    fig.savefig(path, metadata={"Software": None})
    plt.close(fig)


def main(target, golden_dir, fig_path):
    os.makedirs(golden_dir, exist_ok=True)
    c = compute(target)
    write_csv(c, os.path.join(golden_dir, CSV_OUT))
    write_xlsx(c, os.path.join(golden_dir, XLSX_OUT))
    write_png(c, os.path.join(golden_dir, PNG_OUT))
    figs = figures(c)
    with open(fig_path, "w") as fh:
        json.dump(figs, fh, indent=1, sort_keys=True)
        fh.write("\n")
    return c


if __name__ == "__main__":
    task = os.path.dirname(HERE)
    main(os.path.join(task, "target"), os.path.join(task, "golden"), os.path.join(HERE, "figures.json"))
