"""Independent verifier: recompute every graded figure from target/ alone on a separate code path (DuckDB SQL over
the files, no generator module), then compare with the goldens and figures.json.

    python3 task123/generator/verify.py
"""
import csv
import json
import os
import sys

import duckdb
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)


def path_of(target, prefix):
    hits = [f for f in os.listdir(target) if f.startswith(prefix)]
    assert len(hits) == 1, (prefix, hits)
    return os.path.join(target, hits[0]).replace("'", "''")


def recompute(target):
    con = duckdb.connect()
    q = lambda prefix: path_of(target, prefix)
    con.execute("""CREATE TABLE sp AS SELECT PREMISE_ID AS premise, upper(COUNTY) AS county
                   FROM read_csv('%s', all_varchar=true)""" % q("service_points_"))
    con.execute("""CREATE TABLE b AS
        SELECT PREMISE_NO AS premise, strptime(SVC_FROM, '%%m/%%d/%%Y')::DATE AS d0,
               strptime(SVC_TO, '%%m/%%d/%%Y')::DATE AS d1, CAST(THERMS AS DOUBLE) AS therms
        FROM read_csv('%s', all_varchar=true)
        UNION ALL
        SELECT trim(substr(line, 2, 8)), strptime(substr(line, 23, 8), '%%Y%%m%%d')::DATE + 1,
               strptime(substr(line, 31, 8), '%%Y%%m%%d')::DATE, CAST(substr(line, 61, 7) AS DOUBLE)
        FROM read_csv('%s', columns={'line': 'VARCHAR'}, delim='|', header=false, quote='', escape='')
        WHERE substr(line, 1, 1) = 'D'
        UNION ALL
        SELECT PREMISE, SERVICE_START::DATE, SERVICE_END::DATE, CAST(THERMS AS DOUBLE)
        FROM read_csv('%s', all_varchar=true)""" % (q("KVG_RES_BILLING"), q("NNG_R1_READS"), q("PHE_RESIDENTIAL_BILLS")))
    con.execute("""CREATE TABLE shape AS SELECT read_date AS d, avg(usage_ccf) AS w
                   FROM read_parquet('%s') GROUP BY read_date""" % q("PHE_MDMS_DAILY"))
    con.execute("""CREATE TABLE jb AS SELECT row_number() OVER () AS id, * FROM b
                   WHERE d1 >= DATE '2026-01-01' AND d0 <= DATE '2026-01-31'""")
    con.execute("""CREATE TABLE days AS
        SELECT id, premise, therms, CAST(unnest(generate_series(d0, d1, INTERVAL 1 DAY)) AS DATE) AS d FROM jb""")
    missing = con.execute("SELECT count(*) FROM days LEFT JOIN shape USING (d) WHERE shape.w IS NULL").fetchone()[0]
    assert missing == 0, "gas days without an AMI weight: %d" % missing
    jan = dict(con.execute("""
        WITH x AS (SELECT id, premise, max(therms) AS therms,
                          sum(CASE WHEN d BETWEEN DATE '2026-01-01' AND DATE '2026-01-31' THEN w ELSE 0 END) / sum(w) AS f
                   FROM days JOIN shape USING (d) GROUP BY id, premise)
        SELECT county, sum(therms * f) FROM x JOIN sp USING (premise) GROUP BY county""").fetchall())
    tot = sum(jan.values())
    awards = {c: 18640000.0 * v / tot for c, v in jan.items()}

    # Ask 1
    con.execute("""CREATE TABLE kb AS SELECT DISTINCT CAST(ACCOUNT_NO AS BIGINT) AS acct, PREMISE_NO AS premise
                   FROM read_csv('%s', all_varchar=true)""" % q("KVG_RES_BILLING"))
    ws = openpyxl.load_workbook(path_of(target, "Northline_"), read_only=True, data_only=True)["Jobs"]
    jobs = [(str(r[0]), r[1].date().isoformat(), int(r[2])) for r in ws.iter_rows(min_row=6, values_only=True)
            if r[0] and str(r[0]).startswith("NL-")]
    con.execute('CREATE TABLE nl ("Job #" VARCHAR, "Date" VARCHAR, "Acct #" BIGINT)')
    con.executemany("INSERT INTO nl VALUES (?, ?, ?)", jobs)
    k1 = {}
    rows = con.execute("""
        WITH cis AS (SELECT CAST(ACCOUNT_NO AS BIGINT) AS acct,
                            strptime(COMPLETED_TS, '%%m/%%d/%%Y %%H:%%M')::DATE AS d
                     FROM read_csv('%s', all_varchar=true)
                     WHERE REASON_CD = 'NP' AND ORDER_TYPE = 'OFF' AND STATUS = 'COMP'),
             job AS (SELECT CAST("Acct #" AS BIGINT) AS acct, CAST("Date" AS DATE) AS d FROM nl
                     WHERE "Job #" LIKE 'NL-%%'),
             kv AS (SELECT * FROM cis UNION SELECT * FROM job)
        SELECT county, count(*) FROM kv JOIN kb USING (acct) JOIN sp USING (premise)
        WHERE d BETWEEN DATE '2026-01-01' AND DATE '2026-01-31' GROUP BY county""" % q("KVG_SERVICE_ORDERS")).fetchall()
    for c, n in rows:
        k1[c] = k1.get(c, 0) + n
    rows = con.execute("""
        SELECT county, count(*) FROM read_csv('%s', all_varchar=true) a
        JOIN read_csv('%s', all_varchar=true) c USING (ACTIVITY_ID)
        JOIN sp ON sp.premise = a.PREMISE_ID
        WHERE a.ACT_CD = 'DNP' AND a.RESULT_CD = 'C'
          AND CAST(c.COMPLETED_AT AS TIMESTAMP)::DATE BETWEEN DATE '2026-01-01' AND DATE '2026-01-31'
        GROUP BY county""" % (q("NNG_FIELD_ACTIVITY"), q("NNG_CREW_COMPLETIONS"))).fetchall()
    for c, n in rows:
        k1[c] = k1.get(c, 0) + n
    with open(path_of(target, "PHE_WORK_ORDERS")) as fh:
        wo = json.load(fh)["workOrders"]
    con.execute("CREATE TABLE wo (premise VARCHAR, wtype VARCHAR, status VARCHAR, done VARCHAR)")
    con.executemany("INSERT INTO wo VALUES (?, ?, ?, ?)", [(w["premiseId"], w["woType"], w["status"], w["completedAt"])
                                                           for w in wo])
    rows = con.execute("""
        SELECT county, count(*) FROM wo JOIN sp USING (premise)
        WHERE wtype = 'Collections Shutoff' AND status = 'Completed'
          AND CAST(substr(done, 1, 10) AS DATE) BETWEEN DATE '2026-01-01' AND DATE '2026-01-31' GROUP BY county""").fetchall()
    for c, n in rows:
        k1[c] = k1.get(c, 0) + n
    con.execute("""CREATE TABLE mh AS SELECT * FROM read_csv('%s', all_varchar=true)""" % q("PHE_METER_ACCOUNT"))
    rows = con.execute("""
        SELECT county, count(*) FROM read_csv('%s', all_varchar=true) e
        JOIN (SELECT DISTINCT METER_NO, PREMISE_ID FROM mh) m ON m.METER_NO = e.METER_ID
        JOIN sp ON sp.premise = m.PREMISE_ID
        WHERE EVENT_CODE = 'VALVE_CLOSE' AND REASON_CD = 'NP'
          AND (strptime(EVENT_TS_UTC, '%%Y-%%m-%%dT%%H:%%M:%%SZ') - INTERVAL 5 HOUR)::DATE
              BETWEEN DATE '2026-01-01' AND DATE '2026-01-31'
        GROUP BY county""" % q("PHE_HES_METER_EVENTS")).fetchall()
    for c, n in rows:
        k1[c] = k1.get(c, 0) + n

    # Ask 2
    k2 = {}
    rows = con.execute("""
        SELECT county, count(*) FROM read_csv('%s', all_varchar=true) e
        JOIN kb ON kb.acct = CAST(e.ACCOUNT_NO AS BIGINT) JOIN sp USING (premise)
        WHERE strptime(EFFECTIVE_DT, '%%m/%%d/%%Y')::DATE <= DATE '2026-01-31'
          AND (END_DT IS NULL OR strptime(END_DT, '%%m/%%d/%%Y')::DATE >= DATE '2026-01-31')
        GROUP BY county""" % q("KVG_WCP_ENROLLMENTS")).fetchall()
    for c, n in rows:
        k2[c] = k2.get(c, 0) + n
    rows = con.execute("""
        SELECT county, count(*) FROM read_csv('%s', all_varchar=true) f
        JOIN read_csv('%s', all_varchar=true) h ON h.ACCOUNT_NO = f.ACCOUNT_NO
        JOIN sp ON sp.premise = f.PREMISE_ID
        WHERE h.END_DT IS NULL OR CAST(h.END_DT AS DATE) >= DATE '2026-01-31'
        GROUP BY county""" % (q("NNG_HHAP_PREMISE"), q("NNG_ACCOUNT_PREMISE"))).fetchall()
    for c, n in rows:
        k2[c] = k2.get(c, 0) + n
    rows = con.execute("""
        SELECT county, count(*) FROM read_csv('%s', all_varchar=true) p
        JOIN mh ON mh.ACCOUNT_NO = p.ACCOUNT_NO AND CAST(mh.LINKED_FROM AS DATE) <= DATE '2026-01-31'
             AND (mh.LINKED_TO IS NULL OR CAST(mh.LINKED_TO AS DATE) >= DATE '2026-01-31')
        JOIN sp ON sp.premise = mh.PREMISE_ID
        WHERE PLAN_TYPE = 'IBPP' AND CAST(START_DATE AS DATE) <= DATE '2026-01-31'
          AND (END_DATE IS NULL OR CAST(END_DATE AS DATE) >= DATE '2026-01-31')
        GROUP BY county""" % q("PHE_PAYMENT_ARR")).fetchall()
    for c, n in rows:
        k2[c] = k2.get(c, 0) + n
    return jan, awards, k1, k2


def titled(d):
    return {c.title(): v for c, v in d.items()}


def r(x, step):
    return int((x / step) + 0.5) * step


def main(target, golden):
    jan, awards, k1, k2 = (titled(x) for x in recompute(target))
    with open(os.path.join(golden, "jan2026_county_gas_series.csv"), encoding="utf-8") as fh:
        g = {row["county"]: row for row in csv.DictReader(fh)}
    wb = openpyxl.load_workbook(os.path.join(golden, "heating_grant_split_2026_27.xlsx"), data_only=True)
    ws = wb["Split"]
    gx = {ws.cell(row=i, column=2).value: (ws.cell(row=i, column=3).value, ws.cell(row=i, column=4).value)
          for i in range(5, 17)}
    with open(os.path.join(HERE, "figures.json")) as fh:
        figs = json.load(fh)
    bad = []
    for c in sorted(jan):
        t, a = r(jan[c], 10000), r(awards[c], 5000)
        if int(g[c]["residential_therms_jan2026"]) != t or gx[c][0] != t:
            bad.append((c, "therms", t, g[c]["residential_therms_jan2026"], gx[c][0]))
        if gx[c][1] != a or figs["awards_usd"][c] != "{:,}".format(a):
            bad.append((c, "award", a, gx[c][1], figs["awards_usd"][c]))
        if int(g[c]["np_disconnections_completed_jan2026"]) != k1.get(c, 0):
            bad.append((c, "k1", k1.get(c, 0), g[c]["np_disconnections_completed_jan2026"]))
        if int(g[c]["accounts_on_low_income_plan_2026-01-31"]) != k2.get(c, 0):
            bad.append((c, "k2", k2.get(c, 0), g[c]["accounts_on_low_income_plan_2026-01-31"]))
    assert len(jan) == 12, sorted(jan)
    if bad:
        raise SystemExit("verify: %d mismatches: %s" % (len(bad), bad))
    print("verify: 12 counties' January therms and awards, 12 disconnection counts and 12 plan counts recomputed "
          "from target/ alone (DuckDB) and equal to the goldens and figures.json")


if __name__ == "__main__":
    main(os.path.join(TASK, "target"), os.path.join(TASK, "golden"))
