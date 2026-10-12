"""The pack's documents: the State's county energy data standard (edition 3), the Fund's notes on the utilities'
extracts, the planning thread and the export manifest."""
import datetime as dt

from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                           # noqa: E402
from reportlab.lib.pagesizes import LETTER                                 # noqa: E402
from reportlab.lib.styles import ParagraphStyle                            # noqa: E402
from reportlab.lib.units import inch                                       # noqa: E402
from reportlab.pdfbase import pdfmetrics                                   # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                               # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import pack as P    # noqa: E402
import world as W   # noqa: E402

STANDARD = "County_Energy_Data_Standard_Ed3_2025.pdf"
NOTES = "RHEF_data_notes_utility_extracts_2026-04.docx"
THREAD = "RE_ NHAG 2026-27 county split - January 2026 series.eml"
MANIFEST = "export_manifest_2026-04-20.csv"

FONT_DIR = "/usr/share/fonts/truetype/liberation/"
pdfmetrics.registerFont(TTFont("LiberationSerif", FONT_DIR + "LiberationSerif-Regular.ttf"))
pdfmetrics.registerFont(TTFont("LiberationSerif-Bold", FONT_DIR + "LiberationSerif-Bold.ttf"))
pdfmetrics.registerFont(TTFont("LiberationSerif-Italic", FONT_DIR + "LiberationSerif-Italic.ttf"))
SERIF, SERIF_B, SERIF_I = "LiberationSerif", "LiberationSerif-Bold", "LiberationSerif-Italic"
INK = colors.HexColor("#1a1a1a")
GREY = colors.HexColor("#555555")
RULE = colors.HexColor("#9a9a9a")
BODY = ParagraphStyle("body", fontName=SERIF, fontSize=10.6, leading=14.2, textColor=INK, spaceAfter=5)
PARA = ParagraphStyle("para", parent=BODY, leftIndent=34, firstLineIndent=-34)
SUB = ParagraphStyle("sub", parent=BODY, leftIndent=52, firstLineIndent=-18)
H1 = ParagraphStyle("h1", fontName=SERIF_B, fontSize=15, leading=19, textColor=INK, spaceAfter=2, alignment=1)
H1b = ParagraphStyle("h1b", fontName=SERIF, fontSize=11.5, leading=15, textColor=INK, spaceAfter=10, alignment=1)
H2 = ParagraphStyle("h2", fontName=SERIF_B, fontSize=11.4, leading=15, textColor=INK, spaceBefore=9, spaceAfter=4)
META = ParagraphStyle("meta", fontName=SERIF_I, fontSize=9.4, leading=12, textColor=GREY, spaceAfter=10, alignment=1)
CELL = ParagraphStyle("cell", fontName=SERIF, fontSize=9.6, leading=11.8, textColor=INK)
CELLB = ParagraphStyle("cellb", parent=CELL, fontName=SERIF_B)

ISSUER = "Office of Home Energy Assistance"


def _frame(canvas, doc):
    canvas.saveState()
    canvas.setFont(SERIF, 8.4)
    canvas.setFillColor(GREY)
    canvas.drawString(1.0 * inch, 0.62 * inch, "County Energy Data Standard for Home Energy Assistance, Edition 3")
    canvas.drawRightString(LETTER[0] - 1.0 * inch, 0.62 * inch, "Page %d" % doc.page)
    canvas.restoreState()


def _table(rows, widths):
    data = [[Paragraph(str(c), CELLB if i == 0 else CELL) for c in r] for i, r in enumerate(rows)]
    t = Table(data, colWidths=widths, hAlign="LEFT")
    t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LINEABOVE", (0, 0), (-1, 0), 0.8, INK),
                           ("LINEBELOW", (0, 0), (-1, 0), 0.5, INK), ("LINEBELOW", (0, -1), (-1, -1), 0.8, INK),
                           ("TOPPADDING", (0, 0), (-1, -1), 2.4), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.4)]))
    return t


def write_standard(path):
    doc = SimpleDocTemplate(path, pagesize=LETTER, leftMargin=1.0 * inch, rightMargin=1.0 * inch,
                            topMargin=0.9 * inch, bottomMargin=0.95 * inch, title="County Energy Data Standard, Edition 3",
                            author=ISSUER, subject="", creator="", producer="")
    s = [Paragraph("County Energy Data Standard for Home Energy Assistance", H1),
         Paragraph("Edition 3", H1b),
         Paragraph("Issued by the %s. Effective for series filed on or after 1 July 2025. Replaces Edition 2 "
                   "(March 2022)." % ISSUER, META)]
    P_ = lambda n, t: s.append(Paragraph("<b>%s</b>&nbsp;&nbsp;%s" % (n, t), PARA))
    s.append(Paragraph("1. Scope", H2))
    P_("1.1", "This standard governs the county series of residential natural gas use that grant administrators file "
              "with the Office, and the assistance indicators filed with them. Where a program's rules allocate funds "
              "in proportion to a county's residential natural gas use, the use is the series defined here.")
    P_("1.2", "Administrators compile the series from the extracts the gas utilities serving their region provide "
              "under data-sharing agreements. Utility billing conventions differ; the definitions below govern the "
              "series whatever the utility's own conventions.")
    s.append(Paragraph("2. Definitions", H2))
    P_("2.1", "<i>Residential.</i> Accounts billed on the residential rate schedules listed in Appendix A. Accounts on "
              "general service schedules are not residential, including master-metered multi-family buildings.")
    P_("2.2", "<i>Gas day.</i> The period from midnight to midnight local time, dated by the day on which it begins.")
    P_("2.3", "<i>Calendar month.</i> The gas days dated in the month. A month's use is the gas used on those gas "
              "days, whatever the dates on which meters were read or bills were issued.")
    P_("2.4", "<i>Service period.</i> The gas days a bill covers: from the gas day after the prior meter read through "
              "the gas day of the read.")
    P_("2.5", "<i>Therms.</i> Therms as billed by the utility. Where a utility meters in hundreds of cubic feet "
              "(CCF), therms are the billed CCF times the heating value the utility applied to the bill.")
    P_("2.6", "<i>County.</i> The county in which the premise served is located, as recorded in the utility's "
              "service-point records.")
    s.append(Paragraph("3. The county series", H2))
    P_("3.1", "For each county and calendar month, the residential series is the therms of residential use in the "
              "month at premises in the county, summed over every utility serving the county.")
    P_("3.2", "The series is filed within sixty days of the close of the month and revised when a utility re-renders "
              "a bill that covers the month. Each county's value is reported to the nearest 10,000 therms.")
    P_("3.3", "Allocations made in proportion to the series are computed from unrounded county values. Each county's "
              "allocation is rounded as the program's rules direct and is not adjusted to the program total.")
    s.append(Paragraph("4. Assistance indicators", H2))
    P_("4.1", "The administrator files, with each month's series, the indicators in Appendix B for every county in "
              "its region, counted from the utilities' records.")
    P_("4.2", "Indicators are counted at the county of the premise (2.6) and are not rounded.")
    s.append(Paragraph("5. Records", H2))
    P_("5.1", "Administrators keep the extracts from which a filed series was compiled for six years and make them "
              "available to the Office on request.")
    s.append(Spacer(1, 6))
    s.append(Paragraph("Appendix A. Residential rate schedules and low-income payment plans", H2))
    s.append(_table([["Utility", "Residential rate schedules", "General service schedules (not residential)",
                      "Low-income payment plan"],
                     ["Kestrel Valley Gas", "RS-1 Residential Service", "GS-1, GS-2", "Winter Care Plan (WCP)"],
                     ["Northbank Natural Gas", "R1 Residential", "G1, G2, IS", "Home Heat Assistance Plan (HHAP)"],
                     ["Pine Harbor Energy", "RES Residential", "SGS, LGS", "Income-Based Payment Plan (IBPP)"]],
                    [1.35 * inch, 1.45 * inch, 1.75 * inch, 1.95 * inch]))
    s.append(Spacer(1, 8))
    s.append(Paragraph("Appendix B. Assistance indicators", H2))
    s.append(_table([["Indicator", "Definition"],
                     ["Residential disconnections for nonpayment completed",
                      "Residential services disconnected for nonpayment of a utility bill, completed in the month."],
                     ["Residential accounts on a low-income payment plan",
                      "Residential accounts enrolled in the utility's low-income payment plan (Appendix A) on the last "
                      "day of the month. An enrollment counts from its effective date through its end date."]],
                    [2.2 * inch, 4.3 * inch]))
    s.append(Spacer(1, 8))
    s.append(Paragraph("Appendix C. Northern region counties", H2))
    rows = [["Code", "County", "Gas utilities serving residential premises"]]
    serv = {c: [W.UTIL_NAME[W.COUNTIES[W.INTERNAL[c]][0]]] for c in W.DISPLAY_ORDER}
    serv["Tillery"].append(W.UTIL_NAME["KVG"])
    for c in W.DISPLAY_ORDER:
        rows.append([W.DCODE[c], c, ", ".join(serv[c])])
    s.append(_table(rows, [0.7 * inch, 1.3 * inch, 4.5 * inch]))
    doc.build(s, onFirstPage=_frame, onLaterPages=_frame)


# ------------------------------------------------------------------ the Fund's notes (DOCX)
def write_notes(path):
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, Inches
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    for sec in d.sections:
        sec.left_margin = sec.right_margin = Inches(0.9)
        sec.top_margin = sec.bottom_margin = Inches(0.8)
    t = d.add_paragraph()
    r = t.add_run("Ridgeline Home Energy Fund")
    r.bold = True
    r.font.size = Pt(9)
    t = d.add_paragraph()
    r = t.add_run("Data notes: utility extracts for the January 2026 county series")
    r.bold = True
    r.font.size = Pt(14)
    d.add_paragraph("Research team. Compiled 17 April 2026 from the utilities' cover notes and file layouts. "
                    "Field names are the utilities' own.")

    def head(txt):
        p = d.add_paragraph()
        rr = p.add_run(txt)
        rr.bold = True
        rr.font.size = Pt(11.5)

    def para(txt):
        d.add_paragraph(txt)

    def fields(rows, widths=(1.7, 4.9)):
        tb = d.add_table(rows=0, cols=2)
        tb.style = "Table Grid"
        for a, b in rows:
            cells = tb.add_row().cells
            cells[0].text, cells[1].text = a, b
            for c in cells:
                for pp in c.paragraphs:
                    for rr in pp.runs:
                        rr.font.size = Pt(9.5)
        for row in tb.rows:
            row.cells[0].width = Inches(widths[0])
            row.cells[1].width = Inches(widths[1])
        d.add_paragraph()

    head("Kestrel Valley Gas")
    para("Residential (RS-1) premises are read on 21 cycles a month, cycle 01 early in the month and cycle 21 at the "
         "end. The billing extract has one row per billing transaction for reads from 1 October 2025 to 14 April "
         "2026.")
    fields([("BILL_NO", "Bill number. A cancel and its rebill carry the number of the bill they correct."),
            ("TXN_TYPE", "BILL original bill; CNCL cancels the original in full; RBIL the corrected bill."),
            ("SVC_FROM, SVC_TO", "First and last gas day billed."),
            ("DAYS", "Days in the service period."),
            ("READ_TYPE", "A actual read; E estimated read, trued up at the next actual read."),
            ("THERMS", "Therms billed on the transaction (negative on a cancel).")])
    para("Service orders were pulled on 3 February 2026 for the January indicators: every order with an activity "
         "date from 1 December 2025 to the pull.")
    fields([("ORDER_TYPE", "OFF turn-off; ON turn-on."),
            ("REASON_CD", "NP nonpayment; RCN reconnection after nonpayment; CUST customer request; SAFE safety; "
                          "SEAS seasonal."),
            ("STATUS", "COMP completed; UNAB unable to complete (no access); CERR completed in error, reversed; "
                       "CANC cancelled."),
            ("CREATED_TS, COMPLETED_TS", "When the order was keyed; when the field work was completed.")])
    para("Northline Utility Services worked Kestrel Valley's collections orders from 5 to 16 January 2026 under its "
         "field contract. Their job report for the batch came with their February invoice.")
    para("Winter Care Plan enrollments: one row per enrollment, current to 15 April 2026. SERVICE_COUNTY is the "
         "county as recorded on the enrollment. EFFECTIVE_DT is the first day the plan applies; END_DT is blank while "
         "it continues.")

    head("Northbank Natural Gas")
    para("R1 premises are read every second month, route by route: group A in odd months and group B in even months. "
         "The reads file is fixed width, one detail record per read with a header and a trailer record. The service "
         "period runs from the gas day after PREV-READ-DT through CURR-READ-DT. THERMS are CCF times HV-FACTOR, the "
         "heating value of the month of the read, rounded to whole therms.")
    tb = d.add_table(rows=1, cols=5)
    tb.style = "Table Grid"
    for c, h in zip(tb.rows[0].cells, ("Field", "Start", "Length", "Picture", "Content")):
        c.text = h
    for f, start, ln, pic, desc in P.NNG_LAYOUT:
        cells = tb.add_row().cells
        for c, v in zip(cells, (f, str(start), str(ln), pic, desc)):
            c.text = v
    for row in tb.rows:
        for c in row.cells:
            for pp in c.paragraphs:
                for rr in pp.runs:
                    rr.font.size = Pt(9)
    d.add_paragraph("Record length %d. Trailer: record count (9), total CCF (12), total therms (12)." % P.NNG_WIDTH)
    para("Field activity: one row per field activity for December 2025 to February 2026, with crew completions in a "
         "separate file keyed by ACTIVITY_ID.")
    fields([("ACT_CD", "DNP disconnect for nonpayment; RNP reconnect after nonpayment; TOF turn-off at customer "
                       "request; TON turn-on; LKC leak call."),
            ("ACT_DATE", "Activity date."),
            ("RESULT_CD", "C completed; U unable to complete; X cancelled before a crew was sent."),
            ("COMPLETED_AT", "Crew completion file: date and time the crew closed the activity.")])
    para("Account history: every R1 account at each premise with its start and end dates, and the account's Home Heat "
         "Assistance Plan enrollment (PROGRAM_CD, PROGRAM_START_DT, PROGRAM_END_DT) where it has one. HHAP premises: the premises "
         "carrying the Home Heat Assistance Plan indicator at close of business on 31 January 2026, with the account "
         "that enrolled (ACCOUNT_NO) and the start of its enrollment (PROGRAM_START_DT); PROGRAM_STATUS A is active.")

    head("Pine Harbor Energy")
    para("Pine Harbor exchanged the meters at part of its residential premises for AMI meters with service valves "
         "between 6 and 24 October 2025. From the day after the exchange those premises are billed on calendar "
         "months from their daily reads; the other premises are still read on 20 monthly cycles.")
    fields([("BILL_BASIS", "CYCLE read-to-read bill; CALENDAR calendar-month bill from MDMS daily reads."),
            ("SERVICE_START, SERVICE_END", "First and last gas day billed."),
            ("READ_TYPE", "ACT actual read; EST estimated; AMI daily reads.")])
    para("MDMS daily reads: one row per meter and gas day, usage in CCF to two decimals, for the days a residential "
         "account is on the meter, to 14 April 2026. Gas days run midnight to midnight Eastern time; the gas day of "
         "2 November 2025 ran 25 hours and that of 8 March 2026 ran 23. Pine Harbor applies 1.034 therms per CCF "
         "for the 2025-26 heating season.")
    fields([("vee_status", "V validated read; E estimated by the MDMS for a missing interval.")])
    para("Meter and account history: each meter and the accounts linked to it, with the dates of each link. "
         "METER_TYPE AMI-V is an AMI meter with a service valve; DIAL a manually read meter.")
    para("Work orders (JSON): residential work orders scheduled 1 December 2025 to 28 February 2026, as exported "
         "from the field work system.")
    para("Head-end meter events: AMI meter events from 1 December 2025 to 28 February 2026, times in UTC.")
    fields([("EVENT_CODE", "VALVE_CLOSE, VALVE_OPEN, VALVE_CLOSE_FAILED; HIGH_FLOW_ALARM; LOW_BATTERY; COMM_LOSS, "
                           "COMM_RESTORED; TAMPER_ALERT; TIME_SYNC."),
            ("SOURCE", "HES command sent from the head-end system; METER raised by the meter."),
            ("REASON_CD", "On valve events: NP nonpayment; SR customer service request; SF safety shutoff by the "
                          "meter; RC reconnection after payment; RS restoration after a service request.")])
    para("Payment arrangements: every arrangement on a residential account, current to 15 April 2026. PLAN_TYPE "
         "IBPP Income-Based Payment Plan; BUDGET budget billing; DPA deferred payment arrangement. STATUS ACTIVE, "
         "ENDED.")

    head("Service points")
    para("Combined by the research team from the three utilities' service-point files as at 15 April 2026: one row per "
         "premise in the twelve counties, with the utility that serves it. CITY and ZIP are the postal address.")
    d.save(path)


# ------------------------------------------------------------------ the thread (EML)
def write_thread(path):
    body1 = (
        "Douglas,\r\n\r\n"
        "Matthew confirmed this morning that the 2026-27 Northern Heat Assistance Grant for the region is "
        "$18,640,000, split across the twelve counties in proportion to each county's residential natural gas use "
        "in January 2026 under the program rules. He needs the split, with the January county series and its "
        "assistance table, by Friday 15 May so the agency agreements can go out with the September payments.\r\n\r\n"
        "Everything the utilities sent is in the shared folder now, and Alison has updated the manifest.\r\n\r\n"
        "Regina Fowler\r\nHead of Research\r\nRidgeline Home Energy Fund\r\n\r\n")
    quoted = (
        "From: Alison Pacheco <apacheco@ridgelinehef.org>\r\n"
        "Sent: Friday, April 17, 2026 4:52 PM\r\n"
        "To: Regina Fowler <rfowler@ridgelinehef.org>; Douglas Tyler <dtyler@ridgelinehef.org>\r\n"
        "Subject: NHAG 2026-27 county split - January 2026 series\r\n\r\n"
        "Pine Harbor's MDMS export came through today, so we have everything for January. Kestrel Valley's order "
        "extract is the one they pulled on 3 February for the monthly table, and Northline's job sheet came in with "
        "their invoice.\r\n\r\n"
        "Alison\r\n")
    q = "".join("> " + ln + "\r\n" if ln else ">\r\n" for ln in quoted.split("\r\n")[:-1])
    hdr = ("From: Regina Fowler <rfowler@ridgelinehef.org>\r\n"
           "To: Douglas Tyler <dtyler@ridgelinehef.org>\r\n"
           "Cc: Alison Pacheco <apacheco@ridgelinehef.org>\r\n"
           "Subject: RE: NHAG 2026-27 county split - January 2026 series\r\n"
           "Date: Mon, 20 Apr 2026 09:41:17 -0400\r\n"
           "Message-ID: <a41c7e2b90d34f18@mail.ridgelinehef.org>\r\n"
           "In-Reply-To: <7f02bd5531c84e90@mail.ridgelinehef.org>\r\n"
           "MIME-Version: 1.0\r\n"
           "Content-Type: text/plain; charset=\"utf-8\"\r\n"
           "Content-Transfer-Encoding: 8bit\r\n\r\n")
    with open(path, "w", encoding="utf-8", newline="") as fh:
        fh.write(hdr + body1 + q)


# ------------------------------------------------------------------ the manifest
def write_manifest(path, counts):
    F = P.FILES
    rows = [
        (F["kvg_bills"], "Kestrel Valley Gas", "2026-04-15", counts["kvg_bills"], "RS-1 billing transactions, reads 1 Oct 2025 to 14 Apr 2026"),
        (F["kvg_gs"], "Kestrel Valley Gas", "2026-04-15", counts["kvg_gs"], "GS-1 and GS-2 billing transactions, reads 1 Oct 2025 to 14 Apr 2026"),
        (F["kvg_orders"], "Kestrel Valley Gas", "2026-02-03", counts["kvg_orders"], "residential service orders, activity 1 Dec 2025 to 3 Feb 2026"),
        (F["northline"], "Northline Utility Services", "2026-02-09", counts["northline"], "collections jobs worked for Kestrel Valley, 5 to 16 Jan 2026"),
        (F["kvg_wcp"], "Kestrel Valley Gas", "2026-04-15", counts["kvg_wcp"], "Winter Care Plan enrollments"),
        (F["nng_reads"], "Northbank Natural Gas", "2026-04-15", counts["nng_reads"], "R1 reads and billed usage, fixed width (layout in the data notes)"),
        (F["nng_activity"], "Northbank Natural Gas", "2026-04-15", counts["nng_activity"], "residential field activity, Dec 2025 to Feb 2026"),
        (F["nng_crew"], "Northbank Natural Gas", "2026-04-15", counts["nng_crew"], "crew completions for the field activity file"),
        (F["nng_accounts"], "Northbank Natural Gas", "2026-04-15", counts["nng_accounts"], "R1 accounts by premise with start and end dates"),
        (F["nng_flags"], "Northbank Natural Gas", "2026-02-02", counts["nng_flags"], "HHAP premises at close of business 31 Jan 2026"),
        (F["nng_sched"], "Northbank Natural Gas", "2026-04-15", counts["nng_sched"], "2026-27 meter reading schedule by route"),
        (F["phe_bills"], "Pine Harbor Energy", "2026-04-15", counts["phe_bills"], "residential bills, cycle and calendar, service to 14 Apr 2026"),
        (F["phe_mdms"], "Pine Harbor Energy", "2026-04-17", counts["phe_mdms"], "MDMS daily reads, residential AMI meters"),
        (F["phe_meters"], "Pine Harbor Energy", "2026-04-15", counts["phe_meters"], "meter and account links, residential premises"),
        (F["phe_wo"], "Pine Harbor Energy", "2026-04-15", counts["phe_wo"], "residential work orders scheduled Dec 2025 to Feb 2026"),
        (F["phe_events"], "Pine Harbor Energy", "2026-04-15", counts["phe_events"], "head-end meter events, Dec 2025 to Feb 2026"),
        (F["phe_plans"], "Pine Harbor Energy", "2026-04-15", counts["phe_plans"], "payment arrangements on residential accounts"),
        (F["service_points"], "Ridgeline Home Energy Fund", "2026-04-15", counts["service_points"], "the three utilities' service-point files of 15 Apr 2026, combined; premises in the twelve counties"),
        (STANDARD, ISSUER, "2025-06-02", "", "Edition 3, effective 1 Jul 2025"),
        (NOTES, "Ridgeline Home Energy Fund", "2026-04-17", "", "field notes compiled from the utilities' cover notes"),
        (THREAD, "Ridgeline Home Energy Fund", "2026-04-20", "", "planning thread"),
    ]
    rows.sort(key=lambda r: r[0].lower())
    P.write_text(path, P._csv(rows, ["file", "provided_by", "extracted", "rows", "covers"]))
