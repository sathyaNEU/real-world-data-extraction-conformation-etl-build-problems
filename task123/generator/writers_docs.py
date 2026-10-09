"""Writers for the documents of the pack, and the container clean-up every binary goes through
(fixed in-fiction timestamps and authors, no writer signature, deterministic bytes)."""
import datetime as dt
import os
import re
import sys
import zipfile

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

ROUND_RULES = "SGF_round_rules_rev2026-06.pdf"
CUTOVER = "screen_cutover_standard_2026.docx"
BUDGET = "trustees_budget_minute_2026-27_extract.pdf"
NOTICE = "portal_form_change_notice_2024-11.pdf"
FIELD_GUIDE = "warehouse_field_guide.md"
PROVENANCE = "warehouse_extract_record_2026-10-07.md"
THREAD = "grants_team_thread_oct2026.txt"

ORG = "Ashworth Pascoe Trust"


# ----------------------------------------------------------------------------- PDF

def _styles():
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontName="Helvetica", fontSize=10, leading=13.5,
                          spaceAfter=5)
    h1 = ParagraphStyle("h1", parent=ss["Normal"], fontName="Helvetica-Bold", fontSize=14, leading=18,
                        spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=ss["Normal"], fontName="Helvetica-Bold", fontSize=10.5, leading=14,
                        spaceBefore=7, spaceAfter=3)
    small = ParagraphStyle("small", parent=body, fontSize=8.5, leading=11, textColor=colors.HexColor("#444444"))
    return body, h1, h2, small


def _pdf(path, title, story, author, footer):
    def on_page(c, d):
        c.saveState()
        c.setFont("Helvetica", 7.5)
        c.setFillColor(colors.HexColor("#555555"))
        c.drawString(20 * mm, 12 * mm, footer)
        c.drawRightString(190 * mm, 12 * mm, f"Page {d.page}")
        c.restoreState()
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm,
                            topMargin=18 * mm, bottomMargin=20 * mm, title=title, author=author,
                            subject="", creator=ORG, invariant=1)
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)


def write_round_rules(path):
    body, h1, h2, small = _styles()
    P = lambda t, s=body: Paragraph(t, s)
    st = [P(ORG, small), P("Steady Ground Fund: round rules", h1),
          P("Adopted by the trustees on 12 November 2020. Rule 2 amended by the trustees on 16 June 2026.", small),
          Spacer(1, 6)]
    rules = [
        ("1 Purpose", ["The Steady Ground Fund makes twelve-month stabilisation offers to community "
                       "organisations that hold an operating grant from the Trust and whose income has fallen, "
                       "so that a fall in income does not by itself close a service the Trust funds."]),
        ("2 Rounds", ["A round is run for each census date. The census dates are 31 March each year and, "
                      "from 2026, 30 September."]),
        ("3 Who is scored", ["3.1 Every organisation holding a current operating grant from the Trust at the "
                             "census is in scope.",
                             "3.2 An organisation is scored where its returns cover both twelve-month periods "
                             "the screen compares.",
                             "3.3 An organisation is scored once, whatever the number of grants it holds with "
                             "the Trust."]),
        ("4 The screen", ["4.1 For each organisation the screen compares its twelve-month income at the census "
                          "with the twelve months before.",
                          "4.2 The fall is the twelve months before less twelve-month income. It is stated in "
                          "dollars and as a percentage of the twelve months before.",
                          "4.3 An organisation whose fall is 10 per cent or more is offered. An organisation "
                          "whose income rose is not offered."]),
        ("5 Offers and the rate", ["5.1 An offer is the organisation's fall in dollars multiplied by the round's "
                                   "rate, rounded to the nearest whole dollar.",
                                   "5.2 An offer below $15,000 is raised to $15,000. An offer above $150,000 is "
                                   "held to $150,000.",
                                   "5.3 The rate is set in cents per dollar of fall, to two decimal places. It is "
                                   "the highest rate at which the offers together do not exceed the pot for the "
                                   "round. Any remainder stays in the Fund.",
                                   "5.4 Each round is scored afresh. An offer still being paid from an earlier "
                                   "round does not change an organisation's screen or its offer."]),
        ("6 The pot", ["The trustees set the pot for each round in the annual budget."]),
        ("7 Payment", ["Offers are paid in twelve monthly instalments. The first instalment is for June after a "
                       "March census and for December after a September census."]),
        ("8 Approval", ["The head of grants puts each round to the trustees. The trustees approve the rate and "
                        "the offers. An offer is made to the organisation, not to a grant, and is not "
                        "transferable."]),
    ]
    for head, paras in rules:
        st.append(P(head, h2))
        for t in paras:
            st.append(P(t))
    _pdf(path, "Steady Ground Fund round rules", st, ORG,
         f"{ORG}  |  Steady Ground Fund  |  Round rules, as amended 16 June 2026")


def write_budget_minute(path):
    body, h1, h2, small = _styles()
    P = lambda t, s=body: Paragraph(t, s)
    st = [P(ORG, small), P("Minutes of the meeting of trustees, 16 June 2026", h1),
          P("Extract: item 6, 2026-27 budget", small), Spacer(1, 6),
          P("<b>Present:</b> Wiremu Roberts (chair) and four trustees. <b>In attendance:</b> Mia Hart (finance "
            "manager), Marie Griffin (head of grants)."),
          P("6.1 The finance manager presented the 2026-27 budget, circulated with the papers on 9 June. "
            "Trustees asked about the timing of the first in-house Steady Ground round and the administration "
            "line. The head of grants confirmed the September round would come to the November meeting."),
          P("6.2 <b>Resolved</b> that the 2026-27 budget be approved as tabled, with grants expenditure as follows:")]
    tbl = Table([["Operating grants", "$7,450,000"], ["Project grants", "$310,000"],
                 ["Steady Ground Fund, September 2026 round", "$820,000"],
                 ["Steady Ground Fund, March 2027 round", "to be set at the February 2027 meeting"],
                 ["Grants administration and data", "$265,000"]], colWidths=[105 * mm, 55 * mm])
    tbl.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", 9.5),
                             ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#999999")),
                             ("BOTTOMPADDING", (0, 0), (-1, -1), 4), ("TOPPADDING", (0, 0), (-1, -1), 3)]))
    st += [tbl, Spacer(1, 6), P("Moved Wiremu Roberts, seconded. Carried."),
           P("6.3 The finance manager will report spending against the budget at each meeting.")]
    _pdf(path, "Trustees minute 16 June 2026, item 6", st, ORG, f"{ORG}  |  Trustees' minutes  |  16 June 2026  |  Extract")


def write_notice(path):
    body, h1, h2, small = _styles()
    P = lambda t, s=body: Paragraph(t, s)
    st = [P(f"{ORG}  |  Grantee portal", small),
          P("Quarterly financial return: changes from the December 2024 return", h1),
          P("Issued 12 November 2024 to all grantees by the grants team", small), Spacer(1, 6),
          P("From the return for the quarter ending 31 December 2024 the quarterly financial return changes in "
            "three ways. Your financial year and the year-to-date basis of the return do not change."),
          P("1 Government grants and contracts", h2),
          P("Income under contracts with central or local government, which you have reported inside "
            "‘Fees, sales and service contracts’ and shown on the memo line ‘of which government "
            "service contracts’, is now reported with government grants on one line, ‘Government grants "
            "and contracts’. ‘Fees and sales’ then holds fees and sales from everyone else."),
          P("2 Money from the Trust", h2),
          P("‘Grants from non-government sources’ now carries a memo line, ‘of which received from "
            "Ashworth Pascoe Trust’. Count Trust money in the quarter it reaches your bank account."),
          P("3 Last year's figures", h2),
          P("The return now shows last year's figures for the same period beside this year's, under the new "
            "lines. They are shown for information."),
          P("Line codes", h2)]
    rows = [["Until the September 2024 return (QFR-16)", "From the December 2024 return (QFR-24)"],
            ["TOT_INC  Total income", "TOT_REV  Total revenue"],
            ["GOV_GRT  Government grants", "GOV_GRC  Government grants and contracts"],
            ["FEE_SVC  Fees, sales and service contracts", "TRD_SAL  Fees and sales"],
            ["FEE_SVC_GOV  of which government service contracts", "now inside GOV_GRC"],
            ["DON_BEQ  Donations and bequests", "DON_BEQ  Donations and bequests"],
            ["GRT_OTH  Other grants and sponsorship", "GRT_NGO  Grants from non-government sources"],
            ["", "GRT_NGO_APT  of which received from Ashworth Pascoe Trust"],
            ["INV_INC  Investment income", "INV_REV  Investment revenue"],
            ["OTH_INC  Other income", "OTH_REV  Other revenue"]]
    tbl = Table(rows, colWidths=[85 * mm, 85 * mm])
    tbl.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", 8.5),
                             ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8.5),
                             ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#999999")),
                             ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    st += [tbl, Spacer(1, 6),
           P("The short form (QFR-16S, now QFR-24S) keeps its three lines: total, donations and bequests, and "
             "all other income."),
           P("Questions to the grants team through the portal message centre.", small)]
    _pdf(path, "Quarterly financial return changes, December 2024", st, ORG,
         f"{ORG}  |  Grantee portal notice  |  November 2024")


# ----------------------------------------------------------------------------- DOCX

def write_cutover(path):
    import docx
    from docx.shared import Pt
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    d.add_heading("Steady Ground screen: cutover standard", level=1)
    p = d.add_paragraph()
    p.add_run("Owner: ").bold = True
    p.add_run("Liam Bryant, data lead.  ")
    p.add_run("Approved: ").bold = True
    p.add_run("Marie Griffin, head of grants, 18 May 2026.  ")
    p.add_run("Version: ").bold = True
    p.add_run("1.0")
    sections = [
        ("1 Background", "Ledgerwood Analytics ran the Steady Ground screen for the Trust under contract for the "
                         "six March rounds from 2021 to 2026 and issued a screen pack for each round. The contract "
                         "ended with the March 2026 round. From the September 2026 round the screen is run in-house "
                         "from the Trust's warehouse."),
        ("2 Scope", "This standard sets when the in-house screen may be used for a round. It does not change the "
                    "round rules."),
        ("3 The published runs", "The six packs Ledgerwood issued to the Trust, SGF_screen_run_2021-03 to "
                                 "SGF_screen_run_2026-03, are the published runs. They are kept in the grants team "
                                 "folder as issued."),
        ("4 Reproduction", "The in-house screen may be used for a round only once it gives back every grantee row, "
                           "every offer and each round's rate in all six published March runs."),
        ("5 Records", "The data lead keeps the replay result with the round papers."),
        ("6 Review", "This standard is reviewed after the September 2026 round."),
    ]
    for h, t in sections:
        d.add_heading(h, level=2)
        d.add_paragraph(t)
    d.core_properties.author = "Liam Bryant"
    d.core_properties.last_modified_by = "Liam Bryant"
    d.core_properties.title = "Steady Ground screen: cutover standard"
    d.core_properties.revision = 3
    d.core_properties.created = dt.datetime(2026, 5, 11, 9, 20)
    d.core_properties.modified = dt.datetime(2026, 5, 18, 15, 5)
    d.core_properties.comments = ""
    d.save(path)


# ----------------------------------------------------------------------------- Markdown and text

FIELD_GUIDE_TEXT = """# Warehouse field guide: grants data

Maintained by Liam Bryant (data lead). Last updated 6 October 2026.

This guide describes the fields in the warehouse extracts the grants team uses. It describes what each field holds, not how to use it.

## portal_return_lines (grantee portal export)

One row per return version, line and column.

| Field | Meaning |
|---|---|
| return_id | The portal's identifier for one quarterly financial return: one grant reference, one quarter. Every version of a return shares it. |
| grant_ref | The grant the return is filed under. An organisation files a return under each grant it holds. |
| period_end | Last day of the quarter the return covers. |
| year_end | Last day of the organisation's financial year that the year-to-date figures run within. |
| form | QFR-16 up to the September 2024 return, QFR-24 from the December 2024 return. The S suffix is the short form for small grantees. |
| version_no | 1 for the first submission, then each resubmission in order. |
| version_status | accepted, rejected (failed the grants team's review) or withdrawn (withdrawn by the grantee before review). Only accepted versions are part of the record. |
| submitted_at, accepted_at | New Zealand time. accepted_at is blank unless the version was accepted. |
| line_code | See the line codes below. |
| column | YTD is the financial year to date at period_end. PY is the same period of the previous financial year as shown on the return. |
| amount | Whole New Zealand dollars. |

A quarter's record is its own return.

Line codes, QFR-16: TOT_INC total income; GOV_GRT government grants; FEE_SVC fees, sales and service contracts; FEE_SVC_GOV of which government service contracts (memo, inside FEE_SVC); DON_BEQ donations and bequests; GRT_OTH other grants and sponsorship; INV_INC investment income; OTH_INC other income.

Line codes, QFR-24: TOT_REV total revenue; GOV_GRC government grants and contracts; TRD_SAL fees and sales; DON_BEQ donations and bequests; GRT_NGO grants from non-government sources; GRT_NGO_APT of which received from Ashworth Pascoe Trust (memo, inside GRT_NGO); INV_REV investment revenue; OTH_REV other revenue.

Short forms (QFR-16S, QFR-24S): the total line, DON_BEQ, and OTH_INC or OTH_REV for all other income. A line the form does not carry was not reported.

Memo lines are part of the line above them and are not added to the total.

## grants_register (sheets Grants, Variations, Steady Ground offers)

| Field | Meaning |
|---|---|
| grant_ref | Operating grants APT-OG, project grants APT-PG. |
| programme | Operating grant or Project grant. |
| charity_no | The organisation's registration number with Charities Services. |
| organisation | The organisation's name as the Trust holds it. |
| sector, district | The grants team's classification. |
| balance_date | The organisation's financial year end. |
| start_date, end_date | Start of the grant and end of its current term. |
| annual_amount | Annual amount for the current term. |
| status | Active or Ended. |

Variations lists each renewal with its effective date and the annual amount before and after. Steady Ground offers lists every offer made in a round, with the month its first instalment pays for.

## charities_register_returns_extract (register match)

One row per annual return received by Charities Services for an organisation in the Trust's book, matched on charity number.

| Field | Meaning |
|---|---|
| match_id | Warehouse row identifier. |
| charity_no | As keyed in the match. Compare after trimming spaces and upper-casing. |
| registered_name | The name on the register. |
| year_end | End of the financial year the annual return covers. |
| date_received | Date Charities Services received the annual return. |
| return_tier | Reporting tier. |
| total_gross_income | Total gross income for the year, as filed. |
| govt_grants_contracts, donations_bequests, trading_sales, grants_other, investment_income, other_income | The warehouse's mapping of the revenue categories in the annual return: government grants and contracts; donations and bequests; fees, trading and sales; other grants, including the Trust's; investment income; other income. |

## trust_payment_run (Trust bank payments)

One row per payment line from the finance system.

| Field | Meaning |
|---|---|
| payment_id | Finance system payment identifier. |
| batch | Payment batch, named by its date. |
| value_date | Date the payment reached the payee's bank account. |
| charity_no, payee | The organisation paid. |
| grant_ref | The operating grant, project grant or Steady Ground offer the payment is made under. |
| programme | Operating grant, Project grant or Steady Ground Fund. |
| instalment_for | The month the instalment pays for under the grant or offer letter. |
| amount | Whole New Zealand dollars. |
| payment_status | paid, or returned: the payee's bank returned the payment and it was reissued. |
| reissue_of | For a reissue, the payment_id it replaces. |

## Ledgerwood screen packs

SGF_screen_run_YYYY-03.xlsx, one per March round from 2021 to 2026, as issued by Ledgerwood Analytics. Sheets Screen and Round.
"""


def write_text(path, text):
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write(text)


THREAD_TEXT = """Grants team / Steady Ground: message export
Exported 8 October 2026 09:12 NZDT by Liam Bryant

[Tue 6 Oct 2026 08:47] Marie Griffin
Morning all. The September round goes to the trustees on 11 November. Liam, when does the warehouse extract get pulled?

[Tue 6 Oct 2026 09:05] Liam Bryant
Tomorrow morning. Portal returns, the grants register, the register match and the payment run, all as at 7 October. They'll go in the round folder with the field guide. The cutover standard is in there too.

[Tue 6 Oct 2026 09:21] Tony Hughes
This round is really for the groups the March government cut hit. That's who the trustees will expect to see on the list.

[Tue 6 Oct 2026 09:38] Mia Hart
The pot for the round is in the June budget minute. I've put a copy in the folder.

[Tue 6 Oct 2026 10:02] Marie Griffin
Thanks Mia. Andrew, can you confirm what you're handing over from your side?

[Tue 6 Oct 2026 13:15] Andrew Knox
Hi Marie. As agreed when the contract wound up, we won't be running September. The six March packs we issued are the record of the screen, and they're in your folder exactly as we sent them. Happy to answer questions on the packs themselves until the end of October.

[Wed 7 Oct 2026 07:58] Liam Bryant
Extract is in the folder.

[Wed 7 Oct 2026 08:30] Wiremu Roberts
Marie, the trustees will want one rate and the list of offers, with the replay alongside. Please keep the paper short.

[Wed 7 Oct 2026 08:44] Marie Griffin
Will do. Draft to you by 28 October.

[Wed 7 Oct 2026 09:10] Tony Hughes
Happy to look over the list before it goes.
"""


def provenance_text(rows_info):
    """The folder's provenance record. rows_info: {file: row count} for the delimited files."""
    t = f"""# Extract record: Steady Ground September 2026 round

Prepared by Liam Bryant (data lead), 7 October 2026, for the grants team's round folder.

## About this folder

This folder is a constructed case. The Ashworth Pascoe Trust, its Steady Ground Fund, Ledgerwood Analytics, the Canterbury Funders' Data Group, every organisation and person named, and every figure in these files were built for this case. Charity numbers, grant references and payment identifiers follow the formats of the systems named and do not identify real registered charities, grants or payments. Charities Services and its register are real; the register match here is the Trust's own constructed extract and is not a download from the public register. All files in the folder may be used and shared under CC BY 4.0.

## Files

| File | What it is | Coverage | How it was produced |
|---|---|---|---|
| portal_return_lines_2018q3_2026q2.csv | Grantee portal export of quarterly financial returns ({rows_info['spine']:,} rows) | Every version of every return for quarters ending 30 September 2018 to 30 June 2026, all statuses, with submission and acceptance times | Portal reporting export, 7 October 2026, unedited |
| grants_register_20261007.xlsx | Grants register: grants, renewals and Steady Ground offers | Grants current at any time since 1 July 2018, as at 7 October 2026 | Warehouse snapshot, 7 October 2026 |
| charities_register_returns_extract_20261007.csv | Register match: annual returns for organisations in the Trust's book ({rows_info['register']:,} rows) | Returns received by Charities Services up to 7 October 2026, for financial years ending on or after 1 July 2018, from the first year each organisation filed with the Trust | Warehouse match on charity number, 7 October 2026 |
| SGF_screen_run_2021-03.xlsx to SGF_screen_run_2026-03.xlsx | Ledgerwood Analytics' screen packs for the six March rounds | One pack per round | Copied from the grants team folder as issued |
| SGF_round_rules_rev2026-06.pdf | Steady Ground Fund round rules | As amended 16 June 2026 | Copy of the trustees' approved rules |
| screen_cutover_standard_2026.docx | Cutover standard for the in-house screen | Version 1.0, 18 May 2026 | Data team document |
| trustees_budget_minute_2026-27_extract.pdf | Extract of the trustees' minute of 16 June 2026 | Item 6, the 2026-27 budget | Copy supplied by the finance manager |
| portal_form_change_notice_2024-11.pdf | Notice to grantees of the December 2024 return changes | Issued 12 November 2024 | Copy of the portal notice |
| trust_payment_run_2018-07_to_2026-09.csv | The Trust's payments to grantees ({rows_info['payrun']:,} rows) | Value dates 1 July 2018 to 30 September 2026 | Finance system export, 7 October 2026 |
| warehouse_field_guide.md | Field meanings for the warehouse extracts | As at 6 October 2026 | Data team document |
| grants_team_thread_oct2026.txt | Grants team messages about the round | 6 and 7 October 2026 | Message export, 8 October 2026 |
| canterbury_community_income_survey_2025.xlsx | Canterbury Funders' Data Group community sector income survey 2025 | Financial years ending in 2025 | Copy of the published workbook |
| grantee_capacity_ratings_2026.csv | The grants team's governance and capacity ratings of grantees ({rows_info['ratings']:,} rows) | Reviews from September 2025 to June 2026 | Export from the grants team's review tracker |

## Notes

- Amounts are whole New Zealand dollars throughout.
- Times in the portal export are New Zealand time.
- Nothing in the folder has been edited after extraction.
"""
    return t


# ----------------------------------------------------------------------------- containers

FIXED_ZIP_TIME = (2026, 10, 7, 9, 0, 0)


def normalize_ooxml(path, author, created, modified=None, app="Microsoft Office Word", words=None):
    """Rewrite docProps (author, dates, application), drop the template thumbnail, and repack every
    entry with a fixed timestamp, so the file carries no writer signature or build-clock date and
    is byte-stable."""
    modified = modified or created
    with zipfile.ZipFile(path) as z:
        infos = z.infolist()
        data = {i.filename: z.read(i.filename) for i in infos}
    ciso = created.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    miso = modified.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    a = author.encode()
    if "docProps/core.xml" in data:
        b = data["docProps/core.xml"]
        b = re.sub(rb"(<dc:creator[^>]*>)[^<]*(</dc:creator>)", rb"\g<1>" + a + rb"\g<2>", b)
        b = re.sub(rb"(<cp:lastModifiedBy[^>]*>)[^<]*(</cp:lastModifiedBy>)", rb"\g<1>" + a + rb"\g<2>", b)
        b = re.sub(rb"(<dcterms:created[^>]*>)[^<]*(</dcterms:created>)", rb"\g<1>" + ciso + rb"\g<2>", b)
        b = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)", rb"\g<1>" + miso + rb"\g<2>", b)
        b = re.sub(rb"<dc:description[^>]*>[^<]*</dc:description>", b"", b)
        data["docProps/core.xml"] = b
    if "docProps/app.xml" in data:
        b = data["docProps/app.xml"]
        b = re.sub(rb"<Application>[^<]*</Application>", b"<Application>" + app.encode() + b"</Application>", b)
        b = re.sub(rb"<AppVersion>[^<]*</AppVersion>", b"<AppVersion>16.0000</AppVersion>", b)
        if words is not None:
            b = re.sub(rb"<Words>[^<]*</Words>", b"<Words>%d</Words>" % words, b)
            b = re.sub(rb"<Characters>[^<]*</Characters>", b"<Characters>%d</Characters>" % (words * 6), b)
            b = re.sub(rb"<TotalTime>[^<]*</TotalTime>", b"<TotalTime>38</TotalTime>", b)
            b = re.sub(rb"<Lines>[^<]*</Lines>", b"<Lines>%d</Lines>" % (words // 11), b)
            b = re.sub(rb"<Paragraphs>[^<]*</Paragraphs>", b"<Paragraphs>%d</Paragraphs>" % (words // 40 + 1), b)
        data["docProps/app.xml"] = b
    drop = {n for n in data if n.startswith("docProps/thumbnail")}
    if drop and "_rels/.rels" in data:
        data["_rels/.rels"] = re.sub(rb"<Relationship [^>]*thumbnail[^>]*/>", b"", data["_rels/.rels"])
    if drop and "[Content_Types].xml" in data:
        data["[Content_Types].xml"] = re.sub(rb'<Default Extension="jpeg"[^>]*/>', b"",
                                             data["[Content_Types].xml"])
    tmp = path + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
        for i in infos:
            if i.filename in drop:
                continue
            zi = zipfile.ZipInfo(i.filename, date_time=FIXED_ZIP_TIME)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            zi.create_system = 0
            zo.writestr(zi, data[i.filename])
    os.replace(tmp, path)


def normalize_pdf(path, when, producer=ORG):
    """reportlab's invariant mode fixes the dates at 2000-01-01 and the ID; move the dates to the
    document's own date and replace the producer string, keeping every byte offset."""
    b = open(path, "rb").read()
    n0 = len(b)
    stamp = when.strftime("D:%Y%m%d%H%M%S").encode()
    b = re.sub(rb"D:20000101000000", stamp, b)
    m = re.search(rb"/Producer \(([^)]*)\)", b)
    if m:
        old = m.group(1)
        new = producer.encode()
        assert len(new) <= len(old)
        b = b[:m.start(1)] + new + b" " * 0 + b[m.end(1):]
        # keep offsets: pad after the closing parenthesis with spaces
        pad = len(old) - len(new)
        k = m.start(1) + len(new) + 1
        b = b[:k] + b" " * pad + b[k:]
    for tok in (rb"ReportLab Generated PDF document", rb"ReportLab", rb"reportlab"):
        for mm_ in list(re.finditer(rb"%[^\n]*" + tok + rb"[^\n]*", b)):
            line = mm_.group(0)
            b = b.replace(line, b"%" + b" " * (len(line) - 1))
        b = re.sub(tok, lambda x: b"x" * len(x.group(0)), b)
    assert len(b) == n0, "PDF length moved"
    open(path, "wb").write(b)
