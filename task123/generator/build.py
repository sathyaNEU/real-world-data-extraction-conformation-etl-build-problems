#!/usr/bin/env python3
"""task123 generator: builds the Steady Ground evidence pack into <out>/target and <out>/metadata.json,
asserting every property the design note's assertion plan lists.

    python3 build.py --out /path/to/task123 [--record /path/to/record.json]
"""
import argparse
import datetime as dt
import importlib.util
import json
import os
import re
import shutil
import sys

sys.dont_write_bytecode = True
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

from common import MARCH_CENSUSES, SEPT_CENSUS, SEPT_POT, POTS, pct1  # noqa: E402
from design import build as build_world  # noqa: E402
from tune import nudges_for  # noqa: E402
from asserts_world import (Checks, check_world, corpus_checks, flips_check,  # noqa: E402
                           convergence_checks, clean_data_checks, T_SET)
from asserts_asks import check_asks, separation  # noqa: E402
import writers_data as WD  # noqa: E402
import writers_docs as WDOC  # noqa: E402

REPO = os.path.abspath(os.path.join(HERE, "..", ".."))
SCRUB = os.path.join(REPO, ".claude/skills/reduce-house-fixes/scripts/scrub_producer_metadata.py")
EXPORT_TIME = dt.datetime(2026, 10, 7, 9, 0, 0)
DISTRACTORS = [WD.SURVEY, WD.RATINGS]


def load_scrub():
    spec = importlib.util.spec_from_file_location("scrub_producer_metadata", SCRUB)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def tuned_world():
    nud = {}
    W = build_world()
    for _ in range(8):
        n, rem, over = nudges_for(W)
        if not n:
            break
        for k, v in n.items():
            nud[k] = nud.get(k, 0) + v
        W = build_world(nudges=nud)
    return W


def write_pack(W, target):
    if os.path.isdir(target):
        shutil.rmtree(target)
    os.makedirs(target)
    p = lambda f: os.path.join(target, f)
    rows = WD.portal_rows(W)
    W["portal_rows"] = rows
    info = {"spine": WD.write_spine(p(WD.SPINE), rows)}
    WD.write_grants_register(p(WD.GRANTS), W)
    info["register"] = WD.write_register(p(WD.REGISTER), W)
    for c in MARCH_CENSUSES:
        WD.write_pack(p(WD.pack_name(c.year)), W, c)
    info["payrun"] = WD.write_payrun(p(WD.PAYRUN), W)
    WD.write_survey(p(WD.SURVEY))
    info["ratings"] = WD.write_ratings(p(WD.RATINGS), W)
    WDOC.write_round_rules(p(WDOC.ROUND_RULES))
    WDOC.write_budget_minute(p(WDOC.BUDGET))
    WDOC.write_notice(p(WDOC.NOTICE))
    WDOC.write_cutover(p(WDOC.CUTOVER))
    WDOC.write_text(p(WDOC.FIELD_GUIDE), WDOC.FIELD_GUIDE_TEXT)
    WDOC.write_text(p(WDOC.THREAD), WDOC.THREAD_TEXT)
    WDOC.write_text(p(WDOC.PROVENANCE), WDOC.provenance_text(info))
    # containers: in-fiction authors and dates, no writer signature, fixed zip timestamps
    scrub = load_scrub()
    pdf_dates = {WDOC.ROUND_RULES: "2026-06-16", WDOC.BUDGET: "2026-06-17", WDOC.NOTICE: "2024-11-12"}
    for f, d in pdf_dates.items():
        scrub.scrub_pdf(p(f), "Ashworth Pascoe Trust", d, "2018-07-01", "2026-10-08")
    issued = {2021: dt.datetime(2021, 4, 16, 10), 2022: dt.datetime(2022, 4, 14, 10),
              2023: dt.datetime(2023, 4, 18, 10), 2024: dt.datetime(2024, 4, 16, 10),
              2025: dt.datetime(2025, 4, 15, 10), 2026: dt.datetime(2026, 4, 14, 10)}
    for c in MARCH_CENSUSES:
        WDOC.normalize_ooxml(p(WD.pack_name(c.year)), "Ledgerwood Analytics", issued[c.year],
                             app="Microsoft Excel")
    WDOC.normalize_ooxml(p(WD.GRANTS), "Liam Bryant", EXPORT_TIME, app="Microsoft Excel")
    WDOC.normalize_ooxml(p(WD.SURVEY), "Canterbury Funders' Data Group", dt.datetime(2026, 5, 20, 11),
                         app="Microsoft Excel")
    words = len(file_text(p(WDOC.CUTOVER)).split()) // 2
    WDOC.normalize_ooxml(p(WDOC.CUTOVER), "Liam Bryant", dt.datetime(2026, 5, 11, 9, 20),
                         dt.datetime(2026, 5, 18, 15, 5), app="Microsoft Office Word", words=words)
    stamp = EXPORT_TIME.timestamp()
    for f in sorted(os.listdir(target)):
        os.utime(p(f), (stamp, stamp))
    return info


def file_text(path):
    """Plain text of any shipped file, for the sweeps."""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".csv", ".md", ".txt"):
        return open(path, encoding="utf-8").read()
    if ext in (".xlsx", ".docx"):
        out = []
        with zipfile.ZipFile(path) as z:
            for n in z.namelist():
                if n.endswith(".xml"):
                    out.append(re.sub(r"<[^>]+>", " ", z.read(n).decode("utf-8", "ignore")))
        return "\n".join(out)
    if ext == ".pdf":
        from pypdf import PdfReader
        return "\n".join(pg.extract_text() for pg in PdfReader(path).pages)
    return ""


def check_files(W, C, target, info):
    files = sorted(os.listdir(target))
    p = lambda f: os.path.join(target, f)
    fmts = sorted({os.path.splitext(f)[1].lower() for f in files})
    # A51: input gates
    C.ok("A51", len(files) == 19, f"{len(files)} files in the pack")
    C.ok("A51", len(fmts) >= 3, f"{len(fmts)} formats: {', '.join(fmts)}")
    C.ok("A51", info["spine"] >= 60000, f"spine {info['spine']:,} rows (gate 25,000)")
    C.ok("A51", all(d in files for d in DISTRACTORS), "both neighbouring files present in the pack")
    texts = {f: file_text(p(f)) for f in files}
    C.ok("A51", not any("distractor" in f.lower() or "distractor" in t.lower() for f, t in texts.items()),
         "no file name or line in the pack names a distractor")
    # A20 (read back): the written packs give back the generator's T screens exactly
    import openpyxl
    for c in MARCH_CENSUSES:
        wb = openpyxl.load_workbook(p(WD.pack_name(c.year)), read_only=True, data_only=True)
        ws = wb["Screen"]
        got = [r for r in ws.iter_rows(min_row=5, values_only=True) if r[0]]
        want = [(W["by"][r["org"]].cc, W["by"][r["org"]].name, r["cur"], r["prior"], r["fall"], pct1(r["pct"]),
                 r["offer"] or None) for r in W["corpus"][c]["rows"]]
        rate = [r for r in wb["Round"].iter_rows(values_only=True) if r[0] and r[0].startswith("Rate")][0][1]
        C.ok("A20", got == want and round(rate * 100) == W["corpus"][c]["rate"],
             f"pack {c.year} written as computed: {len(got)} rows, rate {rate:.2f}")
    # A39: no shipped sentence or header states a window end, a fourth-quarter source or knowledge time
    bad = re.compile(r"twelve months to|months to 3[01]|year to december|four quarters|fourth quarter|"
                     r"trailing|knowledge time|as held|as at the census|step(ped)? back|provisional|"
                     r"management (return|figure|accounts)|audited quarter|window|balance.date|"
                     r"nine.month year|short(ened)? year|transitional year", re.I)
    hits = [(f, m.group(0)) for f, t in texts.items() for m in [bad.search(t)] if m]
    C.ok("A39", not hits, f"no window, fourth-quarter, knowledge-time or balance-date statement anywhere ({hits})")
    # A40: no shipped artifact ranks September's grantees
    offer_files = [f for f in files if f.endswith(".xlsx") and "Offer ($)" in texts[f]]
    C.ok("A40", sorted(offer_files) == sorted(WD.pack_name(c.year) for c in MARCH_CENSUSES),
         "only the six March packs carry offers; nothing scores the September census")
    C.ok("A40", not any("30 September 2026" in t and "Offer" in t for t in texts.values()),
         "no file pairs the September census with offers")
    # single-statement invariant: each load-bearing fact in exactly one document
    docs = [f for f in files if os.path.splitext(f)[1] in (".pdf", ".docx", ".md", ".txt")]
    facts = {"the September pot": r"\$?560,000", "the floor": r"\$15,000", "the cap": r"\$150,000",
             "the line": r"10 per cent or more", "the reproduction clause": r"gives back every grantee row",
             "the pay day": r"20th\s+of\s+the\s+month\s+before\s+the\s+month\s+it\s+is\s+for",
             "the recency clause": r"ending\s+on\s+or\s+after\s+the\s+census\s+before\s+it",
             "a quarter's own return": r"quarter's record is its own return",
             "accepted versions only": r"Only accepted versions are part of the record"}
    for name, pat in facts.items():
        where = [f for f in docs if re.search(pat, texts[f])]
        C.ok("SS", len(where) == 1, f"{name} stated once, in {where}")
    # social layer quotes no figure; no em dash in any document
    thread = texts[WDOC.THREAD]
    C.ok("SS", "$" not in thread and not re.search(r"\d{1,3},\d{3}", thread), "no figure quoted in the thread")
    C.ok("SS", not any("\u2014" in texts[f] for f in docs), "no em dash in any document")
    # generation tells
    totals = [sum(r["offer"] for r in W["corpus"][c]["rows"]) for c in MARCH_CENSUSES]
    sep_total = sum(r["offer"] for r in W["sept"]["rows"])
    C.ok("TELL", not any(t % 1000 == 0 for t in totals + [sep_total]), f"no offer total on a round thousand {totals + [sep_total]}")
    counts = [len(W["corpus"][c]["rows"]) for c in MARCH_CENSUSES]
    C.ok("TELL", len(set(counts)) == 6, "pack row counts all differ")
    half = [(c.year, r["org"]) for c in MARCH_CENSUSES for r in W["corpus"][c]["rows"]
            if r["offer"] and (W["corpus"][c]["rate"] * r["fall"]) % 10000 == 5000]
    half += [(2026, r["org"]) for r in W["sept"]["rows"] if r["offer"] and (W["sept"]["rate"] * r["fall"]) % 10000 == 5000]
    C.ok("TELL", not half, "no offer lands exactly on a half dollar, so every rounding convention agrees")
    # the extract record's row counts collide with no graded figure (a count that equals a screen
    # figure reads as a leaked answer however unrelated the two are)
    G = W["golden_asks"]
    graded = {abs(v) for r in W["sept"]["rows"] for v in (r["cur"], r["prior"], r["fall"], r["offer"],
                                                          G[r["org"]]["K1"], G[r["org"]]["K2"])
              if v is not None}
    graded |= {sep_total, W["sept"]["rate"], len(W["sept"]["rows"]), W["sept"]["n_offers"]}
    stated = {k: info[k] for k in ("spine", "register", "payrun", "ratings")}
    clash = {k: v for k, v in stated.items() if v in graded}
    C.ok("TELL", not clash, f"no row count in the extract record equals a graded figure ({stated}; clash {clash})")
    # container hygiene: the house audit finds no writer signature and no out-of-band date
    import subprocess
    r = subprocess.run([sys.executable, SCRUB, target, "--floor", "2018-07-01", "--ceiling", "2026-10-08"],
                       capture_output=True, text=True)
    C.ok("H1", r.returncode == 0 and "clean" in r.stdout, f"scrub audit: {r.stdout.strip()[:120]}")


def write_metadata(out_dir, files, info):
    meta = {
        "task": "task123",
        "title": "Steady Ground Fund, September 2026 stabilisation offers",
        "domain": "Nonprofit & Grant-making",
        "subdomain": "grantee financial health",
        "objective": "Data Extraction & Conformation (ETL)",
        "as_of": "2026-10-08",
        "deliverables": ["steady_ground_sep2026_offers.docx", "steady_ground_sep2026_screen.csv",
                         "steady_ground_sep2026_offers.png"],
        "input_files": [{"path": f, "format": os.path.splitext(f)[1].lstrip(".")} for f in files],
        "large_file": {"path": WD.SPINE, "rows": info["spine"]},
        "distractor_files": DISTRACTORS,
    }
    with open(os.path.join(out_dir, "metadata.json"), "w", encoding="utf-8", newline="\n") as f:
        json.dump(meta, f, indent=2)
        f.write("\n")


def record(W, C, info, path):
    T = W["sept"]
    G = W["golden_asks"]
    rows = []
    for r in T["rows"]:
        o = W["by"][r["org"]]
        rows.append({"key": r["org"], "cc": o.cc, "name": o.name, "cur": r["cur"], "prior": r["prior"],
                     "fall": r["fall"], "pct": round(r["pct"], 4), "offer": r["offer"],
                     "K1": G[r["org"]]["K1"], "K2": G[r["org"]]["K2"]})
    S = W["S"]
    rec = {
        "rate_hc": T["rate"], "total": T["total"], "scored": len(T["rows"]),
        "offers": [x for x in rows if x["offer"]], "rows": rows,
        "rungs": {k: {"rate_hc": S[k]["rate"], "offered": sorted(W["by"][o].name for o in
                                                                 {r["org"] for r in S[k]["rows"] if r["offer"]}),
                      "n_rows": len(S[k]["rows"])} for k in ("R0", "R1", "R2", "R3", "R4", "R5", "R6")},
        "cells": {k: S[k]["rate"] for k in S if k not in ("T", "R0", "R1", "R2", "R3", "R4", "R5", "R6")},
        "gaps": {o.key: [[str(x) for x in g] for g in o.gaps] for o in W["orgs"] if o.gaps},
        "cofund": {o.key: o.cofund for o in W["orgs"] if o.cofund},
        "movers": {o.key: o.name for o in W["orgs"] if o.cal_change},
        "corpus": {c.year: {"rows": len(W["corpus"][c]["rows"]), "offers": W["corpus"][c]["n_offers"],
                            "rate_hc": W["corpus"][c]["rate"],
                            "total": sum(r["offer"] for r in W["corpus"][c]["rows"])} for c in MARCH_CENSUSES},
        "rivals": {k: {"rows": v["rows"], "offers_fail": v["off"], "rates_fail": v["rate"]}
                   for k, v in W["rival_misses"].items()},
        "lazy": W["lazy"], "pair": W["pair"], "composed_worst": list(W["composed_worst"][:3]),
        "separation": W["separation"], "info": info, "assertions": C.n, "ids": dict(C.ids),
        "keys": {o.key: o.name for o in W["orgs"]},
    }
    with open(path, "w") as f:
        json.dump(rec, f, indent=1, default=str)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", required=True, help="task folder to build into (target/ and metadata.json)")
    ap.add_argument("--record", help="where to write the build record (JSON)")
    ap.add_argument("--log", help="where to write every assertion line")
    a = ap.parse_args()
    W = tuned_world()
    C = Checks()
    check_world(W, C)
    corpus_checks(W, C)
    flips_check(W, C)
    convergence_checks(W, C)
    clean_data_checks(W, C)
    check_asks(W, C, W["S"])
    target = os.path.join(a.out, "target")
    info = write_pack(W, target)
    separation(W, C, W["portal_rows"])
    check_files(W, C, target, info)
    files = sorted(os.listdir(target))
    write_metadata(a.out, files, info)
    C.ok("A51", json.load(open(os.path.join(a.out, "metadata.json")))["distractor_files"] == DISTRACTORS,
         "metadata.json names the two distractors")
    if a.record:
        record(W, C, info, a.record)
    if a.log:
        with open(a.log, "w") as fh:
            fh.write("\n".join(C.lines) + "\n")
    print(f"task123 generator: {C.n} assertions green across {len(C.ids)} assertion ids")
    print(f"September rate {W['sept']['rate']/100:.2f} cents, {W['sept']['n_offers']} offers, "
          f"{len(W['sept']['rows'])} scored; pack {len(files)} files, spine {info['spine']:,} rows")


if __name__ == "__main__":
    main()
