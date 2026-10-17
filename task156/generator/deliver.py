"""The goldens, from the shipped files only: moderation_pool_2027.xlsx (each site's 2027 need, the request and its
distance from finance's envelope, the three tables, the reviewer pools behind the need and the chart's monthly
figures), council_pool_request_2027.pptx (the request slide and the monthly hours chart), submission.md and
generator/figures.json (every figure the write-up states, as it states it)."""
from __future__ import annotations

import io
import json
import os

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import calc as C
import common as K
import postprocess as PP

XLSX = "moderation_pool_2027.xlsx"
PPTX = "council_pool_request_2027.pptx"
AUTHOR = "Moderation Operations"
CREATED = "2026-10-13T14:10:00Z"
MODIFIED = "2026-10-14T20:36:00Z"
NAME = {s: K.SITE[s]["name"] for s in K.SITES}
MONTHS_ALL = C.HIST_MONTHS + C.FWD_MONTHS                  # 2024-01 .. 2027-12
SAT_CUT = 5                                                 # per cent routed: the inbox outran the volunteers
ENV_MONTH = K.ENVELOPE_2027 * K.FTE_HOURS / 12.0


def fmt(x, dp=0):
    v = K.round_half_up(x, dp)
    return format(v, ",.%df" % dp) if dp else format(int(v), ",")


def pct(x):
    return fmt(x, 1)


# ------------------------------------------------------------------------------------------------- the figures
def compute(target):
    X = C.compute(target)
    D = X.D
    M = C.MIDX
    lo, hi = M[(2024, 9)], M[(2026, 8)]
    ex = sum(D.exits[s][lo:hi + 1].sum() for s in K.SITES)
    risk = sum(D.R[s][lo:hi + 1].sum() for s in K.SITES)
    X.exit_net = ex / risk
    X.exit_site = {s: C.exit_rate(D, s)[0] for s in K.SITES}
    X.arr = {s: C.arrivals(D, s) for s in K.SITES}
    X.joins_2024 = int(sum(D.joins[s][M[(2024, 1)]:M[(2024, 12)] + 1].sum() for s in K.SITES))
    X.joins_since = int(sum(D.joins[s][M[(2025, 1)]:M[(2026, 9)] + 1].sum() for s in K.SITES))
    X.R_sep26 = {s: int(D.R[s][M[(2026, 9)]]) for s in K.SITES}
    X.pool = {s: C.project_pool(D, s) for s in K.SITES}
    idx27 = [k for k, (y, m) in enumerate(C.FWD_MONTHS) if y == 2027]
    X.idx27 = idx27
    X.plan27, X.flags27, X.inbox27, X.cap27, X.routed27, X.short = {}, {}, {}, {}, {}, {}
    for s in K.SITES:
        items, flag = C.ratios(D, s)
        per_q = sum(items.values())
        plan_fwd = np.array([D.outlook[s][m - 10] if y == 2026 else D.plan[s][m - 1] for (y, m) in C.FWD_MONTHS])
        A = plan_fwd * per_q
        cap = X.pool[s] * D.pace
        share = np.maximum(D.floor, 1 - cap / A)
        X.plan27[s] = D.plan[s].sum()
        X.flags27[s] = flag * D.plan[s].sum()
        X.inbox27[s] = A[idx27].sum()
        X.cap27[s] = cap[idx27].sum()
        X.routed27[s] = (A * share)[idx27].sum()
        X.short[s] = [C.FWD_MONTHS[k] for k in range(len(C.FWD_MONTHS)) if cap[k] < A[k]]
    # the sites by where their volunteer cover stands: short already in 2026, short from the forecast, covered
    w26 = slice(M[(2026, 1)], M[(2026, 9)] + 1)
    share26 = {s: D.rout_all[s][w26].sum() / D.ent_all[s][w26].sum() for s in K.SITES}
    X.already = [s for s in K.SITES if share26[s] > 2 * D.floor]
    X.newly = [s for s in K.SITES if s not in X.already and any(y == 2027 for y, m in X.short[s])]
    X.covered = [s for s in K.SITES if not X.short[s]]
    assert sorted(X.already + X.newly + X.covered) == sorted(K.SITES)
    X.newly_from = min(min(X.short[s]) for s in X.newly)
    X.share26 = share26
    X.monthly = {s: np.concatenate([X.hist_hours[s], X.fwd_hours[s]]) for s in K.SITES}
    return X


def names(codes):
    n = [NAME[s] for s in codes]
    return n[0] if len(n) == 1 else ", ".join(n[:-1]) + " and " + n[-1]


# ------------------------------------------------------------------------------------------------- the workbook
THIN = Side(style="thin", color="A3ABB5")
HEAD = PatternFill("solid", fgColor="DCE3EA")
TOT = PatternFill("solid", fgColor="F1F4F7")


def _head(ws, row, cols, widths=None, height=None):
    for c, v in enumerate(cols, 1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.font = Font(bold=True)
        cell.fill = HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.border = Border(bottom=THIN)
    if widths:
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = w
    if height:
        ws.row_dimensions[row].height = height


def _title(ws, title, lines, merge_to):
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=13)
    for k, t in enumerate(lines):
        c = ws.cell(row=2 + k, column=1, value=t)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells("A%d:%s%d" % (2 + k, merge_to, 2 + k))
        ws.row_dimensions[2 + k].height = 15 * (1 + len(t) // 120)


def write_xlsx(path, X):
    D = X.D
    wb = Workbook()
    ws = wb.active
    ws.title = "Need 2027"
    _title(ws, "The nine sites need %d paid moderator FTE in 2027, %d above finance's envelope" % (X.request, X.gap), [
        "Need as the charter defines it (section 4): the hours the paid team will spend in 2027 on each site's moderator "
        "flags and on the community review items no volunteer has reviewed 72 hours after they entered the inbox, each "
        "at the 2027 handle-time standard. One FTE is 1,600 hours; each site to the nearest whole FTE; the request is "
        "the sum of the nine.",
        "Exports pulled 8 October 2026, data through 30 September 2026; questions from the 2027 question plan issued "
        "18 September 2026. Method and sources on the Notes sheet."], "H")
    _head(ws, 5, ["Site", "Questions planned, 2027", "Moderator flags handled, 2027", "Review inbox items, 2027",
                  "Volunteer review capacity, 2027", "Items reaching the paid team, 2027", "Paid review hours, 2027",
                  "Need (FTE)"], [24, 13, 13, 13, 14, 14, 13, 9], 45)
    r = 6
    for s in K.SITES:
        vals = [NAME[s], round(X.plan27[s]), round(X.flags27[s]), round(X.inbox27[s]), round(X.cap27[s]),
                round(X.routed27[s]), round(X.need_h[s]), X.need_round[s]]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            if 2 <= c <= 7:
                cell.number_format = "#,##0"
        r += 1
    tot = ["Network", round(sum(X.plan27.values())), round(sum(X.flags27.values())), round(sum(X.inbox27.values())),
           round(sum(X.cap27.values())), round(sum(X.routed27.values())), round(sum(X.need_h.values())), X.request]
    for c, v in enumerate(tot, 1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = Font(bold=True)
        cell.fill = TOT
        cell.border = Border(top=THIN)
        if 2 <= c <= 7:
            cell.number_format = "#,##0"
    r += 2
    for k, v, bold in (("Council request for 2027 (FTE), the sum of the nine site needs", X.request, True),
                       ("Finance envelope for 2027 (FTE)", K.ENVELOPE_2027, False),
                       ("Request above the envelope (FTE)", X.gap, True),
                       ("Pool funded for 2026 (FTE)", K.POOL_2026, False)):
        a = ws.cell(row=r, column=1, value=k)
        ws.merge_cells("A%d:G%d" % (r, r))
        b = ws.cell(row=r, column=8, value=v)
        a.font = Font(bold=bold)
        b.font = Font(bold=bold)
        r += 1
    ws.freeze_panes = "B6"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    w2 = wb.create_sheet("Reviewer pools")
    _title(w2, "Every site's reviewer pool is losing about %d per cent a month, with almost no new reviewers" % (
        K.round_half_up(100 * X.exit_net)), [
        "Active reviewers made at least one review in the month (reviewer roster). Each pool is carried forward month "
        "by month at its own lapse rate, with the new reviewers it has been getting; every reviewer at %d reviews a "
        "month, the pace reviewers keep when the inbox outruns them." % K.round_half_up(D.pace)], "F")
    _head(w2, 4, ["Site", "Active reviewers, September 2026", "Lapsing a month, September 2024 to August 2026",
                  "New reviewers a month, October 2025 to September 2026", "Active reviewers, December 2027",
                  "Months short of the inbox, October 2026 to December 2027"], [24, 14, 16, 17, 14, 17], 58)
    r = 5
    for s in K.SITES:
        vals = [NAME[s], X.R_sep26[s], X.exit_site[s], round(X.arr[s], 2), round(X.pool[s][-1]), len(X.short[s])]
        for c, v in enumerate(vals, 1):
            cell = w2.cell(row=r, column=c, value=v)
            if c == 3:
                cell.number_format = "0.00%"
            elif c == 4:
                cell.number_format = "0.00"
            elif c in (2, 5):
                cell.number_format = "#,##0"
        r += 1
    tot = ["Network", sum(X.R_sep26.values()), X.exit_net, round(sum(X.arr.values()), 2),
           round(sum(X.pool[s][-1] for s in K.SITES)), None]
    for c, v in enumerate(tot, 1):
        cell = w2.cell(row=r, column=c, value=v)
        cell.font = Font(bold=True)
        cell.fill = TOT
        cell.border = Border(top=THIN)
        cell.number_format = {2: "#,##0", 3: "0.00%", 4: "0.00", 5: "#,##0"}.get(c, "General")
    w2.cell(row=r + 2, column=1, value="New reviewers across the network: %s in 2024, %s from January 2025 to September "
            "2026." % (fmt(X.joins_2024), fmt(X.joins_since)))
    w2.merge_cells("A%d:F%d" % (r + 2, r + 2))
    w2.freeze_panes = "B5"

    w3 = wb.create_sheet("Flag response")
    slow = [s for s in K.SITES if all(X.ask_a[(s, m)] < 95.0 for m in (7, 8, 9))]
    _title(w3, "Moderator flags answered within 24 hours, July to September 2026: %s answered fewer" % names(slow), [
        "Share of the flags raised in the month (UTC) that a paid moderator first acted on within 24 hours of the flag "
        "being raised, by the flag service standard. Flags the platform closed because the post was deleted are left "
        "out. First action from the decision-review log where a decision was reviewed, otherwise the moderator tool."],
        "D")
    _head(w3, 4, ["Site", "July 2026", "August 2026", "September 2026"], [24, 12, 12, 14])
    for k, s in enumerate(K.SITES):
        w3.cell(row=5 + k, column=1, value=NAME[s])
        for j, m in enumerate((7, 8, 9)):
            c = w3.cell(row=5 + k, column=2 + j, value=K.round_half_up(X.ask_a[(s, m)], 1))
            c.number_format = '0.0"%"'

    w4 = wb.create_sheet("Hours worked")
    grew = [s for s in K.SITES if X.ask_b[(s, 3)] > 1.3 * X.ask_b[(s, 2)]]
    _title(w4, "Hours worked by each site's paid moderators, 2026: %s rose in the third quarter, when six moderators "
               "moved to them" % names(grew), [
        "Approved WFM shift lines (a corrected week's superseded lines left out). Each shift to the site the moderator "
        "worked for on that shift, from the HR assignment history, and to the quarter in which it started, on Toronto "
        "time."], "D")
    _head(w4, 4, ["Site", "Q1 2026", "Q2 2026", "Q3 2026"], [24, 12, 12, 12])
    for k, s in enumerate(K.SITES):
        w4.cell(row=5 + k, column=1, value=NAME[s])
        for j, q in enumerate((1, 2, 3)):
            c = w4.cell(row=5 + k, column=2 + j, value=int(K.round_half_up(X.ask_b[(s, q)])))
            c.number_format = "#,##0"

    w5 = wb.create_sheet("Questions vs 2022")
    order = sorted(K.SITES, key=lambda s: X.ask_c[s])
    assert all(35 < X.ask_c[s] < 45 for s in order[:2]) and all(70 < X.ask_c[s] < 80 for s in order[-2:])
    _title(w5, "Questions asked January to September 2026 as a percentage of the same months of 2022: %s kept about two "
               "fifths, %s about three quarters" % (names(order[:2]), names(order[-2:])), [
        "Each question once, in the month it was first asked (UTC), on the site it is on today, as the posts warehouse "
        "counts. 2026 from the warehouse. 2022 from the old review system's workbook, which counted a question moved "
        "between sites at both sites: each move in the migration log puts the question on the site it moved to, in the "
        "month it was asked. Graphics & Shaders questions are in Game Development."], "B")
    _head(w5, 4, ["Site", "January to September 2026 as % of 2022"], [24, 20], 30)
    for k, s in enumerate(K.SITES):
        w5.cell(row=5 + k, column=1, value=NAME[s])
        c = w5.cell(row=5 + k, column=2, value=K.round_half_up(X.ask_c[s], 1))
        c.number_format = '0.0"%"'

    w6 = wb.create_sheet("Monthly hours")
    _title(w6, "Paid review hours by month and site, January 2024 to December 2027 (the deck's chart)", [
        "Moderator flags handled and community items reaching the paid team, each at the handle-time standard in force "
        "(close votes 5.0 minutes to December 2026, 2.0 from January 2027). Actual to September 2026, forecast from "
        "October 2026."], "M")
    _head(w6, 4, ["Month", "Basis"] + [NAME[s] for s in K.SITES] + ["Network", "2027 envelope"],
          [10, 9] + [11] * 9 + [10, 10], 45)
    for i, (y, m) in enumerate(MONTHS_ALL):
        row = 5 + i
        w6.cell(row=row, column=1, value="%s %d" % (K.MON3[m], y))
        w6.cell(row=row, column=2, value="Actual" if (y, m) <= (2026, 9) else "Forecast")
        tot = 0
        for j, s in enumerate(K.SITES):
            v = int(K.round_half_up(X.monthly[s][i]))
            tot += X.monthly[s][i]
            w6.cell(row=row, column=3 + j, value=v).number_format = "#,##0"
        w6.cell(row=row, column=12, value=int(K.round_half_up(tot))).number_format = "#,##0"
        w6.cell(row=row, column=13, value=int(K.round_half_up(ENV_MONTH))).number_format = "#,##0"
    w6.freeze_panes = "C5"

    w7 = wb.create_sheet("Notes")
    w7.column_dimensions["A"].width = 24
    w7.column_dimensions["B"].width = 112
    notes = [
        ("Need", "The charter's definition (section 4): each site's moderator flags a paid moderator handles and its "
                 "community review items that reach the paid team in 2027, each at the 2027 handle-time standard, in "
                 "FTE of 1,600 hours, each site to the nearest whole FTE. The request is the sum of the nine site needs."),
        ("How items reach the team", "Every community item goes first to the site's volunteer reviewers and reaches "
                                     "the paid team when no volunteer has reviewed it 72 hours after it entered the "
                                     "inbox (charter, section 3). So a site's routed items are what its inbox brings "
                                     "beyond what its reviewers can review, not a fixed share of its inbox."),
        ("Questions", "The 2027 plan by site and month (question_plan_2027.xlsx, Plan 2027). October to December 2026 "
                      "from the plan's Outlook sheet, used only for the chart."),
        ("Inbox items and flags", "Each site's review inbox items by type and its handled flags per question asked, "
                                  "January 2024 to September 2026 (the inbox export's 8 October pull, the moderator "
                                  "tool's daily summary and the posts warehouse), times its planned questions."),
        ("Reviewer pace", "About %d reviews a month per active reviewer: volunteer reviews over active reviewers in "
                          "the %d site-months, 2021 to September 2026, in which more than %d per cent of the items "
                          "entering the inbox reached the paid team (the old review system's workbook to 2023; the "
                          "inbox export and the reviewer roster from 2024). When the inbox runs dry, reviewers review "
                          "only what arrives, so those months show the work, not the pace." % (
                              K.round_half_up(D.pace), D.n_sat, SAT_CUT)),
        ("Reviewer pools", "Each site's active reviewers in September 2026, carried forward month by month less the "
                           "share that lapses each month (lapses over active reviewer-months, September 2024 to August "
                           "2026, the last full month of lapses the roster shows; about %d per cent on every site) plus "
                           "the new reviewers a month of the last twelve months. Almost no member has been granted "
                           "review access since late 2024." % K.round_half_up(100 * X.exit_net)),
        ("Items reaching the team", "Each month, the inbox items beyond the pool's capacity. Where a pool has room to "
                                    "spare, the %.1f per cent of items no volunteer picks up (the routed share in "
                                    "site-months with fewer than one item in 60 routed), as now on %s." % (
                                        100 * D.floor, names(X.covered))),
        ("Pricing", "2027 standards: flag 9.0 minutes, close vote 2.0, first post 4.0, low quality 3.5, suggested edit "
                    "2.0, each routed item at its site's mix of inbox items. The chart's history and its October to "
                    "December 2026 months at the standards in force to December 2026 (close vote 5.0)."),
        ("Where cover runs out", "%s have sent the paid team only the items no volunteer picks up since January 2024; at the "
                                 "pace and lapse rate above, their pools fall short of the inbox from %s %d and in "
                                 "most months of 2027. %s are short already. %s stay covered through 2027." % (
                                     names(X.newly), K.MONTH_NAMES[X.newly_from[1]], X.newly_from[0],
                                     names(X.already), names(X.covered))),
        ("Flag response", "The flag service standard: flags raised in the month on UTC, a paid moderator's first "
                          "action within 24 hours, flags closed because the post was deleted left out. The flag "
                          "export's handled time is when the flag was last marked handled, which a decision review "
                          "moves; the decision-review log keeps the first action. Flags raised on 15 August are in "
                          "both parts of the export and are counted once."),
        ("Hours worked", "Approved WFM lines only. The WFM export shows each moderator's current site on every line; "
                         "six moderators changed site in July and August 2026, so each shift goes to the site the HR "
                         "assignment history has for its start day. Infrastructure & DevOps appears in WFM as Ops & "
                         "Cloud."),
        ("Questions vs 2022", "The old workbook counted a question moved between sites at its source in the month it "
                              "was asked and again at its destination in the month it moved. The migration log puts "
                              "each moved question on its destination in the month it was asked, as the warehouse "
                              "counts today."),
    ]
    _head(w7, 1, ["Item", "Note"])
    for i, (k, v) in enumerate(notes, 2):
        w7.cell(row=i, column=1, value=k).font = Font(bold=True)
        c = w7.cell(row=i, column=2, value=v)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        w7.cell(row=i, column=1).alignment = Alignment(vertical="top")
        w7.row_dimensions[i].height = 15 * (1 + len(v) // 108)
    wb.save(path)
    PP.normalize_ooxml(path, AUTHOR, AUTHOR, CREATED, MODIFIED, excel=True)


# ------------------------------------------------------------------------------------------------- the chart
# categorical slots in fixed order (validated adjacent, light surface); the ninth band repeats slot 1 with a texture
SLOTS = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300", "#6250d6", "#e34948", "#2a78d6"]
INK, INK2, MUTED, GRID, AXIS = "#0b0b0b", "#52514e", "#898781", "#e1e0d9", "#c3c2b7"


def chart_png(X):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    from matplotlib.patches import Patch
    for f in ("/usr/share/fonts/truetype/crosextra/Carlito-Regular.ttf",
              "/usr/share/fonts/truetype/crosextra/Carlito-Bold.ttf"):
        font_manager.fontManager.addfont(f)
    plt.rcParams["font.family"] = "Carlito"
    plt.rcParams["hatch.linewidth"] = 1.1
    plt.rcParams["svg.hashsalt"] = "task156"
    n = len(MONTHS_ALL)
    xs = np.arange(n)
    ys = [X.monthly[s] for s in K.SITES]
    fig = plt.figure(figsize=(12.4, 5.15), dpi=150)
    ax = fig.add_axes([0.065, 0.13, 0.74, 0.83])
    k0 = MONTHS_ALL.index((2026, 10))
    ax.axvspan(k0 - 0.5, n - 0.5, color="#f2f1ec", zorder=0, lw=0)
    polys = ax.stackplot(xs, ys, colors=SLOTS, edgecolor="white", linewidth=0.7, zorder=2)
    polys[-1].set_hatch("////")
    top = sum(ys)
    ymax = 1000 * np.ceil(top.max() / 1000.0 + 1.0)
    ax.axhline(ENV_MONTH, color=INK, lw=1.5, ls=(0, (5, 3)), zorder=4)
    ax.text(0.3, ENV_MONTH + ymax * 0.012, "2027 envelope: %d FTE, %s hours a month" % (K.ENVELOPE_2027,
            fmt(ENV_MONTH)), fontsize=10.5, color=INK, va="bottom", ha="left", zorder=5)
    i27 = MONTHS_ALL.index((2027, 1))
    yb = ymax * 0.905
    ax.plot([i27 - 0.4, n - 0.6], [yb, yb], color=INK2, lw=1.0, zorder=5)
    for xx in (i27 - 0.4, n - 0.6):
        ax.plot([xx, xx], [yb, yb - ymax * 0.02], color=INK2, lw=1.0, zorder=5)
    ax.text((i27 - 0.4 + n - 0.6) / 2, yb + ymax * 0.012, "2027: %d FTE of need" % X.request, ha="center",
            va="bottom", fontsize=10.5, color=INK, fontweight="bold", zorder=6)
    ax.text(k0 - 0.2, ymax * 0.975, "Actual", fontsize=10, color=INK2, ha="right", va="top")
    ax.text(k0 + 0.2, ymax * 0.975, "Forecast", fontsize=10, color=INK2, ha="left", va="top")
    ax.set_xlim(-0.5, n - 0.5)
    ax.set_ylim(0, ymax)
    ticks = [k for k, (y, m) in enumerate(MONTHS_ALL) if m in (1, 4, 7, 10)]
    ax.set_xticks(ticks)
    ax.set_xticklabels(["%s\n%d" % (K.MON3[m], y) if m == 1 else K.MON3[m] for (y, m) in
                        [MONTHS_ALL[k] for k in ticks]], fontsize=9.5, color=INK2)
    ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, p: format(int(v), ",")))
    ax.tick_params(axis="y", labelsize=9.5, colors=INK2, length=0)
    ax.tick_params(axis="x", length=3, color=AXIS)
    ax.set_ylabel("Paid review hours a month", fontsize=10.5, color=INK2)
    ax.grid(axis="y", color=GRID, lw=0.7, zorder=1)
    ax.set_axisbelow(True)
    for sp in ("top", "right", "left"):
        ax.spines[sp].set_visible(False)
    ax.spines["bottom"].set_color(AXIS)
    handles = []
    for s, col in reversed(list(zip(K.SITES, SLOTS))):
        h = Patch(facecolor=col, edgecolor="white", hatch="////" if s == K.SITES[-1] else None, label=NAME[s])
        handles.append(h)
    ax.legend(handles=handles, frameon=False, fontsize=10, loc="upper left", bbox_to_anchor=(1.01, 1.0),
              title="Site (top band first)", title_fontsize=10, alignment="left", handlelength=1.6,
              handleheight=1.1, labelspacing=0.55)
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150, metadata={"Software": None})
    plt.close(fig)
    return buf.getvalue()


# ------------------------------------------------------------------------------------------------- the deck
def write_pptx(path, X, png):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    ink = RGBColor(0x1F, 0x23, 0x29)
    grey = RGBColor(0x52, 0x51, 0x4E)
    council = RGBColor(0x3B, 0x4A, 0x6B)
    pale = RGBColor(0xEE, 0xF0, 0xF3)
    white = RGBColor(0xFF, 0xFF, 0xFF)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    def text(sl, x, y, w, h, paras, align=None):
        tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        for k, (t, size, bold, col) in enumerate(paras):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.text = t
            p.font.size, p.font.bold = Pt(size), bold
            p.font.color.rgb = col
            p.font.name = "Calibri"
            if align is not None:
                p.alignment = align
            if k:
                p.space_before = Pt(8)
        return tb

    def footer(sl, n):
        text(sl, 0.6, 7.02, 10.0, 0.3, [("Millrind moderation operations  |  2027 paid moderation pool request  |  "
                                         "October 2026", 9.5, False, grey)])
        text(sl, 11.9, 7.02, 0.8, 0.3, [(str(n), 9.5, False, grey)], align=PP_ALIGN.RIGHT)

    # slide 1: the request
    s1 = prs.slides.add_slide(prs.slide_layouts[6])
    text(s1, 0.6, 0.4, 12.1, 1.3, [
        ("Request %d paid moderator FTE for 2027" % X.request, 32, True, ink),
        ("%d FTE above finance's 2027 envelope of %d FTE (%d FTE funded for 2026)" % (
            X.gap, K.ENVELOPE_2027, K.POOL_2026), 18, False, grey)])
    rows = len(K.SITES) + 2
    shape = s1.shapes.add_table(rows, 2, Inches(0.6), Inches(1.95), Inches(4.6), Inches(0.37 * rows))
    tbl = shape.table
    tbl.first_row, tbl.horz_banding = True, False
    tbl.columns[0].width, tbl.columns[1].width = Inches(3.3), Inches(1.3)
    cells = [("Site", "2027 need (FTE)")] + [(NAME[s], str(X.need_round[s])) for s in K.SITES] + \
            [("Request", str(X.request))]
    for i, (a, b) in enumerate(cells):
        for j, v in enumerate((a, b)):
            c = tbl.cell(i, j)
            c.text = v
            c.fill.solid()
            c.fill.fore_color.rgb = council if i == 0 else (pale if i == rows - 1 else white)
            p = c.text_frame.paragraphs[0]
            p.font.size = Pt(13)
            p.font.name = "Calibri"
            p.font.bold = i in (0, rows - 1)
            p.font.color.rgb = white if i == 0 else ink
            if j == 1:
                p.alignment = PP_ALIGN.RIGHT
    paras = [
        ("Why the need rises", 17, True, ink),
        ("Every community review item goes to the volunteer reviewers first and reaches the paid team only when no "
         "volunteer has reviewed it in 72 hours. A reviewer clears about %d items a month when the inbox outruns "
         "the pool." % K.round_half_up(X.D.pace), 14, False, ink),
        ("Each site's pool loses about %d per cent of its reviewers a month, and almost no member has been granted "
         "review access since late 2024 (%s new reviewers in 2024, %s since)." % (
             K.round_half_up(100 * X.exit_net), fmt(X.joins_2024), fmt(X.joins_since)), 14, False, ink),
        ("%s, which have sent the paid team almost nothing since January 2024, run short of reviewers from %s %d. %s "
         "are short already; %s stay covered." % (names(X.newly), K.MONTH_NAMES[X.newly_from[1]], X.newly_from[0],
                                                 names(X.already), names(X.covered)), 14, False, ink),
        ("Priced at the 2027 standards (close votes at 2.0 minutes, down from 5.0), the nine sites need %d FTE." %
         X.request, 14, False, ink),
    ]
    text(s1, 5.75, 1.9, 7.0, 4.6, paras)
    text(s1, 0.6, 6.42, 12.1, 0.5, [(
        "Need as the council charter defines it (section 4), on the 2027 question plan of 18 September 2026; one FTE "
        "is 1,600 hours. Data through 30 September 2026, pulled 8 October 2026. Workings: %s." % XLSX,
        10.5, False, grey)])
    footer(s1, 1)

    # slide 2: the chart
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    text(s2, 0.6, 0.3, 12.1, 1.2, [
        ("2027 needs %d FTE of paid review hours against an envelope of %d" % (X.request, K.ENVELOPE_2027), 26, True,
         ink),
        ("Paid review hours a month by site, at the handle-time standards in force (close votes at 2.0 minutes from "
         "January 2027, 5.0 before). The rise through 2027 is volunteer cover running out on more sites.", 14, False,
         grey)])
    s2.shapes.add_picture(io.BytesIO(png), Inches(0.55), Inches(1.5), width=Inches(12.2))
    text(s2, 0.6, 6.6, 12.1, 0.4, [(
        "Moderator flags handled and community items reaching the paid team. Actual to September 2026; forecast from "
        "October 2026 on the question plan and each site's reviewer pool. Monthly figures: %s, Monthly hours." % XLSX,
        10.5, False, grey)])
    footer(s2, 2)
    buf = io.BytesIO()
    prs.save(buf)
    with open(path, "wb") as fh:
        fh.write(buf.getvalue())
    PP.normalize_ooxml(path, AUTHOR, AUTHOR, CREATED, MODIFIED, title="2027 paid moderation pool request")


# ------------------------------------------------------------------------------------------------- the write-up
def site_list(vals):
    return ", ".join("%s %d" % (NAME[s], vals[s]) for s in K.SITES[:-1]) + " and %s %d" % (
        NAME[K.SITES[-1]], vals[K.SITES[-1]])


def submission(X):
    D = X.D
    pace = int(K.round_half_up(D.pace))
    exit_pc = int(K.round_half_up(100 * X.exit_net))
    s = []
    s += ["## Tags", "",
          "**Domain:** Product Analytics (retention, cohort and lifetime value modeling: each site's volunteer reviewer "
          "pool carried forward on its own monthly lapse rate, and the reviews a retained reviewer delivers when the "
          "inbox outruns the pool).",
          "**Analytical objective:** Forecasting & Predictive Modeling.", ""]
    s += ["## 1. Final Recommendation", "",
          "**The council should request %d paid moderator FTE for 2027, %d above finance's %d-FTE envelope: %s.**" % (
              X.request, X.gap, K.ENVELOPE_2027, site_list(X.need_round)),
          "Items reach the paid team only when the volunteers fall short, so each site's 2027 routed items are the "
          "inbox beyond what its shrinking reviewer pool can review at about %d reviews a reviewer a month, which puts "
          "%s short of reviewers after routing almost nothing since January 2024. Not the last twelve months' paid hours "
          "per question carried onto the plan, which prices only what those sites routed while their pools still "
          "covered the inbox; not a trend in each site's routed share or paid hours, which extends the past and misses "
          "the months the pools fall below the inbox; and not a reviewer pace read from 2026 throughput, which on the "
          "covered sites measures how little work arrived rather than how much reviewers can do." % (
              pace, names(X.newly)), ""]
    s += ["## 2. Critical Components", "",
          "1. A volunteer reviewer completes about **%d** reviews a month when the inbox outruns the site's pool, "
          "measured in the site-months in which more than %d per cent of entering items reached the paid team." % (
              pace, SAT_CUT),
          "2. Every site's reviewer pool loses about **%d** per cent of its active reviewers a month, and almost no new "
          "reviewers have arrived since late 2024." % exit_pc,
          "3. %s fall short of their inbox in 2027, joining %s; %s stay covered." % (
              names(X.newly), names(X.already), names(X.covered)), ""]
    s += ["## 3. Step-by-Step Solution", "",
          "1. Stacked review_queues_monthly_2021-2023.xlsx (old site names through site_register.csv) with "
          "review_inbox_daily_2024-01-01_2026-09-30.csv and the active reviewers in reviewer_roster_2026-10-08.csv: "
          "over the site-months in which more than %d per cent of entering items were routed to moderation, volunteer "
          "reviews over active reviewers give about %d reviews a month." % (SAT_CUT, pace),
          "2. In reviewer_roster_2026-10-08.csv, lapses over active reviewer-months from September 2024 to August 2026 "
          "give each site's monthly lapse rate, about %d per cent on every site, and the last twelve months' new "
          "reviewers give each site fewer than one a month." % exit_pc,
          "3. Carried each site's September 2026 active reviewers forward month by month at its lapse rate plus its new "
          "reviewers, at %d reviews a reviewer a month, against the inbox of each month in 2027 (the questions in "
          "question_plan_2027.xlsx times the site's inbox items per question since January 2024), and routed the items "
          "beyond that capacity to the paid team (where a pool has room, the %.1f per cent no volunteer picks up): %s "
          "fall short of their inbox." % (pace, 100 * D.floor, names(X.newly)),
          "4. Priced each site's handled flags per question since January 2024 "
          "(modtool_flag_summary_daily_2024-01-01_2026-09-30.csv) times the plan, and its routed items, at "
          "handle_time_standards_2027.pdf, and divided by the charter's 1,600 hours, each site to the nearest whole "
          "FTE: %d FTE in all." % X.request,
          "5. Against finance_envelope_2027_paid_moderation.pdf the request of %d FTE is %d above the %d-FTE envelope." %
          (X.request, X.gap, K.ENVELOPE_2027), ""]
    s += ["## 4. Deliverable Answers", "", "### %s" % XLSX, "",
          "1. 2027 need by site, whole FTE:"]
    for site in K.SITES:
        s.append("   - %s: %d" % (NAME[site], X.need_round[site]))
    s += ["2. Council request for 2027: %d FTE." % X.request,
          "3. Above finance's %d-FTE envelope by %d FTE." % (K.ENVELOPE_2027, X.gap),
          "4. Share of the moderator flags raised in each month that a moderator first acted on within 24 hours:"]
    for site in K.SITES:
        s.append("   - %s: July 2026 %s%%, August 2026 %s%%, September 2026 %s%%" % (
            NAME[site], pct(X.ask_a[(site, 7)]), pct(X.ask_a[(site, 8)]), pct(X.ask_a[(site, 9)])))
    s.append("5. Hours each site's paid moderators worked:")
    for site in K.SITES:
        s.append("   - %s: Q1 2026 %s hours, Q2 2026 %s hours, Q3 2026 %s hours" % (
            NAME[site], fmt(X.ask_b[(site, 1)]), fmt(X.ask_b[(site, 2)]), fmt(X.ask_b[(site, 3)])))
    s.append("6. January to September 2026 questions as a percentage of January to September 2022:")
    for site in K.SITES:
        s.append("   - %s: %s%%" % (NAME[site], pct(X.ask_c[site])))
    s += ["", "### %s" % PPTX, "",
          "7. Slide 1 states the request: %d paid moderator FTE for 2027, %d above the %d-FTE envelope, with each "
          "site's need." % (X.request, X.gap, K.ENVELOPE_2027),
          "8. Slide 2 charts the network's monthly paid review hours from January 2024 to December 2027, one stacked "
          "band per site, the forecast from October 2026 shaded, with finance's envelope (%d FTE, %s hours a month) "
          "drawn across." % (K.ENVELOPE_2027, fmt(ENV_MONTH)), ""]
    return "\n".join(s)


def figures(X):
    D = X.D
    return {
        "call": {"request_fte": X.request, "envelope_fte": K.ENVELOPE_2027, "above_envelope_fte": X.gap,
                 "need_fte": {NAME[s]: X.need_round[s] for s in K.SITES}},
        "components": {"reviews_a_month": int(K.round_half_up(D.pace)), "saturation_cut_pct": SAT_CUT,
                       "lapse_pct_a_month": int(K.round_half_up(100 * X.exit_net)),
                       "newly_short": [NAME[s] for s in X.newly], "already_short": [NAME[s] for s in X.already],
                       "covered": [NAME[s] for s in X.covered]},
        "steps": {"never_picked_up_pct": "%.1f" % (100 * D.floor), "fte_hours": int(K.FTE_HOURS),
                  "newly_short_from": "%s %d" % (K.MONTH_NAMES[X.newly_from[1]], X.newly_from[0])},
        "answers": {
            "need_fte": {NAME[s]: X.need_round[s] for s in K.SITES},
            "ask_a_pct": {"%s %s 2026" % (NAME[s], K.MONTH_NAMES[m]): pct(X.ask_a[(s, m)])
                          for s in K.SITES for m in (7, 8, 9)},
            "ask_b_hours": {"%s Q%d 2026" % (NAME[s], q): fmt(X.ask_b[(s, q)]) for s in K.SITES for q in (1, 2, 3)},
            "ask_c_pct": {NAME[s]: pct(X.ask_c[s]) for s in K.SITES},
            "chart_envelope_hours_a_month": fmt(ENV_MONTH),
        },
    }


def main(target, golden, fig_path, sub_path, B=None):
    os.makedirs(golden, exist_ok=True)
    X = compute(target)
    write_xlsx(os.path.join(golden, XLSX), X)
    png = chart_png(X)
    write_pptx(os.path.join(golden, PPTX), X, png)
    with open(sub_path, "w", encoding="utf-8") as fh:
        fh.write(submission(X))
    with open(fig_path, "w", encoding="utf-8") as fh:
        json.dump(figures(X), fh, indent=2)
        fh.write("\n")
    return X
