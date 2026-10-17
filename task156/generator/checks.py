"""task156 checks: the call and its corridor, the three rules executed on the generator's clean truth, the shipped HOW
sentences executed literally, every semantic fork row, the closure axes that can move, the ladder and its out-of-sample
record, the three asks with their stops and devices, the separation line, the realism sweep, the input gates and the
distractors. Run by build.py after the goldens are written; exits 1 on any failure.

    python3 task156/generator/checks.py"""
from __future__ import annotations

import csv
import datetime as dt
import json
import math
import os
import re
import shutil
import sys
import tempfile
import zipfile

import numpy as np
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import asks as AK           # noqa: E402
import build as BLD         # noqa: E402
import calc as C            # noqa: E402
import common as K          # noqa: E402
import deliver as DL        # noqa: E402
import docs as DOC          # noqa: E402
import records as R         # noqa: E402
import world as WD          # noqa: E402

TASK = os.path.dirname(HERE)
RESULTS = []
GOLD = {"web": 5, "mobile": 2, "data": 7, "databases": 5, "devops": 5, "gamedev": 2, "security": 4, "embedded": 3,
        "scicomp": 2}
RIVALS = {"naive12": 23, "naive9": 23, "share_trend": 25, "hours_trend": 28, "pool_trend_all": 31,
          "pace_censored": 53}
CONVERGENT_EXCLUDE = set(RIVALS) | {"pace_censored_trend_all"}
# the ask figures as built (a parameter change that moves one fails here first)
EXPECT = {
    "A": {"data 7": "97.8", "data 8": "97.3", "data 9": "97.4", "databases 7": "97.1", "databases 8": "97.5",
          "databases 9": "96.4", "devops 7": "98.0", "devops 8": "96.8", "devops 9": "96.7", "embedded 7": "91.4",
          "embedded 8": "87.7", "embedded 9": "87.1", "gamedev 7": "97.6", "gamedev 8": "96.8", "gamedev 9": "96.2",
          "mobile 7": "97.2", "mobile 8": "96.4", "mobile 9": "96.3", "scicomp 7": "85.6", "scicomp 8": "90.0",
          "scicomp 9": "81.4", "security 7": "91.8", "security 8": "89.3", "security 9": "88.1", "web 7": "97.9",
          "web 8": "97.0", "web 9": "97.1"},
    "B": {"data 1": 1725, "data 2": 1719, "data 3": 1509, "databases 1": 1643, "databases 2": 1631,
          "databases 3": 1280, "devops 1": 1624, "devops 2": 1616, "devops 3": 1352, "embedded 1": 898,
          "embedded 2": 870, "embedded 3": 1531, "gamedev 1": 751, "gamedev 2": 622, "gamedev 3": 535,
          "mobile 1": 879, "mobile 2": 827, "mobile 3": 540, "scicomp 1": 706, "scicomp 2": 683, "scicomp 3": 1382,
          "security 1": 1220, "security 2": 1387, "security 3": 2104, "web 1": 2475, "web 2": 2893, "web 3": 2260},
    "C": {"data": "55.5", "databases": "58.4", "devops": "60.2", "embedded": "76.0", "gamedev": "53.8",
          "mobile": "42.1", "scicomp": "76.8", "security": "72.0", "web": "39.7"},
}
MAIN_FILES = [R.F_INBOX, R.F_INBOX_PRE, R.F_OLD, R.F_ROSTER, R.F_FLAGSUM, R.F_POSTS, R.F_PLAN, R.F_REGISTER,
              DOC.F_CHARTER, DOC.F_STANDARDS, DOC.F_ENVELOPE]
ASK_FILES = [R.F_FLAGS1, R.F_FLAGS2, R.F_REVIEWS, R.F_WFM, R.F_ASSIGN, R.F_MOVES, DOC.F_FLAGSTD]


def ok(label, cond, detail=""):
    RESULTS.append((label, bool(cond), detail))


def rows(target, f):
    with open(os.path.join(target, f), newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def rnd(fte):
    return {s: int(K.round_half_up(v)) for s, v in fte.items()}


def req(fte):
    return sum(rnd(fte).values())


def nine(fte):
    return " ".join(str(v) for v in rnd(fte).values())


def pdf_text(path):
    import pdfplumber
    with pdfplumber.open(path) as pdf:
        return "\n".join(pg.extract_text() or "" for pg in pdf.pages)


def open_text(target, f):
    """Every text cell of a shipped workbook, joined."""
    wb = openpyxl.load_workbook(os.path.join(target, f), read_only=True)
    return "\n".join(str(c) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for c in row
                     if isinstance(c, str))


def need_of(D, **kw):
    """Each site's 2027 need in FTE on a loaded bundle, the golden construction with any of its parameters swapped."""
    pace = kw.pop("pace", None)
    floor = kw.pop("floor", None)
    out = {}
    for s in K.SITES:
        out[s] = C.fte(C.site_need(D, s, pace=pace if pace is not None else D.pace,
                                   floor=floor if floor is not None else D.floor, **kw)[1])
    return out


def loaded(target):
    D = C.load(target)
    D.pace, D.n_sat = C.saturated_pace(D)
    D.floor = C.floor_share(D)
    return D


def linked_copy(target, tmp, skip=(), replace=None):
    """A copy of the bundle in tmp: every file symlinked except those skipped or replaced (replace maps a file name to
    a writer taking the destination path)."""
    os.makedirs(tmp, exist_ok=True)
    for f in os.listdir(target):
        if f in skip:
            continue
        if replace and f in replace:
            replace[f](os.path.join(tmp, f))
        else:
            os.symlink(os.path.abspath(os.path.join(target, f)), os.path.join(tmp, f))
    return tmp


# ------------------------------------------------------------------------------------------------- ask variants
def ask_a_variant(target, first="review", drop_closed=True, dedupe=True, drop_reviewed=False):
    flags = []
    for f in (R.F_FLAGS1, R.F_FLAGS2):
        flags += rows(target, f)
    if dedupe:
        flags = list({r["flag_id"]: r for r in flags}.values())
    rev = {r["flag_id"]: r["first_action_at_utc"] for r in rows(target, R.F_REVIEWS)}
    num, den = {}, {}
    for r in flags:
        if drop_closed and r["outcome"] == "closed_post_deleted":
            continue
        if drop_reviewed and r["flag_id"] in rev:
            continue
        a = dt.datetime.fromisoformat(r["raised_at_utc"])
        b = dt.datetime.fromisoformat(rev[r["flag_id"]] if (first == "review" and r["flag_id"] in rev)
                                      else r["handled_at_utc"])
        k = (r["site"], a.month)
        den[k] = den.get(k, 0) + 1
        num[k] = num.get(k, 0) + (b - a <= dt.timedelta(hours=24))
    return {k: 100.0 * num[k] / den[k] for k in den}, num, den


def ask_b_variant(target, site_from="hr", keep_superseded=False, by="start", map_label=True):
    reg = {r["wfm_site_label"]: r["site_code"] for r in rows(target, R.F_REGISTER) if r["wfm_site_label"]}
    if not map_label:
        reg = {K.SITE[c]["name"]: c for c in K.SITES}
    name_to_code = {K.SITE[c]["name"]: c for c in K.SITES}
    hist = {}
    for r in rows(target, R.F_ASSIGN):
        hist.setdefault(r["employee_no"], []).append((r["effective_from"], r["effective_to"] or "9999-12-31",
                                                      name_to_code[r["site"]]))
    out = {}
    for r in rows(target, R.F_WFM):
        if r["line_status"] != "approved" and not keep_superseded:
            continue
        d = r["shift_date"] if by == "start" else r["end_local"][:10]
        if site_from == "hr":
            site = [c for a, b, c in hist[r["employee_no"]] if a <= d <= b][0]
        else:
            site = reg.get(r["site"])
            if site is None:
                continue
        y, m = int(d[:4]), int(d[5:7])
        if y != 2026 or m > 9:
            continue
        k = (site, (m - 1) // 3 + 1)
        out[k] = out.get(k, 0.0) + float(r["paid_hours"])
    return out


def ask_c_variant(target, X, merger=True, mode="golden", restated=True):
    """2026 questions over 2022's: the golden's restatement on the warehouse's basis, the workbook as it stands, the
    copies taken off (each question once but on the site it was asked on), each move applied in the month it moved, or
    the old job read as counting each question once; each with the 2022 sheet's questions as restated in March 2023 or
    as first reported."""
    D = X.D
    old_map = dict(C.OLD_TO_CODE)
    if not merger:
        old_map[K.GRAPHICS_OLD] = None
    wb = {s: 0 for s in K.SITES}
    for (old, y, m), c in D.old.items():
        if y == 2022 and m <= 9 and old_map.get(old):
            wb[old_map[old]] += D.old_restated.get((old, y, m), c["q"]) if restated else c["q"]
    adj = {s: 0 for s in K.SITES}
    for r in rows(target, R.F_MOVES):
        a = dt.datetime.fromisoformat(r["asked_at_utc"])
        b = dt.datetime.fromisoformat(r["moved_at_utc"])
        src, dst = old_map.get(r["from_site"]), old_map.get(r["to_site"])
        win = lambda t: t.year == 2022 and t.month <= 9
        if mode == "golden":
            if win(a):
                if src:
                    adj[src] -= 1
                if dst:
                    adj[dst] += 1
            if win(b) and dst:
                adj[dst] -= 1
        elif mode == "copies_off":
            if win(b) and dst:
                adj[dst] -= 1
        elif mode == "by_move_month":
            if win(b) and src:
                adj[src] -= 1            # the original leaves its source when it moves; the copy stands for it
        elif mode == "net_move":         # the old job read as counting each question once, on the site posted to
            if win(a):
                if src:
                    adj[src] -= 1
                if dst:
                    adj[dst] += 1
    q22 = {s: wb[s] + adj[s] for s in K.SITES}
    return {s: 100.0 * X.q26[s] / q22[s] for s in K.SITES}


def miss(stop, gold, keys, dp):
    return [k for k in keys if K.round_half_up(stop.get(k, 0.0), dp) == K.round_half_up(gold[k], dp)]


# ------------------------------------------------------------------------------------------------- the back-test
def backtest(D):
    """Fit at September 2025, score the routed items of October 2025 to September 2026 against the records: the
    capacity construction and each rival, each with what it could see at the origin."""
    M = C.MIDX
    o = M[(2025, 9)]
    fut = list(range(o + 1, o + 13))
    num = den = 0.0
    for (old, y, m), c in D.old.items():
        if c["ent"] and c["rout"] / c["ent"] > 0.05:
            num += c["vol"]
            den += c["R"]
    for s in K.SITES:
        for i in range(o + 1):
            if D.ent_all[s][i] and D.rout_all[s][i] / D.ent_all[s][i] > 0.05:
                num += D.rev[s][i]
                den += D.R[s][i]
    pace = num / den
    fn = fd = 0.0
    for s in K.SITES:
        for i in range(o + 1):
            if D.ent_all[s][i] and D.rout_all[s][i] / D.ent_all[s][i] < 1 / 60.0:
                fn += D.rout_all[s][i]
                fd += D.ent_all[s][i]
    floor = fn / fd
    out = {k: {} for k in ("capacity", "share", "share_trend", "pool_trend_all", "censored")}
    actual = {}
    lo_x, hi_x = M[(2024, 8)], M[(2025, 7)]
    old_code = C.OLD_TO_CODE
    for s in K.SITES:
        ipq = D.ent_all[s][:o + 1].sum() / D.q[s][:o + 1].sum()
        A = np.array([D.q[s][k] * ipq for k in fut])
        actual[s] = D.rout_all[s][fut].sum()
        h = D.exits[s][lo_x:hi_x + 1].sum() / D.R[s][lo_x:hi_x + 1].sum()
        inf = D.joins[s][o - 11:o + 1].mean()
        x, pool = D.R[s][o], []
        for _ in fut:
            x = x * (1 - h) + inf
            pool.append(x)
        pool = np.array(pool)
        out["capacity"][s] = (A * np.maximum(floor, 1 - pool * pace / A)).sum()
        ts_ = np.arange(0, o + 1)
        b = np.polyfit(ts_, np.log(D.R[s][:o + 1]), 1)
        pt = np.exp(np.polyval(b, np.array(fut)))
        out["pool_trend_all"][s] = (A * np.maximum(floor, 1 - pt * pace / A)).sum()
        w = slice(M[(2025, 1)], o + 1)
        pc = D.rev[s][w].sum() / D.R[s][w].sum()
        out["censored"][s] = (A * np.maximum(floor, 1 - pool * pc / A)).sum()
        l12 = slice(o - 11, o + 1)
        s1 = D.rout_all[s][l12].sum() / D.ent_all[s][l12].sum()
        out["share"][s] = s1 * A.sum()
        # the twelve months before: October to December 2023 from the old workbook, January to September 2024 new
        e0 = sum(c["ent"] for (old, y, m), c in D.old.items() if old_code[old] == s and y == 2023 and m >= 10)
        r0 = sum(c["rout"] for (old, y, m), c in D.old.items() if old_code[old] == s and y == 2023 and m >= 10)
        e0 += D.ent_all[s][M[(2024, 1)]:M[(2024, 9)] + 1].sum()
        r0 += D.rout_all[s][M[(2024, 1)]:M[(2024, 9)] + 1].sum()
        out["share_trend"][s] = max(floor, 2 * s1 - r0 / e0) * A.sum()
    err = {k: (sum(v.values()) - sum(actual.values())) / sum(actual.values()) for k, v in out.items()}
    site_err = {s: (out["capacity"][s] - actual[s]) / actual[s] for s in K.SITES}
    return err, site_err, pace


# ------------------------------------------------------------------------------------------------- the run
def run(target, golden, B, X, meta_path):
    W = B.W
    D = X.D
    M = C.MIDX
    rt = C.routes(X)
    F = json.load(open(os.path.join(HERE, "figures.json"), encoding="utf-8"))
    sub = open(os.path.join(TASK, "submission.md"), encoding="utf-8").read()
    ptxt = open(os.path.join(TASK, "prompt.txt"), encoding="utf-8").read()
    texts = {}
    for f in sorted(os.listdir(target)):
        p = os.path.join(target, f)
        if f.endswith(".pdf"):
            texts[f] = pdf_text(p)
        elif f.endswith(".md"):
            texts[f] = open(p, encoding="utf-8").read()
    flat = re.sub(r"\s+", " ", "\n".join(texts.values()))

    # ------------------------------------------------------------------ the call
    ok("CALL 35 FTE: Frontend & Web 5, Mobile Apps 2, Data & ML 7, Databases 5, Infrastructure & DevOps 5, Game "
       "Development 2, Security 4, Embedded & Hardware 3, Scientific Computing 2; 9 above the 26-FTE envelope",
       X.need_round == GOLD and X.request == 35 and X.gap == 9 and K.ENVELOPE_2027 == 26,
       "%s (%.2f)" % (X.need_round, X.total_fte))
    conv = [k for k in rt if k not in CONVERGENT_EXCLUDE]
    bad = {k: nine(rt[k]["fte"]) for k in conv if rnd(rt[k]["fte"]) != GOLD}
    ok("CALL every defensible construction files the same nine site needs (%d constructions: exit rate over 12, 18 or "
       "24 months or its one-sigma ends, the pool trended over the last 12 to 21 months, arrivals dropped, annual "
       "netting, saturation cuts 1 to 15 per cent, the pace from 2021 to 2023 or 2024 to 2026 only, ratios from 2025, "
       "no floor)" % len(conv), not bad, str(bad))
    clr = {}
    for s in K.SITES:
        vals = [rt[k]["fte"][s] for k in conv]
        clr[s] = min(min(vals) - (GOLD[s] - 0.5), (GOLD[s] + 0.5) - max(vals))
    ok("CALL across every construction, the one-sigma exit rates included, each site stays 0.15 FTE or more inside its "
       "bin", min(clr.values()) >= 0.15, ", ".join("%s %.2f" % (s, v) for s, v in clr.items()))
    ps, fs = {}, {}
    for s in K.SITES:
        num = den = 0.0
        for (old, y, m), c in D.old.items():
            if C.OLD_TO_CODE[old] == s and c["ent"] and c["rout"] / c["ent"] > 0.05:
                num += c["vol"]
                den += c["R"]
        fn = fd = 0.0
        for i in range(len(C.HIST_MONTHS)):
            sh = D.rout_all[s][i] / D.ent_all[s][i]
            if sh > 0.05:
                num += D.rev[s][i]
                den += D.R[s][i]
            if sh < 1 / 60.0:
                fn += D.rout_all[s][i]
                fd += D.ent_all[s][i]
        ps[s] = num / den if den else D.pace
        fs[s] = fn / fd if fd else D.floor
    own = {s: C.fte(C.site_need(D, s, pace=ps[s], floor=fs[s])[1]) for s in K.SITES}
    ok("CALL each site's own pace and own floor (the pooled pace where a site never ran short) file the same nine",
       rnd(own) == GOLD and sum(1 for s in K.SITES if ps[s] != D.pace) >= 5,
       "%s; paces %s" % (nine(own), ", ".join("%s %.1f" % (s, v) for s, v in ps.items())))
    # the lapse rate read by lapse date over the whole roster (the reviewers whose access lapsed in January and February
    # 2024 without a review on the platform are on it), against the golden's window of last reviews
    lap = {s: 0 for s in K.SITES}
    pre = 0
    for r_ in rows(target, R.F_ROSTER):
        if r_["access_status"] == "lapsed" and "2024-01-01" <= r_["lapsed_on"] <= "2026-09-30":
            lap[r_["site"]] += 1
            pre += not r_["first_review_on_platform"]
    risk = {s: D.R[s].sum() for s in K.SITES}
    h_ld = {s: lap[s] / risk[s] for s in K.SITES}
    ld = {s: C.fte(C.site_need(D, s, pool=C.project_pool(D, s, hz_mult=h_ld[s] / C.exit_rate(D, s)[0]))[1])
          for s in K.SITES}
    ok("CALL the lapse rate read by lapse date over January 2024 to September 2026 (the %d reviewers who lapsed in "
       "January and February 2024 without a platform review included) is %.2f to %.2f per cent and files the same "
       "nine; the golden reads lapses by last review, September 2024 to August 2026" % (
           pre, 100 * min(h_ld.values()), 100 * max(h_ld.values())),
       rnd(ld) == GOLD and pre > 300 and all(0.0195 <= v <= 0.0215 for v in h_ld.values()) and
       "a reviewer whose last review fell in March lapses on 1 May." in flat, nine(ld))
    # routed items priced at the routed mix the records show instead of the inbox's mix
    rmix = {}
    for s in K.SITES:
        tot = sum(D.rout[s][t].sum() for t in K.ITEMS)
        rmix[s] = {t: D.rout[s][t].sum() / tot for t in K.ITEMS}

    def need_mix(s, mix):
        items, flag = C.ratios(D, s)
        per_q = sum(items.values())
        pool = C.project_pool(D, s)
        mins = 0.0
        for k, (y, m) in enumerate(C.FWD_MONTHS):
            if y == 2027:
                q = D.plan[s][m - 1]
                A = q * per_q
                sh = max(D.floor, 1 - pool[k] * D.pace / A)
                mins += flag * q * C.STD27["flag"] + A * sh * sum(mix[t] * C.STD27[t] for t in K.ITEMS)
        return mins / 60.0 / K.FTE_HOURS
    rm = {s: need_mix(s, rmix[s]) for s in K.SITES}
    ok("CALL the routed items priced at each site's own routed mix since January 2024 instead of its inbox mix: the "
       "same nine", rnd(rm) == GOLD, nine(rm))
    slack_share = {}
    for t in K.ITEMS:
        num = den = 0.0
        for s in K.SITES:
            for i in range(len(C.HIST_MONTHS)):
                if D.rout_all[s][i] / D.ent_all[s][i] < 1 / 60.0:
                    num += D.rout[s][t][i]
                    den += D.ent[s][t][i]
        slack_share[t] = num / den
    ok("REALISM item types route differently where the volunteers keep up (the items nobody picks up differ by type): "
       "%s" % ", ".join("%s %.2f%%" % (t, 100 * v) for t, v in slack_share.items()),
       max(slack_share.values()) / min(slack_share.values()) > 1.8)
    # the chart's actual months: paid review hours as the operations notes define them, month by month
    defn = ("each moderator flag a paid moderator handled and each community item routed to the team, at the handle-time "
            "standard in force for its year (close votes at 5.0 minutes to the end of 2026, 2.0 from 2027), in the month "
            "the flag was raised or the item entered the inbox, on UTC.")
    wbm = openpyxl.load_workbook(os.path.join(golden, DL.XLSX), data_only=True)["Monthly hours"]
    month_ok = True
    for i in range(len(C.HIST_MONTHS)):
        mins = sum(D.flags[s][i] * K.STD_2026["flag"] + sum(D.rout[s][t][i] * K.STD_2026[t] for t in K.ITEMS)
                   for s in K.SITES)
        month_ok &= wbm.cell(row=5 + i, column=12).value == int(K.round_half_up(mins / 60.0))
    ok("CHART the operations notes define paid review hours once, and the chart's 33 actual months are those hours "
       "exactly (the workbook's Monthly hours sheet, month by month); the prompt grades the chart on its content",
       defn in flat and month_ok and "each site stacked and the envelope drawn across" in ptxt)
    i27 = [K.month_index(2027, m) for m in range(1, 13)]
    real = {s: WD.paid_hours_month(W, s, K.STD_2027)[i27].sum() / K.FTE_HOURS for s in K.SITES}
    ok("CALL the forecast is the world's own 2027: the generator's realised 2027 hours give the same nine site needs",
       rnd(real) == GOLD, "%s (%.2f)" % (nine(real), sum(real.values())))
    paces = {"cut %d%%" % c: C.saturated_pace(D, c / 100.0)[0] for c in (1, 2, 3, 5, 8, 10, 15, 20)}
    paces["2021 to 2023 only"] = C.saturated_pace(D, use_new=False)[0]
    paces["2024 to 2026 only"] = C.saturated_pace(D, use_old=False)[0]
    mr = []
    for (old, y, m), c in D.old.items():
        if c["ent"] and c["rout"] / c["ent"] > 0.05:
            mr.append(c["vol"] / c["R"])
    for s in K.SITES:
        for i in range(len(C.HIST_MONTHS)):
            if D.ent_all[s][i] and D.rout_all[s][i] / D.ent_all[s][i] > 0.05:
                mr.append(D.rev[s][i] / D.R[s][i])
    paces["mean of site-month ratios"] = float(np.mean(mr))
    ok("CALL the pace reviewers keep when the inbox outruns them rounds to 40 a month under every saturation cut, era "
       "and averaging (component 1)", all(39.5 <= v < 40.5 for v in paces.values()) and
       F["components"]["reviews_a_month"] == 40, ", ".join("%s %.2f" % kv for kv in paces.items()))
    ok("CALL each site's monthly lapse rate rounds to 2 per cent, and so does the network's (component 2)",
       all(0.015 <= X_ < 0.025 for X_ in (C.exit_rate(D, s)[0] for s in K.SITES)) and 0.015 <= X.exit_net < 0.025
       and F["components"]["lapse_pct_a_month"] == 2, "%.4f" % X.exit_net)
    ok("CALL almost no new reviewers since late 2024: fewer than one a month at every site over the last twelve months",
       all(C.arrivals(D, s) < 1.0 for s in K.SITES) and X.joins_since < 0.05 * X.joins_2024,
       "%d in 2024, %d since" % (X.joins_2024, X.joins_since))
    a27 = {s: X.inbox27[s] for s in K.SITES}
    c27 = {s: X.cap27[s] for s in K.SITES}
    grp_ok = (X.newly == ["data", "databases", "devops", "gamedev"] and X.already == ["security", "embedded",
                                                                                          "scicomp"] and
              X.covered == ["web", "mobile"] and all(c27[s] < a27[s] for s in X.newly + X.already) and
              all(c27[s] > a27[s] for s in X.covered))
    ok("CALL component 3 holds monthly and over the year: Data & ML, Databases, Infrastructure & DevOps and Game "
       "Development short of their 2027 inbox, Security, Embedded & Hardware and Scientific Computing short already, "
       "Frontend & Web and Mobile Apps covered", grp_ok, "%s %s %s" % (X.newly, X.already, X.covered))
    for mode in ("trend_all", "trend12"):
        sh = []
        for s in K.SITES:
            pool = C.project_pool(D, s, mode)
            items, _ = C.ratios(D, s)
            plan_fwd = np.array([D.outlook[s][m - 10] if y == 2026 else D.plan[s][m - 1] for (y, m) in C.FWD_MONTHS])
            A = plan_fwd * sum(items.values())
            sh.append((s, any(pool[k] * D.pace < A[k] for k in range(len(A)) if C.FWD_MONTHS[k][0] == 2027)))
        grp_ok &= [s for s, x in sh if x] == X.newly + X.already
    ok("CALL component 3 holds with the pool trended instead of carried at its lapse rate", grp_ok)

    # ------------------------------------------------------------------ the three rules on the clean truth
    n_roster = len(rows(target, R.F_ROSTER))
    n_inbox = len(rows(target, R.F_INBOX))
    ok("R1a the decisive computation reads row-level records: %s reviewer access spells and %s site-day-item rows of "
       "the inbox" % (format(n_roster, ","), format(n_inbox, ",")), n_roster > 15000 and n_inbox > 30000)
    heads = []
    for f in sorted(os.listdir(target)):
        p = os.path.join(target, f)
        if f.endswith(".csv"):
            with open(p, encoding="utf-8") as fh:
                heads.append((f, fh.readline()))
        elif f.endswith(".xlsx"):
            wb = openpyxl.load_workbook(p, read_only=True)
            for ws in wb.worksheets:
                for row in ws.iter_rows(min_row=1, max_row=6, values_only=True):
                    heads.append((f, " ".join(str(c) for c in row if c is not None)))
    hit = [(f, h) for f, h in heads if re.search(r"capacit|\bpace\b|cover|churn|lapse rate|exit|forecast|need\b|fte|"
                                                 r"projected|shortfall", h, re.I)]
    ok("R1b no shipped column or sheet heading holds a capacity, a pace, a cover, a lapse rate, a forecast or a need",
       not hit, str(hit[:4]))
    truth_ok = True
    for s in K.SITES:
        for f_, w_ in (("entered", D.ent_all[s]), ("routed", D.rout_all[s])):
            wm = WD.monthly(W, s, f_)[K.month_index(2024, 1):K.month_index(2026, 9) + 1]
            truth_ok &= bool(np.abs(wm - w_).max() < 0.6 * 31 * 4)
        truth_ok &= bool(np.array_equal(D.R[s], W.sites[s].Rm[K.month_index(2024, 1):K.month_index(2026, 9) + 1]))
    ok("R1c the main rows are the clean truth: the inbox's monthly items and the roster's active reviewers are the "
       "world's (routed to the day's rounding)", truth_ok)
    D_new = loaded(target)
    D_new.pace = C.saturated_pace(D, use_old=False)[0]
    D_old = loaded(target)
    D_old.pace = C.saturated_pace(D, use_new=False)[0]
    ok("R2a the platform change mishandled (the old workbook left out, or the daily export left out of the pace): the "
       "same nine", rnd(need_of(D_new)) == GOLD and rnd(need_of(D_old)) == GOLD,
       "%s / %s" % (nine(need_of(D_new)), nine(need_of(D_old))))
    D_g = loaded(target)
    D_g.old = {k: v for k, v in D.old.items() if k[0] != K.GRAPHICS_OLD}
    D_g.pace = C.saturated_pace(D_g)[0]
    ok("R2b Graphics & Shaders' old rows left out (the merger unmapped): the same nine", rnd(need_of(D_g)) == GOLD,
       nine(need_of(D_g)))
    utc_cols = [c for f in (R.F_INBOX, R.F_FLAGSUM) for c in open(os.path.join(target, f)).readline().strip().split(",")
                if "date" in c]
    ok("R2c every date on the call's path is a UTC day (%s) or a month; the Toronto clock lives only in the WFM export"
       % ", ".join(utc_cols), utc_cols == ["entry_date_utc", "date_utc"] and
       "start_local" in open(os.path.join(target, R.F_WFM)).readline())
    tmp = tempfile.mkdtemp(prefix="chk_", dir=TASK)
    try:
        def prelim_inbox(dst):
            keep = [r for r in rows(target, R.F_INBOX) if r["entry_date_utc"] < "2026-09-01"]
            keep += rows(target, R.F_INBOX_PRE)
            with open(dst, "w", newline="", encoding="utf-8") as fh:
                w = csv.DictWriter(fh, fieldnames=R.INBOX_COLS)
                w.writeheader()
                w.writerows(keep)
        t2 = linked_copy(target, os.path.join(tmp, "prelim"), replace={R.F_INBOX: prelim_inbox})
        Dp = loaded(t2)
        ok("R2d / SF12 the preliminary 2 October pull read for September in place of the 8 October pull: the same nine",
           rnd(need_of(Dp)) == GOLD, nine(need_of(Dp)))
        n_pre = len(rows(target, R.F_INBOX_PRE))
        keys = {(r["site"], r["entry_date_utc"], r["item_type"]) for r in rows(target, R.F_INBOX)}
        dup = sum(1 for r in rows(target, R.F_INBOX_PRE) if (r["site"], r["entry_date_utc"], r["item_type"]) in keys)
        ok("R2e both September pulls stacked put %d site-day-item keys in twice (the key check finds every one)" % dup,
           dup == n_pre == 1080)

        # ------------------------------------------------------------------ separation and distractors
        t3 = linked_copy(target, os.path.join(tmp, "main_only"), skip=set(ASK_FILES) | {R.F_PAGEVIEWS,
                                                                                         DOC.F_SURVEY})
        D3 = loaded(t3)
        ok("SEPARATION the call reads no per-flag export, decision review, WFM line, assignment, migration move or flag "
           "service standard: deleted, the same nine to the last digit",
           all(abs(need_of(D3)[s] - X.need_fte[s]) < 1e-9 for s in K.SITES))
        D_q = loaded(target)
        for k in D_q.old:
            D_q.old[k] = dict(D_q.old[k], q=0)
        ok("SEPARATION the old workbook's question column (Ask C's device) is off the call's path: zeroed, the same "
           "needs", all(abs(need_of(D_q)[s] - X.need_fte[s]) < 1e-9 for s in K.SITES))
        meta = json.load(open(meta_path, encoding="utf-8"))
        files = sorted(os.listdir(target))
        t4 = linked_copy(target, os.path.join(tmp, "nodistr"), skip=set(meta["distractor_files"]))
        X4 = C.compute(t4)
        ok("GATE the call and all three asks unchanged with the distractors deleted (%s)" % ", ".join(
            meta["distractor_files"]),
           X4.need_round == X.need_round and X4.ask_a == X.ask_a and X4.ask_b == X.ask_b and X4.ask_c == X.ask_c)
        ok("GATE two or more distractors named in metadata.json, shipped, none named as one",
           len(meta["distractor_files"]) >= 2 and all(d in files for d in meta["distractor_files"]) and
           not any("distract" in f for f in files))
    finally:
        shutil.rmtree(tmp)

    def truth_need(construct):
        out = {}
        for s in K.SITES:
            S = W.sites[s]
            per_q = sum(S.item_rate.values())
            fr = S.flag_rate * (1 - WD.FLAG_CLOSED_SHARE)
            mix = {t: S.item_rate[t] / per_q for t in K.ITEMS}
            out[s] = construct(s, S, per_q, fr, mix) / K.FTE_HOURS
        return out

    def easy(s, S, per_q, fr, mix):
        hs = WD.paid_hours_month(W, s, K.STD_2027)
        q = WD.monthly(W, s, "q")
        lo, hi = K.month_index(2025, 10), K.month_index(2026, 9)
        return hs[lo:hi + 1].sum() / q[lo:hi + 1].sum() * sum(S.plan.values())

    def decisive(s, S, per_q, fr, mix):
        x, hrs = S.Rm[K.month_index(2026, 9)], 0.0
        for (y, m) in C.FWD_MONTHS:
            x *= (1 - K.CHURN)
            if y == 2027:
                A = S.plan[m] * per_q
                sh = max(sum(mix[t] * K.SKIP_BY_TYPE[t] for t in K.ITEMS), 1 - x * K.KAPPA / A)
                hrs += (S.plan[m] * fr * K.STD_2027["flag"] + A * sh * sum(mix[t] * K.STD_2027[t] for t in K.ITEMS)) / 60
        return hrs

    TE = truth_need(easy)
    TD = truth_need(decisive)
    ok("R3a the analyst test on the clean truth: each site's last twelve months of paid hours per question at the 2027 "
       "standards, carried onto the plan, file 23 FTE", req(TE) == 23, "%s (%.2f)" % (nine(TE), sum(TE.values())))
    ok("R3b the decisive step alone on the clean truth (each pool carried from September 2026 at the true lapse rate, "
       "each reviewer at the true pace, each month's inbox beyond capacity routed) files the golden's nine",
       rnd(TD) == GOLD, "%s (%.2f)" % (nine(TD), sum(TD.values())))
    ok("R3b each lower rung's single correction stays wrong: the routed share trended 25, the paid hours trended 28, the "
       "pool trended over the whole record 31",
       [req(rt[k]["fte"]) for k in ("share_trend", "hours_trend", "pool_trend_all")] == [25, 28, 31])
    rep = 0
    seen = set()
    for s in K.SITES:
        e = W.sites[s].items["first_post"][W.day_index[K.NEW_START]:W.day_index[K.HIST_END] + 1]
        for i in range(len(e) - 6):
            key = tuple(int(v) for v in e[i:i + 7])
            rep += key in seen
            seen.add(key)
    ok("R3c no copied block: no seven-day run of a site's first posts repeats anywhere in the inbox", rep == 0, rep)
    plants = re.findall(r"\b\d+(?:\.\d+)? (?:paid )?(?:moderator )?FTE\b", flat)
    allowed = {"%d FTE" % K.ENVELOPE_2027, "%d FTE" % K.POOL_2026}
    ok("R3c no document states a 2027 need, a capacity or a routed share: the only FTE figures are the envelope and "
       "the 2026 pool", set(plants) <= allowed, plants)

    # ------------------------------------------------------------------ HOW, executed literally
    how = {
        "need": "A site's need for a year is the time the paid team will spend that year on the site's moderator "
                "flags and on the site's community review items that reach the team, each item priced at the "
                "handle-time standard the council has set for that year.",
        "flags": "A moderator flag counts when a paid moderator handles it.",
        "year": "An item counts in the year it entered the review inbox, and a flag in the year it was raised, both on "
                "UTC.",
        "future": "The year's need includes the items and flags on posts and accounts that do not yet exist when the "
                  "need is set.",
        "fte": "One FTE is 1,600 hours of need a year.",
        "rounding": "Each site's need is stated in whole FTE, to the nearest whole FTE, and the council's request for "
                    "the year is the sum of the nine site needs.",
        "handoff": "reaches the paid team when no volunteer has reviewed it 72 hours after it entered the site's "
                   "review inbox.",
    }
    missing = [k for k, v in how.items() if v not in flat]
    ok("HOW the charter's seven sentences fixing what is counted are quoted from the pack (population, outcome, window, "
       "unit, rounding, the hand-off)", not missing, missing)
    lit = {s: C.fte(C.naive(D, s, 12)) for s in K.SITES}
    ok("HOW executed word for word on the records (the charter's need for the last twelve months at the 2027 "
       "standards, per question, on the plan) files 23, never the golden", req(lit) == 23 and rnd(lit) != GOLD,
       nine(lit))
    cue = re.findall(r"capacit|\bpace\b|lapse rate|churn|attrition|reviewers? (?:will|per month)|pipeline|"
                     r"privilege|reviews a month|cover\b|shortfall|retention", flat, re.I)
    ok("HOW no shipped sentence states a capacity, a pace, a lapse rate, a pipeline or how many items volunteers will "
       "leave", not cue, cue)
    ok("HOW the prompt asks the need as the charter defines it under the 2027 standards and plan, nothing on method",
       "counted as the charter defines need under the 2027 handle-time standards and question plan" in ptxt and
       not re.search(r"volunteer|reviewer|capacit|trend|share of arrivals", ptxt, re.I))

    # ------------------------------------------------------------------ the semantic fork table
    alt_fte = {h: sum(int(K.round_half_up(X.need_h[s] / h)) for s in K.SITES) for h in (1950.0, 2080.0)}
    ok("SF1 FTE as 1,950 or 2,080 paid hours files %d or %d; the charter's FTE sentence rules it out" % (
        alt_fte[1950.0], alt_fte[2080.0]), how["fte"] in flat and all(v < 35 for v in alt_fte.values()))
    wfm_h = sum(X.ask_b.values()) * 12 / 9 / K.FTE_HOURS
    ok("SF2 need as hours worked (the WFM lines' %.1f FTE a year) is not the charter's need, which prices the items "
       "that reach the team" % wfm_h, how["need"] in flat and abs(wfm_h - 35) > 3)
    D_sh = loaded(target)
    for s in K.SITES:
        for t in K.ITEMS:
            a = D_sh.rout[s][t].copy()
            D_sh.rout[s][t] = a * 0.9 + np.concatenate([[a[0]], a[:-1]]) * 0.1
        D_sh.rout_all[s] = sum(D_sh.rout[s][t] for t in K.ITEMS)
    ok("SF3 items counted by the month they reach the team (three days after entry) instead of the month they entered: "
       "the same nine", rnd(need_of(D_sh)) == GOLD, nine(need_of(D_sh)))
    D_f = loaded(target)
    raised = {s: np.zeros(len(C.HIST_MONTHS)) for s in K.SITES}
    for r_ in rows(target, R.F_FLAGSUM):
        raised[r_["site"]][M[C.ym_of(r_["date_utc"])]] += int(r_["flags_raised"])
    D_f.flags = raised
    ok("SF4 the flags the platform closed counted in (flags raised): a larger request, %d; the charter's flag "
       "sentence rules it out" % req(need_of(D_f)), req(need_of(D_f)) > 35 and
       "Flags the platform closes because the post was deleted before a moderator acted cost the team no time and do "
       "not count." in flat)
    g_new = [r_ for f in (R.F_INBOX, R.F_FLAGSUM, R.F_POSTS, R.F_ROSTER) for r_ in rows(target, f)
             if r_["site"] in ("graphics", K.GRAPHICS_OLD)]
    ok("SF5 Graphics & Shaders has no row after June 2023 in any 2024 to 2026 file, so a tenth site changes nothing",
       not g_new)
    D_26 = loaded(target)
    std_keep = dict(C.STD27)
    C.STD27.update(K.STD_2026)
    try:
        n26 = req(need_of(D_26))
    finally:
        C.STD27.clear()
        C.STD27.update(std_keep)
    ok("SF6 the 2026 standards (close 5.0 minutes) file %d; the prompt and the standards' effective dates rule them out"
       % n26, n26 > 35 and "From 1 January 2027" in flat)
    D_tr = loaded(target)
    tr_plan = {}
    for s in K.SITES:
        lo, hi = M[(2025, 1)], M[(2026, 9)]
        y_ = np.array([D.q[s][i] / K.SEAS[C.HIST_MONTHS[i][1]] for i in range(lo, hi + 1)])
        b = np.polyfit(np.arange(lo, hi + 1), np.log(y_), 1)
        tr_plan[s] = np.array([math.exp(np.polyval(b, M[(2026, 9)] + 3 + m)) * K.SEAS[m] for m in range(1, 13)])
    D_tr.plan = tr_plan
    sf7 = need_of(D_tr)
    ok("SF7 a trend of each site's own questions in place of the plan: %d (%s); the prompt names the plan" % (
        req(sf7), nine(sf7)), "question plan" in ptxt and R.F_PLAN in os.listdir(target))
    once = int(K.round_half_up(sum(X.need_h.values()) / K.FTE_HOURS))
    ok("SF8 the network's hours over 1,600 rounded once: %d, the same request; the charter fixes the rounding path" %
       once, once == 35 and how["rounding"] in flat)
    ok("SF9 how far from the envelope: the request less the envelope, 9 FTE above, whole FTE by the prompt's rounding "
       "sentence; the prompt asks how far above or below it sits",
       X.gap == 9 and "Give FTE in whole numbers" in ptxt and "how far above or below finance's envelope" in ptxt and
       "Above finance's 26-FTE envelope by 9 FTE." in sub)
    ok("SF10 the items on posts and accounts not yet written are in the need: the charter says so and the plan's "
       "questions carry them", how["future"] in flat)
    ok("SF11 the chart's history at the standards in force then: the request is computed on 2027 only, so restating the "
       "history moves no FTE", all(abs(C.fte(C.site_need(D, s)[1]) - X.need_fte[s]) < 1e-12 for s in K.SITES))
    ok("SF12 see R2d: the 8 October pull is the complete one, and the preliminary pull files the same nine",
       "the 8 October pull is the complete one" in flat)

    # ------------------------------------------------------------------ closure axes that can move
    inbox = rows(target, R.F_INBOX)
    ok("CLOSURE 15 no duplicate key in the inbox export, the flag summary, the posts warehouse or the roster",
       len({(r_["site"], r_["entry_date_utc"], r_["item_type"]) for r_ in inbox}) == len(inbox) and
       len({(r_["site"], r_["date_utc"]) for r_ in rows(target, R.F_FLAGSUM)}) == len(rows(target, R.F_FLAGSUM)) and
       len({(r_["account_id"], r_["site"], r_["access_granted"]) for r_ in rows(target, R.F_ROSTER)}) == n_roster)
    ok("CLOSURE 12 the 8 October pull is complete: no item open on any row, and entered is reviewed plus routed plus "
       "open on every row", all(r_["open_at_pull"] == "0" and int(r_["entered"]) == int(r_["reviewed_by_volunteer"]) +
                                int(r_["routed_to_moderation"]) for r_ in inbox))
    bad19 = []
    for r_ in rows(target, R.F_ROSTER):
        y, m = int(r_["last_review"][:4]), int(r_["last_review"][5:7])
        lap = dt.date(y + (m + 1) // 12, (m + 1) % 12 + 1, 1)
        if r_["access_status"] == "lapsed" and r_["lapsed_on"] != lap.isoformat():
            bad19.append(r_["account_id"])
        if r_["access_status"] == "active" and (r_["lapsed_on"] or r_["last_review"] < "2026-09-01"):
            bad19.append(r_["account_id"])
    ok("CLOSURE 19 the roster's status and dates agree on every row (lapsed on the first day of the second month after "
       "the last review; active only with a review in September 2026 or later)", not bad19, bad19[:5])
    ok("CLOSURE 13 monthly and annual netting of the inbox against capacity file the same nine",
       rnd(rt["annual"]["fte"]) == GOLD)

    # ------------------------------------------------------------------ the ladder
    lad = {k: req(rt[k]["fte"]) for k in RIVALS}
    ok("LADDER the rivals file their named requests: last twelve months per question 23 (nine months 23), routed share "
       "trended 25, paid hours trended 28, pool trended over the whole record 31, pace from 2026 throughput 53",
       lad == RIVALS, str(lad))
    ok("LADDER every rival sits 4 FTE or more from the golden and differs from it at two sites or more",
       all(abs(v - 35) >= 4 for v in lad.values()) and
       all(sum(rnd(rt[k]["fte"])[s] != GOLD[s] for s in K.SITES) >= 2 for k in RIVALS),
       ", ".join("%s %s" % (k, nine(rt[k]["fte"])) for k in RIVALS))
    ok("LADDER the golden is bracketed: every share, trend and pool-trend route below it, the censored pace above",
       max(lad[k] for k in ("naive12", "share_trend", "hours_trend", "pool_trend_all")) < 35 < lad["pace_censored"])
    ok("LADDER without the capacity no share or trend moves Data & ML, Databases, Infrastructure & DevOps or Game "
       "Development off its last twelve months' rate by a whole FTE",
       all(abs(rt[k]["fte"][s] - rt["naive12"]["fte"][s]) < 1.0 for k in ("share_trend", "hours_trend")
           for s in ("data", "databases", "devops", "gamedev")))
    err, site_err, bt_pace = backtest(D)
    ok("BACKTEST fitted at September 2025, the capacity construction misses October 2025 to September 2026's routed "
       "items by %.1f per cent network-wide (worst site %.0f per cent)" % (100 * err["capacity"],
                                                                        100 * max(abs(v) for v in site_err.values())),
       abs(err["capacity"]) < 0.10 and max(abs(v) for v in site_err.values()) < 0.20)
    ok("BACKTEST every rival misses by more than a quarter: share carried %.0f, share trended %.0f, pool trended over "
       "the record %.0f, censored pace %+.0f per cent" % tuple(100 * err[k] for k in (
           "share", "share_trend", "pool_trend_all", "censored")),
       all(abs(err[k]) > 0.25 for k in ("share", "share_trend", "pool_trend_all", "censored")))

    # ------------------------------------------------------------------ Ask A
    A = X.ask_a
    keysA = sorted(A)
    clrA = min(K.bin_clearance(v, 0.1) for v in A.values())
    _, numA0, denA0 = ask_a_variant(target)
    ties = [k for k in keysA if (1000 * numA0[k]) % denA0[k] == 0 and ((1000 * numA0[k]) // denA0[k]) % 10 == 5]
    if EXPECT:
        ok("ASK A the 27 shares as built", all("%.1f" % K.round_half_up(A[k], 1) == EXPECT["A"]["%s %d" % k]
                                               for k in keysA))
    ok("ASK A every share is an exact ratio of counts the filed basis fixes, and none falls on a half step, so half-up "
       "and half-even rounding agree (nearest edge %.3f of a step)" % clrA,
       not ties and all(abs(100.0 * numA0[k] / denA0[k] - A[k]) < 1e-9 for k in keysA), str(ties))
    exact24 = 0
    rev = {r_["flag_id"]: r_ for r_ in rows(target, R.F_REVIEWS)}
    flags = {r_["flag_id"]: r_ for f in (R.F_FLAGS1, R.F_FLAGS2) for r_ in rows(target, f)}
    for fid, r_ in flags.items():
        b = rev[fid]["first_action_at_utc"] if fid in rev else r_["handled_at_utc"]
        exact24 += dt.datetime.fromisoformat(b) - dt.datetime.fromisoformat(r_["raised_at_utc"]) == dt.timedelta(
            hours=24)
    ok("ASK A no flag's first action falls exactly 24 hours after it was raised, so within and inside agree",
       exact24 == 0)
    s1, _, _ = ask_a_variant(target, first="export", drop_closed=False, dedupe=False)
    s2, _, _ = ask_a_variant(target, first="export")
    s3, _, _ = ask_a_variant(target, drop_reviewed=True)
    s4, _, _ = ask_a_variant(target, drop_closed=False)
    s5, _, _ = ask_a_variant(target, dedupe=False)
    ok("ASK A S1 every exported row as it stands (duplicates and closed flags in, the export time as first action) "
       "misses all 27", not miss(s1, A, keysA, 1))
    ok("ASK A S2 the cracker (duplicates and closed flags out, the export time as first action) misses all 27",
       not miss(s2, A, keysA, 1), "min gap %.2f" % min(abs(s2[k] - A[k]) for k in keysA))
    hit3 = miss(s3, A, keysA, 1)
    ok("ASK A S3 the reviewed flags dropped from both counts (the over-correction) misses %d of 27" % (27 - len(hit3)),
       27 - len(hit3) >= 12, str(hit3))
    ok("ASK A HA1 the closed flags kept raises every share (all 27 miss)", not miss(s4, A, keysA, 1) and
       all(s4[k] > A[k] for k in keysA))
    hit5 = miss(s5, A, keysA, 1)
    ok("ASK A HA2 the 15 August rows counted twice moves only August figures, and only where they fall across a "
       "rounding edge (%d of 9 August figures move)" % (9 - len([k for k in hit5 if k[1] == 8])),
       all(k[1] == 8 for k in set(keysA) - set(hit5)))
    closed_r = {fid: r_ for fid, r_ in rev.items() if r_["review_closed_at_utc"]}
    open_r = {fid: r_ for fid, r_ in rev.items() if not r_["review_closed_at_utc"]}
    re_ok = all(flags[fid]["handled_at_utc"] == r_["review_closed_at_utc"] and r_["first_action_at_utc"] <
                r_["review_closed_at_utc"] for fid, r_ in closed_r.items())
    re_ok &= all(flags[fid]["handled_at_utc"] == r_["first_action_at_utc"] for fid, r_ in open_r.items())
    ok("ASK A fair: every closed review's flag carries the review's close as its export time, every review still open "
       "at the export its first action, and the review log keeps the first action (%d reviews, %d open)" % (
           len(rev), len(open_r)), re_ok and len(closed_r) > 1000)
    late_t = [r_ for f in (R.F_FLAGS1, R.F_FLAGS2, R.F_REVIEWS) for r_ in rows(target, f)
              if max(v for k, v in r_.items() if k.endswith("_utc") and v) > "2026-10-08 05:38:21"]
    ok("ASK A every time in the flag exports and the review log precedes the export run on 8 October", not late_t,
       str(late_t[:2]))
    ok("ASK A the 24-hour flag response is the standard's share, the post-deleted closures out of both counts; the "
       "prompt names the measure",
       "for each site and month, the share of the flags raised in the month, on UTC, that were answered within 24 hours."
       in flat and "left out of both counts." in flat and "24-hour flag response" in ptxt)
    ok("ASK A the flag service standard files the basis once, and no file says the export time moves on review",
       "first action on it came within 24 hours of the flag being raised" in flat and
       not re.search(r"marked handled again|re-?stamp|overwrit", flat, re.I))

    # ------------------------------------------------------------------ Ask B
    Bh = X.ask_b
    keysB = sorted(Bh)
    if EXPECT:
        ok("ASK B the 27 quarter totals as built", all(int(K.round_half_up(Bh[k])) == EXPECT["B"]["%s %d" % k]
                                                       for k in keysB))
    ok("ASK B every quarter total is a whole number of hours (whole-hour shift lines), so no total sits on a rounding "
       "edge", all(abs(v - round(v)) < 1e-9 for v in Bh.values()))
    b1 = ask_b_variant(target, site_from="export", keep_superseded=True)
    b2 = ask_b_variant(target, site_from="export")
    b3 = ask_b_variant(target, by="end")
    b4 = ask_b_variant(target, keep_superseded=True)
    b5 = ask_b_variant(target, site_from="export", map_label=False)
    ok("ASK B S1 the export's site with the superseded lines kept misses all 27", not miss(b1, Bh, keysB, 0))
    ok("ASK B S2 the cracker (superseded out, Ops & Cloud mapped, by start day, the export's site) misses all 27",
       not miss(b2, Bh, keysB, 0), "min gap %.0f" % min(abs(b2.get(k, 0) - Bh[k]) for k in keysB))
    moved3 = [k for k in keysB if round(b3.get(k, 0)) != round(Bh[k])]
    ok("ASK B S3 shifts put on their end day move the quarter-end nights (%d figures miss)" % len(moved3),
       len(moved3) >= 4, str(moved3))
    moved4 = [k for k in keysB if round(b4.get(k, 0)) != round(Bh[k])]
    ok("ASK B HB2 the superseded May lines kept move the second quarter at the corrected moderators' sites only",
       moved4 and all(k[1] == 2 for k in moved4), str(moved4))
    ok("ASK B HB1 the WFM label Ops & Cloud left unmapped drops Infrastructure & DevOps' three quarters",
       all(("devops", q) not in b5 for q in (1, 2, 3)))
    hist = {}
    for r_ in rows(target, R.F_ASSIGN):
        hist.setdefault(r_["employee_no"], []).append((r_["effective_from"], r_["effective_to"] or "9999-12-31"))
    cover = all(sum(a <= r_["shift_date"] <= b for a, b in hist[r_["employee_no"]]) == 1
                for r_ in rows(target, R.F_WFM))
    movers = sum(1 for v in hist.values() if len(v) > 1)
    ok("ASK B fair: every shift line's start day falls in exactly one assignment of its moderator (%d moved in July and "
       "August)" % movers, cover and movers == 6)
    ok("ASK B the ops notes file the basis once (the site worked for on the shift, the quarter it started, Toronto "
       "time) and nothing says the export's site is the current one",
       "A shift's hours belong to the site the moderator worked for on that shift, and to the quarter in which the "
       "shift started, on Toronto time." in flat and not re.search(
           r"current (?:site|record)|moderators? (?:were |was )?(?:moved|reassigned)|changed site|site column", flat,
           re.I))

    # ------------------------------------------------------------------ Ask C
    Cq = X.ask_c
    if EXPECT:
        ok("ASK C the nine percentages as built", all("%.1f" % K.round_half_up(Cq[s], 1) == EXPECT["C"][s]
                                                      for s in K.SITES))
    clrC = min(K.bin_clearance(v, 0.1) for v in Cq.values())
    ok("ASK C every percentage sits 0.1 of its rounding step or more from the edge", clrC >= 0.1, "%.3f" % clrC)
    c1 = ask_c_variant(target, X, merger=False, mode="workbook")
    c2 = ask_c_variant(target, X, mode="workbook")
    c3 = ask_c_variant(target, X, mode="copies_off")
    c4 = ask_c_variant(target, X, mode="by_move_month")
    ok("ASK C S1 the workbook as it stands, Graphics left out, misses all nine", not miss(c1, Cq, K.SITES, 1))
    ok("ASK C S2 the workbook's 2022 used as it stands on the old job's basis (the merger applied, no restatement on "
       "the warehouse's basis) misses all nine",
       not miss(c2, Cq, K.SITES, 1), "min gap %.2f" % min(abs(c2[s] - Cq[s]) for s in K.SITES))
    h3 = miss(c3, Cq, K.SITES, 1)
    h4 = miss(c4, Cq, K.SITES, 1)
    ok("ASK C S3 the copies taken off (each question once, but on the site it was asked on) misses %d of 9" % (
        9 - len(h3)), 9 - len(h3) >= 6, str(h3))
    ok("ASK C C1 each move applied in the month it moved instead of the month asked files the same nine figures (no "
       "move straddles the window's edges enough to cross a rounding step)", len(h4) == 9, str(h4))
    truth22 = {s: 0 for s in K.SITES}
    for s in K.SITES:
        Sx = W.sites[s]
        truth22[s] = int(sum(Sx.q[W.day_index[d]] for d in K.days(dt.date(2022, 1, 1), dt.date(2022, 9, 30))))
    ok("ASK C fair: the 2022 counts rebuilt from the workbook and the migration log are the world's questions by the "
       "site they live on, to the question", all(X.q22[s] == truth22[s] for s in K.SITES),
       str({s: X.q22[s] - truth22[s] for s in K.SITES}))
    ok("ASK C the warehouse notes file the basis once and the workbook's own note says what it counted",
       "Questions are counted once each, in the month the question was first asked (UTC), on the site the question is "
       "on today." in flat)
    old_job = ("The old statistics job counted a question on every site it was posted to. A question moved to another "
               "site is posted there anew under a new question id, so the job counted it on its first site in the month "
               "it was asked and on its new site in the month it moved.")
    c5 = ask_c_variant(target, X, mode="net_move")
    ok("ASK C the reading that the old job counted each question once on the site it was posted to (moves netted in "
       "the month asked, no copies taken off) misses all nine, and the warehouse notes' sentence on the old job rules "
       "it out", not miss(c5, Cq, K.SITES, 1) and old_job in flat, "min gap %.2f" % min(abs(c5[s] - Cq[s])
                                                                                   for s in K.SITES))
    ok("ASK C the reading with no restatement (the workbook's 2022 as reported, S2) is ruled out by the warehouse "
       "notes: comparisons with 2021 to 2023 are made on the warehouse's basis",
       "Comparisons with 2021 to 2023 are made on the warehouse's basis." in flat and not miss(c2, Cq, K.SITES, 1))
    cg = ask_c_variant(target, X, mode="golden")
    ok("ASK C the golden's construction (restated 2022 sheet, the old job's double count undone on the warehouse's "
       "basis, the merger) reproduces the nine figures exactly", all(abs(cg[s] - Cq[s]) < 1e-9 for s in K.SITES))
    cl = ask_c_variant(target, X, mode="golden", restated=False)
    hl = miss(cl, Cq, K.SITES, 1)
    ok("ASK C C2 device: the careful-but-literal path (the counting rule, the comparison basis, the merger and the "
       "migration log all applied, the 2022 sheet's questions as first reported) misses %d of 9 (nearest %.2f points)"
       % (9 - len(hl), min(abs(cl[s] - Cq[s]) for s in K.SITES)), len(hl) == 0 and
       min(abs(cl[s] - Cq[s]) for s in K.SITES) >= 0.15, str(hl))
    shares = {}
    for (old, y, m), v in X.D.old_restated.items():
        shares.setdefault(old, [0, 0])
        shares[old][0] += v - X.D.old[(old, y, m)]["q"]
        shares[old][1] += v
    sh = {o: a / b for o, (a, b) in shares.items()}
    ok("ASK C C2 fair and honest: the March 2023 rerun restates May to August 2022 for every old site, by %.1f to %.1f "
       "per cent of their questions (the API's share differs by site), on the job's own basis, and touches no review "
       "item, reviewer or flag" % (100 * min(sh.values()), 100 * max(sh.values())),
       len(sh) == 10 and max(sh.values()) / min(sh.values()) > 1.6 and all(v > 0 for v in sh.values()) and
       "its questions replace those above for these months" in open_text(target, R.F_OLD))

    # ------------------------------------------------------------------ realism sweep
    by_site = {}
    for r_ in rows(target, R.F_POSTS):
        by_site.setdefault(r_["site"], []).append(int(r_["questions"]))
    runs = [s for s, v in by_site.items() if any(v[i] == v[i + 1] for i in range(len(v) - 1)) or
            len({(v[i], v[i + 1]) for i in range(len(v) - 1)}) < len(v) - 1]
    heads_ = [int(sum(v[-9:])) for v in by_site.values()] + [int(sum(D.plan[s])) for s in K.SITES] + \
        [sum(X.R_sep26.values()), n_roster]
    allq = [x for v in by_site.values() for x in v]
    ok("REALISM no site's monthly question count repeats month to month or in a two-month run, and no headline total "
       "(each site's January to September 2026 questions, its 2027 plan, the reviewer counts) is a round thousand",
       not runs and not any(h % 1000 == 0 for h in heads_),
       "%d coincident values among %d site-months, as Poisson counts of this size give" % (
           len(allq) - len(set(allq)), len(allq)))
    figsB = [int(round(v)) for v in Bh.values()]
    ok("REALISM the 27 quarter totals are all different", len(set(figsB)) == len(figsB))
    _, numA, denA = ask_a_variant(target)
    ok("REALISM the 27 flag counts behind Ask A are all different", len(set(denA.values())) == len(denA))
    persons = set(re.findall(r"\b([A-Z][a-z]+ [A-Z][a-z]+)\b", "\n".join(texts.values())))
    firsts = {n.split()[0] for n in K.PEOPLE.values()}
    named = {n for n in persons if n.split()[0] in firsts}
    ok("REALISM every person named in the documents is a drawn persona", named <= set(K.PEOPLE.values()),
       sorted(named - set(K.PEOPLE.values())))
    dash = [f for f in os.listdir(target) if f.endswith((".csv", ".md")) and re.search(
        "[—–]", open(os.path.join(target, f), encoding="utf-8").read())]
    dash += [f for f, t in texts.items() if re.search("[—–]", t)]
    gtext = ""
    for f in os.listdir(golden):
        with zipfile.ZipFile(os.path.join(golden, f)) as z:
            for n in z.namelist():
                if re.match(r"(ppt/slides/slide\d+|xl/sharedStrings|xl/worksheets/sheet\d+)\.xml$", n):
                    gtext += " ".join(re.findall(r"<(?:a:t|t)(?: [^>]*)?>([^<]*)<", z.read(n).decode("utf-8")))
    ok("REALISM no em or en dash in the pack, the goldens or the submission",
       not dash and not re.search("[—–]", sub + gtext + ptxt), dash)
    ok("REALISM every shipped file carries the file list's morning as its time",
       all(abs(os.path.getmtime(os.path.join(target, f)) - BLD.MTIME) < 1 for f in os.listdir(target)))
    late = []
    for f in os.listdir(target):
        p = os.path.join(target, f)
        if f.endswith(".xlsx"):
            with zipfile.ZipFile(p) as z:
                core = z.read("docProps/core.xml").decode()
            if re.search(r"<dcterms:modified[^>]*>([^<]+)<", core).group(1)[:10] > K.ASOF.isoformat():
                late.append(f)
        elif f.endswith(".pdf"):
            raw = open(p, "rb").read()
            m_ = re.search(rb"/CreationDate \(D:(\d{8})", raw)
            if not m_ or m_.group(1).decode() > K.ASOF.strftime("%Y%m%d"):
                late.append(f)
    dated = re.findall(r"\b(\d{1,2}) (October|November|December) 2026\b", flat)
    late += ["text: %s %s 2026" % d for d in dated if (d[1] != "October" or int(d[0]) > K.ASOF.day)]
    ok("REALISM no document, workbook or sentence is dated after the file list (9 October 2026)", not late, late)
    fl = texts[DOC.F_FILELIST]
    unlisted = [f for f in sorted(os.listdir(target)) if f != DOC.F_FILELIST and f not in fl]
    counts_ok = all(("| %s |" % format(n_, ",")) in fl or n_ is None for f, n_ in B.n.items()
                    if f in (R.F_INBOX, R.F_ROSTER, R.F_MOVES))
    ok("REGISTER the file list names every shipped file with its row count", not unlisted and counts_ok, unlisted)

    # ------------------------------------------------------------------ input gates
    files = sorted(os.listdir(target))
    fmts = {f.rsplit(".", 1)[1] for f in files}
    big = max(B.n.values())
    ok("GATE 10 or more files, 3 or more formats, one of 25,000 or more rows",
       len(files) >= 10 and len(fmts) >= 3 and big >= 25000, "%d files, %s, %d rows" % (len(files), sorted(fmts), big))

    # ------------------------------------------------------------------ the goldens agree with figures.json
    wb = openpyxl.load_workbook(os.path.join(golden, DL.XLSX), data_only=True)
    ws = wb["Need 2027"]
    need_rows = {ws.cell(row=i, column=1).value: ws.cell(row=i, column=8).value for i in range(6, 16)}
    tail = {ws.cell(row=i, column=1).value: ws.cell(row=i, column=8).value for i in range(17, 21)}
    ok("GOLDEN the workbook's site needs, request and distance from the envelope are the figures'",
       all(need_rows[DL.NAME[s]] == F["call"]["need_fte"][DL.NAME[s]] for s in K.SITES) and
       need_rows["Network"] == F["call"]["request_fte"] and
       tail["Request above the envelope (FTE)"] == F["call"]["above_envelope_fte"])
    wa, wbb, wc = wb["Flag response"], wb["Hours worked"], wb["Questions vs 2022"]
    okA = all("%.1f" % wa.cell(row=5 + i, column=2 + j).value ==
              F["answers"]["ask_a_pct"]["%s %s 2026" % (DL.NAME[s], K.MONTH_NAMES[m])]
              for i, s in enumerate(K.SITES) for j, m in enumerate((7, 8, 9)))
    okB = all(format(wbb.cell(row=5 + i, column=2 + j).value, ",") ==
              F["answers"]["ask_b_hours"]["%s Q%d 2026" % (DL.NAME[s], q)]
              for i, s in enumerate(K.SITES) for j, q in enumerate((1, 2, 3)))
    okC = all("%.1f" % wc.cell(row=5 + i, column=2).value == F["answers"]["ask_c_pct"][DL.NAME[s]]
              for i, s in enumerate(K.SITES))
    ok("GOLDEN the workbook's three tables are the figures'", okA and okB and okC)
    with zipfile.ZipFile(os.path.join(golden, DL.PPTX)) as z:
        s1x = z.read("ppt/slides/slide1.xml").decode()
        media = [n for n in z.namelist() if n.startswith("ppt/media/")]
        big_png = max((z.getinfo(n).file_size for n in media), default=0)
    ok("GOLDEN slide 1 states the request and its distance from the envelope; slide 2 carries the rendered chart",
       "Request 35 paid moderator FTE for 2027" in s1x and "9 FTE above finance" in s1x and big_png > 50000)
    ok("GOLDEN the submission states the call and every figure from figures.json (figures_check runs separately)",
       "**The council should request 35 paid moderator FTE for 2027, 9 above finance's 26-FTE envelope" in sub)

    # the pace stated loosely everywhere a grader reads it; the envelope compared network-wide only
    blocks = sub.split("## 4. Deliverable Answers")[0]
    gold_txt = gtext + "\n" + "\n".join(str(c) for ws in wb.worksheets for row in ws.iter_rows(values_only=True)
                                         for c in row if isinstance(c, str))
    paces_said = re.findall(r"(about )?\b40(?:\.\d+)? (?:reviews|items)", blocks + gold_txt, re.I)
    ok("WRITEUP the volunteer pace appears only as about 40 reviews a month, never a finer figure, in the graded blocks "
       "and both goldens (the call holds from a pace of about 39 to 41.3)",
       paces_said and all(a for a in paces_said) and not re.search(r"\b(?:39|40|41)\.\d+ reviews", blocks + gold_txt),
       str(len(paces_said)))
    ws_need = wb["Need 2027"]
    need_heads = [ws_need.cell(row=5, column=c).value for c in range(1, 9)]
    env_cells = [c for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for c in row
                 if isinstance(c, str) and "envelope" in c.lower()]
    ok("WRITEUP the envelope is compared network-wide only: request less 26 is 9 above; no per-site envelope column or "
       "share in either golden",
       not any("envelope" in str(h).lower() for h in need_heads) and X.gap == X.request - K.ENVELOPE_2027 == 9 and
       "Request above the envelope (FTE)" in env_cells and "9 FTE above finance's 2027 envelope of 26 FTE" in s1x and
       not any(re.search(r"envelope (?:share|by site|per site)|site'?s? (?:share of the )?envelope", c, re.I)
               for c in env_cells), str(env_cells))
    fails = [x for x in RESULTS if not x[1]]
    for label, good, detail in RESULTS:
        print("%s  %s%s" % ("PASS" if good else "FAIL", label, ("  [%s]" % detail) if detail else ""))
    ids = set(re.findall(r"\bSF\d+\b", " ".join(x[0] for x in RESULTS)))
    print("checks: %d assertions, %d failed; fork rows asserted: %s" % (len(RESULTS), len(fails),
                                                                       ", ".join(sorted(ids, key=lambda x: int(x[2:])))))
    return fails


DLX = None


def main():
    global DLX
    target = os.path.join(TASK, "target")
    golden = os.path.join(TASK, "golden")
    B = BLD.world()
    B.n = {f: len(rows(target, f)) for f in os.listdir(target) if f.endswith(".csv")}
    DLX = DL.compute(target)
    X = DLX
    fails = run(target, golden, B, X, os.path.join(TASK, "metadata.json"))
    if fails:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
