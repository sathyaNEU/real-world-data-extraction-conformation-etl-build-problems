"""task156 records: the world shaped into the exports Millrind's systems write.

The review platform (from January 2024): the daily inbox export by site, entry day and item type, pulled on 8 October
2026, with a preliminary September pull of 2 October; the reviewer roster. The old review system (2021 to 2023): a
monthly workbook, one sheet a year, by the old site names. The moderator tool: the daily flag summary and the per-flag
export for July to September 2026 in two parts, with the decision-review log. The posts warehouse (from January
2024): questions by site and month. Planning: the 2027 question plan. The WFM system: the paid team's shift lines and
the HR assignment history. The moderators' migration log, 2021 to 2023. Web analytics: page views."""
from __future__ import annotations

import datetime as dt
import math

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import asks as AK
import common as K
import world as WD

F_INBOX = "review_inbox_daily_2024-01-01_2026-09-30.csv"
F_INBOX_PRE = "review_inbox_daily_2026-09_pulled_2026-10-02.csv"
F_OLD = "review_queues_monthly_2021-2023.xlsx"
F_ROSTER = "reviewer_roster_2026-10-08.csv"
F_FLAGSUM = "modtool_flag_summary_daily_2024-01-01_2026-09-30.csv"
F_POSTS = "posts_questions_by_site_month_2024-01_2026-09.csv"
F_PLAN = "question_plan_2027.xlsx"
F_REGISTER = "site_register.csv"
F_FLAGS1 = "modtool_flags_2026-07-01_to_2026-08-15.csv"
F_FLAGS2 = "modtool_flags_2026-08-15_to_2026-09-30.csv"
F_REVIEWS = "flag_decision_reviews_2026-07_to_2026-09.csv"
F_WFM = "wfm_shift_lines_2026-01-01_2026-09-30.csv"
F_ASSIGN = "moderator_site_assignments.csv"
F_MOVES = "question_moves_2021-2023.csv"
F_PAGEVIEWS = "pageviews_by_site_month_2021-01_2026-09.csv"

ITEM_CODE = {"close": "close_vote", "first_post": "first_post", "low_quality": "low_quality", "edit": "suggested_edit"}
OLD_SITE = {c: K.SITE[c]["old"] for c in K.SITES}
OLD_SITE["graphics"] = K.GRAPHICS_OLD


def ts(t):
    return t.strftime("%Y-%m-%d %H:%M:%S")


# ------------------------------------------------------------------------------------------------- inbox export
INBOX_COLS = ["site", "entry_date_utc", "item_type", "entered", "reviewed_by_volunteer", "routed_to_moderation",
              "open_at_pull", "pulled_at_utc"]


def inbox_cells(W):
    """Integer cells by (site, day, item type) over the platform's record: entered, reviewed, routed."""
    out = {}
    for s in K.SITES:
        S = W.sites[s]
        for di, d in enumerate(W.all_days):
            if d < K.NEW_START or d > K.HIST_END:
                continue
            for t in K.ITEMS:
                e = int(S.items[t][di])
                rt = int(K.round_half_up(S.rout[t][di]))
                rt = max(0, min(rt, e))
                out[(s, d, t)] = (e, e - rt, rt)
    return out


def inbox_rows(W, cells):
    rows = []
    for d in K.days(K.NEW_START, K.HIST_END):
        for s in K.SITES:
            for t in K.ITEMS:
                e, rv, rt = cells[(s, d, t)]
                rows.append([s, d.isoformat(), ITEM_CODE[t], e, rv, rt, 0, "2026-10-08 06:12:44"])
    return rows


def inbox_prelim_rows(W, cells):
    """The 2 October pull of September: items entering on 29 and 30 September not yet past their 72 hours are
    open (those a volunteer has already reviewed are counted)."""
    rows = []
    for d in K.days(dt.date(2026, 9, 1), K.HIST_END):
        for s in K.SITES:
            for t in K.ITEMS:
                e, rv, rt = cells[(s, d, t)]
                if d >= dt.date(2026, 9, 29):
                    rows.append([s, d.isoformat(), ITEM_CODE[t], e, rv, 0, rt, "2026-10-02 05:47:09"])
                else:
                    rows.append([s, d.isoformat(), ITEM_CODE[t], e, rv, rt, 0, "2026-10-02 05:47:09"])
    return rows


# ------------------------------------------------------------------------------------------------- roster
ROSTER_COLS = ["account_id", "site", "access_granted", "first_review_on_platform", "last_review", "access_status",
               "lapsed_on"]


def _month_day(r, y, m, lo=1, hi=None):
    hi = hi or K.days_in_month(y, m)
    return dt.date(y, m, int(r.integers(lo, hi + 1)))


def _granted_on(r, j):
    y, m = WD.month_ym(j)
    return _month_day(r, y, m, 1, 20) if j >= 0 else _month_day(r, y, m)


def _account_id(r, granted, used):
    """Network account ids: older accounts hold smaller ids."""
    base = 180_000 + (granted - dt.date(2008, 1, 1)).days * 2_400
    aid = int(base + int(r.integers(0, 2_300)))
    while aid in used:
        aid += 1
    used.add(aid)
    return aid


def roster_rows(W):
    """One row per reviewer whose review access had not lapsed before 1 January 2024 (the platform's start).
    A reviewer is active in a month from the month of their first review on the platform to the month of their last;
    access lapses after a calendar month with no reviews, on the first day of the month after it."""
    t24 = K.month_index(2024, 1)
    tH = K.month_index(2026, 9)
    rows = []
    used = set()
    for s in K.SITES:
        P = W.sites[s].pool
        r = K.rng("roster", s)
        for i in range(len(P["kappa"])):
            j, L = P["join"][i], P["last"][i]
            if L is not None and L < t24:
                continue
            if j > tH:
                continue
            first_m = max(j, t24)
            last_m = L if (L is not None and L <= tH) else tH
            fy, fm = WD.month_ym(first_m)
            ly, lm = WD.month_ym(last_m)
            granted = _granted_on(r, j)
            if first_m == t24 and j < t24:
                first = _month_day(r, 2024, 1, 1, 6)
            else:
                first = max(granted, _month_day(r, fy, fm))
            last = _month_day(r, ly, lm)
            if last < first:
                last = first
            if (ly, lm) == (2026, 9):
                status, lapsed = "active", ""
                last = _month_day(r, 2026, 9, 1, 30)
                if last < first:
                    last = first
            else:
                status = "lapsed"
                ny, nm = (ly + (lm >= 11), (lm + 1) % 12 + 1)
                lapsed = dt.date(ny, nm, 1).isoformat()
            rows.append([_account_id(r, granted, used), s, granted.isoformat(), first.isoformat(), last.isoformat(),
                         status, lapsed])
    # reviewers whose access was still open on 1 January 2024 but who made no review on the platform: their last review
    # fell in November or December 2023, so access lapsed on 1 January or 1 February 2024
    for s in K.SITES:
        P = W.sites[s].pool
        r = K.rng("roster_pre", s)
        for i in range(len(P["kappa"])):
            j, L = P["join"][i], P["last"][i]
            if L is None or L not in (t24 - 2, t24 - 1) or j > L:
                continue
            granted = _granted_on(r, j)
            ly, lm = WD.month_ym(L)
            last = max(granted, _month_day(r, ly, lm))
            lapsed = dt.date(2024, 1 if lm == 11 else 2, 1).isoformat()
            rows.append([_account_id(r, granted, used), s, granted.isoformat(), "", last.isoformat(), "lapsed", lapsed])
    rows.sort(key=lambda x: x[0])
    return rows


# ------------------------------------------------------------------------------------------------- flag summary
FLAGSUM_COLS = ["site", "date_utc", "flags_raised", "closed_post_deleted", "handled_by_moderator"]


def flagsum_rows(W):
    rows = []
    for d in K.days(K.NEW_START, K.HIST_END):
        di = W.day_index[d]
        for s in K.SITES:
            S = W.sites[s]
            n, c = int(S.flags[di]), int(S.flag_closed[di])
            rows.append([s, d.isoformat(), n, c, n - c])
    return rows


# ------------------------------------------------------------------------------------------------- posts warehouse
POSTS_COLS = ["site", "month", "questions"]


def posts_rows(W):
    rows = []
    for (y, m) in K.months(K.NEW_START, K.HIST_END):
        for s in K.SITES:
            S = W.sites[s]
            q = sum(S.q[W.day_index[d]] for d in K.days(dt.date(y, m, 1), dt.date(y, m, K.days_in_month(y, m))))
            rows.append([s, "%d-%02d" % (y, m), int(q)])
    return rows


def warehouse_month(W, s, y, m):
    S = W.sites[s]
    return int(sum(S.q[W.day_index[d]] for d in K.days(dt.date(y, m, 1), dt.date(y, m, K.days_in_month(y, m)))))


# ------------------------------------------------------------------------------------------------- old workbook
def _split(r, x, share):
    a = int(r.binomial(int(x), share)) if x > 0 else 0
    return a, int(x) - a


def old_workbook_cells(W, moves):
    """Per old site and month 2021 to 2023: the counts the old system reported. Game Development's records before July
    2023 are split between Game Development and Graphics & Shaders, which shared nothing but the topic."""
    r = K.rng("oldsplit")
    cells = {}
    months = K.months(K.D0, K.OLD_END)
    # live questions by old site and month (the posts warehouse's basis, applied to the old sites)
    live = {}
    for (y, m) in months:
        for s in K.SITES:
            q = sum(W.sites[s].q[W.day_index[d]] for d in K.days(dt.date(y, m, 1), dt.date(y, m, K.days_in_month(y, m))))
            if s == "gamedev" and dt.date(y, m, 1) < K.MERGER:
                g, rest = _split(r, q, K.GRAPHICS_SHARE)
                live[("graphics", y, m)] = g
                live[("gamedev", y, m)] = rest
            else:
                live[(s, y, m)] = int(q)
    # the moves: in by asked month, out by asked month, copies in by moved month
    mi_ask, mo_ask, cp_mov = {}, {}, {}
    for mv in moves:
        a, b = mv["asked"], mv["moved"]
        mi_ask[(mv["dst"], a.year, a.month)] = mi_ask.get((mv["dst"], a.year, a.month), 0) + 1
        mo_ask[(mv["src"], a.year, a.month)] = mo_ask.get((mv["src"], a.year, a.month), 0) + 1
        if b.date() <= K.OLD_END:
            cp_mov[(mv["dst"], b.year, b.month)] = cp_mov.get((mv["dst"], b.year, b.month), 0) + 1
    for (y, m) in months:
        t = K.month_index(y, m)
        olds = ["gamedev", "graphics"] if dt.date(y, m, 1) < K.MERGER else ["gamedev"]
        for s in K.SITES:
            S = W.sites[s]
            mdays = [W.day_index[d] for d in K.days(dt.date(y, m, 1), dt.date(y, m, K.days_in_month(y, m)))]
            ent = {x: int(sum(S.items[x][di] for di in mdays)) for x in K.ITEMS}
            rout = {x: int(K.round_half_up(sum(S.rout[x][di] for di in mdays))) for x in K.ITEMS}
            rev = {x: ent[x] - rout[x] for x in K.ITEMS}
            vol = int(K.round_half_up(sum(S.reviewed_on[di] for di in mdays)))
            flags = int(sum(S.flags[di] - S.flag_closed[di] for di in mdays))
            R = int(S.Rm[t])
            targets = olds if s == "gamedev" else [s]
            if len(targets) == 1:
                parts = {targets[0]: dict(ent=ent, rev=rev, rout=rout, vol=vol, flags=flags, R=R)}
            else:
                g = {}
                rest = {}
                for k in ("ent", "rout"):
                    src = ent if k == "ent" else rout
                    g[k], rest[k] = {}, {}
                    for x in K.ITEMS:
                        a, b = _split(r, src[x], K.GRAPHICS_SHARE)
                        if k == "rout":
                            a, b = min(a, g["ent"][x]), src[x] - min(a, g["ent"][x])
                            if b > rest["ent"][x]:
                                b = rest["ent"][x]
                                a = src[x] - b
                        g[k][x], rest[k][x] = a, b
                g["rev"] = {x: g["ent"][x] - g["rout"][x] for x in K.ITEMS}
                rest["rev"] = {x: rest["ent"][x] - rest["rout"][x] for x in K.ITEMS}
                gv, rv_ = _split(r, vol, K.GRAPHICS_SHARE)
                gf, rf = _split(r, flags, K.GRAPHICS_SHARE)
                gR, rR = _split(r, R, K.GRAPHICS_SHARE)
                parts = {"graphics": dict(ent=g["ent"], rev=g["rev"], rout=g["rout"], vol=gv, flags=gf, R=gR),
                         "gamedev": dict(ent=rest["ent"], rev=rest["rev"], rout=rest["rout"], vol=rv_, flags=rf, R=rR)}
            for old, P in parts.items():
                q_wb = (live[(old, y, m)] - mi_ask.get((old, y, m), 0) + mo_ask.get((old, y, m), 0)
                        + cp_mov.get((old, y, m), 0))
                P["questions"] = q_wb
                P["reported"] = q_wb
                P["live"] = live[(old, y, m)]
                cells[(old, y, m)] = P
    # the tag-sync fault of 9 May to 22 August 2022: questions posted through the API stayed out of the statistics
    # job's question counts until the job was rerun on 14 March 2023. The 2022 sheet keeps the counts as first
    # reported; the rerun's restated counts (the job's own basis, complete) follow the sheet's notes.
    rt = K.rng("tagsync")
    for old in OLD_ORDER:
        for m, frac in TAGSYNC_MONTHS.items():
            P = cells[(old, 2022, m)]
            missed = int(rt.binomial(P["questions"], TAGSYNC_API[old] * frac))
            P["reported"] = P["questions"] - missed
    return cells


OLD_HEAD_2021 = ["Site", "Month", "Questions", "Close votes in", "Close votes reviewed", "Close votes to moderators",
                 "First posts in", "First posts reviewed", "First posts to moderators", "Low quality in",
                 "Low quality reviewed", "Low quality to moderators", "Active reviewers", "Volunteer reviews",
                 "Moderator flags handled"]
OLD_HEAD = OLD_HEAD_2021[:12] + ["Suggested edits in", "Suggested edits reviewed", "Suggested edits to moderators"] + \
    OLD_HEAD_2021[12:]
OLD_ORDER = ["web", "mobile", "data", "databases", "devops", "gamedev", "graphics", "security", "embedded", "scicomp"]
# the tag-sync fault: each old site's share of questions posted through the API, and the share of each month it ran
TAGSYNC_API = {"web": 0.019, "mobile": 0.024, "data": 0.034, "databases": 0.027, "devops": 0.036, "gamedev": 0.017,
               "graphics": 0.021, "security": 0.022, "embedded": 0.018, "scicomp": 0.029}
TAGSYNC_MONTHS = {5: 0.74, 6: 1.0, 7: 1.0, 8: 0.71}


def write_old_workbook(path, cells):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="B7BEC8")
    head_fill = PatternFill("solid", fgColor="DCE3EC")
    for y in (2021, 2022, 2023):
        ws = wb.create_sheet(str(y))
        head = OLD_HEAD_2021 if y == 2021 else OLD_HEAD
        ws["A1"] = "Millrind community review: monthly queue report"
        ws["A1"].font = Font(name="Calibri", size=13, bold=True)
        ws["A2"] = ("Calendar year %d. Months on UTC. Source: review system monthly statistics job (%s)."
                    % (y, "final for the year" if y < 2023 else "final; the review system was retired on 31 December 2023"))
        ws["A2"].font = Font(name="Calibri", size=9, italic=True, color="5F6670")
        row = 4
        for c, h in enumerate(head, start=1):
            cell = ws.cell(row=row, column=c, value=h)
            cell.font = Font(name="Calibri", size=10, bold=True)
            cell.fill = head_fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=thin)
        ws.row_dimensions[row].height = 30
        totals = [0] * (len(head) - 2)
        for old in OLD_ORDER:
            for m in range(1, 13):
                if (old, y, m) not in cells:
                    continue
                P = cells[(old, y, m)]
                row += 1
                vals = [OLD_SITE[old], "%s %d" % (K.MON3[m], y), P["reported"]]
                for x in (["close", "first_post", "low_quality"] + ([] if y == 2021 else ["edit"])):
                    vals += [P["ent"][x], P["rev"][x], P["rout"][x]]
                vals += [P["R"], P["vol"], P["flags"]]
                for c, v in enumerate(vals, start=1):
                    ws.cell(row=row, column=c, value=v).font = Font(name="Calibri", size=10)
                for k, v in enumerate(vals[2:]):
                    totals[k] += v
        row += 1
        ws.cell(row=row, column=1, value="All sites").font = Font(name="Calibri", size=10, bold=True)
        ws.cell(row=row, column=2, value="Year %d" % y).font = Font(name="Calibri", size=10, bold=True)
        for k, v in enumerate(totals):
            ws.cell(row=row, column=3 + k, value=v).font = Font(name="Calibri", size=10, bold=True)
        ws.cell(row=row + 2, column=1, value=(
            "Questions: questions posted on the site in the month. Active reviewers: members with one or more reviews in "
            "the month. Volunteer reviews: reviews completed in the month. To moderators: items handed to the moderation "
            "queue.")).font = Font(name="Calibri", size=9, italic=True, color="5F6670")
        if y == 2023:
            ws.cell(row=row + 3, column=1, value=(
                "Graphics & Shaders was folded into Game Development on 1 July 2023; from July the Game Development rows "
                "carry both.")).font = Font(name="Calibri", size=9, italic=True, color="5F6670")
        if y == 2022:
            r0 = row + 5
            ws.cell(row=r0, column=1, value="Questions restated, May to August 2022").font = Font(
                name="Calibri", size=10, bold=True)
            ws.cell(row=r0 + 1, column=1, value=(
                "Statistics job rerun on 14 March 2023 for the tag-sync fault of 9 May to 22 August 2022, which left "
                "questions posted through the API out of the question counts above. The rerun counts on the same basis "
                "as the table above, and its questions replace those above for these months; review items, reviewers "
                "and flags were not affected.")).font = Font(name="Calibri", size=9, italic=True, color="5F6670")
            mh = ["Site"] + ["%s 2022" % K.MON3[m] for m in TAGSYNC_MONTHS]
            for c, h in enumerate(mh, start=1):
                cell = ws.cell(row=r0 + 2, column=c, value=h)
                cell.font = Font(name="Calibri", size=10, bold=True)
                cell.fill = head_fill
            for k, old in enumerate(OLD_ORDER):
                ws.cell(row=r0 + 3 + k, column=1, value=OLD_SITE[old]).font = Font(name="Calibri", size=10)
                for c, m in enumerate(TAGSYNC_MONTHS, start=2):
                    ws.cell(row=r0 + 3 + k, column=c, value=cells[(old, 2022, m)]["questions"]).font = Font(
                        name="Calibri", size=10)
        ws.column_dimensions["A"].width = 24
        ws.column_dimensions["B"].width = 10
        for c in range(3, len(head) + 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 12
        ws.freeze_panes = "C5"
    wb.save(path)


# ------------------------------------------------------------------------------------------------- the plan
def write_plan(path, W, plan):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Plan 2027"
    ws["A1"] = "Millrind 2027 question plan, by site and month"
    ws["A1"].font = Font(name="Calibri", size=13, bold=True)
    ws["A2"] = "Planning & Analytics, issued 18 September 2026. Questions asked, counted as the posts warehouse counts them."
    ws["A2"].font = Font(name="Calibri", size=9, italic=True, color="5F6670")
    head = ["Site code", "Site"] + [K.MON3[m] + " 2027" for m in range(1, 13)] + ["Year 2027"]
    fill = PatternFill("solid", fgColor="DCE3EC")
    for c, h in enumerate(head, start=1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.font = Font(name="Calibri", size=10, bold=True)
        cell.fill = fill
    row = 4
    tot = [0] * 13
    for s in K.SITES:
        row += 1
        vals = [s, K.SITE[s]["name"]] + [plan[s][m] for m in range(1, 13)] + [sum(plan[s].values())]
        for c, v in enumerate(vals, start=1):
            ws.cell(row=row, column=c, value=v).font = Font(name="Calibri", size=10)
        for k in range(13):
            tot[k] += vals[2 + k]
    row += 1
    ws.cell(row=row, column=2, value="Network").font = Font(name="Calibri", size=10, bold=True)
    for k in range(13):
        ws.cell(row=row, column=3 + k, value=tot[k]).font = Font(name="Calibri", size=10, bold=True)
    ws.column_dimensions["A"].width = 11
    ws.column_dimensions["B"].width = 24
    for c in range(3, 16):
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 10
    ws3 = wb.create_sheet("Outlook Oct-Dec 2026")
    ws3["A1"] = "Outlook for the rest of 2026, issued with the plan"
    ws3["A1"].font = Font(name="Calibri", size=11, bold=True)
    head3 = ["Site code", "Site", "Oct 2026", "Nov 2026", "Dec 2026"]
    for c, h in enumerate(head3, start=1):
        cell = ws3.cell(row=3, column=c, value=h)
        cell.font = Font(name="Calibri", size=10, bold=True)
        cell.fill = fill
    for k, s in enumerate(K.SITES):
        vals = [s, K.SITE[s]["name"]] + [W.sites[s].outlook[m] for m in (10, 11, 12)]
        for c, v in enumerate(vals, start=1):
            ws3.cell(row=4 + k, column=c, value=v).font = Font(name="Calibri", size=10)
    ws3.column_dimensions["B"].width = 24
    ws2 = wb.create_sheet("2026 actual")
    ws2["A1"] = "Questions asked January to August 2026 (posts warehouse, pulled 14 September 2026)"
    ws2["A1"].font = Font(name="Calibri", size=11, bold=True)
    head2 = ["Site code", "Site"] + [K.MON3[m] + " 2026" for m in range(1, 9)]
    for c, h in enumerate(head2, start=1):
        cell = ws2.cell(row=3, column=c, value=h)
        cell.font = Font(name="Calibri", size=10, bold=True)
        cell.fill = fill
    for k, s in enumerate(K.SITES):
        vals = [s, K.SITE[s]["name"]] + [warehouse_month(W, s, 2026, m) for m in range(1, 9)]
        for c, v in enumerate(vals, start=1):
            ws2.cell(row=4 + k, column=c, value=v).font = Font(name="Calibri", size=10)
    ws2.column_dimensions["B"].width = 24
    wb.save(path)


# ------------------------------------------------------------------------------------------------- register
REGISTER_COLS = ["site_code", "site_name", "review_system_name_to_2023", "wfm_site_label", "launched", "status",
                 "status_note"]


def register_rows():
    launched = {"web": "2012-03-05", "mobile": "2012-09-17", "data": "2014-04-28", "databases": "2012-11-12",
                "devops": "2013-06-03", "gamedev": "2013-02-11", "security": "2013-10-21", "embedded": "2015-05-18",
                "scicomp": "2016-01-25"}
    rows = []
    for s in K.SITES:
        rows.append([s, K.SITE[s]["name"], K.SITE[s]["old"], AK.WFM_SITE_LABEL[s], launched[s], "live",
                     "absorbed graphics (Graphics & Shaders) on 2023-07-01" if s == "gamedev" else ""])
    rows.append(["graphics", K.GRAPHICS_OLD, K.GRAPHICS_OLD, "", "2014-08-04", "merged",
                 "folded into gamedev on 2023-07-01; questions, answers and reviewers moved to Game Development"])
    return rows


# ------------------------------------------------------------------------------------------------- per-flag export
FLAG_COLS = ["flag_id", "site", "post_id", "post_type", "reason", "flagger_account_id", "raised_at_utc",
             "handled_at_utc", "handled_by", "outcome"]
REVIEW_COLS = ["flag_id", "site", "first_action_at_utc", "first_action_by", "first_outcome", "review_opened_at_utc",
               "review_closed_at_utc", "reviewed_by", "final_outcome", "opened_because"]


def flag_export_rows(flags):
    """Both export parts. The moderator tool marks a flag handled when a moderator acts on it, and again when a
    decision review closes, so a reviewed flag's handled time is the review's."""
    def row(f):
        handled = f["review"]["closed"] if (f["review"] and f["review"]["closed"]) else f["first"]
        return [f["flag_id"], f["site"], f["post_id"], f["post_type"], f["reason"], f["flagger"], ts(f["raised"]),
                ts(handled), f["by"], f["outcome"]]
    p1 = [row(f) for f in flags if f["raised"].date() <= dt.date(2026, 8, 15)]
    p2 = [row(f) for f in flags if f["raised"].date() >= dt.date(2026, 8, 15)]
    return p1, p2


def review_rows(flags):
    out = []
    for f in flags:
        if not f["review"]:
            continue
        R = f["review"]
        out.append([f["flag_id"], f["site"], ts(f["first"]), f["by"], R["first_outcome"], ts(R["opened"]),
                    ts(R["closed"]) if R["closed"] else "", R["reviewer"] or "", R["final"] or "", R["why"]])
    out.sort(key=lambda x: x[5])
    return out


# ------------------------------------------------------------------------------------------------- WFM and HR
WFM_COLS = ["line_id", "employee_no", "site", "shift_date", "start_local", "end_local", "paid_hours", "line_status",
            "approved_on"]
ASSIGN_COLS = ["employee_no", "site", "effective_from", "effective_to"]


def wfm_rows(mods):
    """The WFM system's shift lines, Toronto time, one per worked shift; the site is the moderator's site on their
    current record. The week of 4 to 10 May was re-sent for eight moderators: the original lines stay, superseded."""
    fixes = AK.corrected_lines(mods)
    r = K.rng("wfmlines")
    lines = []
    for m in mods:
        cur = AK.site_on(m, K.HIST_END)
        for sh in m.shifts:
            d = sh["day"]
            appr = d + dt.timedelta(days=(7 - d.weekday()) % 7 + 1)
            base = dict(emp=m.emp, site=AK.WFM_SITE_LABEL[cur], day=d, start=sh["start"], hours=sh["hours"],
                        status="approved", appr=appr)
            if m.emp in fixes and d in fixes[m.emp]:
                orig = dict(base, status="superseded")
                lines.append(orig)
                new_h = fixes[m.emp][d]
                corr = dict(base, hours=new_h, appr=dt.date(2026, 5, 20))
                lines.append(corr)
            else:
                lines.append(base)
    lines.sort(key=lambda x: (x["day"], x["start"], x["emp"], x["status"] != "superseded"))
    lid = 8_800_000 + int(r.integers(100, 900))
    out = []
    for L in lines:
        lid += int(r.integers(1, 4))
        end = L["start"] + dt.timedelta(hours=L["hours"] + 0.5)
        out.append([lid, L["emp"], L["site"], L["day"].isoformat(), L["start"].strftime("%Y-%m-%d %H:%M"),
                    end.strftime("%Y-%m-%d %H:%M"), "%.2f" % L["hours"], L["status"], L["appr"].isoformat()])
    return out


def assign_rows(mods):
    rows = []
    for m in mods:
        segs = m.assign
        for k, (eff, s) in enumerate(segs):
            to = (segs[k + 1][0] - dt.timedelta(days=1)).isoformat() if k + 1 < len(segs) else ""
            start = eff if k else dt.date(2025, 1, 1) if m.emp < "QM04200" else dt.date(2025, 6, 2)
            rows.append([m.emp, K.SITE[s]["name"], start.isoformat(), to])
    rows.sort(key=lambda x: (x[0], x[2]))
    return rows


# ------------------------------------------------------------------------------------------------- moves
MOVE_COLS = ["question_id", "new_question_id", "from_site", "to_site", "asked_at_utc", "moved_at_utc"]


def move_rows(moves):
    return [[mv["qid_old"], mv["qid_new"], OLD_SITE[mv["src"]], OLD_SITE[mv["dst"]], ts(mv["asked"]), ts(mv["moved"])]
            for mv in moves]


# ------------------------------------------------------------------------------------------------- page views
PV_COLS = ["site", "month", "page_views", "unique_visitors"]
PV_DROP = {"web": 0.41, "mobile": 0.38, "data": 0.33, "databases": 0.29, "devops": 0.27, "gamedev": 0.34,
           "security": 0.22, "embedded": 0.19, "scicomp": 0.17}


def pageview_rows(W):
    r = K.rng("pageviews")
    rows = []
    for (y, m) in K.months(K.D0, K.HIST_END):
        for s in K.SITES:
            S = W.sites[s]
            mid = dt.date(y, m, 15)
            x = (K.month_frac(mid) - K.month_frac(dt.date(2023, 9, 1))) / 4.5
            lv = 1.0 - PV_DROP[s] * K.logistic(x)
            pv = S.base * 31.0 * lv * (0.92 + 0.08 * K.SEAS[m]) * math.exp(float(r.normal(0, 0.02)))
            uv = pv / (2.35 + 0.15 * float(r.random()))
            rows.append([s, "%d-%02d" % (y, m), int(pv), int(uv)])
    return rows
