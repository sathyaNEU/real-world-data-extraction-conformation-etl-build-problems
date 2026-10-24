"""The golden deliverables: the settlement workbook and the billing system's credit memo import file."""
import csv

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import world as W

HEAD = ["Partner ID", "District", "State", "Order Form", "Eligible Students Assessed", "Guarantee Benchmark (%)",
        "Outcomes Guarantee", "Credit Percentage (%)", "2025-26 Subscription Fee", "Outcomes Credit",
        "Live PL Hours Contracted", "Live PL Hours Delivered", "PL Shortfall Hours", "PL Credit", "Total Credit",
        "Credit Applied To"]
WIDTH = (10, 44, 6, 13, 10, 11, 12, 10, 13, 12, 10, 10, 10, 11, 12, 16)


def workbook(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Settlement 2025-26"
    bold = Font(bold=True)
    ws["A1"] = "Lumenfold Learning, Inc.: 2025-26 outcomes guarantee and professional learning settlement"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = ("Partnerships Finance, for the settlement statements due 30 November 2026. Credits issue on 30 November "
                "2026 against each district's open 2026-27 subscription invoice, or by check.")
    ws.append([])
    ws.append(HEAD)
    hdr = 4
    fill = PatternFill("solid", fgColor="DCE6F1")
    for c in ws[hdr]:
        c.font = bold
        c.fill = fill
        c.alignment = Alignment(wrap_text=True, vertical="bottom")
    ws.row_dimensions[hdr].height = 42
    for x in sorted(rows, key=lambda r: (r["state"], r["name"])):
        ws.append([x["partner_id"], x["name"], x["state"], x["order_form"], x["assessed"], x["benchmark"],
                   x["result"], x["credit_pct"], x["fee"], x["og"], x["pd_contract"], x["pd_delivered"], x["pd_short"],
                   x["pl"], x["og"] + x["pl"],
                   (x["invoice"] or "Refund by check") if x["og"] + x["pl"] > 0 else None])
    last = ws.max_row
    ws.append([None, "Total, %d districts" % len(rows), None, None, "=SUM(E%d:E%d)" % (hdr + 1, last), None,
               '=COUNTIF(G%d:G%d,"Not met")&" not met"' % (hdr + 1, last), None, "=SUM(I%d:I%d)" % (hdr + 1, last),
               "=SUM(J%d:J%d)" % (hdr + 1, last), "=SUM(K%d:K%d)" % (hdr + 1, last), "=SUM(L%d:L%d)" % (hdr + 1, last),
               "=SUM(M%d:M%d)" % (hdr + 1, last), "=SUM(N%d:N%d)" % (hdr + 1, last), "=SUM(O%d:O%d)" % (hdr + 1, last),
               None])
    top = Side(style="thin")
    for c in ws[ws.max_row]:
        c.font = bold
        c.border = Border(top=top)
    for row in ws.iter_rows(min_row=hdr + 1, max_row=ws.max_row):
        row[4].number_format = "#,##0"
        row[5].number_format = "0.0"
        row[7].number_format = "0"
        for k in (8, 9, 13, 14):
            row[k].number_format = "#,##0.00"
        for k in (10, 11, 12):
            row[k].number_format = "0.0"
    for i, w in enumerate(WIDTH, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C%d" % (hdr + 1)
    ws.print_title_rows = "%d:%d" % (hdr, hdr)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    n = wb.create_sheet("Basis")
    lines = [
        ("Outcomes guarantee (MSA section 9)", True),
        ("Eligible Students: the 1 May 2026 roster records flagged eligible in the guarantee eligibility file, one "
         "record per student; Eligible Students Assessed counts those with a submitted spring 2026 session.", False),
        ("Outcome Share: each district's share of its Eligible Students Assessed at or above Proficient for their grade "
         "(500, 520, 540), estimated from their own scored responses with a latent regression on district, grade, "
         "prior state-test band and the three status flags (2024 item parameters fixed), as the mean posterior "
         "probability of being at or above the grade's level. The contractor's plausible values come from a model "
         "with no district term and are not used for district results.", False),
        ("A district whose Outcome Share is below its Guarantee Benchmark is credited its Credit Percentage of the "
         "2025-26 Subscription Fee.", False),
        ("Professional learning (MSA section 8)", True),
        ("Live sessions (virtual live and on site) with status Delivered and a session date in the district's "
         "2025-26 Subscription Term, each at the time it ran as the scheduling system records it, hours and minutes "
         "after the point (every offering is scheduled in whole hours; a 1.40 is an hour and forty minutes); joint "
         "sessions count in full for each district attending; the district's total to the nearest half hour.", False),
        ("PL Credit: USD 250 for each hour of the shortfall against the hours on the order form (MSA 8.3).", False),
        ("Credits (MSA section 10.1)", True),
        ("Applied to the open 2026-27 subscription invoice on the bill-to account with the earliest due date in the "
         "register of 20 November 2026; refunded by check where none is open.", False),
        ("Sources: order_forms_2025_26.xlsx, guarantee_eligibility_2025_26.csv, the LSOA session and response "
         "exports, caldrith_item_parameters_2024.csv, district_roster_uploads_spring2026.xlsx, "
         "pd_sessions_2025_26.csv, billing_invoices_2026_27.csv.", False),
    ]
    for t, b in lines:
        n.append([t])
        n.cell(n.max_row, 1).font = Font(bold=b)
        n.cell(n.max_row, 1).alignment = Alignment(wrap_text=True)
    n.column_dimensions["A"].width = 120
    wb.save(path)


def memos_csv(path, memos):
    with open(path, "w", newline="") as f:
        w = csv.writer(f, lineterminator="\r\n")
        w.writerow(["memo_reference", "customer_account", "reason_code", "credit_amount", "apply_to_invoice",
                    "refund_method", "credit_date"])
        for m in memos:
            w.writerow([m["memo_reference"], m["customer_account"], m["reason_code"], "%.2f" % m["credit_amount"],
                        m["apply_to_invoice"], m["refund_method"], m["credit_date"]])
