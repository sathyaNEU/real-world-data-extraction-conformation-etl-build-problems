"""The golden deliverables: the board's elections workbook (each practice's projected 2026 network measure against its
threshold and its election, the level 5 paid share, notes) and the of each practice's projected figure against its threshold; returns every figure the write-up states."""
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import matplotlib.ticker  # noqa: E402,F401
import openpyxl  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

import golden as G  # noqa: E402

BOLD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="DCE6F1")
TOTAL_FILL = PatternFill("solid", fgColor="F2F2F2")
ENTER_FILL = PatternFill("solid", fgColor="E2EFDA")
WORDS = {9: "Nine", 10: "Ten", 8: "Eight", 11: "Eleven", 7: "Seven"}


def _head(ws, row):
    for c in ws[row]:
        c.font = BOLD
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")


def short(name):
    return name.split()[0]


def listing(xs):
    xs = list(xs)
    return xs[0] if len(xs) == 1 else ", ".join(xs[:-1]) + " and " + xs[-1]


def level5_counts(p):
    """Each practice's established-patient visits and the visits paid at level 5, on the golden's paid levels."""
    tot = {}
    for prac, payer, code, lvl in G.paid_levels(p, "golden"):
        o = tot.setdefault(prac, [0, 0])
        o[0] += 1
        o[1] += int(lvl == 5)
    return {k: tuple(v) for k, v in tot.items()}


def write(path, p):
    g = G.golden(p)
    order = p["practices"]
    enter = [k for k in order if g[k]["enters"]]
    stay = [k for k in order if not g[k]["enters"]]
    a_counts = level5_counts(p)
    a = G.ask_a(p)
    held = G.open_stays(p, 2026)
    rates = G.stay_rates(p)
    dflt = G.default(p)
    runout = G.golden(p, stays=False)
    r26, hc = G.run_members(p, "run2026"), {}
    for mk in held:
        hc[r26[mk]] = hc.get(r26[mk], 0) + 1
    figs = dict(call=dict(enter=enter, stay=stay), ask_a={},
                parameters=dict(threshold={}, benchmark={}, position_pct={},
                                steps=dict(default_enter=[k for k in order if dflt[k]["enters"]],
                                           runout_enter=[k for k in order if runout[k]["enters"]],
                                           held=[[k, hc[k]] for k in order if k in hc],
                                           open_2025=len(G.open_stays(p, 2025)))),
                constants=dict(order=order, trend=G.TREND, threshold_share=98, cap=int(G.CAP), codes=["99211", "99213", "99214", "99215"],
                               sections=["2", "3", "3.4", "4.1"]))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Elections 2027"
    ws.append(["%s of the thirteen practices enter the 2027 two-sided track: %s" % (WORDS[len(enter)],
                                                                                   listing([short(k) for k in enter]))])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append(["Board meeting of 29 October 2026, Resolution 2026-11. Each practice's 2026 network measure on the payers' "
               "final 2026 attribution, projected from the October interim runs, against 98 per cent of its 2026 "
               "benchmark; USD per attributed member month."])
    ws.append([])
    ws.append(["Practice", "Practice TIN", "2025 network measure, final 2025 runs", "2026 benchmark (x 1.052)",
               "Threshold (98% of benchmark)", "Interim measure, October 2026 runs", "2025 run-out at the same cut",
               "Members in hospital on 31 August, carried at the cap", "Projected 2026 network measure",
               "Projected against threshold", "Election"])
    _head(ws, 4)
    tin = {v: k for k, v in p["tin2prac"].items()}
    for k in order:
        x = g[k]
        gap = x["projected"] / x["threshold"] - 1
        ws.append([k, "%s-%s" % (tin[k][:2], tin[k][2:]), round(x["f2025"], 2), round(x["benchmark"], 2),
                   round(x["threshold"], 2), round(x["interim"], 2), round(x["runout"], 4), x["open_at_cut"] or None,
                   round(x["projected"], 2), gap, "Two-sided" if x["enters"] else "One-sided"])
        if x["enters"]:
            for c in ws[ws.max_row]:
                c.fill = ENTER_FILL
        figs["parameters"]["threshold"][k] = round(x["threshold"], 2)
        figs["parameters"]["benchmark"][k] = round(x["benchmark"], 2)
        figs["parameters"]["position_pct"][k] = round(100 * gap, 1)
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row):
        for c in (row[2], row[3], row[4], row[5], row[8]):
            c.number_format = "#,##0.00"
        row[6].number_format = "0.0000"
        row[7].number_format = "0"
        row[9].number_format = "+0.0%;-0.0%;0.0%"
    for i, w in enumerate([31, 12, 14, 13, 13, 14, 12, 14, 14, 12, 11], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[4].height = 58
    ws.freeze_panes = "B5"

    la = wb.create_sheet("Level 5 visits")
    la.append(["Established-patient visits by the practices' clinicians paid at level 5, January to August 2026"])
    la["A1"].font = BOLD
    la.append(["Visits billed 99211 to 99215 in person or by telehealth, at the claim's current version, by the clinician's "
               "practice on the date of service; the paid level is the level of the payer's 2026 rate for the visit's "
               "setting that the claim's allowed amount reflects, every line of the claim counted."])
    la.append([])
    la.append(["Practice", "Established-patient visits", "Paid at level 5", "Share paid at level 5"])
    _head(la, 4)
    for k in order:
        n, n5 = a_counts[k]
        la.append([k, n, n5, n5 / n])
        figs["ask_a"][k] = a[k]
    tn = sum(v[0] for v in a_counts.values())
    t5 = sum(v[1] for v in a_counts.values())
    la.append(["Network", tn, t5, t5 / tn])
    for c in la[la.max_row]:
        c.font = BOLD
        c.fill = TOTAL_FILL
    for row in la.iter_rows(min_row=5):
        row[1].number_format = "#,##0"
        row[2].number_format = "#,##0"
        row[3].number_format = "0.0%"
    for i, w in enumerate([31, 14, 12, 12], 1):
        la.column_dimensions[get_column_letter(i)].width = w
    la.row_dimensions[4].height = 30
    la.freeze_panes = "B5"

    kinds = {}
    for mk, (kind, first, days) in held.items():
        kinds.setdefault((kind, first), []).append(days)
    count = {1: "one", 2: "two", 3: "three", 4: "four", 5: "five", 7: "Seven"}
    stays_line = "; ".join("%s in %s from %s, %d days to the cut" % (
        count[len(v)], {"NICU": "a neonatal unit", "ICU": "intensive care"}[kd], f.strftime("%-d %B"), v[0])
        for (kd, f), v in sorted(kinds.items(), key=lambda kv: kv[0][1]))
    no = wb.create_sheet("Notes")
    notes = [
        "Prepared by Payer Performance Analytics for the Board's meeting of 29 October 2026; elections due with the "
        "payers by 30 October 2026.",
        "Network measure (Performance Standard, section 3): allowed amounts, at each claim's current version, for covered "
        "medical and pharmacy services incurred in the year and first paid by 31 March of the next year, each member's "
        "year capped at USD 250,000, over the members' months of coverage, pooled across Merrinlake, Kesselcrest and "
        "Pemberglen.",
        "Benchmark (section 3.4): the 2025 network measure on the payers' final 2025 runs, rounded to the cent, times "
        "1.052, rounded to the cent. Threshold: 98 per cent of the benchmark; a practice enters at or under it "
        "(Resolution 2026-11).",
        "Interim measure (section 4.1): claims incurred January to August 2026 paid by each payer's October feed, "
        "divided by the payer's completion factor for the month incurred, for the members on the October runs.",
        "Run-out: the October runs attribute only the members whose qualifying visit had been paid by 30 September, and "
        "the final runs will add the members whose first qualifying visit of 2026 comes later, for every month of their "
        "coverage. Each practice's interim measure is carried by its own 2025 run-out at the same cut: its 2025 network "
        "measure over the measure, January to August 2025, of the members the standard's attribution assigns to it on "
        "qualifying services incurred by 31 August 2025 and paid by 30 September 2025.",
        "Members in hospital at the cut: a hospital bills a stay when it ends, so the October feeds hold a stay still "
        "running on 31 August only as its attending physicians' daily visits. %s (%s) had passed the USD 250,000 cap by "
        "the cut at the lowest per-day amount any closed stay of 30 days or more of the same kind was allowed (USD %s a day "
        "in neonatal care, USD %s in intensive care); each is carried at the cap with its 2026 months of coverage. No "
        "stay of the kind was running at the 2025 cut, and the completion factors leave out claims over USD 100,000 "
        "(Appendix C), so neither the run-out nor the factors carry these stays." % (
            "%s members" % count[len(held)], stays_line,
            "{:,.0f}".format(rates["NICU"]), "{:,.0f}".format(rates["ICU"])),
        "Level 5 paid share: from 1 March 2026 Kesselcrest's level review pays some level 5 visits at its level 4 rate, "
        "the visit's line keeping the level 5 rate and the claim taking the difference on a line with no procedure code, "
        "so such a visit is level 4; Merrinlake allows an E&M billed with modifier 25 beside a minor procedure by the "
        "same clinician that day at half its rate, so such a visit is at the level whose rate is twice its allowed amount "
        "(a halved level 5 equals Merrinlake's level 3 rate), while a modifier 25 visit beside a preventive visit is "
        "paid in full.",
        "Prices are USD per attributed member month, rounded to the cent; the elections are made on unrounded figures.",
    ]
    for t_ in notes:
        no.append([t_])
        no.cell(no.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    no.column_dimensions["A"].width = 112
    wb.save(path("workbook"))

    chart(path("chart"), p, g, enter)
    return figs


def chart(path, p, g, enter):
    """Each practice's projected 2026 measure against its threshold, as the per cent above or below it, the interim
    measure as an open marker beside it; ordered by the projection."""
    green, grey, ink, ink2, line = "#2e7d4f", "#9a9893", "#0b0b0b", "#52514e", "#b9b8b3"
    names = sorted(p["practices"], key=lambda k: g[k]["projected"] / g[k]["threshold"])
    y = list(range(len(names)))[::-1]
    proj = [100 * (g[k]["projected"] / g[k]["threshold"] - 1) for k in names]
    inter = [100 * (g[k]["interim"] / g[k]["threshold"] - 1) for k in names]
    fig, ax = plt.subplots(figsize=(8.8, 5.4), dpi=150)
    cols = [green if k in enter else grey for k in names]
    ax.barh(y, proj, height=0.56, color=cols, zorder=2)
    dots, = ax.plot(inter, y, linestyle="none", marker="o", markersize=5.5, markerfacecolor="white",
                    markeredgecolor=ink2, markeredgewidth=1.2, zorder=3)
    ax.axvline(0, color=ink, linewidth=1.0, zorder=4)
    lo, hi = min(proj + inter), max(proj + inter)
    ax.set_xlim(lo - 3, hi + 3)
    ax.text(0.3, max(y) + 0.75, "threshold: 98% of the 2026 benchmark", fontsize=7.6, color=ink, va="bottom")
    for yi, v in zip(y, proj):
        if v >= 0:                       # inside the bar, clear of the interim marker beyond it
            ax.text(v - 0.3, yi, "%+.1f%%" % v, va="center", ha="right", fontsize=7.8, color="white", fontweight="bold")
        else:
            ax.text(v - 0.35, yi, "%+.1f%%" % v, va="center", ha="right", fontsize=7.8, color=ink)
    ax.set_yticks(y)
    ax.set_yticklabels(names, fontsize=8.6, color=ink)
    ax.set_ylim(min(y) - 0.6, max(y) + 1.2)
    ax.xaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: "%+.0f%%" % v if v else "0%"))
    ax.set_xlabel("Per cent above (+) or below (-) the practice's threshold", fontsize=8.6, color=ink2)
    ax.tick_params(axis="x", labelsize=8, colors=ink2)
    ax.tick_params(axis="y", length=0)
    ax.grid(axis="x", color="#e4e3df", linewidth=0.8, zorder=0)
    for s_ in ("top", "right", "left"):
        ax.spines[s_].set_visible(False)
    ax.spines["bottom"].set_color(line)
    import matplotlib.patches as mpatches
    handles = [mpatches.Patch(color=green, label="Projected 2026 measure, enters the two-sided track"),
               mpatches.Patch(color=grey, label="Projected 2026 measure, stays one-sided"), dots]
    dots.set_label("Interim measure on the October runs")
    ax.legend(handles=handles, loc="lower left", bbox_to_anchor=(-0.38, 1.04), ncol=3, frameon=False, fontsize=7.4,
              handlelength=1.3, borderaxespad=0, columnspacing=1.2)
    fig.suptitle("%s practices clear their 2026 thresholds; the stays at Mallowcrest and Garrowdale keep them out"
                 % WORDS[len(enter)], x=0.012, y=0.975, ha="left", fontsize=10.2, color=ink)
    fig.text(0.012, 0.012, "Larrowbrook Physician Network, Payer Performance Analytics. Projected from the October 2026 "
             "runs by each practice's 2025 run-out; members in hospital on 31 August at the USD 250,000 cap.",
             fontsize=6.6, color=ink2)
    fig.subplots_adjust(left=0.27, right=0.97, top=0.84, bottom=0.13)
    fig.savefig(path, facecolor="white", metadata={"Software": None, "Title": "Practice cost against benchmark 2026",
                                                   "Author": "Larrowbrook Physician Network"})
    plt.close(fig)
