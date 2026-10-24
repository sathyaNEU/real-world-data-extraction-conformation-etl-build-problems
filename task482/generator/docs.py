"""The pack's documents: the cap liability policy (docx), the December 2026 census report, the FY2026 rate sheet and
the revenue cycle file notes (pdf), and two workbooks the analysis does not need (quality measures, the 2027 budget)."""
import datetime as dt

import docx
from docx.shared import Pt
import openpyxl
from openpyxl.styles import Font
import reportlab.rl_config
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import world as W
import records as R

ORG = "Wendleglen Hospice & Palliative Care"
reportlab.rl_config.invariant = 1


def _core(d, author, when):
    cp = d.core_properties
    cp.author = author
    cp.last_modified_by = author
    cp.created = when
    cp.modified = when
    cp.revision = 4
    cp.title = ""
    cp.comments = ""


POLICY = [
    ("h", "1. Purpose and scope"),
    ("p", "Medicare limits the payments each hospice provider number receives for a cap year, 1 October to 30 September, "
          "to its aggregate cap (42 CFR 418.309). At each year-end close Wendleglen estimates the repayment each of its "
          "provider numbers will owe for the cap year that ended on the preceding 30 September, accrues it, and the "
          "external auditors test the estimate. This policy covers the eight provider numbers in the provider register, "
          "Fallowmere Hospice (provider number 361798) included: Wendleglen took over its provider agreement on 1 March "
          "2025 together with its records, and its cap years before and after that date are one provider number's."),
    ("h", "2. The count and the cap"),
    ("p", "2.1 A provider number's aggregate cap for FY2026 is the FY2026 cap amount, USD 35,361.44 per beneficiary under "
          "the FY2026 hospice payment rule, times the number of Medicare beneficiaries counted for it under the "
          "proportional method."),
    ("p", "2.2 A Medicare beneficiary is counted for a provider number if the provider number gave them hospice care on any "
          "day of the cap year. The beneficiary counts as the fraction that their days of hospice care at the provider "
          "number in the cap year are of their days of hospice care in every hospice and in every year."),
    ("p", "2.3 A beneficiary's days of hospice care in every hospice and year are not all known at the close. The estimate "
          "takes them as the days on file through 31 December plus the days the beneficiary is expected to receive after "
          "that date, given what Wendleglen knows of the beneficiary at the close, and the fraction is the beneficiary's "
          "days at the provider number in the cap year divided by that total."),
    ("p", "2.4 What a beneficiary is expected to receive is read from the experience since 1 October 2016 of the "
          "beneficiaries cared for by the provider number at which Wendleglen cared for the beneficiary last, in the "
          "beneficiary's diagnosis group (the four groups of the monthly census report)."),
    ("p", "2.5 The days on file are the days of hospice care in the Medicare benefit periods extract, which holds every "
          "benefit period of every beneficiary the provider numbers have cared for since 1 October 2016, at any hospice "
          "and in any year, as Medicare holds them at the data date."),
    ("h", "3. Payments and the repayment"),
    ("p", "3.1 A provider number's payments for the cap year are the Medicare payments for its claims with statement dates "
          "in the cap year, net of reversals, before the 2 per cent sequestration reduction. Interest paid on claims "
          "paid late is not a payment for care and is left out."),
    ("p", "3.2 The repayment is estimated as the contractor determines it: the payments less the aggregate cap, the "
          "difference reduced by 2 per cent for the sequestration already withheld from the payments. A provider number "
          "whose payments are below its cap owes nothing."),
    ("h", "4. The inpatient cap"),
    ("p", "Under 42 CFR 418.302(f) the days of general inpatient and inpatient respite care Medicare pays a provider number "
          "for in a cap year may not exceed 20 per cent of its days of hospice care for Medicare beneficiaries in the "
          "year. The year-end review reports each provider number's inpatient days and inpatient share beside the "
          "aggregate cap estimate."),
    ("h", "5. Records"),
    ("p", "The estimate is built from the Medicare benefit periods extract, Wendleglen's EMR admissions export (Wendleglen's "
          "programs since 1 October 2016 and Fallowmere from 1 March 2025), Fallowmere's census register as exported by "
          "the seller at the change of ownership, the HICN to MBI crosswalk for records keyed before the MBI, the claim "
          "lines and the remittance register. Diagnosis groups follow the census report's code list."),
    ("h", "6. Presentation"),
    ("p", "Repayments are reported in USD millions to one decimal, each from its unrounded estimate, with their total. The "
          "workpapers keep the schedule of beneficiaries behind each count."),
]


def policy_docx(path):
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    d.add_heading(ORG, level=1)
    d.add_paragraph("Finance Policy FIN-14: Medicare hospice cap liability")
    d.add_paragraph("Owner: Controller. Approved by the Finance Committee, 18 November 2026. Supersedes the revision of "
                    "September 2023.")
    for kind, text in POLICY:
        if kind == "h":
            d.add_heading(text, level=2)
        else:
            d.add_paragraph(text)
    _core(d, "B. Graham", dt.datetime(2026, 11, 18, 9, 12))
    d.save(path)


def _pdf(path, title, story_fn, author):
    doc = SimpleDocTemplate(path, pagesize=letter, leftMargin=0.7 * inch, rightMargin=0.7 * inch,
                            topMargin=0.7 * inch, bottomMargin=0.7 * inch, title=title, author=author, creator=author,
                            subject=title, producer="Wendleglen reporting")
    ss = getSampleStyleSheet()
    story = story_fn(ss)
    doc.build(story)


def _table(rows, widths=None, fs=7.5):
    t = Table(rows, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", fs), ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", fs),
                           ("LINEBELOW", (0, 0), (-1, 0), 0.5, colors.black), ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                           ("TOPPADDING", (0, 0), (-1, -1), 2)]))
    return t


def census_pdf(path, people, closed_2026, adc, admissions, discharges):
    """The programs' operating report for December 2026: census, admissions, discharges, lengths of stay by group."""
    def story(ss):
        s = [Paragraph("%s: Monthly Census Report, December 2026" % ORG, ss["Title"]),
             Paragraph("Prepared by Clinical Operations, 6 January 2027. All payers. Calendar 2026 unless stated.", ss["Normal"]),
             Spacer(1, 8)]
        rows = [["Program", "Provider no.", "ADC Dec 2026", "ADC 2026", "Admissions 2026", "Deaths 2026", "Live discharges 2026"]]
        for code in W.CODES:
            ccn = R.CCN[code]
            rows.append([R.PROGRAM[code], ccn, "%.1f" % adc[ccn][0], "%.1f" % adc[ccn][1], "%d" % admissions[ccn],
                         "%d" % discharges[ccn][0], "%d" % discharges[ccn][1]])
        s += [Paragraph("Census and flow", ss["Heading3"]), _table(rows), Spacer(1, 10)]
        rows = [["Program", "Cancer", "Dementia", "Heart and lung", "Other", "All"]]
        for code in W.CODES:
            ccn = R.CCN[code]
            vals = []
            alls = []
            for g in W.GRP:
                x = closed_2026.get((ccn, W.GROUP_NAME[g]), [])
                alls += x
                vals.append("%.1f" % (sum(x) / len(x)) if x else "")
            vals.append("%.1f" % (sum(alls) / len(alls)) if alls else "")
            rows.append([R.PROGRAM[code]] + vals)
        s += [Paragraph("Average length of stay of 2026 discharges, days, by diagnosis group (Medicare)", ss["Heading3"]),
              _table(rows), Spacer(1, 6),
              Paragraph("Length of stay is counted from admission to the program to discharge from it. Diagnosis groups follow "
                        "the census code list (diagnosis_groups.csv).", ss["Normal"])]
        return s
    _pdf(path, "Monthly Census Report December 2026", story, "Clinical Operations")


def rates_pdf(path):
    def story(ss):
        s = [Paragraph("FY2026 Medicare hospice payment rates and limits", ss["Title"]),
             Paragraph("Reimbursement summary of the FY2026 hospice payment rule (CMS-1835-F, 90 FR 37404, 5 August 2025), "
                       "for claims with dates of service 1 October 2025 to 30 September 2026. National rates for hospices "
                       "submitting the required quality data, before the wage index adjustment.", ss["Normal"]),
             Spacer(1, 8)]
        rows = [["Level of care", "Revenue code", "Rate (USD)", "Unit", "Labor share"],
                ["Routine home care, days 1 to 60", "0651", "230.83", "day", "66.0%"],
                ["Routine home care, days 61 and later", "0651", "181.94", "day", "66.0%"],
                ["Continuous home care", "0652", "69.76", "hour (8 hours a day minimum)", "75.2%"],
                ["Inpatient respite care", "0655", "532.48", "day", "61.0%"],
                ["General inpatient care", "0656", "1,199.86", "day", "63.5%"]]
        s += [_table(rows, fs=8.5), Spacer(1, 10)]
        s += [Paragraph("Each rate is paid as the labor share times the wage index of the location of care plus the rest of the "
                        "rate. Routine home care days are counted from the start of hospice care; the count runs on across a "
                        "gap of 60 days or fewer between elections.", ss["Normal"]), Spacer(1, 6),
              Paragraph("Aggregate cap amount, cap year ending 30 September 2026: USD 35,361.44 per beneficiary.", ss["Normal"]),
              Paragraph("Inpatient limit: general inpatient and respite days may not exceed 20 per cent of a hospice's days of "
                        "care for Medicare beneficiaries in the cap year (42 CFR 418.302(f)).", ss["Normal"]),
              Paragraph("Payments are reduced by 2 per cent under sequestration (claim adjustment reason code 253).", ss["Normal"])]
        return s
    _pdf(path, "FY2026 hospice rates", story, "Reimbursement")


def file_notes_pdf(path, n_extract):
    def story(ss):
        body = ss["Normal"]
        s = [Paragraph("Revenue cycle file notes: cap year-end files", ss["Title"]),
             Paragraph("Revenue Cycle, January 2027. Layouts and field notes for the files pulled for the FY2026 cap review.", body),
             Spacer(1, 8), Paragraph("medicare_hospice_benefit_periods_20270105.txt", ss["Heading3"]),
             Paragraph("Fixed width, one header record (H), one detail record (D) per benefit period per hospice, one trailer "
                       "record (T) with the count of detail records (%s in this pull). Pulled 5 January 2027 for every "
                       "beneficiary the eight provider numbers have cared for since 1 October 2016; data through 31 December "
                       "2026." % format(n_extract, ","), body)]
        rows = [["Field", "Positions", "Notes"],
                ["REC_TYPE", "1", "H, D or T"], ["MBI", "2-12", "Medicare beneficiary identifier"],
                ["BP_NUM", "13-15", "benefit period number in the beneficiary's Medicare record"],
                ["ELECT_DT", "16-23", "date of the election the benefit period belongs to (YYYYMMDD)"],
                ["PROV_CCN", "24-29", "provider number of the hospice for the eight provider numbers; for another hospice the "
                                      "pull's own key for it (X and five digits), its provider number withheld"],
                ["BP_START", "30-37", "first day of the period at this hospice"],
                ["BP_END", "38-45", "last day of the period; the date of death; or the date of a discharge, revocation or "
                                    "transfer; blank while the period is open at the data date"],
                ["BP_DAYS", "46-49", "days of hospice care in the record, through the data date for an open period"],
                ["DOD", "50-57", "date of death where Medicare has one by the data date"]]
        s += [_table(rows, widths=[0.9 * inch, 0.8 * inch, 5.2 * inch], fs=7.5), Spacer(1, 8)]
        s += [Paragraph("wendleglen_emr_admissions.csv", ss["Heading3"]),
              Paragraph("One row per admission to an Wendleglen program (Fallowmere from 1 March 2025; its patients on service "
                        "that day carried over with their admission dates, admission type Conversion). medicare_id is the "
                        "identifier captured at admission (the HICN before April 2018). discharge_date is blank for a patient "
                        "on service at 31 December 2026. All payers.", body),
              Paragraph("fallowmere_census_register_20250228.xlsx", ss["Heading3"]),
              Paragraph("The seller's admission and discharge register, run at the change of ownership. Disch Status is the "
                        "UB-04 patient status: 01 discharged home or to self care, 30 still a patient, 40 expired at home, 41 "
                        "expired in a medical facility, 42 expired, place unknown, 50 discharged to hospice (home), 51 "
                        "discharged to hospice (medical facility). Prim Dx is stored without its decimal point.", body),
              Paragraph("wendleglen_claim_lines_fy2026.csv", ss["Heading3"]),
              Paragraph("One row per revenue line of each claim the billing system submitted through 31 December 2026 for "
                        "statement months October 2025 to September 2026. type_of_bill 0811 admission through discharge, 0812 first claim of a "
                        "continuing stay, 0813 continuing, 0814 last, 0817 replacement of the claim in prior_claim_key, 0818 "
                        "void of it. Revenue codes 0651 routine home care, 0655 inpatient respite and 0656 general inpatient: a "
                        "line for each run of days at the level, line_date the run's first day and units its days; 0652 "
                        "continuous home care: a line per day, units are 15-minute increments. clearinghouse_status is the "
                        "clearinghouse's response to the claim file.", body),
              Paragraph("hicn_mbi_crosswalk.csv", ss["Heading3"]),
              Paragraph("One row per HICN carried by an EMR admission or a census register record, with the MBI of the same "
                        "beneficiary.", body),
              Paragraph("diagnosis_groups.csv", ss["Heading3"]),
              Paragraph("The census report's code list: one row per ICD-10-CM code with its description and census group.", body),
              Paragraph("provider_numbers.json", ss["Heading3"]),
              Paragraph("The provider register: each provider number with its program, the FY2026 wage index applied to its "
                        "claims and the date Wendleglen took it on.", body),
              Paragraph("remittance_register_fy2026.csv", ss["Heading3"]),
              Paragraph("The 835 remittances posted from November 2025 to 31 December 2026 for the FY2026 claims. record_type CLP is "
                        "a claim payment (claim_status 1) or the reversal of an earlier payment (22); paid_amount is the amount "
                        "paid and adj_carc_253 the sequestration withheld (claim adjustment reason code 253). record_type PLB "
                        "is a provider-level adjustment; L6 is interest on a claim paid late.", body)]
        return s
    _pdf(path, "Revenue cycle file notes", story, "Revenue Cycle")


def quality_xlsx(path):
    rng = R.stream(8)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "HQRP 2026"
    hdr = ["Provider number", "Program", "Period", "Hospice Care Index (of 10)", "Visits in last days of life (%)",
           "HIS comprehensive assessment (%)", "CAHPS willing to recommend (%)", "CAHPS overall rating (%)"]
    ws.append(hdr)
    for c in ws[1]:
        c.font = Font(bold=True)
    for code in W.CODES:
        for per in ("CY2025 Q1-Q4", "CY2026 Q1-Q2"):
            ws.append([R.CCN[code], R.PROGRAM[code], per, int(rng.integers(6 if code == "FAL" else 8, 11)),
                       round(float(rng.uniform(38, 62)), 1), round(float(rng.uniform(86, 99)), 1),
                       round(float(rng.uniform(78, 92)), 1), round(float(rng.uniform(76, 90)), 1)])
    wb.properties.creator = "Quality"
    wb.properties.created = dt.datetime(2026, 11, 3, 14, 5)
    wb.properties.modified = dt.datetime(2026, 11, 3, 14, 5)
    wb.save(path)


def budget_xlsx(path, adc):
    rng = R.stream(9)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2027 budget"
    ws.append(["Program", "Provider number", "Budget ADC 2027", "Admissions 2027", "Medicare share of days (%)",
               "Net patient revenue 2027 (USD)"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for code in W.CODES:
        ccn = R.CCN[code]
        a = adc[ccn][1] * float(rng.uniform(0.98, 1.06))
        ws.append([R.PROGRAM[code], ccn, round(a, 1), int(a * 365 / float(rng.uniform(60, 110))),
                   round(float(rng.uniform(88, 95)), 1), int(round(a * 365 * float(rng.uniform(205, 225)), -3))])
    wb.properties.creator = "Finance"
    wb.properties.created = dt.datetime(2026, 10, 21, 10, 30)
    wb.properties.modified = dt.datetime(2026, 12, 2, 16, 48)
    wb.save(path)
