"""Writers for the pack's data files, each in its source system's own layout."""
import csv
import datetime as dt

import numpy as np
import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

import world as W
import pd as PD


def iso(d):
    return d.isoformat() if d else ""


def write_csv(path, header, rows):
    with open(path, "w", newline="") as f:
        w = csv.writer(f, lineterminator="\r\n")
        w.writerow(header)
        for row in rows:
            w.writerow(row)
    return len(rows)


def sessions_csv(path, rows):
    rows = sorted(rows, key=lambda x: (x["started"], x["session_id"]))
    return write_csv(path, ["session_id", "test_taker_id", "form_id", "session_status", "started_at_utc",
                            "submitted_at_utc", "items_presented"],
                     [(x["session_id"], x["test_taker_id"], x["form_id"], x["status"], x["started"], x["submitted"],
                       x["items"]) for x in rows])


def responses_csv(path, rows):
    return write_csv(path, ["session_id", "item_seq", "item_id", "response_status", "score"],
                     [(s, q, it, st, "" if sc is None else sc) for s, q, it, st, sc in rows])


def item_params_csv(path, item_id):
    rows = []
    for m in range(W.N_MOD):
        for k in range(W.N_ITEM):
            rows.append((item_id[(m, k)], "M%d" % (m + 1), k + 1, "2PL", "%.3f" % W.A[m, k], "%.3f" % W.B[m, k],
                         "1.0", "2024 base year"))
    return write_csv(path, ["item_id", "module", "position", "model", "a_discrimination", "b_difficulty",
                            "scaling_constant_D", "calibration"], rows)


def pv_csv(path, v, pv):
    order = np.argsort(v["tt"])
    rows = []
    for i in order:
        dd = W.DISTRICTS[v["dist"][i]]
        rows.append((v["tt"][i], dd["partner_id"], W.GRADES[v["g"][i]], dd["state"]) +
                    tuple("%.1f" % pv[m, i] for m in range(5)))
    return write_csv(path, ["test_taker_id", "partner_id", "grade", "state", "pv1", "pv2", "pv3", "pv4", "pv5"], rows)


def eligibility_csv(path, rows):
    rows = sorted(rows, key=lambda x: (x["partner_id"], x["student_id"]))
    return write_csv(path, ["partner_id", "student_id", "grade", "enrolled_on", "withdrawn_on",
                            "rostered_on_2026_05_01", "lessons_completed_to_2026_05_01", "guarantee_eligible"],
                     [(x["partner_id"], x["student_id"], x["grade"], iso(x["enrolled_on"]), iso(x["withdrawn_on"]),
                       x["rostered_2026_05_01"], x["lessons"], x["eligible"]) for x in rows])


def crosswalk_csv(path, rows):
    rows = sorted(rows, key=lambda x: (x["test_taker_id"], x["student_id"]))
    return write_csv(path, ["student_id", "test_taker_id", "partner_id", "linked_on", "link_source"],
                     [(x["student_id"], x["test_taker_id"], x["partner_id"], iso(x["linked_on"]), x["source"])
                      for x in rows])


def roster_xlsx(path, rows):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    head = ["Partner ID", "Student ID", "Grade", "Economically Disadvantaged", "English Learner",
            "Students with Disabilities", "Prior State Test", "Prior Test Level", "Upload Date"]
    for st in W.STATES:
        ws = wb.create_sheet(W.STATE_NAME[st])
        ws.append(head)
        for c in ws[1]:
            c.font = Font(bold=True)
        mine = sorted((x for x in rows if x["state"] == st), key=lambda x: (x["partner_id"], x["student_id"]))
        for x in mine:
            ws.append([v if v != "" else None for v in (x["partner_id"], x["student_id"], x["grade"], x["econ"],
                                                        x["el"], x["swd"], x["prior_test"], x["prior_level"],
                                                        x["upload_date"])])
        for i, w in enumerate((11, 11, 7, 14, 10, 12, 26, 22, 12), 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for c in ws["I"][1:]:
            c.number_format = "yyyy-mm-dd"
        ws.freeze_panes = "A2"
    wb.save(path)
    return len(rows)


def pd_csv(path, rows):
    return write_csv(path, ["booking_id", "session_id", "offering_code", "offering_title", "partner_id", "account_name",
                            "session_date", "start_time_local", "delivery", "status", "hours", "facilitator",
                            "teachers_attended"],
                     [(x["booking_id"], x["session_id"], x["offering_code"], x["offering_title"], x["partner_id"],
                       x["account_name"], iso(x["session_date"]), x["start_time"], x["delivery"], x["status"],
                       x["hours"], x["facilitator"], x["teachers"]) for x in rows])


def invoices_csv(path, inv):
    return write_csv(path, ["invoice_number", "customer_account", "bill_to_name", "invoice_date", "due_date",
                            "service_period_start", "service_period_end", "description", "invoice_amount",
                            "amount_paid", "balance_due", "status", "paid_date"],
                     [(x["invoice_number"], x["customer_account"], x["bill_to"], iso(x["invoice_date"]),
                       iso(x["due_date"]), iso(x["service_start"]), iso(x["service_end"]), x["description"],
                       "%.2f" % x["amount"], "%.2f" % (x["amount"] - x["balance"]), "%.2f" % x["balance"],
                       x["status"], iso(x["paid_on"])) for x in inv])


def order_forms_xlsx(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Order forms 2025-26"
    head = ["Order Form", "Partner ID", "District", "State", "Signed", "Subscription Term Start",
            "Subscription Term End", "Licensed Seats", "Price per Seat (USD)", "Subscription Fee (USD)", "Billing",
            "Live PL Hours", "Guarantee Benchmark (%)", "Credit Percentage (%)", "Bill-to Account"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
    for dd in sorted(W.DISTRICTS, key=lambda x: x["order_form"]):
        ws.append([dd["order_form"], dd["partner_id"], dd["name"], dd["state"], dd["signed"], dd["term_start"],
                   dd["term_end"], dd["seats"], dd["seat_price"], dd["fee"], dd["billing"], dd["pd_hours"],
                   dd["benchmark"], dd["credit_pct"], dd["account"]])
    widths = (14, 10, 44, 6, 11, 12, 12, 9, 10, 12, 24, 9, 11, 10, 12)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for k in (4, 5, 6):
            row[k].number_format = "yyyy-mm-dd"
        row[8].number_format = "#,##0.00"
        row[9].number_format = "#,##0.00"
        row[12].number_format = "0.0"
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = "A1:O%d" % (len(W.DISTRICTS) + 1)
    n2 = wb.create_sheet("Notes")
    for line in ("Order forms signed for the 2025-26 Subscription Term under Master Subscription Agreement v4.0.",
                 "Subscription Fee is the annual fee before any credit (MSA section 1.2).",
                 "Live PL Hours, Guarantee Benchmark and Credit Percentage are as signed (MSA sections 8 and 9).",
                 "Bill-to Account is the customer account the District's subscription is invoiced to.",
                 "Source: Salesforce contract records, exported 2 November 2026 by Partnerships Operations."):
        n2.append([line])
    n2.column_dimensions["A"].width = 110
    wb.save(path)
    return len(W.DISTRICTS)


def survey_csv(path):
    """Spring teacher implementation survey (a distractor: implementation self-reports by teacher)."""
    r = np.random.default_rng(W.SEED + 41)
    rows = []
    rid = 30000
    for dd in W.DISTRICTS:
        k = int(r.integers(4, 19))
        for _ in range(k):
            rid += int(r.integers(1, 9))
            rows.append(("R%d" % rid, dd["partner_id"], int(r.choice([6, 7, 8])),
                         (dt.date(2026, 4, 20) + dt.timedelta(days=int(r.integers(0, 30)))).isoformat(),
                         int(r.integers(1, 6)), int(r.integers(1, 6)), int(r.integers(1, 6)),
                         int(r.choice([0, 1, 2, 3, 4, 5, 6])), r.choice(["Weekly", "2-3 times a week", "Daily",
                                                                          "Less than weekly"])))
    return write_csv(path, ["response_id", "partner_id", "grade_taught", "submitted_on", "q1_lessons_as_designed",
                            "q2_students_engaged", "q3_lab_simulations_used", "q4_pl_sessions_attended",
                            "q5_frequency_of_use"], rows)


def renewal_xlsx(path):
    """Partnerships' renewal forecast for 2026-27 (a distractor)."""
    r = np.random.default_rng(W.SEED + 43)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Renewals 2026-27"
    ws.append(["Partner ID", "District", "State", "2025-26 ARR (USD)", "Renewal Stage", "Renewal Probability",
               "Forecast 2026-27 ARR (USD)", "Account Health", "CSM"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for dd in W.DISTRICTS:
        p = float(np.round(r.uniform(0.55, 0.98), 2))
        stage = "Closed won" if dd["j"] not in (11, 33, 50) else "Closed lost"
        if stage == "Closed lost":
            p = 0.0
        ws.append([dd["partner_id"], dd["name"], dd["state"], dd["fee"], stage, p if stage != "Closed won" else 1.0,
                   round(dd["fee"] * (1.03 if stage == "Closed won" else 0.0), 2),
                   r.choice(["Green", "Green", "Yellow", "Red"]), "CSM-%02d" % int(r.integers(1, 9))])
    for i, w in enumerate((10, 44, 6, 14, 12, 12, 16, 12, 8), 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        row[3].number_format = "#,##0.00"
        row[5].number_format = "0%"
        row[6].number_format = "#,##0.00"
    ws.freeze_panes = "C2"
    wb.save(path)
    return len(W.DISTRICTS)
