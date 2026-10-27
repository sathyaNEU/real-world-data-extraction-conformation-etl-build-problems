"""Cut the pack: python3 task310/generator/build.py [target dir]. Writes every input file under target/.
Byte-identical on every run."""
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import params as P  # noqa: E402
import postprocess  # noqa: E402
import asks as A  # noqa: E402
import records as R  # noqa: E402
import world as W  # noqa: E402

TASK = os.path.dirname(HERE)
DISTRACTORS = ["tariffs", "survey"]
CONTEXT = ["velkora"]
DATA = ["register", "systems", "velkora", "ondrava", "meter", "connections", "survey"]
ASKS = ["der", "standing", "booster", "telemetry", "pilot_sites", "accreditation", "certificates", "amendments"]
DOCS = ["guideline", "standard", "guide", "tariffs"]
PACK = DATA + ASKS + DOCS


def world():
    w = W.build()
    R.attach(w)
    R.apply_events(w)
    return w


def systems_book(path, df):
    wb = Workbook()
    ws = wb.active
    ws.title = "Systems"
    for j, col in enumerate(df.columns):
        ws.cell(1, j + 1, col).font = Font(bold=True)
    for i, row in enumerate(df.itertuples(index=False)):
        for j, v in enumerate(row):
            ws.cell(i + 2, j + 1, v.item() if hasattr(v, "item") else v)
    widths = {"A": 11, "B": 11, "C": 11, "D": 10, "E": 17, "F": 9, "G": 14, "H": 20, "I": 14, "J": 16, "K": 11, "L": 36}
    for c, wdt in widths.items():
        ws.column_dimensions[c].width = wdt
    wb.save(path)
    postprocess.scrub(path, "SEO-PANEL", "2026-02-16")
    postprocess.finish_office(path, "2024-09-02T22:41:00Z", "2026-02-16T00:33:00Z", P.ORG,
                              (P.PEOPLE["panel"], P.PEOPLE["panel"]))


def sheet_book(path, df, title, widths, stamp, created, modified, people):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    for j, col in enumerate(df.columns):
        ws.cell(1, j + 1, col).font = Font(bold=True)
    for i, row in enumerate(df.itertuples(index=False)):
        for j, v in enumerate(row):
            v = v.item() if hasattr(v, "item") else v
            if v is None or (isinstance(v, float) and v != v):
                v = None
            ws.cell(i + 2, j + 1, v)
    for c, wdt in widths.items():
        ws.column_dimensions[c].width = wdt
    wb.save(path)
    postprocess.scrub(path, "SEO-PROG", stamp)
    postprocess.finish_office(path, created, modified, P.ORG, people)


def build(target):
    os.makedirs(target, exist_ok=True)
    p = lambda k: os.path.join(target, P.FILES[k])
    w = world()
    R.panel_register(w).to_csv(p("register"), index=False)
    systems_book(p("systems"), R.system_register(w))
    R.velkora(w).to_parquet(p("velkora"), index=False)
    R.ondrava(w).to_csv(p("ondrava"), index=False)
    R.meter(w).to_parquet(p("meter"), index=False)
    R.connection_register(w).to_csv(p("connections"), index=False)
    R.household_survey(w).to_csv(p("survey"), index=False)
    aw = A.ask_worlds(w)
    prem = aw["prem"]
    A.der_register(prem).to_csv(p("der"), index=False)
    A.standing_data(prem).to_csv(p("standing"), index=False)
    bc = A.booster_claims(prem)
    sheet_book(p("booster"), bc, "Claims", {"A": 11, "B": 11, "C": 15, "D": 34, "E": 18, "F": 9, "G": 13, "H": 11,
                                             "I": 10, "J": 11}, "2024-02-12", "2021-02-01T23:10:00Z",
               "2024-02-12T04:22:00Z", ("Booster scheme team", "Booster scheme team"))
    A.telemetry(aw["sites"]).to_parquet(p("telemetry"), index=False)
    sheet_book(p("pilot_sites"), A.pilot_register(aw["sites"]), "Pilot sites", {"A": 11, "B": 16, "C": 11, "D": 15,
                                                                                "E": 11, "F": 22}, "2025-08-04",
               "2024-03-01T01:15:00Z", "2025-08-04T06:02:00Z", (P.PEOPLE["director"], "Pilot coordinator"))
    A.accreditation_register(aw["firms"]).to_csv(p("accreditation"), index=False)
    aw["certs"].to_csv(p("certificates"), index=False)
    aw["am"].to_csv(p("amendments"), index=False)
    w["asks"] = aw
    import docs
    docs.build(target, P.FILES)
    return w


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(TASK, "target")
    build(out)
