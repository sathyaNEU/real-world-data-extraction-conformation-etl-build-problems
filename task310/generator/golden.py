"""task310 goldens, written from figures.json: the round schedule workbook battery_rounds_2027_2030.xlsx, the chart
round_yield_by_region.png and the brief minister_brief_battery_rounds.pdf."""
import os

import matplotlib
import matplotlib.ticker
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from reportlab.lib.units import mm  # noqa: E402
from reportlab.platypus import Image, Paragraph, Spacer  # noqa: E402

import docs as D  # noqa: E402
import params as P  # noqa: E402
import postprocess  # noqa: E402

XLSX = "battery_rounds_2027_2030.xlsx"
PNG = "round_yield_by_region.png"
PDF = "minister_brief_battery_rounds.pdf"
HEAD = PatternFill("solid", fgColor="DDE8E4")
NAMES = [P.REGIONS[r] for r in P.REGION_ORDER]


def _hdr(ws, row, cols):
    for j, h in enumerate(cols):
        c = ws.cell(row, j + 1, h)
        c.font = Font(bold=True)
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="top")


def _in_order(f):
    return sorted(f["schedule"], key=lambda n: f["schedule"][n])


def workbook(path, f):
    wb = Workbook()
    ws = wb.active
    ws.title = "Rounds"
    ws["A1"] = "Home Battery Programme: regional rounds 2027 to 2030"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = "Programme valuation under guideline v1.2 s3, 2025 panel readings. Prepared for the minister's brief."
    _hdr(ws, 4, ["Round", "Region", "Energy a programme battery delivers to the home (kWh a year)",
                 "Round of 2,400 batteries (MWh a year)"])
    for i, n in enumerate(_in_order(f)):
        r = 5 + i
        ws.cell(r, 1, f["schedule"][n])
        ws.cell(r, 2, n)
        ws.cell(r, 3, f["kwh_per_battery"][n]).number_format = "#,##0"
        ws.cell(r, 4, f["round_mwh"][n]).number_format = "#,##0"
    ws["A10"] = ("Each region's round is 2,400 times the average energy the programme battery would have delivered "
                 "to its 40 panel homes' general-supply use over 2025, simulated hour by hour from an empty battery. "
                 "Homes on export-limited agreements are valued on the solar their arrays would have produced, rebuilt "
                 "from their region's unlimited homes, not on the output their inverters held to the limit.")
    ws["A10"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A10:D13")
    for col, wdt in zip("ABCD", (9, 22, 30, 24)):
        ws.column_dimensions[col].width = wdt
    ws.row_dimensions[4].height = 46
    ws.freeze_panes = "A5"

    ws = wb.create_sheet("Eligible homes")
    _hdr(ws, 1, ["Region", "Homes eligible for the 2027 round on 30 September 2026"])
    for i, n in enumerate(NAMES):
        ws.cell(2 + i, 1, n)
        ws.cell(2 + i, 2, f["ask_a"][n]).number_format = "#,##0"
    ws.cell(7, 1, "Guideline s2: residential, solar connected throughout the 12 months to the census date, no battery "
                  "recorded, no Battery Booster rebate paid at the premises.")
    for col, wdt in zip("AB", (22, 36)):
        ws.column_dimensions[col].width = wdt
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("Pilot 2025")
    _hdr(ws, 1, ["Quarter", "Energy the pilot batteries delivered to their homes (MWh)"])
    for i, (q, v) in enumerate(f["ask_b"].items()):
        ws.cell(2 + i, 1, q)
        ws.cell(2 + i, 2, v).number_format = "#,##0.0"
    ws.cell(7, 1, "Norvane telemetry: each battery's lifetime counter read at the quarter ends; where a home's two units "
                  "report one battery meter, that meter is counted once.")
    for col, wdt in zip("AB", (12, 40)):
        ws.column_dimensions[col].width = wdt
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("Installers 2025")
    _hdr(ws, 1, ["Region", "Batteries installed in 2025 by installers accredited to 31 December 2027"])
    for i, n in enumerate(NAMES):
        ws.cell(2 + i, 1, n)
        ws.cell(2 + i, 2, f["ask_c"][n]).number_format = "#,##0"
    ws.cell(7, 1, "Installation certificates after withdrawals and installer corrections; an installation lodged twice "
                  "(the same module serials) counted once; installers Current with expiry on or after 31 December "
                  "2027; region from the installation's suburb.")
    for col, wdt in zip("AB", (22, 40)):
        ws.column_dimensions[col].width = wdt
    ws.freeze_panes = "A2"
    wb.save(path)
    postprocess.scrub(path, "SEO-PROG", "2026-11-12")
    postprocess.finish_office(path, "2026-11-09T22:41:00Z", "2026-11-12T05:17:00Z", P.ORG,
                              (P.PEOPLE["director"], P.PEOPLE["director"]))


def chart(path, f):
    TEAL, PALE, INK = "#1f6f63", "#a9c9c2", "#222222"
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 9})
    fig, ax = plt.subplots(figsize=(7.0, 3.7), dpi=150)
    order = _in_order(f)
    vals = [f["round_mwh"][n] for n in order]
    cols = [TEAL if n == f["call"] else PALE for n in order]
    xs = range(len(order))
    ax.bar(list(xs), vals, width=0.6, color=cols)
    for x, v in zip(xs, vals):
        ax.text(x, v + 70, "{:,}".format(v), ha="center", va="bottom", fontsize=8.5, color=INK)
    ax.set_xticks(list(xs))
    ax.set_xticklabels(["%d\n%s" % (f["schedule"][n], n) for n in order])
    ax.set_ylabel("MWh a year")
    ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: "{:,.0f}".format(v)))
    ax.set_ylim(0, max(vals) * 1.15)
    ax.set_title("A %s round of 2,400 batteries would deliver the most, about %s MWh a year"
                 % (f["call"], "{:,}".format(f["round_mwh"][f["call"]])), fontsize=10, loc="left")
    fig.text(0.01, 0.01, "Programme battery 10 kWh, 5 kW, simulated hour by hour on the 2025 panel readings; export-"
             "limited homes valued on their rebuilt solar.", fontsize=7, color="#555555")
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout(rect=(0, 0.03, 1, 1))
    fig.savefig(path, metadata={"Software": None})
    plt.close(fig)


def brief(path, f, png):
    body, h1, h2, small = D._styles()
    order = _in_order(f)
    call, ru = f["call"], f["runner_up"]
    kw, mw = f["kwh_per_battery"], f["round_mwh"]
    title = "Home Battery Programme: %s takes the 2027 round" % call
    s = [Paragraph(title, h1),
         Paragraph("Brief for the Minister for Energy. %s, %s. 12 November 2026." % (P.PEOPLE["director"],
                                                                                   P.TITLES["director"]), small),
         Paragraph("<b>Recommendation.</b> Run the 2027 round in %s, then %s in 2028, %s in 2029 and %s in 2030. The "
                   "2027 round would deliver about %s MWh a year to its homes, against about %s MWh for a round in %s, "
                   "the next best." % (call, order[1], order[2], order[3], "{:,}".format(mw[call]),
                                       "{:,}".format(mw[ru]), ru), body),
         D._table([["Round", "Region", "kWh a battery a year", "MWh a round a year"]] +
                  [[str(f["schedule"][n]), n, "{:,}".format(kw[n]), "{:,}".format(mw[n])] for n in order],
                  [18, 48, 40, 40], right=(2, 3)),
         Spacer(1, 6),
         Paragraph("<b>Why %s first.</b> Most of its panel homes are in Riverbend Estate and Quarry Hill Estate, where "
                   "Southmark limits export to 0.5 kW. Their inverters hold output to the home's use plus that limit, so "
                   "the 2025 readings show only part of what their arrays produce. A battery takes the energy the limit "
                   "holds back, and these families use it in the evening. Valued on the solar their arrays would have "
                   "produced, a battery there delivers about %s kWh a year, the most of any region; on the readings as "
                   "recorded it would look like the least." % (call, "{:,}".format(kw[call])), body),
         Image(png, width=160 * mm, height=160 * mm * 3.7 / 7.0),
         Spacer(1, 6),
         Paragraph("<b>Readiness.</b> On 30 September 2026, %s homes in %s were eligible for its round under the "
                   "guideline, and installers accredited through 2027 fitted %d batteries there in 2025. The 2024 pilot "
                   "batteries delivered %.1f MWh to their homes in the first quarter of 2025 and %.1f MWh in the "
                   "fourth." % ("{:,}".format(f["ask_a"][call]), call, f["ask_c"][call], f["ask_b"]["Q1 2025"],
                                f["ask_b"]["Q4 2025"]), body),
         Paragraph("Detail by region, the eligible homes, the pilot by quarter and the 2025 installations are in "
                   "battery_rounds_2027_2030.xlsx.", small)]
    D._pdf(path, title, "Ministerial brief", s,
           "%s, Home Battery Programme" % P.ORG, "2026-11-12", P.ORG, "SEO-DOCS")


def write(outdir, f):
    os.makedirs(outdir, exist_ok=True)
    workbook(os.path.join(outdir, XLSX), f)
    chart(os.path.join(outdir, PNG), f)
    brief(os.path.join(outdir, PDF), f, os.path.join(outdir, PNG))
