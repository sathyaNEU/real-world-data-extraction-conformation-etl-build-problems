"""task310 documents: the programme guideline for the 2027 to 2030 regional rounds, Southmark Networks' connection
standard for micro-embedded generation, and the monitoring panel's 2025 data guide (pdf); the retail tariff schedule
(xlsx, a distractor)."""
import os

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import params as P
import postprocess

# first and last save times on each paper's date, and the local offset that day (AEDT +11, AEST +10)
SAVED = {"2026-08-18": ("093114", "161842", "+10'00'"), "2024-11-26": ("101503", "152227", "+11'00'"),
         "2026-03-09": ("083940", "142611", "+11'00'"), "2025-06-30": ("110207", "134455", "+10'00'"),
         "2026-02-16": ("091250", "113318", "+11'00'"), "2026-11-12": ("084512", "171903", "+11'00'")}


def _styles():
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontName="Helvetica", fontSize=9.4, leading=13.0, spaceAfter=5)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName="Helvetica-Bold", fontSize=13.2, leading=17, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=10.2, leading=14,
                        spaceBefore=8, spaceAfter=4, keepWithNext=1)
    small = ParagraphStyle("small", parent=body, fontSize=8.1, leading=10.6, textColor=colors.HexColor("#444444"))
    return body, h1, h2, small


def _footer(text):
    def foot(canvas, d):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.3)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(18 * mm, 10 * mm, text)
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, "Page %d" % d.page)
        canvas.restoreState()
    return foot


def _table(rows, widths, size=8.4, right=()):
    cell = ParagraphStyle("cell", fontName="Helvetica", fontSize=size, leading=size + 2.2)
    cellr = ParagraphStyle("cellr", parent=cell, alignment=2)
    head = ParagraphStyle("head", parent=cell, fontName="Helvetica-Bold")
    out = []
    for r, row in enumerate(rows):
        out.append([Paragraph(str(v), head if r == 0 else (cellr if j in right else cell)) for j, v in enumerate(row)])
    t = Table(out, colWidths=[w * mm for w in widths], repeatRows=1, hAlign="LEFT")
    t.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.HexColor("#2F4F4F")),
                           ("LINEBELOW", (0, 1), (-1, -1), 0.25, colors.HexColor("#B9C4C4")),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E6EEEE")),
                           ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    return t


def _pdf(path, title, subject, story, foot, stamp, author, producer):
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm,
                            bottomMargin=17 * mm, title=title, author=author, creator=author, subject=subject,
                            invariant=1)
    doc.build(story, onFirstPage=_footer(foot), onLaterPages=_footer(foot))
    postprocess.scrub(path, producer, stamp)
    postprocess.pdf_clock(path, *SAVED[stamp])


def guideline(path):
    body, h1, h2, small = _styles()
    names = sorted(P.REGIONS.values())
    reg = ", ".join(names[:-1]) + " and " + names[-1]
    s = [Paragraph("Home Battery Programme: regional rounds 2027 to 2030", h1),
         Paragraph("Programme guideline, version 1.2. %s. Approved by %s, %s, 18 August 2026."
                   % (P.ORG, P.PEOPLE["director"], P.TITLES["director"]), small),
         Paragraph("1. The rounds", h2),
         Paragraph("The programme funds 2,400 home batteries a year for four years. Each year's round is held in one "
                   "region of Southmark Networks' area, so that over 2027 to 2030 each of %s has one round. A round's "
                   "batteries are the programme battery set out in Appendix A, fitted to eligible homes in the round's "
                   "region and paid for by a grant towards the installed price." % reg, body),
         Paragraph("2. Eligible homes", h2),
         Paragraph("A home is eligible for a round if, on the round's census date, it is a residential premises where "
                   "rooftop solar has been connected to Southmark Networks' network throughout the twelve months "
                   "before that date, under a micro-embedded generation agreement that is in force on it; no battery "
                   "is recorded against its connection; and no Battery Booster rebate (the state's 2021 to 2023 "
                   "scheme) has been paid for a battery at the premises. A round's census date is 30 September of the "
                   "year before the round.", body),
         Paragraph("3. What a round is worth", h2),
         Paragraph("A region's round is valued at 2,400 times the average, over the homes of the %s in that region, of "
                   "the energy the programme battery would have delivered to the home's own use in 2025 had it been "
                   "fitted at the start of that year. The panel's homes were drawn at random from each region's "
                   "homes with rooftop solar in 2024 (the panel's data guide), so each panel home stands for its "
                   "region's solar homes with equal weight. The valuation runs hour by hour over the panel's 2025 readings, "
                   "with the battery empty at the start of the year." % P.PANEL, body),
         Paragraph("4. The order of the rounds", h2),
         Paragraph("The rounds run one region a year from 2027 to 2030, in descending order of what each region's round "
                   "is worth under section 3: the region whose round is worth most has the 2027 round.", body),
         Paragraph("5. Grants and installers", h2),
         Paragraph("Grants are paid to the installer on the Office's acceptance of the installation certificate. Only "
                   "installers whose battery accreditation is current for the whole of the round's year may install "
                   "programme batteries.", body),
         Paragraph("Appendix A. The programme battery", h2),
         _table([["Item", "Specification"],
                 ["Type", "AC-coupled lithium iron phosphate battery with its own battery inverter, fitted beside the "
                          "home's existing solar inverter"],
                 ["Usable energy", "%.1f kWh" % P.BAT_USABLE_KWH],
                 ["Power", "%.1f kW charging and %.1f kW discharging" % (P.BAT_POWER_KW, P.BAT_POWER_KW)],
                 ["Efficiency", "charging stores %d per cent of the energy taken in; discharging delivers %d per cent of "
                                "the energy drawn from storage" % (round(P.BAT_CHARGE_EFF * 100),
                                                                   round(P.BAT_DISCHARGE_EFF * 100))],
                 ["Operating mode", "self-consumption: charges only from the home's own solar beyond the home's "
                                    "general-supply use at the time, and discharges only to that use; no charging from "
                                    "the grid and no export from the battery"],
                 ["Circuits", "connected to the home's general supply; not connected to a controlled-load circuit"],
                 ["Connection", "fitted under a new connection application to Southmark Networks; the home's existing "
                                "micro-embedded generation agreement and its export limit continue to apply"]],
                [34, 140]),
         Spacer(1, 4),
         Paragraph("Appendix B. The 2024 pilot", h2),
         Paragraph("In 2024 the Office fitted 600 programme batteries, 300 in Sandhaven Coast and 300 in Eastbrook Metro, "
                   "under the pilot round. The battery maker's monitoring service reports each pilot battery's daily "
                   "operation to the Office; those reports, not the panel, are the record of what the pilot batteries "
                   "delivered.", body)]
    _pdf(path, "Home Battery Programme: regional rounds 2027 to 2030", "Programme guideline", s,
         "%s, Home Battery Programme guideline v1.2" % P.ORG, "2026-08-18", P.ORG, "SEO-DOCS")


def standard(path):
    body, h1, h2, small = _styles()
    s = [Paragraph("ES-114 Connection of micro-embedded generation", h1),
         Paragraph("%s technical standard, revision 7, 26 November 2024. Owner: %s, %s." %
                   (P.NETWORK, P.PEOPLE["connections"], P.TITLES["connections"]), small),
         Paragraph("1. Scope", h2),
         Paragraph("This standard applies to inverter-connected generation and storage of up to 30 kVA connected at "
                   "low voltage, including rooftop solar and batteries at residential premises. It sets the technical "
                   "requirements a connection must meet and the export limit that applies to it.", body),
         Paragraph("2. Agreement types", h2),
         _table([["Code", "Agreement", "Export limit"],
                 ["MEG-STD", "Standard micro-embedded generation", "5 kW per phase"],
                 ["MEG-L15", "Limited export, offered where a feeder study found the standard limit unsafe", "1.5 kW per site"],
                 ["MEG-CZ05", "Constrained zone (schedule 2): all new connections from 1 July 2023", "0.5 kW per phase"]],
                [24, 110, 40]),
         Spacer(1, 4),
         Paragraph("3. Export limits", h2),
         Paragraph("3.1 An export limit applies to the active power a site sends to the network, measured at the "
                   "connection point. It does not limit energy generated and used, or stored, on the site.", body),
         Paragraph("3.2 A site complies with its limit through the export control function of its inverter or of a "
                   "site controller, configured at commissioning and acting on the power measured at the connection "
                   "point.", body),
         Paragraph("3.3 The export limit recorded in Southmark's connection register is the limit in force. A change of "
                   "agreement takes effect from the date shown in the register.", body),
         Paragraph("4. Adding a battery", h2),
         Paragraph("4.1 A battery added to an existing solar connection needs a new connection application. The site's "
                   "agreement and its export limit continue to apply at the connection point.", body),
         Paragraph("4.2 From 1 July 2022 every battery connected to Southmark's network must be recorded in the "
                   "distributed energy resources register at commissioning. Batteries connected before that date are "
                   "recorded where the customer or installer has notified Southmark.", body),
         Paragraph("5. Schedule 2: constrained zones", h2),
         Paragraph("Zone substations whose feeders carry reverse power flows above their voltage limits at midday: "
                   "Riverbend ZS (feeders RB-11 and RB-12, supplying Riverbend Estate and Quarry Hill Estate). The "
                   "schedule is reviewed each year.", body)]
    _pdf(path, "ES-114 Connection of micro-embedded generation", "Technical standard", s,
         "%s ES-114 rev 7" % P.NETWORK, "2024-11-26", P.NETWORK, "SMN-ENG")


def data_guide(path):
    body, h1, h2, small = _styles()
    s = [Paragraph("%s: 2025 data release" % P.PANEL, h1),
         Paragraph("Data guide. %s, %s, %s. 9 March 2026." % (P.ORG, P.PEOPLE["panel"], P.TITLES["panel"]), small),
         Paragraph("1. The panel", h2),
         Paragraph("The panel holds 160 homes, 40 in each of Southmark Networks' four regions. In 2024 the Office drew "
                   "each region's 40 homes at random from the region's homes with rooftop solar connected to "
                   "Southmark's network, and each home consented to share its inverter and meter data for 2025. Homes are "
                   "identified by panel ID; the register gives each home's region, network meter number (NMI) and "
                   "inverter.", body),
         Paragraph("2. Files", h2),
         _table([["File", "Holds"],
                 ["hemp_panel_register_2025.csv", "one row per home: region, suburb, NMI, inverter brand and reference"],
                 ["hemp_solar_systems_2025.xlsx", "each home's array and inverter, one row per configuration with the "
                                                  "dates it applied"],
                 ["velkora_portal_hourly_2025.parquet", "Velkora inverters: energy delivered in each hour, in Wh"],
                 ["ondrava_cloud_counters_2025.csv", "Ondrava inverters: the inverter's lifetime yield counter at the "
                                                      "end of each hour, in Wh"],
                 ["southmark_meter_halfhours_2025.parquet", "Southmark's revenue-meter reads for each home's NMI"],
                 ["southmark_connection_register_panel.csv", "each NMI's micro-embedded generation agreement"]],
                [62, 112]),
         Spacer(1, 4),
         Paragraph("3. Clocks", h2),
         Paragraph("Southmark's meter reads are in market time, which is UTC+10 all year, and each read is stamped at "
                   "the end of its half hour. The inverter portals use the local clock, which moves to daylight saving "
                   "time (UTC+11) from the first Sunday in October to the first Sunday in April. Velkora stamps each "
                   "hour at its start and gives the UTC offset; Ondrava stamps each counter reading at the end of the "
                   "hour and gives no offset. Each portal exports its own 2025 local calendar year.", body),
         Paragraph("4. Meter channels and quality", h2),
         Paragraph("E1 is energy imported on the home's general supply, E2 is energy imported on a controlled-load "
                   "circuit (off-peak hot water, switched by the network) where the home has one, and B1 is energy "
                   "exported; all in kWh for the half hour. Quality A is an actual read, E an estimate made at a meter "
                   "exchange and S a value substituted after a communications fault; E and S values are Southmark's "
                   "final settlement values. A meter exchange gives the NMI a new meter serial from the exchange.", body),
         Paragraph("5. Inverters", h2),
         Paragraph("Each portal reports the inverter's AC output. When an Ondrava inverter is replaced, the new "
                   "inverter reports under its own serial and its counter starts from zero; the system register "
                   "records the change.", body),
         Paragraph("6. Use", h2),
         Paragraph("The panel is the Office's record of how solar homes use and produce energy hour by hour. It is "
                   "not weighted: each region's 40 homes are a simple random sample of that region's solar homes "
                   "in 2024.", body)]
    _pdf(path, "Home Energy Monitoring Panel 2025 data release", "Data guide", s,
         "%s, panel data guide 2025" % P.ORG, "2026-03-09", P.ORG, "SEO-DOCS")


def tariffs(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Standing offers 2025-26"
    ws.cell(1, 1, "Standing offer retail tariffs, 1 July 2025 to 30 June 2026 (cents per kWh, inc GST)").font = \
        Font(bold=True)
    head = ["Network region", "General supply peak (3pm to 9pm)", "General supply off-peak", "Controlled load",
            "Daily supply charge (cents)", "Solar feed-in (minimum)"]
    for j, v in enumerate(head):
        ws.cell(3, j + 1, v).font = Font(bold=True)
    vals = {"WEN": (41.8, 27.1, 19.6, 118.4, 4.1), "BEL": (40.2, 26.3, 19.2, 112.9, 4.1),
            "YAR": (43.5, 28.0, 20.4, 127.6, 3.6), "TAR": (44.9, 28.9, 20.9, 131.2, 3.6)}
    for i, r in enumerate(sorted(P.REGION_ORDER, key=lambda k: P.REGIONS[k])):
        ws.cell(4 + i, 1, P.REGIONS[r])
        for j, v in enumerate(vals[r]):
            ws.cell(4 + i, 2 + j, v)
    ws.column_dimensions["A"].width = 22
    for c in "BCDEF":
        ws.column_dimensions[c].width = 20
    wb.save(path)
    postprocess.scrub(path, "SEO-DOCS", "2025-06-30")
    postprocess.finish_office(path, "2025-06-27T09:12:00Z", "2025-06-30T03:44:00Z", P.ORG, ("Office pricing", "Office pricing"))


def build(target, files):
    guideline(os.path.join(target, files["guideline"]))
    standard(os.path.join(target, files["standard"]))
    data_guide(os.path.join(target, files["guide"]))
    tariffs(os.path.join(target, files["tariffs"]))
