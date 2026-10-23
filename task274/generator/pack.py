"""Read the cut pack back into the records the analysis runs on: each examined adult's final primary result with its
county, weight, analyzer, reagent lot and the calibration in force when it was measured; the bridging run's results;
the October check; the supplement products and labels; the sample persons and the screener rosters."""
from __future__ import annotations

import csv
import datetime as dt
import os
import re
import subprocess

import openpyxl

import common as K

DST0, DST1 = dt.datetime(2025, 3, 9, 9, 0), dt.datetime(2025, 11, 2, 8, 0)   # UTC instants of the 2025 changes


def rows(path):
    with open(path, newline="", encoding="utf-8") as f:
        r = list(csv.reader(f))
    return r[0], r[1:]


def utc_to_local(s):
    t = dt.datetime.strptime(s, "%Y-%m-%dT%H:%M:%SZ")
    # Mountain time; the 2024 change back was 3 November 2024 08:00 UTC
    if DST0 <= t < DST1:
        return t - dt.timedelta(hours=6)
    return t - dt.timedelta(hours=7)


def value(txt):
    if txt.startswith("<"):
        return float(txt[1:]) - 0.05
    if txt.startswith(">"):
        return float(txt[1:]) + 0.05
    return float(txt)


def check_from_pdf(path):
    txt = subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True).stdout
    lab, targets = [], []
    for m in re.finditer(r"VD-\S+-\d\d\s+([\d.]+)\s+([\d.]+)\s+([+-][\d.]+)", txt):
        targets.append(float(m.group(1)))
        lab.append(float(m.group(2)))
    mean = float(re.search(r"Mean\s+([+-][\d.]+)", txt).group(1))
    lot = re.search(r"Reagent lot\s+(\S+)", txt).group(1)
    sn = re.search(r"Instrument\s+(\S+)", txt).group(1)
    day = re.search(r"Date tested\s+(\d+ \w+ \d{4})", txt).group(1)
    return dict(lab=lab, targets=targets, mean=mean, lot=lot, serial=sn,
                date=dt.datetime.strptime(day, "%d %B %Y").date())


class Read:
    def __init__(self, target):
        F = K.FILES
        p = lambda k: os.path.join(target, F[k])
        h, rs = rows(p("persons"))
        i = {c: j for j, c in enumerate(h)}
        codes = [c for c, _ in K.COUNTIES]
        self.persons = {int(r[i["SEQN"]]): dict(county=codes.index(r[i["COUNTY_CODE"]]), w=float(r[i["WTMEC2025"]]))
                        for r in rs}
        h, rs = rows(p("tracking"))
        i = {c: j for j, c in enumerate(h)}
        self.barcode = {}
        for r in rs:
            self.barcode[r[i["BARCODE"]].replace("-", "")] = int(r[i["SEQN"]])
        h, rs = rows(p("instruments"))
        self.serial = {r[0]: r[1] for r in rs}
        # calibrations from the event log, local time, per analyzer serial
        h, rs = rows(p("events"))
        i = {c: j for j, c in enumerate(h)}
        cals = []
        for r in rs:
            if r[i["EVENT_TYPE"]] == "CALIBRATION" and r[i["OUTCOME"]] == "ACCEPTED":
                cals.append(dict(serial=r[i["INSTRUMENT_SN"]], time=utc_to_local(r[i["EVENT_UTC"]]),
                                 lot=r[i["REAGENT_LOT"]], calot=r[i["CALIBRATOR_LOT"]]))
        cals.sort(key=lambda c: (c["serial"], c["time"]))
        for k, c in enumerate(cals):
            c["id"] = k
        self.cals = cals
        h, rs = rows(p("results"))
        i = {c: j for j, c in enumerate(h)}
        self.results = []
        for r in rs:
            if r[i["RESULT_STATUS"]] != "F":
                continue
            seqn = self.barcode[r[i["SPECIMEN_ID"]]]
            t = dt.datetime.strptime(r[i["RESULTED_AT"]], "%m/%d/%Y %H:%M")
            self.results.append(dict(seqn=seqn, value=value(r[i["RESULT"]]), code=r[i["INSTRUMENT"]],
                                     serial=self.serial[r[i["INSTRUMENT"]]], lot=r[i["REAGENT_LOT"]], time=t))
        h, rs = rows(p("bridge"))
        i = {c: j for j, c in enumerate(h)}
        self.bridge = {}
        for r in rs:
            seqn = self.barcode[r[i["SPECIMEN_ID"]]]
            self.bridge[seqn] = dict(value=value(r[i["RESULT"]]), serial=self.serial[r[i["INSTRUMENT"]]],
                                     lot=r[i["REAGENT_LOT"]],
                                     time=dt.datetime.strptime(r[i["RESULTED_AT"]], "%m/%d/%Y %H:%M"))
        self.check = check_from_pdf(p("check"))
        self.target = target

    def cal_at(self, serial, t):
        best = None
        for c in self.cals:
            if c["serial"] == serial and c["time"] <= t:
                best = c
        return best

    def records(self):
        """The analysis records: R, pairs, chk (analysis.py's shapes)."""
        lots = sorted({r["lot"] for r in self.results})
        sers = [K.ANALYZERS[0][1], K.ANALYZERS[1][1]]
        by = {s: [c for c in self.cals if c["serial"] == s] for s in sers}
        import bisect
        times = {s: [c["time"] for c in by[s]] for s in sers}
        R = []
        for r in self.results:
            k = bisect.bisect_right(times[r["serial"]], r["time"]) - 1
            c = by[r["serial"]][k]
            pp = self.persons[r["seqn"]]
            R.append(dict(seqn=r["seqn"], county=pp["county"], w=pp["w"], value=r["value"],
                          analyzer=sers.index(r["serial"]), lot=lots.index(r["lot"]), cal=c["id"],
                          month=r["time"].month))
        pairs = {s: b["value"] for s, b in self.bridge.items()}
        chk_t = dt.datetime.combine(self.check["date"], dt.time(12, 0))
        # the check ran on the day it names; the calibration in force that day on its analyzer
        cc = [c for c in by[self.check["serial"]] if c["time"].date() <= self.check["date"]][-1]
        chk = dict(mean=self.check["mean"], lab=self.check["lab"], targets=self.check["targets"], cal=cc["id"],
                   lot=lots.index(self.check["lot"]), analyzer=sers.index(self.check["serial"]), month=chk_t.month)
        return R, pairs, chk

    # ------------------------------------------------------------ the asks
    def supplements(self):
        F = K.FILES
        h, rs = rows(os.path.join(self.target, F["supplements"]))
        wb = openpyxl.load_workbook(os.path.join(self.target, F["labels"]), read_only=True)
        prod = {}
        for r in list(wb["Products"].iter_rows(values_only=True))[1:]:
            prod[r[0]] = dict(serving=float(r[4]))
        vd = {}
        for r in list(wb["Ingredients"].iter_rows(values_only=True))[1:]:
            name = str(r[1]).lower()
            if "vitamin d" in name or "calciferol" in name:
                vd[r[0]] = float(r[2]) * (40.0 if str(r[3]).lower() == "mcg" else 1.0)
        out = []
        for r in rs:
            seqn, line, pid, name, days, qty = int(r[0]), int(r[1]), r[2], r[3], int(r[4]), float(r[5])
            if pid == "GENERIC":
                m = re.search(r"([\d,]+)\s*(IU|MCG)", name)
                iu = float(m.group(1).replace(",", "")) * (40.0 if m.group(2) == "MCG" else 1.0) if m else 0.0
            else:
                iu = vd.get(pid, 0.0) / prod[pid]["serving"]
            out.append(dict(seqn=seqn, line=line, pid=pid, iu=iu, days=days, qty=qty))
        return out

    def status(self):
        F = K.FILES
        h, rs = rows(os.path.join(self.target, F["status"]))
        i = {c: j for j, c in enumerate(h)}
        codes = [c for c, _ in K.COUNTIES]
        out = []
        for r in rs:
            out.append(dict(seqn=int(r[i["SEQN"]]), county=codes.index(r[i["COUNTY_CODE"]]), screener=r[i["SCREENER_ID"]],
                            int_status=r[i["INT_STATUS"]], exam=r[i["EXAM_STATUS"]],
                            wtint=float(r[i["WTINT2025"]]) if r[i["WTINT2025"]] else None))
        h, rs = rows(os.path.join(self.target, F["rosters"]))
        i = {c: j for j, c in enumerate(h)}
        scr = {}
        for r in rs:
            scr[r[i["SCREENER_ID"]]] = (r[i["DWELLING_ID"]], r[i["SCREEN_DATE"]])
        return out, scr
