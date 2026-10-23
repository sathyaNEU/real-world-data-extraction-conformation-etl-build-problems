"""The pack's data files: the laboratory information system's results and bridging run, the specimen tracking
events, the examined persons with their exam weights, the analyzers' event log, the instrument register, the county
codes, the supplement products reported and the label database extract, the screener rosters and the sample persons'
status."""
from __future__ import annotations

import csv
import datetime as dt

import numpy as np

import common as K
import fieldwork as FW
import specimens as SPM
import world as W

RACE_TXT = {"NHW": "NON-HISPANIC WHITE", "NHB": "NON-HISPANIC BLACK", "HSP": "HISPANIC", "OTH": "OTHER OR MULTIPLE"}
SESSION = {0: "AM", 1: "PM", 2: "EVE"}


def write_csv(path, head, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\r\n")
        w.writerow(head)
        w.writerows(rows)


def us_time(t):
    return t.strftime("%m/%d/%Y %H:%M")


def iso(t):
    return t.strftime("%Y-%m-%d %H:%M")


class Pack:
    def __init__(self, w, seed=K.SEED):
        self.w = w
        self.seed = seed
        P = w["P"]
        SPM.labels(seed, P)
        SPM.accessions(seed, w)
        self.fw = FW.fieldwork(seed, P)
        self.cat = FW.catalog(seed)
        self.dsq = FW.products(seed, self.fw["SP"], self.cat)
        self.voided = self._voided()

    # ------------------------------------------------------------ laboratory
    def _voided(self):
        """Two of the year's runs whose controls failed after release: the LIS voided them and the vials were run
        again the next working day (the results the world holds are the repeats)."""
        r = W.rng(self.seed, "track")
        r = np.random.default_rng(int(r.integers(1 << 30)) + 23)
        runs = {}
        for p in self.w["P"]:
            if p.get("result") is None:
                continue
            key = (p["run_time"].date(), p["analyzer"], p["run_time"].hour // 3)
            runs.setdefault(key, []).append(p)
        keys = sorted(k for k, v in runs.items() if len(v) >= 25 and k[0].weekday() in (1, 2, 3, 4))
        pick = [keys[int(i)] for i in r.choice(len(keys), size=2, replace=False)]
        out = []
        for key in sorted(pick):
            vs = sorted(runs[key], key=lambda p: p["run_time"])
            d = key[0] - dt.timedelta(days=1)
            while d.weekday() >= 5 or d in K.HOLIDAYS:
                d -= dt.timedelta(days=1)
            t0 = dt.datetime.combine(d, dt.time(int(r.integers(10, 14)), int(r.integers(0, 60))))
            hi = float(r.uniform(0.07, 0.12))
            for j, p in enumerate(vs):
                v = round(p["result"] * float(np.exp(hi + r.normal(0, 0.02))), 1)
                out.append(dict(p=p, time=t0 + dt.timedelta(seconds=84 * j + int(r.integers(0, 40))), value=v))
        return out

    def results_rows(self):
        head = ["ACCESSION_NO", "SPECIMEN_ID", "ORDER_CODE", "TEST_NAME", "RESULT", "UNITS", "REF_RANGE",
                "RESULT_STATUS", "INSTRUMENT", "REAGENT_LOT", "RESULTED_AT", "VERIFIED_BY", "COMMENT"]
        r = W.rng(self.seed, "track")
        r = np.random.default_rng(int(r.integers(1 << 30)) + 29)
        rep = {v["p"]["seqn"] for v in self.voided}
        lotno = [lot[0] for lot in K.REAGENT_LOTS]
        rows = []
        for x in SPM.lis_results(self.seed, self.w):
            rows.append((x["time"], [x["accession"], x["barcode"], "VITD25", "25-OH VITAMIN D, TOTAL", x["result"],
                                     "ng/mL", "30.0-100.0", "F", x["analyzer"], x["lot"], us_time(x["time"]),
                                     SPM.STAFF_LAB[int(r.integers(len(SPM.STAFF_LAB)))],
                                     "REPEAT ANALYSIS" if x["seqn"] in rep else ""]))
        for v in self.voided:
            p = v["p"]
            a = K.ANALYZERS[p["analyzer"]][0]
            cal = W.cal_at(self.w["cals"], p["analyzer"], v["time"])
            rows.append((v["time"], [p["accession"], p["label"] + "01", "VITD25", "25-OH VITAMIN D, TOTAL",
                                     "%.1f" % v["value"], "ng/mL", "30.0-100.0", "X", a, lotno[cal["lot"]],
                                     us_time(v["time"]), SPM.STAFF_LAB[int(r.integers(len(SPM.STAFF_LAB)))],
                                     "RUN VOIDED - QC FAILURE AFTER RELEASE"]))
        rows.sort(key=lambda t: (t[0], t[1][0]))
        return head, [x[1] for x in rows]

    def bridge_rows(self):
        head = ["ACCESSION_NO", "SPECIMEN_ID", "ORDER_CODE", "TEST_NAME", "RESULT", "UNITS", "RESULT_STATUS",
                "INSTRUMENT", "REAGENT_LOT", "RESULTED_AT", "VERIFIED_BY", "COMMENT"]
        rows = []
        for x in SPM.lis_bridge(self.seed, self.w):
            rows.append([x["accession"], x["barcode"], "VITD25", "25-OH VITAMIN D, TOTAL", x["result"], "ng/mL", "F",
                         x["analyzer"], x["lot"], us_time(x["time"]), "KLY", "MHNES BRIDGING SAMPLE"])
        return head, rows

    def tracking_rows(self):
        head = ["EVENT_ID", "BARCODE", "SEQN", "VIAL", "EVENT", "EVENT_TIME", "LOCATION", "REFERENCE", "STAFF"]
        rows = []
        for x in SPM.tracking(self.seed, self.w):
            rows.append([x[0], x[1], x[2], x[3], x[4], iso(x[5]), x[6], x[7], x[8]])
        return head, rows

    def persons_rows(self):
        head = ["SEQN", "COUNTY_CODE", "EXAM_DATE", "EXAM_SESSION", "AGE_YR", "SEX", "RACE_ETH", "BLOOD_DRAWN",
                "WTMEC2025"]
        rows = []
        for p in sorted(self.w["P"], key=lambda p: p["seqn"]):
            rows.append([p["seqn"], p["ccode"], p["exam"].isoformat(), SESSION[p["session"]], p["age"], p["sex"],
                         RACE_TXT[p["race"]], "Y" if p.get("blood") else "N", "%.2f" % p["wtmec"]])
        return head, rows

    def event_rows(self):
        head = ["EVENT_UTC", "INSTRUMENT_SN", "EVENT_TYPE", "ASSAY", "REAGENT_LOT", "CALIBRATOR_LOT", "QC_LEVEL",
                "OUTCOME", "DETAIL"]
        rows = []
        extra = []
        for v in self.voided:
            p = v["p"]
            extra.append((v["time"] + dt.timedelta(minutes=95), K.ANALYZERS[p["analyzer"]][1]))
        seen = set()
        ev = SPM.event_log(self.seed, self.w)
        for t, sn in sorted(set((e[0].replace(minute=0, second=0), e[1]) for e in [(W.to_utc(t), sn) for t, sn in extra])):
            if (t, sn) in seen:
                continue
            seen.add((t, sn))
        lotno = [lot[0] for lot in K.REAGENT_LOTS]
        for t, sn in extra:
            an = [a for a, (_, s) in enumerate(K.ANALYZERS) if s == sn][0]
            cal = W.cal_at(self.w["cals"], an, t)
            ev.append([W.to_utc(t), sn, "QC", "25OHD", lotno[cal["lot"]], "", "1", "FAIL", "2-2S"])
            ev.append([W.to_utc(t + dt.timedelta(minutes=1)), sn, "QC", "25OHD", lotno[cal["lot"]], "", "2", "FAIL",
                       "2-2S"])
        ev.sort(key=lambda e: (e[0], e[1], e[2]))
        for e in ev:
            rows.append([e[0].strftime("%Y-%m-%dT%H:%M:%SZ")] + e[1:])
        return head, rows

    def instrument_rows(self):
        head = ["LIS_CODE", "SERIAL_NO", "MODEL", "INSTALLED", "DEPARTMENT", "NOTES"]
        return head, [["IA1", K.ANALYZERS[0][1], "Varnholt iQ 820", "2021-03-15", "CORE LAB - IMMUNOASSAY", ""],
                      ["IA2", K.ANALYZERS[1][1], "Varnholt iQ 820", "2022-08-02", "CORE LAB - IMMUNOASSAY", ""],
                      ["CH1", "VHC-55-20711", "Varnholt cx 550", "2020-11-09", "CORE LAB - CHEMISTRY", ""],
                      ["CH2", "VHC-55-21896", "Varnholt cx 550", "2021-06-21", "CORE LAB - CHEMISTRY",
                       "REPLACED CH0 2021"]]

    def county_rows(self):
        head = ["COUNTY_CODE", "COUNTY_NAME", "MEC_STAND_START", "MEC_STAND_END", "ADULTS_20PLUS_JUL2025"]
        rows = []
        for c, (code, name) in enumerate(K.COUNTIES):
            rows.append([code, name, K.STANDS[c][0].isoformat(), K.STANDS[c][1].isoformat(), K.ADULTS[c]])
        return head, rows

    # ------------------------------------------------------------ the asks' files
    def dsq_rows(self):
        head = ["SEQN", "DSQ_LINE", "PRODUCT_ID", "PRODUCT_NAME_REPORTED", "DAYS_TAKEN_30D", "UNITS_PER_DAY"]
        rows = []
        for x in self.dsq:
            p = x["p"]
            if x["generic"]:
                iu = FW._iu_per_unit(p)
                if p["vunit"] == "mcg":
                    s = "%d MCG" % int(round(iu / 40))
                else:
                    s = "%s IU" % format(int(iu), ",")
                pid, name = "GENERIC", "VITAMIN D3 %s %s" % (s, p["form"].upper())
            else:
                pid, name = p["id"], "%s %s" % (p["brand"], p["name"])
            rows.append([x["seqn"], x["line"], pid, name, x["days"], x["qty"]])
        rows.sort(key=lambda r: (r[0], r[1]))
        return head, rows

    def roster_rows(self):
        head = ["SCREENER_ID", "DWELLING_ID", "COUNTY_CODE", "SCREEN_DATE", "INTERVIEWER", "LINE", "AGE", "SEX",
                "RELATIONSHIP", "SELECTED", "SEQN"]
        sc = {s["id"]: s for s in self.fw["screenings"]}
        rows = []
        for x in self.fw["rosters"]:
            s = sc[x["screener"]]
            rows.append([x["screener"], x["dwelling"], K.COUNTIES[x["county"]][0], s["date"].isoformat(),
                         s["interviewer"], x["line"], x["age"], x["sex"], x["rel"], "Y" if x["selected"] else "N",
                         x["seqn"] if x["selected"] else ""])
        rows.sort(key=lambda r: (r[2], r[3], r[0], r[5]))
        return head, rows

    def status_rows(self):
        head = ["SEQN", "COUNTY_CODE", "SCREENER_ID", "INT_STATUS", "INT_DATE", "EXAM_STATUS", "EXAM_DATE",
                "CASE_CLOSED", "WTINT2025"]
        rows = []
        for s in sorted(self.fw["SP"], key=lambda s: s["seqn"]):
            rows.append([s["seqn"], K.COUNTIES[s["county"]][0], s["screener"], s["int_status"],
                         s["int_date"].isoformat() if s["int_date"] else "", s["exam_status"],
                         s["exam_date"].isoformat() if s["exam_date"] else "", s["closed"].isoformat(),
                         "%.2f" % s["wtint"] if s["wtint"] is not None else ""])
        return head, rows


def write_labels(path, cat):
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.append(["DSLD_ID", "BRAND", "PRODUCT_NAME", "FORM", "SERVING_SIZE", "SERVING_UNIT"])
    for p in cat:
        ws.append([p["id"], p["brand"], p["name"], p["form"], p["serving"], p["serving_unit"]])
    w2 = wb.create_sheet("Ingredients")
    w2.append(["DSLD_ID", "INGREDIENT", "AMOUNT_PER_SERVING", "UNIT"])
    for p in cat:
        if p["vd"] is not None:
            w2.append([p["id"], p["vname"], p["vd"], p["vunit"]])
        for nm, amt, u in p["extra"]:
            w2.append([p["id"], nm, amt, u])
    for sh, widths in ((ws, {"A": 11, "B": 16, "C": 36, "D": 10, "E": 13, "F": 13}),
                       (w2, {"A": 11, "B": 34, "C": 20, "D": 13})):
        for col, wd in widths.items():
            sh.column_dimensions[col].width = wd
        for cell in sh[1]:
            cell.font = Font(bold=True)
        sh.freeze_panes = "A2"
    w3 = wb.create_sheet("About")
    for line in ["Label database extract for the products reported in the 2025 wave's dietary supplement questionnaire",
                 "Amounts are per serving as printed on the label; SERVING_SIZE is the number of units in one serving.",
                 "Extracted 2026-02-17 by the Center for Health Statistics."]:
        w3.append([line])
    w3.column_dimensions["A"].width = 110
    wb.save(path)
