"""The golden deliverable, computed from the cut pack: vitamin_d_status_2025.xlsx, and generator/figures.json with every
figure the submission states.

Sheets: Counties (each county's examined adults with a final result, the share below 20 ng/mL on the reference scale
and whether the programme opens there first), Supplements, Response, and Standardisation (each calibration the year's
results were measured under: its analyzer, start, lots, results, bridging pairs, shift against the bridging run and
bias on the reference scale)."""
from __future__ import annotations

import json
import os

import numpy as np

import analysis as A
import asks as AS
import common as K
import pack as PK
import postprocess as PP

HERE = os.path.dirname(os.path.abspath(__file__))


def compute(target):
    rd = PK.Read(target)
    R, pairs, chk = rd.records()
    g, info = A.golden(R, pairs, chk)
    G = A.designated(g)
    n = [0] * A.NC
    for r in R:
        n[r["county"]] += 1
    cal_rows = []
    res_n = {}
    for r in R:
        res_n[r["cal"]] = res_n.get(r["cal"], 0) + 1
    code = {K.ANALYZERS[0][1]: "IA1", K.ANALYZERS[1][1]: "IA2"}
    for c in rd.cals:
        if c["id"] not in res_n:
            continue
        sh = info["shift"][c["id"]]
        cal_rows.append(dict(analyzer=code[c["serial"]], start=c["time"].strftime("%Y-%m-%d %H:%M"), lot=c["lot"],
                             calot=c["calot"], results=res_n[c["id"]], pairs=info["n"][c["id"]],
                             shift=round(sh, 4), bias=round(100 * (np.exp(info["bridge_bias"] + sh) - 1), 1)))
    fam = A.family(R, pairs, chk)
    margin = min(A.margin(v) for v in fam.values())
    used = sorted({r["cal"] for r in R})
    on_ia1 = sum(1 for c in used if rd.cals[c]["serial"] == K.ANALYZERS[0][1])
    lines = rd.supplements()
    status, scr = rd.status()
    a = AS.ask_a(status, lines)
    b = AS.ask_b(status, scr)
    names = [nm for _, nm in K.COUNTIES]
    wsum = sum(r["w"] for r in R)
    lb = [info["anchor"] + info["shift"][r["cal"]] - info["shift"][chk["cal"]] for r in R]
    dist_share = sum(r["w"] for r, b in zip(R, lb) if r["value"] / np.exp(b) < A.CUTOFF) / wsum
    avg = AS.supplement_avg(lines)
    iw = [s for s in status if s["int_status"] in AS.INT_DONE and s["wtint"] is not None]
    tw = sum(s["wtint"] for s in iw)
    dist_a = {t: round(100.0 * sum(s["wtint"] for s in iw if avg.get(s["seqn"], 0.0) >= t - 1e-9) / tw + 1e-9, 1)
              for t in (1000.0, 2000.0)}
    dist_b = dict(interview=round(100.0 * sum(b["interviewed"]) / sum(b["eligible"]) + 1e-9, 1),
                  exam=round(100.0 * sum(b["examined"]) / sum(b["eligible"]) + 1e-9, 1))
    voided = sum(1 for _ in open(os.path.join(target, K.FILES["results"]), encoding="utf-8") if ",X," in _)
    fig = dict(
        call=dict(counties=[names[i] for i in G]),
        counties={names[i]: dict(adults_with_result=n[i], share_below_20=round(100 * float(g[i]), 1),
                                 share_whole=int(round(100 * float(g[i]))), opens_first=i in G) for i in range(A.NC)},
        family_margin_points=round(100 * margin, 1),
        calibrations_used=dict(total=len(used), IA1=on_ia1, IA2=len(used) - on_ia1),
        check=dict(mean_difference_pct=chk["mean"], calibration_start=rd.cals[chk["cal"]]["time"].strftime("%Y-%m-%d %H:%M")),
        bridging=dict(pairs=len(pairs), bias_pct=round(100 * (np.exp(info["bridge_bias"]) - 1), 1),
                      bias_whole=int(round(100 * (np.exp(info["bridge_bias"]) - 1)))),
        calibrations=cal_rows,
        asks=dict(supplements_1000={names[i]: a[1000.0][i] for i in range(A.NC)},
                  supplements_2000={names[i]: a[2000.0][i] for i in range(A.NC)},
                  interview_rr={names[i]: b["interview"][i] for i in range(A.NC)},
                  exam_rr={names[i]: b["exam"][i] for i in range(A.NC)}),
        thresholds_iu=[1000, 2000],
        district=dict(adults_with_result=len(R), share_whole=int(round(100 * dist_share)),
                      supplements_1000=dist_a[1000.0], supplements_2000=dist_a[2000.0],
                      interview_rr=dist_b["interview"], exam_rr=dist_b["exam"], voided_rows=voided,
                      dwellings_screened_twice=len(AS.superseded(status, scr)[1])),
        results=len(R))
    return fig


def write_workbook(path, fig):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook()
    head_fill = PatternFill("solid", fgColor="1F4E5F")
    head_font = Font(bold=True, color="FFFFFF")
    hi = PatternFill("solid", fgColor="E2EFDA")
    top = Border(top=Side(style="thin", color="7F7F7F"))
    D = fig["district"]

    def header(ws, cols, widths):
        ws.append(cols)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
            c = ws.cell(row=ws.max_row, column=i)
            c.fill = head_fill
            c.font = head_font
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left" if i == 1 else "right")
        ws.row_dimensions[ws.max_row].height = 30
        ws.freeze_panes = "B%d" % (ws.max_row + 1)
        ws.print_title_rows = "%d:%d" % (ws.max_row, ws.max_row)

    def fmt(ws, row, formats):
        for col, f in formats.items():
            ws.cell(row=row, column=col).number_format = f

    def total(ws, cells, formats):
        ws.append(cells)
        r = ws.max_row
        for col in range(1, len(cells) + 1):
            ws.cell(row=r, column=col).font = Font(bold=True)
            ws.cell(row=r, column=col).border = top
        fmt(ws, r, formats)

    opens = fig["call"]["counties"]
    ws = wb.active
    ws.title = "Counties"
    ws.append(["%s, %s, %s and %s open the 2027 vitamin D programme: more than 30 per cent of their adults were "
               "below 20 ng/mL in 2025" % tuple(opens)])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append(["MHNES 2025 wave. Examined adults aged 20 and over with a final serum 25(OH)D result, on the reference "
               "measurement procedure's scale (see Standardisation), weighted with WTMEC2025."])
    ws["A2"].font = Font(italic=True, color="595959")
    ws.append([])
    header(ws, ["County", "Examined adults with a result", "Below 20 ng/mL (%)", "Opens first in 2027"],
           [15, 15, 13, 13])
    for name, v in fig["counties"].items():
        ws.append([name, v["adults_with_result"], v["share_whole"], "Yes" if v["opens_first"] else "No"])
        fmt(ws, ws.max_row, {2: "#,##0", 3: "0"})
        ws.cell(row=ws.max_row, column=4).alignment = Alignment(horizontal="right")
        if v["opens_first"]:
            for col in range(1, 5):
                ws.cell(row=ws.max_row, column=col).fill = hi
    total(ws, ["District", D["adults_with_result"], D["share_whole"], None], {2: "#,##0", 3: "0"})
    ws.append([])
    ws.append(["Shares to the nearest whole per cent; a county opens first when its share is above 30 per cent "
               "(board criterion). Every county is at least %.1f points from the line." % fig["family_margin_points"]])
    ws.cell(row=ws.max_row, column=1).font = Font(italic=True, color="595959")

    w2 = wb.create_sheet("Supplements")
    w2.append(["Adults taking 1,000 IU of vitamin D a day or more from supplements, by county"])
    w2["A1"].font = Font(bold=True, size=12)
    w2.append(["Interviewed adults, WTINT2025. Average over the 30 days before the interview (analysis plan 4.2); "
               "a product line appended by label matching beside the line it matched counts once."])
    w2["A2"].font = Font(italic=True, color="595959")
    w2.append([])
    header(w2, ["County", "At least 1,000 IU a day (%)", "At least 2,000 IU a day (%)"], [15, 16, 16])
    for name in fig["asks"]["supplements_1000"]:
        w2.append([name, fig["asks"]["supplements_1000"][name], fig["asks"]["supplements_2000"][name]])
        fmt(w2, w2.max_row, {2: "0.0", 3: "0.0"})
    total(w2, ["District", D["supplements_1000"], D["supplements_2000"]], {2: "0.0", 3: "0.0"})

    w3 = wb.create_sheet("Response")
    w3.append(["Response rates of the 2025 sample persons, by county"])
    w3["A1"].font = Font(bold=True, size=12)
    w3.append(["Unweighted; sample persons aged 20 and over, those found ineligible left out (analysis plan 5.1). "
               "The %d dwellings screened twice count once, at their later screening." % D["dwellings_screened_twice"]])
    w3["A2"].font = Font(italic=True, color="595959")
    w3.append([])
    header(w3, ["County", "Interview response rate (%)", "Examination response rate (%)"], [15, 16, 16])
    for name in fig["asks"]["interview_rr"]:
        w3.append([name, fig["asks"]["interview_rr"][name], fig["asks"]["exam_rr"][name]])
        fmt(w3, w3.max_row, {2: "0.0", 3: "0.0"})
    total(w3, ["District", D["interview_rr"], D["exam_rr"]], {2: "0.0", 3: "0.0"})

    w4 = wb.create_sheet("Standardisation")
    w4.append(["The assay's bias at each calibration the 2025 results were measured under"])
    w4["A1"].font = Font(bold=True, size=12)
    w4.append(["Event 2025-B: mean difference %+.1f per cent on IA1 (lot 43012), under the calibration of %s. The "
               "bridging run reads %+.1f per cent on the reference scale; each calibration's bias is the bridging run's "
               "times its results' mean ratio to their bridging results." % (fig["check"]["mean_difference_pct"],
                                                                            fig["check"]["calibration_start"],
                                                                            fig["bridging"]["bias_pct"])])
    w4["A2"].font = Font(italic=True, color="595959")
    w4.append([])
    header(w4, ["Analyzer", "Calibrated (local time)", "Reagent lot", "Calibrator lot", "Results", "Bridging pairs",
                "Mean log ratio to the bridging run", "Bias vs reference (%)"], [9, 16, 10, 11, 8, 9, 14, 11])
    for c in fig["calibrations"]:
        w4.append([c["analyzer"], c["start"], c["lot"], c["calot"], c["results"], c["pairs"], c["shift"], c["bias"]])
        fmt(w4, w4.max_row, {5: "#,##0", 6: "#,##0", 7: "0.0000", 8: "+0.0;-0.0"})
    total(w4, ["Total", None, None, None, sum(c["results"] for c in fig["calibrations"]),
               sum(c["pairs"] for c in fig["calibrations"]), None, None], {5: "#,##0", 6: "#,##0"})

    w5 = wb.create_sheet("Notes")
    notes = [
        ("Source", "MHNES 2025 wave release of 9 January 2026; LIS results and the December bridging run from MRMC "
                   "Laboratory Medicine; analyzer event log from the iQ 820 middleware."),
        ("Result used", "Each adult's primary vial result released as final. %d result rows from two runs the "
                        "laboratory voided after release are left out; their vials' repeat results are used."
                        % D["voided_rows"]),
        ("Calibration", "A result is measured under the calibration accepted on its analyzer most recently before the "
                        "result was released (event log in UTC, read in Mountain time)."),
        ("Reference scale", "The 20 ng/mL cut-off refers to the reference measurement procedure (analysis plan 3.3)."),
        ("Rounding", "Shares of adults below 20 ng/mL to the nearest whole per cent; supplement percentages and "
                     "response rates to one decimal."),
        ("Prepared", "Center for Health Statistics, for the board of supervisors' meeting of 9 December 2026."),
    ]
    w5.append(["Item", "Note"])
    for c in (1, 2):
        w5.cell(row=1, column=c).font = Font(bold=True)
    for k, v in notes:
        w5.append([k, v])
        w5.cell(row=w5.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        w5.cell(row=w5.max_row, column=1).alignment = Alignment(vertical="top")
    w5.column_dimensions["A"].width = 16
    w5.column_dimensions["B"].width = 100
    for sh in (ws, w2, w3, w4):
        sh.page_setup.orientation = "landscape"
        sh.page_setup.fitToWidth = 1
    wb.save(path)


def write(target, golden, figures=True):
    os.makedirs(golden, exist_ok=True)
    fig = compute(target)
    path = os.path.join(golden, K.GOLDEN)
    write_workbook(path, fig)
    PP.normalize_ooxml(path, K.REQUESTER, K.REQUESTER, "2026-10-26T15:12:00Z", "2026-10-28T09:41:00Z", excel=True)
    if figures:
        with open(os.path.join(HERE, "figures.json"), "w", encoding="utf-8") as f:
            json.dump(fig, f, indent=1)
            f.write("\n")
    return fig
