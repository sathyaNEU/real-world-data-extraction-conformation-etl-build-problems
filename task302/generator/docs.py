"""The documents: the head of supply chain's planning note for the first category on automated store ordering."""
import datetime as dt

from docx import Document
from docx.shared import Pt

import params as P
import postprocess

AUTHOR = "Nikau Fletcher"
ANALYST = "Archie Eglinton"
VENDOR = "Kowhai Retail Systems"


def _doc(author, created, modified, revision):
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    cp = d.core_properties
    cp.author = author
    cp.last_modified_by = author
    cp.created = created
    cp.modified = modified
    cp.revision = revision
    cp.title = ""
    cp.comments = ""
    return d


def plan(path):
    d = _doc(AUTHOR, dt.datetime(2026, 11, 9, 8, 40), dt.datetime(2026, 11, 16, 15, 12), 6)
    d.add_heading("Automated store ordering: the first category", 1)
    d.add_paragraph("Supply chain planning note, 16 November 2026. Nikau Fletcher, Head of Supply Chain.")
    d.add_heading("The change", 2)
    d.add_paragraph(
        "From 1 February 2027 the lines of one category in all 24 stores move onto %s' automated store "
        "ordering. The system reads each store's sales and opening stock every night and raises the next morning's "
        "order for every line, and the distribution centre's trailers deliver it before the store opens. The rest of "
        "the range follows over the next eighteen months, a category at a time." % VENDOR)
    d.add_paragraph(
        "The supply chain team has shortlisted four categories for the first move: bottled water, infant formula, dry "
        "pet food and gluten-free bakery. Infant formula comes to the stores direct from Arataki Nutrition on its "
        "weekly visit today; for the change it would be stocked at the distribution centre and come on the trailer "
        "with the other three.")
    d.add_heading("What going first is worth", 2)
    d.add_paragraph(
        "%s' contract commits the system to removing half of a category's out-of-stock line-days, and this plan "
        "takes that and no more. A line-day is out of stock when the store opened with none of the line on hand, as "
        "the stock system's opening position records it." % VENDOR)
    d.add_paragraph(
        "We value each category by what it would have sold in the 24 stores over the twelve months to 30 June 2026 "
        "had half of its out-of-stock line-days not happened, less what it did sell, at each line's 2025-26 shelf "
        "price, and we take that year as the guide to 2027. The category that adds the most goes first.")
    d.add_heading("The cereal trial", 2)
    d.add_paragraph(
        "%s ran the system for breakfast cereal in our Feilding, Levin and Carterton stores from 2 February 2026. "
        "Its export of the three stores' cereal stock for the thirteen weeks either side of that date is in the "
        "folder." % VENDOR)
    d.add_heading("Deliveries and suppliers", 2)
    d.add_paragraph(
        "The system's orders come on the next morning's trailer, so a trailer that reaches a store after it has "
        "opened puts the delivery on the shelf while customers are in the store. The category that goes first also "
        "changes what the distribution centre takes in from its suppliers each week.")
    d.add_heading("Decision", 2)
    d.add_paragraph(
        "The supply chain team will choose the first category on 25 November, from a short deck with the case and a "
        "workbook with the figures behind it.")
    d.save(path)
    postprocess.scrub(path, P.ORG, "2026-11-16")
    postprocess.finish_office(path, "2026-11-09T08:40:00Z", "2026-11-16T15:12:00Z", P.ORG, person=AUTHOR, revision=6)


