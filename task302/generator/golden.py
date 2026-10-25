"""The goldens: store_ordering_first_category_2027.pptx (the case and the chart of the four candidates) and
store_ordering_2027_workbook.xlsx (the figures behind it)."""
import datetime as dt
import io
import os
import re
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

import params as P
import postprocess
import solve as S

DECK = "store_ordering_first_category_2027.pptx"
XLSX = "store_ordering_2027_workbook.xlsx"
AUTHOR = "Archie Eglinton"
SHORT = {"Bottled water": "Water", "Infant formula": "Formula", "Dry pet food": "Pet food",
         "Gluten-free bakery": "Gluten-free bakery"}


def c(n):
    return "{:,}".format(int(n))


def _text(slide, left, top, width, height, text, size, bold=False, colour=None):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tb.text_frame.word_wrap = True
    tb.text_frame.text = text
    for para in tb.text_frame.paragraphs:
        for run in para.runs:
            run.font.size = Pt(size)
            run.font.bold = bold
            if colour:
                run.font.color.rgb = RGBColor.from_string(colour)
    return tb


def _title(slide, text):
    t = slide.shapes.title
    t.text = text
    t.left, t.top, t.width, t.height = Inches(0.45), Inches(0.3), Inches(12.4), Inches(0.95)
    p = t.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.runs[0].font.size = Pt(26)
    p.runs[0].font.bold = True


def deck(F, path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[5])
    _title(s, "Infant formula goes first: about NZD %s a year more formula sales" % c(F["call_add"]))
    order = sorted(S.CATS, key=lambda k: F["add_nzd"][k])
    cd = CategoryChartData()
    cd.categories = [SHORT[k] for k in order]
    cd.add_series("NZD thousands a year", [round(F["add_nzd"][k] / 1000.0, 1) for k in order])
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.45), Inches(1.35), Inches(8.2), Inches(5.3), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = True
    ch.chart_title.text_frame.text = ("Sales each category adds a year with half its out-of-stock line-days removed "
                                      "(NZD thousands, 2025-26 prices)")
    ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
    va = ch.value_axis
    va.minimum_scale, va.maximum_scale, va.major_unit = 0, 70, 10
    va.has_major_gridlines = False
    va.tick_labels.font.size = Pt(10)
    ch.category_axis.tick_labels.font.size = Pt(11)
    pl = ch.plots[0]
    pl.gap_width = 70
    ser = pl.series[0]
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = RGBColor(0xA6, 0xAE, 0xB8)
    i = order.index(F["call_category"])
    pt = ser.points[i]
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = RGBColor(0x1F, 0x5F, 0x8B)
    dl = pt.data_label
    dl.has_text_frame = True
    dl.text_frame.text = "about %s" % c(F["call_add"])
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.text_frame.paragraphs[0].runs[0].font.size = Pt(11)
    dl.text_frame.paragraphs[0].runs[0].font.bold = True
    o = F["oos_line_days"]
    _text(s, 8.95, 1.45, 3.95, 4.9,
          "Formula's %s out-of-stock line-days all fell on the %s store-days its whole shelf was empty, the weeks "
          "Arataki's visit was missed: nothing else on the shelf took the sale and shoppers did not wait for the "
          "next visit.\n\nWater had the most out-of-stock line-days (%s), but on those days the other water lines sold "
          "most of what the missing line would have.\n\nPet food's shoppers mostly bought the missing bag after the "
          "restock." % (c(o["Infant formula"]), c(F["win_whole_shelf_store_days"]), c(o["Bottled water"])), 13)
    _text(s, 0.45, 6.8, 12.0, 0.35,
          "Point-of-sale and stock system, 24 stores, 1 July 2025 to 30 June 2026. A line-day is out of stock when the "
          "store opened with none on hand. Supply chain analytics, 20 November 2026.", 9, colour="777777")
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    _title(s2, "How the figures were worked")
    _text(s2, 0.45, 1.4, 12.3, 5.0,
          "Each line's expected sales on a day are its sales on the same weekday in the four weeks either side, on days "
          "it was on the shelf, outside the fortnight after its own restock and with no other line of its category "
          "out.\n\nA category's loss over the year is its missing lines' expected sales on their out-of-stock days, "
          "less what its other lines sold above their own expected sales on those days and what each line sold above "
          "its expected sales in the fortnight after it came back.\n\nThe value of going first is half that loss, at "
          "each line's 2025-26 shelf price, as the planning note sets out. The trial, delivery and pallet figures "
          "are in the workbook.", 15)
    prs.core_properties.author = AUTHOR
    prs.core_properties.last_modified_by = AUTHOR
    prs.core_properties.title = "Automated store ordering: the first category"
    prs.core_properties.comments = ""
    prs.core_properties.created = dt.datetime(2026, 11, 18, 9, 30)
    prs.core_properties.modified = dt.datetime(2026, 11, 20, 14, 5)
    prs.core_properties.revision = 5
    prs.save(path)
    _fix_embedded(path)
    postprocess.scrub(path, P.ORG, "2026-11-20")


def _fix_embedded(path):
    """the chart workbook python-pptx embeds carries the build clock: give it the deck's own date and fixed entry
    stamps so two builds are byte-identical"""
    with zipfile.ZipFile(path) as z:
        items = [(i.filename, z.read(i.filename)) for i in z.infolist()]
    out = []
    for n, b in items:
        if n.endswith(".xlsx"):
            with zipfile.ZipFile(io.BytesIO(b)) as zz:
                inner = [(i.filename, zz.read(i.filename)) for i in zz.infolist()]
            buf = io.BytesIO()
            with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zo:
                for m, bb in inner:
                    if m == "docProps/core.xml":
                        bb = re.sub(rb"\d{4}-\d\d-\d\dT\d\d:\d\d:\d\dZ", b"2026-11-20T14:05:00Z", bb)
                        bb = re.sub(rb"<dc:creator>[^<]*</dc:creator>", b"<dc:creator>" + AUTHOR.encode() +
                                    b"</dc:creator>", bb)
                    zi = zipfile.ZipInfo(m, date_time=(1980, 1, 1, 0, 0, 0))
                    zi.compress_type = zipfile.ZIP_DEFLATED
                    zo.writestr(zi, bb)
            b = buf.getvalue()
        if n.startswith("ppt/slideMasters/") and n.endswith(".xml"):
            b = b.replace("\u2013".encode(), "\u2022".encode())
        out.append((n, b))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in out:
            zi = zipfile.ZipInfo(n, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, b)


def workbook(F, path):
    wb = Workbook()
    hdr, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F5F8B")
    ws = wb.active
    ws.title = "Cereal trial"
    ws.append(["Breakfast cereal in the three trial stores: out-of-stock line-days in the 13 weeks after the system "
               "went live (2 February 2026)"])
    ws["A1"].font = Font(bold=True)
    ws.append(["Store", "Store name", "Out-of-stock line-days, 2 Feb to 3 May 2026",
               "Change from the 13 weeks before (%)"])
    for cell in ws[2]:
        cell.font, cell.fill = hdr, fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for s in S.TRIAL_STORES:
        v = F["ask_trial"][str(s)]
        ws.append([s, F["store_names"][str(s)], v["after"], v["change_pct"]])
        ws.cell(row=ws.max_row, column=4).number_format = "0.0"
    ws.append([])
    ws.append(["A line-day is out of stock when the store opened with none of the line on hand."])
    for col, w in zip("ABCD", (8, 22, 24, 22)):
        ws.column_dimensions[col].width = w
    ws.row_dimensions[2].height = 32
    ws2 = wb.create_sheet("Late deliveries")
    ws2.append(["Distribution centre deliveries that reached a store after it opened, 1 July 2025 to 30 June 2026"])
    ws2["A1"].font = Font(bold=True)
    ws2.append(["Island", "Deliveries after opening"])
    for cell in ws2[2]:
        cell.font, cell.fill = hdr, fill
    for g, v in F["ask_late_stops"].items():
        ws2.append([g + " Island", v])
    for col, w in zip("AB", (24, 24)):
        ws2.column_dimensions[col].width = w
    ws3 = wb.create_sheet("Pallets received")
    ws3.append(["Pallets the distribution centre received from each category's suppliers, by month of receipt"])
    ws3["A1"].font = Font(bold=True)
    ws3.append(["Month"] + list(S.DC_CATS))
    for cell in ws3[2]:
        cell.font, cell.fill = hdr, fill
    for m in S.MONTHS:
        ws3.append([dt.datetime.strptime(m, "%Y-%m").strftime("%b %Y")] + [F["ask_pallets"][k][m] for k in S.DC_CATS])
    ws3.append([])
    ws3.append(["Each supplier truck counted once; infant formula comes direct from its supplier."])
    for col, w in zip("ABCD", (12, 16, 16, 20)):
        ws3.column_dimensions[col].width = w
    wb.properties.creator = AUTHOR
    wb.properties.lastModifiedBy = AUTHOR
    wb.properties.created = dt.datetime(2026, 11, 17, 16, 20)
    wb.properties.modified = dt.datetime(2026, 11, 20, 13, 48)
    wb.save(path)
    postprocess.scrub(path, P.ORG, "2026-11-20")
    postprocess.finish_office(path, "2026-11-17T16:20:00Z", "2026-11-20T13:48:00Z", P.ORG, person=AUTHOR)


def write(golden_dir, F, X=None):
    os.makedirs(golden_dir, exist_ok=True)
    deck(F, os.path.join(golden_dir, DECK))
    workbook(F, os.path.join(golden_dir, XLSX))
