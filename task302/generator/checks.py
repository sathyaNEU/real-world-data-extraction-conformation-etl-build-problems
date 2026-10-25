"""Generator assertions: the call, its family, the ladder, the fork table, the three rules on the clean truth, the asks'
devices, separation, the input gates and the stock identity. Returns the number of failures."""
import os

import numpy as np
import pandas as pd

import build as B
import params as P
import solve as S

BAD = []


def check(name, ok, info=""):
    print(("ok   " if ok else "FAIL ") + name + ((" :: " + str(info)) if info != "" else ""))
    if not ok:
        BAD.append(name)


def rnd(x, u):
    return int(round(x / float(u)) * u)


def margin(d):
    v = sorted(d.values(), reverse=True)
    return v[0] / v[1]


def main(target, R=None):
    del BAD[:]
    D, est, add, oos = S.golden(target)
    win = max(add, key=add.get)
    check("call: the golden names infant formula", win == "Infant formula", {k: round(v) for k, v in add.items()})
    check("call: printed to the nearest NZD 10,000 is 60,000", rnd(add[win], 10000) == 60000, round(add[win]))
    check("call: at least NZD 2,000 inside its printed cell", 5000 - abs(add[win] - 60000) >= 2000,
          round(5000 - abs(add[win] - 60000)))
    check("call: formula leads the next category by 1.5 times or more", margin(add) >= 1.5, round(margin(add), 2))
    fam, rungs = S.family(D)
    for name, v in fam.items():
        w = max(v, key=v.get)
        check("family: %s names infant formula at 60,000, 1.1 times the next or more" % name,
              w == "Infant formula" and rnd(v[w], 10000) == 60000 and margin(v) >= 1.1,
              "%s %.0f, x%.2f" % (w, v[w], margin(v)))
    want = {"line-level lost sales": "Bottled water", "net of the other lines on the out days only": "Dry pet food",
            "net of the restock only": "Bottled water", "store-week regression": "Dry pet food"}
    for name, v in rungs.items():
        w = max(v, key=v.get)
        check("ladder: %s names %s, 1.1 times the next or more" % (name, want[name]),
              w == want[name] and margin(v) >= 1.1, "%s %.0f, x%.2f" % (w, v[w], margin(v)))
    # the critical component: formula's line-days out are whole-shelf days
    k = D["cat"] == "Infant formula"
    check("component: formula's out-of-stock line-days all fall on store-days its whole shelf was out",
          int(D["O"][:, k, S.BASE:].all(axis=1).sum()) * int(k.sum()) == oos["Infant formula"],
          "%d store-days, %d line-days" % (int(D["O"][:, k, S.BASE:].all(axis=1).sum()), oos["Infant formula"]))
    # fork table rows
    def named(v):
        w = max(v, key=v.get)
        return w, rnd(v[w], 10000), round(margin(v), 2)

    def adds(DD, **kw):
        e = S.by_category(DD, **kw)
        return {c: e[c]["net"] * P.CUT for c in S.CATS}
    stk = pd.read_csv(os.path.join(target, B.FILES["stock"]))
    s_i = {s_: j for j, s_ in enumerate(D["st"].store_no)}
    l_i = {a: j for j, a in enumerate(D["ln"].article)}
    dd = (pd.to_datetime(stk.date) - P.D0).dt.days.values
    on, adj = np.zeros(D["U"].shape), np.zeros(D["U"].shape)
    on[stk.store_no.map(s_i).values, stk.article.map(l_i).values, dd] = stk.on_hand_open.values
    adj[stk.store_no.map(s_i).values, stk.article.map(l_i).values, dd] = stk.adjustment.values
    shut = D["closed"]
    w, f, m = named(adds(dict(D, O=((on - D["U"] + adj) <= 0) & ~shut[:, None, :])))
    check("SF1 line-days that sold out by the close instead of opening empty: infant formula still first", w == "Infant formula"
          and m >= 1.2, "%s %d, x%.2f (the plan's opening position gives the golden)" % (w, f, m))
    import docx
    plan = " ".join(p_.text for p_ in docx.Document(os.path.join(target, B.FILES["plan"])).paragraphs)
    dic = docx.Document(os.path.join(target, B.FILES["dictionary"]))
    dic_text = " ".join(c_.text for t_ in dic.tables for r_ in t_.rows for c_ in r_.cells)
    check("SF2 the plan values half of the out-of-stock line-days, naming no particular half; the additions are half the loss",
          P.CUT == 0.5 and "removing half of a category's out-of-stock line-days" in plan and
          all(abs(add[c] - 0.5 * est[c]["net"]) < 1e-6 for c in S.CATS))
    w, f, m = named(adds(D, span=2))
    check("SF3 both years averaged instead of the twelve months to 30 June 2026: infant formula still first",
          w == "Infant formula" and m >= 1.5 and "over the twelve months to 30 June 2026" in plan, "%s %d, x%.2f" % (w, f, m))
    lnf = D["ln"]
    pos = pd.read_parquet(os.path.join(target, B.FILES["pos"]))
    fy = pd.to_datetime(pos.trade_date) >= P.BASE_LO
    code = pd.read_csv(os.path.join(target, B.FILES["codes"]), dtype={"item_code": str})
    art = dict(zip(code.item_code, code.article))
    p26 = dict(zip(lnf.article, lnf.shelf_price_fy2026))
    till = pos[fy].assign(p=lambda x: x.item_code.map(art).map(p26))
    check("SF4 the till's 2025-26 value is units at the 2025-26 shelf price: one valuation",
          np.allclose(till.net_sales_nzd, (till.units * till.p).round(2)))
    check("SF5 the shelf price includes GST (the dictionary); ex-GST is a common factor and the same call",
          "Shelf price in NZD including GST" in dic_text, "ex-GST formula %d" % rnd(add["Infant formula"] / 1.15, 10000))
    neg = int((pos.units < 0).sum())
    w, f, m = named(adds(S.load(target, returns=False)))
    check("SF6 the till rows netting to negative units dropped: infant formula at 60,000 still",
          0 < neg < 0.01 * len(pos) and w == "Infant formula" and f == 60000, "%d rows; %s %d, x%.2f" % (neg, w, f, m))
    w, f, m = named(adds(S.load(target, codes="current")))
    check("SF7 the tills' codes mapped by lines.csv alone (current GTINs): infant formula at 60,000 still",
          w == "Infant formula" and f == 60000, "%s %d, x%.2f" % (w, f, m))
    prompt = open(os.path.join(os.path.dirname(os.path.abspath(target)), "prompt.txt")).read()
    check("SF8 the addition is a year's: the prompt asks what it adds a year, the plan values twelve months",
          "add a year" in prompt, "February to December only: %d" % rnd(add["Infant formula"] * 11 / 12, 10000))
    lead = {}
    for j in range(D["U"].shape[0]):
        a_ = adds(dict(D, U=D["U"][j:j + 1], O=D["O"][j:j + 1], closed=D["closed"][j:j + 1]))
        w_ = max(a_, key=a_.get)
        lead[w_] = lead.get(w_, 0) + 1
    check("SF9 the plan's first category adds the most across the chain; store by store formula leads the most stores",
          max(lead, key=lead.get) == "Infant formula" and "The category that adds the most goes first" in plan, lead)
    closed_days = int(D["closed"][:, S.BASE:].sum())
    zero_shut = int(((on == 0) & D["closed"][:, None, :]).sum())
    w, f, m = named(adds(dict(D, closed=np.zeros_like(D["closed"]))))
    check("SF10 the store-days with no sales at all are the two refit weeks, none opening a line at zero; read as trading "
          "days, infant formula at 60,000 still", closed_days == 10 and zero_shut == 0 and w == "Infant formula" and
          f == 60000, "%d store-days, %d zero openings; %s %d" % (closed_days, zero_shut, w, f))
    # three rules on the clean truth
    if R is not None:
        tr, gross_t, sub_t = {}, {}, {}
        for c, (key, cat) in zip(("W", "F", "P", "G"), enumerate(S.CATS)):
            T = R["res"][c]
            sl = slice(S.BASE, None)
            pr = T["price"][None, :, None]
            tr[cat] = (T["lost"][:, :, sl] * pr).sum() * P.CUT
            gross_t[cat] = (T["lam"][:, :, sl] * T["oos"][:, :, sl] * pr).sum() * P.CUT
            sub_t[cat] = ((T["first"][:, :, sl] * T["oos"][:, :, sl] * pr).sum() - (T["sub"][:, :, sl] * pr).sum()) * P.CUT
        check("three rules: on the clean truth the category's own loss names infant formula",
              max(tr, key=tr.get) == "Infant formula", {k_: round(v) for k_, v in tr.items()})
        check("three rules: on the clean truth the lines' own lost sales name bottled water",
              max(gross_t, key=gross_t.get) == "Bottled water", {k_: round(v) for k_, v in gross_t.items()})
        check("three rules: on the clean truth netting the other lines alone names dry pet food",
              max(sub_t, key=sub_t.get) == "Dry pet food", {k_: round(v) for k_, v in sub_t.items()})
    # the asks
    A = S.ask_trial(target)
    t = pd.read_csv(os.path.join(target, B.FILES["trial"]))
    t["date"] = pd.to_datetime(t.date)
    hits = 0
    for st in S.TRIAL_STORES:
        x = t[t.store_no == st]
        b = int((x[x.date < S.TRIAL_LIVE].on_hand == 0).sum())
        a = int((x[x.date >= S.TRIAL_LIVE].on_hand == 0).sum())
        nd = round(100 * (1 - a / b), 1)
        hits += int(a == A[st]["after"]) + int(nd == round(A[st]["down_pct"], 1))
        check("ask A store %d: the trial shares sit away from the contract's half under both readings" % st,
              abs(A[st]["down_pct"] - 50) >= 4 and abs(nd - 50) >= 4, "%.1f and %.1f" % (A[st]["down_pct"], nd))
        check("ask A store %d: at least 0.01 inside its printed cell" % st,
              0.05 - abs(A[st]["down_pct"] - round(A[st]["down_pct"], 1)) >= 0.01, round(A[st]["down_pct"], 3))
    check("ask A: on_hand read as the opening stock throughout misses all six cells", hits == 0, "%d of 6" % hits)
    sells_empty = int(((t.on_hand == 0) & (t.units_sold > 0) & (t.date >= S.TRIAL_LIVE)).sum())
    sells_empty_b = int(((t.on_hand == 0) & (t.units_sold > 0) & (t.date < S.TRIAL_LIVE)).sum())
    check("ask A device: lines sell on days they show nothing on hand only from go-live",
          sells_empty > 20 and sells_empty_b == 0, "%d after, %d before" % (sells_empty, sells_empty_b))
    Bst = S.ask_stops(target)
    s = pd.read_csv(os.path.join(target, B.FILES["stops"]))
    arr = pd.to_datetime(s.arrived)
    m = arr.dt.hour * 60 + arr.dt.minute + arr.dt.second / 60.0
    reg = dict(zip(D["st"].store_no, D["st"].island))
    naive = s[m > 7 * 60].groupby(s.store_no.map(reg)).size()
    gold = Bst.groupby(Bst.index.map(reg)).sum()
    hits = int(sum(int(naive.get(g, 0)) == int(v) for g, v in gold.items()))
    check("ask B: the store list's 7:00 all year misses both islands", hits == 0, "%d of %d" % (hits, len(gold)))
    near = ((m >= 6 * 60 + 20) & (m < 6 * 60 + 40)) | ((m >= 6 * 60 + 50) & (m < 7 * 60 + 10))
    check("ask B: no stop within ten minutes of either opening time", int(near.sum()) == 0, int(near.sum()))
    tr_ = pd.read_csv(os.path.join(target, B.FILES["trading"]))
    fs = pd.to_datetime(tr_.trade_date + " " + tr_.first_sale)
    early = fs < pd.Timestamp("2026-04-01")
    check("ask B evidence: first sales from 6:30 to 31 March 2026 and from 7:00 from 1 April",
          (fs[early].dt.hour * 60 + fs[early].dt.minute).between(390, 399).all() and
          (fs[~early].dt.hour * 60 + fs[~early].dt.minute).between(420, 429).all())
    C = S.ask_pallets(target)
    grn = pd.read_csv(os.path.join(target, B.FILES["grn"]))
    dock = pd.read_csv(os.path.join(target, B.FILES["dock"]))
    po = pd.read_csv(os.path.join(target, B.FILES["po"]))
    art_cat = dict(zip(D["ln"].article, D["ln"].category))
    g = grn.assign(cat=grn.article.map(art_cat), m=pd.to_datetime(grn.received).dt.to_period("M").astype(str),
                   s=grn.po_no.map(dict(zip(po.po_no, po.supplier))),
                   ap=grn.delivery_ref.map(dict(zip(dock.delivery_ref, dock.appointment_no))))
    check("ask C: every dispatch advice received is on a dock appointment", int(g.ap.isna().sum()) == 0,
          int(g.ap.isna().sum()))
    for nm, keys, bar in (("each dispatch advice's pallets once", "delivery_ref", 0), ("each receipt's pallets once",
                          "grn_no", 0), ("each order's pallets once", "po_no", 0), ("every line's pallets", None, 0),
                          ("each supplier-day's pallets once", ["s", "received"], 6)):
        v = (g if keys is None else g.drop_duplicates(keys)).groupby(["cat", "m"]).pallets.sum().unstack().reindex(
            index=S.DC_CATS, columns=S.MONTHS)
        hits = int((v.values == C.values).sum())
        check("ask C: %s matches %d or fewer of the 36 cells" % (nm, bar), hits <= bar, "%d of 36" % hits)
    per = g.groupby("ap").agg(n=("delivery_ref", "nunique"), p=("pallets", "nunique"))
    check("ask C device: one pallet figure on every line of an appointment, two or more dispatch advices on many",
          bool((per.p == 1).all()) and int((per.n > 1).sum()) >= 60, "%d appointments carry two or more"
          % int((per.n > 1).sum()))
    # separation and gates
    check("separation: the call reads the sales, stock, lines, till codes and stores only", True,
          "solve.load opens %s, %s, %s, %s and %s" % (B.FILES["pos"], B.FILES["stock"], B.FILES["lines"],
                                                       B.FILES["codes"], B.FILES["stores"]))
    files = sorted(os.listdir(target))
    fmts = sorted({f.rsplit(".", 1)[1] for f in files})
    check("input gates: 10 or more files", len(files) >= 10, len(files))
    check("input gates: 3 or more formats", len(fmts) >= 3, fmts)
    check("input gates: a file of 25,000 rows or more", len(pos) >= 25000, len(pos))
    check("input gates: 2 or more distractors", len(B.DISTRACTORS) >= 2, B.DISTRACTORS)
    stk = pd.read_csv(os.path.join(target, B.FILES["stock"]))
    stk = stk.sort_values(["store_no", "article", "date"]).reset_index(drop=True)
    pos["article"] = pos.item_code.map(art)
    pos["date"] = pd.to_datetime(pos.trade_date).dt.strftime("%Y-%m-%d")
    u = pos.groupby(["store_no", "article", "date"], as_index=False).units.sum()
    stk = stk.merge(u, on=["store_no", "article", "date"], how="left").fillna({"units": 0})
    stk = stk.sort_values(["store_no", "article", "date"]).reset_index(drop=True)
    g = stk.groupby(["store_no", "article"])
    lhs = g.on_hand_open.shift(-1)
    rhs = stk.on_hand_open - stk.units + stk.adjustment + g.delivered.shift(-1)
    ok = lhs.isna() | (lhs == rhs)
    check("stock identity: every opening position follows from the day before", bool(ok.all()), int((~ok).sum()))
    print("%d failures" % len(BAD))
    return len(BAD)
