"""The pack's documents: the Brannoch Lake Phosphorus Agreement (definitions, Schedule 1 allocations, the
Tributary Restoration Fund, the annual report), the commission's tributary monitoring program, and the gage
contractor's data transmittal."""
import datetime as dt

from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                      # noqa: E402
from reportlab.lib.pagesizes import LETTER                                            # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                       # noqa: E402
from reportlab.lib.units import inch                                                  # noqa: E402
from reportlab.pdfbase import pdfmetrics                                              # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                          # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import world as W   # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/liberation/"
for nm, f in (("LiberationSerif", "Regular"), ("LiberationSerif-Bold", "Bold"), ("LiberationSans", "Regular"),
              ("LiberationSans-Bold", "Bold")):
    pdfmetrics.registerFont(TTFont(nm, FONT_DIR + nm.split("-")[0] + "-" + f + ".ttf"))
INK = colors.HexColor("#1d1f21")
GREY = colors.HexColor("#5a5f66")
BODY = ParagraphStyle("body", fontName="LiberationSerif", fontSize=10.5, leading=14.5, textColor=INK, spaceAfter=7)
NUM = ParagraphStyle("num", parent=BODY, leftIndent=28, firstLineIndent=-28)
SUBNUM = ParagraphStyle("subnum", parent=BODY, leftIndent=46, firstLineIndent=-18)
H1 = ParagraphStyle("h1", fontName="LiberationSerif-Bold", fontSize=13.5, leading=17, textColor=INK, spaceAfter=4, alignment=1)
SUB = ParagraphStyle("sub", fontName="LiberationSerif", fontSize=10, leading=13, textColor=GREY, spaceAfter=14, alignment=1)
H2 = ParagraphStyle("h2", fontName="LiberationSerif-Bold", fontSize=11, leading=15, textColor=INK, spaceBefore=8, spaceAfter=4)
CELL = ParagraphStyle("cell", fontName="LiberationSans", fontSize=8.8, leading=11, textColor=INK)
CELLB = ParagraphStyle("cellb", parent=CELL, fontName="LiberationSans-Bold")
SMALL = ParagraphStyle("small", fontName="LiberationSans", fontSize=8.5, leading=11, textColor=GREY)


def _doc(path, title, author, footer):
    def frame(canvas, doc):
        canvas.saveState()
        canvas.setFont("LiberationSans", 7.5)
        canvas.setFillColor(GREY)
        canvas.drawString(1.0 * inch, 0.6 * inch, footer)
        canvas.drawRightString(LETTER[0] - 1.0 * inch, 0.6 * inch, "Page %d" % doc.page)
        canvas.restoreState()
    d = SimpleDocTemplate(path, pagesize=LETTER, leftMargin=1.0 * inch, rightMargin=1.0 * inch, topMargin=0.9 * inch,
                          bottomMargin=0.95 * inch, title=title, author=author, subject="", creator="", producer="")
    return d, frame


def _table(rows, widths, head=True):
    data = [[Paragraph(str(c), CELLB if (head and i == 0) else CELL) for c in r] for i, r in enumerate(rows)]
    t = Table(data, colWidths=widths, repeatRows=1 if head else 0)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9aa0a6")),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eaed") if head else colors.white),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 4),
                           ("RIGHTPADDING", (0, 0), (-1, -1), 4)]))
    return t


# ---------------------------------------------------------------------------------------------- the agreement
def write_agreement(path, alloc):
    title = "Brannoch Lake Phosphorus Agreement"
    doc, frame = _doc(path, title, W.COMMISSION, "%s  |  Brannoch Lake Phosphorus Agreement  |  Consolidated text, 2025" % W.COMMISSION)
    s = [Paragraph("BRANNOCH LAKE PHOSPHORUS AGREEMENT", H1),
         Paragraph("Among the State, the ten watershed towns and the permitted dischargers of the Brannoch basin, "
                   "administered by the Brannoch Lake Commission. Consolidated text incorporating the amendments adopted "
                   "on 9 December 2025.", SUB)]
    N = lambda n, t: s.append(Paragraph("<b>%s</b>&nbsp;&nbsp;&nbsp;%s" % (n, t), NUM))
    M = lambda n, t: s.append(Paragraph("(%s)&nbsp;&nbsp;%s" % (n, t), SUBNUM))
    s.append(Paragraph("1 Definitions", H2))
    N("1.1", "<i>Water year</i> means the twelve months from 1 October to 30 September, named for the calendar year in "
             "which it ends.")
    N("1.2", "<i>Compliance station</i> means the station listed for a tributary in Schedule 1, at which the "
             "tributary's load is measured.")
    N("1.3", "<i>Total phosphorus</i> means phosphorus in all forms in an unfiltered sample of stream water, expressed as P.")
    N("1.4", "<i>Load</i> means the mass of total phosphorus carried past a tributary's compliance station during a water "
             "year, in kilograms, estimated from the tributary monitoring record: the Commission's samples, the storm "
             "samples the %s takes under its cooperative agreement with the Commission, and the gage record." % W.LAB)
    s.append(Paragraph("2 Allocations", H2))
    N("2.1", "Each tributary's allocation is its share of the lake's annual phosphorus target, set out in Schedule 1.")
    N("2.2", "A tributary's load above its allocation is its load less its allocation. Comparisons under this agreement "
             "are made on unrounded loads.")
    s.append(Paragraph("3 The Tributary Restoration Fund", H2))
    N("3.1", "Each year the Fund takes on three watershed projects for the following calendar year, one in each of the "
             "three tributaries whose load for the water year just closed exceeds its allocation by the most kilograms. "
             "Where fewer than three tributaries exceed their allocations, the Fund takes on projects only in those that do.")
    N("3.2", "The Commission's technical committee confirms the year's tributaries at its January meeting, and the "
             "projects are scoped with the watershed towns by 31 March.")
    s.append(Paragraph("4 The annual tributary report", H2))
    N("4.1", "Beside each tributary's load and allocation, the Commission's annual tributary report gives for each tributary:")
    M("a", "the phosphorus discharged during the water year by permitted facilities whose outfalls are upstream of its "
           "compliance station, from the facilities' discharge monitoring reports, each month's average daily "
           "discharge (the monthly average concentration times the monthly average flow, or the monthly average mass) "
           "times the days in the month; and")
    M("b", "the acres of cover crop lying in its watershed on fields under the %s's contracts for the crop year in "
           "which the water year ends." % W.PROGRAM)
    s.append(Spacer(1, 6))
    s.append(Paragraph("Schedule 1: Tributaries, compliance stations and allocations", H2))
    rows = [["Tributary", "Compliance station", "Allocation (kg of total phosphorus a year)"]]
    for t in W.TRIBS:
        rows.append([t[1], t[2], "{:,.0f}".format(alloc[t[0]])])
    s.append(_table(rows, [2.2 * inch, 1.5 * inch, 2.6 * inch]))
    s.append(Spacer(1, 6))
    s.append(Paragraph("Allocations as amended on 9 December 2025. Compliance stations are the Commission's; the "
                       "register lists the identifiers other organisations report them under.", SMALL))
    s.append(Spacer(1, 14))
    s.append(Paragraph("For the Commission: %s, %s." % W.PEOPLE["director"], BODY))
    doc.build(s, onFirstPage=frame, onLaterPages=frame)


# ------------------------------------------------------------------------------------ the monitoring program
def write_program(path):
    import docx
    from docx.shared import Pt
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    d.add_heading("Tributary monitoring program, water years 2024 to 2026", level=1)
    d.add_paragraph("Brannoch Lake Commission, water quality program. For staff and partners using the program's "
                    "results. Questions to %s, %s." % W.PEOPLE["manager"])
    d.add_heading("Stations", level=2)
    d.add_paragraph("Each of the ten monitored tributaries has one compliance station near its mouth (the register "
                    "lists them). The Commission reports a station under its own identifier (BLC-01 to BLC-10); the "
                    "%s reports the same sites under its own (ACWL-FB and so on)." % W.LAB)
    d.add_heading("Sampling", level=2)
    d.add_paragraph("Commission staff sample every station on the first and third Tuesdays of each month, working the stations "
                    "in a fixed order through the day, whatever the weather. Under a cooperative agreement, the %s "
                    "samples stations during high water when its crews can reach them; its storm samples are reported "
                    "with the Commission's results. Questions on storm sampling to %s, %s." % ((W.LAB,) + W.PEOPLE["lab"]))
    d.add_paragraph("Each sample is a grab sample from the centre of flow. Water temperature is read in the field.")
    d.add_heading("Laboratories and reporting", level=2)
    d.add_paragraph("Commission samples were analysed by %s to 30 June 2025 and by %s from 1 July 2025. %s reports "
                    "total phosphorus as \"Total Phosphorus, mixed forms\" and dissolved phosphorus as orthophosphate, "
                    "both in micrograms per litre; earlier results are \"Phosphorus\" in milligrams per litre, with the "
                    "fraction (total or dissolved) in the sample fraction field. The college's laboratory reports "
                    "\"Phosphorus\" in milligrams per litre throughout." % (W.LAB_OLD, W.LAB_NEW, W.LAB_NEW))
    d.add_paragraph("Results below the reporting limit are reported as not detected, with the limit. Result status is "
                    "Accepted once a result passes review, or Rejected with the reason in the comment.")
    d.add_heading("Quality control", level=2)
    d.add_paragraph("About one Commission sampling visit in twelve includes a field replicate (a second sample taken "
                    "minutes after the first) and about one in twenty a field blank (laboratory water carried through "
                    "the sampling). Both are reported with their activity type.")
    d.add_heading("Discharge", level=2)
    d.add_paragraph("Discharge is recorded every 15 minutes by the Commission's gage network, operated under contract by "
                    "%s, which delivers the approved record after each water year. Each compliance station has a gage "
                    "at the station except Dunmarsh Brook, whose gage is 4.2 km upstream; the register gives both "
                    "drainage areas." % W.GAGER)
    d.save(path)


# -------------------------------------------------------------------------------------- the gage transmittal
def write_transmittal(path, counts):
    L = []
    a = L.append
    a("MERRIFIELD HYDROMETRICS")
    a("Data transmittal")
    a("Client: %s, tributary gage network" % W.COMMISSION)
    a("Date: 14 December 2026")
    a("")
    a("Delivered:")
    a("  merrifield_unit_values_wy2026_approved.csv   %d rows. Approved 15-minute discharge, water year 2026," % counts["uv"])
    a("                                               1 October 2025 to 30 September 2026, all ten gages.")
    a("  merrifield_daily_values_wy2024-wy2025.csv    %d rows. Approved daily mean discharge, water years 2024" % counts["dv"])
    a("                                               and 2025.")
    a("")
    a("Earlier delivery (superseded): provisional 15-minute values for 1 July to 30 September 2026, delivered")
    a("6 October 2026. The approved values above replace them. Two gages' ratings were shifted after the autumn")
    a("field measurements: 01-4420 (Kestle River) and 01-4480 (Ambrey River), from 1 August 2026.")
    a("")
    a("Conventions:")
    a("  Discharge in cubic feet per second, three significant figures.")
    a("  Unit values are stamped in Coordinated Universal Time. Daily values are means over local calendar days")
    a("  (Eastern time).")
    a("  Qualifiers: A approved; P provisional; e estimated. Winter values at gages affected by ice are estimated")
    a("  from the record before and after the ice, nearby gages and winter measurements.")
    a("")
    a("Gages:")
    for t in W.TRIBS:
        a("  %s  %-18s drainage area %6.1f km2" % (t[4], t[1], t[6]))
    a("")
    a("Contact: %s, %s" % W.PEOPLE["hydro"])
    open(path, "w", encoding="utf-8", newline="\r\n").write("\n".join(L) + "\n")
