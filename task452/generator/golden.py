"""The golden from the bundle alone: conform the interpretation campaigns into one reference class per point on the
national forest definition and the reporting period, estimate the area of deforestation on the strata the sample
was drawn from, turn it into the emission reductions offered and the payment, estimate each district's area and the
interval, count the final map's deforestation patches by district (in all and of 5 hectares or more), and write the
monitoring report summary and the conformed reference sample."""
import io
import json
import math
import os
import re

import numpy as np
import pandas as pd

import world as W

PDF_NAME = "er_monitoring_report_2024-2025.pdf"
XLSX_NAME = "reference_sample_conformed_2024-2025.xlsx"
ES_LC = {"bosque denso": 1, "bosque abierto": 2, "arbustal": 3, "plantación forestal": 4, "cultivo": 5, "pasto": 6,
         "urbano o suelo desnudo": 7, "agua o humedal": 8}
ES_CANOPY = {"<10%": "A", "10-30%": "B", "30-60%": "C", ">60%": "D"}
CLASS_NAME = {"D": "deforestation 2024-2025", "F": "stable forest", "N": "non-forest"}
Z90 = 1.6448536269514722


def pdf_text(path):
    from pypdf import PdfReader
    return "\n".join(p.extract_text() or "" for p in PdfReader(path).pages)


def read_strata(target):
    txt = pdf_text(os.path.join(target, "reference_sampling_design_2026-01.pdf"))
    out = {}
    for s in ("S1", "S2", "S3"):
        m = re.search(r"\b%s\s*\n[^\n]*\n\s*([\d,]{6,})\s+([\d,]+)" % s, txt)
        out[s] = (int(m.group(1).replace(",", "")), int(m.group(2).replace(",", "")))
    return out


def reference_class(lc24, cc24, lc25, cc25, loss_year):
    forest24 = lc24 in (1, 2) and cc24 in ("C", "D")
    forest25 = lc25 in (1, 2) and cc25 in ("C", "D")
    if not forest24:
        return "N"
    if forest25:
        return "F"
    return "D" if loss_year in (2024, 2025) else "N"


def read_points(target):
    """the main campaign with every adjudicated field applied"""
    m = pd.read_csv(os.path.join(target, "reference_points_main_campaign_2026.csv"), dtype={"district": str})
    adj = pd.read_excel(os.path.join(target, "adjudication_log_2026-04.xlsx"), header=2)
    m = m.set_index("plot_id")
    for r in adj.itertuples():
        col = {"lc24": "lc_2024", "cc24": "canopy_2024", "lc25": "lc_2025", "cc25": "canopy_2025",
               "loss_year": "loss_year", "loss_month": "loss_month", "confidence": "confidence"}[r.field]
        m.loc[r.plot_id, col] = r.adjudicated_value
    m = m.reset_index()
    m["ref"] = [reference_class(int(a), b, int(c), d, (int(e) if pd.notna(e) else None)) for a, b, c, d, e in
                zip(m.lc_2024, m.canopy_2024, m.lc_2025, m.canopy_2025, m.loss_year)]
    fin = pd.read_csv(os.path.join(target, "reference_points_final_map_classes_2026-04-22.csv"))
    return m.merge(fin, on="plot_id", how="left")


def compute(target):
    strata = read_strata(target)
    pts = read_points(target)
    n = pts.groupby("stratum").size()
    pts["weight"] = [strata[s][0] / n[s] for s in pts.stratum]
    pts["y"] = (pts.ref == "D").astype(float)
    area = float((pts.weight * pts.y).sum())
    var = 0.0
    for s in strata:
        ph = pts.loc[pts.stratum == s, "y"].mean()
        var += strata[s][0] ** 2 * ph * (1 - ph) / (n[s] - 1)
    se = math.sqrt(var)
    by_d = {c: float((pts.weight * pts.y)[pts.district == c].sum()) for c in W.DCODE}
    import methods as M
    er = M.emission_reductions(area)
    stats = pd.read_excel(os.path.join(target, "final_change_map_v2_statistics.xlsx"), header=3)
    conc = pd.read_csv(os.path.join(target, "district_concordance_2019_2025.csv"), dtype=str)
    to19 = dict(zip(conc.district_2025, conc.district_2019))
    prog = set(conc.loc[conc.accounting_area == "program", "district_2025"])
    stats = stats[stats["district_2025"].isin(prog)]
    map_fd = {c: 0.0 for c in W.DCODE}
    for r in stats.itertuples():
        map_fd[to19[r.district_2025]] += float(r[3])
    poly = pd.read_parquet(os.path.join(target, "final_change_map_v2_polygons.parquet"))
    poly = poly[poly.district_2025.isin(prog)].copy()
    poly["d19"] = poly.district_2025.map(to19)
    # a patch delivered in several polygons (tiles, quarters) carries its area on each: each polygon is the share
    # area_ha / patch_ha of one patch, so the shares sum to the number of patches
    poly["share"] = poly.area_ha / poly.patch_ha
    patches_all = {c: int(round(poly.share[poly.d19 == c].sum())) for c in W.DCODE}
    patches = {c: int(round(poly.share[(poly.d19 == c) & (poly.patch_ha >= 5.0)].sum())) for c in W.DCODE}
    small = {c: int(round(poly.share[(poly.d19 == c) & (poly.patch_ha < 1.0)].sum())) for c in W.DCODE}
    return dict(strata=strata, pts=pts, area=area, se=se, lo=area - Z90 * se, hi=area + Z90 * se, by_d=by_d, er=er,
                map_fd=map_fd, patches=patches, patches_all=patches_all, small=small)


def hundred(x):
    return int(100 * math.floor(x / 100 + 0.5))


def ten_thousand(x):
    return int(10_000 * math.floor(x / 10_000 + 0.5))


def fmt_int(v):
    return format(int(v), ",")


def chart_png(X):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams["font.family"] = "DejaVu Sans"
    order = sorted(W.DCODE, key=lambda c: -X["by_d"][c])
    names = [W.DNAME[c] for c in order]
    est = [X["by_d"][c] / 1000 for c in order]
    mp = [X["map_fd"][c] / 1000 for c in order]
    y = np.arange(len(order))
    fig, ax = plt.subplots(figsize=(7.0, 4.0), dpi=150)
    ax.barh(y - 0.2, est, height=0.38, color="#1f5a3d", label="Estimated from the reference sample")
    ax.barh(y + 0.2, mp, height=0.38, color="#a9b8a3", label="Final change map, version 2.0")
    ax.set_yticks(y)
    ax.set_yticklabels(names, fontsize=7.5)
    ax.invert_yaxis()
    ax.set_xlabel("Deforestation 2024-2025 (thousand hectares)", fontsize=8)
    ax.tick_params(axis="x", labelsize=7.5)
    for sp in ("top", "right"):
        ax.spines[sp].set_visible(False)
    ax.legend(loc="lower right", fontsize=7.5, frameon=False)
    ax.set_title("The final map shows %s of the %s hectares the reference sample finds" %
                 ("{:.0f} per cent".format(100 * sum(X["map_fd"].values()) / X["er"]["area_ha"]),
                  fmt_int(X["er"]["area_ha"])), fontsize=9.5, loc="left")
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", metadata={"Software": None})
    plt.close(fig)
    return buf.getvalue()


def write_pdf(X, path):
    from reportlab import rl_config
    rl_config.invariant = 1
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    import docs  # noqa: F401  (registers the fonts)
    INK, GREY, GREEN = colors.HexColor("#1b1d20"), colors.HexColor("#5b6067"), colors.HexColor("#1f5a3d")
    body = ParagraphStyle("b", fontName="LiberationSans", fontSize=9.6, leading=13.6, textColor=INK, spaceAfter=6)
    h1 = ParagraphStyle("h1", fontName="LiberationSans-Bold", fontSize=14, leading=18, textColor=GREEN, spaceAfter=4)
    sub = ParagraphStyle("s", fontName="LiberationSans", fontSize=9, leading=12.5, textColor=GREY, spaceAfter=10)
    h2 = ParagraphStyle("h2", fontName="LiberationSans-Bold", fontSize=10.5, leading=14, textColor=GREEN,
                        spaceBefore=8, spaceAfter=3)
    cell = ParagraphStyle("c", fontName="LiberationSans", fontSize=8.6, leading=11, textColor=INK)
    num = ParagraphStyle("num", parent=cell, alignment=2)
    note = ParagraphStyle("n", fontName="LiberationSans-Italic", fontSize=7.8, leading=10.5, textColor=GREY,
                          spaceBefore=3, spaceAfter=8)

    def frame(canvas, doc):
        canvas.saveState()
        canvas.setFont("LiberationSans", 7.5)
        canvas.setFillColor(GREY)
        canvas.drawString(20 * mm, 12 * mm, "%s  |  Monitoring report 2024-2025, summary" % W.PROGRAM)
        canvas.drawRightString(A4[0] - 20 * mm, 12 * mm, "%d" % doc.page)
        canvas.restoreState()

    er = X["er"]
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=18 * mm,
                            bottomMargin=20 * mm, title="Monitoring report 2024-2025: emission reductions",
                            author=W.LAB, subject="", creator="", producer="")
    lo, hi = ten_thousand(X["lo"]), ten_thousand(X["hi"])
    s = [Paragraph("The program can offer %s tonnes for USD %s" % (fmt_int(er["offered_t"]), fmt_int(er["payment_usd"])),
                   h1),
         Paragraph("%s, monitoring report for 1 January 2024 to 31 December 2025: summary for the %s. %s, %s. 9 June "
                   "2026." % (W.PROGRAM, W.FUND, W.LAB, W.PEOPLE["head"]), sub),
         Paragraph("Emission reductions", h2),
         Paragraph("Deforestation in the accounting area over the reporting period was <b>%s hectares</b> (90 per cent "
                   "confidence interval %s to %s hectares), against a reference level of %s hectares for the two years. "
                   "At 497.0 tonnes of CO2e a hectare the program emitted %s tonnes against reference level emissions "
                   "of %s, for emission reductions of %s tonnes. After the 18 per cent buffer the program can offer "
                   "<b>%s tonnes</b>, for which the %s pays <b>USD %s</b>." %
                   (fmt_int(er["area_ha"]), fmt_int(lo), fmt_int(hi), fmt_int(2 * W.RL_HA_YEAR),
                    fmt_int(er["monitored_t"]), fmt_int(er["rl_t"]), fmt_int(er["er_t"]), fmt_int(er["offered_t"]),
                    W.FUND, fmt_int(er["payment_usd"])), body),
         Paragraph("How the area was estimated", h2),
         Paragraph("The 4,000 reference points were drawn on 20 January 2026 by stratified random sampling on the "
                   "preliminary change map, version 0.9: 400 in its deforestation class, 2,400 in stable forest and "
                   "1,200 in non-forest. The final map, version 2.0, moved most of the preliminary map's false alarms, "
                   "and some real clearings, out of its deforestation class, but a point's weight is set by the "
                   "stratum it was drawn from: %s hectares for a point of the deforestation stratum, %s for stable forest and %s for "
                   "non-forest. The area is the sum of those weights over the %s points the interpreters class as "
                   "deforestation in 2024-2025 on the national forest definition, after quality control." %
                   tuple([fmt_int(round(X["strata"][s][0] / X["strata"][s][1])) for s in ("S1", "S2", "S3")] +
                         [fmt_int(int(X["pts"].y.sum()))]), body),
         Image(io.BytesIO(chart_png(X)), width=170 * mm, height=97 * mm),
         Paragraph("Estimates from the reference sample; the final map's areas by pixel counting. Source: %s reference sample 2024-2025 and final change map version 2.0." % W.NFMS,
                   note),
         PageBreak(), Paragraph("Deforestation by district", h2)]
    rows = [[Paragraph("<b>District</b>", cell), Paragraph("<b>Estimated deforestation (ha)</b>", num),
             Paragraph("<b>Deforestation patches on the final map</b>", num),
             Paragraph("<b>Under 1 ha</b>", num), Paragraph("<b>5 ha or more</b>", num)]]
    for c in W.DCODE:
        rows.append([Paragraph(W.DNAME[c], cell), Paragraph(fmt_int(hundred(X["by_d"][c])), num),
                     Paragraph(fmt_int(X["patches_all"][c]), num), Paragraph(fmt_int(X["small"][c]), num),
                     Paragraph(fmt_int(X["patches"][c]), num)])
    rows.append([Paragraph("<b>Accounting area</b>", cell), Paragraph("<b>%s</b>" % fmt_int(er["area_ha"]), num),
                 Paragraph("<b>%s</b>" % fmt_int(sum(X["patches_all"].values())), num),
                 Paragraph("<b>%s</b>" % fmt_int(sum(X["small"].values())), num),
                 Paragraph("<b>%s</b>" % fmt_int(sum(X["patches"].values())), num)])
    t = Table(rows, colWidths=[42 * mm, 34 * mm, 42 * mm, 24 * mm, 26 * mm])
    t.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, 0), 0.6, GREEN), ("LINEABOVE", (0, -1), (-1, -1), 0.4, GREY),
                           ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 2.4),
                           ("BOTTOMPADDING", (0, 0), (-1, -1), 2.4)]))
    s += [t, Paragraph("District estimates sum the weights of each district's deforestation points; they carry far wider "
                       "intervals than the accounting area's and are reported for the benefit-sharing plan. A patch is a "
                       "connected area of change on the national mosaic, counted once however many tiles and quarters "
                       "its polygons are delivered in.", note),
          Spacer(1, 6), Paragraph("%s<br/>%s" % (W.PEOPLE["head"], W.LAB), body)]
    doc.build(s, onFirstPage=frame, onLaterPages=frame)


def write_xlsx(X, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Reference points"
    hdr = ["plot_id", "stratum", "district_2019", "reference_class_2024_2025", "map_class_v2", "weight_ha"]
    for j, h in enumerate(hdr, 1):
        ws.cell(row=1, column=j, value=h).font = Font(name="Calibri", bold=True, size=10)
    fin = {"DEF": "deforestation 2024-2025", "BOS": "stable forest", "NOB": "non-forest"}
    for i, r in enumerate(X["pts"].sort_values("plot_id").itertuples(), 2):
        vals = [r.plot_id, r.stratum, r.district, CLASS_NAME[r.ref], fin[r.map_class_v2], round(r.weight, 2)]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=v)
            if j == 6:
                c.number_format = "#,##0.00"
    for col, w in zip("ABCDEF", (14, 9, 13, 26, 24, 12)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:F%d" % (len(X["pts"]) + 1)
    s2 = wb.create_sheet("Notes")
    notes = ["Reference sample 2024-2025, conformed: one row per reference point.",
             "reference_class_2024_2025: the point's class for the reporting period on the national forest definition "
             "(monitoring plan, section 3), from the main campaign with the adjudicated quality-control differences "
             "applied.",
             "weight_ha: the area the point stands for, its stratum's area over the stratum's points (sample design, "
             "section 2).",
             "map_class_v2: the point's class on the final change map, version 2.0."]
    for i, t in enumerate(notes, 1):
        s2.cell(row=i, column=1, value=t)
    s2.column_dimensions["A"].width = 110
    wb.save(path)


def figures(X):
    er = X["er"]
    return {
        "offered_t": er["offered_t"], "payment_usd": int(er["payment_usd"]), "area_ha": er["area_ha"],
        "area_unrounded_ha": round(X["area"], 1), "er_t": int(er["er_t"]), "monitored_t": int(er["monitored_t"]),
        "rl_t": int(er["rl_t"]), "rl_ha_year": int(W.RL_HA_YEAR), "rl_two_years_ha": int(2 * W.RL_HA_YEAR),
        "buffer_pct": 18, "price_usd": 5, "ci_level_pct": 90, "forest_canopy_pct": 30,
        "ef_t_per_ha": W.EF, "ci90_low_ha": ten_thousand(X["lo"]), "ci90_high_ha": ten_thousand(X["hi"]),
        "points_deforestation": int(X["pts"].y.sum()),
        "weights_ha": {s: round(X["strata"][s][0] / X["strata"][s][1]) for s in ("S1", "S2", "S3")},
        "district_area_ha": {W.DNAME[c]: hundred(X["by_d"][c]) for c in W.DCODE},
        "patches": {W.DNAME[c]: X["patches_all"][c] for c in W.DCODE},
        "patches_all": sum(X["patches_all"].values()),
        "patches_under_1ha": {W.DNAME[c]: X["small"][c] for c in W.DCODE},
        "patches_under_1ha_all": sum(X["small"].values()),
        "patches_5ha": {W.DNAME[c]: X["patches"][c] for c in W.DCODE},
        "patches_5ha_all": sum(X["patches"].values()),
        "map_fd_ha": int(round(sum(X["map_fd"].values()))),
        "reference_points": 4000,
        "reference_classes": {CLASS_NAME[k]: int((X["pts"].ref == k).sum()) for k in ("D", "F", "N")},
        "final_map_classes": {n: int((X["pts"].map_class_v2 == k).sum()) for k, n in
                              (("DEF", "deforestation 2024-2025"), ("BOS", "stable forest"), ("NOB", "non-forest"))},
    }


def main(target, out_dir, fig_path):
    os.makedirs(out_dir, exist_ok=True)
    X = compute(target)
    write_pdf(X, os.path.join(out_dir, PDF_NAME))
    write_xlsx(X, os.path.join(out_dir, XLSX_NAME))
    import postprocess as PP
    PP.normalize_pdf(os.path.join(out_dir, PDF_NAME), b"D:20260609104512-05'00'", b"D:20260609111938-05'00'")
    PP.normalize_ooxml(os.path.join(out_dir, XLSX_NAME), W.PEOPLE["reference"], W.PEOPLE["head"],
                       "2026-06-02T14:21:00Z", "2026-06-09T15:47:12Z", excel=True, company=W.LAB)
    with open(fig_path, "w") as f:
        json.dump(figures(X), f, indent=1, ensure_ascii=False)
        f.write("\n")
    return X
