"""Independent verifier: reads only the shipped bundle and the goldens, on its own code path (csv, openpyxl, pyarrow and
pypdf readers), and recomputes the conformed reference classes, the area of deforestation on the sampling strata,
the tonnes offered and the payment, the interval, each district's area and the final map's deforestation patches by
district, in all and of 5 ha or more."""
import csv
import math
import os
import re
import sys

MAIN = "reference_points_main_campaign_2026.csv"
ADJ = "adjudication_log_2026-04.xlsx"
DESIGN = "reference_sampling_design_2026-01.pdf"
ERPA = "emission_reductions_payment_agreement_extract.pdf"
POLY = "final_change_map_v2_polygons.parquet"
CONC = "district_concordance_2019_2025.csv"
PDF = "er_monitoring_report_2024-2025.pdf"
XLSX = "reference_sample_conformed_2024-2025.xlsx"
FIELD = {"lc24": "lc_2024", "cc24": "canopy_2024", "lc25": "lc_2025", "cc25": "canopy_2025", "loss_year": "loss_year",
         "loss_month": "loss_month", "confidence": "confidence"}


def text(path):
    from pypdf import PdfReader
    return " ".join(" ".join((p.extract_text() or "").split()) for p in PdfReader(path).pages)


def points(target):
    from openpyxl import load_workbook
    rows = {r["plot_id"]: r for r in csv.DictReader(open(os.path.join(target, MAIN), newline="", encoding="utf-8"))}
    ws = load_workbook(os.path.join(target, ADJ), read_only=True, data_only=True).active
    for r in ws.iter_rows(min_row=4, values_only=True):
        if r[0]:
            rows[r[0]][FIELD[r[1]]] = str(r[4])
    for r in rows.values():
        f24 = r["lc_2024"] in ("1", "2") and r["canopy_2024"] in ("C", "D")
        f25 = r["lc_2025"] in ("1", "2") and r["canopy_2025"] in ("C", "D")
        yr = int(float(r["loss_year"])) if r["loss_year"] else None
        r["cls"] = "N" if not f24 else ("F" if f25 else ("D" if yr in (2024, 2025) else "N"))
    return rows


def strata(target):
    t = text(os.path.join(target, DESIGN))
    out = {}
    for s in ("S1", "S2", "S3"):
        m = re.search(r"%s (?:deforestation 2024-2025|stable forest|non-forest) ([\d,]+) ([\d,]+)" % s, t)
        out[s] = (int(m.group(1).replace(",", "")), int(m.group(2).replace(",", "")))
    return out


def main(target, golden):
    fails = []
    pts = points(target)
    st = strata(target)
    w = {s: a / n for s, (a, n) in st.items()}
    counts = {}
    for r in pts.values():
        counts[r["stratum"]] = counts.get(r["stratum"], 0) + 1
    assert all(counts[s] == st[s][1] for s in st)
    area = sum(w[r["stratum"]] for r in pts.values() if r["cls"] == "D")
    var = 0.0
    for s, (a, n) in st.items():
        p = sum(1 for r in pts.values() if r["stratum"] == s and r["cls"] == "D") / n
        var += a * a * p * (1 - p) / (n - 1)
    se = math.sqrt(var)
    erpa = text(os.path.join(target, ERPA))
    rl = float(re.search(r"reference level is ([\d,]+) hectares", erpa).group(1).replace(",", ""))
    ef = float(re.search(r"([\d.]+) tonnes of carbon dioxide equivalent per hectare", erpa).group(1))
    rep = int(100 * math.floor(area / 100 + 0.5))
    er = (2 * rl - rep) * ef
    offer = max(0, math.floor(er * 0.82))
    pay = offer * 5
    lo, hi = area - 1.6448536 * se, area + 1.6448536 * se
    by = {}
    for r in pts.values():
        if r["cls"] == "D":
            by[r["district"]] = by.get(r["district"], 0.0) + w[r["stratum"]]
    import pyarrow.parquet as pq
    t = pq.read_table(os.path.join(target, POLY)).to_pydict()
    conc = list(csv.DictReader(open(os.path.join(target, CONC), newline="", encoding="utf-8")))
    to19 = {r["district_2025"]: r["district_2019"] for r in conc if r["accounting_area"] == "program"}
    dname = {r["district_2019"]: r["name_2019"] for r in conc}
    every, small, big = {}, {}, {}
    for d, a, ph in zip(t["district_2025"], t["area_ha"], t["patch_ha"]):
        if d in to19:
            every[to19[d]] = every.get(to19[d], 0.0) + a / ph
            small[to19[d]] = small.get(to19[d], 0.0) + (a / ph if ph < 1 else 0.0)
            big[to19[d]] = big.get(to19[d], 0.0) + (a / ph if ph >= 5 else 0.0)
    print("verify: area %.1f ha (reported %d), se %.1f, offered %d t, payment USD %d" % (area, rep, se, offer, pay))
    print("verify: patches %d, under 1 ha %d, of 5 ha or more %d" % (round(sum(every.values())),
                                                                     round(sum(small.values())), round(sum(big.values()))))
    rep_txt = text(os.path.join(golden, PDF))
    for need in (format(offer, ","), format(pay, ","), format(rep, ","), format(int(10_000 * math.floor(lo / 10_000 + 0.5)),
                 ","), format(int(10_000 * math.floor(hi / 10_000 + 0.5)), ",")):
        if need not in rep_txt:
            fails.append("report lacks %s" % need)
    for d, v in by.items():
        if format(int(100 * math.floor(v / 100 + 0.5)), ",") not in rep_txt:
            fails.append("report lacks district %s area %d" % (d, round(v)))
    for d in every:
        if max(abs(v[d] - round(v[d])) for v in (every, small, big)) > 1e-6:
            fails.append("district %s: polygon shares do not sum to whole patches" % d)
        if not re.search(r"%s [\d,]+ %s %s %s\b" % (re.escape(dname[d]), format(int(round(every[d])), ","),
                                                     format(int(round(small[d])), ","),
                                                     format(int(round(big[d])), ",")), rep_txt):
            fails.append("report lacks district %s patches" % d)
    from openpyxl import load_workbook
    ws = load_workbook(os.path.join(golden, XLSX), read_only=True, data_only=True)["Reference points"]
    name = {"D": "deforestation 2024-2025", "F": "stable forest", "N": "non-forest"}
    seen = 0
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[0]:
            continue
        seen += 1
        if r[3] != name[pts[r[0]]["cls"]] or r[1] != pts[r[0]]["stratum"]:
            fails.append("xlsx %s" % r[0])
    if seen != 4000:
        fails.append("xlsx rows %d" % seen)
    if offer != 23_229_780 or rep != 122_900:
        fails.append("call: %d t on %d ha" % (offer, rep))
    for f in fails[:12]:
        print("VERIFY FAIL  " + f)
    print("verify: %s (%d mismatches)" % ("PASS" if not fails else "FAIL", len(fails)))
    if fails:
        raise SystemExit(1)


if __name__ == "__main__":
    here = os.path.dirname(os.path.abspath(__file__))
    task = os.path.dirname(here)
    main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(task, "target"),
         sys.argv[2] if len(sys.argv) > 2 else os.path.join(task, "golden"))
