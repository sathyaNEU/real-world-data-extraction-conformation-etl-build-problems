"""The pack's data files: the two interpretation campaigns' exports, the adjudication log, each reference point's
class on the final change map, the final map's statistics and polygon delivery, the district concordance, the data
dictionary, the export manifest, and two files the monitoring report does not read (the national forest
inventory's tree measurements and the fire hotspot detections)."""
import csv
import datetime as dt
import math

import numpy as np
import pyarrow as pa
import pyarrow.parquet as pq

import interp as I
import mapping as MP
import world as W

FILES = {
    "main": "reference_points_main_campaign_2026.csv",
    "qc": "control_calidad_reinterpretacion_2026.json",
    "adj": "adjudication_log_2026-04.xlsx",
    "intersect": "reference_points_final_map_classes_2026-04-22.csv",
    "stats": "final_change_map_v2_statistics.xlsx",
    "poly": "final_change_map_v2_polygons.parquet",
    "conc": "district_concordance_2019_2025.csv",
    "dict": "data_dictionary.csv",
    "manifest": "export_manifest.csv",
    "design": "reference_sampling_design_2026-01.pdf",
    "erpa": "emission_reductions_payment_agreement_extract.pdf",
    "plan": "monitoring_plan_extract.pdf",
    "frel": "forest_reference_level_2015-2022_summary.pdf",
    "release": "final_change_map_v2_release_note.pdf",
    "nfi": "nfi_tree_measurements_2018-2022.csv",
    "fire": "fire_hotspots_2024-2025.csv",
}


def _csv(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\r\n")
        w.writerow(header)
        for r in rows:
            w.writerow(["" if v is None else v for v in r])
    return len(rows)


REMARKS = {  # the interpreters' free-text remarks by what the point is, most points carrying none
    "deforestation": ([("", 45), ("cleared for pasture", 20), ("cleared, crop rows in 2025 image", 8),
                       ("burnt then cleared", 9), ("access road visible", 7), ("partial cloud in 2024 image, loss clear", 4),
                       ("selective logging then clearing", 4), ("new fence line", 3)]),
    "stable forest": ([("", 66), ("closed canopy", 15), ("riparian strip", 8), ("selective logging trails, canopy kept", 5),
                       ("palm-rich canopy", 3), ("haze in 2025 image, canopy intact", 3)]),
    "farmland": ([("", 55), ("improved pasture", 20), ("rotational crop", 12), ("pasture with scattered trees", 8),
                  ("fallow", 5)]),
    "woodland": ([("", 50), ("scattered trees", 25), ("shrub regrowth", 15), ("savanna woodland", 10)]),
    "plantation": ([("", 35), ("teak stand", 25), ("eucalyptus rows", 20), ("young plantation", 12), ("oil palm", 8)]),
    "other": ([("", 45), ("river", 15), ("village", 15), ("road", 10), ("wetland", 10), ("quarry", 5)]),
    "regrowth": ([("", 40), ("secondary regrowth on old pasture", 35), ("regrowth after abandonment", 25)])}


def write_main(pts, path):
    rng = np.random.default_rng(W.SEED + 31)
    rows = []
    for p in sorted(pts, key=lambda p: p["plot_id"]):
        if p["kind"] == "cleared before 2024":
            rem = "cleared %d, pasture since" % p["cleared"]
        else:
            opts = REMARKS[p["kind"]]
            wts = np.array([w for _, w in opts], dtype=float)
            rem = opts[int(rng.choice(len(opts), p=wts / wts.sum()))][0]
        rows.append((p["plot_id"], p["stratum"], p["district"], p["x"], p["y"], p["operator"],
                     p["interpreted"].strftime("%d/%m/%Y"), p["lc24"], p["cc24"], p["lc25"], p["cc25"], p["plantation"],
                     p["change"], p["loss_year"], p["loss_month"], p["img24"].isoformat(), p["img25"].isoformat(),
                     p["confidence"], rem))
    return _csv(path, ["plot_id", "stratum", "district", "x_ntm", "y_ntm", "operator", "date_interpreted", "lc_2024",
                       "canopy_2024", "lc_2025", "canopy_2025", "plantation", "change", "loss_year", "loss_month",
                       "image_2024", "image_2025", "confidence", "remarks"], rows)


def write_qc(qc, path):
    """the ministry's form tool export: one JSON record per submission, the tool's own _id and submission time first"""
    import json
    rng = np.random.default_rng(W.SEED + 33)
    timed = sorted(((dt.datetime.combine(q["interpreted"], dt.time(8, 0)) +
                     dt.timedelta(minutes=int(rng.integers(0, 600)), seconds=int(rng.integers(0, 60))), q["plot_id"], q)
                    for q in qc), key=lambda r: (r[0], r[1]))
    recs = []
    for k, (t, _, q) in enumerate(timed):
        recs.append({"_id": 418_207 + 3 * k + int(rng.integers(0, 3)),
                     "_submission_time": t.strftime("%Y-%m-%dT%H:%M:%S"),
                     "id_parcela": q["plot_id"].replace("LFM26-", "LFM26/"), "estrato": q["stratum"],
                     "interprete": q["operator"], "fecha": q["interpreted"].isoformat(),
                     "cobertura_2024": I.ES_LC[q["lc24"]], "dosel_2024": I.ES_CANOPY[q["cc24"]],
                     "cobertura_2025": I.ES_LC[q["lc25"]], "dosel_2025": I.ES_CANOPY[q["cc25"]],
                     "cambio": I.ES_CHANGE[q["change"]], "anio_perdida": q["loss_year"], "mes_perdida": q["loss_month"],
                     "confianza": I.ES_CONF[q["confidence"]]})
    with open(path, "w", encoding="utf-8") as f:
        json.dump(recs, f, ensure_ascii=False, indent=1)
        f.write("\n")
    return len(recs)


def write_adjudication(adj, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Adjudication"
    ws.cell(row=1, column=1, value="Reference sample 2024-2025: adjudication of quality-control disagreements").font = \
        Font(name="Calibri", bold=True, size=11)
    hdr = ["plot_id", "field", "main_value", "qc_value", "adjudicated_value", "adjudicator", "date"]
    for j, h in enumerate(hdr, 1):
        ws.cell(row=3, column=j, value=h).font = Font(name="Calibri", bold=True, size=10)
    for i, a in enumerate(sorted(adj, key=lambda a: (a["on"], a["plot_id"], a["field"])), 4):
        vals = [a["plot_id"], a["field"], a["main"], a["qc"], a["adjudicated"], a["by"], a["on"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=v)
            if isinstance(v, dt.date):
                c.number_format = "dd/mm/yyyy"
    for col, w in zip("ABCDEFG", (14, 12, 12, 12, 17, 11, 11)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"
    wb.save(path)
    return len(adj)


def write_intersect(pts, path):
    rows = []
    for p in sorted(pts, key=lambda p: p["plot_id"]):
        d25 = W.D2025[p["district"]]
        if p["district"] == "05" and p["x"] > I.district_boxes()["05"][0] + 0.62 * I.district_boxes()["05"][2]:
            d25 = W.SPLIT_NEW[0]
        rows.append((p["plot_id"], I.FINAL_CODE[p["final"]], d25))
    return _csv(path, ["plot_id", "map_class_v2", "district_2025"], rows)


def write_stats(path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    areas = MP.class_areas_2025()
    names = {W.D2025[c]: n for c, n in W.DNAME.items()}
    names[W.SPLIT_NEW[0]] = W.SPLIT_NEW[1]
    names.update({W.D2025[c]: n for c, n, _ in W.HIGHLANDS})
    wb = Workbook()
    ws = wb.active
    ws.title = "Class areas"
    bold = Font(name="Calibri", bold=True, size=11)
    ws.cell(row=1, column=1, value="Final forest change map 2024-2025, version 2.0 (22 April 2026): class areas by "
                                    "district").font = bold
    ws.cell(row=2, column=1, value="Hectares, by pixel counting on the 30 m grid. Districts on the 2025 district "
                                    "layer.").font = Font(name="Calibri", size=9, italic=True)
    hdr = ["district_2025", "district", "deforestation 2024-2025 (ha)", "stable forest (ha)", "non-forest (ha)",
           "total (ha)"]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(row=4, column=j, value=h)
        c.font = Font(name="Calibri", bold=True, size=10)
        c.alignment = Alignment(wrap_text=True)
    k = 5
    tot = {"FD": 0.0, "FF": 0.0, "FN": 0.0}
    for code in sorted(areas):
        a = areas[code]
        vals = [round(a["FD"], 2), round(a["FF"], 2), round(a["FN"], 2), round(a["FD"] + a["FF"] + a["FN"], 2)]
        ws.cell(row=k, column=1, value=code)
        ws.cell(row=k, column=2, value=names[code])
        for j, v in enumerate(vals, 3):
            ws.cell(row=k, column=j, value=v).number_format = "#,##0.00"
        for f in tot:
            tot[f] += a[f]
        k += 1
    ws.cell(row=k, column=2, value="Country").font = Font(name="Calibri", bold=True, size=10)
    for j, f in enumerate(("FD", "FF", "FN"), 3):
        ws.cell(row=k, column=j, value=round(tot[f], 2)).number_format = "#,##0.00"
    ws.cell(row=k, column=6, value=round(sum(tot.values()), 2)).number_format = "#,##0.00"
    for col, w in zip("ABCDEF", (14, 24, 16, 16, 16, 16)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C5"
    wb.save(path)
    return len(areas)


def write_polygons(path):
    pol = MP.polygons()
    cols = {"polygon_id": [], "tile": [], "district_2025": [], "first_detected": [], "area_ha": [], "patch_ha": [],
            "x_ntm": [], "y_ntm": []}
    for r in pol:
        cols["polygon_id"].append(r["polygon_id"])
        cols["tile"].append(r["tile"])
        cols["district_2025"].append(r["district"])
        cols["first_detected"].append(r["detected"])
        cols["area_ha"].append(round(r["area_pix"] * W.PIXEL_HA, 2))
        cols["patch_ha"].append(round(r["patch_pix"] * W.PIXEL_HA, 2))
        cols["x_ntm"].append(r["x"])
        cols["y_ntm"].append(r["y"])
    schema = pa.schema([("polygon_id", pa.string()), ("tile", pa.string()), ("district_2025", pa.string()),
                        ("first_detected", pa.date32()), ("area_ha", pa.float64()), ("patch_ha", pa.float64()),
                        ("x_ntm", pa.int32()), ("y_ntm", pa.int32())])
    pq.write_table(pa.table(cols, schema=schema), path, compression="snappy", use_dictionary=True,
                   write_statistics=False, version="2.6", data_page_version="1.0")
    return len(pol)


def write_concordance(path):
    rows = []
    for c, n in W.DNAME.items():
        rows.append((c, n, W.D2025[c], n, "program"))
        if c == "05":
            rows.append((c, n, W.SPLIT_NEW[0], W.SPLIT_NEW[1], "program"))
    for c, n, _ in W.HIGHLANDS:
        rows.append((c, n, W.D2025[c], n, "outside"))
    return _csv(path, ["district_2019", "name_2019", "district_2025", "name_2025", "accounting_area"], rows)


def write_nfi(path):
    """the national forest inventory's tree measurements, 2018-2022 (the reference level's emission factor was
    derived from them; the monitoring report reads the factor, not the trees)"""
    rng = np.random.default_rng(W.SEED + 41)
    species = ["CEDO", "SWMA", "CAGU", "VIRO", "TERM", "CEIB", "BROS", "PSEU", "INGA", "CECR", "OCOT", "POUT",
               "LICA", "PROT", "SLOA", "CALO"]
    density = [0.42, 0.53, 0.61, 0.38, 0.66, 0.29, 0.71, 0.58, 0.55, 0.34, 0.47, 0.63, 0.69, 0.52, 0.81, 0.57]
    rows = []
    plots_by_year = {2018: 104, 2019: 131, 2020: 97, 2021: 146, 2022: 132}      # the inventory's measurement cycle
    year_of = [y for y, n in plots_by_year.items() for _ in range(n)]
    for plot in range(1, 611):
        year = year_of[plot - 1]
        n = int(rng.integers(38, 92))
        for t in range(1, n + 1):
            k = int(rng.integers(0, len(species)))
            dbh = round(float(10 + rng.gamma(1.6, 11.0)), 1)
            h = round(float(min(48.0, 1.3 + 2.9 * dbh ** 0.62 * rng.lognormal(0, 0.12))), 1)
            rows.append(("NFI-%04d" % plot, year, t, species[k], dbh, h, density[k]))
    return _csv(path, ["plot", "year", "tree", "species", "dbh_cm", "height_m", "wood_density"], rows)


def write_fire(path):
    """fire hotspot detections over the country, 2024-2025 (thermal anomalies, not change)"""
    rng = np.random.default_rng(W.SEED + 43)
    boxes = MP.district2025_boxes()
    codes = sorted(boxes)
    rows = []
    for _ in range(17_340):
        d = codes[int(rng.integers(0, len(codes)))]
        x0, y0, w, h = boxes[d]
        day = dt.date(2024, 1, 1) + dt.timedelta(days=int(rng.integers(0, 731)))
        m = day.month
        if m in (5, 6, 7, 8, 9, 10) and rng.random() < 0.7:                  # most burning in the dry season
            day = dt.date(day.year, int(rng.choice([1, 2, 3, 4, 12])), min(day.day, 28))
        rows.append((day.isoformat(), "%02d:%02d" % (6 if rng.random() < 0.55 else 18, int(rng.integers(0, 60))),
                     int(x0 + rng.random() * w), int(y0 + rng.random() * h), d,
                     round(float(rng.lognormal(1.9, 0.9)), 1), str(rng.choice(["n", "n", "h", "l"]))))
    rows.sort()
    return _csv(path, ["date", "time_utc", "x_ntm", "y_ntm", "district_2025", "frp_mw", "confidence"], rows)


DICTIONARY = [
    (FILES["main"], "", "Main interpretation campaign, February and March 2026: one row per reference point."),
    ("", "plot_id", "Reference point."),
    ("", "stratum", "Sampling stratum of the point (see the sample design report)."),
    ("", "district", "District of the point on the 2019 district layer."),
    ("", "x_ntm", "Easting, national transverse Mercator grid (metres)."),
    ("", "y_ntm", "Northing, national transverse Mercator grid (metres)."),
    ("", "operator", "Interpreter."),
    ("", "date_interpreted", "Date of interpretation."),
    ("", "lc_2024", "Land cover in January 2024: 1 dense forest, 2 open forest, 3 woodland and shrub, 4 tree "
                    "plantation, 5 cropland, 6 pasture, 7 settlement and bare soil, 8 water and wetland."),
    ("", "canopy_2024", "Tree canopy cover in January 2024: A under 10%, B 10-30%, C 30-60%, D over 60%."),
    ("", "lc_2025", "Land cover in December 2025, codes as lc_2024."),
    ("", "canopy_2025", "Tree canopy cover in December 2025, codes as canopy_2024."),
    ("", "plantation", "Y where the trees are a plantation."),
    ("", "change", "Change seen between the two dates: none, loss, gain."),
    ("", "loss_year", "Year of the first image showing the loss."),
    ("", "loss_month", "Month of the first image showing the loss."),
    ("", "image_2024", "Date of the image read for January 2024."),
    ("", "image_2025", "Date of the image read for December 2025."),
    ("", "confidence", "Interpreter's confidence: H high, M medium, L low."),
    ("", "remarks", "Interpreter's remarks."),
    (FILES["qc"], "", "Quality-control re-interpretation, March and April 2026, exported from the ministry's form "
                      "tool, one record per submission; field names and labels in Spanish."),
    (FILES["adj"], "", "Adjudication of every field the two campaigns recorded differently."),
    (FILES["intersect"], "", "Each reference point's class on the final change map version 2.0: DEF deforestation "
                             "2024-2025, BOS stable forest, NOB non-forest."),
    (FILES["stats"], "", "Class areas of the final change map version 2.0 by 2025 district."),
    (FILES["poly"], "", "Deforestation polygons of the final change map version 2.0, delivered by 10 km tile and by "
                        "quarter of first detection."),
    ("", "polygon_id", "Polygon."),
    ("", "tile", "Delivery tile."),
    ("", "district_2025", "District on the 2025 district layer."),
    ("", "first_detected", "Date of the first observation showing change in the polygon."),
    ("", "area_ha", "Area of the polygon (hectares)."),
    ("", "patch_ha", "Area of the deforestation patch, the connected area of change on the national mosaic (hectares)."),
    ("", "x_ntm", "Easting of the polygon's centroid."),
    ("", "y_ntm", "Northing of the polygon's centroid."),
    (FILES["conc"], "", "Districts on the 2019 and 2025 layers."),
    (FILES["nfi"], "", "National forest inventory tree measurements, 2018-2022."),
    (FILES["fire"], "", "Fire hotspot detections (thermal anomalies), 2024-2025."),
]


UNITS = {"x_ntm": "m", "y_ntm": "m", "area_ha": "ha", "patch_ha": "ha", "date_interpreted": "dd/mm/yyyy",
         "first_detected": "yyyy-mm-dd", "image_2024": "yyyy-mm-dd", "image_2025": "yyyy-mm-dd"}


def write_dictionary(path):
    rows = [(f, c, d, UNITS.get(c, "")) for f, c, d in DICTIONARY]
    return _csv(path, ["file", "column", "meaning", "unit"], rows)


def write_manifest(path, counts):
    rows = [
        (FILES["main"], "Main interpretation campaign", "4,000 reference points", counts.get("main", ""), W.PEOPLE["reference"]),
        (FILES["qc"], "Quality-control re-interpretation", "one point in five", counts.get("qc", ""), "Ministry of Environment"),
        (FILES["adj"], "Adjudication log", "April 2026", counts.get("adj", ""), W.PEOPLE["reference"]),
        (FILES["intersect"], "Reference points on the final map", "version 2.0", counts.get("intersect", ""), W.PEOPLE["maps"]),
        (FILES["stats"], "Final map class areas", "version 2.0, by 2025 district", counts.get("stats", ""), W.PEOPLE["maps"]),
        (FILES["poly"], "Final map deforestation polygons", "version 2.0, whole country", counts.get("poly", ""), W.PEOPLE["maps"]),
        (FILES["conc"], "District concordance", "2019 and 2025 layers", counts.get("conc", ""), "National Statistics Institute"),
        (FILES["design"], "Reference sample design", "January 2026", "", W.PEOPLE["reference"]),
        (FILES["erpa"], "Payment agreement extract", "articles 1 to 9", "", W.MINISTRY),
        (FILES["plan"], "Monitoring plan extract", "sections 3 and 4", "", W.MINISTRY),
        (FILES["frel"], "Forest reference level summary", "2015-2022", "", W.MINISTRY),
        (FILES["release"], "Final map release note", "version 2.0", "", W.PEOPLE["maps"]),
        (FILES["nfi"], "Forest inventory tree measurements", "2018-2022", counts.get("nfi", ""), W.NFMS),
        (FILES["fire"], "Fire hotspots", "2024-2025", counts.get("fire", ""), W.NFMS),
        (FILES["dict"], "Data dictionary", "", "", W.PEOPLE["reference"]),
    ]
    return _csv(path, ["file", "content", "covers", "rows", "prepared_by"], rows)
