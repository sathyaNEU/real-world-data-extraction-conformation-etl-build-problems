"""Independent verifier for task501. Reads only target/, deliverables/ and submission.md; shares no code with the
generator. Rebuilds the walk-forward test and the plan with the csv module and plain dicts, then checks the figures
the submission and the board paper state."""
import csv
import os
import re
import sys
import zipfile
from datetime import date, timedelta

import openpyxl

ROOT = os.path.dirname(os.path.abspath(__file__))
T = os.path.join(ROOT, "target")
D = os.path.join(ROOT, "deliverables")
fails = []


def check(c, msg):
    if not c:
        fails.append(msg)


def dt(s):
    return date.fromisoformat(s) if s else None


def addm(d, n):
    y, m = d.year, d.month + n
    y += (m - 1) // 12
    m = (m - 1) % 12 + 1
    last = (date(y + (m == 12), 1 if m == 12 else m + 1, 1) - timedelta(days=1)).day
    return date(y, m, min(d.day, last))


def ym_end(ym):
    y, m = divmod(ym, 12)
    m += 1
    return date(y + (m == 12), 1 if m == 12 else m + 1, 1) - timedelta(days=1)


# ------------------------------------------------------------- pack
files = [f for f in os.listdir(T) if not f.startswith(".")]
check(len(files) >= 10, "10+ files")
check(len({f.rsplit(".", 1)[-1] for f in files}) >= 3, "3+ formats")
lines = list(csv.DictReader(open(os.path.join(T, "kapitalabrufe_ausschuettungen_2015-2026.csv"), newline="")))
check(len(lines) > 10000, "spine over 10,000 rows")
for r in lines:
    r["amount"] = float(r["amount"]); r["nd"] = dt(r["notice_date"]); r["vd"] = dt(r["value_date"])
issued = [r for r in lines if r["status"] == "ISSUED"]
wb = openpyxl.load_workbook(os.path.join(T, "fondsregister_2026-09-30.xlsx"), read_only=True)
ws = wb["Fonds"]
hdr = [c for c in next(ws.iter_rows(values_only=True))]
reg = {}
for row in ws.iter_rows(min_row=2, values_only=True):
    r = dict(zip(hdr, row))
    for c in ("investment_period_end", "admission_date"):
        r[c] = r[c].date() if hasattr(r[c], "date") else dt(str(r[c])[:10])
    reg[r["fund_id"]] = r
inv = {}
for r in csv.DictReader(open(os.path.join(T, "beteiligungsverzeichnis_2026-09-30.csv"), newline="")):
    inv[(r["fund_id"], r["investment_ref"])] = (dt(r["initial_investment_date"]), dt(r["realisation_date"]))
wb = openpyxl.load_workbook(os.path.join(T, "kapitalkonten_revidiert_2021-2025.xlsx"), read_only=True)
ws = wb["Kapitalkonten"]
hdr = [c for c in next(ws.iter_rows(values_only=True))]
aud = []
for row in ws.iter_rows(min_row=2, values_only=True):
    r = dict(zip(hdr, row)); r["fy"] = r["fy_end"].date() if hasattr(r["fy_end"], "date") else dt(str(r["fy_end"])[:10]); aud.append(r)
nav = list(csv.DictReader(open(os.path.join(T, "nav_register_quartal.csv"), newline="")))
for r in nav:
    r["pe"] = dt(r["period_end"]); r["v"] = float(r["nav_fund_ccy"])
fx = {dt(r["rate_date"]): (float(r["USDCHF"]), float(r["EURCHF"])) for r in csv.DictReader(open(os.path.join(T, "fx_referenzkurse_tagesende.csv")))}


def rate(d, ccy):
    if ccy == "CHF":
        return 1.0
    while d not in fx:
        d -= timedelta(days=1)
    return fx[d][0] if ccy == "USD" else fx[d][1]


PLAN = date(2026, 9, 30)
RESERVE = 75e6
BANDS = [(1, 12), (13, 24), (25, 36), (37, 54), (55, 78), (79, 10 ** 6)]
band = lambda a: next(i for i, (lo, hi) in enumerate(BANDS) if lo <= a <= hi)
mi = lambda d: d.year * 12 + d.month - 1
calls = [r for r in issued if r["notice_type"] == "CAPITAL_CALL"]
dists = [r for r in issued if r["notice_type"] == "DISTRIBUTION"]
first = {}
for r in calls:
    first[r["fund_id"]] = min(first.get(r["fund_id"], mi(r["nd"])), mi(r["nd"]))
cfm, dfm = {}, {}
for r in calls:
    cfm[(r["fund_id"], mi(r["nd"]))] = cfm.get((r["fund_id"], mi(r["nd"])), 0.0) + r["amount"]
for r in dists:
    dfm[(r["fund_id"], mi(r["nd"]))] = dfm.get((r["fund_id"], mi(r["nd"])), 0.0) + r["amount"]
anav = {(a["fund_id"], a["fy"].year): float(a["nav_fy_end"]) for a in aud}
C = {f: float(r["commitment"]) for f, r in reg.items()}
CCY = {f: r["currency"] for f, r in reg.items()}
ST = {f: r["strategy"] for f, r in reg.items()}


def recall(fid, upto):
    tot = 0.0
    months = int(reg[fid]["recycling_months"])
    for r in dists:
        if r["fund_id"] != fid or r["line_type"] != "RETURN_OF_CAPITAL" or r["nd"] > upto:
            continue
        init, real = inv.get((fid, r["investment_ref"]), (None, None))
        if init is None or real is None:
            continue
        if real <= addm(init, months) and r["nd"] <= reg[fid]["investment_period_end"]:
            tot += r["amount"]
    return tot


_u = {}


def unf(fid, upto):
    if (fid, upto) not in _u:
        paid = sum(r["amount"] for r in calls if r["fund_id"] == fid and r["nd"] <= upto)
        _u[(fid, upto)] = max(0.0, C[fid] - paid + recall(fid, upto))
    return _u[(fid, upto)]


ok_rule = sum(abs((C[a["fund_id"]] - sum(r["amount"] for r in calls if r["fund_id"] == a["fund_id"] and r["nd"] <= a["fy"]) + recall(a["fund_id"], a["fy"])) - float(a["unfunded_commitment"])) < 1 for a in aud)
check(ok_rule == len(aud) == 180, f"ledger basis reproduces {ok_rule} of {len(aud)} audited cells")


def age_at(fid, ym, origin):
    s = first.get(fid)
    if s is None or s > origin:
        s = origin + 1
    return ym - s + 1


def pooled(kind, fit_end, fxd):
    num, den = {}, {}
    for fid in reg:
        s = first.get(fid)
        if s is None or s > fit_end:
            continue
        f = rate(fxd, CCY[fid])
        for ym in range(max(2019 * 12, s), fit_end + 1):
            b = band(ym - s + 1)
            key = (ST[fid], b) if kind == "strategy_band" else b
            num[key] = num.get(key, 0.0) + cfm.get((fid, ym), 0.0) * f
            den[key] = den.get(key, 0.0) + C[fid] * f
    return {k: num[k] / den[k] for k in num}


def fund_rates(fit_end):
    out = {}
    for fid in reg:
        s = first.get(fid)
        if s is None or s > fit_end:
            continue
        num, cnt = {}, {}
        for ym in range(s, fit_end + 1):
            b = band(ym - s + 1)
            num[b] = num.get(b, 0.0) + cfm.get((fid, ym), 0.0) / C[fid]
            cnt[b] = cnt.get(b, 0) + 1
        out[fid] = {b: num[b] / cnt[b] for b in num}
    return out


def forecast(origin, kind, horizon=12):
    fxd = ym_end(origin)
    rates = pooled(kind, origin, fxd) if kind != "fund_own" else fund_rates(origin)
    fb = pooled("strategy_band", origin, fxd) if kind == "fund_own" else None
    U = {fid: unf(fid, fxd) for fid in reg if reg[fid]["admission_date"] <= fxd}
    monthly = []
    for h in range(1, horizon + 1):
        ym = origin + h
        tc = 0.0
        for fid in U:
            b = band(max(1, age_at(fid, ym, origin)))
            if kind == "fund_own":
                own = rates.get(fid)
                r = (own[b] if b in own else own[max(own)]) if own else fb.get((ST[fid], b), 0.0)
            elif kind == "strategy_band":
                r = rates.get((ST[fid], b), 0.0)
            else:
                r = rates.get(b, 0.0)
            c = min(r * C[fid], U[fid]); U[fid] -= c; tc += c * rate(fxd, CCY[fid])
        monthly.append(tc)
    return monthly


def realised_monthly(origin, horizon=12):
    fxd = ym_end(origin)
    out = [0.0] * horizon
    for fid in reg:
        if reg[fid]["admission_date"] > fxd:
            continue
        f = rate(fxd, CCY[fid])
        for h in range(1, horizon + 1):
            out[h - 1] += cfm.get((fid, origin + h), 0.0) * f
    return out


ORIGINS = [(2022, 12), (2023, 3), (2023, 6), (2023, 9), (2023, 12), (2024, 3), (2024, 6), (2024, 9), (2024, 12), (2025, 3), (2025, 6), (2025, 9)]
CANDS = ["band", "strategy_band", "fund_own"]
test = {c: dict(f=[], g=[]) for c in CANDS}
actual = []
for (y, m) in ORIGINS:
    o = y * 12 + m - 1
    rm = realised_monthly(o); A = sum(rm); actual.append(A)
    for c in CANDS:
        mo = forecast(o, c); test[c]["f"].append(sum(mo))
        cf = ca = 0.0; gaps = []
        for a_, fm in zip(rm, mo):
            cf += fm; ca += a_; gaps.append(ca - cf)
        test[c]["g"].append(max(0.0, max(gaps)))
for c in CANDS:
    test[c]["e"] = sum(abs(F - A) / A for F, A in zip(test[c]["f"], actual)) / 12
    test[c]["b"] = max(test[c]["g"])
sel = min(CANDS, key=lambda c: (test[c]["e"], CANDS.index(c)))
check(sel == "band", f"walk-forward selects {sel}")
check(test["band"]["e"] < test["strategy_band"]["e"] < test["fund_own"]["e"], "error ordering band < strategy band < manager's own")
o = 2026 * 12 + 8
FXP = {c: rate(PLAN, c) for c in ("CHF", "USD", "EUR")}
cp = forecast(o, sel, horizon=15)
dn, dd = {}, {}
for fid in reg:
    s = first.get(fid)
    if s is None:
        continue
    f = FXP[CCY[fid]]
    for ym in range(max(2022 * 12, s), 2026 * 12 + 9):
        nv = anav.get((fid, ym // 12 - 1))
        if not nv:
            continue
        key = (ST[fid], band(ym - s + 1)); dn[key] = dn.get(key, 0.0) + dfm.get((fid, ym), 0.0) * f; dd[key] = dd.get(key, 0.0) + nv * f
dr = {k: dn[k] / dd[k] for k in dn}
rows = []
for (y, m), tc in zip([(2026, 10), (2026, 11), (2026, 12)] + [(2027, k) for k in range(1, 13)], cp):
    td = sum(dr.get((ST[fid], band(max(1, age_at(fid, y * 12 + m - 1, o)))), 0.0) * anav.get((fid, 2025), 0.0) * FXP[CCY[fid]] for fid in reg)
    rows.append([y, m, tc, td, 0.25 * td, tc - 0.25 * td])
cum = 0.0; fac = []
for y, m, tc, td, cr, net in rows:
    if y == 2027:
        cum += net; fac.append(dict(month=m, calls=tc, dists=td, credited=cr, net=net, cum=cum))
peak = max(f["cum"] for f in fac); calls27 = sum(f["calls"] for f in fac); buffer = test[sel]["b"]
req = peak + buffer - RESERVE
TR = (50, 75, 100, 125, 150, 175, 200, 225, 250, 275, 300)
tr = next(t for t in TR if t * 1e6 >= req); below = TR[TR.index(tr) - 1]
check(tr == 225, f"tranche {tr}")
first_past = next(f["month"] for f in fac if f["cum"] + buffer - RESERVE > below * 1e6)
K = lambda x: round(x / 1000)
# ------------------------------------------------------------- bridge and cash
START, END = date(2025, 12, 31), date(2026, 6, 30)
br = {}
for fid in reg:
    s = ST[fid]; b = br.setdefault(s, dict(opening=0.0, calls=0.0, dists=0.0, closing=0.0)); ccy = CCY[fid]
    n0 = [x for x in nav if x["fund_id"] == fid and x["pe"] == START]
    if n0:
        b["opening"] += n0[0]["v"] * rate(START, ccy)
    for x in issued:
        if x["fund_id"] == fid and START < x["vd"] <= END and x["line_type"] != "IN_KIND":
            b["calls" if x["notice_type"] == "CAPITAL_CALL" else "dists"] += x["amount"] * rate(x["vd"], ccy)
    n1 = [x for x in nav if x["fund_id"] == fid and x["pe"] == END]
    if n1:
        b["closing"] += n1[0]["v"] * rate(END, ccy)
    else:
        prev = sorted([x for x in nav if x["fund_id"] == fid and x["pe"] < END], key=lambda x: x["pe"])
        if prev:
            pe, base = prev[-1]["pe"], prev[-1]["v"]
            for x in issued:
                if x["fund_id"] == fid and pe < x["vd"] <= END:
                    base += x["amount"] if x["notice_type"] == "CAPITAL_CALL" else -x["amount"]
            b["closing"] += base * rate(END, ccy)
tot = {kk: sum(b[kk] for b in br.values()) for kk in ("opening", "calls", "dists", "closing")}
br["TOTAL"] = tot
for b in br.values():
    b["vfx"] = b["closing"] - b["opening"] - b["calls"] + b["dists"]
cash = {m: [0.0, 0.0] for m in range(1, 10)}
for x in issued:
    if date(2026, 1, 1) <= x["vd"] <= PLAN and x["line_type"] != "IN_KIND":
        cash[x["vd"].month][0 if x["notice_type"] == "CAPITAL_CALL" else 1] += x["amount"] * rate(x["vd"], CCY[x["fund_id"]])
# ------------------------------------------------------------- the submission and the board paper carry these figures
sub = open(os.path.join(ROOT, "submission.md")).read()
with zipfile.ZipFile(os.path.join(D, "standby_facility_2027_board_paper.docx")) as z:
    docx_text = re.sub(r"<[^>]+>", " ", z.read("word/document.xml").decode("utf8"))
    for n in z.namelist():
        if n.startswith("docProps/"):
            meta = z.read(n).decode("utf8", "ignore")
            check("python-docx" not in meta and "openpyxl" not in meta, f"no writer signature in {n}")


def present(figure, where, label):
    s = f"{figure:,}" if isinstance(figure, int) else figure
    check(s in where, f"{label} {s} present")


for f in fac:
    for kk in ("calls", "dists", "credited", "net", "cum"):
        present(K(f[kk]), sub, f"month {f['month']} {kk}"); present(K(f[kk]), docx_text, f"paper month {f['month']} {kk}")
for i, A in enumerate(actual):
    present(K(A), sub, f"actual {i}"); present(K(A), docx_text, f"paper actual {i}")
    for c in CANDS:
        present(K(test[c]["f"][i]), sub, f"{c} forecast {i}"); present(K(test[c]["f"][i]), docx_text, f"paper {c} forecast {i}")
for c in CANDS:
    present(f"{test[c]['e'] * 100:.1f}", sub, f"{c} error"); present(f"{test[c]['e'] * 100:.1f}", docx_text, f"paper {c} error")
present(K(buffer), sub, "buffer"); present(K(buffer), docx_text, "paper buffer")
for c in CANDS:
    present(K(test[c]["b"]), sub, f"{c} furthest ahead"); present(K(test[c]["b"]), docx_text, f"paper {c} furthest ahead")
present(K(req - below * 1e6), sub, "uncovered below"); present(K(req - below * 1e6), docx_text, "paper uncovered")
present(K(peak), sub, "peak"); present(K(req), sub, "required")
present(round(225e6 * 0.0045), sub, "fee 225"); present(round(200e6 * 0.0045), sub, "fee 200")
for s in ("BO", "GR", "PC", "IN", "TOTAL"):
    for kk in ("opening", "calls", "dists", "vfx", "closing"):
        present(K(br[s][kk]), sub, f"bridge {s} {kk}"); present(K(br[s][kk]), docx_text, f"paper bridge {s} {kk}")
for m, (c, d) in cash.items():
    for v in (c, d, c - d):
        present(K(v), sub, f"cash {m}"); present(K(v), docx_text, f"paper cash {m}")
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]
check("CHF 225" in sub and MONTHS[first_past - 1] in sub, "submission names the tranche and the crossing month")
check("225" in docx_text and MONTHS[first_past - 1] in docx_text, "board paper names the tranche and the crossing month")
subs = [os.path.join(dp, f) for dp, _, fs in os.walk(ROOT) for f in fs if f in ("submission.md", "prompt.md")]
check(len(subs) == 2, "exactly one submission.md and one prompt.md in the tree")
print(f"verify_pack: {len(fails)} failures")
for f in fails:
    print("  FAIL", f)
print("tranche", tr, "required", round(req / 1e6, 3), "peak", round(peak / 1e6, 3), "buffer", round(buffer / 1e6, 3), "first past", first_past, "errors", {c: round(test[c]["e"], 4) for c in CANDS}, "ahead", {c: round(test[c]["b"] / 1e6, 2) for c in CANDS}, "rule", ok_rule)
sys.exit(1 if fails else 0)
