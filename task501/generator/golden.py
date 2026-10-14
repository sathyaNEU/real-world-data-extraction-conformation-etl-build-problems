"""The goldens, computed from the shipped pack only: the call and the six reproduction numbers (callmodel on the call's
files), the three asks (asks on the register and case files), written into winter_measles_2027.xlsx and
measles_readiness_2027.pptx, with generator/figures.json and the task's metadata.json."""
import io
import json
import os

import numpy as np

import asks as A
import bundle as B
import callmodel as CM
import pack as PK
import post as PP
import world as W

XLSX = "winter_measles_2027.xlsx"
PPTX = "measles_readiness_2027.pptx"
AUTHOR = W.PEOPLE["manager"]
SCEN = ["none"] + W.OPTION_KEYS
SCEN_LABEL = {"none": "No catch-up", "A": "Option A", "B": "Option B", "C": "Option C", "D": "Option D", "E": "Option E"}


def compute(target):
    inp = B.call_inputs(target)
    R = CM.golden(inp)
    r1 = {k: CM.round1(v) for k, v in R.items()}
    below = [k for k in W.OPTION_KEYS if R[k] < 1.0]
    assert R["none"] >= 1.0 and below == ["B"], (R, below)
    N = inp.N()
    S = inp.groups(inp.sus_age())
    share = S / N
    region_share = S.sum() / N.sum()
    pool = inp.with_matrix(inp.matrix(None))
    R_all = CM.golden(pool)
    Ps = inp.matrix(CM.summer_diary)
    parts, _ = inp.survey
    f = B.ask_frames(target)
    a1 = A.ask1(f["people"])
    a2 = A.ask2(f["people"], f["events"], f["log"], f["merges"])
    a3 = A.ask3(f["cases"])
    return dict(inp=inp, R=R, r1=r1, call="B", share_1519=float(share[3]), region_share=float(region_share),
                immune_pct=float(100 * (1 - region_share)), R_avg=CM.R_average(inp), R_all=R_all, ask1=a1, ask2=a2,
                ask3=a3, made={k: sum(p["made"] for p in inp.options[k]) for k in W.OPTION_KEYS},
                own_term=float(inp.P[3, 3] / inp.P[3].sum()), own_summer=float(Ps[3, 3] / Ps[3].sum()),
                n_term=sum(1 for q in parts if CM.term_diary(q[2])), n_summer=sum(1 for q in parts if CM.summer_diary(q[2])))


def pct1(x):
    return "%.1f" % (np.floor(x * 1000 + 0.5) / 10)


def pct0(x):
    return "%d" % int(np.floor(x * 100 + 0.5))


def figures(c):
    return {
        "call": "Option B, the catch-up for those born 2008 to 2012",
        "reproduction_number": {SCEN_LABEL[k]: "%.1f" % c["r1"][k] for k in SCEN},
        "reproduction_number_unrounded": {SCEN_LABEL[k]: round(c["R"][k], 4) for k in SCEN},
        "R0": 15,
        "region_immune_per_cent": pct1(1 - c["region_share"]),
        "herd_immunity_threshold_per_cent": "93.3",
        "default_R0_times_susceptible_share": "%.1f" % CM.round1(c["R_avg"]),
        "default_every_diary_reproduction_number": "%.1f" % CM.round1(c["R_all"]["none"]),
        "susceptible_share_15_19_per_cent": pct1(c["share_1519"]),
        "susceptible_share_region_per_cent": pct1(c["region_share"]),
        "diaries_term": "{:,}".format(c["n_term"]),
        "diaries_summer": "{:,}".format(c["n_summer"]),
        "people_made_immune": {k: "{:,}".format(c["made"][k]) for k in W.OPTION_KEYS},
        "ask1_register_residents": {k: "{:,}".format(c["ask1"][k]) for k in W.OPTION_KEYS},
        "ask2_no_valid_dose": {k: "{:,}".format(c["ask2"][k]["none"]) for k in W.OPTION_KEYS},
        "ask2_one_valid_dose": {k: "{:,}".format(c["ask2"][k]["one"]) for k in W.OPTION_KEYS},
        "ask3_confirmed_cases": {k: str(c["ask3"][k]["confirmed"]) for k in W.OPTION_KEYS},
        "ask3_imported": {k: str(c["ask3"][k]["imported"]) for k in W.OPTION_KEYS},
        "option_birth_years": {k: "%d to %d" % W.OPTIONS[k]["years"] for k in W.OPTION_KEYS},
        "assessment_date": "1 January 2028",
        "register_as_at": "30 June 2027",
        "case_window": "1 January 2019 to 30 June 2027",
        "dates": ["1 July 2024", "28 days", "2 to 26 May 2023", "3 July to 18 August 2023"],
        "age_group_teenagers": "15 to 19",
        "citations": ["PS-07 3.1", "PS-07 3.2", "PS-07 3.3", "PS-07 3.4", "PS-07 3.5", "PS-07 3.6", "guidelines 12.5"],
    }


# --- the workbook ------------------------------------------------------------------------------------------------
def write_xlsx(path, c):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Call"
    navy = "1F3A5F"
    hdr_fill = PatternFill("solid", fgColor="DCE6EF")
    pick_fill = PatternFill("solid", fgColor="E3F1E6")
    thin = Side(style="thin", color="9AA6AD")
    ws["A1"] = "Only the catch-up for those born 2008 to 2012 keeps measles below one this winter"
    ws["A1"].font = Font(bold=True, size=14, color=navy)
    ws["A2"] = ("Winter 2027-28 assessment under PS-07 section 3. Regional Immunisation Office, for the Director of Public "
                "Health, 14 July 2027.")
    ws["A2"].font = Font(italic=True, size=10, color="5B6670")
    ws["A4"] = "Recommendation"
    ws["A4"].font = Font(bold=True, size=11, color=navy)
    ws["A5"] = ("Fund option B, the MMR catch-up for those born 2008 to 2012 (aged 15 to 19 on 1 January 2028). It is the "
                "only option that brings measles' reproduction number for winter 2027-28 below one: %s with it, against "
                "%s with no catch-up and %s under each of the other four options."
                % ("%.1f" % c["r1"]["B"], "%.1f" % c["r1"]["none"], "%.1f" % c["r1"]["A"]))
    ws["A5"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A5:F5")
    ws.row_dimensions[5].height = 48
    heads = ["Scenario", "Birth years targeted", "Expected people made immune", "Reproduction number, winter 2027-28",
             "Below one"]
    r0 = 7
    for j, h in enumerate(heads, start=1):
        cell = ws.cell(row=r0, column=j, value=h)
        cell.font = Font(bold=True)
        cell.fill = hdr_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=thin)
    ws.row_dimensions[r0].height = 32
    for i, k in enumerate(SCEN, start=1):
        r = r0 + i
        ws.cell(row=r, column=1, value=SCEN_LABEL[k])
        ws.cell(row=r, column=2, value=None if k == "none" else "%d to %d" % W.OPTIONS[k]["years"])
        if k != "none":
            ws.cell(row=r, column=3, value=c["made"][k]).number_format = "#,##0"
        v = ws.cell(row=r, column=4, value=c["r1"][k])
        v.number_format = "0.0"
        ws.cell(row=r, column=5, value="Yes" if c["R"][k] < 1 else "No").alignment = Alignment(horizontal="center")
        if k == "B":
            for j in range(1, 6):
                ws.cell(row=r, column=j).fill = pick_fill
                ws.cell(row=r, column=j).font = Font(bold=True)
    notes = [
        "Reproduction number: the dominant eigenvalue of the next-generation matrix, built from the National Social Contact "
        "Survey 2023 diaries kept in school term (2 to 26 May 2023, contacts per day by age band) and each age band's "
        "susceptible share on 1 January 2028, scaled so that a wholly susceptible region gives the R0 of 15 in PS-07 3.2. "
        "The winter is school term apart from the Christmas and February breaks, and PS-07 3.1 needs the figure below "
        "one throughout.",
        "Susceptible shares: the national immunity profile (May 2027) by year of birth, weighted by the department's "
        "projection of residents at 1 January 2028 (PS-07 3.3, 3.4).",
        "Each option takes the options paper's expected people made immune off the susceptibles of its birth years "
        "(PS-07 3.5); option E as the paper splits it between those born 1993 to 1997 and 1998 to 2002.",
        "On the region's average immunity alone (%s per cent immune, above the %s per cent herd immunity threshold for R0 "
        "15) the figure would be %s, and on every diary in the survey %s. But two thirds of the diaries were kept from 3 "
        "July to 18 August 2023, in the school summer holidays, when the 15 to 19-year-olds have under a fifth of their "
        "contacts with each other, against more than half in term; %s per cent of them are susceptible, against %s per "
        "cent across the region." % (
            pct1(1 - c["region_share"]), "93.3", "%.1f" % CM.round1(c["R_avg"]), "%.1f" % CM.round1(c["R_all"]["none"]),
            pct1(c["share_1519"]), pct1(c["region_share"])),
    ]
    rn = r0 + len(SCEN) + 2
    ws.cell(row=rn, column=1, value="Notes").font = Font(bold=True, color=navy)
    for i, t in enumerate(notes, start=1):
        cell = ws.cell(row=rn + i, column=1, value=t)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=rn + i, start_column=1, end_row=rn + i, end_column=6)
        ws.row_dimensions[rn + i].height = 46
    for col, wdt in zip("ABCDEF", (16, 20, 18, 22, 11, 14)):
        ws.column_dimensions[col].width = wdt

    ws2 = wb.create_sheet("Options")
    ws2["A1"] = "What each option involves"
    ws2["A1"].font = Font(bold=True, size=14, color=navy)
    ws2["A2"] = ("Register: regional immunisation register at 30 June 2027 (extract run 3 July 2027). Cases: notifiable "
                 "disease system, cases resident in the region, extract of 6 July 2027.")
    ws2["A2"].font = Font(italic=True, size=10, color="5B6670")
    heads2 = ["Option", "Born", "People the register holds as living in the region, 30 June 2027",
              "Of whom no valid MMR dose", "Of whom exactly one valid dose",
              "Confirmed measles cases, rash onset 1 January 2019 to 30 June 2027", "Of which imported"]
    for j, h in enumerate(heads2, start=1):
        cell = ws2.cell(row=4, column=j, value=h)
        cell.font = Font(bold=True)
        cell.fill = hdr_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=thin)
    ws2.row_dimensions[4].height = 62
    for i, k in enumerate(W.OPTION_KEYS, start=1):
        r = 4 + i
        vals = [k, "%d to %d" % W.OPTIONS[k]["years"], c["ask1"][k], c["ask2"][k]["none"], c["ask2"][k]["one"],
                c["ask3"][k]["confirmed"], c["ask3"][k]["imported"]]
        for j, v in enumerate(vals, start=1):
            cell = ws2.cell(row=r, column=j, value=v)
            if isinstance(v, int):
                cell.number_format = "#,##0"
    rn = 4 + len(W.OPTION_KEYS) + 2
    notes2 = [
        "Living in the region: active registrations with an address in the State whose local health office, read in the "
        "scheme in force when it was recorded, is in the region; LH27 recorded before 1 July 2024 is Monaghan West "
        "(circular 03/2024). Records with a UK postcode, pupils living in Northern Ireland, are not living in the region.",
        "Valid doses: MMR events re-pointed through the register's merges to the registration they now belong to, less "
        "GP return lines the query log shows replaced or withdrawn, counted under the guidelines' 12.5 (on or after the "
        "first birthday, at least 28 days after the previous valid dose), given on or before 30 June 2027.",
        "Cases: confirmed cases, one per person and illness: the 2019 outbreak's cases back-entered to the outbreak "
        "module in May 2021 (legacy number in the comments) count once with their migrated records, and a case "
        "notified twice (the same date of birth, sex and rash onset) once. Option E includes option D's birth years.",
    ]
    ws2.cell(row=rn, column=1, value="Notes").font = Font(bold=True, color=navy)
    for i, t in enumerate(notes2, start=1):
        cell = ws2.cell(row=rn + i, column=1, value=t)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws2.merge_cells(start_row=rn + i, start_column=1, end_row=rn + i, end_column=7)
        ws2.row_dimensions[rn + i].height = 44
    for col, wdt in zip("ABCDEFG", (9, 14, 22, 15, 15, 22, 11)):
        ws2.column_dimensions[col].width = wdt
    ws2.freeze_panes = "C5"
    for sh in (ws, ws2):
        sh.page_setup.orientation = "landscape"
        sh.page_setup.fitToWidth = 1
        sh.sheet_properties.pageSetUpPr.fitToPage = True
    ws2.print_title_rows = "4:4"
    wb.save(path)
    PP.normalize_ooxml(path, AUTHOR, AUTHOR, "2027-07-13T13:05:00Z", "2027-07-14T09:42:00Z", excel=True)


# --- the deck -----------------------------------------------------------------------------------------------------------
def chart_png(c):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams["font.family"] = "DejaVu Sans"
    fig, ax = plt.subplots(figsize=(8.6, 4.3), dpi=150)
    labels = ["No catch-up", "A\nborn 2013-17", "B\nborn 2008-12", "C\nborn 2003-07", "D\nborn 1993-97", "E\nborn 1993-2002"]
    vals = [c["r1"][k] for k in SCEN]
    cols = ["#9aa6ad" if k == "none" else ("#2e7d4f" if k == "B" else "#5b7fa6") for k in SCEN]
    bars = ax.bar(range(6), vals, color=cols, width=0.62)
    for b_, v in zip(bars, vals):
        ax.text(b_.get_x() + b_.get_width() / 2, v + 0.03, "%.1f" % v, ha="center", va="bottom", fontsize=11,
                fontweight="bold", color="#1d2329")
    ax.axhline(1.0, color="#b23a3a", linestyle="--", linewidth=1.4)
    ax.text(-0.42, 1.5, "Dashed line at 1.0: below it, measles cannot keep spreading", color="#b23a3a", fontsize=9.5,
            ha="left", va="center")
    ax.annotate("only\noption\nbelow one", xy=(2, vals[2] / 2), ha="center", va="center", fontsize=8.5,
                color="white", fontweight="bold", linespacing=1.3)
    ax.set_xticks(range(6))
    ax.set_xticklabels(labels, fontsize=9.5)
    ax.set_ylim(0, 1.6)
    ax.set_ylabel("Reproduction number, winter 2027-28", fontsize=10)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.tick_params(axis="y", labelsize=9)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", metadata={"Software": None})
    plt.close(fig)
    return buf.getvalue()


def write_pptx(path, c):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Pt
    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)
    blank = prs.slide_layouts[6]
    navy = RGBColor(0x1F, 0x3A, 0x5F)
    grey = RGBColor(0x5B, 0x66, 0x70)

    def textbox(slide, x, y, w, h, paras, size=18, bold_first=False, color=navy):
        tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, t in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            run = p.add_run()
            run.text = t
            run.font.size = Pt(size if i == 0 or not bold_first else size - 4)
            run.font.bold = bool(bold_first and i == 0)
            run.font.color.rgb = color if i == 0 else RGBColor(0x1D, 0x23, 0x29)
            p.space_after = Pt(8)
        return tb

    def footer(slide, n, source):
        textbox(slide, 600000, 6420000, 10000000, 300000, [source], size=10, color=grey)
        textbox(slide, 11000000, 6420000, 700000, 300000, ["%d / 5" % n], size=10, color=grey)

    # 1: the call
    s = prs.slides.add_slide(blank)
    textbox(s, 600000, 520000, 11000000, 700000, ["Winter 2027-28 measles readiness"], size=30, bold_first=True)
    textbox(s, 600000, 1350000, 11000000, 2200000, [
        "Recommendation: fund option B, the MMR catch-up for those born 2008 to 2012",
        "It is the only option that brings measles' reproduction number for the winter below one: %s with it, against %s "
        "with no catch-up and %s under each of the other options." % ("%.1f" % c["r1"]["B"], "%.1f" % c["r1"]["none"],
                                                                   "%.1f" % c["r1"]["A"]),
        "Without it, one case in the region would on average start a chain that keeps growing through the winter."],
        size=24, bold_first=True)
    textbox(s, 600000, 3900000, 11000000, 1300000, [
        "Decision for the executive management team, 21 July 2027",
        "Fund option B for delivery from September to November 2027 through the school immunisation team, further "
        "education colleges and GP letters (EUR %s in the options paper)." % "{:,}".format(W.OPTIONS["B"]["cost"])],
        size=16, bold_first=True)
    footer(s, 1, "Regional Immunisation Office, Oriel Health Region")
    # 2: the chart
    s = prs.slides.add_slide(blank)
    textbox(s, 600000, 350000, 11000000, 600000, ["Only option B brings the winter reproduction number below one"],
            size=26, bold_first=True)
    png = chart_png(c)
    s.shapes.add_picture(io.BytesIO(png), Emu(900000), Emu(1150000), width=Emu(10300000))
    footer(s, 2, "Dominant eigenvalue of the next-generation matrix: National Social Contact Survey 2023 term-time "
                 "diaries, national immunity profile May 2027, projection at 1 January 2028, scaled to R0 15 (PS-07 3.2).")
    # 3: why the teenagers
    s = prs.slides.add_slide(blank)
    textbox(s, 600000, 350000, 11000000, 600000, ["Why the 15 to 19-year-olds decide it"], size=26, bold_first=True)
    textbox(s, 600000, 1250000, 11000000, 4500000, [
        "The region is %s per cent immune, above the %s per cent that R0 15 needs, so on the average no catch-up looks "
        "needed (%s)." % (pct1(1 - c["region_share"]), "93.3", "%.1f" % CM.round1(c["R_avg"])),
        "All the contact survey's diaries together give %s, but two thirds were kept in July and August with the schools "
        "closed. The winter runs in term." % ("%.1f" % CM.round1(c["R_all"]["none"])),
        "In the May diaries, kept in term, more than half of the contacts of 15 to 19-year-olds are with others of that "
        "age; in the summer diaries, under a fifth.",
        "The cohort born 2008 to 2012, whose uptake fell in the 2010-12 vaccine scare, is 15 to 19 on 1 January 2028 and "
        "%s per cent susceptible, against %s per cent across the region. In term, chains of infection run inside that "
        "group, so only option B takes the reproduction number below one; D and E reach the least immune adults and the "
        "most people and leave it at %s." % (pct1(c["share_1519"]), pct1(c["region_share"]), "%.1f" % c["r1"]["D"])],
        size=17)
    footer(s, 3, "National Social Contact Survey 2023 public use file; national immunity profile, May 2027")
    # 4: what each option involves
    s = prs.slides.add_slide(blank)
    textbox(s, 600000, 350000, 11000000, 600000, ["What each option involves"], size=26, bold_first=True)
    rows, cols = 6, 6
    tbl_shape = s.shapes.add_table(rows, cols, Emu(600000), Emu(1250000), Emu(11000000), Emu(3600000))
    tbl = tbl_shape.table
    from pptx.oxml.ns import qn
    tbl_shape._element.graphic.graphicData.tbl.tblPr.set("bandRow", "0")
    for j, wdt in enumerate((2500000, 2100000, 1500000, 1600000, 1800000, 1500000)):
        tbl.columns[j].width = Emu(wdt)
    heads = ["Option", "On the register, living in the region", "No valid MMR dose", "Exactly one valid dose",
             "Confirmed cases since 2019", "Of which imported"]
    for j, h in enumerate(heads):
        cell = tbl.cell(0, j)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = navy
        run = cell.text_frame.paragraphs[0].runs[0]
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for i, k in enumerate(W.OPTION_KEYS, start=1):
        vals = ["%s (born %d to %d)" % ((k,) + W.OPTIONS[k]["years"]), "{:,}".format(c["ask1"][k]),
                "{:,}".format(c["ask2"][k]["none"]), "{:,}".format(c["ask2"][k]["one"]), str(c["ask3"][k]["confirmed"]),
                str(c["ask3"][k]["imported"])]
        for j, v in enumerate(vals):
            cell = tbl.cell(i, j)
            cell.text = v
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(0xE3, 0xF1, 0xE6) if k == "B" else RGBColor(0xF4, 0xF6, 0xF7)
            run = cell.text_frame.paragraphs[0].runs[0]
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(0x1D, 0x23, 0x29)
            if j:
                cell.text_frame.paragraphs[0].alignment = 3
    textbox(s, 600000, 5050000, 11000000, 1300000, [
        "Register at 30 June 2027; valid doses under the national guidelines (12.5). Cases with rash onset from 1 January "
        "2019 to 30 June 2027. Option E includes option D's birth years."], size=12, color=grey)
    footer(s, 4, "Regional immunisation register extract of 3 July 2027; notifiable disease system extract of 6 July 2027")
    # 5: method and sources
    s = prs.slides.add_slide(blank)
    textbox(s, 600000, 350000, 11000000, 600000, ["Method and sources"], size=26, bold_first=True)
    textbox(s, 600000, 1250000, 11000000, 4800000, [
        "PS-07 section 3: the reproduction number below one throughout the winter, December to February, with R0 15 over "
        "the winter; population and immunity on 1 January 2028 from the national profile; each option counted for the "
        "people the options paper expects it to make immune.",
        "Next-generation matrix: contacts per day by age band from the survey's term-time diaries (2 to 26 May 2023) "
        "times each band's susceptible share, scaled so a wholly susceptible region gives 15; the reproduction number "
        "is its dominant eigenvalue, which has to stay below one through the term weeks of the winter (PS-07 3.1).",
        "Register: addresses in Northern Ireland are outside the region; local health office codes read in the scheme "
        "in force when recorded; events through the merge table; GP return lines replaced or withdrawn removed; doses "
        "counted under guidelines 12.5.",
        "Cases: notifiable disease system extract of 6 July 2027; the 2019 outbreak's back-entered records and cases "
        "notified twice counted once."],
        size=16)
    footer(s, 5, "PS-07 Measles winter planning standard, version 3.0; options paper of 9 June 2027")
    prs.save(path)
    PP.normalize_ooxml(path, AUTHOR, AUTHOR, "2027-07-14T10:12:00Z", "2027-07-14T16:20:00Z",
                       title="Winter 2027-28 measles readiness")


def build(target, golden_dir, figures_path, w=None):
    os.makedirs(golden_dir, exist_ok=True)
    c = compute(target)
    write_xlsx(os.path.join(golden_dir, XLSX), c)
    write_pptx(os.path.join(golden_dir, PPTX), c)
    figs = figures(c)
    with open(figures_path, "w", encoding="utf-8") as f:
        json.dump(figs, f, indent=1, ensure_ascii=False)
        f.write("\n")
    return c


def metadata(path, target, counts):
    files = sorted(os.listdir(target))
    fmts = sorted({f.rsplit(".", 1)[1] for f in files})
    big = max(((k, v) for k, v in counts.items() if v), key=lambda kv: kv[1])
    meta = {
        "task": "task501",
        "title": "Which autumn 2027 MMR catch-up, if any, keeps measles' winter reproduction number in Oriel Health "
                 "Region below one, with each option's register residents, unvaccinated and part-vaccinated people and "
                 "confirmed cases since 2019",
        "domain": "Epidemiology & Public Health",
        "subdomain": "infectious-disease transmission modelling (a measles catch-up appraised on the next-generation "
                     "matrix's dominant eigenvalue)",
        "objective": "Forecasting & Predictive Modeling",
        "prompt_shape": "07 (grid of cells): the winter reproduction number under no catch-up and each of five options, "
                        "with each option's register residents, no-dose and one-dose counts, confirmed and imported cases",
        "as_of": "2027-07-12",
        "deliverables": [XLSX, PPTX],
        "distractor_files": [PK.FILES[k] for k in PK.DISTRACTORS],
        "input_gates": {"files": len(files), "formats": fmts, "largest_file": PK.FILES[big[0]], "largest_file_rows": big[1]},
        "sources": [{"files": "all files under target/",
                     "origin": "constructed for this task by task501/generator/build.py (seeded, deterministic) around a "
                               "fictional regional health authority (Oriel Health Region), its immunisation register and "
                               "case records, a fictional national contact survey and immunity profile; no real person, "
                               "patient, practice or case is depicted, and the place names are used as geography only.",
                     "date": "2027-07-12", "license": "CC0-1.0 (original work)"}],
        "created": "2026-10-12",
        "generator": "task501/generator/build.py",
    }
    with open(path, "w", encoding="utf-8") as f:
        json.dump(meta, f, indent=2, ensure_ascii=False)
        f.write("\n")
