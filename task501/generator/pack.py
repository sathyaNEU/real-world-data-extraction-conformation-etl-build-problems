"""Write the evidence pack's data files: the profile, the projection and the survey workbook (the call's inputs),
the register extracts and the query log, the code list, the case export and the laboratory results (the asks'
inputs), the stock return (a distractor) and the transmittal list."""
import csv
import datetime as dt
import os

import callsim as CS
import world as W

FILES = {
    "profile": "measles_immunity_profile_oriel_2027-05.csv",
    "projection": "population_projection_oriel_2028-01-01.csv",
    "participants": "nscs_2023_participants.csv",
    "contacts": "nscs_2023_contacts.csv",
    "codebook": "nscs_2023_public_file_codebook.xlsx",
    "options": "mmr_catch_up_options_autumn_2027.pdf",
    "standard": "ps07_measles_winter_planning_standard.pdf",
    "people": "register_people_active_born_1993-2017.csv",
    "events": "register_mmr_events_born_1993-2017.csv",
    "merges": "register_merges_born_1993-2017.csv",
    "queries": "gp_return_queries_mmr_2005-2027.csv",
    "lho": "lho_codes_national_2024-07.csv",
    "circular": "lho_revision_circular_2024-04.pdf",
    "notes": "register_extract_notes_2027-07.docx",
    "guidelines": "immunisation_guidelines_ch12_mmr_excerpt.pdf",
    "cases": "nds_measles_cases_oriel_2019-2027.csv",
    "lab": "nmrl_measles_results_2019-2027.csv",
    "fields": "nds_case_export_fields.txt",
    "uptake": "national_mmr_uptake_report_2027q1.pdf",
    "stock": "mmr_stock_return_2027-06-30.csv",
    "transmittal": "transmittal_list_2027-07-12.csv",
}
DISTRACTORS = ("uptake", "stock")
CALL_FILES = ("profile", "projection", "participants", "contacts", "codebook", "options", "standard")
ASK_FILES = ("people", "events", "merges", "queries", "lho", "circular", "notes", "guidelines", "cases", "lab", "fields")


def iso(d):
    return d.isoformat() if d else ""


def irl(d):
    return d.strftime("%d/%m/%Y") if d else ""


def write_csv(path, header, rows, crlf=False):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\r\n" if crlf else "\n")
        w.writerow(header)
        for r in rows:
            w.writerow(r)
    return len(rows)


# --- the call's inputs ------------------------------------------------------------------------------------------------
def write_profile(path):
    v, basis = W.profile_values()
    rows = []
    for y in range(2027, 1926, -1):
        label = "1927 and earlier" if y == 1927 else str(y)
        rows.append([label, "%.1f" % v[y], W.PROFILE_NOTE_BASIS[basis[y]]])
    return write_csv(path, ["year_of_birth", "per_cent_immune", "basis"], rows)


def write_projection(path, reg):
    rows = []
    for c in W.COUNTIES:
        for a in CS.AGES:
            age = "100 and over" if a == 100 else str(a)
            yob = "1927 and earlier" if a == 100 else str(2027 - a)
            rows.append([c, age, yob, int(reg[c][a])])
    return write_csv(path, ["county", "age_at_1_january_2028", "year_of_birth", "persons"], rows)


OCCUPATIONS = {0: ["Pre-school"], 1: ["School pupil"], 2: ["School pupil"], 3: ["School pupil", "Student", "Employed"],
               4: ["Student", "Employed", "Unemployed"]}
DURATIONS = ["Less than 5 minutes", "5 to 15 minutes", "15 minutes to 1 hour", "1 to 4 hours", "4 hours or more"]
FREQUENCIES = ["Daily or almost daily", "About once a week", "About once a month", "Less often", "First time"]
SETTING_COLS = ["at_home", "at_school", "at_work", "at_leisure", "in_transport", "elsewhere"]


def _occupation(rng, band, age):
    if band in OCCUPATIONS and band < 3:
        return OCCUPATIONS[band][0]
    if band == 3:
        return "School pupil" if age <= 17 else ("Student" if rng.random() < 0.55 else "Employed")
    if band == 4:
        u = rng.random()
        return "Student" if u < 0.42 else ("Employed" if u < 0.9 else "Unemployed")
    if band >= 13:
        return "Retired" if rng.random() < 0.86 else "Employed"
    u = rng.random()
    if band == 12:
        return "Employed" if u < 0.62 else ("Retired" if u < 0.82 else "Home duties")
    return "Employed" if u < 0.8 else ("Home duties" if u < 0.89 else ("Unemployed" if u < 0.95 else "Student"))


def _household(rng, band):
    if band <= 3:
        return int(rng.choice([2, 3, 4, 5, 6, 7], p=[0.08, 0.22, 0.34, 0.24, 0.09, 0.03]))
    if band <= 5:
        return int(rng.choice([1, 2, 3, 4, 5], p=[0.12, 0.30, 0.25, 0.22, 0.11]))
    if band <= 10:
        return int(rng.choice([1, 2, 3, 4, 5, 6], p=[0.10, 0.22, 0.20, 0.28, 0.15, 0.05]))
    return int(rng.choice([1, 2, 3, 4], p=[0.30, 0.50, 0.14, 0.06]))


SETTING_OF = {"home": "at_home", "school": "at_school", "childcare": "at_school", "work": "at_work",
              "leisure": "at_leisure"}
PHYSICAL = {"at_home": 0.72, "at_school": 0.34, "at_work": 0.18, "at_leisure": 0.38, "in_transport": 0.06,
            "elsewhere": 0.12}
DURATION_P = {"at_home": [0.04, 0.08, 0.20, 0.30, 0.38], "at_school": [0.10, 0.18, 0.28, 0.30, 0.14],
              "at_work": [0.16, 0.22, 0.26, 0.24, 0.12], "at_leisure": [0.12, 0.20, 0.32, 0.30, 0.06],
              "in_transport": [0.55, 0.30, 0.12, 0.03, 0.0], "elsewhere": [0.48, 0.28, 0.16, 0.06, 0.02]}
FREQUENCY_P = {"at_home": [0.86, 0.08, 0.03, 0.02, 0.01], "at_school": [0.78, 0.14, 0.03, 0.03, 0.02],
               "at_work": [0.62, 0.20, 0.06, 0.06, 0.06], "at_leisure": [0.30, 0.34, 0.14, 0.12, 0.10],
               "in_transport": [0.10, 0.10, 0.08, 0.30, 0.42], "elsewhere": [0.12, 0.16, 0.12, 0.28, 0.32]}


def survey_rows(parts, contacts):
    """The public use file's rows: participants renumbered in a random order (the release order), their contacts with
    sex, settings, physical contact, duration and frequency drawn on the texture stream."""
    import numpy as np
    rng = np.random.default_rng(W.SURVEY_TEXTURE_SEED)
    order = rng.permutation(len(parts))
    newid = {}
    for k, i in enumerate(order):
        newid[parts[i]["part"]] = 230001 + k
    prow = []
    for q in parts:
        band = q["band"]
        sex = "F" if rng.random() < (0.55 if band >= 13 else 0.5) else "M"
        by = "Parent or guardian" if band <= 1 else ("Self, with a parent's help" if band == 2 else "Self")
        prow.append([newid[q["part"]], W.GROUP_LABELS[band], sex, _household(rng, band), _occupation(rng, band, q["age"]),
                     by, q["date"].isoformat(), q["date"].strftime("%A")])
    prow.sort(key=lambda r: r[0])
    by_part = {}
    for c in contacts:
        by_part.setdefault(c[0], []).append(c)
    crow = []
    cid = 0
    for q in sorted(parts, key=lambda q: newid[q["part"]]):
        cs = by_part.get(q["part"], [])
        perm = rng.permutation(len(cs)) if cs else []
        for i in perm:
            _, cb, setting, age = cs[i]
            col = SETTING_OF.get(setting)
            if col is None:
                col = "in_transport" if rng.random() < 0.3 else "elsewhere"
            flags = {k: "N" for k in SETTING_COLS}
            flags[col] = "Y"
            if rng.random() < 0.06:
                extra = SETTING_COLS[int(rng.integers(0, len(SETTING_COLS)))]
                flags[extra] = "Y"
            u = rng.random()
            csex = "" if u < 0.018 else ("F" if u < 0.509 else "M")
            phys = "Y" if rng.random() < PHYSICAL[col] else "N"
            dur = DURATIONS[int(rng.choice(5, p=DURATION_P[col]))]
            freq = FREQUENCIES[int(rng.choice(5, p=FREQUENCY_P[col]))]
            cid += 1
            crow.append([cid, newid[q["part"]], "Unknown" if cb < 0 else W.GROUP_LABELS[cb], csex]
                        + [flags[k] for k in SETTING_COLS] + [phys, dur, freq])
    return prow, crow


def write_participants(path, prow):
    return write_csv(path, ["part_id", "age_band", "sex", "household_size", "occupation", "diary_completed_by",
                            "diary_date", "diary_day"], prow)


def write_contacts(path, crow):
    return write_csv(path, ["contact_id", "part_id", "contact_age_band", "contact_sex"] + SETTING_COLS
                     + ["physical", "duration", "frequency"], crow)


def write_codebook(path, nat, n_parts, n_contacts):
    import openpyxl
    from openpyxl.styles import Alignment, Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "About"
    (p1, a1, b1, q1), (p2, a2, b2, q2) = W.SURVEY_PERIODS
    about = [
        ("National Social Contact Survey 2023", True),
        ("Public use file, release 1.0, %s" % W.SURVEY_RELEASE.strftime("%-d %B %Y"), False),
        ("", False),
        ("The survey asked a sample of people living in Ireland to record in a diary everyone they had contact with on "
         "one assigned day: a conversation of at least three words in each other's presence, or any skin-to-skin "
         "contact. A person met more than once on the day is recorded once.", False),
        ("Fieldwork: diaries were kept for days from %s to %s and from %s to %s. Participants were recruited to a quota "
         "in each five-year age band in each fieldwork period (%d per band in the first period, %d in the second), and "
         "diary days were assigned so that each day of the week carries one seventh of each band's diaries in each "
         "period." % (a1.strftime("%-d %B"), b1.strftime("%-d %B %Y"), a2.strftime("%-d %B"), b2.strftime("%-d %B %Y"),
                       q1, q2), False),
        ("%s participants returned a diary, listing %s contacts." % (format(n_parts, ","), format(n_contacts, ",")), False),
        ("Children under 10 were reported by a parent or guardian; participants aged 10 to 14 completed the diary with a "
         "parent's help.", False),
        ("Ages of participants and of their contacts are released in five-year bands, with 70 and over in one band. A "
         "contact whose age the participant could not estimate is coded Unknown. The diary had room for 45 contacts.",
         False),
        ("A contact can be recorded in more than one setting. School includes college and childcare.", False),
        ("Geography below the State is not released in the public file.", False),
        ("", False),
        ("Files: nscs_2023_participants.csv (one row per participant) and nscs_2023_contacts.csv (one row per contact "
         "listed), joined on part_id. The Population sheet gives the statistics office's estimate of the population of "
         "the State at 1 January 2023 by age band.", False),
        ("Please cite as: National Social Contact Survey 2023, public use file, release 1.0 (2024).", False),
    ]
    for i, (t, b) in enumerate(about, start=1):
        if t:
            c = ws.cell(row=i, column=1, value=t)
            c.font = Font(bold=b, size=12 if b else 11)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 120
    f1 = wb.create_sheet("Participants file")
    f1.append(["Field", "Description"])
    for r in [("part_id", "Participant number."), ("age_band", "Age band of the participant on the diary day."),
              ("sex", "F or M."), ("household_size", "People living in the participant's household, the participant "
                                    "included."),
              ("occupation", "Main activity: Pre-school, School pupil, Student, Employed, Unemployed, Home duties or "
                             "Retired."),
              ("diary_completed_by", "Self; Self, with a parent's help; or Parent or guardian."),
              ("diary_date", "The day the diary was kept (yyyy-mm-dd)."), ("diary_day", "Day of the week of diary_date.")]:
        f1.append(list(r))
    f2 = wb.create_sheet("Contacts file")
    f2.append(["Field", "Description"])
    for r in [("contact_id", "Contact number."), ("part_id", "The participant who listed the contact."),
              ("contact_age_band", "Age band of the contact, as estimated by the participant; Unknown if not given."),
              ("contact_sex", "F or M; blank if not given."),
              ("at_home, at_school, at_work, at_leisure, in_transport, elsewhere",
               "Y where the contact took place in that setting; school includes college and childcare."),
              ("physical", "Y for any skin-to-skin contact."),
              ("duration", "Total time spent with the contact on the day, in five bands."),
              ("frequency", "How often the participant usually sees the contact.")]:
        f2.append(list(r))
    f3 = wb.create_sheet("Population")
    f3.append(["Age band", "Population, 1 January 2023"])
    Np = CS.published_population(nat)
    for k, lab in enumerate(W.GROUP_LABELS):
        f3.append([lab, int(Np[k])])
    f3.append(["All ages", int(Np.sum())])
    for sh in (f1, f2, f3):
        for c in sh[1]:
            c.font = Font(bold=True)
        sh.column_dimensions["A"].width = 34
        sh.column_dimensions["B"].width = 96
    f3.column_dimensions["B"].width = 28
    wb.save(path)


# --- the asks' inputs ---------------------------------------------------------------------------------------------------
def write_people(path, people, active):
    rows = []
    for p in sorted((p for p in people if p.status in active), key=lambda p: p.pid):
        rows.append([p.pid, iso(p.dob), p.sex, p.lho, iso(p.lho_recorded), p.postcode, iso(p.registered),
                     p.shown_practice])
    return write_csv(path, ["person_number", "date_of_birth", "sex", "lho_code", "lho_recorded", "postcode",
                            "registered_on", "gp_practice"], rows, crlf=True)


def write_events(path, events):
    rows = []
    for e in events:
        rows.append([e.eid, e.pid, "MMR", e.product, iso(e.date), e.given_by, e.source_ref, iso(e.loaded)])
    return write_csv(path, ["event_number", "person_number", "vaccine", "product", "date_given", "given_by",
                            "source_ref", "loaded_on"], rows, crlf=True)


def write_merges(path, merges):
    rows = [[m["retired"], m["surviving"], iso(m["merged_on"]), m["reason"]] for m in merges]
    return write_csv(path, ["retired_number", "surviving_number", "merged_on", "reason"], rows, crlf=True)


def write_queries(path, log):
    rows = [[r["query_ref"], r["line_ref"], r["practice"], iso(r["raised"]), r["reason"], r["outcome"] if r["closed"]
             else "Open", r["replaced_by"], iso(r["closed"])] for r in log]
    return write_csv(path, ["query_ref", "line_ref", "practice", "raised_on", "reason", "outcome", "replaced_by",
                            "closed_on"], rows)


def write_lho(path):
    rows = []
    for code in sorted(set(W.REGION_LHO_NEW) | set(W.OTHER_LHO)):
        if code in W.REGION_LHO_NEW:
            name, county = W.REGION_LHO_NEW[code]
            region = W.REGION
        else:
            name, county, region = W.OTHER_LHO[code]
        rows.append([code, name, county, region, "2024-07-01"])
    return write_csv(path, ["lho_code", "lho_name", "county", "health_region", "in_force_from"], rows)


def write_cases(path, cases):
    rows = []
    for c in cases:
        rows.append([c["case_id"], c["legacy_ref"], irl(c["notified"]), irl(c["onset_date"]), irl(c["date_of_birth"]),
                     c["sex"], c["county"], c["classification"], c["importation"], c["outbreak_ref"], c["epi_link"],
                     c["hospitalised"], c["mmr_doses_reported"], c["comments"]])
    return write_csv(path, ["case_id", "legacy_ref", "notified", "onset_date", "date_of_birth", "sex", "county",
                            "classification", "importation", "outbreak_ref", "epi_link", "hospitalised",
                            "mmr_doses_reported", "comments"], rows, crlf=True)


def write_lab(path, lab):
    rows = [[r["specimen_id"], r["case_id"], irl(r["specimen_date"]), r["specimen_type"], r["test"], r["result"],
             r["genotype"]] for r in lab]
    return write_csv(path, ["specimen_number", "case_id", "specimen_date", "specimen_type", "test", "result",
                            "genotype"], rows, crlf=True)


def write_stock(path, stock):
    rows = [[r["store"], r["product"], r["batch"], iso(r["expiry"]), r["doses_on_hand"], iso(r["counted_on"])]
            for r in stock]
    return write_csv(path, ["store", "product", "batch", "expiry", "doses_on_hand", "counted_on"], rows)


TRANSMITTAL = [
    ("profile", "Measles immunity profile by year of birth, Oriel Health Region, projected to 1 January 2028",
     "National immunisation programme", "2027-05-14"),
    ("projection", "Projected usual residents at 1 January 2028 by county and single year of age",
     "Department of Public Health, health intelligence", "2027-03-22"),
    ("participants", "National Social Contact Survey 2023, public use file: participants", "National Social Contact Survey",
     "2024-03-08"),
    ("contacts", "National Social Contact Survey 2023, public use file: contacts", "National Social Contact Survey",
     "2024-03-08"),
    ("codebook", "National Social Contact Survey 2023, public use file: codebook", "National Social Contact Survey",
     "2024-03-08"),
    ("options", "Autumn 2027 MMR catch-up: options paper", "Regional Immunisation Office", "2027-06-09"),
    ("standard", "PS-07 Measles winter planning standard (version 3.0)", "Department of Public Health", "2026-03-10"),
    ("people", "Register extract: people with an active registration, born 1993 to 2017", "Regional immunisation register",
     "2027-06-30"),
    ("events", "Register extract: MMR events for registrations born 1993 to 2017", "Regional immunisation register",
     "2027-06-30"),
    ("merges", "Register extract: merged duplicate registrations, born 1993 to 2017", "Regional immunisation register",
     "2027-06-30"),
    ("queries", "Query log for GP return lines with an MMR dose, from January 2005", "Regional Immunisation Office",
     "2027-07-03"),
    ("lho", "Local health office codes in force from 1 July 2024", "National health service planning unit", "2024-07-01"),
    ("circular", "Local health office areas: revision of 1 July 2024 (circular)", "National health service planning unit",
     "2024-04-12"),
    ("notes", "Register extract notes", "Regional immunisation register", "2027-07-05"),
    ("guidelines", "National immunisation guidelines, chapter 12 (MMR), sections 12.4 to 12.6",
     "National immunisation programme", "2025-11-24"),
    ("cases", "Notifiable disease system: measles cases resident in Oriel Health Region, notified from 2019",
     "Department of Public Health, surveillance", "2027-07-06"),
    ("lab", "National measles reference laboratory results for those cases", "National measles reference laboratory",
     "2027-07-06"),
    ("fields", "Case export field notes", "Department of Public Health, surveillance", "2025-03-03"),
    ("uptake", "MMR uptake at 24 months and 5 years, quarter 1 2027", "National immunisation programme", "2027-06-20"),
    ("stock", "MMR stock return at 30 June 2027", "Regional Immunisation Office", "2027-06-30"),
]


def write_transmittal(path, rows_by_key):
    rows = []
    for key, desc, source, asof in TRANSMITTAL:
        n = rows_by_key.get(key)
        rows.append([FILES[key], desc, source, asof, "" if n is None else n])
    return write_csv(path, ["file", "description", "source", "as_of", "rows"], rows)
