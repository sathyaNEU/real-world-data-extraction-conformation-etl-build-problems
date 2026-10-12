"""Read the shipped pack back: the call's inputs into callmodel.CallInputs, the asks' files into frames. The golden
and the checks compute from these, never from simulation state."""
import csv
import os

import numpy as np
import pandas as pd

import callmodel as CM
import pack as PK
import world as W


def path(target, key):
    return os.path.join(target, PK.FILES[key])


def call_inputs(target):
    prof = {}
    with open(path(target, "profile"), encoding="utf-8") as f:
        for r in csv.DictReader(f):
            y = 1927 if r["year_of_birth"].startswith("1927") else int(r["year_of_birth"])
            prof[y] = float(r["per_cent_immune"])
    persons, yob = {}, {}
    with open(path(target, "projection"), encoding="utf-8") as f:
        for r in csv.DictReader(f):
            a = 100 if r["age_at_1_january_2028"].startswith("100") else int(r["age_at_1_january_2028"])
            y = 1927 if r["year_of_birth"].startswith("1927") else int(r["year_of_birth"])
            persons[a] = persons.get(a, 0) + int(r["persons"])
            yob[a] = y
    import openpyxl
    wb = openpyxl.load_workbook(path(target, "survey"), data_only=True)
    t2 = wb["Table 2"]
    rows = list(t2.iter_rows(min_row=3, max_row=2 + W.NG, values_only=True))
    P = np.array([[float(v) for v in r[1:1 + W.NG]] for r in rows])
    t1 = wb["Table 1"]
    natpop = np.array([float(r[3]) for r in t1.iter_rows(min_row=3, max_row=2 + W.NG, values_only=True)])
    opts = options_from_paper(target)
    return CM.CallInputs(prof, persons, yob, P, natpop, opts)


def options_from_paper(target):
    """The options paper's table read back from the PDF text: per option and year-of-birth span, the people, the
    estimated susceptible and the expected people made immune."""
    import re
    from pypdf import PdfReader
    text = " ".join(p.extract_text() for p in PdfReader(path(target, "options")).pages)
    text = re.sub(r"\s+", " ", text)
    spans = list(re.finditer(r"(\d{4}) to (\d{4}) (\d{1,2}) to (\d{1,2}) ", text))
    rows = []
    for k, m in enumerate(spans):
        chunk = text[m.end(): spans[k + 1].start() if k + 1 < len(spans) else m.end() + 600]
        r = re.search(r"(\d{1,3}(?:,\d{3})+) (\d{1,3}(?:,\d{3})*) (\d{2})% (\d{1,3}(?:,\d{3})*)", chunk)
        if not r:
            continue
        num = lambda x: int(x.replace(",", ""))
        rows.append(dict(years=(int(m.group(1)), int(m.group(2))), people=num(r.group(1)), susceptible=num(r.group(2)),
                         made=num(r.group(4))))
    out = {k: [] for k in W.OPTION_KEYS}
    seen = set()
    for r in rows:
        y = r["years"]
        if y == (1993, 1997) and "D" not in seen:
            out["D"].append(r)
            seen.add("D")
        elif y in ((1993, 1997), (1998, 2002)):
            out["E"].append(r)
        else:
            k = next(k for k in "ABC" if W.OPTIONS[k]["years"] == y)
            out[k].append(r)
    return out


def ask_frames(target):
    people = pd.read_csv(path(target, "people"), dtype=str).rename(columns={"person_number": "person_id"})
    events = pd.read_csv(path(target, "events"), dtype=str).rename(columns={"person_number": "person_id"})
    merges = pd.read_csv(path(target, "merges"), dtype=str)
    log = pd.read_csv(path(target, "queries"), dtype=str, keep_default_na=False)
    cases = pd.read_csv(path(target, "cases"), dtype=str, keep_default_na=False)
    for c in ("notified", "onset_date", "date_of_birth"):
        cases[c] = pd.to_datetime(cases[c], format="%d/%m/%Y").dt.strftime("%Y-%m-%d")
    lab = pd.read_csv(path(target, "lab"), dtype=str, keep_default_na=False)
    return dict(people=people, events=events, merges=merges, log=log, cases=cases, lab=lab)
