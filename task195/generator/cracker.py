"""The cracker's path, run on the cut pack.

The cracker landed the call, ran the hygiene battery, executed every rule the asking sentences point at
and matched every control. For each ask this module computes what that path returns and scores it
against the golden within half a unit of the stated rounding. It is measurement, not argument.
"""
import csv
import os
import statistics
import sys

import openpyxl

import golden as G
import world as W


def _rows(path, delim=","):
    with open(path, encoding="utf-8") as fh:
        return list(csv.DictReader((l for l in fh if not l.startswith("#")), delimiter=delim))


def ask1(target):
    """Counts the register's own specimen identifiers per report: the crosswalk is a separate file and
    nothing in the register points at it."""
    per = {}
    for r in _rows(os.path.join(target, W.FILES["panel"])):
        if r["report_date"][:7] in ("2027-04", "2027-05"):
            d = per.setdefault(r["report_ref"], {"rows": 0, "spec": set()})
            d["rows"] += 1
            d["spec"].add(r["specimen_id"])
    return {k: {"specimens": len(v["spec"]), "difference": v["rows"] - len(v["spec"])}
            for k, v in per.items()}


def ask2(target):
    """Reads the warehouse extract as it stands, applies the asking sentence's rule for a reviewed
    specimen and counts specimens released in each month."""
    rel, flag = {}, {}
    seen_r, seen_f = set(), set()
    for r in _rows(os.path.join(target, W.FILES["history"])):
        if r["release_status"] != "RELEASED" or not r["released_date"]:
            continue
        month, key = r["released_date"][:7], (r["released_date"][:7], r["specimen_id"])
        if key not in seen_r:
            seen_r.add(key)
            rel[month] = rel.get(month, 0) + 1
        call = r["review_outcome"] if r["call_source"] == "REVIEWED" else r["hetx_call"]
        if call in ("FLAG", "FLAGGED") and key not in seen_f:
            seen_f.add(key)
            flag[month] = flag.get(month, 0) + 1
    return {m: round(1000.0 * flag.get(m, 0) / rel[m], 1) for m in rel}


def ask3(target):
    """Takes the prepared level from the workbook's own prepared-levels sheet, level by level."""
    wb = openpyxl.load_workbook(os.path.join(target, W.FILES["dilution"]), data_only=True)
    prepared, results = {}, {}
    for row in wb["Prepared levels"].iter_rows(min_row=2, values_only=True):
        prepared.setdefault(str(row[1]).strip().upper(), float(row[3]))
    for row in wb["Replicate results"].iter_rows(min_row=2, values_only=True):
        if row[4] in (None, ""):
            continue
        results.setdefault(str(row[0]).strip().upper(), []).append(float(row[4]))
    out = {}
    for lv, _label, _frac in W.DILUTION:
        vals = sorted(results[lv])
        out[lv] = {"median_recovery": round(statistics.median(vals) / prepared[lv], 3),
                   "within_fifth": sum(1 for v in vals if abs(v - prepared[lv]) <= 0.2 * prepared[lv])}
    return out


def score(target, g=None):
    """Per ask: the figures the cracker's path returns, and how many land within half a unit of the
    golden at the rounding the asking sentence states."""
    g = g or G.compute(target)
    rows = []
    a1 = ask1(target)
    hits = 0
    for w in g["panel_weeks"]:
        got = a1.get(w["report_ref"], {})
        hits += sum(1 for k in ("specimens", "difference") if got.get(k) == w[k])
    rows.append(("A1", 2 * len(g["panel_weeks"]), hits))
    a2 = ask2(target)
    hits = sum(1 for b in g["baseline"] if abs(a2.get(b["month"], -1) - round(b["per_thousand"], 1)) < 0.05)
    rows.append(("A2", len(g["baseline"]), hits))
    a3 = ask3(target)
    hits = 0
    for d in g["dilution"]:
        got = a3.get(d["level"], {})
        hits += 1 if abs(got.get("median_recovery", -1) - round(d["median_recovery"], 3)) < 0.0005 else 0
        hits += 1 if got.get("within_fifth") == d["within_fifth"] else 0
    rows.append(("A3", 2 * len(g["dilution"]), hits))
    return rows


def ceiling(rows):
    """supplemental-stumping Part 0: C = 38 + 7 + 55 Lc, S = r + 7 + 55 Ls with r about 5."""
    total = sum(n for _a, n, _h in rows)
    kept = sum(h for _a, _n, h in rows)
    share = kept / float(total)
    c = 38 + 7 + 55 * share
    s = 5 + 7 + 55 * share
    return total, kept, share, (c + s) / 2.0, 100.0 * share


if __name__ == "__main__":
    tgt = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "target")
    rows = score(tgt)
    for ask, n, h in rows:
        print("%s: %d figures, cracker keeps %d, misses %d" % (ask, n, h, n - h))
    total, kept, share, pair, mirror = ceiling(rows)
    print("total %d figures, kept %d (%.1f per cent), pair ceiling %.1f, mirror leakage %.1f"
          % (total, kept, 100 * share, pair, mirror))
