"""The two things the screening lead puts in the validation file."""
import os

import world as W

BLUE = "#1f4e79"
GREY = "#8a8f98"
AMBER = "#b5651d"


def workbook(path, g):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    t = g["table"]
    wb = Workbook()
    thin = Side(style="thin", color="FFB9BCC1")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    head = Font(bold=True, size=10)
    fill = PatternFill("solid", fgColor="FFE6EAF0")

    def header(ws, row, names):
        for j, h in enumerate(names):
            c = ws.cell(row=row, column=1 + j, value=h)
            c.font = head
            c.fill = fill
            c.border = box
            c.alignment = Alignment(horizontal="left" if j == 0 else "right", wrap_text=True)

    ws = wb.active
    ws.title = "Change decision"
    ws["A1"] = ("Screen change adopted at the %s cutoff for the quarter from 1 July 2027" % g["call"]
                if g["call"] != "hold" else
                "No listed cutoff meets both tests: the change is not made this quarter")
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("KVAL-07 4.1 on the parallel run of receipts 1 March to 31 May 2027, validation "
                "extract 21 June 2027.")
    ws["A2"].font = Font(size=9, italic=True)
    ws["A4"] = "Decision"
    ws["A4"].font = head
    ws["B4"] = ("Replace the screen in service at the %s cutoff" % g["call"] if g["call"] != "hold"
                else "Do not make the change; continue the parallel run")
    ws["A5"] = "Increase at that cutoff"
    ws["A5"].font = head
    ws["B5"] = ("%d confirmed contamination events, %.2f percentage points, 95 per cent interval "
                "%.2f to %.2f points" % (g["chosen"]["gain"], g["chosen"]["gain_pp"],
                                         g["chosen"]["lo_paired"], g["chosen"]["hi_paired"]))
    ws["A6"] = "False-call rate"
    ws["A6"].font = head
    ws["B6"] = ("%.2f per cent against %.2f per cent for the screen in service"
                % (g["chosen"]["false_rate_pct"], t["inc_false_rate_pct"]))
    ws["A7"] = "Comparison"
    ws["A7"].font = head
    ws["B7"] = ("{:,} specimens, of which {:,} confirmed contamination events and {:,} panel-clean; "
                "the screen in service detects {:,} of the events and flags {:,} of the clean "
                "specimens".format(g["comparison"], t["n_mixed"], t["n_clean"], t["inc_detected"],
                                   t["inc_false"]))
    header(ws, 9, ["Cutoff", "Events detected", "Increase, specimens", "Increase, percentage points",
                   "False-call rate, per cent", "Increase established at 95 per cent",
                   "False-call rate not higher"])
    r = 10
    for row in t["cutoffs"]:
        vals = ["%.3f" % row["cutoff"], row["detected"], row["gain"], round(row["gain_pp"], 2),
                round(row["false_rate_pct"], 2), "Yes" if row["established"] else "No",
                "Yes" if row["ceiling_ok"] else "No"]
        for j, v in enumerate(vals):
            c = ws.cell(row=r, column=1 + j, value=v)
            c.border = box
            c.alignment = Alignment(horizontal="left" if j == 0 else "right")
            if j in (3, 4):
                c.number_format = "0.00"
            if "%.3f" % row["cutoff"] == g["call"]:
                c.font = Font(bold=True)
        r += 1
    ws.cell(row=r + 1, column=1,
            value="Screen in service: {:,} of {:,} events detected, {:,} of {:,} panel-clean "
                  "specimens flagged ({:.2f} per cent).".format(
                      t["inc_detected"], t["n_mixed"], t["inc_false"], t["n_clean"],
                      t["inc_false_rate_pct"])).font = Font(size=9)
    ws.cell(row=r + 2, column=1,
            value="Source: qc_screen_calls_gsqa and gsqb, mixture_fraction_rerun_v34, "
                  "identity_panel_register_2027.").font = Font(size=9, italic=True)

    ws2 = wb.create_sheet("Panel deliveries")
    ws2["A1"] = "Identity panel verdicts against the partner laboratory's weekly reports"
    ws2["A1"].font = Font(bold=True, size=12)
    ws2["A2"] = "April and May 2027 reports, register as extracted 21 June 2027."
    ws2["A2"].font = Font(size=9, italic=True)
    header(ws2, 4, ["Report", "Week ending", "Specimens with a verdict", "Difference against the report's total"])
    for i, w in enumerate(g["panel_weeks"]):
        for j, v in enumerate([w["report_ref"], w["report_date"], w["specimens"], w["difference"]]):
            c = ws2.cell(row=5 + i, column=1 + j, value=v)
            c.border = box
            c.alignment = Alignment(horizontal="left" if j < 2 else "right")
    ws2.cell(row=6 + len(g["panel_weeks"]), column=1,
             value="The report's own total counts results issued, so a result sent again and a second "
                   "aliquot typed again each add to it.").font = Font(size=9)

    ws3 = wb.create_sheet("Baseline, screen in service")
    ws3["A1"] = "Screen in service: flagged specimens per thousand released"
    ws3["A1"].font = Font(bold=True, size=12)
    ws3["A2"] = "The twelve reporting months to February 2027, from the quality-control warehouse extract."
    ws3["A2"].font = Font(size=9, italic=True)
    header(ws3, 4, ["Reporting month", "Specimens released", "Flagged", "Per thousand released"])
    for i, b in enumerate(g["baseline"]):
        for j, v in enumerate([b["month"], b["released"], b["flagged"], round(b["per_thousand"], 1)]):
            c = ws3.cell(row=5 + i, column=1 + j, value=v)
            c.border = box
            c.alignment = Alignment(horizontal="left" if j == 0 else "right")
            if j == 1:
                c.number_format = "#,##0"
            if j == 3:
                c.number_format = "0.0"

    ws4 = wb.create_sheet("Dilution panel")
    ws4["A1"] = "Analytical sensitivity: recovery by prepared level, 2027Q1 panel"
    ws4["A1"].font = Font(bold=True, size=12)
    ws4["A2"] = "Nine replicates at each level, prepared 11 March 2027; voided replicates are not counted."
    ws4["A2"].font = Font(size=9, italic=True)
    header(ws4, 3, ["Level", "Prepared level, fraction", "Replicates reported",
                    "Median recovery of the prepared level", "Replicates within a fifth of it"])
    for i, d in enumerate(g["dilution"]):
        for j, v in enumerate([d["level"], d["prepared"], d["replicates"],
                               round(d["median_recovery"], 3), d["within_fifth"]]):
            c = ws4.cell(row=4 + i, column=1 + j, value=v)
            c.border = box
            c.alignment = Alignment(horizontal="left" if j == 0 else "right")
            if j == 1:
                c.number_format = "0.0000"
            if j == 3:
                c.number_format = "0.000"

    ws3.cell(row=6 + len(g["baseline"]), column=1,
             value="Source: qc_screen_history_2026-02-15_2027-02-28 and the reporting pack's "
                   "release_volumes_monthly_2026-03_2027-02.").font = Font(size=9, italic=True)
    ws4.cell(row=5 + len(g["dilution"]), column=1,
             value="Source: dilution_panel_2027q1, prepared levels against replicate "
                   "results.").font = Font(size=9, italic=True)
    ws2.cell(row=7 + len(g["panel_weeks"]), column=1,
             value="Source: identity_panel_register_2027 against "
                   "identity_panel_weekly_reports_2027.").font = Font(size=9, italic=True)
    ws5 = wb.create_sheet("Notes")
    for line in [
            ["Validation file extract, KVAL-07 issue 4, section 5.1"],
            ["Prepared by the screening lead for the laboratory quality review of 30 June 2027."],
            ["The comparison is the specimens KVAL-07 3.2 describes: an accepted read, a mixture "
             "fraction on the accepted read set, and a panel verdict of MIXED or CONCORDANT reported "
             "on or before 21 June 2027."],
            ["Specimens received under a contract agreement are outside the validation scope (3.1)."],
            ["A confirmed contamination event is a specimen whose panel verdict is MIXED (3.3)."],
            ["Increases are stated in percentage points of the %d confirmed contamination events."
             % t["n_mixed"]],
            ["False-call rates are stated on the %d panel-clean specimens of the comparison (3.5)."
             % t["n_clean"]],
            ["Verdicts reported after the extract date are carried into the next parallel run (5.2)."]]:
        ws5.append(line)
    for sheet, widths in ((ws, [11, 15, 18, 20, 18, 22, 20]), (ws2, [16, 14, 22, 30]),
                          (ws3, [18, 20, 12, 20]), (ws4, [9, 20, 18, 30, 26]), (ws5, [104])):
        for i, wd in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(i)].width = wd
    ws.freeze_panes = "A10"
    ws2.freeze_panes = "A5"
    ws3.freeze_panes = "A5"
    ws4.freeze_panes = "A4"
    ws.print_area = "A1:G%d" % (r + 2)
    wb.save(path)


def chart(path, g):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    t = g["table"]
    rows = t["cutoffs"]
    labels = ["%.3f" % r["cutoff"] for r in rows]
    gains = [r["gain_pp"] for r in rows]
    lo = [r["gain_pp"] - r["lo_paired"] for r in rows]
    hi = [r["hi_paired"] - r["gain_pp"] for r in rows]
    fig, ax = plt.subplots(figsize=(7.6, 4.3), dpi=170)
    for i, r in enumerate(rows):
        chosen = "%.3f" % r["cutoff"] == g["call"]
        colour = BLUE if chosen else (AMBER if not r["ceiling_ok"] else GREY)
        ax.errorbar(i, gains[i], yerr=[[lo[i]], [hi[i]]], fmt="o", color=colour,
                    markersize=8 if chosen else 6, elinewidth=2.0 if chosen else 1.3,
                    capsize=5, zorder=3)
    ax.axhline(0, color="#333333", linewidth=1.0, zorder=2)
    ax.annotate("no increase", xy=(len(rows) - 0.35, 0), xytext=(len(rows) - 0.35, 0.45),
                fontsize=8, color="#333333", ha="right")
    ch = labels.index(g["call"])
    ax.annotate("adopted: %s cutoff, %.2f points (%.2f to %.2f)"
                % (g["call"], gains[ch], rows[ch]["lo_paired"], rows[ch]["hi_paired"]),
                xy=(ch, gains[ch]), xytext=(ch + 0.18, gains[ch] - 3.1), fontsize=8.5,
                color=BLUE, arrowprops=dict(arrowstyle="-", color=BLUE, linewidth=0.8))
    breach = [i for i, r in enumerate(rows) if not r["ceiling_ok"]]
    if breach:
        ax.annotate("false-call rate above the screen in service", xy=(breach[0], gains[breach[0]]),
                    xytext=(breach[0] - 1.65, gains[breach[0]] + 1.9), fontsize=8, color=AMBER,
                    arrowprops=dict(arrowstyle="-", color=AMBER, linewidth=0.8))
    ax.set_xticks(range(len(rows)))
    ax.set_xticklabels(labels)
    ax.set_xlim(-0.6, len(rows) - 0.2)
    ax.set_xlabel("Mixture fraction cutoff listed in KVAL-07 Table 2")
    ax.set_ylabel("Increase over the screen in service,\npercentage points of confirmed events")
    ax.set_title("Only the 0.030 cutoff both clears the interval test and holds the false-call rate"
                 if g["call"] == "0.030" else "Increase by cutoff, parallel run", fontsize=11,
                 loc="left", color="#1a1a1a")
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(axis="y", color="#e2e4e8", linewidth=0.7)
    ax.set_axisbelow(True)
    fig.text(0.01, 0.015, "Parallel run, receipts 1 March to 31 May 2027, %d confirmed contamination "
                          "events. Intervals at 95 per cent." % t["n_mixed"], fontsize=7.5,
             color="#5b6068")
    fig.tight_layout(rect=(0, 0.035, 1, 1))
    fig.savefig(path, facecolor="white", metadata={"Software": None})
    plt.close(fig)
