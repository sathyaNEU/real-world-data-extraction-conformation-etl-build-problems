"""Every graded figure, computed from the shipped bundle and nothing else."""
import csv
import datetime as dt
import math
import os
import statistics

import openpyxl

import world as W

Z = 1.959963985
Z1 = 1.644853627


def _rows(path, delim=","):
    with open(path, encoding="utf-8") as fh:
        return list(csv.DictReader((l for l in fh if not l.startswith("#")), delimiter=delim))


def _num(s):
    return None if s in ("", None) else float(s.replace(",", "."))


def read_pack(target):
    qc = {}
    for key in ("qc_a", "qc_b"):
        for r in _rows(os.path.join(target, W.FILES[key]), ";"):
            if r["read_accepted"] != "Y":
                continue
            qc[r["specimen_id"]] = {"project": r["project_code"], "agreement": r["agreement_ref"],
                                    "received": dt.date.fromisoformat(r["received_date"]),
                                    "flag": r["hetx_call"] == "FLAG", "het": _num(r["het_excess_pct"]),
                                    "call_source": r["call_source"], "plate": r["plate_id"]}
    mf = {}
    for r in _rows(os.path.join(target, W.FILES["mf"])):
        if r["read_set_accepted"] == "Y":
            mf[r["specimen_id"]] = float(r["mixture_fraction"])
    panel, panel_rows = {}, _rows(os.path.join(target, W.FILES["panel"]))
    for r in panel_rows:
        if dt.date.fromisoformat(r["report_date"]) > W.AS_OF:
            continue
        panel.setdefault(r["specimen_id"], []).append(r)
    return qc, mf, panel, panel_rows


def comparison(target):
    """KVAL-07 3.1 and 3.2, executed word for word."""
    qc, mf, panel, _ = read_pack(target)
    projects = {p[0] for p in W.PROJECTS}
    out = []
    for sid, q in sorted(qc.items()):
        if not (W.RUN_FROM <= q["received"] <= W.RUN_TO):
            continue
        if q["agreement"] or q["project"] not in projects:
            continue
        if sid not in mf:
            continue
        verdicts = {r["verdict"] for r in panel.get(sid, [])}
        if not verdicts & {"MIXED", "CONCORDANT"}:
            continue
        out.append({"specimen_id": sid, "flag": q["flag"], "estimate": mf[sid],
                    "mixed": "MIXED" in verdicts, "plate": q["plate"], "het": q["het"],
                    "call_source": q["call_source"]})
    return out


def table(rows, service_flag=lambda r: r["flag"]):
    """Per listed cutoff, the detection counts, the discordance and the false-call rate."""
    mixed = [r for r in rows if r["mixed"]]
    clean = [r for r in rows if not r["mixed"]]
    n = len(mixed)
    inc_det = sum(1 for r in mixed if service_flag(r))
    inc_false = sum(1 for r in clean if service_flag(r))
    out = []
    for cut in W.CUTOFFS:
        det = sum(1 for r in mixed if r["estimate"] >= cut)
        b = sum(1 for r in mixed if r["estimate"] >= cut and not service_flag(r))
        c = sum(1 for r in mixed if r["estimate"] < cut and service_flag(r))
        fp = sum(1 for r in clean if r["estimate"] >= cut)
        gain = det - inc_det
        diff = gain / n
        se_paired = math.sqrt(b + c - (b - c) ** 2 / n) / n
        p1, p2 = inc_det / n, det / n
        se_unpaired = math.sqrt(p1 * (1 - p1) / n + p2 * (1 - p2) / n)
        out.append({"cutoff": cut, "detected": det, "gain": gain, "gain_pp": 100.0 * diff,
                    "only_candidate": b, "only_service": c,
                    "false_calls": fp, "false_rate_pct": 100.0 * fp / len(clean),
                    "se_paired": se_paired, "se_unpaired": se_unpaired,
                    "lo_paired": 100.0 * (diff - Z * se_paired), "hi_paired": 100.0 * (diff + Z * se_paired),
                    "lo_unpaired": 100.0 * (diff - Z * se_unpaired),
                    "hi_unpaired": 100.0 * (diff + Z * se_unpaired),
                    "z_paired": diff / se_paired if se_paired else 0.0,
                    "z_unpaired": diff / se_unpaired if se_unpaired else 0.0,
                    "established": diff - Z * se_paired > 0,
                    "ceiling_ok": fp / len(clean) <= inc_false / len(clean)})
    return {"n_mixed": n, "n_clean": len(clean), "inc_detected": inc_det, "inc_false": inc_false,
            "inc_false_rate_pct": 100.0 * inc_false / len(clean),
            "inc_detect_pct": 100.0 * inc_det / n, "cutoffs": out}


def decide(tab, key="established"):
    """KVAL-07 4.1: the highest listed cutoff where both conditions hold, else the hold."""
    for row in tab["cutoffs"]:
        if row["ceiling_ok"] and row[key]:
            return "%.3f" % row["cutoff"]
    return "hold"


def wilson(k, n, z=Z):
    p = k / n
    d = 1 + z * z / n
    centre = (p + z * z / (2 * n)) / d
    half = (z / d) * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n))
    return centre - half, centre + half


def rungs(tab):
    """The four pipelines of the ladder, each on the same rows."""
    inc_lo, inc_hi = wilson(tab["inc_detected"], tab["n_mixed"])
    out = {}
    for row in tab["cutoffs"]:
        row["marginal_lo"], row["marginal_hi"] = wilson(row["detected"], tab["n_mixed"])
        row["separated"] = row["marginal_lo"] > inc_hi
    out["rung0_sign"] = next(("%.3f" % r["cutoff"] for r in tab["cutoffs"]
                              if r["ceiling_ok"] and r["gain"] > 0), "hold")
    out["rung1_overlap"] = next(("%.3f" % r["cutoff"] for r in tab["cutoffs"]
                                 if r["ceiling_ok"] and r["separated"]), "hold")
    out["rung2_unpaired"] = next(("%.3f" % r["cutoff"] for r in tab["cutoffs"]
                                  if r["ceiling_ok"] and r["gain_pp"] / 100.0 - Z * r["se_unpaired"] > 0), "hold")
    out["rung3_paired"] = decide(tab)
    out["service_interval"] = (100.0 * inc_lo, 100.0 * inc_hi)
    return out


def ask_panel(target, panel_rows):
    """Ask 1: the distinct vials the results in each April and May weekly report came from, and the
    excess of results issued over vials. The register's identifier is an accession; the vial's own
    reference is the biobank's, and one vial booked in twice carries two accessions."""
    biobank = {r["specimen_id"]: r["biobank_ref"]
               for r in _rows(os.path.join(target, W.FILES["crosswalk"]), "\t")}
    # the biobank's register covers every accession the panel register holds, so the vial a result
    # came from is never guessed from the accession
    missing = {r["specimen_id"] for r in panel_rows} - set(biobank)
    assert not missing, "the crosswalk does not cover %d register accessions" % len(missing)
    per = {}
    for r in panel_rows:
        day = dt.date.fromisoformat(r["report_date"])
        if day.month in (4, 5) and day.year == 2027:
            per.setdefault(r["report_ref"], {"rows": 0, "spec": set(), "date": r["report_date"]})
            per[r["report_ref"]]["rows"] += 1
            per[r["report_ref"]]["spec"].add(biobank[r["specimen_id"]])
    out = []
    for ref in sorted(per):
        d = per[ref]
        out.append({"report_ref": ref, "report_date": d["date"], "results_issued": d["rows"],
                    "specimens": len(d["spec"]), "difference": d["rows"] - len(d["spec"])})
    return out


def ask_baseline(target):
    """Ask 2: the screen in service's flagged specimens per thousand released, by release month.
    A corrected result is released again, so a specimen released twice in a month is one specimen."""
    amend = {}
    for a in _rows(os.path.join(target, W.FILES["amendments"])):
        amend[(a["specimen_id"], a["field"])] = a["value_after"]
    flagged, released = {}, {}
    seen_rel, seen_flag = set(), set()
    with open(os.path.join(target, W.FILES["history"]), encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            if r["release_status"] != "RELEASED" or not r["released_date"]:
                continue
            released_date = amend.get((r["specimen_id"], "released_date"), r["released_date"])
            r = dict(r,
                     hetx_call=amend.get((r["specimen_id"], "hetx_call"), r["hetx_call"]),
                     review_outcome=amend.get((r["specimen_id"], "review_outcome"),
                                              r["review_outcome"]))
            month = released_date[:7]
            key = (month, r["specimen_id"])
            if key not in seen_rel:
                seen_rel.add(key)
                released[month] = released.get(month, 0) + 1
            flagged_call = (r["review_outcome"] if r["call_source"] == "REVIEWED" else r["hetx_call"])
            if flagged_call in ("FLAG", "FLAGGED") and key not in seen_flag:
                seen_flag.add(key)
                flagged[month] = flagged.get(month, 0) + 1
    months = [r["reporting_month"] for r in _rows(os.path.join(target, W.FILES["volumes"]))]
    return [{"month": m, "released": released.get(m, 0), "flagged": flagged.get(m, 0),
             "per_thousand": 1000.0 * flagged.get(m, 0) / released[m]} for m in months]


def ask_dilution(target):
    """Ask 3: the dilution panel's median recovery and the replicates within a fifth of the level."""
    wb = openpyxl.load_workbook(os.path.join(target, W.FILES["dilution"]), data_only=True)
    prepared, stocks = {}, {}
    for row in wb["Prepared levels"].iter_rows(min_row=2, values_only=True):
        prepared[(str(row[0]).strip(), str(row[1]).strip().upper())] = float(row[3])
    results = {}
    for row in wb["Replicate results"].iter_rows(min_row=2, values_only=True):
        if row[4] in (None, ""):
            continue
        key = str(row[0]).strip().upper()
        stocks[key] = str(row[1]).strip()
        results.setdefault(key, []).append(float(row[4]))
    out = []
    for lv, _label, _frac in W.DILUTION:
        vals = sorted(results[lv])
        level = prepared[(stocks[lv], lv)]
        within = sum(1 for v in vals if abs(v - level) <= 0.2 * level)
        out.append({"level": lv, "prepared": level, "replicates": len(vals),
                    "median_recovery": statistics.median(vals) / level, "within_fifth": within})
    return out


def compute(target):
    qc, mf, panel, panel_rows = read_pack(target)
    rows = comparison(target)
    tab = table(rows)
    lad = rungs(tab)
    g = {"comparison": len(rows), "table": tab, "rungs": lad,
         "call": lad["rung3_paired"], "panel_weeks": ask_panel(target, panel_rows),
         "baseline": ask_baseline(target), "dilution": ask_dilution(target), "rows": rows}
    chosen = next(r for r in tab["cutoffs"] if "%.3f" % r["cutoff"] == g["call"])
    g["chosen"] = chosen
    g["runner_up"] = next(r for r in tab["cutoffs"]
                          if r["ceiling_ok"] and r["established"] and r["cutoff"] < chosen["cutoff"])
    return g


if __name__ == "__main__":
    import sys
    g = compute(sys.argv[1] if len(sys.argv) > 1 else
                os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "target"))
    t = g["table"]
    print("comparison %d, mixed %d, clean %d, service detects %d (%.2f%%), false %d (%.2f%%)"
          % (g["comparison"], t["n_mixed"], t["n_clean"], t["inc_detected"], t["inc_detect_pct"],
             t["inc_false"], t["inc_false_rate_pct"]))
    for r in t["cutoffs"]:
        print("  %.3f det %3d gain %+3d (%+.2f pp) [%.2f, %.2f] zp %.3f zu %.3f false %.2f%% %s %s"
              % (r["cutoff"], r["detected"], r["gain"], r["gain_pp"], r["lo_paired"], r["hi_paired"],
                 r["z_paired"], r["z_unpaired"], r["false_rate_pct"],
                 "established" if r["established"] else "not established",
                 "ceiling ok" if r["ceiling_ok"] else "CEILING BREACH"))
    print("rungs", g["rungs"])
    print("call", g["call"], "runner-up %.3f" % g["runner_up"]["cutoff"])
