"""The golden deliverables from the cut pack: the Unit's provisional release workbook (the 2025-26 table against the
Population Policy's allowances, the back series for 2018-19 to 2022-23 and notes) and the chart for the Council paper.
Every figure is computed once here and returned for figures.json."""
import math

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import openpyxl  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

import golden as G  # noqa: E402

OVER, UNDER, INK, INK2, SURF = "#eb6834", "#2a78d6", "#0b0b0b", "#52514e", "#fcfcfb"
YEARS = ["%d-%02d" % (y, (y + 1) % 100) for y in G.BACK_YEARS]


def r50(x):
    return int(50 * math.floor(x / 50.0 + 0.5))


def compute(P=None):
    P = P or G.Pack()
    gold = G.golden(P)
    c25 = P.cohort(G.Y25)
    c24 = P.cohort(G.Y24)
    f = {"call": G.designated(gold)}
    for s in P.sectors:
        rows = [r for r in c25 if r["sector"] == s]
        st = [G.state(r) for r in rows]
        f["sector.%s.arrivals" % s] = len(rows)
        f["sector.%s.long_at_cut" % s] = st.count("long")
        f["sector.%s.short_at_cut" % s] = st.count("short")
        f["sector.%s.open_at_cut" % s] = st.count("open")
        f["sector.%s.provisional" % s] = round(gold[s], 1)
        f["sector.%s.published" % s] = r50(gold[s])
        f["sector.%s.allowance" % s] = G.ALLOWANCE[s]
        f["sector.%s.over" % s] = gold[s] > G.ALLOWANCE[s]
    for s in ("CON", "HSC", "RET"):
        for tag, coh in (("2024_25", c24), ("2025_26", c25)):
            rows = [r for r in coh if r["sector"] == s]
            ret = [r for r in rows if P.history(r)]
            f["returning.%s.%s.share_pct" % (s, tag)] = round(100.0 * len(ret) / len(rows), 0)
        ret = [r for r in c24 if r["sector"] == s and P.history(r)]
        first = [r for r in c24 if r["sector"] == s and not P.history(r)]
        f["returning.%s.2024_25.long_pct" % s] = round(100.0 * sum(map(G.resolved_long, ret)) / len(ret), 0)
        f["first.%s.2024_25.long_pct" % s] = round(100.0 * sum(map(G.resolved_long, first)) / len(first), 0)
    for s in ("CON", "HSC", "RET"):
        f["returning.%s.2025_26.open" % s] = sum(1 for r in c25 if r["sector"] == s and G.state(r) == "open"
                                               and P.history(r))
    f["history_window_days"] = 365
    easy = G.completion(P, lambda r, m: (r["sector"], m))
    f["easy.call"] = G.designated(easy)
    for s in ("CON", "HSC", "RET"):
        f["easy.%s.provisional" % s] = round(easy[s], 0)
    f["arrivals.2025_26"] = len(c25)
    bt = G.backtests(P)
    f["backtest.golden"] = round(bt["golden"][0], 3)
    f["backtest.sector_months"] = round(bt["sector_months"][0], 3)
    f["backtest.flat_sector"] = round(bt["flat_sector"][0], 3)
    f["backtest.scored"] = bt["golden"][1]
    A, B = G.back_series(P)
    for y, tag in zip(G.BACK_YEARS, YEARS):
        for s in P.sectors:
            f["back.arrivals.%s.%s" % (tag, s)] = A[y][s]
            f["back.long.%s.%s" % (tag, s)] = B[y][s]
        f["back.arrivals.%s.total" % tag] = sum(A[y].values())
        f["back.long.%s.total" % tag] = sum(B[y].values())
    return P, gold, f


def _widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def workbook(path, P, gold, f):
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="E8EDF2")
    thin = Side(style="thin", color="7F8C99")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Provisional 2025-26"
    ws["A1"] = ("Provisional long-term immigration by employment sector, year to 30 June 2026, against the Population "
                "Policy's 2025-26 allowances")
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = ("Esklett Statistics Unit, for the Council of Ministers' meeting of 15 October 2026. Population Register "
                "cut at 31 August 2026, extract of 30 September 2026.")
    hdr = ["Code", "Employment sector", "Arrivals 2025-26", "Reached twelve months by the cut",
           "Left within twelve months", "Open at the cut", "Provisional long-term immigration (nearest 50)",
           "Allowance 2025-26", "Exceeds allowance"]
    ws.append([])
    ws.append(hdr)
    for c in ws[4]:
        c.font = bold
        c.fill = head_fill
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.border = Border(bottom=thin)
    tot = [0, 0, 0, 0, 0.0]
    for s in P.sectors:
        row = [s, P.names[s], f["sector.%s.arrivals" % s], f["sector.%s.long_at_cut" % s],
               f["sector.%s.short_at_cut" % s], f["sector.%s.open_at_cut" % s], f["sector.%s.published" % s],
               f["sector.%s.allowance" % s], "Yes" if f["sector.%s.over" % s] else "No"]
        ws.append(row)
        for i in range(4):
            tot[i] += row[2 + i]
        tot[4] += gold[s]
    ws.append([None, "All sectors"] + tot[:4] + [r50(tot[4]), sum(G.ALLOWANCE.values()), None])
    for c in ws[ws.max_row]:
        c.font = bold
        c.border = Border(top=thin)
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row):
        for c in row[2:8]:
            c.number_format = "#,##0"
        row[8].alignment = Alignment(horizontal="center")
    ws.append([])
    ws.append(["Exceeding their allowances: " + ", ".join(P.names[s] for s in P.sectors if s in f["call"]) + "."])
    ws.cell(ws.max_row, 1).font = bold
    ws.freeze_panes = "C5"
    _widths(ws, [7, 36, 11, 14, 12, 10, 16, 11, 11])
    ws.print_title_rows = "4:4"

    wb2 = wb.create_sheet("Back series")
    wb2["A1"] = "Long-term immigration by employment sector, years to 30 June 2019 to 2023 (final, persons)"
    wb2["A1"].font = Font(bold=True, size=12)
    wb2["A2"] = ("Compiled on the 2026 definitions from the legacy register's archive (to 12 March 2023) and the current "
                 "system's extract of 30 September 2026.")
    for title, key in (("Arrivals registered for work", "arrivals"), ("Long-term immigration", "long")):
        wb2.append([])
        wb2.append([title])
        wb2.cell(wb2.max_row, 1).font = bold
        wb2.append(["Code", "Employment sector"] + YEARS)
        for c in wb2[wb2.max_row]:
            c.font = bold
            c.fill = head_fill
            c.border = Border(bottom=thin)
        first = wb2.max_row + 1
        for s in P.sectors:
            wb2.append([s, P.names[s]] + [f["back.%s.%s.%s" % (key, t, s)] for t in YEARS])
        wb2.append([None, "All sectors"] + [f["back.%s.%s.total" % (key, t)] for t in YEARS])
        for c in wb2[wb2.max_row]:
            c.font = bold
            c.border = Border(top=thin)
        for row in wb2.iter_rows(min_row=first, max_row=wb2.max_row):
            for c in row[2:]:
                c.number_format = "#,##0"
    wb2.freeze_panes = "C4"
    _widths(wb2, [7, 36, 10, 10, 10, 10, 10])

    wn = wb.create_sheet("Notes")
    notes = [
        ("Definitions", "Long-term immigration: arrivals in the year to 30 June, by arrival date, registered for work, "
                        "whose stay lasts twelve months or more (leaving date on or after the first anniversary of "
                        "arrival). Each stay is classified on its own length (methods note, sections 1.1 and 1.2)."),
        ("Sector", "The activity declared at the arrival registration: 2014 activity codes before 1 January 2026, "
                   "the 2026 sector list from then, each 2014 code mapped to its 2026 sector."),
        ("Provisional figures", "Stays that reached their anniversary, or left before it, by 31 August 2026 are "
                                "counted as they are. Each stay still open at the cut is counted at the share of the "
                                "2024-25 arrivals of its sector and migration history still on the island the same "
                                "whole months after arriving whose stay reached twelve months; migration history is "
                                "whether the same registration number had a stay end in the 365 days before the "
                                "arrival."),
        ("Choice of estimate", "Estimated the same way from the 2023-24 arrivals, these chances scored %.3f on the "
                               "2024-25 arrivals open at 31 August 2025 (mean squared difference from the outcome, "
                               "%s stays), against %.3f for the share by sector and months alone and %.3f for the "
                               "sector's share of all arrivals (methods note, section 3.3)." % (
                                   f["backtest.golden"], "{:,}".format(f["backtest.scored"]),
                                   f["backtest.sector_months"], f["backtest.flat_sector"])),
        ("Rounding", "Provisional figures are published to the nearest 50 (methods note, section 3.4); the comparison "
                     "with each allowance is made on the unrounded figure. Back series figures are whole persons."),
        ("Back series", "The legacy archive's records, each record open at the move with the leaving date of its "
                        "TRANSFER record, a stay registered at two parish offices counted once (its two records "
                        "identical in every declared particular and closed by one leaving notice), with the current "
                        "system's own arrivals."),
        ("Allowances", "Population Policy 2024 to 2028, Table 1 (CO/POP/24/03). A sector over its allowance has its "
                       "2027 work permit quota reduced by the excess (section 3.2)."),
        ("Prepared", "Esklett Statistics Unit, 7 October 2026. Contact: Jean Grant, Chief Statistician.")]
    wn.append(["Item", "Note"])
    for c in wn[1]:
        c.font = bold
        c.fill = head_fill
    for k, v in notes:
        wn.append([k, v])
    for row in wn.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        row[0].alignment = Alignment(vertical="top")
    _widths(wn, [20, 110])
    wb.save(path)


def chart(path, P, gold, f):
    order = sorted(P.sectors, key=lambda s: gold[s])
    fig, ax = plt.subplots(figsize=(9.6, 5.6), dpi=150)
    fig.patch.set_facecolor(SURF)
    ax.set_facecolor(SURF)
    y = range(len(order))
    for i, s in enumerate(order):
        over = gold[s] > G.ALLOWANCE[s]
        ax.barh(i, gold[s], height=0.56, color=OVER if over else UNDER, zorder=2)
        a = G.ALLOWANCE[s]
        ax.plot([a, a], [i - 0.38, i + 0.38], color=INK, lw=2.0, zorder=3, solid_capstyle="butt")
        lab = "%d" % f["sector.%s.published" % s]
        ax.text(max(gold[s], a) + 6, i, "%s  (allowance %d)" % (lab, a), va="center", ha="left", fontsize=8.6,
                color=INK if over else INK2)
    ax.set_yticks(list(y))
    ax.set_yticklabels([P.names[s] for s in order], fontsize=9, color=INK)
    ax.set_xlabel("Provisional long-term immigration, year to 30 June 2026 (persons; labels to the nearest 50)",
                  fontsize=9, color=INK2)
    ax.set_xlim(0, max(max(gold.values()), max(G.ALLOWANCE.values())) * 1.32)
    ax.grid(axis="x", color="#d9d8d3", lw=0.6, zorder=0)
    for sp in ("top", "right", "left"):
        ax.spines[sp].set_visible(False)
    ax.spines["bottom"].set_color("#a9a8a2")
    ax.tick_params(axis="x", colors=INK2, labelsize=8.5)
    ax.tick_params(axis="y", length=0)
    short = {"HOS": "Hospitality", "FIN": "finance", "DIG": "digital", "CON": "construction", "HSC": "care",
             "RET": "retail", "AGR": "agriculture", "OTH": "public services"}
    over = [short[s] for s in P.sectors if s in f["call"]]
    over[0] = over[0][0].upper() + over[0][1:]
    ax.set_title("%s exceed their 2025-26 allowances" % (", ".join(over[:-1]) + " and " + over[-1]),
                 loc="left", fontsize=11.5, color=INK, pad=22, fontweight="bold")
    ax.text(0, 1.02, "Bars: provisional figure on the 31 August 2026 cut. Black tick: Population Policy allowance. "
                     "Orange: over the allowance.", transform=ax.transAxes, fontsize=8.3, color=INK2)
    fig.text(0.01, 0.012, "Source: Population Register extract of 30 September 2026; Population Policy 2024 to 2028, "
                          "Table 1. Esklett Statistics Unit.", fontsize=7.6, color=INK2)
    fig.tight_layout(rect=(0, 0.03, 1, 1))
    fig.savefig(path, facecolor=SURF, metadata={"Software": None})
    plt.close(fig)


def write(gpath):
    P, gold, f = compute()
    workbook(gpath("workbook"), P, gold, f)
    chart(gpath("chart"), P, gold, f)
    return f
