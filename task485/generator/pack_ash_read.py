"""The three ash contractors' exports read back from the cut pack into the rows ceiling.py scores: each row's site
through the Directorate's customer cross-reference and the permit register, its day of 2026, its quantity, and what
its own text says it carried and where."""
import csv
import datetime as dt
import os

import openpyxl

import world as W

D0 = dt.date(2026, 1, 1)


def _day(s):
    return (dt.date.fromisoformat(str(s)[:10]) - D0).days


def _sites(target):
    firm = {f: s for s, f, _ in W.SITES}
    ws = openpyxl.load_workbook(os.path.join(target, "permitted_boilers_register_2026.xlsx"), read_only=True)["Register"]
    permit_site = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0]:
            permit_site[r[0]] = firm[r[1]]
    names = {}
    with open(os.path.join(target, "ash_contractor_customers.csv"), newline="", encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            ss = {permit_site[p.strip()] for p in r["Permit Numbers"].split(";")}
            assert len(ss) == 1
            names[(r["Contractor"], r["Customer as recorded"])] = ss.pop()
    return names


def _kind(text):
    t = text.lower()
    if "bottom ash" in t:
        return "bottom"
    if "fly ash" in t or "pfa" in t:
        return "fly"
    return "other"


def read(target):
    names = _sites(target)
    moor = []
    ws = openpyxl.load_workbook(os.path.join(target, "whinmoss_invoice_lines_2026.xlsx"), read_only=True)["Invoice lines"]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[0]:
            continue
        svc = r[5]
        dest = "other"
        if _kind(svc) != "other":
            dest = "landfill" if "Galscar" in svc else ("cement" if "Velthwaite" in svc else "blocks")
        moor.append(dict(day=_day(r[1]), site=names[("Whinmoss Waste Haulage", r[3])], service=svc, qty=float(r[6]),
                         unit=r[7], ref=r[10], kind=_kind(svc), dest=dest))
    tol = []
    with open(os.path.join(target, "carnhope_jobs_2026.csv"), newline="", encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            tol.append(dict(job=r["Job No"], day=_day(r["Date"]), site=names[("Carnhope Skips", r["Site"])],
                            net_kg=int(r["Net kg"]), kind=_kind(r["Material"]), dest=r["Destination"]))
    gre = []
    ws = openpyxl.load_workbook(os.path.join(target, "pellrigg_dockets_2026.xlsx"), read_only=True)["Dockets"]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[0]:
            continue
        gre.append(dict(job=r[0], day=_day(r[1]), site=names[("Pellrigg and Sons Haulage", r[2])], net_t=float(r[5]),
                        kind=_kind(r[3]), dest=r[4]))
    return dict(moorside=moor, tolland=tol, greave=gre)
