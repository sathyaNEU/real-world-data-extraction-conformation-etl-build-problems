"""The golden from the shipped files alone: each ward's electors on the register to be published on 1 December 2032,
every dwelling aged to its number of canvasses on that register on the build-up the 2022 to 2025 registers show, and
the wards whose electors per councillor lie more than 10 per cent from the district's average."""
import csv
import datetime as dt
import math
import os
from collections import defaultdict

import openpyxl

D = dt.date
HORIZON = 2032
COUNCIL_SIZE = 43
SETTLED_K = 7          # a dwelling at its seventh canvass or later holds what the settled stock holds


def first_canvass(day):
    """the first canvass of a property in use from day: canvassed that year if in use on 31 August"""
    return day.year if day <= D(day.year, 8, 31) else day.year + 1


def canvass_k(day, Y):
    """canvasses a dwelling in use from day has had by the register published 1 December Y (0: in use after 31
    August of Y, on rolling registration only); None if not in use by 31 October Y"""
    if day > D(Y, 10, 31):
        return None
    return Y - first_canvass(day) + 1


def load(target):
    p = lambda n: os.path.join(target, n)
    wards = {}
    with open(p("wards_councillors_2014_order.csv"), newline="") as f:
        for r in csv.DictReader(f):
            wards[r["ward_code"]] = (r["ward_name"], int(r["councillors"]))
    llpg = {}
    with open(p("llpg_extract_2026-10-12.csv"), newline="") as f:
        for r in csv.DictReader(f):
            if r["LOGICAL_STATUS"] != "1" or r["BLPU_STATE"] != "2":
                continue
            cls = r["CLASSIFICATION"]
            if cls not in ("RD02", "RD03", "RD04", "RD06"):
                continue
            llpg[int(r["UPRN"])] = dict(type="flat" if cls == "RD06" else "house", ward=r["WARD_CODE"],
                                        inuse=D.fromisoformat(r["BLPU_STATE_DATE"]))
    reg = defaultdict(dict)
    with open(p("ems_property_electorate_2022-2025.csv"), newline="") as f:
        for r in csv.DictReader(f):
            Y = int(r["register_date"][:4])
            reg[Y][int(r["uprn"])] = int(r["local_government_electors"])
    comp, comp_by_site = {}, defaultdict(lambda: [0, 0])
    wb = openpyxl.load_workbook(p("housing_completions_monitoring_2014-2026.xlsx"), read_only=True)
    for i, r in enumerate(wb["Completions"].iter_rows(values_only=True)):
        if i < 4 or r[0] is None:
            continue
        typ = "flat" if r[7] in ("Flat", "Maisonette") else "house"
        comp[int(r[5])] = dict(type=typ, done=r[10].date(), site=r[0])
        comp_by_site[r[0]][1 if typ == "flat" else 0] += 1
    traj = []
    wb = openpyxl.load_workbook(p("housing_trajectory_2026-09-30.xlsx"), read_only=True)
    rows = list(wb["Trajectory"].iter_rows(values_only=True))
    hdr = list(rows[3])
    name_to_code = {v[0]: k for k, v in wards.items()}
    for r in rows[4:]:
        if r[0] is None:
            continue
        rec = dict(zip(hdr, r))
        traj.append(rec)
    return dict(wards=wards, llpg=llpg, reg=reg, comp=comp, traj=traj, name_to_code=name_to_code, hdr=hdr,
                comp_by_site={k: tuple(v) for k, v in comp_by_site.items()})


def electra_rows(target, Y):
    """the Electra property electorate report for the Y register: (2014 polling district, property reference, TOTAL,
    ATTAINERS, OVERSEAS, LOCAL_ONLY)"""
    out = []
    with open(os.path.join(target, "ELECTRA_PROPERTY_ELECTORATE_REV_%d.txt" % Y), encoding="utf-8") as f:
        rows = f.read().splitlines()
    i = [n for n, r in enumerate(rows) if r.startswith("PD|")][0]
    for r in rows[i + 1:]:
        p = r.split("|")
        out.append((p[0], int(p[1]), int(p[4]), int(p[5]), int(p[6]), int(p[7])))
    return out


def notice_rows(target):
    """Electra's notices of alteration: (notice date, 2014 polling district, property reference, action, franchise,
    attains 18)"""
    out = []
    with open(os.path.join(target, "ELECTRA_NOTICES_OF_ALTERATION_2021-01_2022-03.txt"), encoding="utf-8") as f:
        rows = f.read().splitlines()
    i = [n for n, r in enumerate(rows) if r.startswith("NOTICE|")][0]
    for r in rows[i + 1:]:
        p = r.split("|")
        dd, mm, yy = p[0].split("/")
        out.append((D(int(yy), int(mm), int(dd)), p[1], int(p[3]), p[5], p[6], p[7]))
    return out


def old_pd_ward(target, wards):
    """each 2014-scheme polling district's ward, through the 2023 review's schedule (wards were not changed)"""
    name_to_code = {v[0]: k for k, v in wards.items()}
    out = {}
    with open(os.path.join(target, "polling_district_review_2023_schedule.csv"), newline="") as f:
        for r in csv.DictReader(f):
            out[r["replaces_district"]] = name_to_code[r["ward"]]
    return out


def ask_a(target, wards, rule="golden"):
    """each ward's local government electors on the register published 1 December 2021. Electra's 2021 report is the
    register it held when it was switched off on 4 March 2022: the 2021 register with the notices of January to March
    2022 applied, so the golden backs those notices' local government electors out (an entry with no F and no
    attainment date; K electors are local government electors)"""
    pdw = old_pd_ward(target, wards)
    e = defaultdict(int)
    for pd, ref, tot, at, ov, lo in electra_rows(target, 2021):
        e[pdw[pd]] += tot if rule == "total" else tot - at - ov
    if rule in ("total", "export"):
        return dict(e)
    for day, pd, ref, act, fr, a18 in notice_rows(target):
        if rule != "all_notices" and day < D(2022, 1, 1):
            continue
        lg = fr != "F" and not a18
        if rule == "all_rows":
            lg = True
        elif rule == "no_k" and fr == "K":
            lg = False
        if not lg:
            continue
        if act == "ADD":
            e[pdw[pd]] -= 1
        elif rule != "adds_only":
            e[pdw[pd]] += 1
    return dict(e)


def settled_ratio(L, Y=2025):
    """electors per dwelling on the Y register at the dwellings at their seventh canvass or later, by type"""
    num, den = defaultdict(float), defaultdict(int)
    for u, e in L["reg"][Y].items():
        a = L["llpg"].get(u)
        if a is None:
            continue
        k = canvass_k(a["inuse"], Y)
        if k is None or k < SETTLED_K:
            continue
        num[a["type"]] += e
        den[a["type"]] += 1
    return {t: num[t] / den[t] for t in den}


def buildup(L, S, years=(2022, 2023, 2024, 2025)):
    """electors per dwelling at each canvass count over the settled level of the dwelling's type, pooled over the new
    system's registers; k from 0 (rolling registration only) to SETTLED_K - 1"""
    num, den = defaultdict(float), defaultdict(int)
    for Y in years:
        for u, e in L["reg"][Y].items():
            a = L["llpg"].get(u)
            if a is None or u not in L["comp"]:
                continue
            k = canvass_k(a["inuse"], Y)
            if k is None or k >= SETTLED_K:
                continue
            num[k] += e / S[a["type"]]
            den[k] += 1
    return {k: num[k] / den[k] for k in sorted(den)}


def future_units(L, mix="remaining"):
    """the trajectory's dwellings by ward, type and month in use: each monitoring year's dwellings spread evenly
    through its months (the 2026/27 column over October to March), each year's mix the mix the site still has to
    build (its houses and flats less those the completions schedule shows completed); mix="whole" takes the site's
    whole mix"""
    out = []
    cols = [c for c in L["hdr"] if isinstance(c, str) and (c[:4].isdigit() and "/" in c[:7])]
    for rec in L["traj"]:
        ward = L["name_to_code"][rec["Ward"]]
        h, f = rec["Houses"] or 0, rec["Flats"] or 0
        dh, df = L["comp_by_site"].get(rec["Site ref"], (0, 0))
        rh, rf = (h - dh, f - df) if mix == "remaining" else (h, f)
        share_f = rf / float(rh + rf) if (rh + rf) else 0.0
        for col in cols:
            n = rec[col] or 0
            if not n or col.startswith("2033"):
                continue
            y0 = int(col[:4])
            months = [(y0, m) for m in range(10 if col.startswith("2026/27") else 4, 13)] + [(y0 + 1, m) for m in range(1, 4)]
            for (yy, mm) in months:
                day = D(yy, mm, 15)
                for typ, sh in (("house", 1.0 - share_f), ("flat", share_f)):
                    if sh > 0:
                        out.append((ward, typ, day, n * sh / len(months)))
    return out


def forecast(L, S=None, c=None, Y=HORIZON):
    S = S or settled_ratio(L)
    c = c or buildup(L, S)
    e = defaultdict(float)
    seen = set()
    for u, n in L["reg"][2025].items():
        a = L["llpg"].get(u)
        if a is None:
            continue
        seen.add(u)
        k25 = canvass_k(a["inuse"], 2025)
        if k25 is not None and k25 >= SETTLED_K:
            e[a["ward"]] += n
        else:
            k = canvass_k(a["inuse"], Y)
            e[a["ward"]] += S[a["type"]] * (1.0 if k >= SETTLED_K else c[k])
    # dwellings in use after the 2025 register's last alteration (from the gazetteer and the completions monitoring)
    for u, a in L["llpg"].items():
        if u in seen or a["inuse"] <= D(2025, 10, 31):
            continue
        k = canvass_k(a["inuse"], Y)
        if k is None:
            continue
        e[a["ward"]] += S[a["type"]] * (1.0 if k >= SETTLED_K else c[k])
    for ward, typ, day, n in future_units(L):
        k = canvass_k(day, Y)
        if k is None:
            continue
        e[ward] += n * S[typ] * (1.0 if k >= SETTLED_K else c[k])
    return dict(e), S, c


def ward_totals(L, Y):
    """each ward's local government electors on the Y register (new system), by the gazetteer's ward"""
    out = defaultdict(int)
    for u, n in L["reg"][Y].items():
        out[L["llpg"][u]["ward"]] += n
    return dict(out)


def default_path(L, Y=HORIZON):
    """the default route: each dwelling on the 2025 register at its 2025 count, every dwelling since and to come at the
    settled stock's electors per dwelling by type"""
    S = settled_ratio(L)
    e = defaultdict(float)
    for u, n in L["reg"][2025].items():
        e[L["llpg"][u]["ward"]] += n
    for u, a in L["llpg"].items():
        if u in L["reg"][2025] or a["inuse"] <= D(2025, 10, 31) or canvass_k(a["inuse"], Y) is None:
            continue
        e[a["ward"]] += S[a["type"]]
    for ward, typ, day, n in future_units(L):
        if canvass_k(day, Y) is not None:
            e[ward] += n * S[typ]
    return dict(e)


def trajectory_years(L, first=2026, last=2031):
    """the trajectory's dwellings in the monitoring years first/first+1 to last/last+1 (the 2026/27 column from 1
    October)"""
    cols = [c for c in L["hdr"] if isinstance(c, str) and c[:4].isdigit() and first <= int(c[:4]) <= last]
    return sum((r[c] or 0) for r in L["traj"] for c in cols)


def switch_off_net(target, wards):
    """the January to March 2022 notices' net local government electors in Electra's 2021 report, by ward"""
    g, x = ask_a(target, wards), ask_a(target, wards, "export")
    d = {w: x[w] - g[w] for w in g}
    return dict(net=sum(d.values()), low=min(d.values()), high=max(d.values()), by_ward=d)


def verdicts(L, e):
    tot = sum(e.values())
    avg = tot / COUNCIL_SIZE
    v = {w: 100.0 * (e[w] / L["wards"][w][1] / avg - 1.0) for w in L["wards"]}
    return v, avg, tot


if __name__ == "__main__":
    import sys
    target = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "target")
    L = load(target)
    e, S, c = forecast(L)
    v, avg, tot = verdicts(L, e)
    print("settled", {t: round(x, 4) for t, x in S.items()}, "build-up", {k: round(x, 3) for k, x in c.items()})
    print("2032 total %.0f, average %.1f a councillor" % (tot, avg))
    print("outside:", sorted(w for w, x in v.items() if abs(x) > 10))
    print(" ".join("%s %+.1f" % (w, v[w]) for w in L["wards"]))
    A = ask_a(target, L["wards"])
    print("ask A, 2021 register:", A, "district", sum(A.values()))
    for r in ("export", "total", "adds_only", "all_rows", "all_notices", "no_k"):
        o = ask_a(target, L["wards"], r)
        print("  %-12s keeps %2d of %d" % (r, sum(1 for k in A if o.get(k) == A[k]), len(A)))
