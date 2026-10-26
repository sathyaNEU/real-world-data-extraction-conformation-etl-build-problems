"""An independent recomputation of every graded figure from the shipped files alone (csv, json and the Policy's
numbers as the PDF prints them), sharing no code with golden.py, checked against figures.json and the golden workbook.

    python3 task485/generator/verify.py
"""
import collections
import csv
import datetime as dt
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)
T = os.path.join(TASK, "target")
GOLD = os.path.join(TASK, "golden")
D = dt.date
CUT = D(2026, 8, 31)
OK = []


def say(label, ok, detail=""):
    OK.append(ok)
    print("%s  %-60s %s" % ("PASS" if ok else "FAIL", label, detail))


def iso(s):
    return D(*map(int, s.split("-"))) if s else None


def dmy(s):
    if not s:
        return None
    d, m, y = map(int, s.split("/"))
    return D(y, m, d)


def anniv(a):
    try:
        return a.replace(year=a.year + 1)
    except ValueError:
        return D(a.year + 1, 3, 1)


def whole_months(a, b):
    m = (b.year - a.year) * 12 + (b.month - a.month)
    return m - 1 if b.day < a.day else m


def allowances():
    """Table 1 of the Policy as printed: each sector code followed by its name and four allowances, 2025-26 second"""
    from pypdf import PdfReader
    pdf = PdfReader(os.path.join(T, "population_policy_2024_2028_sector_allowances.pdf"))
    words = "\n".join(p.extract_text() for p in pdf.pages).split("\n")
    out = {}
    for i, w in enumerate(words):
        w = w.strip()
        if w in ("HOS", "AGR", "CON", "FIN", "HSC", "RET", "DIG", "OTH") and w not in out:
            nums = [x.strip() for x in words[i + 1:i + 7] if re.fullmatch(r"\d+", x.strip())]
            if len(nums) >= 4:
                out[w] = int(nums[1])
    return out


def main():
    sc = json.load(open(os.path.join(T, "sector_classification_2026.json")))
    to26 = {c["code"]: c["sector_2026"] for c in sc["activity_codes_2014"]}
    to26.update({c["code"]: c["code"] for c in sc["sectors_2026"]})
    sectors = [c["code"] for c in sc["sectors_2026"]]
    rows = list(csv.DictReader(open(os.path.join(T, "population_register_extract_20260930.csv"))))
    stays_of = collections.defaultdict(list)
    for r in rows:
        stays_of[r["reg_no"]].append((iso(r["arrival_date"]), iso(r["leaving_date"])))
    xw0 = {r["legacy_no"]: r["reg_no"] for r in csv.DictReader(open(os.path.join(T, "registration_number_crosswalk_2023.csv")))}
    for a in csv.DictReader(open(os.path.join(T, "population_register_legacy_archive_2014_2023.csv"))):
        if a["CLOSURE"] == "LEFT":
            stays_of[xw0[a["LEGACY_NO"]]].append((dmy(a["DATE_OF_ARRIVAL"]), dmy(a["DATE_OF_LEAVING"])))

    def returning(r):
        a = iso(r["arrival_date"])
        return any(lv is not None and lv < a and (a - lv).days <= 365 for (ar, lv) in stays_of[r["reg_no"]]
                   if ar != a or lv != iso(r["leaving_date"]))

    def year_rows(y0):
        lo, hi = D(y0, 7, 1), D(y0 + 1, 6, 30)
        return [r for r in rows if r["record_origin"] == "ARRIVAL" and lo <= iso(r["arrival_date"]) <= hi]

    ref = collections.defaultdict(lambda: [0, 0])
    for r in year_rows(2024):
        a, lv = iso(r["arrival_date"]), iso(r["leaving_date"])
        reached = lv is None or lv >= anniv(a)
        stayed = 12 if lv is None else whole_months(a, lv)
        for m in range(0, min(stayed, 12) + 1):
            c = ref[(to26[r["activity_code"]], returning(r), m)]
            c[0] += 1
            c[1] += reached
    prov = {s: 0.0 for s in sectors}
    for r in year_rows(2025):
        a, lv, s = iso(r["arrival_date"]), iso(r["leaving_date"]), to26[r["activity_code"]]
        if lv is not None and lv < anniv(a):
            continue
        if anniv(a) <= CUT:
            prov[s] += 1
            continue
        n, k = ref[(s, returning(r), whole_months(a, CUT))]
        prov[s] += k / n if n else 0.0
    allow = allowances()
    say("Policy Table 1 read from the PDF", len(allow) == 8, json.dumps(allow))
    over = sorted(s for s in sectors if prov[s] > allow[s])
    f = json.load(open(os.path.join(HERE, "figures.json")))
    say("the call", over == f["call"], ",".join(over))
    for s in sectors:
        say("provisional %s" % s, abs(round(prov[s], 1) - f["sector.%s.provisional" % s]) < 1e-9,
            "%.1f against %s" % (prov[s], allow[s]))
    # the back series: the archive's records, those open at the move with their TRANSFER record's leaving and notice
    # dates, a stay registered twice counted once, and the current system's own arrivals
    xw = {r["legacy_no"]: r["reg_no"] for r in csv.DictReader(open(os.path.join(T, "registration_number_crosswalk_2023.csv")))}
    tr = {r["reg_no"]: (iso(r["leaving_date"]), iso(r["notice_received"])) for r in rows if r["record_origin"] == "TRANSFER"}
    seen, stays = set(), []
    for a in csv.DictReader(open(os.path.join(T, "population_register_legacy_archive_2014_2023.csv"))):
        if a["CLOSURE"] == "LEFT":
            lv, nt = dmy(a["DATE_OF_LEAVING"]), dmy(a["CLOSED_ON"])
        else:
            lv, nt = tr[xw[a["LEGACY_NO"]]]
        if lv is not None:
            key = (a["DATE_OF_ARRIVAL"], a["SEX"], a["YOB"], a["NATIONALITY"], a["ACTIVITY"], lv, nt)
            if key in seen:
                continue
            seen.add(key)
        stays.append((dmy(a["DATE_OF_ARRIVAL"]), to26[a["ACTIVITY"]], lv))
    for r in rows:
        if r["record_origin"] == "ARRIVAL":
            stays.append((iso(r["arrival_date"]), to26[r["activity_code"]], iso(r["leaving_date"])))
    labels = ["2018-19", "2019-20", "2020-21", "2021-22", "2022-23"]
    for y0, lab in zip(range(2018, 2023), labels):
        lo, hi = D(y0, 7, 1), D(y0 + 1, 6, 30)
        ca = collections.Counter(s for a, s, lv in stays if lo <= a <= hi)
        cl = collections.Counter(s for a, s, lv in stays if lo <= a <= hi and (lv is None or lv >= anniv(a)))
        for s in sectors:
            say("back series %s %s arrivals and long-term" % (lab, s),
                ca[s] == f["back.arrivals.%s.%s" % (lab, s)] and cl[s] == f["back.long.%s.%s" % (lab, s)],
                "%d, %d" % (ca[s], cl[s]))
    # the methods note's test: 2023-24 arrivals as the year before, the 2024-25 arrivals open at 31 August 2025
    cut25, ext25 = D(2025, 8, 31), D(2025, 9, 30)
    for name, use_history in (("golden", True), ("sector_months", False)):
        t = collections.defaultdict(lambda: [0, 0])
        for r in year_rows(2023):
            a, lv = iso(r["arrival_date"]), iso(r["leaving_date"])
            reached = lv is None or lv >= anniv(a)
            stayed = 12 if lv is None else whole_months(a, lv)
            for m in range(0, min(stayed, 12) + 1):
                c = t[(to26[r["activity_code"]], returning(r) if use_history else None, m)]
                c[0] += 1
                c[1] += reached
        err, n = 0.0, 0
        for r in year_rows(2024):
            a, lv, nt = iso(r["arrival_date"]), iso(r["leaving_date"]), iso(r["notice_received"])
            if iso(r["registered_on"]) > ext25 or anniv(a) <= cut25 or (lv is not None and lv <= cut25 and nt <= ext25):
                continue
            k, rr = t[(to26[r["activity_code"]], returning(r) if use_history else None, whole_months(a, cut25))]
            err += (rr / k - (lv is None or lv >= anniv(a))) ** 2
            n += 1
        say("the methods note's test, %s" % name, abs(round(err / n, 3) - f["backtest.%s" % name]) < 1e-9,
            "%.4f on %d stays" % (err / n, n))
    # the golden workbook carries the same figures
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(GOLD, "provisional_long_term_immigration_2025_26.xlsx"))
    ws = wb["Provisional 2025-26"]
    cells = {r[0].value: r for r in ws.iter_rows(min_row=5, max_row=12)}
    for s in sectors:
        say("workbook row %s" % s, cells[s][6].value == 50 * int(prov[s] / 50.0 + 0.5) and
            (cells[s][8].value == "Yes") == (prov[s] > allow[s]) and cells[s][7].value == allow[s])
    wb2 = wb["Back series"]
    blocks, cur = {}, None
    for r in wb2.iter_rows(min_row=1, values_only=True):
        if r[0] in ("Arrivals registered for work", "Long-term immigration"):
            cur = "arrivals" if r[0].startswith("Arrivals") else "long"
        elif cur and r[0] in sectors:
            blocks[(cur, r[0])] = list(r[2:7])
    say("workbook back series", all(blocks[(k, s)] == [f["back.%s.%s.%s" % (k, lab, s)] for lab in labels]
                                    for k in ("arrivals", "long") for s in sectors))
    n = len(OK)
    print("%d of %d" % (sum(OK), n))
    return all(OK)


if __name__ == "__main__":
    sys.exit(0 if main() else 1)
